"""Strict import preview, no Excel formulas executed and no imported price auto-approved."""
import csv
import io
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook
from .engine import number, DomainError
from .parsers import validate_file

ALIASES = {
    'code': ['清单编码', '项目编码', '编码', 'code'],
    'name': ['项目名称', '清单名称', '名称', 'name'],
    'features': ['项目特征', '特征描述', '特征', 'features'],
    'unit': ['计量单位', '单位', 'unit'],
    'quantity': ['工程量', '数量', 'quantity'],
    'node': ['工程层级', '单位工程', 'node'],
    'reference_unit_price': ['综合单价', '单价', 'unit_price'],
}


def preview_boq(filename, raw):
    ext = Path(filename).suffix.lower()
    grids = []
    if ext in ('.csv', '.tsv'):
        try:
            text = raw.decode('utf-8-sig')
        except UnicodeDecodeError:
            text = raw.decode('gb18030')
        grids = [('导入表', list(csv.reader(io.StringIO(text), delimiter='\t' if ext == '.tsv' else ',')))]
    elif ext == '.xlsx':
        validate_file(ext, raw)
        w = load_workbook(io.BytesIO(raw), read_only=True, data_only=False)
        for s in w:
            rows = []
            for i, row in enumerate(s.iter_rows()):
                if i > 5002:
                    raise DomainError('每张导入表最多5000条明细')
                rows.append([{'formula': c.value} if c.data_type == 'f' else c.value for c in row[:50]])
            grids.append((s.title, rows))
        w.close()
    else:
        raise DomainError('清单导入支持 XLSX、UTF-8/GB18030 CSV、TSV')
    records, errors, sheets = [], [], []
    for sheet, grid in grids:
        header_index, mapping = None, {}
        for i, row in enumerate(grid[:25]):
            labels = [str(v or '').strip() for v in row]
            mapped = {key: next((j for j, v in enumerate(labels) if v in aliases), None) for key, aliases in ALIASES.items()}
            if all(mapped[k] is not None for k in ['code', 'name', 'unit', 'quantity']):
                header_index, mapping = i, {k: j for k, j in mapped.items() if j is not None}
                break
        if header_index is None:
            sheets.append({'sheet': sheet, 'state': 'skipped', 'reason': '未找到编码、名称、单位、工程量四个必需列'})
            continue
        sheets.append({'sheet': sheet, 'state': 'mapped', 'columns': mapping})
        for index, row in enumerate(grid[header_index+1:], header_index+2):
            if not any(v is not None and str(v).strip() for v in row):
                continue
            values = {k: (row[j] if j < len(row) else '') for k, j in mapping.items()}
            try:
                if any(isinstance(v, dict) for v in values.values()):
                    raise DomainError('包含公式，请先在来源软件中核实并转为值后导入')
                if not all(str(values.get(k) or '').strip() for k in ['code', 'name', 'unit']):
                    raise DomainError('编码、名称和单位不能为空；汇总行请移到单独工作表')
                quantity = number(values.get('quantity'))
                record = {k: str(v if v is not None else '').strip() for k, v in values.items()}
                record.update(quantity=str(quantity), sheet=sheet, row=index, state='draft')
                record.setdefault('features', '')
                record.setdefault('node', '导入工程 / 待核定单位工程')
                if record.get('reference_unit_price'):
                    record['reference_unit_price'] = str(number(record['reference_unit_price']))
                records.append(record)
            except DomainError as exc:
                errors.append({'sheet': sheet, 'row': index, 'message': str(exc)})
            if len(records)+len(errors) > 5000:
                raise DomainError('单次导入最多5000条，请拆分文件')
    if not records and not errors:
        raise DomainError('没有找到可导入的清单明细；请使用包含清单编码、项目名称、单位、工程量的表格')
    keys = {}
    duplicates = []
    for row in records:
        key = (row['code'], row['node'])
        if key in keys:
            duplicates.append({'sheet': row['sheet'], 'row': row['row'], 'message': '相同层级存在重复编码，确认是否为不同特征清单', 'code': row['code']})
        keys[key] = True
    return {'records': records, 'errors': errors, 'warnings': duplicates, 'sheets': sheets,
            'notice': '历史单价仅作为参考保留，不作为当前授权资源价；工程量和清单均需独立复核。'}
