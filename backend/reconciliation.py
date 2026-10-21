"""Same-component independent quantity baseline, never reconciles a known winning total."""
import csv
import hashlib
import io
from decimal import Decimal, localcontext
from pathlib import Path
from fastapi import APIRouter, Depends, File, UploadFile, Form, HTTPException
from openpyxl import load_workbook
from . import main as core, store
from .engine import number, quantity_number, DomainError
from .parsers import validate_file

router = APIRouter(prefix='/api/v1')


def quantity_hash(p):
    return store.digest([{'id':t['id'], 'quantity':t['quantity'], 'unit':t['unit'], 'evidence':t['evidence'],
                          'measurement_basis':t['measurement_basis']} for t in p['takeoffs']])


def context_hash(p):
    return store.digest([{k:v for k,v in t.items() if k != 'state'} for t in p['takeoffs']])


def matches_context(p, record):
    return record['input_hash']==quantity_hash(p) and (not record.get('calculation_hash') or record['calculation_hash']==context_hash(p))


def baseline_rows(filename, raw):
    if Path(filename).suffix.lower() == '.xlsx':
        validate_file('.xlsx', raw)
        w = load_workbook(io.BytesIO(raw), read_only=True, data_only=False)
        grid = []
        for i,row in enumerate(w.active.iter_rows()):
            if i > 5000: raise DomainError('底稿最多5000条')
            if any(c.data_type == 'f' for c in row): raise DomainError('专业底稿须为核定数值，不能使用未核验公式缓存')
            grid.append([c.value for c in row])
        w.close()
    elif Path(filename).suffix.lower() == '.csv':
        try: text = raw.decode('utf-8-sig')
        except UnicodeDecodeError: text = raw.decode('gb18030')
        grid = list(csv.reader(io.StringIO(text)))
    else: raise DomainError('专业底稿支持 CSV 或 XLSX（第一张表）')
    required = ['构件名称', '单位', '基准工程量', '绝对容差', '相对容差']
    if not grid or not all(k in grid[0] for k in required):
        raise DomainError('底稿需包含：构件名称、单位、基准工程量、绝对容差、相对容差；可选计算书ID')
    out = []
    for rowno,row in enumerate(grid[1:], 2):
        if not any(v is not None and str(v).strip() for v in row): continue
        values = dict(zip(grid[0], row))
        if not values.get('构件名称') or not values.get('单位'): raise DomainError(f'第{rowno}行缺构件名称或单位')
        expected, absolute, relative = [number(values.get(k)) for k in required[2:]]
        if relative > 1: raise DomainError('相对容差使用0至1小数，不能以中标总价差异作为算量容差')
        out.append({'name':str(values['构件名称']), 'unit':str(values['单位']), 'expected':str(expected),
                    'absolute_tolerance':str(absolute), 'relative_tolerance':str(relative), 'takeoff_id':str(values.get('计算书ID') or ''), 'row':rowno})
    if not out or len(out)>5000: raise DomainError('底稿需1至5000条独立构件明细')
    return out


@router.post('/projects/{pid}/takeoffs/reconcile')
async def reconcile(pid:str, file:UploadFile=File(...), basis:str=Form(...), a=Depends(core.actor)):
    core.edit(a)
    if not basis.strip(): raise HTTPException(422,'须填写底稿提供人、同构件范围与计量规则依据')
    raw=await file.read(10*1024*1024+1)
    if not raw or len(raw)>10*1024*1024: raise HTTPException(413,'底稿大小须为1字节至10MB')
    try:
        rows=baseline_rows(file.filename or '',raw)
    except DomainError:
        raise
    except Exception:
        raise HTTPException(422,'底稿文件结构损坏或格式不符，请检查原文件')
    with store.transaction() as c,localcontext() as ctx:
        ctx.prec=80
        p=core.project(c,pid,a)
        if not p['takeoffs']: raise HTTPException(409,'请先形成当前工程量计算书')
        used,results=set(),[]
        for row in rows:
            candidates=[t for t in p['takeoffs'] if (t['id']==row['takeoff_id'] if row['takeoff_id'] else t['name']==row['name']) and t['unit']==row['unit']]
            if len(candidates)!=1:
                results.append({**row,'passed':False,'message':'未唯一匹配同构件及同单位；请提供计算书ID，不能跨项目/跨口径比较'})
                continue
            t=candidates[0]
            if t['id'] in used: raise HTTPException(422,'同一计算书重复出现在底稿，不能重复覆盖')
            used.add(t['id'])
            actual,expected=quantity_number(t['quantity']),Decimal(row['expected'])
            delta=actual-expected
            limit=max(Decimal(row['absolute_tolerance']),abs(expected)*Decimal(row['relative_tolerance']))
            results.append({**row,'takeoff_id':t['id'],'actual':str(actual),'delta':str(delta),'allowed_delta':str(limit),
                            'relative_delta':str(abs(delta)/abs(expected)) if expected else None,'passed':abs(delta)<=limit,
                            'message':'在专业指定容差内' if abs(delta)<=limit else '超出容差，须按尺寸/扣减/工法/规则/修约归因'})
        missing=[t['id'] for t in p['takeoffs'] if t['id'] not in used]
        sha=hashlib.sha256(raw).hexdigest();folder=store.ROOT/'originals';folder.mkdir(exist_ok=True)
        if not (folder/sha).exists(): (folder/sha).write_bytes(raw)
        doc={'id':store.uid(),'name':Path(file.filename).name,'hash':sha,'version':1,'size':len(raw),'specialty':'合同计价','parse_state':'needs_review','parsed':True,'uploaded_at':store.now(),'message':'专业独立工程量底稿，待核查来源及口径'}
        p['documents'].append(doc)
        e={'id':store.uid(),'document_id':doc['id'],'document_hash':sha,'page':None,'text':basis,'method':'independent_baseline','bbox':None}
        p['evidence'].append(e)
        record={'id':store.uid(),'name':doc['name'],'input_hash':quantity_hash(p),'calculation_hash':context_hash(p),'evidence_id':e['id'],'basis':basis,'rows':results,
                'missing_takeoffs':missing,'passed':not missing and all(r['passed'] for r in results),'state':'draft','created_at':store.now()}
        p['reconciliations'].append(record)
        core.invalidate(p,'quotes')
        store.audit(c,p,a,'QUANTITY_RECONCILED',{'id':record['id'],'input_hash':record['input_hash'],'passed':record['passed'],'missing':missing})
        store.save(c,p)
    return record
