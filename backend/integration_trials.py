"""Evidence-backed, human-run round-trip trials. Never enables a vendor adapter."""
import copy
import hashlib
import io
import json
import re
from datetime import date
from decimal import Decimal, InvalidOperation
from pathlib import Path

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from openpyxl import Workbook, load_workbook
from . import main as core, store
from .engine import precise
from .parsers import validate_file
from .task_execution import original_bytes

router = APIRouter(prefix='/api/v1')
HEADERS = ['对象类型', '对象标识', '字段', '值', '单位', '比较方式', '绝对容差']
NOTICE = '仅对人工整理并注明来源的字段对照表进行比较。软件运行及授权由所附原件和独立人员核查；未调用厂商软件，不证明原生模型完整性，不启用导入/回导接口，也不替代工程成果审批。'
FIELDS = ['name', 'provider', 'product', 'software_version', 'return_software_version', 'adapter_version', 'input_format', 'output_format',
          'legal_scope', 'execution_summary', 'model_fidelity', 'fee_fidelity', 'failure_recovery', 'fallback', 'tested_on']
FIELD_LABELS = dict(zip(FIELDS, ['试验名称','软件厂商','软件产品','源软件完整版本','回导软件完整版本','适配/整理工具版本','源导出格式','回导文件格式','授权条件与本次使用范围','实际运行步骤及回导结果','模型/层级/清单关系保真','定额/资源/措施/税费保真','失败及重试恢复','不支持字段与回退方案','试验日期']))
LABELS = {'equal':'一致', 'within_tolerance':'容差内差异', 'changed':'值变化', 'missing':'字段丢失', 'unit_mismatch':'单位变化', 'added':'新增字段'}


def text(value, label, limit=2000):
    if not isinstance(value, str) or not value.strip() or len(value) > limit:
        raise HTTPException(422, label+'必填且不能超长')
    return value.strip()


def decimal(value):
    try: n = Decimal(str(value))
    except (InvalidOperation, ValueError): raise HTTPException(422, '数值比较须使用有限十进制数') from None
    if not n.is_finite() or abs(n) > Decimal('1e18') or len(n.as_tuple().digits) > 18 or n.as_tuple().exponent < -18:
        raise HTTPException(422, '对照值最多18位有效数字和18位小数，绝对值不超过1e18')
    return n


def parse_table(doc, raw, baseline):
    if Path(doc['name']).suffix.lower() != '.xlsx': raise HTTPException(422, '字段对照表须使用下载模板的XLSX文件')
    try:
        validate_file('.xlsx', raw)
        workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=False)
    except Exception: raise HTTPException(422, '字段对照表无法读取或包含不支持内容') from None
    try:
        if '字段对照' not in workbook.sheetnames: raise HTTPException(422, '需要名为“字段对照”的工作表')
        sheet = workbook['字段对照']
        if sheet.max_row and sheet.max_row > 5001 or sheet.max_column and sheet.max_column > 7:
            raise HTTPException(422, '字段对照表最多5000条明细且仅包含模板的7列')
        sheet.reset_dimensions()
        rows = sheet.iter_rows(max_col=8)
        header = next(rows, ())
        if [c.value for c in header[:7]] != HEADERS or len(header)>7 and header[7].value is not None: raise HTTPException(422, '字段对照表头与模板不一致')
        records, seen = [], set()
        for index, cells in enumerate(rows, 2):
            if index > 5001: raise HTTPException(422, '字段对照表最多5000条明细')
            if any(c.data_type in ('f', 'e') for c in cells): raise HTTPException(422, '对照表不能包含公式或错误单元格，请核实后转为值')
            if len(cells)>7 and cells[7].value is not None: raise HTTPException(422, '对照表存在模板之外的多余列')
            values = [str(c.value) if c.value is not None else '' for c in cells[:7]]
            if not any(v.strip() for v in values): continue
            values += ['']*(7-len(values))
            kind, oid, field = [text(values[i], '对象类型/标识/字段', 200) for i in range(3)]
            value = values[3]  # Preserve whitespace in text values; do not silently normalize results.
            if len(value) > 4000 or len(values[4]) > 100: raise HTTPException(422, '字段值或单位超长')
            key = (kind, oid, field)
            if key in seen: raise HTTPException(422, '同一对象同一字段不能重复，子目和资源请使用不同对象标识')
            seen.add(key)
            mode, tolerance = values[5].strip(), values[6].strip()
            if baseline:
                if mode not in ('文本', '数值'): raise HTTPException(422, '基准字段须指定文本或数值比较方式')
                if mode == '数值':
                    decimal(value)
                    if not tolerance or decimal(tolerance) < 0: raise HTTPException(422, '数值字段须显式填写非负绝对容差，无容差填0')
                elif tolerance not in ('', '0'): raise HTTPException(422, '文本比较不适用数值容差')
            records.append({'key':list(key), 'value':value, 'unit':values[4], 'mode':mode if baseline else '',
                            'tolerance':tolerance if baseline else '', 'row':index})
        if not records: raise HTTPException(422, '字段对照表至少需要一条明细')
        return records
    finally: workbook.close()


@precise
def compare(baseline, returned, run_status):
    lookup = {tuple(x['key']):x for x in returned}
    originals = {tuple(x['key']) for x in baseline}
    rows, counts = [], {k:0 for k in LABELS}
    for old in baseline:
        new = lookup.get(tuple(old['key']))
        delta, error = None, ''
        if new is None: status = 'missing'
        elif new['unit'] != old['unit']: status = 'unit_mismatch'
        elif old['mode'] == '文本': status = 'equal' if old['value'] == new['value'] else 'changed'
        else:
            try:
                delta = str(decimal(new['value'])-decimal(old['value']))
                status = 'equal' if Decimal(delta) == 0 else 'within_tolerance' if abs(Decimal(delta)) <= decimal(old['tolerance']) else 'changed'
            except HTTPException:
                status, error = 'changed', '返回值不是受支持的十进制数'
        counts[status] += 1
        rows.append({'key':old['key'], 'before':old['value'], 'after':new['value'] if new else None, 'before_unit':old['unit'],
                     'after_unit':new['unit'] if new else None, 'mode':old['mode'], 'tolerance':old['tolerance'], 'delta':delta,
                     'status':status, 'error':error, 'source_row':old['row'], 'returned_row':new['row'] if new else None})
    for new in returned:
        if tuple(new['key']) not in originals:
            counts['added'] += 1
            rows.append({'key':new['key'], 'before':None, 'after':new['value'], 'before_unit':None, 'after_unit':new['unit'],
                         'mode':'', 'tolerance':'', 'delta':None, 'status':'added', 'error':'', 'source_row':None, 'returned_row':new['row']})
    matched = len(baseline)-counts['missing']
    return {'rows':rows, 'counts':counts, 'baseline_fields':len(baseline), 'returned_fields':len(returned), 'matched_fields':matched,
            'coverage_percent':str((Decimal(matched)*100/len(baseline)).quantize(Decimal('.01'))),
            'passed':run_status == 'imported' and not any(counts[k] for k in ['changed','missing','unit_mismatch','added']),
            'coverage_notice':'覆盖率仅为列示基准字段的同标识匹配率；单位变化和值变化单列，不代表软件完整字段覆盖率。'}


def latest(p, key): return next((r for r in reversed(p['integration_trials']) if r['trial_key'] == key), None)
def content_hash(record): return store.digest({k:v for k,v in record.items() if k not in ('state','content_hash')})


def prepare(p, body):
    data = {k:text(body.get(k), FIELD_LABELS[k], 200 if k in FIELDS[:8] else 2000) for k in FIELDS}
    if data['software_version'] != data['return_software_version']: raise HTTPException(422, '往返验证须回到同一软件版本；不同版本请另建迁移试验')
    try: date.fromisoformat(data['tested_on'])
    except ValueError: raise HTTPException(422, '请填写有效的试验日期') from None
    key = text(body.get('trial_key'), '试验标识', 80)
    if not re.fullmatch(r'[A-Za-z0-9_-]{1,80}', key): raise HTTPException(422, '试验标识仅使用字母、数字、横线或下划线')
    run_status = body.get('run_status')
    if run_status not in ('imported', 'failed'): raise HTTPException(422, '请记录软件回导成功或失败')
    proofs, parsed, checked = [], {}, {}
    for role in ['baseline', 'returned']:
        did = body.get(role+'_document_id')
        if role == 'returned' and run_status == 'failed' and not did:
            parsed[role] = []
            continue
        doc = core.find(p, 'documents', did)
        if doc.get('superseded'): raise HTTPException(409, '对照文件已被新版本替代')
        raw = original_bytes(doc);checked[doc['hash']] = raw
        parsed[role] = parse_table(doc, raw, role == 'baseline')
        proofs.append({'role':role, 'document':{k:copy.deepcopy(doc[k]) for k in ['id','name','hash','version']}})
    for role in ['license', 'execution']:
        ev = core.require_evidence(p, body.get(role+'_evidence_id'));doc = core.find(p,'documents',ev['document_id'])
        if ev['document_hash'] != doc['hash']: raise HTTPException(409, '试验依据与原件关联不一致')
        if doc['hash'] not in checked: checked[doc['hash']] = original_bytes(doc)
        proofs.append({'role':role, 'document':{k:copy.deepcopy(doc[k]) for k in ['id','name','hash','version']}, 'evidence':copy.deepcopy(ev)})
    data.update(trial_key=key, run_status=run_status, proofs=proofs, comparison=compare(parsed['baseline'],parsed['returned'],run_status),
                schema='field-roundtrip-1.0', notice=NOTICE)
    return data


def validate(p, record, verify_bytes=True):
    if content_hash(record) != record['content_hash']: raise HTTPException(409, '试验记录内容校验失败')
    if latest(p, record['trial_key'])['id'] != record['id']: raise HTTPException(409, '已有较新的同组试验记录，历史版本只供查阅')
    checked = set()
    for proof in record['proofs']:
        doc = core.find(p, 'documents', proof['document']['id'])
        if doc.get('superseded') or any(doc[k] != value for k,value in proof['document'].items()): raise HTTPException(409, '试验文件或原始依据版本已变化')
        if proof.get('evidence') and core.require_evidence(p, proof['evidence']['id']) != proof['evidence']: raise HTTPException(409, '试验原始依据内容已变化')
        if not (store.ROOT/'originals'/doc['hash']).is_file(): raise HTTPException(409, '试验原件缺失')
        if verify_bytes and doc['hash'] not in checked: original_bytes(doc);checked.add(doc['hash'])
    if verify_bytes:
        docs = {x['role']:core.find(p,'documents',x['document']['id']) for x in record['proofs']}
        expected = parse_table(docs['baseline'],original_bytes(docs['baseline']),True)
        returned = parse_table(docs['returned'],original_bytes(docs['returned']),False) if 'returned' in docs else []
        if compare(expected,returned,record['run_status']) != record['comparison']: raise HTTPException(409, '试验对照结果与原件不一致')


@router.get('/projects/{pid}/integration-trials/template')
def template(pid:str, a=Depends(core.actor)):
    with store.connect() as c: core.project(c,pid,a)
    w=Workbook();s=w.active;s.title='字段对照';s.append(HEADERS);s.freeze_panes='A2'
    for col in 'ABCDEFG': s.column_dimensions[col].width=24
    notes=w.create_sheet('填写说明')
    for line in [NOTICE,'基准表须逐行指定文本/数值比较方式；数值容差须明确填写，无容差填0。返回表只读取前5列，最后2列不影响对照。',
                 '对象类型、对象标识和字段组成唯一标识。请覆盖实际需要验证的层级、构件、清单、定额、资源、换算、措施和税费字段。',
                 '单位须一致，不自动换算。文本值保留空格、大小写和空值差异。请把来源软件公式经人工核实后转成静态值。',
                 '仅解析“字段对照”页；其他页作为说明。人工整理不得以删掉丢失字段来制造通过；原始软件导出及回导日志请上传为运行依据。']:
        notes.append([line])
    out=io.BytesIO();w.save(out)
    return StreamingResponse(io.BytesIO(out.getvalue()),media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',headers={'Content-Disposition':'attachment; filename="roundtrip-field-template.xlsx"'})


@router.post('/projects/{pid}/integration-trials/preview')
def preview(pid:str,body:dict,a=Depends(core.actor)):
    core.edit(a)
    with store.connect() as c: data=prepare(core.project(c,pid,a),body)
    return {**data,'preview_hash':store.digest(data)}


@router.post('/projects/{pid}/integration-trials')
def submit(pid:str,body:dict,a=Depends(core.actor)):
    core.edit(a);request_id=text(body.get('request_id'),'请求标识',100)
    with store.transaction() as c:
        p=core.project(c,pid,a);request_hash=store.digest(body)
        duplicate=next((r for r in p['integration_trials'] if r['request_id']==request_id),None)
        if duplicate:
            if duplicate['request_hash']!=request_hash:raise HTTPException(409,'请求标识已用于不同试验内容')
            return duplicate
        data=prepare(p,body);old=latest(p,data['trial_key'])
        if body.get('previous_hash','')!=(store.digest(old) if old else ''):raise HTTPException(409,'试验版本已变化，请刷新后重新编制')
        if body.get('preview_hash')!=store.digest(data):raise HTTPException(409,'预览或原件已变化，请重新核对')
        record={**data,'id':store.uid(),'version':old['version']+1 if old else 1,'state':'draft','author':a['username'],
                'created_at':store.now(),'request_id':request_id,'request_hash':request_hash}
        record['content_hash']=content_hash(record)
        p['integration_trials'].append(record)
        store.audit(c,p,a,'INTEGRATION_TRIAL_SUBMITTED',{'id':record['id'],'trial_key':record['trial_key'],'version':record['version'],'counts':record['comparison']['counts']})
        store.save(c,p)
        return record


@router.get('/projects/{pid}/integration-trials/status')
def status(pid:str,a=Depends(core.actor)):
    with store.connect() as c:
        p=core.project(c,pid,a);items=[]
        for record in p['integration_trials']:
            try: validate(p,record);issue=''
            except HTTPException as exc: issue=str(exc.detail)
            items.append({'record':record,'object_hash':store.digest(record),'source_valid':not issue,'issue':issue,
                          'review_valid':not issue and record['state']=='approved','can_import':False,'can_export':False})
        return {'items':items,'notice':NOTICE}


@router.get('/projects/{pid}/integration-trials/{tid}/exports')
def export(pid:str,tid:str,format:str='json',a=Depends(core.actor)):
    if format not in ('json','xlsx'):raise HTTPException(422,'试验记录支持JSON和Excel')
    with store.transaction() as c:
        p=core.project(c,pid,a);r=core.find(p,'integration_trials',tid)
        if content_hash(r)!=r['content_hash']:raise HTTPException(409,'试验记录内容校验失败')
        try:validate(p,r);issue=''
        except HTTPException as exc:issue=str(exc.detail)
        approved=not issue and r['state']=='approved'
        label='往返试验记录已复核（限定列示字段）' if approved else '往返试验草稿 · 未经有效复核'
        reviews=[x for x in p['reviews'] if x['collection']=='integration_trials' and x['object_id']==tid]
        if format=='json':
            payload=json.dumps({'status':label,'issue':issue,'can_import':False,'can_export':False,'record':r,'reviews':reviews},ensure_ascii=False,indent=2).encode();media='application/json';suffix='json'
        else:
            from .reports import export_report
            tables=[('往返试验说明',['字段','内容'],[['状态',label],['警示',issue],['边界',NOTICE]]+[[FIELD_LABELS[k],r[k]] for k in FIELDS]+[['试验标识',r['trial_key']],['版本',str(r['version'])],['软件回导结果',r['run_status']],['字段对照通过',str(r['comparison']['passed'])],['列示字段匹配率',r['comparison']['coverage_percent']+'%']]),
                    ('逐字段对照',['对象类型','对象标识','字段','基准值 / 单位','返回值 / 单位','差量 / 容差','结果'],
                     [[*x['key'],str(x['before'])+' / '+str(x['before_unit']),str(x['after'])+' / '+str(x['after_unit']),str(x['delta'])+' / '+x['tolerance'],LABELS[x['status']]+(' / '+x['error'] if x['error'] else '')] for x in r['comparison']['rows']]),
                    ('试验原始依据',['用途','原件','校验值','版本','证据'],[[x['role'],x['document']['name'],x['document']['hash'],str(x['document']['version']),x.get('evidence',{}).get('id','对照文件全文')] for x in r['proofs']]),
                    ('试验复核记录',['决定','复核人','时间','意见'],[[x['decision'],x['reviewer'],x['timestamp'],x['reason']] for x in reviews])]
            payload,media,suffix=export_report(p,None,'xlsx',False,table_data=(label,tables))
        sha=hashlib.sha256(payload).hexdigest();folder=store.ROOT/'exports';folder.mkdir(exist_ok=True);path=folder/sha
        if path.exists() and hashlib.sha256(path.read_bytes()).hexdigest()!=sha:raise HTTPException(409,'试验留档文件校验失败')
        if not path.exists():path.write_bytes(payload)
        artifact={'id':store.uid(),'purpose':'integration_trial','trial_id':tid,'trial_key':r['trial_key'],'trial_version':r['version'],'format':format,
                  'quote_id':None,'quote_version':None,'hash':sha,'approved':approved,'timestamp':store.now(),'actor':a['username'],'media':media,'suffix':suffix}
        c.execute('INSERT INTO artifacts VALUES(?,?,?,?)',(artifact['id'],pid,a['tenant'],store.canonical(artifact)));store.audit(c,p,a,'INTEGRATION_TRIAL_EXPORTED',artifact)
        return StreamingResponse(io.BytesIO(payload),media_type=media,headers={'Content-Disposition':'attachment; filename="QiangQuote-roundtrip-V'+str(r['version'])+'.'+suffix+'"','X-Artifact-Hash':sha})
