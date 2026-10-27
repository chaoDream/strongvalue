"""Synthetic convex windows: independent clipping, quadrature and workflow gates."""
import copy
import io
import math
import random
import json
from decimal import Decimal, localcontext
import pytest
from openpyxl import load_workbook
from pypdf import PdfReader
from backend import store
from backend.engine import DomainError
from backend.foundation_horizontal_formwork import calculate
from backend.foundation_polygon_windows import VERSION
from backend.foundation_horizontal_formwork_reports import POLYGON_TITLES
from tests.test_foundation_horizontal_windows import body as rect_body, source, window, FRAME, check
from tests.test_foundation_hole_end_formwork import source as end_source, body as end_body
from tests.test_foundation_step_formwork import body as step_body
from tests.test_foundation_openings import hole
from tests.test_foundation_voids import rectangle
from tests.test_foundation_recesses import interval
from tests.test_foundation_hole_end_windows import setup as end_setup
from tests.test_foundation_round_step_windows import adaptive
from tests.test_route_takeoff import evaluate
from tests.test_post_foundation import equal_body, fixed_body, fixed_added, equal_added, fixed_adjusted, equal_adjusted, combined
from tests.test_polygon_takeoff import post, mutate
from tests.test_workflow import get, approve, upload, priced_project, create

DIAMOND = [(-100,0),(0,-100),(100,0),(0,100)]


def polygon(pts=DIAMOND,eid='test',unit='mm'):
    params = {a+str(i):str(v) for i,p in enumerate(pts,1) for a,v in zip(('x','y'),p)}
    return dict(shape='convex_polygon',parameters=params,evidence={k:eid for k in params},unit=unit,
                coordinate_basis=FRAME,coordinate_evidence_id=eid,scope_confirmed=True)


def body(s,eid='test',face='bottom',pts=DIAMOND):
    b = end_body(s,eid,face) if face.startswith('hole') else step_body(s,eid) if face=='step' else rect_body(s,eid,face)
    b['contact_ranges'][0]['plan_range'] = polygon(pts,eid)
    return b


def replay(r):
    # Rebuild each intermediate from the preceding expressions and original
    # input dimensions; never trust the stored derived values as an oracle.
    dims={k:v for k,v in r['dimensions'].items() if '.derived.' not in k}
    for d in r['contact_ranges']:
        for definition in d['window_expression_definitions']:
            actual=evaluate(definition['formula_ast'],dims)
            assert actual==Decimal(definition['value'])
            dims[definition['dimension']]=str(actual)
    assert dims.keys()==r['dimensions'].keys() and all(Decimal(v)==Decimal(r['dimensions'][k]) for k,v in dims.items())
    check(r)
    with localcontext() as ctx:
        ctx.prec=80
        for d in r['contact_ranges']:
            for trace in [d['polygon_gross_trace'],d['polygon_net_trace']]+[h.get('polygon_cut_trace',h['curved_intersection']) for h in d['apertures']]:
                for q in trace['terms']:
                    assert evaluate(q['formula_ast'],r['dimensions'])==Decimal(q['area'])
                    for k,n in q['formula_asts'].items():assert evaluate(n,r['dimensions'])==Decimal(q[k])
            for q in d['polygon_net_trace']['partitions']:
                for k,n in q['formula_asts'].items():assert evaluate(n,r['dimensions'])==Decimal(q[k])
                for pair in [q['included']]+q['excluded']+q['remaining']:
                    for k,n in pair['formula_asts'].items():assert evaluate(n,r['dimensions'])==Decimal(pair[k])


@pytest.mark.parametrize('shape',['rectangular','circular','stepped_rectangular','stepped_circular'])
@pytest.mark.parametrize('face',['top','bottom','hole_top','hole_bottom'])
@pytest.mark.parametrize('make',[equal_body,fixed_body,fixed_added,equal_added,fixed_adjusted,equal_adjusted,combined,lambda s:combined(s,equal=True)])
def test_all_shapes_ends_final_counts_immutable_and_replay(shape,face,make):
    s=end_source(shape,make) if face.startswith('hole') else source(shape,make)
    b=body(s,face=face);before=copy.deepcopy((s,b));r=calculate(b,s)
    assert before==(s,b) and r['rule_version']==VERSION
    d=r['contact_ranges'][0]
    assert Decimal(d['window_area'])==Decimal('.02')
    with localcontext() as ctx:
        ctx.prec=80
        assert Decimal(r['gross'])==Decimal(d['single_gross_area'])*Decimal(d['total_count'])
    replay(r)


@pytest.mark.parametrize('shape',['stepped_rectangular','stepped_circular'])
@pytest.mark.parametrize('reverse',[False,True])
def test_platform_slanted_window_upper_coverage_and_boundary_order(shape,reverse):
    s=source(shape);pts=[(150,-70),(270,-30),(240,70),(120,30)]
    if reverse:pts.reverse()
    r=calculate(body(s,face='step',pts=pts),s);d=r['contact_ranges'][0]
    assert Decimal(d['single_gross_area'])>0 and Decimal(d['window_upper_area'])>0
    replay(r)


def clip(poly,box):
    # Independent Sutherland-Hodgman polygon/rectangle clipping.
    for axis,limit,above in [(0,box[0],True),(0,box[1],False),(1,box[2],True),(1,box[3],False)]:
        out=[]
        for a,b in zip(poly,poly[1:]+poly[:1]):
            inside=lambda p:p[axis]>=limit if above else p[axis]<=limit
            if inside(a):out.append(a)
            if inside(a)!=inside(b):
                t=(limit-a[axis])/(b[axis]-a[axis]);out.append(tuple(a[k]+t*(b[k]-a[k]) for k in (0,1)))
        poly=out
    return abs(sum(a[0]*b[1]-a[1]*b[0] for a,b in zip(poly,poly[1:]+poly[:1])))/2


def quadrature(pts,disks,boxes=()):
    lo,hi=min(p[0] for p in pts),max(p[0] for p in pts)
    cuts=sorted({lo,hi,*[p[0] for p in pts],*[v for x,y,r in disks for v in (x-r,x,x+r) if lo<v<hi]})
    def f(x):
        ys=[]
        for a,b in zip(pts,pts[1:]+pts[:1]):
            if a[0]!=b[0] and min(a[0],b[0])<=x<=max(a[0],b[0]):ys.append(a[1]+(x-a[0])*(b[1]-a[1])/(b[0]-a[0]))
        if len(ys)<2:return 0.
        bottom,top=min(ys),max(ys)
        for cx,cy,r in disks:
            if abs(x-cx)>=r:return 0.
            h=math.sqrt(max(0,r*r-(x-cx)**2));bottom=max(bottom,cy-h);top=min(top,cy+h)
        for x1,x2,y1,y2 in boxes:
            if not x1<=x<=x2:return 0.
            bottom=max(bottom,y1);top=min(top,y2)
        return max(0.,top-bottom)
    return sum(adaptive(f,a,b,1e-13) for a,b in zip(cuts,cuts[1:]))


def test_seeded_60_slanted_triangles_independent_rectangle_clipping():
    rng=random.Random(10701);s=source()
    for _ in range(60):
        x,y=rng.randint(-350,150),rng.randint(-350,150)
        pts=[(x,y),(x+rng.randint(20,400),y+rng.randint(0,30)),(x+rng.randint(0,30),y+rng.randint(50,400))]
        r=calculate(body(s,pts=pts),s);d=r['contact_ranges'][0]
        expected=(clip(pts,(-200,200,-250,250))-clip(pts,(100,200,-30,30)))/1e6
        assert abs(float(d['single_net_area'])-expected)<2e-16
    replay(r)


def test_shared_expression_graph_size_and_all_input_replay():
    s=end_source('stepped_circular',holes=[{**hole(x='190'),'vertical_range':interval('0','200')}])
    r=calculate(body(s,face='hole_top',pts=[(165,0),(190,-30),(220,0),(190,30)]),s)
    assert len(json.dumps(r,ensure_ascii=False).encode())<1500000
    replay(r)


@pytest.mark.parametrize('delta',['0','0.000000001'])
def test_exact_slanted_tangent_and_tiny_positive_circle_cap(delta):
    s=source('circular');x=str(Decimal('500')-Decimal(delta))
    r=calculate(body(s,face='top',pts=[(0,375),(x,0),(500,500)]),s)
    assert (Decimal(r['contact_ranges'][0]['single_gross_area'])>0)==(delta!='0');replay(r)


def test_twelve_vertex_limit_and_compact_replay():
    pts=[(0,300),(150,260),(260,150),(300,0),(260,-150),(150,-260),(0,-300),(-150,-260),(-260,-150),(-300,0),(-260,150),(-150,260)]
    s=source('circular');r=calculate(body(s,pts=pts),s)
    assert len(r['contact_ranges'][0]['window_vertices'])==12 and len(json.dumps(r).encode())<3000000;replay(r)


@pytest.mark.parametrize('face',['bottom','step','hole_top'])
def test_seeded_24_slanted_circle_intersections_independent_quadrature(face):
    rng=random.Random(10702)
    for _ in range(8):
        cx,cy=rng.randint(140,210),rng.randint(-25,25)
        pts=[(cx-65,cy-20),(cx-5,cy-65),(cx+65,cy+20),(cx+5,cy+65)]
        hs=[{**hole(x='190',diameter='40'),'vertical_range':interval('0','200')}]
        s=end_source('stepped_circular',holes=hs)
        r=calculate(body(s,face=face,pts=pts),s);d=r['contact_ranges'][0];p=[(x/1000,y/1000) for x,y in pts]
        lower=quadrature(p,[(0,0,.3)]);upper=quadrature(p,[(0,0,.2)])
        h=quadrature(p,[(.19,0,.02)]);common=quadrature(p,[(.19,0,.02),(0,0,.2)])
        expected=lower-h if face=='bottom' else lower-upper-h+common if face=='step' else common
        assert abs(float(d['single_net_area'])-expected)<4e-11
    replay(r)


@pytest.mark.parametrize('unit,factor',[('mm','1'),('cm','.1'),('dm','.01'),('m','.001'),('km','.000001')])
def test_units_low_precision_and_both_orientations(unit,factor):
    s=source('circular');b=body(s);expected=calculate(b,s)
    c=b['contact_ranges'][0]['plan_range'];c['unit']=unit
    c['parameters']={k:str(Decimal(v)*Decimal(factor)) for k,v in c['parameters'].items()}
    with localcontext() as ctx:ctx.prec=6;r=calculate(b,s)
    assert r['gross']==expected['gross'] and r['quantity']==expected['quantity'];replay(r)


def test_actual_polygon_separation_shared_diagonal_tiny_overlap_and_rect_mixed():
    s=source();b=body(s,pts=[(0,0),(100,0),(0,100)])
    d=copy.deepcopy(b['contact_ranges'][0]);d['name']='共享斜边';d['plan_range']=polygon([(100,0),(100,100),(0,100)]);b['contact_ranges'].append(d)
    r=calculate(b,s);assert r['horizontal_window_checks'][0]['state']=='boundary_touch';replay(r)
    d['plan_range']=polygon([(100,0),('100','100'),('0','99.999999999')])
    with pytest.raises(DomainError,match='重叠'):calculate(b,s)
    d['plan_range']=window(x1='-100',x2='0',y1='0',y2='100');calculate(b,s)
    d['plan_range']=window(x1='-100',x2='0.000000001',y1='0',y2='100')
    with pytest.raises(DomainError,match='重叠'):calculate(b,s)


@pytest.mark.parametrize('pts',[
    [(0,0),(1,0)],[(0,0),(1,0),(1,0)],[(0,0),(1,0),(2,0)],
    [(0,0),(100,0),(30,30),(0,100)],[(0,0),(100,100),(0,100),(100,0)],
    [(0,100),(59,-81),(-95,31),(95,31),(-59,-81)],
    [(0,0),(100,0),(100,100),(50,100),(0,100)],
])
def test_invalid_concave_repeated_collinear_or_crossed_vertices_blocked(pts):
    s=source()
    with pytest.raises(DomainError):calculate(body(s,pts=pts),s)


@pytest.mark.parametrize('fault',['extra','missing','blank_proof','nan','bool','bad_unit','unconfirmed','truthy','bad_shape','count','foreign_key'])
def test_strict_schema_no_inferred_coordinates(fault):
    s=source();b=body(s);c=b['contact_ranges'][0]['plan_range']
    if fault=='extra':c['angle']='45'
    elif fault=='missing':del c['evidence']['x1']
    elif fault=='blank_proof':c['evidence']['x1']=' '
    elif fault=='nan':c['parameters']['x1']='NaN'
    elif fault=='bool':c['parameters']['x1']=True
    elif fault=='bad_unit':c['unit']='m²'
    elif fault=='unconfirmed':c['scope_confirmed']=False
    elif fault=='truthy':c['scope_confirmed']=1
    elif fault=='bad_shape':c['shape']='arbitrary'
    elif fault=='count':c['parameters']={f'{a}{i}':str(i) for i in range(1,14) for a in ('x','y')}
    else:c['parameters']['x0']=c['parameters'].pop('x1')
    with pytest.raises(DomainError):calculate(b,s)


@pytest.mark.parametrize('gap',['0','0.000000001'])
def test_internal_end_full_coverage_or_positive_gap_not_tolerance_zero(gap):
    v=str(Decimal('25')+Decimal(gap));hs=[{**hole(x='0',diameter='80'),'vertical_range':interval('100','300')},
       {**rectangle(name='H2',x='-'+v,length='50',width='100'),'vertical_range':interval('300','500')},
       {**rectangle(name='H3',x=v,length='50',width='100'),'vertical_range':interval('300','500')}]
    s=end_source(holes=hs);r=calculate(body(s,face='hole_top',pts=[(-50,0),(0,-50),(50,0),(0,50)]),s);d=r['contact_ranges'][0]
    assert ('polygon_zero_contact' in d)==(gap=='0')
    assert (Decimal(d['single_net_area'])==0)==(gap=='0');replay(r)


def test_ten_adjacent_voids_zero_partition_preserves_hole_identity():
    hs=[{**rectangle(x='0',length='80',width='80'),'vertical_range':interval('100','300')}]
    hs += [{**rectangle(name='相接H'+str(i+2),x=str(-45+i*10),length='10',width='100'),'vertical_range':interval('300','500')} for i in range(10)]
    s=end_source(holes=hs);r=calculate(body(s,face='hole_top',pts=[(-50,0),(0,-50),(50,0),(0,50)]),s);d=r['contact_ranges'][0]
    assert Decimal(d['single_net_area'])==0 and d['polygon_zero_contact']['net_trace']['partitions']
    assert all(p['profile'].startswith('实际孔H') for q in d['polygon_net_trace']['partitions'] for p in q['excluded']);replay(r)


def setup(clients,pid=None,zero=False):
    e,r,pid,s,b,eid=end_setup(clients,pid,zero)
    b['contact_ranges'][0]['plan_range']=polygon([(-50,0),(0,-50),(50,0),(0,50)] if zero else [(165,0),(190,-30),(220,0),(190,30)],eid)
    return e,r,pid,s,b,eid


def test_api_readonly_claims_proofs_roles_tenants_history_and_source_stale(clients):
    e,r,pid,s,b,eid=setup(clients);before=get(e,pid)
    post(e,pid,b,'/components/preview');assert get(e,pid)==before
    post(r,pid,b,expected=403);assert clients['other'].get(f'/api/v1/projects/{pid}').status_code==404
    foreign=create(e)['id'];upload(e,foreign,'另一项目合成依据');bad=copy.deepcopy(b);bad['contact_ranges'][0]['plan_range']['evidence']['x1']=get(e,foreign)['evidence'][-1]['id'];post(e,pid,bad,expected=404)
    c=post(e,pid,b).json();approve(r,pid,'components',c);c=next(x for x in get(e,pid)['components'] if x['id']==c['id']);post(e,pid,b,expected=409)
    other=copy.deepcopy(b);other['contact_ranges'][0]['plan_range']=window(eid,'500','600','0','100');post(e,pid,other)
    b['object_hash']=store.digest(c);b['contact_ranges'][0]['plan_range']['parameters']['x1']='160'
    url=f'/api/v1/projects/{pid}/components/'+c['id'];res=e.put(url,json=b);assert res.status_code==200,res.text
    assert get(e,pid)['component_history'][-1]==c and e.put(url,json=b).status_code==409
    approve(r,pid,'components',res.json())
    s['object_hash']=store.digest(s);s['foundation_items'][0]['levels'][1]['parameters']['diameter']='420'
    res=e.put(f'/api/v1/projects/{pid}/components/'+s['id'],json=s);assert res.status_code==200,res.text
    assert next(x for x in get(e,pid)['components'] if x['id']==c['id'])['state']=='stale'


@pytest.mark.parametrize('fault',['vertex','gross','net_trace','edge','ast','frame','book'])
def test_saved_tamper_blocked_on_independent_review(clients,fault):
    e,r,pid,s,b,eid=setup(clients);c=post(e,pid,b).json()
    if fault=='book':approve(r,pid,'components',c);post(e,pid,{},'/takeoffs/calculate')
    def forge(p):
        c=p['components'][-1];d=c['preview']['contact_ranges'][0]
        if fault=='vertex':c['contact_ranges'][0]['plan_range']['parameters']['x1']='150'
        elif fault=='gross':d['single_gross_area']='0'
        elif fault=='net_trace':d['polygon_net_trace']['terms']=[]
        elif fault=='edge':d['polygon_gross_trace']['terms'][0]['upper_boundary']='forged'
        elif fault=='ast':d['formula_asts']['window_area']={'constant':'999'}
        elif fault=='frame':d['plan_frame_hash']='forged'
        else:p['takeoffs'][-1]['quantity']='999'
    mutate(pid,forge);collection='takeoffs' if fault=='book' else 'components';approve(r,pid,collection,get(e,pid)[collection][-1],expected=409)


@pytest.mark.parametrize('gate',['preview','save','review','calculate','book_review'])
def test_original_bytes_checked_at_all_five_gates(clients,gate):
    e,r,pid,s,b,eid=setup(clients)
    if gate in ('review','calculate','book_review'):c=post(e,pid,b).json()
    if gate in ('calculate','book_review'):approve(r,pid,'components',c)
    if gate=='book_review':t=post(e,pid,{},'/takeoffs/calculate').json()[-1]
    doc=get(e,pid)['documents'][-1];(store.ROOT/'originals'/doc['hash']).write_bytes(b'corrupted-synthetic-file')
    if gate=='preview':post(e,pid,b,'/components/preview',409)
    elif gate=='save':post(e,pid,b,expected=409)
    elif gate=='review':approve(r,pid,'components',c,expected=409)
    elif gate=='calculate':post(e,pid,{},'/takeoffs/calculate',409)
    else:approve(r,pid,'takeoffs',t,expected=409)


@pytest.mark.parametrize('fmt',['json','xlsx','pdf'])
@pytest.mark.parametrize('zero',[False,True])
def test_exports_preserve_polygon_and_old_quote_isolation(clients,fmt,zero):
    e,r,pid,_,q=priced_project(clients);e,r,pid,s,b,eid=setup(clients,pid,zero)
    c=post(e,pid,b).json();approve(r,pid,'components',c)
    for t in post(e,pid,{},'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
    post(e,pid,{},'/boq/generate');url=f'/api/v1/projects/{pid}/exports';res=e.get(url,params=dict(format=fmt,unpriced=True));assert res.status_code==200,res.text
    titles=set(POLYGON_TITLES if zero else POLYGON_TITLES[:2])
    if fmt=='json':assert titles<={t[0] for t in res.json()['tables']}
    elif fmt=='xlsx':assert titles<=set(load_workbook(io.BytesIO(res.content)).sheetnames)
    else:assert all(any(t in p.extract_text() for p in PdfReader(io.BytesIO(res.content)).pages) for t in titles)
    assert not any(t[0] in titles for t in e.get(url,params=dict(format='json',quote_id=q['id'])).json()['tables'])
