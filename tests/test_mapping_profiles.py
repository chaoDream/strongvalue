"""Synthetic mapping compatibility and provenance, never official code data."""
import copy
import io
import json
from datetime import date, timedelta
from decimal import Decimal
import pytest
from openpyxl import load_workbook
from pypdf import PdfReader
from backend import store
from backend.engine import GEOMETRIES
from backend.compound_takeoff import STRUCTURED_KINDS
from backend.mapping_types import TYPES, selected_profile
from tests.test_workflow import create, get, upload, approve
from tests.test_boq_mapping import mapping_project, selections


CASES = [(k, p['id']) for k in STRUCTURED_KINDS for p in TYPES[k]['profiles']]


def rule_body(p, kind, profile, eid):
    item = next(v for v in TYPES[kind]['profiles'] if v['id'] == profile)
    return {'name': '合成口径映射（非官方）', 'code': 'TEST-PROFILE', 'item_name': '合成映射项目（非官方编码）',
            'features': '仅用于合成软件验收，须人工核定材料、工法及特征', 'standard': '合成清单标准', 'edition': 'test-2026',
            'license': '仅合成测试使用', 'evidence_id': eid, 'kind': kind, 'measurement_profile': profile, 'unit': item['unit'],
            'keywords': ['合成'], 'discipline': p['scope'][0], 'region': p['region'],
            'valid_from': (date.today()-timedelta(days=30)).isoformat(), 'valid_until': (date.today()+timedelta(days=365)).isoformat()}


def design_project(clients, kind='chamber_plate', profile='cover'):
    if kind == 'compound':
        from tests.test_compound_takeoff import payload
        e, r = clients['estimator'], clients['reviewer']; pid = create(e, ['结构'] if profile == 'mass' else None)['id']
        upload(e, pid, '仅合成尺寸或配筋表：长度2m、宽3m、厚0.2m；钢筋下料2m、3根、每米质量0.5kg/m。', specialty='结构' if profile == 'mass' else '建筑')
        eid = get(e, pid)['evidence'][0]['id']; b = payload(eid)
        g, params = {'length': ('pipe', {'length': '2'}), 'area': ('area', {'length': '2', 'width': '3'}),
                     'volume': ('slab', {'length': '2', 'width': '3', 'thickness': '.2'}),
                     'mass': ('rebar_schedule', {'length': '2', 'count': '3', 'unit_weight': '.5'})}[profile]
        b['measurement_parts'] = [{'name': '合成项', 'operation': 'add', 'kind': g, 'parameters': params, 'multiplicity': '1',
                                   'measurement_basis': '仅合成数学口径，无专业效力', 'evidence': {k: eid for k in [*params, 'multiplicity', 'measurement_basis']}}]
    elif kind in ('planar_regions', 'polygon_region', 'spatial_regions', 'design_route'):
        import importlib
        module = importlib.import_module('tests.test_'+{'planar_regions': 'planar_takeoff', 'polygon_region': 'polygon_takeoff', 'spatial_regions': 'spatial_takeoff', 'design_route': 'route_takeoff'}[kind])
        e, r, pid, b = module.setup(clients)
        if kind in ('planar_regions','polygon_region'): b = module.payload(get(e, pid)['evidence'][0]['id'], profile)
    else:
        import importlib
        module = importlib.import_module('tests.test_'+{'surface_faces': 'surface_takeoff', 'station_trench': 'trench_takeoff',
             'volume_balance': 'volume_balance', 'pipe_envelope': 'pipe_envelope', 'route_bedding': 'bedding_takeoff',
             'circular_chamber': 'chamber_takeoff', 'chamber_plate': 'plate_openings', 'polygon_railing': 'railing_takeoff', 'railing_material':'railing_material'}[kind])
        if kind=='polygon_railing' and profile!='length':
            from tests.test_route_railing import setup as route_railing_setup
            values=route_railing_setup(clients,basis='spatial' if profile=='spatial_length' else 'horizontal')
        else:values=module.setup(clients)
        e, r, pid = values[:3]; b = values[-2]
        if kind in ('surface_faces', 'circular_chamber'): b['purpose'] = profile
        if kind == 'chamber_plate': b['plate_role'] = profile
    base = '/api/v1/projects/'+pid
    response = e.post(base+'/components', json=b); assert response.status_code == 200, response.text
    cid = response.json()['id']
    upload(e, pid, '仅合成编码依据：TEST-PROFILE、TEST-OTHER是不同计量口径的软件测试编码，非官方条目，无真实计价效力。', name='mapping-code.txt', specialty='合同计价')
    for c in get(e, pid)['components']: approve(r, pid, 'components', c)
    response = e.post(base+'/takeoffs/calculate'); assert response.status_code == 200, response.text
    for t in response.json(): approve(r, pid, 'takeoffs', t)
    response = e.post(base+'/boq/generate'); assert response.status_code == 200, response.text
    p = get(e, pid); b = next(b for b in p['boq'] if next(t for t in p['takeoffs'] if t['id'] == b['takeoff_id'])['component_id'] == cid)
    body = rule_body(p, kind, profile, p['evidence'][-1]['id'])
    assert e.patch(base+'/settings', json={'boq_standard': body['standard'], 'boq_edition': body['edition'], 'reason': '仅合成适用标准'}).status_code == 200
    return e, r, pid, cid, b, body


def save_rule(e, r, pid, body):
    response = e.post('/api/v1/projects/'+pid+'/boq-mappings', json=body); assert response.status_code == 200, response.text
    rule = response.json(); approve(r, pid, 'boq_mappings', rule)
    return get(e, pid)['boq_mappings'][-1]


def preview(e, pid, bid):
    response = e.get('/api/v1/projects/'+pid+'/boq/mapping-preview'); assert response.status_code == 200, response.text
    return next(row for row in response.json()['rows'] if row['boq']['id'] == bid)


def selection(row, rule=None):
    rule = rule or row['candidates'][0]['rule']
    return {'boq_id': row['boq']['id'], 'boq_hash': row['boq_hash'], 'rule_id': rule['id'], 'rule_hash': store.digest(rule)}


def test_catalog_keeps_primitive_forms_separate_and_covers_all_kinds(clients):
    e = clients['estimator']
    types = e.get('/api/v1/boq-mapping-types').json()
    assert set(types) == set(GEOMETRIES) | set(STRUCTURED_KINDS)
    assert e.get('/api/v1/takeoff-types').json() == GEOMETRIES
    assert sum(len(v['profiles']) for v in types.values()) == 33
    for kind, v in types.items():
        for p in v['profiles']: assert selected_profile({'kind': kind, 'measurement_profile': p['id'], 'unit': p['unit']}) == p


@pytest.mark.parametrize('kind,profile', CASES)
def test_each_structured_profile_review_preview_apply_and_export(clients, kind, profile):
    e, r, pid, cid, b, body = design_project(clients, kind, profile); base = '/api/v1/projects/'+pid
    rule = save_rule(e, r, pid, body); row = preview(e, pid, b['id'])
    assert row['state'] == 'unique' and row['sources'][0]['measurement_profile'] == profile
    assert row['sources'][0]['component_id'] == cid and row['sources'][0]['unit'] == body['unit']
    assert e.post(base+'/boq/apply-mappings', json={'selections': [selection(row)]}).status_code == 200
    current = next(v for v in get(e, pid)['boq'] if v['id'] == b['id'])
    assert current['quantity'] == b['quantity'] and current['code'] == body['code'] and current['state'] == 'draft'
    assert current['mapping_ref']['hash'] == store.digest(rule)
    approve(r, pid, 'boq', current)
    result = e.get(base+'/exports?unpriced=true&format=json'); assert result.status_code == 200, result.text
    table = next(t for t in result.json()['tables'] if t[0] == '清单映射适用范围')
    assert table[2] and TYPES[kind]['label'] in table[2][0][2]
    assert next(p['label'] for p in TYPES[kind]['profiles'] if p['id'] == profile) in table[2][0][2]


@pytest.mark.parametrize('kind,first,other', [('circular_chamber', 'wall_material', 'outer_envelope'), ('chamber_plate', 'base', 'cover'), ('surface_faces', 'formwork', 'waterproof'), ('polygon_railing','spatial_length','horizontal_projection'), ('polygon_railing','length','spatial_length')])
def test_same_unit_different_purpose_never_matches(clients, kind, first, other):
    e, r, pid, cid, b, body = design_project(clients, kind, first); base = '/api/v1/projects/'+pid
    wrong = save_rule(e, r, pid, {**body, 'measurement_profile': other})
    row = preview(e, pid, b['id']); assert row['state'] == 'unmatched' and not row['source_issue']
    assert e.post(base+'/boq/apply-mappings', json={'selections': [selection(row, wrong)]}).status_code == 409
    save_rule(e, r, pid, {**body, 'code': 'TEST-OTHER'})
    assert preview(e, pid, b['id'])['state'] == 'unique'


@pytest.mark.parametrize('delta', [{'measurement_profile': None}, {'measurement_profile': []}, {'measurement_profile': 'standard'},
    {'measurement_profile': 'base', 'unit': 'm²'}, {'kind': 'invented'}, {'unit': 'kg'}, {'kind': []}])
def test_invalid_profile_or_unit_is_rejected(clients, delta):
    e, r, pid, cid, b, body = design_project(clients)
    response = e.post('/api/v1/projects/'+pid+'/boq-mappings', json={**body, **delta}); assert response.status_code == 422, response.text
    assert not get(e, pid)['boq_mappings']
    body.pop('measurement_profile')
    assert e.post('/api/v1/projects/'+pid+'/boq-mappings', json=body).status_code == 422


def test_legacy_primitive_rule_without_new_fields_still_works(clients):
    e, r, pid, body = mapping_project(clients); base = '/api/v1/projects/'+pid
    rule = e.post(base+'/boq-mappings', json=body).json()
    with store.transaction() as c:
        p = json.loads(c.execute('SELECT data FROM projects WHERE id=?', (pid,)).fetchone()['data']); old = p['boq_mappings'][0]; old.pop('schema'); old.pop('source_stamp'); store.save(c, p)
    approve(r, pid, 'boq_mappings', rule)
    assert e.post(base+'/boq/apply-mappings', json={'selections': selections(e, pid)}).status_code == 200


def test_split_merge_contributions_and_mixed_profile_merge(clients):
    e, r, pid, cid, b, body = design_project(clients, 'chamber_plate', 'cover'); base = '/api/v1/projects/'+pid
    rule = save_rule(e, r, pid, body)
    q = Decimal(b['quantity']); response = e.post(base+'/boq/'+b['id']+'/split', json={'parts': [{'name': '合成分一', 'quantity': str(q/2)}, {'name': '合成分二', 'quantity': str(q-q/2)}]})
    assert response.status_code == 200, response.text
    parts = response.json(); response = e.post(base+'/boq/merge', json={'ids': [x['id'] for x in parts], 'name': '合成同源合并'})
    assert response.status_code == 200, response.text
    merged = response.json(); row = preview(e, pid, merged['id'])
    assert row['state'] == 'unique' and len(row['sources']) == 1 and Decimal(row['sources'][0]['quantity']) == q
    # A different m³ source must prevent the entire merged item from matching.
    p = get(e, pid); chamber = next(x for x in p['boq'] if x['unit'] == 'm³' and x['id'] != merged['id'])
    mixed = e.post(base+'/boq/merge', json={'ids': [merged['id'], chamber['id']], 'name': '合成混合口径合并'}).json()
    row = preview(e, pid, mixed['id']); assert row['state'] == 'unmatched' and len(row['sources']) == 2
    assert e.post(base+'/boq/apply-mappings', json={'selections': [selection(row, rule)]}).status_code == 409


@pytest.mark.parametrize('change', ['book_quantity', 'component_hash', 'book_unit', 'unapproved', 'negative_contribution', 'external', 'book_formula', 'comp_preview', 'original'])
def test_changed_sources_block_preview_and_atomic_apply(clients, change):
    e, r, pid, cid, b, body = design_project(clients); base = '/api/v1/projects/'+pid
    save_rule(e, r, pid, body); row = preview(e, pid, b['id']); chosen = selection(row)
    if change == 'original':
        p = get(e, pid); source = next(c for c in p['components'] if c['id'] == cid)
        eid = next(iter(source['evidence'].values())); evidence = next(v for v in p['evidence'] if v['id'] == eid)
        doc = next(d for d in p['documents'] if d['id'] == evidence['document_id']); (store.ROOT/'originals'/doc['hash']).write_bytes(b'corrupted synthetic original')
    else:
        with store.transaction() as c:
            p = json.loads(c.execute('SELECT data FROM projects WHERE id=?', (pid,)).fetchone()['data']); book = next(t for t in p['takeoffs'] if t['id'] == b['takeoff_id']); comp = next(c for c in p['components'] if c['id'] == cid)
            if change == 'book_quantity': book['quantity'] = '999'
            elif change == 'component_hash': book['component_hash'] = 'stale'
            elif change == 'book_unit': book['unit'] = 'm²'
            elif change == 'unapproved': comp['state'] = 'draft'
            elif change == 'external': book['source_kind'] = 'imported_boq'
            elif change == 'book_formula': book['formula_ast'] = {'constant': '999'}
            elif change == 'comp_preview': comp['preview']['quantity'] = '999'; book['component_hash'] = store.digest(comp)
            elif change == 'negative_contribution':
                line = next(x for x in p['boq'] if x['id'] == b['id']); line['quantity_links'] = [{'takeoff_id': book['id'], 'quantity': '-1'}, {'takeoff_id': book['id'], 'quantity': str(Decimal(b['quantity'])+1)}]
            store.save(c, p)
    with store.connect() as c: before = json.loads(c.execute('SELECT data FROM projects WHERE id=?', (pid,)).fetchone()['data'])['boq']
    row = preview(e, pid, b['id'])
    assert row['state'] == 'unmatched' and row['source_issue']
    assert e.post(base+'/boq/apply-mappings', json={'selections': [chosen]}).status_code == 409
    with store.connect() as c: assert json.loads(c.execute('SELECT data FROM projects WHERE id=?', (pid,)).fetchone()['data'])['boq'] == before


def test_mapping_original_stamp_review_and_exports_are_current(clients):
    e, r, pid, cid, b, body = design_project(clients); base = '/api/v1/projects/'+pid
    rule = save_rule(e, r, pid, body); row = preview(e, pid, b['id'])
    assert e.post(base+'/boq/apply-mappings', json={'selections': [selection(row)]}).status_code == 200
    for format in ['xlsx', 'pdf']:
        response = e.get(base+'/exports?unpriced=true&format='+format); assert response.status_code == 200, response.text
        if format == 'xlsx': assert '清单映射适用范围' in load_workbook(io.BytesIO(response.content)).sheetnames
        else: assert '清单映射适用范围' in ''.join(p.extract_text() for p in PdfReader(io.BytesIO(response.content)).pages)
    with store.transaction() as c:
        p = json.loads(c.execute('SELECT data FROM projects WHERE id=?', (pid,)).fetchone()['data']); evidence = next(x for x in p['evidence'] if x['id'] == body['evidence_id']); evidence['text'] += ' 依据文本发生变化'; store.save(c, p)
    assert preview(e, pid, b['id'])['state'] == 'unmatched'
    approve(r, pid, 'boq_mappings', rule, expected=409)
    approve(r, pid, 'boq', b, expected=409)
    # An explicit revision can bind the new current evidence; it starts as draft.
    response = e.post(base+'/boq-mappings', json={**rule, 'previous_id': rule['id']}); assert response.status_code == 200, response.text
    assert response.json()['state'] == 'draft' and response.json()['source_stamp'] != rule['source_stamp']


def test_structured_mapping_permissions_ambiguity_versions_and_atomic_batch(clients):
    e, r, pid, cid, b, body = design_project(clients); base = '/api/v1/projects/'+pid
    assert r.post(base+'/boq-mappings', json=body).status_code == 403
    assert clients['other'].get(base+'/boq/mapping-preview').status_code == 404
    assert clients['other'].post(base+'/boq-mappings', json=body).status_code == 404
    first = save_rule(e, r, pid, body); save_rule(e, r, pid, {**body, 'code': 'TEST-OTHER'})
    row = preview(e, pid, b['id']); assert row['state'] == 'ambiguous'
    assert r.post(base+'/boq/apply-mappings', json={'selections': [selection(row)]}).status_code == 403
    before = copy.deepcopy(get(e, pid)['boq'])
    assert e.post(base+'/boq/apply-mappings', json={'selections': [selection(row), {**selection(row), 'boq_hash': 'old'}]}).status_code == 422
    assert get(e, pid)['boq'] == before
    assert e.post(base+'/boq/apply-mappings', json={'selections': [selection(row, first)]}).status_code == 200
    revised = e.post(base+'/boq-mappings', json={**body, 'previous_id': first['id'], 'features': '合成新版特征'}).json(); approve(r, pid, 'boq_mappings', revised)
    assert next(v for v in get(e, pid)['boq'] if v['id'] == b['id'])['state'] == 'stale'


def test_distinct_rows_stale_selection_aborts_entire_batch(clients):
    e, r, pid, cid, b, body = design_project(clients); base = '/api/v1/projects/'+pid
    save_rule(e, r, pid, body)
    chamber_rule = rule_body(get(e, pid), 'circular_chamber', 'outer_envelope', body['evidence_id'])
    save_rule(e, r, pid, chamber_rule)
    rows = e.get(base+'/boq/mapping-preview').json()['rows']; chosen = [selection(row) for row in rows if row['candidates']]
    assert len(chosen) == 2
    chosen[-1]['rule_hash'] = 'stale'; before = copy.deepcopy(get(e, pid)['boq'])
    assert e.post(base+'/boq/apply-mappings', json={'selections': chosen}).status_code == 409
    assert get(e, pid)['boq'] == before


def test_all_independent_merged_sources_must_match_keywords_and_profile(clients):
    e, r, pid, cid, b, body = design_project(clients); base = '/api/v1/projects/'+pid
    source = copy.deepcopy(next(c for c in get(e, pid)['components'] if c['id'] == cid))
    for k in ['id', 'object_hash', 'revision']: source.pop(k, None)
    source['name'] = '合成另一井盖板'
    response = e.post(base+'/components', json=source); assert response.status_code == 200, response.text
    other_id = response.json()['id']; approve(r, pid, 'components', response.json())
    response = e.post(base+'/takeoffs/calculate'); assert response.status_code == 200, response.text
    for t in response.json(): approve(r, pid, 'takeoffs', t)
    response = e.post(base+'/boq/generate'); assert response.status_code == 200, response.text
    p = get(e, pid); book_ids = {t['id'] for t in p['takeoffs'] if t.get('component_id') in [cid, other_id]}
    boards = [v for v in p['boq'] if v['takeoff_id'] in book_ids]
    merged = e.post(base+'/boq/merge', json={'ids': [v['id'] for v in boards], 'name': '合成两份独立盖板合并'}).json()
    save_rule(e, r, pid, body)
    save_rule(e, r, pid, {**body, 'code': 'TEST-OTHER', 'keywords': ['偏心']})
    row = preview(e, pid, merged['id'])
    assert row['state'] == 'unique' and len(row['sources']) == 2
    assert row['candidates'][0]['rule']['code'] == 'TEST-PROFILE'
    assert sum(Decimal(s['quantity']) for s in row['sources']) == Decimal(merged['quantity'])


def test_code_original_integrity_blocks_review_apply_and_boq_review(clients):
    e, r, pid, cid, b, body = design_project(clients); base = '/api/v1/projects/'+pid
    rule = save_rule(e, r, pid, body); row = preview(e, pid, b['id']); chosen = selection(row)
    assert e.post(base+'/boq/apply-mappings', json={'selections': [chosen]}).status_code == 200
    p = get(e, pid); doc = next(d for d in p['documents'] if d['name'] == 'mapping-code.txt')
    (store.ROOT/'originals'/doc['hash']).write_bytes(b'bad code original')
    assert preview(e, pid, b['id'])['state'] == 'unmatched'
    approve(r, pid, 'boq_mappings', rule, expected=409)
    approve(r, pid, 'boq', b, expected=409)
    assert any('编码映射需重新核对' in issue for issue in get(e, pid)['blockers'])


def test_structured_mapping_quote_snapshot_survives_rule_revision(clients):
    e, r, pid, cid, b, body = design_project(clients, 'compound', 'volume'); base = '/api/v1/projects/'+pid
    rule = save_rule(e, r, pid, body)
    assert e.post(base+'/boq/apply-mappings', json={'selections': [selection(preview(e, pid, b['id']))]}).status_code == 200
    eid = get(e, pid)['evidence'][0]['id']
    resource = {'name': '合成资源', 'type': '材料', 'unit': 'm³', 'consumption': '1', 'price': '100', 'source': '合成报价',
                'license': '仅测试', 'period': '2026-10', 'valid_until': body['valid_until'], 'region': body['region'], 'tax_mode': '不含税', 'evidence_id': eid}
    response = e.put(base+'/boq/'+b['id'], json={'quotas': [{'code': 'TEST-Q', 'edition': 'test-v1', 'source': '合成', 'license': '仅测试', 'evidence_id': eid, 'coefficient': '1', 'resources': [resource]}]})
    assert response.status_code == 200, response.text
    approve(r, pid, 'boq', response.json())
    rules = e.put(base+'/rules', json={'name': '合成取费', 'version': 'v1', 'region': body['region'], 'standard': '仅合成', 'source': '合成依据',
             'license': '仅测试', 'evidence_id': eid, 'valid_from': body['valid_from'], 'valid_until': body['valid_until'], 'fees': []}).json()
    approve(r, pid, 'rules', rules)
    cov = e.post(base+'/coverage', json={'specialty': '建筑', 'reason': '仅合成软件验收', 'evidence_id': eid}).json(); approve(r, pid, 'coverage', cov)
    response = e.post(base+'/quotes/calculate'); assert response.status_code == 200, response.text
    q = response.json(); assert q['total'] == '120.00' and q['snapshot']['boq_mappings'][0]['measurement_profile'] == 'volume'
    response = e.get(base+'/exports?format=json&quote_id='+q['id']); assert response.status_code == 200, response.text
    previous = copy.deepcopy(q['snapshot'])
    revised = e.post(base+'/boq-mappings', json={**body, 'previous_id': rule['id'], 'features': '合成新特征'}).json(); approve(r, pid, 'boq_mappings', revised)
    assert get(e, pid)['quotes'][0]['snapshot'] == previous
    assert e.post(base+'/quotes/calculate').status_code == 409
