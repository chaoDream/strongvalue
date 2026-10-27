"""Synthetic separated blocks, independent clipping and preserved review gates."""
import copy
import io
import random
from decimal import Decimal,localcontext
from fractions import Fraction
import pytest
from openpyxl import load_workbook
from pypdf import PdfReader
from backend import store
from backend.engine import DomainError
from backend.foundation_horizontal_formwork import calculate
from backend.foundation_multi_windows import VERSION,METHOD,part_config,separation,validate
from backend.foundation_horizontal_windows import checks
from tests.test_foundation_polygon_windows import polygon,body as pb,replay,clip,quadrature
from tests.test_foundation_perforated_windows import excluded,setup as old_setup
from tests.test_foundation_horizontal_windows import source,window
from tests.test_foundation_hole_end_formwork import source as end_source
from tests.test_foundation_openings import hole
from tests.test_foundation_voids import rectangle
from tests.test_foundation_recesses import interval
from tests.test_post_foundation import equal_body,fixed_body,fixed_added,equal_added,fixed_adjusted,equal_adjusted,combined
from tests.test_polygon_takeoff import post,mutate
from tests.test_workflow import get,approve,upload,create,priced_project
from tests.test_route_takeoff import evaluate

LEFT=[(-100,-100),(-5,-100),(-5,100),(-100,100)]
RIGHT=[(5,-100),(100,-100),(100,100),(5,100)]
INNER=[(-60,-15),(-30,-15),(-30,15),(-60,15)]


def part(pts,eid='test',name='合成净块',holes=()):
    return {**excluded(pts,eid,name),'measurement_basis':'合成本块独立支模，无块间空区或附加数量；仅软件验证','holes':[excluded(p,eid,'合成W'+str(i)) for i,p in enumerate(holes,1)]}


def config(parts=None,eid='test',unit='mm'):
    c=polygon(eid=eid,unit=unit);c={k:v for k,v in c.items() if k not in ('parameters','evidence')}
    c.update(shape='multi_polygon',parts=parts if parts is not None else [part(LEFT,eid,'合成左块',[INNER]),part(RIGHT,eid,'合成右块')]);return c


def body(s,eid='test',face='bottom',parts=None):
    b=pb(s,eid,face);b['contact_ranges'][0]['plan_range']=config(parts,eid);return b


def replay_blocks(r):
    replay(r)
    with localcontext() as ctx:
        ctx.prec=80
        for d in r['contact_ranges']:
            for p in d['window_parts']:
                for k,n in p['formula_asts'].items():assert evaluate(n,r['dimensions'])==Decimal(p[k])
                for ring in p.get('window_rings',[]):assert evaluate(ring['formula_ast'],r['dimensions'])==Decimal(ring['area'])


@pytest.mark.parametrize('shape',['rectangular','circular','stepped_rectangular','stepped_circular'])
@pytest.mark.parametrize('face',['top','bottom','hole_top','hole_bottom'])
@pytest.mark.parametrize('make',[equal_body,fixed_body,fixed_added,equal_added,fixed_adjusted,equal_adjusted,combined,lambda s:combined(s,equal=True)])
def test_shapes_ends_final_counts_no_input_mutation_and_replay(shape,face,make):
    s=end_source(shape,make) if face.startswith('hole') else source(shape,make)
    b=body(s,face=face);before=copy.deepcopy((s,b));r=calculate(b,s);d=r['contact_ranges'][0]
    assert before==(s,b) and r['rule_version']==VERSION
    assert Decimal(d['window_area'])==Decimal('.0371') and len(d['window_parts'])==2
    assert d['window_part_checks']==[dict(left='B1',right='B2',relation='strictly_disjoint')]
    replay_blocks(r)


@pytest.mark.parametrize('shape,face',[(s,f) for s in ['rectangular','circular','stepped_rectangular','stepped_circular'] for f in ['top','bottom','hole_bottom','hole_top']]+[('stepped_rectangular','step'),('stepped_circular','step')])
def test_single_block_references_match_combined_gross_hole_net_and_no_extra_count(shape,face):
    s=end_source(shape) if face.startswith('hole') else source(shape)
    parts=None
    if face=='step':parts=[part([(150,-100),(280,-100),(280,-5),(150,-5)],name='下块'),part([(150,5),(280,5),(280,100),(150,100)],name='上块')]
    b=body(s,face=face,parts=parts);r=calculate(b,s);d=r['contact_ranges'][0];children=[]
    for p in b['contact_ranges'][0]['plan_range']['parts']:
        child=copy.deepcopy(b);child['contact_ranges'][0]['plan_range']=part_config(b['contact_ranges'][0]['plan_range'],p)
        children.append(calculate(child,s)['contact_ranges'][0])
    with localcontext() as ctx:
        ctx.prec=80
        for k in ['single_gross_area','single_opening_area','single_net_area']:
            assert abs(Decimal(d[k])-sum((Decimal(c[k]) for c in children),Decimal(0)))<Decimal('1e-75')
        assert abs(Decimal(d['total_net_area'])-sum((Decimal(c['total_net_area']) for c in children),Decimal(0)))<Decimal('1e-73')
    replay_blocks(r)


def test_40_seeded_blocks_with_inner_diamonds_independent_fraction_clip():
    rng=random.Random(11001);s=source()
    for _ in range(40):
        parts=[]
        for i,(x,X) in enumerate([(-300,-20),(20,300)],1):
            y=rng.randint(-300,-50);Y=rng.randint(50,300);outer=[(x,y),(X,y),(X,Y),(x,Y)]
            cx=(x+X)//2;cy=(y+Y)//2;rx=rng.randint(5,60);ry=rng.randint(5,40)
            parts.append(part(outer,name='合成B'+str(i),holes=[[(cx-rx,cy),(cx,cy-ry),(cx+rx,cy),(cx,cy+ry)]]))
        r=calculate(body(s,parts=parts),s);d=r['contact_ranges'][0]
        def measured(p):
            pts=[tuple(map(Fraction,p)) for p in polygon_vertices(p)]
            return clip(pts,tuple(map(Fraction,(-200,200,-250,250))))-clip(pts,tuple(map(Fraction,(100,200,-30,30))))
        expected=sum(measured(p)-sum(measured(h) for h in p['holes']) for p in parts)/1000000
        assert abs(float(d['single_net_area'])-float(expected))<2e-16
    replay_blocks(r)


def polygon_vertices(p):
    return [(p['parameters']['x'+str(i)],p['parameters']['y'+str(i)]) for i in range(1,len(p['parameters'])//2+1)]


@pytest.mark.parametrize('offset',[-350,-200,-80,0,80,200,350])
def test_curved_blocks_independent_meter_adaptive_integral(offset):
    s=source('circular');s['foundation_items'][0].pop('through_holes',None)
    low=[(-200+offset,-250),(200+offset,-250),(200+offset,-10),(-200+offset,-10)]
    high=[(-200+offset,10),(200+offset,10),(200+offset,250),(-200+offset,250)]
    inner=[(-30+offset,-100),(30+offset,-100),(30+offset,-40),(-30+offset,-40)]
    parts=[part(low,name='下块',holes=[inner]),part(high,name='上块')]
    r=calculate(body(s,face='top',parts=parts),s);d=r['contact_ranges'][0]
    meters=lambda pts:[(x/1000,y/1000) for x,y in pts]
    expected=sum(quadrature(meters(pts),[(0,0,.3)]) for pts in [low,high])-quadrature(meters(inner),[(0,0,.3)])
    assert abs(float(d['single_net_area'])-expected)<3e-11;replay_blocks(r)


@pytest.mark.parametrize('unit,factor',[('mm','1'),('cm','.1'),('dm','.01'),('m','.001'),('km','.000001')])
@pytest.mark.parametrize('reverse',[False,True])
def test_common_units_orientation_and_low_caller_precision(unit,factor,reverse):
    s=source();b=body(s);c=b['contact_ranges'][0]['plan_range'];c['unit']=unit
    for p in c['parts']:
        for ring in [p]+p['holes']:
            pts=polygon_vertices(ring)
            if reverse:pts.reverse()
            ring['parameters']={a+str(i):str(Decimal(str(v))*Decimal(factor)) for i,xy in enumerate(pts,1) for a,v in zip(('x','y'),xy)}
    with localcontext() as ctx:ctx.prec=6;r=calculate(b,s)
    assert Decimal(r['contact_ranges'][0]['window_area'])==Decimal('.0371');replay_blocks(r)


def test_four_blocks_48_points_and_eight_holes_limits():
    poly=[(-20,-10),(-16,-16),(-10,-20),(10,-20),(16,-16),(20,-10),(20,10),(16,16),(10,20),(-10,20),(-16,16),(-20,10)]
    parts=[part([(x+cx,y) for x,y in poly],name='合成B'+str(cx)) for cx in [-120,-40,40,120]]
    s=source();r=calculate(body(s,parts=parts),s);assert len(r['contact_ranges'][0]['window_parts'])==4;replay_blocks(r)
    parts[0]['holes']=[excluded([(-122,-2),(-118,-2),(-120,2)])]
    with pytest.raises(DomainError,match='48'):validate(config(parts))
    parts=[]
    for i,cx in enumerate([-120,-40,40,120],1):
        outer=[(cx-30,-40),(cx+30,-40),(cx+30,40),(cx-30,40)]
        holes=[[(cx-4,y-4),(cx+4,y-4),(cx,y+4)] for y in [-20,0]]
        parts.append(part(outer,name='合成B'+str(i),holes=holes))
    r=calculate(body(s,parts=parts),s);replay_blocks(r)
    parts[0]['holes'].append(excluded([(-124,16),(-116,16),(-120,24)],name='合成W3'))
    with pytest.raises(DomainError,match='8'):validate(config(parts))


@pytest.mark.parametrize('pts',[
    [(-5,-100),(50,-100),(50,100),(-5,100)],
    [(-5,100),(50,100),(50,150),(-5,150)],
    [(-10,-100),(50,-100),(50,100),(-10,100)],
    [(-90,-50),(-20,-50),(-20,50),(-90,50)],LEFT])
def test_blocks_touch_overlap_or_contain_are_rejected(pts):
    with pytest.raises(DomainError):validate(config([part(LEFT,name='A'),part(pts,name='B')]))


def test_concave_bay_block_allowed_but_nested_island_in_inner_hole_rejected():
    from tests.test_foundation_simple_windows import C_SHAPE
    c=config([part(C_SHAPE,name='C形'),part([(0,-20),(80,-20),(80,20),(0,20)],name='凹口内另一块')]);validate(c)
    c=config([part([(-100,-100),(100,-100),(100,100),(-100,100)],name='带孔大块',holes=[[(-40,-40),(40,-40),(40,40),(-40,40)]]),part([(-10,-10),(10,-10),(10,10),(-10,10)],name='嵌套岛')])
    with pytest.raises(DomainError,match='嵌套'):validate(c)


@pytest.mark.parametrize('gap',['0','0.000000001'])
def test_exact_block_gap_no_tolerance(gap):
    x=str(Decimal(-5)+Decimal(gap));c=config([part(LEFT,name='A'),part([(x,-100),(100,-100),(100,100),(x,100)],name='B')])
    if gap=='0':
        with pytest.raises(DomainError):validate(c)
    else:validate(c)


@pytest.mark.parametrize('fault',['one','five','null','extra','part_unit','part_z','part_extra','part_holes','duplicate','missing_basis','missing_evidence','confirm','unit','manual_area'])
def test_strict_schema_and_block_specific_evidence(fault):
    c=config()
    if fault=='one':c['parts']=c['parts'][:1]
    elif fault=='five':c['parts']*=3
    elif fault=='null':c['parts']=None
    elif fault=='extra':c['parameters']={}
    elif fault=='part_unit':c['parts'][1]['unit']='mm'
    elif fault=='part_z':c['parts'][1]['z']='0'
    elif fault=='part_extra':c['parts'][1]['multiplicity']='2'
    elif fault=='part_holes':c['parts'][1]['holes']={}
    elif fault=='duplicate':c['parts'][1]['name']=c['parts'][0]['name']
    elif fault=='missing_basis':c['parts'][1]['evidence_id']=''
    elif fault=='missing_evidence':c['parts'][1]['evidence']['x1']=''
    elif fault=='confirm':c['scope_confirmed']=1
    elif fault=='unit':c['unit']='px'
    else:c['area']='100'
    with pytest.raises(DomainError):validate(c)


def test_100_seeded_claims_against_independent_fraction_clip():
    rng=random.Random(11002);a={'plan_range':config()}
    for _ in range(100):
        x,y=rng.randint(-140,140),rng.randint(-140,140);box=tuple(map(Fraction,(x,x+rng.randint(1,120),y,y+rng.randint(1,120))))
        def area(p):return clip([tuple(map(Fraction,p)) for p in polygon_vertices(p)],box)
        expected=sum(area(p)-sum(area(h) for h in p['holes']) for p in a['plan_range']['parts'])>0
        assert separation(a,{'plan_range':window('test',*[str(v) for v in box])})['overlap']==expected
    b=dict(index=2,name='间隙',material_index=1,face='top',plan_range=window('test','-5','5','-100','100'))
    a.update(index=1,name='两块',material_index=1,face='top');q=checks([a,b])[0]
    assert q['method']==METHOD and q['left_index']==1 and q['partitions']==separation(a,b)['partitions']


@pytest.mark.parametrize('gap',['0','0.000000001'])
def test_all_blocks_required_for_zero_proof_tiny_positive_gap_not_zero(gap):
    v=str(Decimal(25)+Decimal(gap));hs=[{**hole(x='0',diameter='80'),'vertical_range':interval('100','300')},
      {**rectangle(name='H2',x='-'+v,length='50',width='100'),'vertical_range':interval('300','500')},
      {**rectangle(name='H3',x=v,length='50',width='100'),'vertical_range':interval('300','500')}]
    s=end_source(holes=hs);parts=[part([(-50,-50),(50,-50),(50,-5),(-50,-5)],name='下块'),part([(-50,5),(50,5),(50,50),(-50,50)],name='上块')]
    r=calculate(body(s,face='hole_top',parts=parts),s);d=r['contact_ranges'][0]
    assert ('polygon_zero_contact' in d)==(gap=='0') and (Decimal(d['single_net_area'])==0)==(gap=='0')
    assert {q['block_index'] for q in d['polygon_net_trace']['partitions']}=={1,2};replay_blocks(r)


def test_empty_first_block_does_not_hide_second_block_or_fill_gap():
    s=source();parts=[part([(-600,-100),(-500,-100),(-500,100),(-600,100)],name='实体外空块'),part(RIGHT,name='有效块')]
    r=calculate(body(s,parts=parts),s);d=r['contact_ranges'][0]
    assert Decimal(d['window_parts'][0]['single_gross_area'])==0 and Decimal(d['single_net_area'])==Decimal('.019')
    assert Decimal(d['x_start'])==Decimal('-.6') and Decimal(d['x_end'])==Decimal('.1');replay_blocks(r)


def test_one_zero_block_cannot_certify_whole_zero_and_block_order_does_not_change_quantity():
    hs=[{**hole(x='0',diameter='80'),'vertical_range':interval('100','300')},
        {**rectangle(name='H2',x='-25',length='50',width='100'),'vertical_range':interval('300','500')}]
    s=end_source(holes=hs);b=body(s,face='hole_top');r=calculate(b,s);d=r['contact_ranges'][0]
    assert 'polygon_zero_contact' in d['window_parts'][0] and 'polygon_zero_contact' not in d['window_parts'][1]
    assert 'polygon_zero_contact' not in d and Decimal(d['single_net_area'])>0
    b['contact_ranges'][0]['plan_range']['parts'].reverse();reverse=calculate(b,s)
    assert reverse['quantity']==r['quantity'] and 'polygon_zero_contact' not in reverse['contact_ranges'][0]
    replay_blocks(r);replay_blocks(reverse)


def setup(clients,pid=None,zero=False):
    e,r,pid,s,b,eid=old_setup(clients,pid,zero)
    low=[(-50,-50),(50,-50),(50,-5),(-50,-5)] if zero else [(170,-30),(210,-30),(210,-5),(170,-5)]
    high=[(-50,5),(50,5),(50,50),(-50,50)] if zero else [(170,5),(210,5),(210,30),(170,30)]
    inner=[(-20,-20),(-10,-20),(-10,-10),(-20,-10)] if zero else [(193,-17),(196,-17),(196,-15),(193,-15)]
    b['contact_ranges'][0]['plan_range']=config([part(low,eid,'合成下块',[inner]),part(high,eid,'合成上块')],eid)
    return e,r,pid,s,b,eid


def test_workflow_readonly_roles_foreign_block_history_concurrency_and_source_stale(clients):
    e,r,pid,s,b,eid=setup(clients);before=get(e,pid);post(e,pid,b,'/components/preview');assert get(e,pid)==before
    post(r,pid,b,expected=403);assert clients['other'].get(f'/api/v1/projects/{pid}').status_code==404
    foreign=create(e)['id'];upload(e,foreign,'异项目合成');bad=copy.deepcopy(b);bad['contact_ranges'][0]['plan_range']['parts'][1]['evidence_id']=get(e,foreign)['evidence'][-1]['id'];post(e,pid,bad,expected=404)
    c=post(e,pid,b).json();approve(r,pid,'components',c);c=next(x for x in get(e,pid)['components'] if x['id']==c['id'])
    b['object_hash']=store.digest(c);b['contact_ranges'][0]['plan_range']['parts'][1]['parameters']['x1']='171'
    res=e.put(f'/api/v1/projects/{pid}/components/'+c['id'],json=b);assert res.status_code==200,res.text
    assert e.put(f'/api/v1/projects/{pid}/components/'+c['id'],json=b).status_code==409
    assert get(e,pid)['component_history'][-1]==c;approve(r,pid,'components',res.json())
    s['object_hash']=store.digest(s);s['foundation_items'][0]['levels'][1]['parameters']['diameter']='420'
    res=e.put(f'/api/v1/projects/{pid}/components/'+s['id'],json=s);assert res.status_code==200,res.text;approve(r,pid,'components',res.json())
    assert next(x for x in get(e,pid)['components'] if x['id']==c['id'])['state']=='stale'


def test_cross_record_gap_inner_region_allowed_overlap_and_wrong_frame_rejected(clients):
    e,r,pid,s,b,eid=setup(clients);post(e,pid,b)
    other=copy.deepcopy(b);other['name']='间隙另项';other['contact_ranges'][0]['plan_range']=window(eid,'170','210','-5','5')
    bad=copy.deepcopy(other);bad['contact_ranges'][0]['plan_range']['parameters']['y_end']='5.000000001';before=get(e,pid);post(e,pid,bad,expected=409);assert get(e,pid)==before
    post(e,pid,other)
    other['name']='内孔空区另项';other['contact_ranges'][0]['plan_range']=window(eid,'193','196','-17','-15');post(e,pid,other)
    bad=copy.deepcopy(other);bad['contact_ranges'][0]['plan_range']['coordinate_basis']='另一方向';post(e,pid,bad,expected=409)


@pytest.mark.parametrize('fault',['block_input','block_basis','block_quantity','whole_quantity','trace_index','topology','child_proof','whole_proof','expression'])
def test_tamper_block_or_aggregate_rejected(clients,fault):
    e,r,pid,s,b,eid=setup(clients,zero=True);c=post(e,pid,b).json()
    def forge(p):
        c=p['components'][-1];d=c['preview']['contact_ranges'][0]
        if fault=='block_input':c['contact_ranges'][0]['plan_range']['parts'][1]['parameters']['x1']='-49'
        elif fault=='block_basis':c['contact_ranges'][0]['plan_range']['parts'][1]['measurement_basis']='篡改支模范围'
        elif fault=='block_quantity':d['window_parts'][1]['single_gross_area']='999'
        elif fault=='whole_quantity':d['window_area']='999'
        elif fault=='trace_index':d['polygon_net_trace']['partitions'][0]['block_index']=99
        elif fault=='topology':d['window_part_checks']=[]
        elif fault=='child_proof':d['window_parts'][1]['polygon_zero_contact']['net_trace']['partitions'][0]['remaining']=[{}]
        elif fault=='whole_proof':d['polygon_zero_contact']['net_trace']['partitions'][0]['remaining']=[{}]
        else:d['window_expression_definitions'][0]['value']='999'
    mutate(pid,forge);approve(r,pid,'components',c,expected=409)


@pytest.mark.parametrize('gate',['preview','save','review','calculate','book_review'])
def test_second_block_original_bytes_all_five_gates(clients,gate):
    e,r,pid,s,b,eid=setup(clients)
    upload(e,pid,'合成第二净块支模范围及坐标的独立原件','second-block-basis.txt')
    for c in get(e,pid)['components']:approve(r,pid,'components',c)
    b['source_component_hash']=e.get(f'/api/v1/projects/{pid}/foundation-horizontal-formwork-candidates',params={'component_id':s['id']}).json()['source_component_hash']
    block=b['contact_ranges'][0]['plan_range']['parts'][1];block['evidence_id']=get(e,pid)['evidence'][-1]['id'];block['evidence']={k:block['evidence_id'] for k in block['evidence']}
    if gate in ('review','calculate','book_review'):c=post(e,pid,b).json()
    if gate in ('calculate','book_review'):approve(r,pid,'components',c)
    if gate=='book_review':t=post(e,pid,{},'/takeoffs/calculate').json()[-1]
    doc=get(e,pid)['documents'][-1];(store.ROOT/'originals'/doc['hash']).write_bytes(b'corrupted-synthetic-second-block-original')
    if gate=='preview':post(e,pid,b,'/components/preview',409)
    elif gate=='save':post(e,pid,b,expected=409)
    elif gate=='review':approve(r,pid,'components',c,expected=409)
    elif gate=='calculate':post(e,pid,{},'/takeoffs/calculate',409)
    else:approve(r,pid,'takeoffs',t,expected=409)


@pytest.mark.parametrize('fmt',['json','xlsx','pdf'])
@pytest.mark.parametrize('zero',[False,True])
def test_three_formats_and_old_quote_isolation(clients,fmt,zero):
    from backend.foundation_multi_window_reports import TITLES
    e,r,pid,_,q=priced_project(clients);e,r,pid,s,b,eid=setup(clients,pid,zero);c=post(e,pid,b).json();approve(r,pid,'components',c)
    for t in post(e,pid,{},'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
    post(e,pid,{},'/boq/generate');url=f'/api/v1/projects/{pid}/exports';res=e.get(url,params=dict(format=fmt,unpriced=True));assert res.status_code==200,res.text
    titles=set(TITLES[:5] if zero else TITLES[:4])
    if fmt=='json':assert titles<={t[0] for t in res.json()['tables']}
    elif fmt=='xlsx':assert titles<=set(load_workbook(io.BytesIO(res.content)).sheetnames)
    else:assert all(any(title in page.extract_text() for page in PdfReader(io.BytesIO(res.content)).pages) for title in titles)
    assert not any(t[0] in titles for t in e.get(url,params=dict(format='json',quote_id=q['id'])).json()['tables'])
