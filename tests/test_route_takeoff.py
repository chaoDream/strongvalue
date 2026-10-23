"""Synthetic design routes, not real design dimensions or professional acceptance."""
import copy
import io
import json
import math
import random
from decimal import Decimal, ROUND_HALF_UP, localcontext
import pytest
from openpyxl import load_workbook
from pypdf import PdfReader
from backend import store
from backend.engine import DomainError
from backend.route_takeoff import calculate
from tests.test_workflow import create, get, upload, approve, priced_project


def point(name, x, y, z, eid='test', unit='m'):
    return {'name': name, 'parameters': dict(zip(['x','y','z'], map(str,[x,y,z]))), 'unit':unit, 'evidence': {k:eid for k in ['x','y','z']}}


def payload(eid='test'):
    return {'name':'合成设计管线（非真实工程）','kind':'design_route','parameters':{},'quantity_type':'length',
            'route_points':[point('起点',0,0,0,eid),point('坡段终点',3000,4000,12000,eid,'mm'),point('立管终点',3,4,17,eid)],
            'coordinate_basis':'共同设计原点，X东Y北Z向上，同一高程基准；仅合成测试', 'coordinate_evidence_id':eid,
            'measurement_basis':'按已确认设计中心线顺序连接，不计弯头管件附加量；仅合成数学验收','basis_evidence_id':eid,
            'multiplicity':'1','multiplicity_evidence_id':eid,'rounding_digits':3,'rounding_evidence_id':eid}


def setup(clients):
    e,r=clients['estimator'],clients['reviewer'];pid=create(e)['id']
    upload(e,pid,'仅合成管线：共同设计XYZ基准。起点(0,0,0)m，坡段终点(3000,4000,12000)mm，立管终点(3,4,17)m。3-4-12斜段13m+竖段5m=18m，重复1，总量四舍五入3位。不计管件。')
    return e,r,pid,payload(get(e,pid)['evidence'][0]['id'])


def evaluate(node, dimensions):
    if 'dimension' in node:return Decimal(dimensions[node['dimension']])
    if 'constant' in node:return Decimal(node['constant'])
    values=[evaluate(a,dimensions) for a in node['args']]
    if node['op']=='divide':return values[0]/values[1]
    if node['op']=='sqrt':return values[0].sqrt()
    if node['op']=='round':return values[0].quantize(Decimal(1).scaleb(-node['digits']),rounding=ROUND_HALF_UP)
    if node['op']=='add':return sum(values,Decimal(0))
    if node['op']=='subtract':return values[0]-values[1]
    result=Decimal(1)
    for value in values:result*=value
    return result


def test_true_3d_slope_vertical_mixed_units_and_ast():
    body=payload();before=copy.deepcopy(body);r=calculate(body)
    assert body==before and r['quantity']=='18.000' and r['unit']=='m'
    assert [Decimal(s['length']) for s in r['route_segments']]==[13,5]
    assert [Decimal(s['horizontal_length']) for s in r['route_segments']]==[5,0]
    assert r['route_points'][1]['normalized_parameters']=={'x':'3','y':'4','z':'12'}
    assert r['unit_conversions'][3]['original_value']=='3000' and r['unit_conversions'][3]['factor']=='0.001'
    with localcontext() as ctx:
        ctx.prec=80
        assert evaluate(r['formula_ast'],r['dimensions'])==18
        for s in r['route_segments']:assert evaluate(s['formula_ast'],r['dimensions'])==Decimal(s['length'])
    assert len(r['source_evidence'])==13


def test_round_only_total_after_repetition_and_explicit_zero_digits():
    b=payload();b['route_points']=[point('a',0,0,0),point('b',1,1,0)];b['multiplicity']='2';b['rounding_digits']=2
    r=calculate(b);assert r['quantity']=='2.83' and len(r['route_segments'][0]['length'])>70
    with localcontext() as ctx:
        ctx.prec=80
        assert evaluate(r['formula_ast'],r['dimensions'])==Decimal('2.83')
        assert Decimal(r['quantity'])-Decimal(r['gross'])==Decimal(r['route_summary']['rounding_adjustment'])
    b['rounding_digits']=0;assert calculate(b)['quantity']=='3'
    b['route_points']=[point('a',0,0,0),point('b','2.5',0,0)];b['multiplicity']='1'
    assert calculate(b)['quantity']=='3'


def test_signed_coordinates_translation_reversal_and_units():
    b=payload();b['route_points']=[point('a',-1,-2,-3),point('b',2,2,9),point('c',2,2,14)]
    assert calculate(b)['quantity']=='18.000'
    b['route_points'].reverse();assert calculate(b)['quantity']=='18.000'
    for p in b['route_points']:
        p['unit']='cm';p['parameters']={k:str((Decimal(v)+100)*100) for k,v in p['parameters'].items()}
    assert calculate(b)['quantity']=='18.000'


def test_seeded_routes_match_independent_float_distance_and_ast():
    rng=random.Random(2901)
    for _ in range(35):
        b=payload();coords=[tuple(rng.randint(-100,100) for _ in range(3)) for _ in range(8)]
        b['route_points']=[point(str(i),*v) for i,v in enumerate(coords)];b['rounding_digits']=9
        r=calculate(b);expected=sum(math.dist(a,c) for a,c in zip(coords,coords[1:]))
        assert math.isclose(float(r['quantity']),expected,rel_tol=0,abs_tol=1e-9)
        with localcontext() as ctx:
            ctx.prec=80
            assert evaluate(r['formula_ast'],r['dimensions'])==Decimal(r['quantity'])


def test_crossing_projection_does_not_connect_or_shorten_3d_route():
    b=payload();b['route_points']=[point('a',0,0,0),point('b',2,2,0),point('c',0,2,2),point('d',2,0,2)]
    r=calculate(b);assert len(r['route_segments'])==3 and r['quantity']=='8.485'


@pytest.mark.parametrize('case',['too_few','too_many','duplicate_name','same_point','closed','overlap','missing_z','missing_z_proof','unknown_axis','unknown_field','nan','unit','precision','count','rounding_missing','rounding_float','rounding_bool','rounding_range','missing_rounding_proof','mode','parameters'])
def test_bad_routes_are_rejected(case):
    b=payload();p=b['route_points'][0]
    if case=='too_few':b['route_points']=b['route_points'][:1]
    elif case=='too_many':b['route_points']*=35
    elif case=='duplicate_name':b['route_points'][1]['name']=p['name']
    elif case=='same_point':b['route_points'][1]['parameters']=dict(p['parameters'])
    elif case=='closed':b['route_points'][2]['parameters']=dict(p['parameters'])
    elif case=='overlap':b['route_points']=[point('a',0,0,0),point('b',3,3,3),point('c',1,1,1)]
    elif case=='missing_z':del p['parameters']['z']
    elif case=='missing_z_proof':del p['evidence']['z']
    elif case=='unknown_axis':p['parameters']['elevation']='0'
    elif case=='unknown_field':p['scale']='100'
    elif case=='nan':p['parameters']['z']='NaN'
    elif case=='unit':p['unit']='m³'
    elif case=='precision':p['unit']='mm';p['parameters']['z']='.000000000000000001'
    elif case=='count':b['multiplicity']='1.1'
    elif case=='rounding_missing':del b['rounding_digits']
    elif case=='rounding_float':b['rounding_digits']=2.0
    elif case=='rounding_bool':b['rounding_digits']=True
    elif case=='rounding_range':b['rounding_digits']=13
    elif case=='missing_rounding_proof':b['rounding_evidence_id']=''
    elif case=='mode':b['quantity_type']='area'
    else:b['parameters']={'length':'1'}
    with pytest.raises(DomainError):calculate(b)


def test_disjoint_and_subdivided_collinear_segments_are_allowed():
    b=payload();b['route_points']=[point('a',0,0,0),point('b',1,1,1),point('c',2,2,2)]
    assert len(calculate(b)['route_segments'])==2
    b['route_points']=[point('a',0,0,0),point('b',1,0,0),point('c',1,1,0),point('d',3,1,0),point('e',3,0,0),point('f',4,0,0)]
    assert calculate(b)['quantity']=='6.000'


def test_workflow_boq_exports_and_no_write_preview(clients):
    e,r,pid,b=setup(clients);base='/api/v1/projects/'+pid;before=get(e,pid)
    preview=e.post(base+'/components/preview',json=b);assert preview.status_code==200,preview.text
    assert get(e,pid)==before
    response=e.post(base+'/components',json=b);assert response.status_code==200,response.text
    c=response.json();assert e.post(base+'/takeoffs/calculate').status_code==409
    approve(r,pid,'components',c)
    response=e.post(base+'/takeoffs/calculate');assert response.status_code==200,response.text
    book=response.json()[0];approve(r,pid,'takeoffs',book)
    assert e.post(base+'/boq/generate').json()[0]['quantity']=='18.000'
    j=e.get(base+'/exports?unpriced=true&format=json').json()['unpriced']['takeoffs'][0]
    assert j['route_segments']==book['route_segments'] and j['formula_ast']==book['formula_ast']
    response=e.get(base+'/exports?unpriced=true&format=xlsx');assert response.status_code==200,response.text
    wb=load_workbook(io.BytesIO(response.content));assert {'设计管线长度汇总','管线设计坐标','三维管段计算','算量原始单位换算'}<=set(wb.sheetnames)
    assert all(c.data_type!='f' for sheet in wb for row in sheet for c in row)
    response=e.get(base+'/exports?unpriced=true&format=pdf');assert response.status_code==200,response.text
    text='\n'.join(p.extract_text() for p in PdfReader(io.BytesIO(response.content)).pages)
    assert '设计管线长度汇总' in text and '三维管段计算' in text and '18.000' in text


def test_concurrent_revision_atomicity_and_history(clients):
    e,r,pid,b=setup(clients);base='/api/v1/projects/'+pid
    c=e.post(base+'/components',json=b).json();approve(r,pid,'components',c)
    book=e.post(base+'/takeoffs/calculate').json()[0];approve(r,pid,'takeoffs',book)
    assert e.put(base+'/components/'+c['id'],json=b).status_code==409
    b['object_hash']=e.get(base+'/hashes').json()['components'][c['id']]
    bad=copy.deepcopy(b);bad['route_points'][2]['parameters']['z']='12';before=get(e,pid)
    assert e.put(base+'/components/'+c['id'],json=bad).status_code==422 and get(e,pid)==before
    b['route_points'][2]['parameters']['z']='18'
    response=e.put(base+'/components/'+c['id'],json=b);assert response.status_code==200,response.text
    assert response.json()['revision']==2 and response.json()['preview']['quantity']=='19.000'
    assert get(e,pid)['takeoffs'][0]['state']=='stale'
    assert e.put(base+'/components/'+c['id'],json=b).status_code==409
    changed={'kind':'pipe','parameters':{'length':'1','deduction':'0'},'object_hash':e.get(base+'/hashes').json()['components'][c['id']]}
    assert e.put(base+'/components/'+c['id'],json=changed).status_code==409


@pytest.mark.parametrize('gate',['preview','save','component_review','calculate','book_review'])
def test_corrupt_original_blocks_all_gates(clients,gate):
    e,r,pid,b=setup(clients);base='/api/v1/projects/'+pid
    if gate not in ('preview','save'):c=e.post(base+'/components',json=b).json()
    if gate in ('calculate','book_review'):approve(r,pid,'components',c)
    if gate=='book_review':book=e.post(base+'/takeoffs/calculate').json()[0]
    doc=get(e,pid)['documents'][0];(store.ROOT/'originals'/doc['hash']).write_bytes(b'corrupt')
    if gate=='preview':assert e.post(base+'/components/preview',json=b).status_code==409
    elif gate=='save':assert e.post(base+'/components',json=b).status_code==409
    elif gate=='component_review':approve(r,pid,'components',c,expected=409)
    elif gate=='calculate':assert e.post(base+'/takeoffs/calculate').status_code==409
    else:approve(r,pid,'takeoffs',book,expected=409)


def test_permissions_foreign_proof_and_replaced_original(clients):
    e,r,pid,b=setup(clients);base='/api/v1/projects/'+pid
    assert r.post(base+'/components/preview',json=b).status_code==403
    assert r.post(base+'/components',json=b).status_code==403
    assert clients['other'].post(base+'/components',json=b).status_code==404
    pid2=create(e)['id'];upload(e,pid2,'其他工程')
    bad=copy.deepcopy(b);bad['rounding_evidence_id']=get(e,pid2)['evidence'][0]['id']
    assert e.post(base+'/components',json=bad).status_code==404
    c=e.post(base+'/components',json=b).json();approve(r,pid,'components',c)
    upload(e,pid,'新管线坐标版本',name='fixture.txt')
    assert get(e,pid)['components'][0]['state']=='stale'
    approve(r,pid,'components',c,expected=409)


@pytest.mark.parametrize('target',['component','book'])
def test_tampered_saved_computation_cannot_be_approved(clients,target):
    e,r,pid,b=setup(clients);base='/api/v1/projects/'+pid
    c=e.post(base+'/components',json=b).json()
    if target=='book':approve(r,pid,'components',c);e.post(base+'/takeoffs/calculate')
    with store.transaction() as db:
        p=json.loads(db.execute('SELECT data FROM projects WHERE id=?',(pid,)).fetchone()['data'])
        if target=='component':p['components'][0]['preview']['route_segments'][0]['length']='999'
        else:p['takeoffs'][0]['route_segments'][0]['length']='999'
        store.save(db,p)
    collection='components' if target=='component' else 'takeoffs'
    approve(r,pid,collection,get(e,pid)[collection][0],expected=409)


def test_new_route_preserves_immutable_quote_snapshot(clients):
    e,r,pid,_,q=priced_project(clients);before=copy.deepcopy(q)
    response=e.post('/api/v1/projects/'+pid+'/components',json=payload(get(e,pid)['evidence'][0]['id']))
    assert response.status_code==200,response.text
    p=get(e,pid);assert p['quotes'][0]==before and q['id'] in p['stale_quotes']
