import copy
import io
import pytest
from pypdf import PdfReader
from openpyxl import load_workbook
from backend import store,main as core,ocr
from tests.test_ocr_drawing_candidates import fixture,mutate,result
from tests.test_ocr_line_locations import selected
from tests.test_drawing_candidates import read,adopt
from tests.test_workflow import get,approve,priced_project
from tests.test_ocr import pdf,submit


def body_for(e,**extra):
    return {'expected_line':copy.deepcopy(e['ocr']['lines'][0]),'evidence_hash':store.digest(e),
            'text':'图号：J-02','reason':'仅合成测试：对照扫描原图核对识别错误','original_page_checked':True,**extra}


def save(c,pid,e,body,expected=200):
    response=c.post(f'/api/v1/projects/{pid}/evidence/{e["id"]}/ocr-lines/1/transcriptions',json=body)
    assert response.status_code==expected,response.text
    return response.json()


def test_append_only_idempotent_manual_evidence_never_changes_raw_ocr_or_quantities(clients,monkeypatch):
    c,r,pid,d,_=fixture(clients,monkeypatch);e,_=selected(c,pid);before=get(c,pid)
    body=body_for(e);first=save(c,pid,e,body);second=save(c,pid,e,body)
    assert not first['duplicate'] and second['duplicate'] and first['evidence']==second['evidence']
    record=first['evidence'];trace=record['ocr_transcription']
    assert record['method']=='manual' and record['state']=='draft' and record['bbox']==e['ocr']['lines'][0]['bbox']
    assert trace['original_line']['text']=='图号：J-01' and trace['text']=='图号：J-02' and trace['created_by']=='estimator'
    assert trace['source_location']['evidence_hash']==store.digest(e)
    after=get(c,pid);assert next(v for v in after['evidence'] if v['id']==e['id'])==e
    assert len(after['evidence'])==len(before['evidence'])+1
    for key in ['facts','components','takeoffs','boq','quotes','reviews']:assert before[key]==after[key]
    candidate=next(v for v in read(c,pid)['candidates'] if v['evidence_id']==record['id'])
    assert candidate['transcription']['original_line']['text']=='图号：J-01'
    sheet=adopt(c,pid,candidate);assert sheet['state']=='draft'
    approve(c,pid,'drawing_sheets',sheet,expected=403);approve(r,pid,'drawing_sheets',sheet)
    assert [a['action'] for a in c.get(f'/api/v1/projects/{pid}/audit').json()].count('OCR_LINE_TRANSCRIBED')==1


@pytest.mark.parametrize('extra',[
    {'original_page_checked':False},{'original_page_checked':'true'},{'original_page_checked':1},
    {'text':'   '},{'reason':'  '},{'text':'x'*2001},{'reason':'x'*2001},{'text':42},
    {'state':'approved'},{'document_id':'foreign'}])
def test_invalid_transcription_is_atomic(clients,monkeypatch,extra):
    c,r,pid,d,_=fixture(clients,monkeypatch);e,_=selected(c,pid);before=get(c,pid)
    save(c,pid,e,body_for(e,**extra),422)
    assert get(c,pid)==before


def test_roles_missing_hash_and_stale_version(clients,monkeypatch):
    c,r,pid,d,_=fixture(clients,monkeypatch);e,_=selected(c,pid);body=body_for(e)
    save(r,pid,e,body,403);save(clients['other'],pid,e,body,404)
    save(c,pid,e,{k:v for k,v in body.items() if k!='evidence_hash'},422)
    save(c,pid,e,{**body,'evidence_hash':'0'*64},409)


@pytest.mark.parametrize('change',['source_task','source_text','original_bytes','record_text','record_box','cycle','source_hash','empty_trace'])
def test_changed_source_cannot_be_reused_for_facts_or_candidate_adoption(clients,monkeypatch,change):
    c,r,pid,d,_=fixture(clients,monkeypatch);e,_=selected(c,pid)
    record=save(c,pid,e,body_for(e))['evidence']
    def action(p,ev,j):
        saved=next(v for v in p['evidence'] if v['id']==record['id'])
        if change=='source_task':j['state']='failed'
        elif change=='source_text':ev['text']+=' different'
        elif change=='record_text':saved['text']='changed'
        elif change=='record_box':saved['bbox']=[.2,.2,.2,.2]
        elif change=='cycle':saved['ocr_transcription']['source_location']['evidence_id']=saved['id']
        elif change=='empty_trace':saved['ocr_transcription']={}
        elif change=='source_hash':saved['ocr_transcription']['source_location']['evidence_hash']='invalid'
    if change=='original_bytes':(store.ROOT/'originals'/d['hash']).write_bytes(b'bad original')
    else:mutate(pid,action)
    response=c.post(f'/api/v1/projects/{pid}/facts',json={'key':'lot','label':'标段','value':'合成标段','evidence_id':record['id']})
    assert response.status_code==409,response.text
    assert not get(c,pid)['facts']
    if change!='original_bytes':assert not any(v['evidence_id']==record['id'] for v in read(c,pid)['candidates'])


def test_different_human_transcriptions_are_retained_without_choosing_winner(clients,monkeypatch):
    c,r,pid,d,_=fixture(clients,monkeypatch);e,_=selected(c,pid)
    a=save(c,pid,e,body_for(e))['evidence'];b=save(c,pid,e,body_for(e,text='图号：J-03'))['evidence']
    assert a['id']!=b['id']
    rows=read(c,pid)['candidates'];assert {a['id'],b['id']}<=set(v['evidence_id'] for v in rows)
    assert get(c,pid)['drawing_sheets']==[]


def test_transcription_export_and_quote_snapshot_remain_frozen(clients,monkeypatch):
    c,r,pid,_,old=priced_project(clients);base=f'/api/v1/projects/{pid}'
    monkeypatch.setattr(ocr,'capability',lambda:{'available':True})
    monkeypatch.setattr(ocr,'recognize',lambda *a:result('图号：J-01\n图名：合成扫描图'))
    d=c.post(base+'/documents',data={'specialty':'建筑'},files={'file':('scan.pdf',pdf(),'application/pdf')}).json()['document']
    submit(c,base+'/documents/'+d['id']+'/ocr',d);e,_=selected(c,pid)
    record=save(c,pid,e,body_for(e))['evidence']
    for coll in ['components','takeoffs','boq']:
        for row in get(c,pid)[coll]:approve(r,pid,coll,row)
    response=c.post(base+'/quotes/calculate');assert response.status_code==200,response.text
    q=response.json();frozen=copy.deepcopy(q)
    save(c,pid,e,body_for(e,text='图号：J-03',reason='后续人工抄录（合成）'))
    data=c.get(base+'/exports',params={'format':'json','quote_id':q['id']}).json()
    assert data['quote']==frozen and '后续人工抄录' not in str(data['tables'])
    table=next(t for t in data['tables'] if t[0]=='扫描原行人工抄录')
    assert '图号：J-01' in str(table) and '图号：J-02' in str(table)
    x=c.get(base+'/exports',params={'format':'xlsx','quote_id':q['id']});wb=load_workbook(io.BytesIO(x.content))
    assert '扫描原行人工抄录' in wb.sheetnames
    raw=c.get(base+'/exports',params={'format':'pdf','quote_id':q['id']}).content
    assert '扫描原行人工抄录' in ''.join(page.extract_text() for page in PdfReader(io.BytesIO(raw)).pages)
    assert not any(e.get('ocr_transcription') for e in old['snapshot']['evidence'])
