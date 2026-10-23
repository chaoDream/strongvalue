"""Synthetic per-stock kerf capacity and end-to-end traceability."""
import copy,io,random
from decimal import Decimal,localcontext
from collections import Counter
import pytest
from openpyxl import load_workbook
from pypdf import PdfReader
from backend import store,stock_kerf,mixed_stock_cutting as mixed
from backend.engine import DomainError
from backend.railing_material import calculate
from tests.test_stock_processing import body as processing_body,setup as processing_setup,group,geometry,member,equal_body,fixed_body,fixed_added,equal_added,fixed_adjusted,equal_adjusted,combined,local_body,path_body
from tests.test_multi_stock_cutting import body as multi_body
from tests.test_limited_stock_cutting import limits
from tests.test_mixed_stock_cutting import config as mixed_config
from tests.test_route_takeoff import evaluate
from tests.test_polygon_takeoff import post,mutate
from tests.test_workflow import get,approve,upload,create,priced_project

TITLES={'逐规格每件锯缝','逐规格锯缝范围依据'}


def kerfs(eid='test'):
 return dict(alternatives=[dict(kerf_per_piece='13',evidence_id=eid),dict(kerf_per_piece='20',evidence_id=eid)],scope_confirmed=True,measurement_basis='合成S1每件3mm，S2每件13mm，S3每件20mm，含最后一件，不重复计净下料或端切；本说明优先于组及端切说明的通用锯缝，非真实工程',evidence_id=eid)


def body(s,eid='test'):
 b=processing_body(s,eid);b['name']='合成逐规格锯缝（非真实工程）';c=b['mixed_cutting'][0]['stock_choices'];c.update(kerfs=kerfs(eid),measurement_basis='合成同规格方管A，三种有据定尺可按需供应，加工及锯缝按逐规格原件，不增加净量和报价');return b


def setup(clients,pid=None):
 e,r,pid,s,b,eid=processing_setup(clients,pid)
 upload(e,pid,'合成逐规格锯缝软件验收，非真实工程。净长短各4件3500/2400mm，3000g/m，净70.8kg。S1为6000mm端切10/10mm、每件3mm，留用50mm。S2为3523mm端切4/6mm、每件13mm，留用10mm，长件恰好占满可用3513mm。S3为2423mm端切1/2mm、每件20mm，留用21mm，短件恰好占满2420mm。含最后一件各计一次，未计净下料或端切，优先于原通用锯缝。选S2/S3各4根，购入23.784m71.352kg，净23.6m锯缝.132m端切.052m余料0。S2改13.000000001mm后长件超限，改选四根6m24m72kg，净量不变须重新复核。相同参数可与本次有限供应上限组合，仅方案不预留或入库。','stock-kerf-basis.txt')
 for c in get(e,pid)['components']:approve(r,pid,'components',c)
 s=get(e,pid)['components'][-1];eid=get(e,pid)['evidence'][-1]['id'];return e,r,pid,s,body(s,eid),eid


def test_actual_kerfs_positions_asts_and_net_unchanged():
 s=geometry();b=body(s);before=copy.deepcopy((s,b));v=calculate(b,s);g=v['mixed_stock_cutting']['groups'][0];old=calculate(processing_body(s),s)
 assert (s,b)==before and v['quantity']==old['quantity']=='70.800' and v['material_items']==old['material_items'] and v['formula_ast']==old['formula_ast'] and v['rule_version']=='railing-material-mass-1.18'
 assert g['stock_selection']['kerf_independent'] and g['stock_selection']['processing_independent'] and g['proof']['kerf_model']=='per_stock_per_piece' and g['proof']['status']=='length_optimal_search'
 assert [o['stock_count'] for o in g['stock_selection']['options']]==[0,4,4] and g['total']['kerf_length']=='0.132' and g['total']['trim_length']=='0.052' and g['total']['retained_length']==g['total']['scrap_length']==g['balance_difference']=='0'
 for o in g['stock_selection']['options']:assert evaluate(o['kerf']['formula_ast'],v['dimensions'])==Decimal(o['kerf']['kerf_per_piece'])
 for bar in g['bars']:
  p=bar['pieces'][0];assert (p['from_s'],p['to_s'],p['kerf_to_s'])==(('0.004','3.504','3.517') if bar['stock_choice_index']==2 else ('0.001','2.401','2.421'))
  for k,ast in bar['formula_asts'].items():assert evaluate(ast,v['dimensions'])==Decimal(bar['values'][k])
 for k,ast in g['total_formula_asts'].items():assert evaluate(ast,v['dimensions'])==Decimal(g['total'][k])


@pytest.mark.parametrize('k,counts',[('12.999999999',[0,4,4]),('13',[0,4,4]),('13.000000001',[4,0,0])])
def test_exact_capacity_micro_boundary(k,counts):
 s=geometry();b=body(s);b['mixed_cutting'][0]['stock_choices']['kerfs']['alternatives'][0]['kerf_per_piece']=k;g=group(s,b);assert [o['stock_count'] for o in g['stock_selection']['options']]==counts


def test_zero_kerf_and_uniform_end_trims_without_processing():
 s=geometry();b=body(s);c=b['mixed_cutting'][0]['stock_choices'];c.pop('processing');c['kerfs']['alternatives'][0]['kerf_per_piece']='0';c['kerfs']['alternatives'][1]['kerf_per_piece']='0';g=group(s,b)
 assert 'processing_independent' not in g['stock_selection'] and g['total']['kerf_length']=='0' and g['total']['purchased_mass']=='71.352'


def test_finite_combination_capacity_and_zero_supply():
 s=geometry();b=body(s);c=b['mixed_cutting'][0]['stock_choices'];c.update(supply_confirmed=False,supply_limits=limits());g=group(s,b)
 assert g['stock_selection']['supply_limited'] and g['total']['purchased_mass']=='71.676' and g['total']['kerf_length']=='0.078' and g['total']['retained_length']=='0.148'
 c['kerfs']['alternatives'][0]['kerf_per_piece']='13.000000001'
 with pytest.raises(DomainError,match='已由容量约束或穷尽搜索证明'):group(s,b)
 c['supply_limits']['quantities'][0]['quantity']='4';c['supply_limits']['quantities'][1]['quantity']='0';c['supply_limits']['quantities'][2]['quantity']='0';assert [o['stock_count'] for o in group(s,b)['stock_selection']['options']]==[4,0,0]


@pytest.mark.parametrize('fault',['null','list','empty','extra','missing','false','truthy','no-basis','no-proof','bool-proof','rows-null','rows-dict','rows-few','rows-many','row-null','row-list','row-extra','row-missing','negative','nan','infinite','bool','precision','blank-proof','bad-proof','no-piece','no-long'])
def test_strict_schema(fault):
 s=geometry();b=body(s);c=b['mixed_cutting'][0]['stock_choices'];p=c['kerfs'];r=p['alternatives'][0]
 if fault=='null':c['kerfs']=None
 elif fault=='list':c['kerfs']=[]
 elif fault=='empty':c['kerfs']={}
 elif fault=='extra':p['quantity']='1'
 elif fault=='missing':p.pop('scope_confirmed')
 elif fault=='false':p['scope_confirmed']=False
 elif fault=='truthy':p['scope_confirmed']=1
 elif fault=='no-basis':p['measurement_basis']=''
 elif fault=='no-proof':p['evidence_id']=''
 elif fault=='bool-proof':p['evidence_id']=True
 elif fault=='rows-null':p['alternatives']=None
 elif fault=='rows-dict':p['alternatives']={}
 elif fault=='rows-few':p['alternatives'].pop()
 elif fault=='rows-many':p['alternatives'].append(r)
 elif fault=='row-null':p['alternatives'][0]=None
 elif fault=='row-list':p['alternatives'][0]=[]
 elif fault=='row-extra':r['trim_start']='1'
 elif fault=='row-missing':r.pop('kerf_per_piece')
 elif fault=='negative':r['kerf_per_piece']='-1'
 elif fault=='nan':r['kerf_per_piece']='NaN'
 elif fault=='infinite':r['kerf_per_piece']='Infinity'
 elif fault=='bool':r['kerf_per_piece']=True
 elif fault=='precision':r['kerf_per_piece']='.1234567890123456789'
 elif fault=='blank-proof':r['evidence_id']=''
 elif fault=='bad-proof':r['evidence_id']=False
 elif fault=='no-piece':r['kerf_per_piece']='2000'
 else:b['mixed_cutting'][0]['kerf_per_piece']='3000';r['kerf_per_piece']='14'
 with pytest.raises(DomainError):calculate(b,s)


@pytest.mark.parametrize('unit,factor',[('m',Decimal('.001')),('mm',Decimal('1')),('cm',Decimal('.1')),('dm',Decimal('.01')),('km',Decimal('.000001'))])
def test_units_and_low_context(unit,factor):
 s=geometry();b=body(s);c=b['mixed_cutting'][0];c['length_unit']=unit
 for k in ('stock_length','kerf_per_piece','trim_start','trim_end','reusable_min_length'):c[k]=str(Decimal(c[k])*factor)
 for o in c['stock_choices']['alternatives']:o['stock_length']=str(Decimal(o['stock_length'])*factor)
 for o in c['stock_choices']['processing']['alternatives']:
  for k in ('trim_start','trim_end','reusable_min_length'):o[k]=str(Decimal(o[k])*factor)
 for o in c['stock_choices']['kerfs']['alternatives']:o['kerf_per_piece']=str(Decimal(o['kerf_per_piece'])*factor)
 with localcontext() as ctx:ctx.prec=6;v=calculate(b,s)
 assert v==calculate(b,s) and v['mixed_stock_cutting']['total']['kerf_length']=='0.132'


@pytest.mark.parametrize('make',[equal_body,fixed_body,fixed_added,equal_added,fixed_adjusted,equal_adjusted,combined,lambda s:combined(s,equal=True),local_body,path_body])
def test_post_modes_and_spacing_unchanged(make):
 s=geometry();b=make(s);b['material_items'][0]['name']='合成立柱';b['material_items'].append({**member('合成短料','cut_pieces'),'quantity':'2','cut_length':'2400'});old=calculate(copy.deepcopy(b),s);g=mixed_config(members=['合成立柱','合成短料']);g['stock_choices']=body(s)['mixed_cutting'][0]['stock_choices'];b['mixed_cutting']=[g];v=calculate(b,s);assert v['quantity']==old['quantity'] and v['material_items']==old['material_items']


def test_independent_oracle_variable_kerf_and_trim():
 rng=random.Random(8501)
 def oracle(cuts,stocks,trims,kerfs,counts):
  best=None;bins=[];used=[0]*len(stocks)
  def walk(i):
   nonlocal best
   if i==len(cuts):
    v=(sum(stocks[j] for j,_ in bins),len(bins));best=v if best is None else min(best,v);return
   for k,(j,load) in enumerate(bins):
    if load+cuts[i]+kerfs[j]<=stocks[j]-trims[j]:bins[k]=(j,load+cuts[i]+kerfs[j]);walk(i+1);bins[k]=(j,load)
   for j,stock in enumerate(stocks):
    if used[j]<counts[j] and cuts[i]+kerfs[j]<=stock-trims[j]:used[j]+=1;bins.append((j,cuts[i]+kerfs[j]));walk(i+1);bins.pop();used[j]-=1
  walk(0);return best
 for _ in range(100):
  stocks=sorted(rng.sample(range(5,16),3));trims=[rng.randint(0,2) for _ in stocks];kerfs=[rng.randint(0,3) for _ in stocks];cuts=sorted([rng.randint(1,10) for _ in range(rng.randint(2,6))],reverse=True)
  for counts in (None,[rng.randint(0,3) for _ in stocks]):
   expected=oracle(cuts,stocks,trims,kerfs,counts or [len(cuts)]*3)
   if expected is None:
    with pytest.raises(DomainError,match='已由容量约束或穷尽搜索证明'):stock_kerf.pack(cuts,stocks,trims,kerfs,counts)
   else:
    plan,p=stock_kerf.pack(cuts,stocks,trims,kerfs,counts);assert (p['purchased_length'],p['stock_count'])==expected and p['status']!='length_feasible_unproven'
    assert all(sum(cuts[i]+kerfs[j] for i in ps)<=stocks[j]-trims[j] for j,ps in plan) and (counts is None or all(Counter(j for j,_ in plan)[j]<=counts[j] for j in range(3)))


def test_max_200_pieces_five_candidates():
 s=geometry();b=body(s)
 for m in b['material_items']:m['quantity']='50'
 c=b['mixed_cutting'][0]['stock_choices'];c['alternatives']+=[dict(stock_length=x,evidence_id='test') for x in ('5926','8000')];c['processing']['alternatives']*=2;c['kerfs']['alternatives']*=2;assert group(s,b)['required_piece_count']==200


@pytest.mark.parametrize('fmt',['json','xlsx','pdf'])
def test_workflow_export_net_boq_and_old_quote(clients,fmt):
 e,r,pid,_,quote=priced_project(clients);e,r,pid,s,b,eid=setup(clients,pid);before=get(e,pid);post(e,pid,b,'/components/preview');assert get(e,pid)==before;c=post(e,pid,b).json();approve(r,pid,'components',c)
 for t in post(e,pid,{},'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
 assert any(t['quantity']=='70.800' for t in post(e,pid,{},'/boq/generate').json());url=f'/api/v1/projects/{pid}/exports';res=e.get(url,params=dict(format=fmt,unpriced=True));assert res.status_code==200,res.text
 if fmt=='json':assert TITLES<={t[0] for t in res.json()['tables']}
 elif fmt=='xlsx':assert TITLES<=set(load_workbook(io.BytesIO(res.content)).sheetnames)
 else:assert all(t in ''.join(p.extract_text() for p in PdfReader(io.BytesIO(res.content)).pages) for t in TITLES)
 assert not any(t[0] in TITLES for t in e.get(url,params=dict(format='json',quote_id=quote['id'])).json()['tables'])


@pytest.mark.parametrize('fault',['raw','proof','scope','normalized','ast','position','loss','book'])
def test_tamper_rebuild(clients,fault):
 e,r,pid,s,b,eid=setup(clients);c=post(e,pid,b).json()
 if fault=='book':approve(r,pid,'components',c);post(e,pid,{},'/takeoffs/calculate')
 def forge(p):
  c=p['components'][-1];g=c['preview']['mixed_stock_cutting']['groups'][0]
  if fault=='raw':c['mixed_cutting'][0]['stock_choices']['kerfs']['alternatives'][0]['kerf_per_piece']='12'
  elif fault=='proof':c['mixed_cutting'][0]['stock_choices']['kerfs']['alternatives'][0]['evidence_id']='forged'
  elif fault=='scope':c['mixed_cutting'][0]['stock_choices']['kerfs']['scope_confirmed']=False
  elif fault=='normalized':g['stock_selection']['options'][1]['kerf']['kerf_per_piece']='0.012'
  elif fault=='ast':g['stock_selection']['options'][1]['kerf']['formula_ast']={'constant':'0.012'}
  elif fault=='position':g['bars'][0]['pieces'][0]['kerf_to_s']='3.516'
  elif fault=='loss':g['total']['kerf_length']='0.13'
  else:p['takeoffs'][-1]['mixed_stock_cutting']['groups'][0]['proof']['kerf_model']='forged'
 mutate(pid,forge);coll='takeoffs' if fault=='book' else 'components';approve(r,pid,coll,get(e,pid)[coll][-1],expected=404 if fault=='proof' else 422 if fault=='scope' else 409)


def test_same_net_revision_history_removal_concurrency_source(clients):
 e,r,pid,s,b,eid=setup(clients);c=post(e,pid,b).json();approve(r,pid,'components',c)
 for t in post(e,pid,{},'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
 post(e,pid,{},'/boq/generate');before=get(e,pid);old=before['components'][-1];new=copy.deepcopy(b);new['mixed_cutting'][0]['stock_choices']['kerfs']['alternatives'][0]['kerf_per_piece']='13.000000001';new['object_hash']=store.digest(old);url=f'/api/v1/projects/{pid}/components/'+old['id'];res=e.put(url,json=new);assert res.status_code==200,res.text
 p=get(e,pid);assert p['component_history'][-1]==old and p['boq']==before['boq'] and all(t['state']=='stale' for t in p['takeoffs']) and res.json()['preview']['quantity']=='70.800' and res.json()['preview']['mixed_stock_cutting']['total']['purchased_mass']=='72';assert e.put(url,json=new).status_code==409;approve(r,pid,'components',res.json())
 current=get(e,pid)['components'][-1];new['object_hash']=store.digest(current);new['mixed_cutting'][0]['stock_choices'].pop('kerfs');res=e.put(url,json=new);assert res.status_code==200 and res.json()['preview']['rule_version']=='railing-material-mass-1.17';approve(r,pid,'components',res.json())
 from tests.test_railing_takeoff import interval
 s['object_hash']=store.digest(s);s['edge_decisions'][0].update(mode='ranges',ranges=[interval(0,10,eid)]);res=e.put(f'/api/v1/projects/{pid}/components/'+s['id'],json=s);assert res.status_code==200;approve(r,pid,'components',res.json());assert get(e,pid)['components'][-1]['state']=='stale'


@pytest.mark.parametrize('gate',['preview','save','review','calculate','book_review'])
def test_original_bytes_gates(clients,gate):
 e,r,pid,s,b,eid=setup(clients)
 if gate in ('review','calculate','book_review'):c=post(e,pid,b).json()
 if gate in ('calculate','book_review'):approve(r,pid,'components',c)
 if gate=='book_review':t=post(e,pid,{},'/takeoffs/calculate').json()[-1]
 d=get(e,pid)['documents'][-1];(store.ROOT/'originals'/d['hash']).write_bytes(b'corrupt')
 if gate=='preview':post(e,pid,b,'/components/preview',409)
 elif gate=='save':post(e,pid,b,expected=409)
 elif gate=='review':approve(r,pid,'components',c,expected=409)
 elif gate=='calculate':post(e,pid,{},'/takeoffs/calculate',409)
 else:approve(r,pid,'takeoffs',t,expected=409)


def test_roles_tenant_foreign_parameter_and_scope_evidence(clients):
 e,r,pid,s,b,eid=setup(clients);post(r,pid,b,expected=403);assert clients['other'].get(f'/api/v1/projects/{pid}').status_code==404
 other=create(e)['id'];upload(e,other,'异项目');fid=get(e,other)['evidence'][-1]['id'];original=copy.deepcopy(b);b['mixed_cutting'][0]['stock_choices']['kerfs']['alternatives'][0]['evidence_id']=fid;post(e,pid,b,expected=404)
 b=original;b['mixed_cutting'][0]['stock_choices']['kerfs']['evidence_id']=fid;post(e,pid,b,expected=404)


def test_budget_unproven_and_long_names_pdf(clients,monkeypatch):
 monkeypatch.setattr(mixed,'SEARCH_LIMIT',0);e,r,pid,s,b,eid=setup(clients);b['name']='合成'+('长名称验收'*39)+'末';b['mixed_cutting'][0]['name']='合成'+('长混排组名'*39)+'末';c=post(e,pid,b).json();assert c['preview']['mixed_stock_cutting']['groups'][0]['proof']['status']=='length_feasible_unproven';approve(r,pid,'components',c)
 for t in post(e,pid,{},'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
 res=e.get(f'/api/v1/projects/{pid}/exports',params=dict(format='pdf',unpriced=True));assert res.status_code==200,res.text


@pytest.mark.parametrize('weight,unit',[('3','kg/m'),('3000','g/m'),('.003','kg/mm'),('3','g/mm')])
def test_weight_units(weight,unit):
 s=geometry();b=body(s)
 for m in b['material_items']:m.update(unit_weight=weight,weight_unit=unit)
 assert group(s,b)['total']['kerf_mass']=='0.396'


def test_search_finds_solution_after_all_greedy_seeds_fail_and_unknown_budget():
 args=([10,8,5,4,4,4,2],[6,8,16],[2,1,0],[3,2,0],[0,1,2])
 plan,p=stock_kerf.pack(*args);assert p['initial_purchased_length'] is None and p['status']=='length_optimal_search' and p['purchased_length']==40 and p['stock_count']==3
 with pytest.raises(DomainError,match='不能据此判断供应不足'):stock_kerf.pack(*args,limit=0)
 with pytest.raises(DomainError,match='已由容量约束或穷尽搜索证明'):stock_kerf.pack(*args[:-1],[0,1,1])


def test_same_length_secondary_count_and_bound():
 plan,p=stock_kerf.pack([3,3],[4,8],[0,0],[1,0]);assert p['purchased_length']==8 and p['stock_count']==1
 plan,p=stock_kerf.pack([3,3],[4,6],[0,0],[1,0]);assert p['status']=='length_optimal_bound' and p['purchased_length']==6 and p['stock_count']==1


def test_each_piece_on_same_bar_uses_chosen_kerf_including_final_piece():
 s=geometry();b=body(s);c=b['mixed_cutting'][0]['stock_choices'];c['alternatives'][0]['stock_length']='5918';c['kerfs']['alternatives'][0]['kerf_per_piece']='4';g=group(s,b);assert [o['stock_count'] for o in g['stock_selection']['options']]==[0,4,0]
 for bar in g['bars']:
  a,z=bar['pieces'];assert (a['from_s'],a['to_s'],a['kerf_to_s'])==('0.004','3.504','3.508') and (z['from_s'],z['to_s'],z['kerf_to_s'])==('3.508','5.908','5.912') and bar['offcut_length']=='0'
 assert g['total']['kerf_length']=='0.032' and g['total']['purchased_length']=='23.672'
