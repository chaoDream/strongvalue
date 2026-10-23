import copy
import io
import pytest
from openpyxl import load_workbook
from pypdf import PdfReader
from backend import main as core,store,ocr
from tests.test_ocr import setup,output,submit,pdf
from tests.test_drawing_candidates import read,adoption,adopt
from tests.test_workflow import get,approve,priced_project


def result(text):
    return {**output(),'lines':[{'text':line,'bbox':[.1,.1+i*.05,.7,.04],'confidence':.3 if i==0 else .9} for i,line in enumerate(text.splitlines())]}


def fixture(clients,monkeypatch,text='图号：J-01\n图名：合成扫描平面\n版次：A\n视图类型：平面图'):
    e,pid,d,url=setup(clients,monkeypatch)
    monkeypatch.setattr(ocr,'recognize',lambda *a:result(text))
    assert submit(e,url,d).status_code==200
    return e,clients['reviewer'],pid,d,read(e,pid)['candidates'][0]


def mutate(pid,action):
    with store.transaction() as conn:
        p=core.project(conn,pid,{'tenant':'local','role':'estimator','username':'estimator'})
        ev=next(e for e in p['evidence'] if e['method']=='local_ocr');j=next(j for j in p['jobs'] if j['id']==ev['ocr']['job_id'])
        action(p,ev,j);store.save(conn,p)


def test_source_lines_low_confidence_and_read_only_candidates(clients,monkeypatch):
    e,r,pid,d,c=fixture(clients,monkeypatch)
    before=get(e,pid);assert read(e,pid)==read(r,pid) and get(e,pid)==before
    assert c['source_method']=='local_ocr' and c['ocr']['field_lines']['drawing_no'][0]['confidence']==.3
    assert c['ocr']['field_lines']['title'][0]['line']==2 and c['suggestions']['drawing_no']=='J-01'
    assert c['warnings'] and not before['drawing_sheets']


@pytest.mark.parametrize('checked',[None,False,'true',1,[],{}])
def test_adoption_requires_strict_original_page_confirmation(clients,monkeypatch,checked):
    e,r,pid,d,c=fixture(clients,monkeypatch)
    adopt(e,pid,c,422,ocr_source_checked=checked)
    assert not get(e,pid)['drawing_sheets']


def test_independent_review_records_source_and_corrections(clients,monkeypatch):
    e,r,pid,d,c=fixture(clients,monkeypatch,text='图号：J-01\n图名：合成扫描平面')
    sheet=adopt(e,pid,c,ocr_source_checked=True,edition='A',view_type='plan',title='原页核对后的合成名称')
    trace=sheet['candidate_trace'];assert sheet['identity_origin']=='ocr_labels' and sheet['state']=='draft'
    assert trace['ocr_verification']['checked_by']=='estimator' and trace['ocr_verification']['fields']==['drawing_no','title','edition','view_type']
    assert set(trace['corrections'])=={'title','edition','view_type'}
    approve(e,pid,'drawing_sheets',sheet,expected=403)
    approve(r,pid,'drawing_sheets',sheet)
    assert not get(e,pid)['drawing_register']['issues']
    adopt(e,pid,c,409,ocr_source_checked=True)


def test_different_ocr_values_not_resolved_by_confidence(clients,monkeypatch):
    e,r,pid,d,c=fixture(clients,monkeypatch,text='图号：J-01\n图号：J-02\n图名：合成扫描平面')
    assert 'drawing_no' not in c['suggestions'] and c['conflicting_fields']==['drawing_no']
    assert len(c['ocr']['field_lines']['drawing_no'])==2
    sheet=adopt(e,pid,c,ocr_source_checked=True,drawing_no='J-02')
    assert sheet['candidate_trace']['corrections']['drawing_no']['adopted']=='J-02'


@pytest.mark.parametrize('change', ['job_state','job_evidence','job_page','engine','revision','coordinate','text','lines','original'])
def test_changed_ocr_provenance_cannot_be_adopted(clients,monkeypatch,change):
    e,r,pid,d,c=fixture(clients,monkeypatch)
    def action(p,ev,j):
        if change=='job_state':j['state']='failed'
        elif change=='job_evidence':j['evidence_id']='foreign'
        elif change=='job_page':j['page']=2
        elif change=='engine':ev['ocr']['engine']='other-engine'
        elif change=='revision':ev['ocr']['revision']='3'
        elif change=='coordinate':ev['ocr']['coordinate_space']='other-space'
        elif change=='text':ev['text']+='\n图号：J-02'
        elif change=='lines':ev['ocr']['lines'][0]['bbox']=[0,0,2,.1]
    if change=='original':(store.ROOT/'originals'/d['hash']).write_bytes(b'changed')
    else:mutate(pid,action)
    adopt(e,pid,c,409,ocr_source_checked=True)
    assert not get(e,pid)['drawing_sheets']


def test_approved_sheet_invalidated_if_ocr_task_binding_changes(clients,monkeypatch):
    e,r,pid,d,c=fixture(clients,monkeypatch);sheet=adopt(e,pid,c,ocr_source_checked=True)
    approve(r,pid,'drawing_sheets',sheet)
    mutate(pid,lambda p,ev,j:j.update(evidence_id='other'))
    assert '扫描识别' in get(e,pid)['drawing_register']['source_issues']['drawing_sheets'][sheet['id']]
    approve(r,pid,'drawing_sheets',sheet,expected=409)


def test_roles_and_cross_tenant_cannot_adopt_machine_candidate(clients,monkeypatch):
    e,r,pid,d,c=fixture(clients,monkeypatch)
    adopt(r,pid,c,403,ocr_source_checked=True)
    adopt(clients['other'],pid,c,404,ocr_source_checked=True)


def test_ocr_trace_exported_in_all_formats_and_frozen(clients,monkeypatch):
    e,r,pid,_,old=priced_project(clients);base=f'/api/v1/projects/{pid}'
    monkeypatch.setattr(ocr,'capability',lambda:{'available':True})
    monkeypatch.setattr(ocr,'recognize',lambda *a:result('图号：J-01\n图名：扫描合成平面\n版次：A\n视图类型：平面图'))
    doc=e.post(base+'/documents',data={'specialty':'建筑'},files={'file':('ocr.pdf',pdf(),'application/pdf')}).json()['document']
    assert submit(e,base+'/documents/'+doc['id']+'/ocr',doc).status_code==200
    c=read(e,pid)['candidates'][0];sheet=adopt(e,pid,c,ocr_source_checked=True);approve(r,pid,'drawing_sheets',sheet)
    for coll in ['components','takeoffs','boq']:
        for row in get(e,pid)[coll]:approve(r,pid,coll,row)
    q=e.post(base+'/quotes/calculate').json();frozen=copy.deepcopy(q)
    mutate(pid,lambda p,ev,j:ev['ocr']['lines'][0].update(confidence=.1))
    data=e.get(base+'/exports',params={'format':'json','quote_id':q['id']}).json()
    assert data['quote']==frozen
    table=next(t for t in data['tables'] if t[0]=='扫描标题栏字段核对')
    assert '置信度 0.3' in str(table) and 'estimator' in str(table)
    x=e.get(base+'/exports',params={'format':'xlsx','quote_id':q['id']});wb=load_workbook(io.BytesIO(x.content))
    assert '扫描标题栏字段核对' in wb.sheetnames
    raw=e.get(base+'/exports',params={'format':'pdf','quote_id':q['id']}).content
    text=''.join(page.extract_text() for page in PdfReader(io.BytesIO(raw)).pages)
    assert '扫描标题栏字段核对' in text and '置信度' in text and 'estimator' in text
    assert not old['snapshot']['drawing_sheets']
