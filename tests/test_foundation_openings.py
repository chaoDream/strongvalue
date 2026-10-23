"""Vertical foundation hole containment, disjointness, net balances and provenance."""
import copy,io,random
from decimal import Decimal,localcontext,ROUND_HALF_UP
import pytest
from openpyxl import load_workbook
from pypdf import PdfReader
from backend import store
from backend.compound_takeoff import PI
from backend.engine import DomainError
from backend.post_foundation import calculate
from tests.test_post_foundation import body as rect_body,source,equal_body,fixed_body,fixed_added,equal_added,fixed_adjusted,equal_adjusted,combined
from tests.test_foundation_steps import body as steps_body
from tests.test_foundation_round import body as round_body,setup as round_setup
from tests.test_polygon_takeoff import post,mutate
from tests.test_route_takeoff import evaluate
from tests.test_workflow import get,approve,upload,create,priced_project

SHAPES=('rectangular','stepped_rectangular','circular','stepped_circular')
QUANTITIES=dict(zip(SHAPES,('2.607','2.343','3.699','2.317')))
TITLES={'基础贯通孔毛扣净','基础贯通孔原始依据','基础贯通孔校验','基础贯通孔范围与常量'}


def hole(name='合成H1',x='-50',y='0',diameter='40',eid='test'):
 return dict(name=name,parameters=dict(x=x,y=y,diameter=diameter),measurement_basis='合成竖直圆孔自基础底贯穿全部层，实际孔径，独立扣一次，非真实工程',evidence={k:eid for k in ('x','y','diameter','measurement_basis')})


def config(eid='test'):
 return dict(unit='mm',coordinate_basis='合成以各层共同基础中心为局部XY原点，X沿基础长度正向、Y沿宽度正向，孔自基础底贯穿全部层，不从立柱或锚栓推断',coordinate_evidence_id=eid,scope_confirmed=True,holes=[hole(eid=eid),hole('合成H2','50',eid=eid)])


def body(s,eid='test',shape='stepped_circular'):
 b=rect_body(s,eid) if shape=='rectangular' else steps_body(s,eid) if shape=='stepped_rectangular' else round_body(s,eid,shape)
 b.update(name='合成基础贯通圆孔扣减（非真实工程）',material='合成同类C30基础，各层同一材料，仅软件验收',measurement_basis='合成22个独立基础，每基础两孔，实际孔径及共同局部XY逐项据原件核定，竖直贯通全部层，孔在每层内且互不重叠，孔数层数不乘基础个数，来源重复一次，无未列扣减、基础交叠或既有清单重复')
 d=b['foundation_items'][0];d['measurement_basis']=b['measurement_basis'];d['through_holes']=config(eid);return b


def setup(clients,pid=None,shape='stepped_circular'):
 e,r,pid,m,_,old_eid=round_setup(clients,pid)
 upload(e,pid,'合成基础贯通孔软件验收，非真实工程。已审最终立柱单组11根、重复2共22个基础，原材料79.2kg保持。矩形400×500×600mm、矩形台阶下600×500×净高200mm上400×300×净高400mm、圆柱实际直径600mm净高600mm、圆台阶下实际直径600mm净高200mm上直径400mm净高400mm，各层同心同轴紧接，不外伸。每基础两竖直圆孔，共44孔，孔心局部XY以各层共同中心为原点，X沿长度正向、Y沿宽度正向，H1(-50,0)、H2(50,0)mm，实际孔径40mm，自基础底贯穿全部层共600mm，孔与孔不重叠且位于每层内；合计扣减.01056πm3，矩形净2.607m3、矩形台阶净2.343m3、圆柱净3.699m3、圆台阶净2.317m3，净总体积修约3位，原π常量保留。两孔径均改50mm，扣减.0165πm3，圆台阶净2.298m3。圆台阶H2孔心X改180mm、直径40mm与上层直径400mm内相切；180.000000001mm越界须拒。H2孔心X=-10mm与H1相切，-10.000000001mm重叠须拒。相切只是几何许可，非净距结构验收。不含盲孔凹槽、斜孔、预埋实体、土方钢筋模板或基础间交叠，不重复既有清单。TEST-HOLES是仅合成软件映射编码，非官方条目。','foundation-holes-basis.txt')
 for c in get(e,pid)['components']:approve(r,pid,'components',c)
 m=get(e,pid)['components'][-1];eid=get(e,pid)['evidence'][-1]['id'];return e,r,pid,m,body(m,eid,shape),eid


@pytest.mark.parametrize('shape',SHAPES)
def test_four_shapes_balances_asts_and_unchanged_source(shape):
 s=source();b=body(s,shape=shape);before=copy.deepcopy((s,b));r=calculate(b,s);d=r['foundation_items'][0];assert (s,b)==before and r['quantity']==QUANTITIES[shape] and r['rule_version']=='railing-post-foundation-1.3' and r['foundation_summary']['total_count']=='22' and r['foundation_opening_summary']['total_hole_count']=='44'
 with localcontext() as ctx:
  ctx.prec=80;assert Decimal(r['gross'])-Decimal(r['deduction'])==Decimal(r['foundation_summary']['unrounded_quantity']);assert Decimal(r['deduction'])==Decimal(PI)*Decimal('.01056');assert evaluate(r['formula_ast'],r['dimensions'])==Decimal(r['quantity'])
  for k,a in d['formula_asts'].items():assert evaluate(a,r['dimensions'])==Decimal(d[{'single_volume':'single_foundation_volume'}.get(k,k)])
  for v in d.get('foundation_levels',[]):
   for k,a in v['formula_asts'].items():assert evaluate(a,r['dimensions'])==Decimal(v[k])
  for h in d['opening_holes']:
   assert evaluate(h['formula_ast'],r['dimensions'])==Decimal(h['single_volume']) and evaluate(h['height_ast'],r['dimensions'])==Decimal('.6')
   for c in h['containments']:
    for k,a in c['formula_asts'].items():assert evaluate(a,r['dimensions'])==Decimal(c['margins'][k])>=0
  for c in d['opening_checks']:assert evaluate(c['formula_ast'],r['dimensions'])==Decimal(c['margin'])>=0


@pytest.mark.parametrize('shape',SHAPES)
@pytest.mark.parametrize('make',[equal_body,fixed_body,fixed_added,equal_added,fixed_adjusted,equal_adjusted,combined,lambda s:combined(s,equal=True)])
def test_final_post_count_modes(shape,make):
 s=source(make);r=calculate(body(s,shape=shape),s)
 with localcontext() as ctx:ctx.prec=80;assert Decimal(r['deduction'])==Decimal(PI)*Decimal('.00048')*r['foundation_summary']['single_group_count']*2


@pytest.mark.parametrize('shape',SHAPES)
@pytest.mark.parametrize('unit,factor',[('m',Decimal('.001')),('mm',Decimal(1)),('cm',Decimal('.1')),('dm',Decimal('.01')),('km',Decimal('.000001'))])
def test_independent_hole_units_and_low_context(shape,unit,factor):
 s=source();b=body(s,shape=shape);c=b['foundation_items'][0]['through_holes'];c['unit']=unit
 for h in c['holes']:
  for k in h['parameters']:h['parameters'][k]=str(Decimal(h['parameters'][k])*factor)
 with localcontext() as ctx:ctx.prec=6;r=calculate(b,s)
 assert r==calculate(b,s) and r['quantity']==QUANTITIES[shape]


@pytest.mark.parametrize('shape',SHAPES)
@pytest.mark.parametrize('delta,allowed',[('0',True),('-.000000001',True),('.000000001',False)])
def test_layer_containment_exact_boundary(shape,delta,allowed):
 s=source();b=body(s,shape=shape);c=b['foundation_items'][0]['through_holes'];c['holes']=[hole(x=str(Decimal('280' if shape=='circular' else '180')+Decimal(delta)))];r=None
 if allowed:r=calculate(b,s);assert r['foundation_opening_summary']['total_hole_count']=='22' and r['foundation_items'][0]['opening_holes'][0]['containments'][-1]['state']==('tangent' if delta=='0' else 'inside')
 else:
  with pytest.raises(DomainError,match='越出'):calculate(b,s)


@pytest.mark.parametrize('x,allowed',[('-10',True),('-9.999999999',True),('-10.000000001',False)])
def test_pair_disjointness_exact_boundary(x,allowed):
 s=source();b=body(s);b['foundation_items'][0]['through_holes']['holes'][1]['parameters']['x']=x
 if allowed:assert calculate(b,s)['foundation_items'][0]['opening_checks'][0]['state']==('tangent' if x=='-10' else 'disjoint')
 else:
  with pytest.raises(DomainError,match='重叠'):calculate(b,s)


def test_small_top_layer_catches_hole_inside_bottom_and_hole_containment():
 s=source();b=body(s);c=b['foundation_items'][0]['through_holes'];c['holes']=[hole(x='250')]
 with pytest.raises(DomainError,match='第2层'):calculate(b,s)
 c['holes']=[hole(diameter='100'),hole('contained','-50',diameter='20')]
 with pytest.raises(DomainError,match='相互包含'):calculate(b,s)


def test_twenty_holes_ten_layers_partial_selection_and_only_total_rounding():
 s=source();s['material_items'].append({**copy.deepcopy(s['material_items'][0]),'name':'合成未选立柱项'});b=body(s);d=b['foundation_items'][0];d['levels']=[dict(parameters=dict(diameter='400',height='100'),evidence=dict(diameter='test',height='test')) for _ in range(10)];d['through_holes']['holes']=[hole('合成孔'+str(i),str(-190+i*20),diameter='10') for i in range(20)];r=calculate(b,s);assert r['foundation_opening_summary']['total_hole_count']=='440' and r['foundation_summary']['unselected_post_item_count']==1 and len(r['foundation_items'][0]['opening_checks'])==190
 with localcontext() as ctx:ctx.prec=80;assert Decimal(r['deduction'])==Decimal(PI)*Decimal('.011')


def test_mixed_holed_and_solid_rows_remain_independent():
 s=source();s['material_items'].append({**copy.deepcopy(s['material_items'][0]),'name':'合成无孔柱'});b=body(s);plain=rect_body(s)['foundation_items'][0];plain['material_index']=2;b['foundation_items'].append(plain);r=calculate(b,s);assert r['foundation_summary']['total_count']=='44' and r['foundation_opening_summary']['total_hole_count']=='44' and 'opening_holes' not in r['foundation_items'][1]
 with localcontext() as ctx:ctx.prec=80;assert Decimal(r['gross'])==Decimal(PI)*Decimal('.748')+Decimal('2.640') and Decimal(r['deduction'])==Decimal(PI)*Decimal('.01056')
 b['foundation_items'].reverse();assert calculate(b,s)==r


def test_seeded_independent_hole_sums():
 rng=random.Random(8901)
 for shape in SHAPES:
  for _ in range(20):
   s=source();b=body(s,shape=shape);d=b['foundation_items'][0];a=rng.randint(1,40);z=rng.randint(1,40);d['through_holes']['holes']=[hole(diameter=str(a)),hole('other','50',diameter=str(z))];digits=rng.randint(0,12);b['rounding_digits']=digits;plain=copy.deepcopy(b);plain['foundation_items'][0].pop('through_holes');old=calculate(plain,s);r=calculate(b,s)
   with localcontext() as ctx:
    ctx.prec=80;expected=Decimal(PI)/4*(Decimal(a*a)+z*z)*Decimal('.6')/1000000*22;net=Decimal(old['gross'])-expected;assert Decimal(r['deduction'])==expected and Decimal(r['quantity'])==net.quantize(Decimal(1).scaleb(-digits),rounding=ROUND_HALF_UP)


@pytest.mark.parametrize('fault',['null','list','extra-config','missing-config','unconfirmed','truthy','coordinate-blank','coordinate-proof','unit','holes-null','holes-dict','empty','twentyone','hole-null','hole-list','extra-hole','missing-hole','duplicate-name','blank-name','params-null','params-extra','params-missing','radius','proof-null','proof-extra','proof-missing','proof-blank','proof-bool','basis-blank','negative','zero','bool','nan','infinite','precision','x-nan','y-bool','negative-outside','covers-layer','overlap'])
def test_strict_hole_schema_and_geometry(fault):
 s=source();b=body(s);d=b['foundation_items'][0];c=d['through_holes'];h=c['holes'][0]
 if fault=='null':d['through_holes']=None
 elif fault=='list':d['through_holes']=[]
 elif fault=='extra-config':c['height']='600'
 elif fault=='missing-config':c.pop('scope_confirmed')
 elif fault=='unconfirmed':c['scope_confirmed']=False
 elif fault=='truthy':c['scope_confirmed']=1
 elif fault=='coordinate-blank':c['coordinate_basis']=' '
 elif fault=='coordinate-proof':c['coordinate_evidence_id']=False
 elif fault=='unit':c['unit']={}
 elif fault=='holes-null':c['holes']=None
 elif fault=='holes-dict':c['holes']={}
 elif fault=='empty':c['holes']=[]
 elif fault=='twentyone':c['holes']=[hole('h'+str(i)) for i in range(21)]
 elif fault=='hole-null':c['holes'][0]=None
 elif fault=='hole-list':c['holes'][0]=[]
 elif fault=='extra-hole':h['count']='2'
 elif fault=='missing-hole':h.pop('name')
 elif fault=='duplicate-name':c['holes'][1]['name']=h['name']
 elif fault=='blank-name':h['name']=''
 elif fault=='params-null':h['parameters']=None
 elif fault=='params-extra':h['parameters']['height']='600'
 elif fault=='params-missing':h['parameters'].pop('diameter')
 elif fault=='radius':h['parameters']['radius']=h['parameters'].pop('diameter')
 elif fault=='proof-null':h['evidence']=None
 elif fault=='proof-extra':h['evidence']['height']='test'
 elif fault=='proof-missing':h['evidence'].pop('x')
 elif fault=='proof-blank':h['evidence']['diameter']=''
 elif fault=='proof-bool':h['evidence']['x']=True
 elif fault=='basis-blank':h['measurement_basis']=''
 elif fault in ('negative','zero','bool','nan','infinite','precision'):h['parameters']['diameter']={'negative':'-1','zero':'0','bool':True,'nan':'NaN','infinite':'Infinity','precision':'.1234567890123456789'}[fault]
 elif fault=='x-nan':h['parameters']['x']='NaN'
 elif fault=='y-bool':h['parameters']['y']=True
 elif fault=='negative-outside':h['parameters']['x']='-500'
 elif fault=='covers-layer':h['parameters'].update(x='0',diameter='400')
 else:c['holes'][1]['parameters']['x']='-50'
 with pytest.raises(DomainError):calculate(b,s)


@pytest.mark.parametrize('shape',SHAPES)
@pytest.mark.parametrize('fmt',['json','xlsx','pdf'])
def test_workflow_exports_net_boq_mapping_and_old_quote(clients,shape,fmt):
 e,r,pid,_,quote=priced_project(clients);e,r,pid,s,b,eid=setup(clients,pid,shape);before=get(e,pid);assert post(e,pid,b,'/components/preview').json()['quantity']==QUANTITIES[shape] and get(e,pid)==before;c=post(e,pid,b).json();approve(r,pid,'components',c)
 for t in post(e,pid,{},'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
 boq=post(e,pid,{},'/boq/generate').json();assert any(v['quantity']==QUANTITIES[shape] and v['unit']=='m³' for v in boq) and any(v['quantity']=='79.200' for v in boq);url=f'/api/v1/projects/{pid}/exports';res=e.get(url,params=dict(format=fmt,unpriced=True));assert res.status_code==200,res.text
 if fmt=='json':assert TITLES<={t[0] for t in res.json()['tables']}
 elif fmt=='xlsx':
  w=load_workbook(io.BytesIO(res.content));assert TITLES<=set(w.sheetnames);assert w['基础贯通孔范围与常量'].max_row>3;assert w['基础贯通孔毛扣净'].row_dimensions[3].height>48
 else:assert all(v in ''.join(p.extract_text() for p in PdfReader(io.BytesIO(res.content)).pages) for v in TITLES)
 assert not any(t[0] in TITLES for t in e.get(url,params=dict(format='json',quote_id=quote['id'])).json()['tables'])
 if fmt=='json':
  from tests.test_mapping_profiles import rule_body,save_rule,preview,selection
  p=get(e,pid);rule={**rule_body(p,'post_foundation','volume',eid),'code':'TEST-HOLES'};assert e.patch(f'/api/v1/projects/{pid}/settings',json=dict(boq_standard=rule['standard'],boq_edition=rule['edition'],reason='仅合成贯通孔净体积口径验收')).status_code==200;save_rule(e,r,pid,rule);bid=next(v['id'] for v in p['boq'] if v['takeoff_id']==p['takeoffs'][-1]['id']);chosen=preview(e,pid,bid);assert chosen['state']=='unique',chosen;assert e.post(f'/api/v1/projects/{pid}/boq/apply-mappings',json=dict(selections=[selection(chosen)])).status_code==200;current=next(v for v in get(e,pid)['boq'] if v['id']==bid);assert current['quantity']==QUANTITIES[shape];approve(r,pid,'boq',current)


@pytest.mark.parametrize('fault',['raw','proof','confirm','coordinate','pi','hole-volume','check','containment','tier-net','tier-deduction','ast','count','book'])
def test_tamper_reconstruction(clients,fault):
 e,r,pid,s,b,eid=setup(clients);c=post(e,pid,b).json()
 if fault=='book':approve(r,pid,'components',c);post(e,pid,{},'/takeoffs/calculate')
 def forge(p):
  c=p['components'][-1];d=c['preview']['foundation_items'][0]
  if fault=='raw':c['foundation_items'][0]['through_holes']['holes'][0]['parameters']['diameter']='50'
  elif fault=='proof':c['foundation_items'][0]['through_holes']['holes'][0]['evidence']['x']='forged'
  elif fault=='confirm':c['foundation_items'][0]['through_holes']['scope_confirmed']=False
  elif fault=='coordinate':c['foundation_items'][0]['through_holes']['coordinate_basis']='changed'
  elif fault=='pi':c['preview']['foundation_pi_constant']='3.14'
  elif fault=='hole-volume':d['opening_holes'][0]['single_volume']='0'
  elif fault=='check':d['opening_checks'][0]['margin']='0'
  elif fault=='containment':d['opening_holes'][0]['containments'][1]['margins']['radial_squared']='0'
  elif fault=='tier-net':d['foundation_levels'][1]['net_level_volume']='0'
  elif fault=='tier-deduction':d['foundation_levels'][1]['single_opening_volume']='0'
  elif fault=='ast':d['formula_asts']['single_volume']={'constant':'1'}
  elif fault=='count':c['preview']['foundation_opening_summary']['total_hole_count']='88'
  else:p['takeoffs'][-1]['deduction']='0'
 mutate(pid,forge);coll='takeoffs' if fault=='book' else 'components';approve(r,pid,coll,get(e,pid)[coll][-1],expected=404 if fault=='proof' else 422 if fault=='confirm' else 409)


@pytest.mark.parametrize('gate',['preview','save','review','calculate','book_review'])
def test_hole_original_bytes_gates(clients,gate):
 e,r,pid,s,b,eid=setup(clients)
 if gate in ('review','calculate','book_review'):c=post(e,pid,b).json()
 if gate in ('calculate','book_review'):approve(r,pid,'components',c)
 if gate=='book_review':t=post(e,pid,{},'/takeoffs/calculate').json()[-1]
 d=get(e,pid)['documents'][-1];(store.ROOT/'originals'/d['hash']).write_bytes(b'corrupt')
 if gate=='preview':post(e,pid,b,'/components/preview',409)
 elif gate=='save':post(e,pid,b,expected=409)
 elif gate=='review':approve(r,pid,'components',c,expected=409)
 elif gate=='calculate':post(e,pid,{},'/takeoffs/calculate',409)
 else:approve(r,pid,'takeoffs',t,expected=409)


def test_revision_exact_history_switch_off_source_and_forged_claim(clients):
 e,r,pid,s,b,eid=setup(clients);c=post(e,pid,b).json();approve(r,pid,'components',c)
 for t in post(e,pid,{},'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
 post(e,pid,{},'/boq/generate');p=get(e,pid);old=p['components'][-1];new=copy.deepcopy(b)
 for h in new['foundation_items'][0]['through_holes']['holes']:h['parameters']['diameter']='50'
 new['object_hash']=store.digest(old);url=f'/api/v1/projects/{pid}/components/'+old['id'];res=e.put(url,json=new);assert res.status_code==200 and res.json()['preview']['quantity']=='2.298';assert get(e,pid)['component_history'][-1]==old and get(e,pid)['boq']==p['boq'] and all(t['state']=='stale' for t in get(e,pid)['takeoffs']);assert e.put(url,json=new).status_code==409;post(e,pid,{**b,'id':old['id']},expected=409);approve(r,pid,'components',res.json())
 new=round_body(s,eid);new['object_hash']=store.digest(get(e,pid)['components'][-1]);res=e.put(url,json=new);assert res.status_code==200 and res.json()['preview']['quantity']=='2.350' and 'foundation_opening_summary' not in res.json()['preview'];approve(r,pid,'components',res.json())
 s['object_hash']=store.digest(s);s['material_items'][0]['post_ranges'][0]['spacing']='3000';res=e.put(f'/api/v1/projects/{pid}/components/'+s['id'],json=s);assert res.status_code==200;approve(r,pid,'components',res.json());assert get(e,pid)['components'][-1]['state']=='stale'


@pytest.mark.parametrize('foreign',['coordinate','diameter','scope'])
def test_role_tenant_and_foreign_hole_evidence(clients,foreign):
 e,r,pid,s,b,eid=setup(clients);post(r,pid,b,expected=403);assert clients['other'].get(f'/api/v1/projects/{pid}').status_code==404;oid=create(e)['id'];upload(e,oid,'异项目孔原件');fid=get(e,oid)['evidence'][-1]['id'];c=b['foundation_items'][0]['through_holes']
 if foreign=='coordinate':c['coordinate_evidence_id']=fid
 else:c['holes'][0]['evidence']['measurement_basis' if foreign=='scope' else 'diameter']=fid
 post(e,pid,b,expected=404)


def test_long_named_twenty_holes_pdf(clients):
 e,r,pid,s,b,eid=setup(clients);b['name']='合成'+('贯通孔长名称'*25);d=b['foundation_items'][0];d['through_holes']['holes']=[hole('合成孔'+str(i)+('原件名称'*10),str(-190+i*20),diameter='10',eid=eid) for i in range(20)];c=post(e,pid,b).json();approve(r,pid,'components',c)
 for t in post(e,pid,{},'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
 res=e.get(f'/api/v1/projects/{pid}/exports',params=dict(format='pdf',unpriced=True));assert res.status_code==200,res.text
