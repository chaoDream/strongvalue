"""Synthetic linear-stock schedules; never an engineering material acceptance."""
import copy,io,json,random
from decimal import Decimal,localcontext,ROUND_HALF_UP
import pytest
from fastapi import HTTPException
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pypdf import PdfReader
from backend import store
from backend.engine import DomainError
from backend.main import app
from backend.railing_material import calculate,snapshot
from backend.railing_takeoff import calculate as railing_quantity,snapshot as geometry_snapshot
from tests.test_route_railing import geometry as route_geometry,payload as railing_payload,setup as railing_setup
from tests.test_railing_takeoff import interval
from tests.test_polygon_takeoff import post,mutate
from tests.test_workflow import get,upload,approve,create,priced_project
from tests.test_route_takeoff import point,evaluate


def geometry(eid='test',basis='spatial'):
 s=route_geometry(eid);b=railing_payload(s,eid,basis);b.update(id='synthetic-railing',revision=1,source_geometry=geometry_snapshot(s));return b


def member(name,mode,eid='test'):
 v=dict(name=name,mode=mode,quantity='2' if mode=='continuous' else '6',unit_weight='2' if mode=='continuous' else '3000',weight_unit='kg/m' if mode=='continuous' else 'g/m',measurement_basis='合成按完整范围两道，或独立六件下料，不推断间距或损耗',evidence={k:eid for k in ('quantity','unit_weight','measurement_basis')})
 if mode=='cut_pieces':v.update(cut_length='1200',length_unit='mm');v['evidence']['cut_length']=eid
 return v


def payload(s,eid='test'):
 return dict(name='合成栏杆线材表（非真实工程）',kind='railing_material',parameters={},quantity_type='mass',source_component_id=s['id'],source_component_hash=store.digest(snapshot(s)),material_items=[member('两道连续扶手', 'continuous',eid),member('独立下料立柱','cut_pieces',eid)],scope_confirmed=True,material='合成同类线材，仅软件流程验收',material_evidence_id=eid,measurement_basis='单组18m×2道×2kg/m+6件×1.2m×3kg/m=93.6kg，来源重复2=187.2kg；不含损耗接头',basis_evidence_id=eid,rounding_digits=3,rounding_evidence_id=eid)


def setup(clients,pid=None):
 e,r,pid,_,b,_=railing_setup(clients,pid)
 upload(e,pid,'合成材料依据：栏杆三维单组18m、重复2。连续扶手2道适用全部已计入范围，2kg/m；立柱独立每组6件、单件实际下料1200mm，每米质量3000g/m，不按间距估算，无损耗接头。单组72+21.6=93.6kg，总187.2kg，仅总质量修约3位。修订独立件数8时总201.6kg。栏杆来源第一段改S0→10m、第二段5m后，单组范围15m，需重新核定材料；同参数总163.2kg。非实际施工或专业签审。','railing-material-basis.txt')
 for c in get(e,pid)['components']:approve(r,pid,'components',c)
 s=post(e,pid,b).json();approve(r,pid,'components',s);s=get(e,pid)['components'][-1];eid=get(e,pid)['evidence'][-1]['id']
 return e,r,pid,s,payload(s,eid),eid


def test_multiple_runs_and_independent_cut_pieces_count_only_once_ast():
 s=geometry();b=payload(s);before=copy.deepcopy((s,b));v=calculate(b,s)
 assert (s,b)==before and v['quantity']=='187.200' and v['unit']=='kg'
 assert [Decimal(r['single_group_stock_length']) for r in v['material_items']]==[36,Decimal('7.2')]
 assert [Decimal(r['single_group_mass']) for r in v['material_items']]==[72,Decimal('21.6')]
 assert v['material_summary']['source_single_length']=='18' and v['material_summary']['multiplicity']=='2'
 with localcontext() as c:
  c.prec=80;assert evaluate(v['formula_ast'],v['dimensions'])==Decimal('187.2')
  for item in v['material_items']:assert evaluate(item['formula_ast'],v['dimensions'])==Decimal(item['single_group_mass'])


@pytest.mark.parametrize('unit,raw',[('kg/m','2'),('g/m','2000'),('kg/mm','.002'),('g/mm','2')])
def test_explicit_unit_weight_conversions(unit,raw):
 s=geometry();b=payload(s);b['material_items'][0].update(unit_weight=raw,weight_unit=unit);assert calculate(b,s)['quantity']=='187.200'


def test_cut_length_units_source_rounding_and_mass_rounding_are_independent():
 s=geometry();b=payload(s);d=b['material_items'][1];d.update(cut_length='120',length_unit='cm')
 assert calculate(b,s)['quantity']=='187.200'
 # A forged display/rounded quantity is never the input to stock length arithmetic.
 s['preview']={'quantity':'999'};b=payload(s);assert calculate(b,s)['quantity']=='187.200'
 s=geometry();src=s['source_geometry'];src['route_points']=[point('A',0,0,0),point('B',1,1,1)];s['source_component_hash']=store.digest(geometry_snapshot(src));s['edge_decisions']=railing_payload(src)['edge_decisions'];s['rounding_digits']=0
 b=payload(s);b['material_items']=b['material_items'][:1];b['rounding_digits']=12
 with localcontext() as c:
  c.prec=80;v=calculate(b,s);assert Decimal(v['quantity'])==(Decimal(3).sqrt()*8).quantize(Decimal('1e-12'))


@pytest.mark.parametrize('mode',['continuous','cut_pieces'])
def test_horizontal_projection_cannot_be_used_as_continuous_stock_length(mode):
 s=geometry(basis='horizontal');b=payload(s);b['material_items']=[member('合成材料',mode)]
 if mode=='continuous':
  with pytest.raises(DomainError,match='水平投影'):calculate(b,s)
 else:assert Decimal(calculate(b,s)['quantity'])==Decimal('43.2')


def test_polygon_horizontal_scope_and_local_gaps_are_respected():
 from tests.test_railing_takeoff import geometry as pg,payload as pp
 src=pg();s=pp(src);s.update(id='polygon-railing',revision=1,source_geometry=geometry_snapshot(src));b=payload(s)
 assert calculate(b,s)['material_summary']['source_single_length']=='14' and Decimal(calculate(b,s)['quantity'])==Decimal('155.2')
 # Excluded edges and the E1 gap never become material length.
 s['edge_decisions'][0]['ranges']=[interval(0,1)];b=payload(s);assert Decimal(calculate(b,s)['quantity'])==Decimal('131.2')


def test_seeded_independent_material_schedule_oracle():
 rng=random.Random(6701)
 for _ in range(30):
  s=geometry();src=s['source_geometry'];k=rng.randint(1,12);src['route_points']=[point('A',0,0,0),point('B',3*k,4*k,12*k)];src['multiplicity']=str(rng.randint(1,8));s['source_component_hash']=store.digest(geometry_snapshot(src));s['edge_decisions']=railing_payload(src)['edge_decisions'];b=payload(s)
  a,z=b['material_items'];a['quantity']=str(rng.randint(1,6));a['unit_weight']=str(Decimal(rng.randint(1,99))/10);z['quantity']=str(rng.randint(1,15));z['cut_length']=str(rng.randint(1,3000));z['unit_weight']=str(rng.randint(1,5000))
  expected=(Decimal(13*k)*Decimal(a['quantity'])*Decimal(a['unit_weight'])+Decimal(z['quantity'])*Decimal(z['cut_length'])/1000*Decimal(z['unit_weight'])/1000)*Decimal(src['multiplicity'])
  assert Decimal(calculate(b,s)['quantity'])==expected.quantize(Decimal('.001'),rounding=ROUND_HALF_UP)


@pytest.mark.parametrize('fault',['empty','too_many','mode','missing','extra','duplicate','missing_proof','extra_proof','blank_basis','zero_count','fraction_count','huge_count','bool_count','zero_weight','negative_weight','nan','density','unit_precision','zero_length','negative_length','area_unit','continuous_cut','mass_mode','parameters','count','scope','str_scope','digits','blank_material','missing_basis','old_hash'])
def test_strict_invalid_inputs(fault):
 s=geometry();b=payload(s);a,z=b['material_items']
 if fault=='empty':b['material_items']=[]
 elif fault=='too_many':b['material_items']*=26
 elif fault=='mode':a['mode']=True
 elif fault=='missing':a.pop('unit_weight')
 elif fault=='extra':a['loss_factor']='1.03'
 elif fault=='duplicate':z['name']=a['name']
 elif fault=='missing_proof':z['evidence'].pop('cut_length')
 elif fault=='extra_proof':a['evidence']['cut_length']='test'
 elif fault=='blank_basis':a['measurement_basis']=' '
 elif fault=='zero_count':a['quantity']='0'
 elif fault=='fraction_count':a['quantity']='1.5'
 elif fault=='huge_count':a['quantity']='1000001'
 elif fault=='bool_count':a['quantity']=True
 elif fault=='zero_weight':a['unit_weight']='0'
 elif fault=='negative_weight':a['unit_weight']='-1'
 elif fault=='nan':a['unit_weight']='NaN'
 elif fault=='density':a['weight_unit']='kg/m³'
 elif fault=='unit_precision':a['unit_weight']='.0000000000000000001'
 elif fault=='zero_length':z['cut_length']='0'
 elif fault=='negative_length':z['cut_length']='-1'
 elif fault=='area_unit':z['length_unit']='m²'
 elif fault=='continuous_cut':a['cut_length']='1'
 elif fault=='mass_mode':b['quantity_type']='length'
 elif fault=='parameters':b['parameters']={'mass':1}
 elif fault=='count':b['multiplicity']='2'
 elif fault=='scope':b['scope_confirmed']=False
 elif fault=='str_scope':b['scope_confirmed']='true'
 elif fault=='digits':b['rounding_digits']=True
 elif fault=='blank_material':b['material']=''
 elif fault=='missing_basis':b['basis_evidence_id']=''
 elif fault=='old_hash':b['source_component_hash']='old'
 with pytest.raises((DomainError,HTTPException)):calculate(b,s)


def test_candidates_preview_review_boq_and_material_profile(clients):
 e,r,pid,s,b,eid=setup(clients);before=get(e,pid);url=f'/api/v1/projects/{pid}/railing-material-candidates?component_id='+s['id'];res=e.get(url);assert res.status_code==200,res.text
 assert res.json()['source_railing']['railing_summary']['single_length']=='18' and get(e,pid)==before
 post(e,pid,b,'/components/preview');assert get(e,pid)==before
 bad=copy.deepcopy(b);bad['material_items'][0]['quantity']='1.5';post(e,pid,bad,expected=422);assert get(e,pid)==before
 c=post(e,pid,{**b,'source_geometry':{'forged':True}}).json();assert c['source_geometry']==snapshot(s)
 post(e,pid,{},'/takeoffs/calculate',409);approve(r,pid,'components',c)
 for t in post(e,pid,{},'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
 assert any(v['unit']=='kg' and Decimal(v['quantity'])==Decimal('187.2') for v in post(e,pid,{},'/boq/generate').json())
 from backend.mapping_types import component_profile
 assert component_profile(c,c['preview'])['id']=='mass'


@pytest.mark.parametrize('fmt',['json','xlsx','pdf'])
def test_exports_and_old_quote_isolation(clients,fmt):
 e,r,pid,_,quote=priced_project(clients);e,r,pid,s,b,eid=setup(clients,pid);c=post(e,pid,b).json();approve(r,pid,'components',c)
 for t in post(e,pid,{},'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
 url=f'/api/v1/projects/{pid}/exports';res=e.get(url,params={'format':fmt,'unpriced':True});assert res.status_code==200,res.text
 titles=['栏杆材料质量汇总','栏杆材料逐项计算','栏杆材料原始参数','栏杆材料计长来源']
 if fmt=='json':assert all(any(t[0]==v for t in res.json()['tables']) for v in titles)
 elif fmt=='xlsx':assert set(titles)<=set(load_workbook(io.BytesIO(res.content)).sheetnames)
 else:assert all(v in ''.join(p.extract_text() for p in PdfReader(io.BytesIO(res.content)).pages) for v in titles)
 old=e.get(url,params={'format':'json','quote_id':quote['id']}).json();assert old['quote']==quote and not any(t[0] in titles for t in old['tables'])


@pytest.mark.parametrize('fault',['row','weight','formula','range','source','stamp','evidence','book'])
def test_saved_and_book_records_are_rebuilt_before_review(clients,fault):
 e,r,pid,s,b,eid=setup(clients);c=post(e,pid,b).json()
 if fault=='book':approve(r,pid,'components',c);post(e,pid,{},'/takeoffs/calculate')
 def forge(p):
  x=p['components'][-1];v=x['preview']
  if fault=='row':v['material_items'][0]['single_group_mass']='999'
  elif fault=='weight':x['material_items'][0]['unit_weight']='999'
  elif fault=='formula':v['formula_ast']={'constant':'187.2'}
  elif fault=='range':v['source_railing']['railing_edges'][0]['selected_length']='999'
  elif fault=='source':x['source_geometry']['edge_decisions'][0]['mode']='exclude'
  elif fault=='stamp':x['source_stamp']='forged'
  elif fault=='evidence':x['evidence']['material.basis']='forged'
  else:p['takeoffs'][-1]['material_items'][0]['normalized_quantity']='999'
 mutate(pid,forge);coll='takeoffs' if fault=='book' else 'components';approve(r,pid,coll,get(e,pid)[coll][-1],expected=409)


def test_revision_history_concurrency_source_range_and_transitive_invalidation(clients):
 e,r,pid,s,b,eid=setup(clients);c=post(e,pid,b).json();approve(r,pid,'components',c)
 for t in post(e,pid,{},'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
 post(e,pid,{},'/boq/generate');old=get(e,pid)['components'][-1];new=copy.deepcopy(b);new['material_items'][1]['quantity']='8';new['object_hash']=store.digest(old);url=f'/api/v1/projects/{pid}/components/'+c['id'];res=e.put(url,json=new);assert res.status_code==200,res.text
 p=get(e,pid);assert p['component_history'][-1]==old and all(t['state']=='stale' for t in p['takeoffs']) and any(v['quantity']=='187.200' for v in p['boq'])
 assert res.json()['preview']['quantity']=='201.600' and e.put(url,json=new).status_code==409;approve(r,pid,'components',res.json())
 for t in post(e,pid,{},'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
 assert any(v['quantity']=='201.600' for v in post(e,pid,{},'/boq/generate').json())
 s['object_hash']=store.digest(s);s['edge_decisions'][0].update(mode='ranges',ranges=[interval(0,10,eid)]);res=e.put(f'/api/v1/projects/{pid}/components/'+s['id'],json=s);assert res.status_code==200,res.text
 assert get(e,pid)['components'][-1]['state']=='stale';approve(r,pid,'components',res.json());post(e,pid,new,'/components/preview',409)
 # A new source binding computes the actual changed range; copied rounded old length is ignored.
 current=get(e,pid)['components'][-2];fresh=payload(current,eid);assert post(e,pid,fresh,'/components/preview').json()['quantity']=='163.200'
 geom=get(e,pid)['components'][0];geom['object_hash']=store.digest(geom);geom['route_points'][1]['parameters']['z']='9';res=e.put(f'/api/v1/projects/{pid}/components/'+geom['id'],json=geom);assert res.status_code==200,res.text
 assert get(e,pid)['components'][-2]['state']=='stale' and get(e,pid)['components'][-1]['state']=='stale'


@pytest.mark.parametrize('gate',['candidate','preview','save','review','calculate','book_review'])
def test_original_bytes_gate_material_and_source_parameters(clients,gate):
 e,r,pid,s,b,eid=setup(clients)
 if gate in ('review','calculate','book_review'):c=post(e,pid,b).json()
 if gate in ('calculate','book_review'):approve(r,pid,'components',c)
 if gate=='book_review':t=post(e,pid,{},'/takeoffs/calculate').json()[-1]
 doc=get(e,pid)['documents'][-1];(store.ROOT/'originals'/doc['hash']).write_bytes(b'corrupt')
 if gate=='candidate':
  # Candidate checks the approved railing geometry, not as-yet-unselected material proofs.
  assert e.get(f'/api/v1/projects/{pid}/railing-material-candidates?component_id='+s['id']).status_code==200
 elif gate=='preview':post(e,pid,b,'/components/preview',409)
 elif gate=='save':post(e,pid,b,expected=409)
 elif gate=='review':approve(r,pid,'components',c,expected=409)
 elif gate=='calculate':post(e,pid,{},'/takeoffs/calculate',409)
 else:approve(r,pid,'takeoffs',t,expected=409)


def test_source_original_corruption_and_empty_scope_are_rejected(clients):
 e,r,pid,s,b,eid=setup(clients);doc=get(e,pid)['documents'][0];(store.ROOT/'originals'/doc['hash']).write_bytes(b'corrupt');assert e.get(f'/api/v1/projects/{pid}/railing-material-candidates?component_id='+s['id']).status_code==409
 s=geometry()
 for d in s['edge_decisions']:d.update(mode='exclude',ranges=[])
 with pytest.raises(DomainError,match='范围为空'):calculate(payload(s),s)


def test_permissions_tenant_foreign_proof_members_and_unapproved_source(clients):
 e,r,pid,s,b,eid=setup(clients);url=f'/api/v1/projects/{pid}/railing-material-candidates?component_id='+s['id'];assert r.get(url).status_code==200 and clients['other'].get(url).status_code==404 and TestClient(app).get(url).status_code==401
 post(r,pid,b,expected=403);other=create(e)['id'];upload(e,other,'其他项目依据');bad=copy.deepcopy(b);bad['material_items'][0]['evidence']['unit_weight']=get(e,other)['evidence'][-1]['id'];post(e,pid,bad,expected=404)
 assert e.get(f'/api/v1/projects/{other}/railing-material-candidates?component_id='+s['id']).status_code==404
 mutate(pid,lambda p:p['components'][-1].update(state='draft'));assert e.get(url).status_code==409
 from tests.test_project_access import restrict
 restrict(clients['admin'],pid,[]);assert e.get(url).status_code==404


def test_large_origin_tiny_stock_lengths_and_low_caller_precision():
 s=geometry();src=s['source_geometry']
 with localcontext() as c:
  c.prec=80
  for p in src['route_points']:p['parameters']={k:str(Decimal('1000000000')+Decimal(v)*Decimal('1e-8')) for k,v in p['parameters'].items()}
 s['source_component_hash']=store.digest(geometry_snapshot(src));s['edge_decisions']=railing_payload(src)['edge_decisions'];b=payload(s);b['material_items']=b['material_items'][:1];b['rounding_digits']=12
 with localcontext() as c:
  c.prec=10;v=calculate(b,s)
 assert Decimal(v['quantity'])==Decimal('144e-8')


def test_normalized_weight_precision_and_source_hash_guard():
 s=geometry();b=payload(s);b['material_items'][0].update(unit_weight='1.23456789012345678',weight_unit='kg/mm')
 v=calculate(b,s)
 with localcontext() as c:
  c.prec=80;assert evaluate(v['formula_ast'],v['dimensions'])==Decimal(v['quantity'])
 b['material_items'][0].update(unit_weight='1e-18',weight_unit='g/m')
 with pytest.raises(DomainError):calculate(b,s)
