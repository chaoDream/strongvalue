"""Synthetic multi-stock purchase-length optimization, evidence and version gates."""
import copy,io,random
from decimal import Decimal,localcontext
from fractions import Fraction
import pytest
from openpyxl import load_workbook
from pypdf import PdfReader
from backend import store,mixed_stock_cutting as mixed,multi_stock_cutting as multi
from backend.engine import DomainError
from backend.railing_material import calculate
from tests.test_mixed_stock_cutting import body as single_body,setup as single_setup,TITLES,geometry,equal_body,fixed_body,fixed_added,equal_added,fixed_adjusted,equal_adjusted,combined,local_body,path_body
from tests.test_railing_takeoff import interval
from tests.test_railing_material import member
from tests.test_route_takeoff import evaluate
from tests.test_polygon_takeoff import post,mutate
from tests.test_workflow import get,approve,upload,create,priced_project

TITLE='多定尺候选与选用'


def choices(eid='test',lengths=('3523','2423')):
 return dict(alternatives=[dict(stock_length=s,evidence_id=eid) for s in lengths],supply_confirmed=True,measurement_basis='合成同规格方管A，全部定尺可按需供应，无库存限量；所有规格沿用每件3mm锯缝和两端各10mm。先最小购入总长度，同长度再最少根数，不计采购价或增加净量，非真实工程',evidence_id=eid)


def body(s,eid='test'):
 b=single_body(s,eid);b['name']='合成多定尺联合选材（非真实工程）';b['mixed_cutting'][0]['stock_choices']=choices(eid);return b


def setup(clients,pid=None):
 e,r,pid,s,b,eid=single_setup(clients,pid)
 upload(e,pid,'合成多定尺软件验收，非真实工程。同材质同截面方管A，每米3000g，长短每组各2件3500/2400mm，来源重复2共8件净23.6m70.8kg。原定尺6000mm，备选3523/2423mm，全部可按需供应，无限量库存；每件3mm锯缝含最后一件，每根两端各10mm未重复，留用下限50mm。目标先最小购入总长度，同长度再最少根数。选择3523mm四根及2423mm四根，零余料，购入23.784m71.352kg，净70.8kg不变；搜索穷尽证明双目标最优。仅6000mm为4根24m72kg。追加5926mm后可选四根共23.704m71.112kg，达到双目标下界。各定尺同规格、同加工口径，未按供应价格计价，不自动入库。','multi-stock-basis.txt')
 for c in get(e,pid)['components']:approve(r,pid,'components',c)
 s=get(e,pid)['components'][-1];eid=get(e,pid)['evidence'][-1]['id'];return e,r,pid,s,body(s,eid),eid


def group(s,b):return calculate(b,s)['mixed_stock_cutting']['groups'][0]


def test_exact_multi_stock_net_invariance_assignment_balances_and_all_asts():
 s=geometry();b=body(s);before=copy.deepcopy((s,b));v=calculate(b,s);g=v['mixed_stock_cutting']['groups'][0];old=single_body(s);base=calculate(old,s)
 assert (s,b)==before and v['quantity']==base['quantity']=='70.800' and v['formula_ast']==base['formula_ast'] and v['material_items']==base['material_items'] and v['rule_version']=='railing-material-mass-1.15'
 assert g['stock_selection']['options']==[dict(index=1,stock_length='6',stock_count=0),dict(index=2,stock_length='3.523',stock_count=4),dict(index=3,stock_length='2.423',stock_count=4)]
 assert g['proof']['status']=='length_optimal_search' and g['proof']['search_exhausted'] and g['proof']['length_lower_bound']=='23.704' and g['stock_count']==8 and g['proof']['purchased_length']=='23.784'
 assert g['total']['purchased_mass']=='71.352' and g['total']['retained_length']==g['total']['scrap_length']==g['balance_difference']=='0' and g['total']['trim_length']=='0.16'
 assert len([k for k in v['source_evidence'] if '.mixed_cutting.' in k])==10
 assert sorted((p['item_index'],p['piece_index']) for bar in g['bars'] for p in bar['pieces'])==[(i,j) for i in (1,2) for j in range(1,5)]
 for bar in g['bars']:
  assert bar['piece_count']==1 and bar['pieces'][0]['from_s']=='0.01' and bar['offcut_disposition']=='none'
  assert Fraction(bar['pieces'][0]['kerf_to_s'])+Fraction('.01')==Fraction(bar['values']['purchased_length'])
  for k,ast in bar['formula_asts'].items():assert evaluate(ast,v['dimensions'])==Decimal(bar['values'][k])
 for k,ast in g['total_formula_asts'].items():assert evaluate(ast,v['dimensions'])==Decimal(g['total'][k])
 for k,ast in v['mixed_stock_cutting']['total_formula_asts'].items():assert evaluate(ast,v['dimensions'])==Decimal(v['mixed_stock_cutting']['total'][k])


def test_secondary_objective_and_both_lower_bounds():
 s=geometry();b=body(s);b['mixed_cutting'][0]['stock_choices']['alternatives'].append(dict(stock_length='5926',evidence_id='test'));g=group(s,b)
 assert g['stock_count']==4 and g['total']['purchased_length']=='23.704' and g['proof']['status']=='length_optimal_bound' and g['stock_selection']['options'][-1]['stock_count']==4
 # Same total purchased length: 2 short bars vs 1 long bar. Secondary objective must choose 1.
 plan,p=multi.pack([5,5],[5,10],0);assert p['purchased_length']==10 and p['stock_count']==1 and plan==[(1,[0,1])]


def test_long_first_fit_improvement_and_limit_zero():
 sizes=[6,5,3,2,2,2];stocks=[10,11];plan,p=multi.pack(sizes,stocks,0);assert p['purchased_length']==20 and p['stock_count']==2 and p['initial_purchased_length']>20
 plan,p=multi.pack(sizes,stocks,0,limit=0);assert p['status']=='length_feasible_unproven' and not p['search_exhausted'] and p['unproven_length_gap']>0
 assert sorted(i for _,ps in plan for i in ps)==list(range(len(sizes)))


def test_primary_may_only_fit_short_piece_and_alternative_fits_long():
 s=geometry();b=body(s);b['mixed_cutting'][0]['stock_length']='2423';b['mixed_cutting'][0]['stock_choices']=choices(lengths=('3523',));g=group(s,b);assert g['total']['purchased_length']=='23.784'


@pytest.mark.parametrize('unit,factor',[('m',Decimal('.001')),('mm',Decimal('1')),('cm',Decimal('.1')),('dm',Decimal('.01')),('km',Decimal('.000001'))])
def test_parameter_units_and_conversion_asts(unit,factor):
 s=geometry();b=body(s);c=b['mixed_cutting'][0];c['length_unit']=unit
 for k in ('stock_length','kerf_per_piece','trim_start','trim_end','reusable_min_length'):c[k]=str(Decimal(c[k])*factor)
 for o in c['stock_choices']['alternatives']:o['stock_length']=str(Decimal(o['stock_length'])*factor)
 assert group(s,b)['total']['purchased_mass']=='71.352'


@pytest.mark.parametrize('weight,unit',[('3','kg/m'),('3000','g/m'),('.003','kg/mm'),('3','g/mm')])
def test_weight_units(weight,unit):
 s=geometry();b=body(s)
 for m in b['material_items']:m.update(unit_weight=weight,weight_unit=unit)
 assert group(s,b)['total']['purchased_mass']=='71.352'


@pytest.mark.parametrize('fault',['null','list','empty','extra','missing','unconfirmed','truthy','no-basis','no-proof','proof-type','alternatives-null','alternatives-dict','alternatives-empty','too-many','row-null','row-list','row-missing','row-extra','row-no-proof','zero','negative','nan','infinite','bool','precision','trim','duplicate-primary','duplicate-alternative','too-short','all-overlong'])
def test_strict_input_source_and_scope_rejections(fault):
 s=geometry();b=body(s);g=b['mixed_cutting'][0];c=g['stock_choices'];a=c['alternatives']
 if fault=='null':g['stock_choices']=None
 elif fault=='list':g['stock_choices']=[]
 elif fault=='empty':g['stock_choices']={}
 elif fault=='extra':c['price']=1
 elif fault=='missing':c.pop('supply_confirmed')
 elif fault=='unconfirmed':c['supply_confirmed']=False
 elif fault=='truthy':c['supply_confirmed']=1
 elif fault=='no-basis':c['measurement_basis']=''
 elif fault=='no-proof':c['evidence_id']=''
 elif fault=='proof-type':c['evidence_id']=False
 elif fault=='alternatives-null':c['alternatives']=None
 elif fault=='alternatives-dict':c['alternatives']={}
 elif fault=='alternatives-empty':c['alternatives']=[]
 elif fault=='too-many':c['alternatives']=a*3
 elif fault=='row-null':a[0]=None
 elif fault=='row-list':a[0]=[]
 elif fault=='row-missing':a[0].pop('stock_length')
 elif fault=='row-extra':a[0]['quantity']=4
 elif fault=='row-no-proof':a[0]['evidence_id']=''
 elif fault=='zero':a[0]['stock_length']='0'
 elif fault=='negative':a[0]['stock_length']='-2'
 elif fault=='nan':a[0]['stock_length']='NaN'
 elif fault=='infinite':a[0]['stock_length']='Infinity'
 elif fault=='bool':a[0]['stock_length']=True
 elif fault=='precision':a[0]['stock_length']='3523.000000000000001'
 elif fault=='trim':a[0]['stock_length']='20'
 elif fault=='duplicate-primary':a[0]['stock_length']='6000.0'
 elif fault=='duplicate-alternative':a[1]['stock_length']='3523'
 elif fault=='too-short':a[0]['stock_length']='200'
 else:
  g['stock_length']='2500';c['alternatives']=[dict(stock_length='2600',evidence_id='test')]
 with pytest.raises(DomainError):calculate(b,s)


def test_exact_threshold_and_micro_length_floats_not_used():
 s=geometry();b=body(s);c=b['mixed_cutting'][0];c['stock_choices']['alternatives'][0]['stock_length']='3522.999999999';g=group(s,b);assert g['stock_selection']['options'][0]['stock_count']>0 and g['stock_selection']['options'][1]['stock_count']==0
 c['stock_choices']['alternatives'][0]['stock_length']='3523.000000001';g=group(s,b);assert g['stock_selection']['options'][1]['stock_count']==4 and Decimal(g['total']['scrap_length'])==Decimal('0.000000000004')


def test_option_order_only_relabels_choice_not_objective():
 s=geometry();b=body(s);g=group(s,b);b['mixed_cutting'][0]['stock_choices']['alternatives'].reverse();h=group(s,b);assert h['total']==g['total'] and h['proof']==g['proof']


def test_seeded_independent_exhaustive_lexicographic_oracle():
 rng=random.Random(8201)
 def oracle(sizes,stocks,trim):
  best=(sum(max(stocks) for _ in sizes),len(sizes));bins=[]
  def visit(i):
   nonlocal best
   if i==len(sizes):best=min(best,(sum(s for s,_ in bins),len(bins)));return
   for j,(stock,load) in enumerate(bins):
    if load+sizes[i]<=stock-trim:bins[j]=(stock,load+sizes[i]);visit(i+1);bins[j]=(stock,load)
   for stock in stocks:
    if sizes[i]<=stock-trim:bins.append((stock,sizes[i]));visit(i+1);bins.pop()
  visit(0);return best
 for _ in range(60):
  trim=rng.randint(0,2);stocks=sorted(rng.sample(range(5,15),rng.randint(2,3)));sizes=sorted([rng.randint(1,max(stocks)-trim) for _ in range(rng.randint(2,6))],reverse=True);plan,p=multi.pack(sizes,stocks,trim)
  assert (p['purchased_length'],p['stock_count'])==oracle(sizes,stocks,trim) and p['status']!='length_feasible_unproven'
  assert all(sum(sizes[i] for i in ps)<=stocks[j]-trim for j,ps in plan) and sorted(i for _,ps in plan for i in ps)==list(range(len(sizes)))


def test_maximum_200_pieces_five_choices_and_low_precision():
 s=geometry();b=body(s)
 for m in b['material_items']:m['quantity']='50'
 b['mixed_cutting'][0]['stock_choices']['alternatives']+= [dict(stock_length=x,evidence_id='test') for x in ('5926','8000')]
 with localcontext() as ctx:ctx.prec=6;v=calculate(b,s)
 assert v==calculate(b,s) and v['mixed_stock_cutting']['groups'][0]['required_piece_count']==200


@pytest.mark.parametrize('fmt',['json','xlsx','pdf'])
def test_workflow_exports_and_old_quote_isolation(clients,fmt):
 e,r,pid,_,quote=priced_project(clients);e,r,pid,s,b,eid=setup(clients,pid);before=get(e,pid);post(e,pid,b,'/components/preview');assert get(e,pid)==before;c=post(e,pid,b).json();assert c['mixed_cutting']==b['mixed_cutting'];approve(r,pid,'components',c)
 for t in post(e,pid,{},'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
 assert any(t['quantity']=='70.800' for t in post(e,pid,{},'/boq/generate').json());url=f'/api/v1/projects/{pid}/exports';res=e.get(url,params=dict(format=fmt,unpriced=True));assert res.status_code==200,res.text
 if fmt=='json':assert TITLES|{TITLE}<={t[0] for t in res.json()['tables']}
 elif fmt=='xlsx':assert TITLES|{TITLE}<=set(load_workbook(io.BytesIO(res.content)).sheetnames)
 else:assert all(t in ''.join(p.extract_text() for p in PdfReader(io.BytesIO(res.content)).pages) for t in TITLES|{TITLE})
 assert not any(t[0]==TITLE for t in e.get(url,params=dict(format='json',quote_id=quote['id'])).json()['tables'])


@pytest.mark.parametrize('fault',['raw','evidence','supply','status','choice','length','ast','total','book'])
def test_saved_multi_stock_reconstruction(clients,fault):
 e,r,pid,s,b,eid=setup(clients);c=post(e,pid,b).json()
 if fault=='book':approve(r,pid,'components',c);post(e,pid,{},'/takeoffs/calculate')
 def forge(p):
  c=p['components'][-1];g=c['preview']['mixed_stock_cutting']['groups'][0]
  if fault=='raw':c['mixed_cutting'][0]['stock_choices']['alternatives'][0]['stock_length']='3600'
  elif fault=='evidence':c['mixed_cutting'][0]['stock_choices']['alternatives'][0]['evidence_id']='forged'
  elif fault=='supply':c['mixed_cutting'][0]['stock_choices']['supply_confirmed']=False
  elif fault=='status':g['proof']['status']='length_feasible_unproven'
  elif fault=='choice':g['bars'][0]['stock_choice_index']=1
  elif fault=='length':g['stock_selection']['options'][1]['stock_length']='3.5'
  elif fault=='ast':g['bars'][0]['formula_asts']['purchased_length']={'constant':'6'}
  elif fault=='total':g['total']['purchased_mass']='72'
  else:p['takeoffs'][-1]['mixed_stock_cutting']['groups'][0]['stock_selection']['options'][0]['stock_count']=4
 mutate(pid,forge);coll='takeoffs' if fault=='book' else 'components';approve(r,pid,coll,get(e,pid)[coll][-1],expected=404 if fault=='evidence' else 422 if fault=='supply' else 409)


def test_same_net_revision_history_concurrency_removal(clients):
 e,r,pid,s,b,eid=setup(clients);c=post(e,pid,b).json();approve(r,pid,'components',c)
 for t in post(e,pid,{},'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
 post(e,pid,{},'/boq/generate');before=get(e,pid);old=before['components'][-1];new=copy.deepcopy(b);new['mixed_cutting'][0]['stock_choices']['alternatives'].append(dict(stock_length='5926',evidence_id=eid));new['object_hash']=store.digest(old);url=f'/api/v1/projects/{pid}/components/'+old['id'];res=e.put(url,json=new);assert res.status_code==200,res.text
 p=get(e,pid);assert p['component_history'][-1]==old and p['boq']==before['boq'] and all(t['state']=='stale' for t in p['takeoffs']) and res.json()['preview']['quantity']=='70.800' and res.json()['preview']['mixed_stock_cutting']['total']['purchased_mass']=='71.112';assert e.put(url,json=new).status_code==409;approve(r,pid,'components',res.json())
 for t in post(e,pid,{},'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
 current=get(e,pid)['components'][-1];new['object_hash']=store.digest(current);new['mixed_cutting'][0].pop('stock_choices');res=e.put(url,json=new);assert res.status_code==200 and res.json()['preview']['rule_version']=='railing-material-mass-1.14' and 'stock_selection' not in res.json()['preview']['mixed_stock_cutting']['groups'][0];assert get(e,pid)['component_history'][-1]==current


@pytest.mark.parametrize('gate',['preview','save','review','calculate','book_review'])
def test_multi_stock_original_bytes_gate(clients,gate):
 e,r,pid,s,b,eid=setup(clients)
 if gate in ('review','calculate','book_review'):c=post(e,pid,b).json()
 if gate in ('calculate','book_review'):approve(r,pid,'components',c)
 if gate=='book_review':t=post(e,pid,{},'/takeoffs/calculate').json()[-1]
 d=get(e,pid)['documents'][-1];(store.ROOT/'originals'/d['hash']).write_bytes(b'corrupt')
 if gate=='preview':post(e,pid,b,'/components/preview',409)
 elif gate=='save':post(e,pid,b,expected=409)
 elif gate=='review':approve(r,pid,'components',c,expected=409)
 elif gate=='calculate':post(e,pid,{},'/takeoffs/calculate',409)
 else:approve(r,pid,'takeoffs',t,expected=409)


def test_roles_tenant_foreign_option_and_supply_evidence(clients):
 e,r,pid,s,b,eid=setup(clients);post(r,pid,b,expected=403);assert clients['other'].get(f'/api/v1/projects/{pid}').status_code==404
 other=create(e)['id'];upload(e,other,'异项目');foreign=get(e,other)['evidence'][-1]['id'];original=copy.deepcopy(b)
 b['mixed_cutting'][0]['stock_choices']['alternatives'][0]['evidence_id']=foreign;post(e,pid,b,expected=404)
 b=original;b['mixed_cutting'][0]['stock_choices']['evidence_id']=foreign;post(e,pid,b,expected=404)


def test_unproven_status_export_with_zero_budget(clients,monkeypatch):
 monkeypatch.setattr(mixed,'SEARCH_LIMIT',0);e,r,pid,s,b,eid=setup(clients);c=post(e,pid,b).json();assert c['preview']['mixed_stock_cutting']['groups'][0]['proof']['status']=='length_feasible_unproven';approve(r,pid,'components',c)
 for t in post(e,pid,{},'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
 res=e.get(f'/api/v1/projects/{pid}/exports',params=dict(format='json',unpriced=True));assert res.status_code==200 and '尚未证明双目标最优' in str(res.json()['tables'])


def test_single_and_multi_groups_mixed_coverage_and_long_names_pdf(clients):
 e,r,pid,s,b,eid=setup(clients);b['name']='合成'+('长名称验收'*39)+'末';b['mixed_cutting'][0]['name']='合成'+('长混排组名'*39)+'末'
 second=copy.deepcopy(b['mixed_cutting'][0]);second.pop('stock_choices');second['name']='单定尺组';second['members']=['长料二','短料二'];b['mixed_cutting'].append(second)
 for old,name in zip(copy.deepcopy(b['material_items']),second['members']):old['name']=name;b['material_items'].append(old)
 c=post(e,pid,b).json();assert c['preview']['mixed_stock_cutting']['covered_item_count']==4 and c['preview']['mixed_stock_cutting']['total']['purchased_mass']=='143.352';approve(r,pid,'components',c)
 for t in post(e,pid,{},'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
 res=e.get(f'/api/v1/projects/{pid}/exports',params=dict(format='pdf',unpriced=True));assert res.status_code==200,res.text


@pytest.mark.parametrize('make',[equal_body,fixed_body,fixed_added,equal_added,fixed_adjusted,equal_adjusted,combined,lambda s:combined(s,equal=True),local_body,path_body])
def test_post_modes_and_spacing_preserved_with_multiple_stocks(make):
 s=geometry();b=make(s);m=b['material_items'][0];m['name']='合成立柱';b['material_items'].append({**member('合成短料','cut_pieces'),'quantity':'2','cut_length':'2400'});old=calculate(copy.deepcopy(b),s)
 from tests.test_mixed_stock_cutting import config
 c=config(members=['合成立柱','合成短料']);c['stock_choices']=choices(lengths=('3000','4000'));b['mixed_cutting']=[c];v=calculate(b,s);assert v['quantity']==old['quantity'] and v['material_items']==old['material_items']


def test_source_revision_invalidates_multi_stock_even_same_cut_lengths(clients):
 e,r,pid,s,b,eid=setup(clients);c=post(e,pid,b).json();approve(r,pid,'components',c);s['object_hash']=store.digest(s);s['edge_decisions'][0].update(mode='ranges',ranges=[interval(0,10,eid)]);res=e.put(f'/api/v1/projects/{pid}/components/'+s['id'],json=s);assert res.status_code==200;approve(r,pid,'components',res.json());assert get(e,pid)['components'][-1]['state']=='stale'
