"""Synthetic rectangular ledge windows, triple clipping and immutable review gates."""
import copy,io,math,random
from decimal import Decimal,localcontext,ROUND_HALF_UP
import pytest
from openpyxl import load_workbook
from pypdf import PdfReader
from backend import store
from backend.engine import DomainError
from backend.railing_material import snapshot
from backend.foundation_horizontal_formwork import calculate
from backend.foundation_step_windows import VERSION
from backend.foundation_horizontal_formwork_reports import STEP_WINDOW_TITLES,WINDOW_TITLES,WINDOW_CURVED_TITLE,tables
from tests.test_foundation_horizontal_windows import window,FRAME,check as window_check
from tests.test_foundation_step_formwork import source,body as step_body,contact,setup as step_setup
from tests.test_foundation_horizontal_formwork import body as end_body
from tests.test_foundation_openings import hole
from tests.test_foundation_voids import rectangle
from tests.test_foundation_recesses import interval
from tests.test_post_foundation import equal_body,fixed_body,fixed_added,equal_added,fixed_adjusted,equal_adjusted,combined
from tests.test_route_takeoff import evaluate
from tests.test_polygon_takeoff import post,mutate
from tests.test_workflow import get,approve,upload,create,priced_project


def body(s,eid='test',bounds=('180','240','-100','100'),level=1):
    b=step_body(s,eid);b.update(name='合成矩形平台局部窗口（非真实工程）',measurement_basis='仅合成矩形平台窗口实际支模，先扣上层共同覆盖，再扣外露平台内孔，未选范围不计。')
    b['contact_ranges']=[{**contact(eid,level),'plan_range':window(eid,*bounds)}];return b


def check(r):
    window_check(r)
    with localcontext() as ctx:
        ctx.prec=80
        for d in r['contact_ranges']:
            for h in d['apertures']:
                for key in ('curved_intersection','window_hole_curved_intersection'):
                    if key in h:
                        for q in h[key]['terms']:assert evaluate(q['formula_ast'],r['dimensions'])==Decimal(q['area'])


@pytest.mark.parametrize('make',[equal_body,fixed_body,fixed_added,equal_added,fixed_adjusted,equal_adjusted,combined,lambda s:combined(s,equal=True)])
@pytest.mark.parametrize('axis',['x','y'])
def test_partition_counts_and_level_geometry_once(make,axis):
    s=source(make=make);before=copy.deepcopy(s);b=body(s,bounds=('-1000','0','-1000','1000') if axis=='x' else ('-1000','1000','-1000','0'));other=copy.deepcopy(b['contact_ranges'][0]);other['name']='合成另半平台';other['plan_range']=window(x1='0' if axis=='x' else '-1000',x2='1000',y1='-1000' if axis=='x' else '0',y2='1000');b['contact_ranges'].append(other);r=calculate(b,s);whole=calculate(step_body(s),s);assert s==before and r['quantity']==whole['quantity'] and r['rule_version']==VERSION;assert r['horizontal_window_checks'][0]['state']=='boundary_touch' and r['horizontal_window_checks'][0]['level_index']==1;check(r)


@pytest.mark.parametrize('unit,factor',[('mm','1'),('cm','.1'),('dm','.01'),('m','.001'),('km','.000001')])
@pytest.mark.parametrize('source_unit,source_factor',[('mm','1'),('cm','.1'),('dm','.01'),('m','.001'),('km','.000001')])
def test_25_source_and_window_units_low_precision(unit,factor,source_unit,source_factor):
    s=source();item=s['foundation_items'][0];item['length_unit']=source_unit
    for t in item['levels']:
        for k in t['parameters']:t['parameters'][k]=str(Decimal(t['parameters'][k])*Decimal(source_factor))
    b=body(s);w=b['contact_ranges'][0]['plan_range'];w['unit']=unit
    for k in w['parameters']:w['parameters'][k]=str(Decimal(w['parameters'][k])*Decimal(factor))
    with localcontext() as ctx:ctx.prec=6;r=calculate(b,s)
    assert r==calculate(b,s) and r['quantity']=='0.123';check(r)


@pytest.mark.parametrize('bounds,gross,cut',[
    (('-1000','1000','-1000','1000'),'.18','.003'),
    (('150','250','-100','100'),'.01','.003'),
    (('180','240','-100','100'),'.008','.0024'),
    (('-100','100','-100','100'),'0','0'),
    (('300','400','-100','100'),'0','0'),
    (('200','300','-100','100'),'.02','.003'),
    (('200','250','-30','30'),'.003','.003'),
    (('200','250','30','100'),'.0035','0'),
    (('-400','-300','-400','400'),'0','0'),
    (('150','250','100','200'),'.0075','0')])
def test_window_clips_lower_upper_and_hole_independently(bounds,gross,cut):
    s=source();b=body(s,bounds=bounds);before=copy.deepcopy((s,b));r=calculate(b,s);d=r['contact_ranges'][0];assert (s,b)==before;assert Decimal(d['single_gross_area'])==Decimal(gross) and Decimal(d['single_opening_area'])==Decimal(cut);check(r)


@pytest.mark.parametrize('axis,sign',[('x',1),('x',-1),('y',1),('y',-1)])
def test_circle_hole_window_and_upper_joint_clip_independent_cap(axis,sign):
    edge=200 if axis=='x' else 150;center=(edge-10)*sign;h=hole(x=str(center) if axis=='x' else '0',y=str(center) if axis=='y' else '0');s=source(holes=[{**h,'vertical_range':interval('0','200')}]);lo,hi=((edge-20,edge+40) if sign>0 else (-edge-40,-edge+20));bounds=(str(lo),str(hi),'-100','100') if axis=='x' else ('-100','100',str(lo),str(hi));r=calculate(body(s,bounds=bounds),s);d=r['contact_ranges'][0];h=d['apertures'][0];cap=.02**2*math.acos(.5)-.01*math.sqrt(.02**2-.01**2);assert abs(float(h['single_area'])-cap)<2e-18 and abs(float(d['single_gross_area'])-.008)<1e-16;assert h['window_hole_curved_intersection'] and h['curved_intersection'];check(r)


def test_200_rectangle_windows_independent_integer_triple_intersection():
    rng=random.Random(10401)
    def clip(*rects):
        return max(0,min(r[1] for r in rects)-max(r[0] for r in rects))*max(0,min(r[3] for r in rects)-max(r[2] for r in rects))
    for _ in range(200):
        box=(*sorted(rng.sample(range(-400,401),2)),*sorted(rng.sample(range(-400,401),2)));s=source();b=body(s,bounds=tuple(map(str,box)));b['rounding_digits']=rng.randrange(13);r=calculate(b,s);lower=(-300,300,-250,250);upper=(-200,200,-150,150);hole_rect=(150,250,-30,30);g=clip(lower,box)-clip(upper,box);cut=clip(hole_rect,box)-clip(hole_rect,upper,box)
        with localcontext() as ctx:ctx.prec=80;assert Decimal(r['quantity'])==(Decimal(g-cut)/1000000*22).quantize(Decimal(1).scaleb(-b['rounding_digits']),rounding=ROUND_HALF_UP)
        check(r)


@pytest.mark.parametrize('lo,hi,active',[('0','200',True),('0','199.999999999',False),('0','200.000000001',True),('200','600',False),('200.000000001','600',False)])
def test_platform_z_uses_lower_side_only_and_threshold_does_not_fill_hole(lo,hi,active):
    s=source(holes=[{**rectangle(x='0'),'vertical_range':interval(lo,hi)}]);r=calculate(body(s,bounds=('-300','300','-250','250')),s);assert bool(r['contact_ranges'][0]['apertures'])==active;assert r['quantity']=='3.960';check(r)


def test_three_levels_same_xy_are_distinct_platforms_and_zero_equal_tiers():
    s=source();third=copy.deepcopy(s['foundation_items'][0]['levels'][1]);third['parameters'].update(length='200',width='100',height='100');s['foundation_items'][0]['levels'].append(third);b=body(s,bounds=('-300','300','-250','250'));other=body(s,bounds=('-300','300','-250','250'),level=2)['contact_ranges'][0];other['name']='合成另一平台';b['contact_ranges'].append(other);r=calculate(b,s);assert r['horizontal_window_checks']==[] and r['quantity']=='6.094';check(r)
    s['foundation_items'][0]['levels'][1]['parameters'].update(length='600',width='500');r=calculate(body(s),s);assert r['quantity']=='0.000';check(r)


@pytest.mark.parametrize('fault',['circle','hole-end','full-before','full-after','overlap','basis','proof','missing-upper','manual-gross'])
def test_unsupported_geometry_and_duplicate_frame_guards(fault):
    s=source();b=body(s);d=b['contact_ranges'][0]
    if fault=='circle':s=source('stepped_circular');b=body(s)
    elif fault=='hole-end':d.update(face='hole_top',hole_index=1);d.pop('level_index')
    elif fault=='full-before':b['contact_ranges'].insert(0,{**contact(),'name':'完整平台'})
    elif fault=='full-after':b['contact_ranges'].append({**contact(),'name':'完整平台'})
    elif fault=='missing-upper':d['level_index']=2
    elif fault=='manual-gross':d['window_lower_area']='1'
    else:
        other=copy.deepcopy(d);other['name']='合成第二窗口';b['contact_ranges'].append(other)
        if fault=='overlap':other['plan_range']['parameters'].update(x_start='239.999999999',x_end='300')
        else:
            other['plan_range']['parameters'].update(x_start='240',x_end='300');other['plan_range']['coordinate_basis' if fault=='basis' else 'coordinate_evidence_id']='异基准'
    with pytest.raises(DomainError):calculate(b,s)


def setup(clients,pid=None):
    e,r,pid,s,_,_=step_setup(clients,pid)
    upload(e,pid,FRAME+'仅合成22个矩形台阶基础，下层600×500×200mm，上层400×300×400mm。H1圆孔Φ40中心(190,0)mm，H2矩形100×60中心(-200,0)mm，两孔Z0至200mm。体积阈值100m²不扣计量混凝土，物理孔仍扣模板，材料79.200kg保持。平台Z200mm窗口X180至240、Y负100至100mm，毛.008m²/个，孔仅扣H1超出上层X200mm的圆弓形；V2另加X负240至负180同Y窗口，毛.008扣H2 .0024m²/个，支模和总量修约3位仅合成，TEST-LEDGE-WINDOW非官方码。','foundation-step-window-basis.txt')
    for c in get(e,pid)['components']:approve(r,pid,'components',c)
    p=get(e,pid);s=next(c for c in p['components'] if c['id']==s['id']);eid=p['evidence'][-1]['id'];s['object_hash']=store.digest(s);item=s['foundation_items'][0];item['through_holes']['deduction_rule'].update(threshold='100',evidence_id=eid);item['through_holes']['holes']=[{**hole(eid=eid,x='190'),'vertical_range':interval('0','200',eid)},{**rectangle(eid=eid,x='-200',name='合成H2'),'vertical_range':interval('0','200',eid)}];s['basis_evidence_id']=eid;res=e.put(f'/api/v1/projects/{pid}/components/'+s['id'],json=s);assert res.status_code==200,res.text;approve(r,pid,'components',res.json());s=next(c for c in get(e,pid)['components'] if c['id']==s['id']);return e,r,pid,s,body(s,eid),eid


def test_api_preview_readonly_cross_record_level_claims_and_frames(clients):
    e,r,pid,s,b,eid=setup(clients);before=get(e,pid);post(e,pid,b,'/components/preview');assert get(e,pid)==before;c=post(e,pid,b).json();approve(r,pid,'components',c);post(e,pid,{**b,'id':c['id']},expected=409);other=body(s,eid,bounds=('-240','-180','-100','100'));post(e,pid,other);post(e,pid,step_body(s,eid),expected=409);bad=copy.deepcopy(other);bad['contact_ranges'][0]['plan_range']['coordinate_basis']='不同基准';post(e,pid,bad,expected=409)


@pytest.mark.parametrize('fault',['lower','upper','full','hole-window','hole-covered','hole-net','window-trace','covered-trace','bounds','ast','book'])
def test_tamper_rebuilt_at_review_gates(clients,fault):
    e,r,pid,s,b,eid=setup(clients);c=post(e,pid,b).json()
    if fault=='book':approve(r,pid,'components',c);post(e,pid,{},'/takeoffs/calculate')
    def forge(p):
        d=p['components'][-1]['preview']['contact_ranges'][0];h=d['apertures'][0]
        if fault in ('lower','upper','full'):d[{'lower':'window_lower_area','upper':'window_upper_area','full':'full_face_area'}[fault]]='0'
        elif fault in ('hole-window','hole-covered','hole-net'):h[{'hole-window':'window_hole_area','hole-covered':'covered_area','hole-net':'single_area'}[fault]]='0'
        elif fault in ('window-trace','covered-trace'):h['window_hole_curved_intersection' if fault=='window-trace' else 'curved_intersection']['terms'][0]['area']='0'
        elif fault=='bounds':d['plan_bounds']['x_start']='0'
        elif fault=='ast':d['formula_asts']['window_upper_area']={'constant':'0'}
        else:p['takeoffs'][-1]['quantity']='99'
    mutate(pid,forge);coll='takeoffs' if fault=='book' else 'components';approve(r,pid,coll,get(e,pid)[coll][-1],expected=409)


@pytest.mark.parametrize('gate',['preview','save','review','calculate','book_review'])
def test_original_integrity_five_gates(clients,gate):
    e,r,pid,s,b,eid=setup(clients)
    if gate in ('review','calculate','book_review'):c=post(e,pid,b).json()
    if gate in ('calculate','book_review'):approve(r,pid,'components',c)
    if gate=='book_review':t=post(e,pid,{},'/takeoffs/calculate').json()[-1]
    doc=get(e,pid)['documents'][-1];(store.ROOT/'originals'/doc['hash']).write_bytes(b'corrupt')
    if gate=='preview':post(e,pid,b,'/components/preview',409)
    elif gate=='save':post(e,pid,b,expected=409)
    elif gate=='review':approve(r,pid,'components',c,expected=409)
    elif gate=='calculate':post(e,pid,{},'/takeoffs/calculate',409)
    else:approve(r,pid,'takeoffs',t,expected=409)


def test_history_permissions_foreign_proof_concurrency_source_stale(clients):
    e,r,pid,s,b,eid=setup(clients);post(r,pid,b,expected=403);assert clients['other'].get(f'/api/v1/projects/{pid}').status_code==404;foreign=create(e)['id'];upload(e,foreign,'异项目窗口依据');bad=copy.deepcopy(b);bad['contact_ranges'][0]['plan_range']['evidence']['x_start']=get(e,foreign)['evidence'][-1]['id'];post(e,pid,bad,expected=404);c=post(e,pid,b).json();approve(r,pid,'components',c)
    for t in post(e,pid,{},'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
    post(e,pid,{},'/boq/generate');before=get(e,pid);old=before['components'][-1];b['object_hash']=store.digest(old);other=body(s,eid,bounds=('-240','-180','-100','100'))['contact_ranges'][0];other['name']='合成第二窗口';b['contact_ranges'].append(other);url=f'/api/v1/projects/{pid}/components/'+old['id'];res=e.put(url,json=b);assert res.status_code==200,res.text;after=get(e,pid);assert after['component_history'][-1]==old and after['boq']==before['boq'] and e.put(url,json=b).status_code==409;approve(r,pid,'components',res.json());s['object_hash']=store.digest(s);s['foundation_items'][0]['levels'][1]['parameters']['length']='420';res=e.put(f'/api/v1/projects/{pid}/components/'+s['id'],json=s);assert res.status_code==200,res.text;approve(r,pid,'components',res.json());assert get(e,pid)['components'][-1]['state']=='stale'


@pytest.mark.parametrize('fmt',['json','xlsx','pdf'])
def test_three_formats_old_quote_excluded_correct_curve_targets(clients,fmt):
    e,r,pid,_,q=priced_project(clients);e,r,pid,s,b,eid=setup(clients,pid);c=post(e,pid,b).json();approve(r,pid,'components',c)
    for t in post(e,pid,{},'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
    post(e,pid,{},'/boq/generate');url=f'/api/v1/projects/{pid}/exports';res=e.get(url,params=dict(format=fmt,unpriced=True));assert res.status_code==200,res.text;titles=set(STEP_WINDOW_TITLES)|set(WINDOW_TITLES)|{WINDOW_CURVED_TITLE}
    if fmt=='json':assert titles<={t[0] for t in res.json()['tables']}
    elif fmt=='xlsx':assert titles<=set(load_workbook(io.BytesIO(res.content)).sheetnames)
    else:assert all(t in ''.join(p.extract_text() for p in PdfReader(io.BytesIO(res.content)).pages) for t in titles)
    assert not any(t[0] in titles for t in e.get(url,params=dict(format='json',quote_id=q['id'])).json()['tables']);assert tables(dict(takeoffs=[]))==[]
    raw=tables(dict(takeoffs=[{**c['preview'],'name':c['name'],'measurement_basis':c['measurement_basis']}]))
    trace=next(rows for title,headers,rows in raw if title==WINDOW_CURVED_TITLE);assert {row[0].split(' / ')[-1] for row in trace}=={'窗口内全孔H1','上层窗口内孔H1'}
