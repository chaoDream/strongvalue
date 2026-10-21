import copy
import io
import json
from datetime import date,timedelta
import pytest
from openpyxl import load_workbook
from backend import store
from tests.test_workflow import priced_project,get,approve


def setup(clients):
    e,r,pid,comp,q=priced_project(clients);base=f'/api/v1/projects/{pid}'
    b=get(e,pid)['boq'][0]
    response=e.post(base+'/boq/'+b['id']+'/split',json={'parts':[{'name':'合成底板甲','quantity':'7'},{'name':'合成底板乙','quantity':'8'}]})
    assert response.status_code==200,response.text
    source,target=response.json()
    for b in [source,target]:approve(r,pid,'boq',b)
    p=get(e,pid);source,target=p['boq'];eid=p['evidence'][0]['id']
    body={'boq_id':source['id'],'object_hash':store.digest(source),'discipline':'建筑','quota_edition':'math-v1',
          'applicability':'合成底板，原消耗口径，非真实专业定额','reuse_basis':'仅本项目软件测试允许复用，不代表商业许可',
          'reuse_evidence_id':eid,'valid_from':(date.today()-timedelta(days=1)).isoformat(),'valid_until':(date.today()+timedelta(days=30)).isoformat()}
    response=e.post(base+'/catalog',json=body);assert response.status_code==200,response.text
    return e,r,pid,source,target,response.json(),body,q


def body_for(e,pid,b,item):
    p=get(e,pid);current=next(x for x in p['boq'] if x['id']==b['id']);item=next(x for x in p['catalog'] if x['id']==item['id'])
    return {'boq_hash':store.digest(current),'object_hash':store.digest(item),'discipline':'建筑','quota_edition':'math-v1'}


def endpoint(pid,b,item,action):return f"/api/v1/projects/{pid}/boq/{b['id']}/quota-candidates/{item['id']}/{action}"


def adopt(e,pid,b,item,extra=None):
    body=body_for(e,pid,b,item);preview=e.post(endpoint(pid,b,item,'preview'),json=body)
    assert preview.status_code==200,preview.text
    body.update(preview_hash=preview.json()['preview_hash'],reason='合成采用测试，条件与许可已核对',confirm_conditions=True,request_id=store.uid())
    body.update(extra or {})
    return e.post(endpoint(pid,b,item,'adopt'),json=body),body,preview.json()


def test_review_preview_adopt_snapshot_reports_and_idempotency(clients):
    e,r,pid,source,target,item,publish,old=setup(clients);base=f'/api/v1/projects/{pid}'
    assert e.post(endpoint(pid,target,item,'preview'),json=body_for(e,pid,target,item)).status_code==409
    assert not e.get(base+'/catalog-status').json()['items'][0]['available']
    approve(r,pid,'catalog',item)
    candidates=e.get(base+'/boq/'+target['id']+'/quota-candidates?discipline=建筑&quota_edition=math-v1').json()
    assert len(candidates)==1 and candidates[0]['adoptable'] and '不代表专业匹配' in candidates[0]['match_reason']
    before=get(e,pid);response,body,preview=adopt(e,pid,target,item)
    assert response.status_code==200,response.text
    assert preview['delta']=='0.00' and preview['after']['total']=='840.08' and '不含任何' in preview['notice']
    p=get(e,pid);b=p['boq'][1]
    assert p['status']=='BOQ_DRAFT' and b['state']=='draft' and b['quotas']==source['quotas'] and b['catalog_ref']['id']==item['id']
    for k in ['code','name','features','quantity','unit','node','takeoff_id']:assert b[k]==target[k]
    assert p['quotes']==before['quotes'] and p['boq'][0]==before['boq'][0]
    assert e.post(endpoint(pid,target,item,'adopt'),json=body).json()==response.json()
    changed={**body,'reason':'changed'}
    assert e.post(endpoint(pid,target,item,'adopt'),json=changed).status_code==409
    approve(r,pid,'boq',b)
    quote=e.post(base+'/quotes/calculate');assert quote.status_code==200,quote.text
    q=quote.json();assert q['snapshot']['quota_adoptions'][0]['candidate']['id']==item['id']
    assert e.post(base+'/quotes/calculate').json()['input_hash']==q['input_hash']
    j=e.get(base+'/exports?format=json').json();assert '定额候选采用台账' in [x[0] for x in j['tables']]
    x=e.get(base+'/exports?format=xlsx');assert '定额候选采用台账' in load_workbook(io.BytesIO(x.content)).sheetnames
    assert e.get(base+'/exports?format=pdf').content.startswith(b'%PDF')
    assert e.get(base+'/exports?format=zip').status_code==200
    assert e.get(base+'/exports?format=json&quote_id='+old['id']).json()['quote']==old


@pytest.mark.parametrize('change',[
    {'discipline':'电气'},{'quota_edition':'math-v2'},{'reuse_basis':''},{'applicability':''},{'reuse_evidence_id':'missing'},
    {'valid_from':'bad'},{'valid_until':'2000-01-01'},{'object_hash':'stale'}])
def test_publication_validation_atomic(clients,change):
    e,r,pid,s,t,item,body,old=setup(clients);before=get(e,pid)
    response=e.post(f'/api/v1/projects/{pid}/catalog',json={**body,**change})
    assert response.status_code in (409,422,404),response.text
    assert get(e,pid)==before


@pytest.mark.parametrize('change',[{'discipline':'电气'},{'quota_edition':'other'},{'boq_hash':'stale'},{'object_hash':'stale'}])
def test_preview_filters_and_hashes(clients,change):
    e,r,pid,s,t,item,body,old=setup(clients);approve(r,pid,'catalog',item)
    response=e.post(endpoint(pid,t,item,'preview'),json={**body_for(e,pid,t,item),**change})
    assert response.status_code==409,response.text
    assert e.get(f'/api/v1/projects/{pid}/boq/{t["id"]}/quota-candidates?quota_edition=bad').json()==[]


def test_source_change_revocation_and_detach_keep_old_quote(clients):
    e,r,pid,s,t,item,body,old=setup(clients);base=f'/api/v1/projects/{pid}';approve(r,pid,'catalog',item)
    response,body,preview=adopt(e,pid,t,item);assert response.status_code==200
    b=get(e,pid)['boq'][1];approve(r,pid,'boq',b)
    q=e.post(base+'/quotes/calculate').json();saved=copy.deepcopy(get(e,pid)['quotes'])
    approve(r,pid,'catalog',item,'reject')
    assert any('定额候选不可用' in x for x in get(e,pid)['blockers'])
    assert e.post(base+'/quotes/calculate').status_code==409
    approve(r,pid,'boq',b,expected=409)
    current=get(e,pid)['boq'][1]
    assert e.put(base+'/boq/'+t['id'],json={'detach_catalog':True,'object_hash':store.digest(current)}).status_code==422
    detached=e.put(base+'/boq/'+t['id'],json={'detach_catalog':True,'reason':'重新独立核查原始来源','object_hash':store.digest(current)})
    assert detached.status_code==200 and 'catalog_ref' not in detached.json()
    assert detached.json()['quotas']==current['quotas'] and detached.json()['state']=='draft'
    assert get(e,pid)['quotes']==saved and len(get(e,pid)['quota_adoptions'])==1


def test_stale_preview_due_to_date_and_source_change(clients):
    e,r,pid,s,t,item,publish,old=setup(clients);base=f'/api/v1/projects/{pid}';approve(r,pid,'catalog',item)
    b=body_for(e,pid,t,item);preview=e.post(endpoint(pid,t,item,'preview'),json=b).json()
    e.patch(base+'/settings',json={'pricing_date':(date.today()+timedelta(days=1)).isoformat(),'reason':'测试日期变化'})
    b.update(preview_hash=preview['preview_hash'],reason='测试',confirm_conditions=True,request_id=store.uid())
    assert e.post(endpoint(pid,t,item,'adopt'),json=b).status_code==409
    source=e.put(base+'/boq/'+s['id'],json={'features':'原清单条件修改'}).json()
    approve(r,pid,'boq',source)
    assert not e.get(base+'/catalog-status').json()['items'][0]['available']
    assert e.get(base+'/boq/'+t['id']+'/quota-candidates').json()==[]


def test_roles_tenant_confirmation_and_no_cycle(clients):
    e,r,pid,s,t,item,publish,old=setup(clients);base=f'/api/v1/projects/{pid}'
    assert r.post(base+'/catalog',json=publish).status_code==403
    assert clients['other'].get(base+'/catalog-status').status_code==404
    assert clients['other'].post(endpoint(pid,t,item,'preview'),json={}).status_code==404
    approve(e,pid,'catalog',item,expected=403);approve(r,pid,'catalog',item)
    response,body,preview=adopt(e,pid,t,item,{'confirm_conditions':False});assert response.status_code==422
    assert r.post(endpoint(pid,t,item,'adopt'),json=body).status_code==403
    response,body,preview=adopt(e,pid,t,item);assert response.status_code==200
    b=get(e,pid)['boq'][1];approve(r,pid,'boq',b)
    assert e.post(base+'/catalog',json={**publish,'boq_id':b['id'],'object_hash':store.digest(get(e,pid)['boq'][1])}).status_code==409
    assert e.post(endpoint(pid,s,item,'preview'),json=body_for(e,pid,s,item)).status_code==409


def test_missing_and_tampered_original_blocks_approval_and_adoption(clients):
    e,r,pid,s,t,item,publish,old=setup(clients);p=get(e,pid);doc=p['documents'][0];path=store.ROOT/'originals'/doc['hash'];raw=path.read_bytes()
    path.write_bytes(b'tampered')
    approve(r,pid,'catalog',item,expected=409)
    path.write_bytes(raw);approve(r,pid,'catalog',item)
    path.unlink()
    assert e.post(endpoint(pid,t,item,'preview'),json=body_for(e,pid,t,item)).status_code==409


def test_unpriced_target_preview_has_no_fabricated_zero_delta(clients):
    e,r,pid,s,t,item,publish,old=setup(clients);base=f'/api/v1/projects/{pid}';approve(r,pid,'catalog',item)
    e.put(base+'/boq/'+t['id'],json={'quotas':[]});approve(r,pid,'boq',get(e,pid)['boq'][0])
    response=e.post(endpoint(pid,t,item,'preview'),json=body_for(e,pid,t,item))
    assert response.status_code==200,response.text
    assert response.json()['before'] is None and response.json()['delta'] is None


def test_manual_edit_preserves_historical_adoption_and_removes_current_binding(clients):
    e,r,pid,s,t,item,publish,old=setup(clients);approve(r,pid,'catalog',item)
    response,body,preview=adopt(e,pid,t,item);assert response.status_code==200
    b=get(e,pid)['boq'][1];quotas=copy.deepcopy(b['quotas']);quotas[0]['coefficient']='1.01'
    result=e.put(f'/api/v1/projects/{pid}/boq/'+t['id'],json={'quotas':quotas,'object_hash':store.digest(b)})
    assert result.status_code==200 and 'catalog_ref' not in result.json()
    assert get(e,pid)['quota_adoptions'][0]['candidate']['quotas'][0]['coefficient']=='1'


def test_nonzero_preview_matches_recalculated_line_without_fees(clients):
    e,r,pid,s,t,item,publish,old=setup(clients);base=f'/api/v1/projects/{pid}'
    quotas=copy.deepcopy(t['quotas']);quotas[0]['coefficient']='1.1'
    result=e.put(base+'/boq/'+t['id'],json={'quotas':quotas});assert result.status_code==200
    for b in get(e,pid)['boq']:approve(r,pid,'boq',b)
    approve(r,pid,'catalog',item)
    response,body,preview=adopt(e,pid,t,item);assert response.status_code==200,response.text
    assert preview['before']['total']=='924.08' and preview['after']['total']=='840.08' and preview['delta']=='-84.00'
    approve(r,pid,'boq',get(e,pid)['boq'][1])
    q=e.post(base+'/quotes/calculate').json()
    assert q['lines'][1]['amount']==preview['after']['total'] and q['total']=='1622.40'


def test_target_changes_after_preview_reject_without_partial_apply(clients):
    e,r,pid,s,t,item,publish,old=setup(clients);base=f'/api/v1/projects/{pid}';approve(r,pid,'catalog',item)
    body=body_for(e,pid,t,item);preview=e.post(endpoint(pid,t,item,'preview'),json=body).json()
    e.put(base+'/boq/'+t['id'],json={'features':'目标特征变更，须重新确认专业匹配'})
    body.update(preview_hash=preview['preview_hash'],reason='测试',confirm_conditions=True,request_id=store.uid())
    before=get(e,pid)
    assert e.post(endpoint(pid,t,item,'adopt'),json=body).status_code==409
    assert get(e,pid)==before


def test_tampered_original_blocks_new_quote_and_report_approval(clients):
    e,r,pid,s,t,item,publish,old=setup(clients);base=f'/api/v1/projects/{pid}';approve(r,pid,'catalog',item)
    response,body,preview=adopt(e,pid,t,item);assert response.status_code==200
    approve(r,pid,'boq',get(e,pid)['boq'][1]);q=e.post(base+'/quotes/calculate').json()
    doc=get(e,pid)['documents'][0];(store.ROOT/'originals'/doc['hash']).write_bytes(b'changed original')
    assert e.post(base+'/quotes/calculate').status_code==409
    approve(r,pid,'quotes',q,expected=409)
    assert get(e,pid)['quotes'][-1]==q
