"""Synthetic edge-local surface partitions, no inferred construction scope."""
import copy
import io
import json
import random
from decimal import Decimal, localcontext
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pypdf import PdfReader
from backend import store
from backend.main import app
from backend.engine import DomainError
from backend.surface_takeoff import boundary, partition, calculate, snapshot
from tests.test_polygon_surface import geometry, setup as source_setup
from tests.test_polygon_takeoff import payload as polygon
from tests.test_surface_takeoff import payload as plain_payload
from tests.test_route_takeoff import evaluate
from tests.test_workflow import get, approve, upload, create, priced_project


def cuts(source,eid='test'):
    face=next(f for f in boundary(source)['faces'] if f.get('ring_index')==2 and f.get('edge_index')==2)
    return [{'face_id':face['face_id'],'u_cuts':[{'value':'200','unit':'cm','evidence_id':eid}],
             'v_cuts':[{'value':'500','unit':'mm','evidence_id':eid}]}]


def payload(source,eid='test',groups=None):
    groups=cuts(source,eid) if groups is None else groups
    result=partition(source,groups);b=plain_payload(source,eid,included=lambda f:False)
    b.update(name='合成斜孔局部防水（非真实工程）',purpose='waterproof',face_subdivisions=groups,
             measurement_basis='合成原边P2至P3长5m，局部S=0至2m、Z=0至0.5m计入，其余排除；仅软件验收',
             face_decisions=[{'face_id':f['face_id'],'included':bool(f.get('parent_index')) and Decimal(f['u_to'])<=2 and Decimal(f['v_to'])<=Decimal('.5'),
                              'basis':'仅合成：按原边起点距离及局部高度核定，不作真实工程依据','evidence_id':eid} for f in result['faces']])
    return b


def setup(clients,pid=None):
    e,r,pid,s,_,eid=source_setup(clients,pid)
    upload(e,pid,'合成侧面做法，非真实工程：三四五孔E2由P2(4,1)至P3(1,5)m，边长5m、厚2m、重复3。沿原边起点距离S=200cm和局部底面起高度Z=500mm分割；仅S0至2m、Z0至0.5m计入防水，其余全部排除；1m²×3=3m²。上下净面不可分区，局部Z非绝对标高。','edge-cut-basis.txt')
    for c in get(e,pid)['components']:approve(r,pid,'components',c)
    s=next(c for c in get(e,pid)['components'] if c['id']==s['id']);eid=get(e,pid)['evidence'][-1]['id']
    return e,r,pid,s,payload(s,eid),eid


def post(e,pid,path,body=None,expected=200):
    res=e.post(f'/api/v1/projects/{pid}'+path,json=body or {});assert res.status_code==expected,res.text;return res


def mutate(pid,action):
    with store.transaction() as c:
        p=json.loads(c.execute('select data from projects where id=?',(pid,)).fetchone()['data']);action(p);store.save(c,p)


def test_two_local_axes_exact_tiles_raw_units_decisions_and_ast():
    s=geometry();b=payload(s);before=copy.deepcopy(b);r=calculate(b,s)
    assert b==before and r['quantity']=='3.0' and r['rule_version']=='polygon-boundary-1.1'
    children=[f for f in r['surface_faces'] if f.get('parent_index')];assert len(children)==4 and len(r['surface_faces'])==12
    assert sorted(Decimal(f['area']) for f in children)==[Decimal('1'),Decimal('1.5'),Decimal('3'),Decimal('4.5')]
    assert {f['parent_index'] for f in children}=={8} and Decimal(r['surface_summary']['total_area'])==132
    assert [(c['original_value'],c['factor'],c['canonical_value'],c['coordinate_axis'],c['coordinate_system']) for c in r['surface_subdivisions']]==[('200','0.01','2','s','edge_local'),('500','0.001','0.5','z','edge_local')]
    assert r['surface_partition_checks'][0]['finite_decimal_tiling'] is True and Decimal(r['surface_partition_checks'][0]['area_difference'])==0
    assert all('原边XY' in f['position'] and '局部Z' in f['position'] for f in children)
    with localcontext() as ctx:
        ctx.prec=80
        for f in r['surface_faces']:assert evaluate(f['formula_ast'],r['dimensions'])==Decimal(f['area'])
        assert evaluate(r['formula_ast'],r['dimensions'])==Decimal(r['quantity'])


@pytest.mark.parametrize('which',['outer','hole'])
@pytest.mark.parametrize('axis',['u','v','both'])
def test_outer_and_hole_edges_allow_each_axis_and_repetition_once(which,axis):
    s=geometry();f=next(f for f in boundary(s)['faces'] if f['normal']==which)
    group={'face_id':f['face_id'],'u_cuts':[],'v_cuts':[]}
    if axis in ('u','both'):group['u_cuts']=[{'value':'1','unit':'m','evidence_id':'test'}]
    if axis in ('v','both'):group['v_cuts']=[{'value':'1','unit':'m','evidence_id':'test'}]
    result=partition(s,[group]);body=payload(s,groups=[group]);body['face_decisions']=[{**d,'included':True} for d in body['face_decisions']]
    out=calculate(body,s);assert Decimal(out['quantity'])==396
    assert len([v for v in result['faces'] if v.get('parent_index')])==(4 if axis=='both' else 2)


def test_random_local_rectangles_match_independent_grid_area_and_cover_parent():
    rng=random.Random(6001);s=geometry();f=next(f for f in boundary(s)['faces'] if f.get('ring_index')==2 and f.get('edge_index')==2)
    for _ in range(25):
        us=sorted(rng.sample(range(1,50),rng.randint(1,6)));vs=sorted(rng.sample(range(1,20),rng.randint(1,6)))
        groups=[{'face_id':f['face_id'],'u_cuts':[{'value':str(x*100),'unit':'mm','evidence_id':'test'} for x in us],
                 'v_cuts':[{'value':str(x*10),'unit':'cm','evidence_id':'test'} for x in vs]}]
        r=partition(s,groups);children=[c for c in r['faces'] if c.get('parent_index')]
        expected=[Decimal(b-a)/10*Decimal(d-c)/10 for a,b in zip([0]+us,us+[50]) for c,d in zip([0]+vs,vs+[20])]
        assert sorted(Decimal(c['area']) for c in children)==sorted(expected)
        assert sum(Decimal(c['area']) for c in children)==10 and Decimal(r['total_area'])==132


def test_cut_order_and_equivalent_units_keep_child_identity():
    s=geometry();a=cuts(s);a[0]['u_cuts']+=[{'value':'1000','unit':'mm','evidence_id':'test'}]
    b=copy.deepcopy(a);b[0]['u_cuts'].reverse();b[0]['v_cuts'][0].update(value='.5',unit='m')
    def geometry_rows(groups):return [{k:f.get(k) for k in ['face_id','parent_face_id','sub_index','area','u_from','u_to','v_from','v_to']} for f in partition(s,groups)['faces']]
    assert geometry_rows(a)==geometry_rows(b)


def test_unpartitioned_legacy_output_and_removal_are_identical():
    s=geometry();body=plain_payload(s);old=calculate(body,s)
    assert calculate({**body,'face_subdivisions':[]},s)==old and old['rule_version']=='polygon-boundary-1.0'
    partial=partition(s,cuts(s));ids={f['face_id'] for f in boundary(s)['faces']}
    assert all(f['face_id'] not in ids for f in partial['faces'] if f.get('parent_index'))
    assert partition(s,[])==boundary(s) and 'surface_partition_checks' not in old


def test_irrational_edge_precision_is_audited_without_adjusting_children():
    s=polygon();s.update(id='irrational',revision=1,thickness='.123',thickness_unit='m')
    f=next(f for f in boundary(s)['faces'] if f.get('ring_index')==2 and f.get('edge_index')==2)
    group={'face_id':f['face_id'],'u_cuts':[{'value':'.3','unit':'m','evidence_id':'test'}],
           'v_cuts':[{'value':'.05','unit':'m','evidence_id':'test'}]}
    r=partition(s,[group]);children=[f for f in r['faces'] if f.get('parent_index')]
    with localcontext() as ctx:
        ctx.prec=80
        for child in children:assert evaluate(child['formula_ast'],r['dimensions'])==Decimal(child['area'])
    with localcontext() as ctx:
        ctx.prec=200
        check=r['surface_partition_checks'][0]
        assert Decimal(check['children_area'])==sum((Decimal(c['area']) for c in children),Decimal(0))
        assert Decimal(check['area_difference'])==Decimal(check['children_area'])-Decimal(f['area'])
        assert abs(Decimal(check['area_difference']))<Decimal('1e-78')
        assert sum(((Decimal(c['u_to'])-Decimal(c['u_from']))*(Decimal(c['v_to'])-Decimal(c['v_from'])) for c in children),Decimal(0))==Decimal(f['edge_length'])*Decimal(f['height'])


@pytest.mark.parametrize('case',['null','not_list','fields','unknown','cap','duplicate_face','empty','too_many','outside_s','outside_z','negative','zero','end','duplicate_coordinate','unit','nan','precision','proof','extra_cut_field','cuts_not_list'])
def test_invalid_local_ranges_and_unsupported_caps_rejected(case):
    s=geometry();groups=cuts(s);g=groups[0];c=g['u_cuts'][0]
    if case=='null':groups=None
    elif case=='not_list':groups={}
    elif case=='fields':g['plane']='invented'
    elif case=='unknown':g['face_id']='unknown'
    elif case=='cap':g['face_id']=boundary(s)['faces'][0]['face_id']
    elif case=='duplicate_face':groups.append(copy.deepcopy(g))
    elif case=='empty':g['u_cuts']=[];g['v_cuts']=[]
    elif case=='too_many':g['u_cuts']*=9
    elif case=='outside_s':c['value']='600'
    elif case=='outside_z':g['v_cuts'][0]['value']='3000'
    elif case=='negative':c['value']='-100'
    elif case=='zero':c['value']='0'
    elif case=='end':c['value']='500'
    elif case=='duplicate_coordinate':g['u_cuts'].append({'value':'2','unit':'m','evidence_id':'test'})
    elif case=='unit':c['unit']='m²'
    elif case=='nan':c['value']='NaN'
    elif case=='precision':c['value']='.0000000000000000001'
    elif case=='proof':c['evidence_id']=' '
    elif case=='extra_cut_field':c['area']='99'
    else:g['u_cuts']='bad'
    with pytest.raises(DomainError):partition(s,groups)


def test_limits_eight_lines_each_axis_and_total_face_cap():
    s=geometry();faces=[f for f in boundary(s)['faces'] if f['geometry_type']=='polygon_side']
    def group(f):return {'face_id':f['face_id'],'u_cuts':[{'value':str(i/10),'unit':'m','evidence_id':'test'} for i in range(1,9)],
                         'v_cuts':[{'value':str(i/10),'unit':'m','evidence_id':'test'} for i in range(1,9)]}
    r=partition(s,[group(faces[0])]);assert len([f for f in r['faces'] if f.get('parent_index')])==81
    # More source edges make the hard 1000-face output bound reachable.
    from tests.test_polygon_takeoff import ring
    s['polygon_rings']=[ring('外轮廓',[(0,0),(30,0),(30,30),(0,30)])]+[ring('孔'+str(i),[(1+i*3,1),(2+i*3,1),(2+i*3,2),(1+i*3,2)],role='hole') for i in range(8)]
    groups=[group(f) for f in boundary(s)['faces'] if f['geometry_type']=='polygon_side']
    with pytest.raises(DomainError,match='1000'):partition(s,groups)


def test_readonly_partition_preview_and_review_boq_workflow(clients):
    e,r,pid,s,b,eid=setup(clients);before=get(e,pid)
    c=post(e,pid,'/surface-face-candidates/partition',b).json();assert get(e,pid)==before and len(c['faces'])==12
    post(r,pid,'/surface-face-candidates/partition',b,403)
    view=post(e,pid,'/components/preview',b).json();assert get(e,pid)==before and Decimal(view['quantity'])==3
    bad={**b,'face_decisions':b['face_decisions'][:-1]};post(e,pid,'/components',bad,422);assert get(e,pid)==before
    comp=post(e,pid,'/components',{**b,'source_geometry':{'fake':'ignored'}}).json();assert comp['source_geometry']==snapshot(s)
    approve(r,pid,'components',comp)
    for t in post(e,pid,'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
    assert any(v['unit']=='m²' and Decimal(v['quantity'])==3 for v in post(e,pid,'/boq/generate').json())


@pytest.mark.parametrize('fault',['missing','bytes','superseded','proof','source_coordinate','source_preview','cut_coordinate','cut_evidence','saved_check','saved_face'])
def test_actual_sources_cuts_and_saved_checks_revalidated(clients,fault):
    e,r,pid,s,b,eid=setup(clients);comp=post(e,pid,'/components',b).json()
    if fault in ('missing','bytes'):
        doc=get(e,pid)['documents'][-1];path=store.ROOT/'originals'/doc['hash']
        if fault=='missing':path.unlink()
        else:path.write_bytes(b'corrupt')
    else:
        def damage(p):
            target=p['components'][-1]
            if fault=='superseded':p['documents'][-1]['superseded']=True
            elif fault=='proof':p['evidence'][-1]['text']+=' changed'
            elif fault=='source_coordinate':p['components'][-2]['polygon_rings'][0]['points'][1]['parameters']['x']='7'
            elif fault=='source_preview':p['components'][-2]['preview']['quantity']='999'
            elif fault=='cut_coordinate':target['face_subdivisions'][0]['u_cuts'][0]['value']='100'
            elif fault=='cut_evidence':target['face_subdivisions'][0]['u_cuts'][0]['evidence_id']='missing'
            elif fault=='saved_check':target['preview']['surface_partition_checks'][0]['area_difference']='999'
            else:target['preview']['surface_faces'][-1]['area']='999'
        mutate(pid,damage)
    current=get(e,pid)['components'][-1]
    # Missing evidence produces 404; all other saved changes conflict or have incomplete decisions.
    res=r.post(f'/api/v1/projects/{pid}/reviews',json={'collection':'components','object_id':current['id'],'object_hash':store.digest(current),'decision':'approve','reason':'仅软件验收'})
    assert res.status_code in (404,409,422),res.text


def test_partition_revision_changes_only_current_object_and_precise_history(clients):
    e,r,pid,s,b,eid=setup(clients);comp=post(e,pid,'/components',b).json();approve(r,pid,'components',comp)
    for t in post(e,pid,'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
    post(e,pid,'/boq/generate');current=get(e,pid)['components'][-1];groups=cuts(s,eid);groups[0]['v_cuts'][0]['value']='1000'
    changed=payload(s,eid,groups);changed['face_decisions']=[{**d,'included':bool(f.get('parent_index')) and Decimal(f['u_to'])<=2 and Decimal(f['v_to'])<=1} for d,f in zip(changed['face_decisions'],partition(s,groups)['faces'])]
    changed['object_hash']=store.digest(current);res=e.put(f'/api/v1/projects/{pid}/components/'+comp['id'],json=changed);assert res.status_code==200,res.text
    p=get(e,pid);assert Decimal(res.json()['preview']['quantity'])==6 and res.json()['revision']==2 and p['component_history'][-1]==current
    assert p['components'][-2]['state']=='approved' and all(t['state']=='stale' for t in p['takeoffs'])
    assert any(v['unit']=='m²' and Decimal(v['quantity'])==3 for v in p['boq'])
    assert e.put(f'/api/v1/projects/{pid}/components/'+comp['id'],json=changed).status_code==409
    approve(r,pid,'components',res.json())
    for t in post(e,pid,'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
    assert any(v['unit']=='m²' and Decimal(v['quantity'])==6 for v in post(e,pid,'/boq/generate').json())


@pytest.mark.parametrize('field',['surface_partition_checks','surface_subdivisions'])
def test_calculation_book_partition_records_cannot_be_forged(clients,field):
    e,r,pid,s,b,eid=setup(clients);comp=post(e,pid,'/components',b).json();approve(r,pid,'components',comp)
    book=post(e,pid,'/takeoffs/calculate').json()[-1]
    def damage(p):
        row=p['takeoffs'][-1][field][0]
        if field=='surface_partition_checks':row['finite_decimal_tiling']=False
        else:row['coordinate_system']='global'
    mutate(pid,damage);approve(r,pid,'takeoffs',book,expected=409)


def test_source_change_rejects_old_partition_and_unknown_original_face(clients):
    e,r,pid,s,b,eid=setup(clients);comp=post(e,pid,'/components',b).json();approve(r,pid,'components',comp)
    old=get(e,pid)['components'][-2];res=e.put(f'/api/v1/projects/{pid}/components/'+s['id'],json={**copy.deepcopy(old),'object_hash':store.digest(old),'thickness':'3'});assert res.status_code==200,res.text
    approve(r,pid,'components',res.json());post(e,pid,'/surface-face-candidates/partition',b,409)
    source=get(e,pid)['components'][-2];new=payload(source,eid);assert new['source_component_hash']!=b['source_component_hash']
    new['face_subdivisions']=b['face_subdivisions'];post(e,pid,'/surface-face-candidates/partition',new,422)


def test_permissions_tenant_cross_project_and_cut_evidence_scope(clients):
    e,r,pid,s,b,eid=setup(clients)
    post(r,pid,'/components',b,403);post(clients['other'],pid,'/surface-face-candidates/partition',b,404);post(TestClient(app),pid,'/surface-face-candidates/partition',b,401)
    second=create(e)['id'];post(e,second,'/surface-face-candidates/partition',b,404)
    from tests.test_project_access import restrict
    restrict(clients['admin'],pid,[]);post(e,pid,'/surface-face-candidates/partition',b,404)


@pytest.mark.parametrize('fmt',['json','xlsx','pdf'])
def test_three_formats_keep_local_axes_and_precision_and_old_quote_isolation(clients,fmt):
    e,r,pid,_,quote=priced_project(clients);e,r,pid,s,b,eid=setup(clients,pid)
    comp=post(e,pid,'/components',b).json();approve(r,pid,'components',comp)
    for t in post(e,pid,'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
    base=f'/api/v1/projects/{pid}/exports';res=e.get(base,params={'format':fmt,'unpriced':True});assert res.status_code==200,res.text
    titles=['表面分区坐标','多边形侧面分区校验','逐面计量决定']
    if fmt=='json':
        data=res.json();assert all(any(t[0]==title for t in data['tables']) for title in titles)
        book=next(t for t in data['unpriced']['takeoffs'] if t.get('surface_summary'));assert book['surface_subdivisions'][0]['coordinate_system']=='edge_local' and book['rule_version']=='polygon-boundary-1.1'
    elif fmt=='xlsx':assert set(titles)<=set(load_workbook(io.BytesIO(res.content)).sheetnames)
    else:assert all(n in ''.join(p.extract_text() for p in PdfReader(io.BytesIO(res.content)).pages) for n in titles)
    old=e.get(base,params={'format':'json','quote_id':quote['id']}).json();assert old['quote']==quote and not any(t[0] in titles for t in old['tables'])
