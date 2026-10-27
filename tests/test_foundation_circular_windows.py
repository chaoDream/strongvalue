"""Synthetic disk windows: independent areas, exact boundary cases and workflow gates."""
import copy,io,math,random
from decimal import Decimal,localcontext
from fractions import Fraction
import pytest
from openpyxl import load_workbook
from pypdf import PdfReader
from backend import store
from backend.engine import DomainError
from backend.compound_takeoff import PI
from backend.foundation_horizontal_formwork import calculate
from backend.foundation_circular_windows import VERSION,METHOD,validate,separation
from backend.foundation_horizontal_windows import checks
from tests.test_foundation_horizontal_windows import source,window,FRAME
from tests.test_foundation_polygon_windows import body as pb,replay
from tests.test_foundation_hole_end_formwork import source as end_source
from tests.test_foundation_round_step_windows import adaptive
from tests.test_foundation_nested_windows import setup as old_setup,config as nested_config,block,box
from tests.test_foundation_multi_windows import config as multi_config,part
from tests.test_foundation_perforated_windows import config as perforated_config
from tests.test_foundation_simple_windows import C_SHAPE
from tests.test_foundation_openings import hole
from tests.test_foundation_voids import rectangle
from tests.test_foundation_recesses import interval
from tests.test_post_foundation import equal_body,fixed_body,fixed_added,equal_added,fixed_adjusted,equal_adjusted,combined
from tests.test_polygon_takeoff import post,mutate
from tests.test_workflow import get,approve,upload,create,priced_project


def config(eid='test',x='120',y='10',diameter='240',unit='mm'):
    p=dict(x=str(x),y=str(y),diameter=str(diameter))
    return dict(shape='circular',parameters=p,evidence={k:eid for k in p},unit=unit,coordinate_basis=FRAME,coordinate_evidence_id=eid,scope_confirmed=True)


def body(s,eid='test',face='bottom',**kwargs):
    b=pb(s,eid,face);b['contact_ranges'][0]['plan_range']=config(eid,**kwargs);return b


def disk(**kw):return dict(plan_range=config(**kw))


@pytest.mark.parametrize('shape',['rectangular','circular','stepped_rectangular','stepped_circular'])
@pytest.mark.parametrize('face',['top','bottom','hole_top','hole_bottom'])
@pytest.mark.parametrize('make',[equal_body,fixed_body,fixed_added,equal_added,fixed_adjusted,equal_adjusted,combined,lambda s:combined(s,equal=True)])
def test_all_shapes_ends_counts_immutable_and_ast_replay(shape,face,make):
    s=end_source(shape,make) if face.startswith('hole') else source(shape,make);b=body(s,face=face);before=copy.deepcopy((s,b));r=calculate(b,s);d=r['contact_ranges'][0]
    assert before==(s,b) and r['rule_version']==VERSION and d['circular_window'] and not d.get('polygon_window')
    with localcontext() as ctx:
        ctx.prec=80;assert Decimal(d['window_area'])==Decimal(PI)*Decimal('.12')**2
        assert Decimal(d['total_net_area'])==Decimal(d['single_net_area'])*Decimal(d['total_count'])
    assert r['horizontal_formwork_pi_constant']==PI;replay(r)


def independent_area(disks,boxes=(),excluded_disks=(),excluded_boxes=()):
    # Adaptive numerical integration in meters, independent of engine AST/antiderivatives.
    lo,hi=max(x-r for x,y,r in disks),min(x+r for x,y,r in disks)
    if lo>=hi:return 0.
    cuts={lo,hi,*[v for x,y,r in list(disks)+list(excluded_disks) for v in (x-r,x,x+r) if lo<v<hi],*[v for x,X,y,Y in list(boxes)+list(excluded_boxes) for v in (x,X) if lo<v<hi]}
    def span(x,ds,bs):
        lows=[];highs=[]
        for cx,cy,r in ds:
            if abs(x-cx)>=r:return None
            h=math.sqrt(max(0.,r*r-(x-cx)**2));lows.append(cy-h);highs.append(cy+h)
        for L,R,B,T in bs:
            if not L<x<R:return None
            lows.append(B);highs.append(T)
        low,high=max(lows),min(highs);return (low,high) if low<high else None
    def f(x):
        base=span(x,disks,boxes)
        if not base:return 0.
        cuts=[]
        for d in excluded_disks:
            s=span(x,[d],[])
            if s:cuts.append(s)
        cuts += [(y,Y) for L,R,y,Y in excluded_boxes if L<x<R]
        remain=[base]
        for l,h in cuts:
            new=[]
            for L,H in remain:
                if h<=L or l>=H:new.append((L,H))
                else:
                    if L<l:new.append((L,l))
                    if h<H:new.append((h,H))
            remain=new
        return sum(H-L for L,H in remain)
    return sum(adaptive(f,a,b,eps=2e-13) for a,b in zip(sorted(cuts),sorted(cuts)[1:]))


@pytest.mark.parametrize('shape',['rectangular','circular','stepped_rectangular','stepped_circular'])
@pytest.mark.parametrize('offset',[-350,-160,0,160,350])
def test_independent_numerical_disk_and_actual_rectangular_hole(shape,offset):
    s=source(shape);s['foundation_items'][0]['through_holes']['holes']=[{**rectangle(x='150',length='100',width='60'),'vertical_range':interval('0','150')}];r=calculate(body(s,x=offset,y='30',diameter='360'),s);d=r['contact_ranges'][0]
    disks=[(offset/1000,.03,.18)];boxes=[]
    if shape in ('rectangular','stepped_rectangular'):boxes=[(-.2,.2,-.25,.25)] if shape=='rectangular' else [(-.3,.3,-.25,.25)]
    else:disks.append((0.,0.,.3))
    expected=independent_area(disks,boxes,excluded_boxes=[(.1,.2,-.03,.03)])
    assert abs(float(d['single_net_area'])-expected)<3e-10;replay(r)


@pytest.mark.parametrize('shape',['stepped_rectangular','stepped_circular'])
@pytest.mark.parametrize('x',['0','170','280','700'])
def test_platform_lower_minus_upper_with_independent_integral(shape,x):
    s=source(shape);s['foundation_items'][0].pop('through_holes',None);r=calculate(body(s,face='step',x=x,y='0',diameter='220'),s);d=r['contact_ranges'][0]
    disks=[(float(x)/1000,0.,.11)]
    expected=independent_area(disks,[(-.3,.3,-.25,.25)],excluded_boxes=[(-.2,.2,-.15,.15)]) if shape=='stepped_rectangular' else independent_area(disks+[(0.,0.,.3)],excluded_disks=[(0.,0.,.2)])
    assert abs(float(d['single_net_area'])-expected)<3e-10;replay(r)


@pytest.mark.parametrize('x,diameter,factor',[('0','100','1'),('200','200','.5'),('0','400','1')])
def test_full_half_and_tangent_circle_on_rectangular_face(x,diameter,factor):
    s=source();s['foundation_items'][0].pop('through_holes',None);d=calculate(body(s,x=x,y='0',diameter=diameter),s)['contact_ranges'][0]
    expected=Decimal(PI)*(Decimal(diameter)/2000)**2*Decimal(factor)
    assert abs(Decimal(d['single_gross_area'])-expected)<Decimal('1e-27')


@pytest.mark.parametrize('unit,factor',[('mm','1'),('cm','.1'),('dm','.01'),('m','.001'),('km','.000001')])
def test_units_low_caller_precision_and_exact_replay(unit,factor):
    s=source();b=body(s);c=b['contact_ranges'][0]['plan_range'];c['unit']=unit;c['parameters']={k:str(Decimal(v)*Decimal(factor)) for k,v in c['parameters'].items()}
    expected=calculate(body(s),s)
    with localcontext() as ctx:ctx.prec=6;r=calculate(b,s)
    assert r['quantity']==expected['quantity'];replay(r)


@pytest.mark.parametrize('fault',['radius','area','extra','missing','missing_evidence','zero','negative','bool','nan','unit','confirm','frame','proof'])
def test_schema_evidence_and_positive_real_diameter(fault):
    c=config()
    if fault=='radius':c['parameters']['radius']=c['parameters'].pop('diameter')
    elif fault=='area':c['area']='100'
    elif fault=='extra':c['holes']=[]
    elif fault=='missing':del c['parameters']['y']
    elif fault=='missing_evidence':c['evidence']['diameter']=''
    elif fault in ('zero','negative','bool','nan'):c['parameters']['diameter']={'zero':'0','negative':'-1','bool':True,'nan':'NaN'}[fault]
    elif fault=='unit':c['unit']='px'
    elif fault=='confirm':c['scope_confirmed']=1
    elif fault=='frame':c['coordinate_basis']=''
    else:c['coordinate_evidence_id']=''
    with pytest.raises(DomainError):validate(c)


@pytest.mark.parametrize('gap',['-0.000000001','0','0.000000001'])
def test_exact_circle_circle_and_axis_rectangle_tangency(gap):
    a=disk(x='0',y='0',diameter='200');b=disk(x=str(Decimal(200)+Decimal(gap)),y='0',diameter='200')
    c=dict(plan_range=window('test',str(Decimal(100)+Decimal(gap)),'200','-100','100'))
    for other in [b,c]:
        for left,right in [(a,other),(other,a)]:
            result=separation(left,right);assert result['overlap']==(Decimal(gap)<0)
            if Decimal(gap)>=0:assert result['state']==('boundary_touch' if gap=='0' else 'separated') and result['method']==METHOD


def test_diagonal_3_4_5_corner_tangent_and_containment():
    a=disk(x='0',y='0',diameter='10');b=dict(plan_range=window('test','3','10','4','10'))
    assert separation(a,b)['state']=='boundary_touch'
    assert separation(a,disk(x='6',y='8',diameter='10'))['state']=='boundary_touch'
    assert separation(a,disk(x='0',y='0',diameter='1'))['overlap']
    assert separation(a,dict(plan_range=window('test','-1','1','-1','1')))['overlap']
    assert separation(a,dict(plan_range=window('test','-20','20','-20','20')))['overlap']


def test_concave_bay_hole_and_nested_island_in_both_orders():
    from tests.test_foundation_simple_windows import config as simple_config
    cases=[(simple_config(C_SHAPE),disk(x='40',y='0',diameter='40'),False),
      (multi_config([part(box(-100,100,-100,100),name='父块',holes=[box(-60,60,-60,60)]),part(box(200,300,200,300),name='远块')]),disk(x='0',y='0',diameter='100'),False),
      (nested_config(),disk(x='0',y='0',diameter='10'),True),
      (nested_config(),disk(x='40',y='0',diameter='20'),False),
      (nested_config(),disk(x='0',y='0',diameter='500'),True)]
    for c,a,expected in cases:
        for left,right in [(a,dict(plan_range=c)),(dict(plan_range=c),a)]:assert separation(left,right)['overlap']==expected


def test_120_seeded_exact_rect_distance_and_witness_replay():
    rng=random.Random(11201)
    for _ in range(120):
        x,y,r=[Fraction(rng.randint(-100,100)) for _ in range(3)];r=abs(r)+1;L,B=rng.randint(-100,100),rng.randint(-100,100);R,T=L+rng.randint(1,80),B+rng.randint(1,80)
        # Distance to solid rectangle, computed independently by axis clamping.
        dx=max(Fraction(L)-x,0,x-R);dy=max(Fraction(B)-y,0,y-T);expected=dx*dx+dy*dy<r*r
        a=disk(x=x,y=y,diameter=2*r);b=dict(plan_range=window('test',str(L),str(R),str(B),str(T)))
        q=separation(a,b);assert q['overlap']==expected
        if not expected:
            for w in q['edge_witnesses']:
                X,Y=Fraction(w['nearest_x']),Fraction(w['nearest_y']);cx,cy=Fraction(q['center_x']),Fraction(q['center_y'])
                assert Fraction(w['distance_squared'])==(X-cx)**2+(Y-cy)**2>=Fraction(w['limit_squared'])


@pytest.mark.parametrize('gap',['0','0.000000001'])
def test_all_fragments_covered_or_tiny_gap_not_zero(gap):
    v=str(Decimal(25)+Decimal(gap));hs=[{**hole(x='0',diameter='80'),'vertical_range':interval('100','300')},{**rectangle(name='H2',x='-'+v,length='50',width='100'),'vertical_range':interval('300','500')},{**rectangle(name='H3',x=v,length='50',width='100'),'vertical_range':interval('300','500')}]
    s=end_source(holes=hs);r=calculate(body(s,face='hole_top',x='0',y='0',diameter='60'),s);d=r['contact_ranges'][0]
    assert ('polygon_zero_contact' in d)==(gap=='0') and (Decimal(d['single_net_area'])==0)==(gap=='0');replay(r)


def setup(clients,pid=None,zero=False):
    e,r,pid,s,b,eid=old_setup(clients,pid,zero);b['contact_ranges'][0]['plan_range']=config(eid,x='0' if zero else '190',y='0',diameter='40' if zero else '60');return e,r,pid,s,b,eid


def test_workflow_roles_foreign_evidence_history_concurrency_stale(clients):
    e,r,pid,s,b,eid=setup(clients);before=get(e,pid);post(e,pid,b,'/components/preview');assert get(e,pid)==before
    post(r,pid,b,expected=403);assert clients['other'].get(f'/api/v1/projects/{pid}').status_code==404
    foreign=create(e)['id'];upload(e,foreign,'异项目合成');bad=copy.deepcopy(b);bad['contact_ranges'][0]['plan_range']['evidence']['diameter']=get(e,foreign)['evidence'][-1]['id'];post(e,pid,bad,expected=404)
    c=post(e,pid,b).json();approve(r,pid,'components',c);c=next(x for x in get(e,pid)['components'] if x['id']==c['id']);b['object_hash']=store.digest(c);b['contact_ranges'][0]['plan_range']['parameters']['diameter']='58'
    res=e.put(f'/api/v1/projects/{pid}/components/'+c['id'],json=b);assert res.status_code==200,res.text;assert e.put(f'/api/v1/projects/{pid}/components/'+c['id'],json=b).status_code==409
    assert get(e,pid)['component_history'][-1]==c;approve(r,pid,'components',res.json())
    s['object_hash']=store.digest(s);s['foundation_items'][0]['levels'][1]['parameters']['diameter']='420';res=e.put(f'/api/v1/projects/{pid}/components/'+s['id'],json=s);assert res.status_code==200;approve(r,pid,'components',res.json());assert next(x for x in get(e,pid)['components'] if x['id']==c['id'])['state']=='stale'


def test_same_and_cross_record_repeats_frames_and_circle_in_residual_hole(clients):
    e,r,pid,s,b,eid=old_setup(clients);post(e,pid,b)
    other=copy.deepcopy(b);other['name']='合成父孔残余圆窗';other['contact_ranges'][0]['plan_range']=config(eid,x='184',y='0',diameter='4');post(e,pid,other)
    bad=copy.deepcopy(other);bad['contact_ranges'][0]['plan_range']['parameters']['x']='194';bad['contact_ranges'][0]['plan_range']['parameters']['y']='10';before=get(e,pid);post(e,pid,bad,expected=409);assert get(e,pid)==before
    bad=copy.deepcopy(other);bad['contact_ranges'][0]['plan_range']['coordinate_basis']='其他XY';post(e,pid,bad,expected=409)
    fresh=copy.deepcopy(other);fresh['name']='重复圆';fresh['contact_ranges']*=2;post(e,pid,fresh,expected=422)


@pytest.mark.parametrize('fault',['input','normalized','quantity','trace','proof','expression'])
def test_tamper_recalculation_rejected(clients,fault):
    e,r,pid,s,b,eid=setup(clients,zero=True);c=post(e,pid,b).json()
    def forge(p):
        c=p['components'][-1];d=c['preview']['contact_ranges'][0]
        if fault=='input':c['contact_ranges'][0]['plan_range']['parameters']['diameter']='41'
        elif fault=='normalized':d['window_circle']['diameter']='999'
        elif fault=='quantity':d['single_net_area']='1'
        elif fault=='trace':d['polygon_net_trace']['partitions']=[]
        elif fault=='proof':d['polygon_zero_contact']['net_trace']['partitions'][0]['remaining']=[{}]
        else:d['window_expression_definitions'][0]['value']='999'
    mutate(pid,forge);before=get(e,pid);approve(r,pid,'components',c,expected=409);assert get(e,pid)==before


@pytest.mark.parametrize('gate',['preview','save','review','calculate','book_review'])
def test_circle_diameter_original_bytes_five_gates(clients,gate):
    e,r,pid,s,b,eid=setup(clients);upload(e,pid,'合成圆形支模窗口直径及圆心，非真实工程','circle-window-basis.txt')
    for c in get(e,pid)['components']:approve(r,pid,'components',c)
    b['source_component_hash']=e.get(f'/api/v1/projects/{pid}/foundation-horizontal-formwork-candidates',params={'component_id':s['id']}).json()['source_component_hash'];b['contact_ranges'][0]['plan_range']['evidence']['diameter']=get(e,pid)['evidence'][-1]['id']
    if gate in ('review','calculate','book_review'):c=post(e,pid,b).json()
    if gate in ('calculate','book_review'):approve(r,pid,'components',c)
    if gate=='book_review':t=post(e,pid,{},'/takeoffs/calculate').json()[-1]
    doc=get(e,pid)['documents'][-1];(store.ROOT/'originals'/doc['hash']).write_bytes(b'corrupted-synthetic-circle-window-original')
    if gate=='preview':post(e,pid,b,'/components/preview',409)
    elif gate=='save':post(e,pid,b,expected=409)
    elif gate=='review':approve(r,pid,'components',c,expected=409)
    elif gate=='calculate':post(e,pid,{},'/takeoffs/calculate',409)
    else:approve(r,pid,'takeoffs',t,expected=409)


@pytest.mark.parametrize('fmt',['json','xlsx','pdf'])
@pytest.mark.parametrize('zero',[False,True])
def test_three_formats_and_old_quote_isolation(clients,fmt,zero):
    from backend.foundation_circular_window_reports import TITLES
    e,r,pid,_,q=priced_project(clients);e,r,pid,s,b,eid=setup(clients,pid,zero);c=post(e,pid,b).json();approve(r,pid,'components',c)
    for t in post(e,pid,{},'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
    post(e,pid,{},'/boq/generate');url=f'/api/v1/projects/{pid}/exports';res=e.get(url,params=dict(format=fmt,unpriced=True));assert res.status_code==200,res.text
    titles=set(TITLES[:3] if zero else TITLES[:2])
    if fmt=='json':assert titles<={t[0] for t in res.json()['tables']}
    elif fmt=='xlsx':assert titles<=set(load_workbook(io.BytesIO(res.content)).sheetnames)
    else:assert all(any(title in page.extract_text() for page in PdfReader(io.BytesIO(res.content)).pages) for title in titles)
    assert not any(t[0] in titles for t in e.get(url,params=dict(format='json',quote_id=q['id'])).json()['tables'])

@pytest.mark.parametrize('offset',['599.999999999','600','600.000000001'])
def test_actual_circle_tiny_positive_lens_is_not_tolerance_zero(offset):
    s=source('circular');s['foundation_items'][0].pop('through_holes',None);r=calculate(body(s,x=offset,y='0',diameter='600'),s);d=r['contact_ranges'][0]
    assert (Decimal(d['single_net_area'])>0)==(Decimal(offset)<600)
    assert r['quantity']=='0.000';assert 'polygon_zero_contact' not in d;replay(r)


def test_coincident_window_and_source_full_circle():
    s=source('circular');s['foundation_items'][0].pop('through_holes',None);r=calculate(body(s,x='0',y='0',diameter='600'),s)
    with localcontext() as ctx:
        ctx.prec=80;assert Decimal(r['contact_ranges'][0]['single_gross_area'])==Decimal(PI)*Decimal('.09')
    replay(r)
