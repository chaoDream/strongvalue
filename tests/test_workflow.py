import io
import json
from datetime import date, timedelta
import pytest
from openpyxl import load_workbook
from pypdf import PdfReader
from backend import store


def create(c, scope=None):
    return c.post('/api/v1/projects', json={'name': '几何测试项目（非真实工程）', 'scope': scope or ['建筑'], 'require_golden_baseline': False}).json()


def get(c, pid):
    return c.get('/api/v1/projects/'+pid).json()


def approve(c, pid, coll, obj, decision='approve', expected=200):
    hashes = c.get(f'/api/v1/projects/{pid}/hashes').json()
    r = c.post(f'/api/v1/projects/{pid}/reviews', json={'collection': coll, 'object_id': obj['id'],
        'object_hash': hashes[coll][obj['id']], 'decision': decision, 'reason': '测试核查依据（非专业签审）'})
    assert r.status_code == expected, r.text
    return r


def upload(c, pid, text='构件底板：长度10m，宽度5m，厚度0.3m；无扣减。', name='fixture.txt', specialty='建筑'):
    r = c.post(f'/api/v1/projects/{pid}/documents', data={'specialty': specialty}, files={'file': (name, text.encode(), 'text/plain')})
    assert r.status_code == 200, r.text
    return r.json()


def setup_boq(clients):
    e, r = clients['estimator'], clients['reviewer']
    p = create(e)
    pid = p['id']
    upload(e, pid)
    p = get(e, pid)
    eid = p['evidence'][0]['id']
    base = f'/api/v1/projects/{pid}'
    comp = e.post(base+'/components', json={'name': '测试底板', 'kind': 'slab',
        'parameters': {'length': '10', 'width': '5', 'thickness': '.3', 'deduction': '0'},
        'evidence': {k: eid for k in ['length', 'width', 'thickness', 'deduction']},
        'measurement_basis': '仅数学单测，无专业计量效力'}).json()
    approve(r, pid, 'components', comp)
    response = e.post(base+'/takeoffs/calculate')
    assert response.status_code == 200, response.text
    takeoff = response.json()[0]
    approve(r, pid, 'takeoffs', takeoff)
    b = e.post(base+'/boq/generate').json()[0]
    return e, r, pid, comp, takeoff, b, eid


def priced_project(clients):
    e, r, pid, comp, t, b, eid = setup_boq(clients)
    base = f'/api/v1/projects/{pid}'
    resource = {'name': '数学测试资源', 'type': '材料', 'unit': 'm³', 'consumption': '1.05', 'price': '100.005',
        'source': '测试报价单', 'license': '仅用于测试', 'period': '2026-10', 'valid_until': (date.today()+timedelta(days=365)).isoformat(),
        'region': '河南省', 'tax_mode': '不含税', 'evidence_id': eid}
    b = e.put(base+'/boq/'+b['id'], json={'code': 'TEST-ONLY', 'features': '测试编码，非真实定额',
        'quotas': [{'code': 'TEST-QUOTA', 'edition': 'math-v1', 'source': '测试数据', 'license': 'test', 'evidence_id': eid, 'resources': [resource], 'coefficient': '1'}]}).json()
    approve(r, pid, 'boq', b)
    rules = e.put(base+'/rules', json={'name': '数学测试规则', 'version': 'v1', 'region': '河南省', 'standard': '测试标准',
        'source': '人工单测', 'license': 'test', 'evidence_id': eid, 'valid_from': (date.today()-timedelta(days=1)).isoformat(),
        'valid_until': (date.today()+timedelta(days=365)).isoformat(), 'fees': [{'name': '数学测试费', 'rate': '.03', 'basis': 'subtotal', 'source': 'test'}]}).json()
    approve(r, pid, 'rules', rules)
    coverage = e.post(base+'/coverage', json={'specialty': '建筑', 'reason': '单测覆盖', 'evidence_id': eid}).json()
    approve(r, pid, 'coverage', coverage)
    response = e.post(base+'/quotes/calculate')
    assert response.status_code == 200, response.text
    return e, r, pid, comp, response.json()


def test_tenant_and_role_isolation(clients):
    p = create(clients['admin'])
    pid = p['id']
    assert clients['other'].get('/api/v1/projects/'+pid).status_code == 404
    assert clients['other'].get('/api/v1/projects/'+pid+'/audit').status_code == 404
    assert clients['reviewer'].post('/api/v1/projects', json={'name': 'forbidden'}).status_code == 403
    r = clients['admin'].post('/api/v1/projects/'+pid+'/reviews', json={})
    assert r.status_code == 403


def test_fact_conflicts_and_specialty_gaps(clients):
    c, r = clients['admin'], clients['reviewer']
    pid = create(c, ['建筑', '结构', '电气'])['id']
    upload(c, pid, '规模500m³/日；工期45日历天；施工中标价2,982,485.80元；监理中标报价1.49%')
    p = get(c, pid)
    values = {f['key']: f['value'] for f in p['facts']}
    assert values == {'capacity': '500', 'duration': '45', 'construction_amount': '2982485.80', 'supervision_rate': '1.49'}
    assert next(g for g in p['gaps'] if g['specialty'] == '结构')['status'] == '缺失'
    upload(c, pid, '规模600m³/日', name='conflict.txt')
    p = get(c, pid)
    fs = [f for f in p['facts'] if f['key'] == 'capacity']
    assert all(f['state'] == 'conflict' for f in fs)
    approve(r, pid, 'facts', fs[0], expected=409)
    approve(r, pid, 'facts', fs[1], 'reject')
    approve(r, pid, 'facts', fs[0])


def test_original_immutability_and_dedup(clients):
    c = clients['admin']
    pid = create(c)['id']
    a = upload(c, pid, '旧图')
    assert upload(c, pid, '旧图')['duplicate']
    b = upload(c, pid, '新图')
    assert b['document']['version'] == 2
    p = get(c, pid)
    assert len(p['documents']) == 2
    old = c.get(f'/api/v1/projects/{pid}/documents/{a["document"]["id"]}/file')
    assert old.content.decode() == '旧图'
    assert p['documents'][0]['superseded']
    assert clients['other'].get(f'/api/v1/projects/{pid}/documents/{a["document"]["id"]}/file').status_code == 404


def test_approval_required_before_calculation_and_stale_hash(clients):
    e,r,pid,comp,t,b,eid = setup_boq(clients)
    assert e.post(f'/api/v1/projects/{pid}/quotes/calculate').status_code == 409
    hashes = r.get(f'/api/v1/projects/{pid}/hashes').json()
    e.put(f'/api/v1/projects/{pid}/boq/{b["id"]}', json={'features': 'changed'})
    response = r.post(f'/api/v1/projects/{pid}/reviews', json={'collection':'boq','object_id':b['id'],
        'object_hash':hashes['boq'][b['id']],'decision':'approve','reason':'test'})
    assert response.status_code == 409


def test_complete_quote_export_and_dimension_invalidation(clients):
    e,r,pid,comp,q = priced_project(clients)
    assert q['subtotal'] == '1575.15'
    assert q['total'] == '1622.40'
    approve(r, pid, 'quotes', q)
    base = f'/api/v1/projects/{pid}'
    for fmt in ['pdf', 'xlsx', 'json']:
        response = e.get(base+'/exports', params={'format': fmt, 'quote_id': q['id']})
        assert response.status_code == 200, response.text[:200]
        assert response.headers.get('x-artifact-hash')
        if fmt == 'xlsx':
            wb = load_workbook(io.BytesIO(response.content), data_only=True)
            assert len(wb.sheetnames) == 10
            assert '审核成果' in wb.worksheets[0]['A1'].value
            assert wb['分部分项清单']['G3'].value == 1575.15
        if fmt == 'pdf':
            pdf = PdfReader(io.BytesIO(response.content))
            assert len(pdf.pages) >= 10
            assert '1622.40' in ''.join(page.extract_text() for page in pdf.pages)
    old_snapshot = json.dumps(q, sort_keys=True)
    comp = {**comp, 'parameters': {**comp['parameters'], 'thickness': '.4'}}
    assert e.put(base+'/components/'+comp['id'], json=comp).status_code == 200
    p = get(e, pid)
    assert q['id'] in p['stale_quotes']
    assert p['report_approval'] is None
    assert all(x['state'] == 'stale' for x in p['takeoffs']+p['boq'])
    assert json.dumps(p['quotes'][0], sort_keys=True) == old_snapshot
    approve(r,pid,'quotes',q,expected=409)
    data = e.get(base+'/exports', params={'format':'json', 'quote_id':q['id']}).json()
    assert '草稿' in data['status']


def test_price_expiry_blocks_and_audit_chain(clients):
    e,r,pid,comp,q = priced_project(clients)
    p = get(e,pid)
    b = p['boq'][0]
    b['quotas'][0]['resources'][0]['valid_until'] = '2000-01-01'
    e.put(f'/api/v1/projects/{pid}/boq/{b["id"]}',json={'quotas':b['quotas']})
    approve(r,pid,'boq',get(e,pid)['boq'][0])
    response = e.post(f'/api/v1/projects/{pid}/quotes/calculate')
    assert response.status_code == 409
    assert '价格过期' in response.text
    events = list(reversed(e.get(f'/api/v1/projects/{pid}/audit').json()))
    prev = ''
    for event in events:
        sha = event.pop('hash')
        assert event['previous_hash'] == prev
        assert store.digest(event) == sha
        prev = sha
    with store.connect() as c:
        with pytest.raises(Exception, match='immutable'):
            c.execute('DELETE FROM audit')


def test_missing_scope_and_rejected_takeoff_block_report(clients):
    e,r,pid,comp,q = priced_project(clients)
    p = get(e,pid)
    approve(r,pid,'takeoffs',p['takeoffs'][0],'reject')
    approve(r,pid,'quotes',q,expected=409)
    assert get(e,pid)['report_approval'] is None


def test_invalid_upload_and_unverified_exchange(clients):
    c = clients['admin']
    pid = create(c)['id']
    base = f'/api/v1/projects/{pid}'
    assert c.post(base+'/documents', data={'specialty':'建筑'}, files={'file':('fake.pdf',b'not-pdf')}).status_code == 422
    assert c.post(base+'/documents', data={'specialty':'建筑'}, files={'file':('evil.exe',b'MZ')}).status_code == 422
    assert c.get(base+'/exports?format=xml').status_code == 422
    assert all(not x['can_export'] for x in c.get('/api/v1/integrations/capabilities').json())


def test_unpriced_export_and_spreadsheet_injection(clients):
    c = clients['admin']
    pid = create(c)['id']
    upload(c,pid,'=HYPERLINK("https://example.invalid")')
    res=c.get(f'/api/v1/projects/{pid}/exports?format=xlsx')
    assert res.status_code==200
    w=load_workbook(io.BytesIO(res.content),data_only=False)
    assert '草稿' in w.worksheets[0]['A1'].value
    assert w['依据与证据']['F3'].data_type=='s'
    assert w['依据与证据']['F3'].value.startswith('=HYPERLINK')


def test_cross_origin_mutation_rejected(clients):
    response=clients['admin'].post('/api/v1/projects',json={'name':'attack'},headers={'Origin':'https://other.invalid'})
    assert response.status_code==403


def test_historic_artifact_survives_source_replacement(clients):
    e,r,pid,comp,q = priced_project(clients)
    approve(r,pid,'quotes',q)
    base=f'/api/v1/projects/{pid}'
    original=e.get(base+'/exports?format=pdf').content
    artifacts=e.get(base+'/artifacts').json()
    assert len(artifacts)==1 and artifacts[0]['approved']
    aid=artifacts[0]['id']
    upload(e,pid,'全新图纸文字')
    historic=e.get(base+'/artifacts/'+aid+'/file')
    assert historic.content==original
    assert clients['other'].get(base+'/artifacts/'+aid+'/file').status_code==404
    assert e.get(base+'/artifacts').json()[0]['hash']==artifacts[0]['hash']


def test_recovery_is_retryable_and_idempotent(clients):
    from backend.recovery import recover_interrupted_jobs
    c=clients['admin']
    pid=create(c)['id']
    upload(c,pid)
    with store.transaction() as connection:
        p=get(c,pid)
        p['documents'][0]['parsed']=False
        p['jobs'][0]['state']='running'
        store.save(connection,p)
    recover_interrupted_jobs()
    p=get(c,pid)
    assert p['jobs'][0]['state']=='failed'
    response=c.post(f'/api/v1/projects/{pid}/jobs/{p["jobs"][0]["id"]}/retry')
    assert response.status_code==200
    p=get(c,pid)
    assert p['jobs'][0]['state']=='completed'
    count=len(p['evidence'])
    from backend.main import parse_job
    parse_job(pid,p['documents'][0]['id'],p['jobs'][0]['id'],{'tenant':'local','username':'admin','role':'admin'})
    assert len(get(c,pid)['evidence'])==count


def test_long_evidence_is_paginated_without_silent_truncation(clients):
    c=clients['admin']
    pid=create(c)['id']
    upload(c,pid,'长'*6000)
    response=c.get(f'/api/v1/projects/{pid}/exports?format=pdf')
    assert response.status_code==200
    pdf=PdfReader(io.BytesIO(response.content))
    text=''.join(page.extract_text() for page in pdf.pages)
    assert text.count('长') >= 6000
    assert len(pdf.pages)>10
