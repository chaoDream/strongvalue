"""Synthetic two-center windows: independent integrals, exact topology and workflow gates."""
import copy,io,random,math
from decimal import Decimal,localcontext
from fractions import Fraction
import pytest
from openpyxl import load_workbook
from pypdf import PdfReader
from backend import store
from backend.compound_takeoff import PI
from backend.engine import DomainError
from backend.foundation_horizontal_formwork import calculate
from backend.foundation_eccentric_windows import VERSION,METHOD,validate,separation,containment
from backend.foundation_horizontal_windows import checks
from tests.test_foundation_circular_windows import body as cb,config as cc,independent_area,replay as old_replay
from tests.test_foundation_annular_windows import config as ac,setup as old_setup
from tests.test_foundation_horizontal_windows import source,window
from tests.test_foundation_hole_end_formwork import source as end_source
from tests.test_foundation_nested_windows import config as nc,block,box
from tests.test_foundation_multi_windows import config as mc,part
from tests.test_foundation_perforated_windows import config as pc
from tests.test_foundation_simple_windows import config as sc,C_SHAPE
from tests.test_foundation_openings import hole
from tests.test_foundation_voids import rectangle
from tests.test_foundation_recesses import interval
from tests.test_post_foundation import equal_body,fixed_body,fixed_added,equal_added,fixed_adjusted,equal_adjusted,combined
from tests.test_polygon_takeoff import post,mutate
from tests.test_workflow import get,approve,upload,create,priced_project
from tests.test_route_takeoff import evaluate


def config(eid='test',x='120',y='10',outer='240',inner_x='140',inner_y='20',inner='80',unit='mm'):
    c=ac(eid,x,y,outer,inner,unit);c['shape']='eccentric_circular';c['parameters'].update(inner_x=str(inner_x),inner_y=str(inner_y));c['evidence'].update(inner_x=eid,inner_y=eid);return c


def body(s,eid='test',face='bottom',**kw):
    b=cb(s,eid,face);b['contact_ranges'][0]['plan_range']=config(eid,**kw);return b


def punctured(**kw):return dict(plan_range=config(**kw))
def disk(**kw):return dict(plan_range=cc(**kw))


def replay(r):
    old_replay(r)
    for d in r['contact_ranges']:
        if not d.get('eccentric_circular_window'):continue
        proof=d['window_inner_containment']
        for k,n in proof['formula_asts'].items():assert Fraction(evaluate(n,r['dimensions']))==Fraction(proof[k])
        assert Fraction(proof['center_distance_squared'])<Fraction(proof['radius_gap_squared'])


@pytest.mark.parametrize('shape',['rectangular','circular','stepped_rectangular','stepped_circular'])
@pytest.mark.parametrize('face',['top','bottom','hole_top','hole_bottom'])
@pytest.mark.parametrize('make',[equal_body,fixed_body,fixed_added,equal_added,fixed_adjusted,equal_adjusted,combined,lambda s:combined(s,equal=True)])
def test_shapes_ends_counts_immutable_and_all_ast_replay(shape,face,make):
    s=end_source(shape,make) if face.startswith('hole') else source(shape,make);b=body(s,face=face);before=copy.deepcopy((s,b));r=calculate(b,s);d=r['contact_ranges'][0]
    assert before==(s,b) and r['rule_version']==VERSION and d['eccentric_circular_window'] and not d.get('annular_window')
    with localcontext() as ctx:
        ctx.prec=80;assert Decimal(d['window_area'])==Decimal(PI)*Decimal('.0128');assert Decimal(d['total_net_area'])==Decimal(d['single_net_area'])*Decimal(d['total_count'])
    assert r['source_evidence']['horizontal_formwork.1.window_inner_basis']=='test';replay(r)


@pytest.mark.parametrize('shape',['rectangular','circular','stepped_rectangular','stepped_circular'])
@pytest.mark.parametrize('face',['top','bottom','hole_bottom','hole_top'])
def test_separate_actual_offset_inner_disk_subtraction_matches_full_gross_cut_net(shape,face):
    s=end_source(shape) if face.startswith('hole') else source(shape);r=calculate(body(s,face=face),s);a=calculate(cb(s,face=face,x='120',y='10',diameter='240'),s);b=calculate(cb(s,face=face,x='140',y='20',diameter='80'),s)
    with localcontext() as ctx:
        ctx.prec=80
        for k in ['single_gross_area','single_opening_area','single_net_area','total_net_area']:
            assert abs(Decimal(r['contact_ranges'][0][k])-(Decimal(a['contact_ranges'][0][k])-Decimal(b['contact_ranges'][0][k])))<Decimal('1e-73')
    replay(r)


@pytest.mark.parametrize('shape',['rectangular','circular','stepped_rectangular','stepped_circular'])
@pytest.mark.parametrize('offset',[-350,-160,0,160,350])
def test_20_independent_numeric_offset_hole_face_and_physical_hole_intersections(shape,offset):
    s=source(shape);s['foundation_items'][0]['through_holes']['holes']=[{**rectangle(x='150',length='100',width='60'),'vertical_range':interval('0','150')}]
    r=calculate(body(s,x=offset,y='30',outer='360',inner_x=offset+45,inner_y='5',inner='160'),s);d=r['contact_ranges'][0];disks=[(offset/1000,.03,.18)];boxes=[]
    if shape in ('rectangular','stepped_rectangular'):boxes=[(-.2,.2,-.25,.25)] if shape=='rectangular' else [(-.3,.3,-.25,.25)]
    else:disks.append((0.,0.,.3))
    expected=independent_area(disks,boxes,excluded_disks=[((offset+45)/1000,.005,.08)],excluded_boxes=[(.1,.2,-.03,.03)])
    assert abs(float(d['single_net_area'])-expected)<3e-10;replay(r)


@pytest.mark.parametrize('shape',['stepped_rectangular','stepped_circular'])
@pytest.mark.parametrize('x',['0','170','280','700'])
def test_offset_inner_disk_platform_independent_integral(shape,x):
    s=source(shape);s['foundation_items'][0].pop('through_holes',None);r=calculate(body(s,face='step',x=x,y='0',outer='220',inner_x=str(Decimal(x)+20),inner_y='15',inner='80'),s);d=r['contact_ranges'][0]
    ds=[(float(x)/1000,0.,.11)];inner=((float(x)+20)/1000,.015,.04)
    expected=independent_area(ds,[(-.3,.3,-.25,.25)],excluded_disks=[inner],excluded_boxes=[(-.2,.2,-.15,.15)]) if shape=='stepped_rectangular' else independent_area(ds+[(0.,0.,.3)],excluded_disks=[inner,(0.,0.,.2)])
    assert abs(float(d['single_net_area'])-expected)<3e-10;replay(r)


@pytest.mark.parametrize('unit,factor',[('mm','1'),('cm','.1'),('dm','.01'),('m','.001'),('km','.000001')])
def test_all_six_parameters_units_and_low_caller_precision(unit,factor):
    s=source();b=body(s);c=b['contact_ranges'][0]['plan_range'];c['unit']=unit;c['parameters']={k:str(Decimal(v)*Decimal(factor)) for k,v in c['parameters'].items()}
    with localcontext() as ctx:ctx.prec=6;r=calculate(b,s)
    assert r['quantity']==calculate(body(s),s)['quantity'];replay(r)


@pytest.mark.parametrize('gap',['-0.000000001','0','0.000000001'])
def test_strict_internal_tangency_no_tolerance(gap):
    c=config(x='0',y='0',outer='200',inner_x=str(Decimal(50)+Decimal(gap)),inner_y='0',inner='100')
    if Decimal(gap)<0:validate(c)
    else:
        with pytest.raises(DomainError,match='严格'):validate(c)


def test_negative_centers_diagonal_3_4_5_containment_and_concentric_equivalence():
    c=config(x='-100',y='-100',outer='200',inner_x='-70',inner_y='-60',inner='100')
    with pytest.raises(DomainError):validate(c)
    c['parameters']['inner_diameter']='99.999999999';validate(c)
    from tests.test_foundation_annular_windows import body as ab
    s=source();r=calculate(body(s,inner_x='120',inner_y='10'),s);old=calculate(ab(s),s)
    assert r['quantity']==old['quantity'] and r['contact_ranges'][0]['single_net_area']==old['contact_ranges'][0]['single_net_area'];replay(r)


@pytest.mark.parametrize('fault',['zero','negative','equal','larger','bool','nan','missing_center','relative_offset','area','extra','evidence','exclude','exclude_proof','confirm','frame','unit'])
def test_strict_schema_explicit_inner_center_and_evidence(fault):
    c=config()
    if fault in ['zero','negative','equal','larger','bool','nan']:c['parameters']['inner_diameter']={'zero':'0','negative':'-1','equal':'240','larger':'241','bool':True,'nan':'NaN'}[fault]
    elif fault=='missing_center':del c['parameters']['inner_x']
    elif fault=='relative_offset':c['parameters']['offset_x']='20'
    elif fault=='area':c['area']='100'
    elif fault=='extra':c['holes']=[]
    elif fault=='evidence':c['evidence']['inner_y']=''
    elif fault=='exclude':c['exclusion_basis']=''
    elif fault=='exclude_proof':c['exclusion_evidence_id']=''
    elif fault=='confirm':c['scope_confirmed']=1
    elif fault=='frame':c['coordinate_basis']=''
    else:c['unit']='px'
    with pytest.raises(DomainError):validate(c)


@pytest.mark.parametrize('gap',['-0.000000001','0','0.000000001'])
def test_other_outer_disk_inner_tangency_and_outer_tangency(gap):
    a=punctured(x='0',y='0',outer='200',inner_x='20',inner_y='10',inner='80');g=Decimal(gap)
    for other,expected in [(disk(x=str(Decimal(40)+g),y='10',diameter='40'),g>0),(disk(x=str(Decimal(120)+g),y='0',diameter='40'),g<0)]:
        for left,right in [(a,other),(other,a)]:
            q=separation(left,right);assert q['overlap']==expected
            if not expected:assert q['state']==('boundary_touch' if gap=='0' else 'separated') and q['method']==METHOD


def test_other_inner_contains_outer_and_crossing_outer_circles_keep_positive_boundary_neighborhood():
    a=punctured(x='0',y='0',outer='40',inner_x='1',inner_y='0',inner='10');b=punctured(x='-20',y='0',outer='200',inner_x='0',inner_y='0',inner='40')
    q=separation(a,b);assert not q['overlap'] and q['branch']=='self_in_other_inner' and q['state']=='boundary_touch'
    # Almost-full inner disks cannot erase the neighborhoods of crossing outer boundaries.
    a=punctured(x='0',y='0',outer='200',inner_x='.1',inner_y='0',inner='199');b=punctured(x='100',y='0',outer='200',inner_x='99.9',inner_y='0',inner='199')
    assert separation(a,b)['overlap'] and separation(b,a)['overlap']


def test_100_seeded_curved_pairs_independent_numerical_boolean_intersections():
    rng=random.Random(11401)
    for _ in range(100):
        x,y,X,Y=[rng.randint(-140,140) for _ in range(4)];R,S=rng.randint(30,100),rng.randint(30,100);r,s=rng.randint(5,R//2),rng.randint(5,S//2)
        ix,iy=x+rng.randint(-5,5),y+rng.randint(-5,5);I,J=X+rng.randint(-5,5),Y+rng.randint(-5,5)
        a=punctured(x=x,y=y,outer=2*R,inner_x=ix,inner_y=iy,inner=2*r);kind=rng.randrange(3)
        if kind==0:b=disk(x=X,y=Y,diameter=2*S);holes=[(ix/1000,iy/1000,r/1000)]
        elif kind==1:b=dict(plan_range=ac(x=X,y=Y,outer=2*S,inner=2*s));holes=[(ix/1000,iy/1000,r/1000),(X/1000,Y/1000,s/1000)]
        else:b=punctured(x=X,y=Y,outer=2*S,inner_x=I,inner_y=J,inner=2*s);holes=[(ix/1000,iy/1000,r/1000),(I/1000,J/1000,s/1000)]
        area=independent_area([(x/1000,y/1000,R/1000),(X/1000,Y/1000,S/1000)],excluded_disks=holes);expected=area>1e-11
        assert separation(a,b)['overlap']==expected and separation(b,a)['overlap']==expected


def test_120_rectangles_independent_two_center_extrema_and_witness_replay():
    rng=random.Random(11402)
    for _ in range(120):
        x,y,ix,iy=map(Fraction,[rng.randint(-100,100) for _ in range(4)]);ix=x+(ix%11)-5;iy=y+(iy%11)-5;r,R=20,100
        L,B=rng.randint(-180,180),rng.randint(-180,180);X,Y=L+rng.randint(1,80),B+rng.randint(1,80)
        minimum=max(L-x,0,x-X)**2+max(B-y,0,y-Y)**2;maximum=max((L-ix)**2,(X-ix)**2)+max((B-iy)**2,(Y-iy)**2);expected=minimum<R*R and maximum>r*r
        q=separation(punctured(x=x,y=y,outer=2*R,inner_x=ix,inner_y=iy,inner=2*r),dict(plan_range=window('test',str(L),str(X),str(B),str(Y))));assert q['overlap']==expected
        if not expected:
            p=q['components'][0];assert Fraction(p['minimum_outer_squared'])*1000000==minimum and Fraction(p['maximum_inner_squared'])*1000000==maximum
            for w in p['edge_witnesses']:
                cx,cy=map(Fraction,q['outer'][:2]);nx,ny=Fraction(w['nearest_x']),Fraction(w['nearest_y']);assert Fraction(w['distance_squared'])==(nx-cx)**2+(ny-cy)**2


def test_separate_net_blocks_hole_and_nested_island_two_center_behavior():
    a=punctured(x='0',y='0',outer='200',inner_x='30',inner_y='0',inner='80')
    c=mc([part(box(25,35,-5,5),name='内圆内'),part(box(150,170,-10,10),name='外圆外')]);q=separation(a,dict(plan_range=c));assert not q['overlap'] and [p['branch'] for p in q['components']]==['inside_inner','outside_outer']
    c=pc(outer=box(-150,150,-150,150),inners=[box(-110,110,-110,110)]);assert not separation(a,dict(plan_range=c))['overlap']
    c=nc([block(box(-150,150,-150,150),holes=[box(-110,110,-110,110)]),block(box(25,35,-5,5),name='偏心内圆岛',parent=dict(block_index=1,hole_index=1))]);assert not separation(a,dict(plan_range=c))['overlap']
    c['parts'][1]['parameters']=part(box(-30,-20,-10,10))['parameters'];assert separation(a,dict(plan_range=c))['overlap']


@pytest.mark.parametrize('gap',['0','0.000000001'])
def test_zero_all_boundaries_or_tiny_remaining_gap(gap):
    v=str(Decimal(25)+Decimal(gap));hs=[{**hole(x='0',diameter='80'),'vertical_range':interval('100','300')},{**rectangle(name='H2',x='-'+v,length='50',width='100'),'vertical_range':interval('300','500')},{**rectangle(name='H3',x=v,length='50',width='100'),'vertical_range':interval('300','500')}]
    s=end_source(holes=hs);r=calculate(body(s,face='hole_top',x='0',y='0',outer='60',inner_x='4',inner_y='3',inner='20'),s);d=r['contact_ranges'][0]
    assert ('polygon_zero_contact' in d)==(gap=='0') and (Decimal(d['single_net_area'])==0)==(gap=='0');replay(r)


def setup(clients,pid=None,zero=False):
    e,r,pid,s,b,eid=old_setup(clients,pid,zero);b['contact_ranges'][0]['plan_range']=config(eid,x='0' if zero else '196',y='0' if zero else '8',outer='40' if zero else '42',inner_x='3' if zero else '201',inner_y='2' if zero else '10',inner='12' if zero else '16');return e,r,pid,s,b,eid


def test_workflow_roles_foreign_inner_coordinate_history_concurrency_and_source_stale(clients):
    e,r,pid,s,b,eid=setup(clients);before=get(e,pid);post(e,pid,b,'/components/preview');assert get(e,pid)==before
    post(r,pid,b,expected=403);assert clients['other'].get(f'/api/v1/projects/{pid}').status_code==404
    foreign=create(e)['id'];upload(e,foreign,'异项目合成');bad=copy.deepcopy(b);bad['contact_ranges'][0]['plan_range']['evidence']['inner_x']=get(e,foreign)['evidence'][-1]['id'];post(e,pid,bad,expected=404)
    c=post(e,pid,b).json();approve(r,pid,'components',c);c=next(x for x in get(e,pid)['components'] if x['id']==c['id']);b['object_hash']=store.digest(c);b['contact_ranges'][0]['plan_range']['parameters']['inner_y']='11'
    res=e.put(f'/api/v1/projects/{pid}/components/'+c['id'],json=b);assert res.status_code==200,res.text;assert e.put(f'/api/v1/projects/{pid}/components/'+c['id'],json=b).status_code==409
    assert get(e,pid)['component_history'][-1]==c;approve(r,pid,'components',res.json())
    s['object_hash']=store.digest(s);s['foundation_items'][0]['levels'][1]['parameters']['diameter']='420';res=e.put(f'/api/v1/projects/{pid}/components/'+s['id'],json=s);assert res.status_code==200;approve(r,pid,'components',res.json());assert next(x for x in get(e,pid)['components'] if x['id']==c['id'])['state']=='stale'


def test_inner_disk_claim_allowed_but_positive_overlap_or_frame_mismatch_refused(clients):
    e,r,pid,s,b,eid=setup(clients);post(e,pid,b);other=copy.deepcopy(b);other['name']='合成偏心内圆另项';other['contact_ranges'][0]['plan_range']=cc(eid,x='201',y='10',diameter='16');post(e,pid,other)
    bad=copy.deepcopy(other);bad['contact_ranges'][0]['plan_range']['parameters']['x']='201.000000001';before=get(e,pid);post(e,pid,bad,expected=409);assert get(e,pid)==before
    bad=copy.deepcopy(other);bad['contact_ranges'][0]['plan_range']['coordinate_basis']='其他基准';post(e,pid,bad,expected=409)


@pytest.mark.parametrize('fault',['input','normalized','containment','containment_ast','quantity','basis','trace','proof','expression'])
def test_saved_input_containment_or_calculation_tamper_rejected(clients,fault):
    e,r,pid,s,b,eid=setup(clients,zero=True);c=post(e,pid,b).json()
    def forge(p):
        c=p['components'][-1];d=c['preview']['contact_ranges'][0]
        if fault=='input':c['contact_ranges'][0]['plan_range']['parameters']['inner_x']='4'
        elif fault=='normalized':d['window_circular_hole']['inner_x']='999'
        elif fault=='containment':d['window_inner_containment']['center_distance_squared']='999'
        elif fault=='containment_ast':d['window_inner_containment']['formula_asts']['center_distance_squared']={'constant':'0'}
        elif fault=='quantity':d['single_net_area']='1'
        elif fault=='basis':c['contact_ranges'][0]['plan_range']['exclusion_basis']='篡改用途'
        elif fault=='trace':d['polygon_net_trace']['partitions']=[]
        elif fault=='proof':d['polygon_zero_contact']['net_trace']['partitions'][0]['remaining']=[{}]
        else:d['window_expression_definitions'][0]['value']='999'
    mutate(pid,forge);before=get(e,pid);approve(r,pid,'components',c,expected=409);assert get(e,pid)==before


@pytest.mark.parametrize('gate',['preview','save','review','calculate','book_review'])
def test_inner_coordinate_original_bytes_five_gates(clients,gate):
    e,r,pid,s,b,eid=setup(clients);upload(e,pid,'合成内圆心XY原件，非真实工程','eccentric-center-basis.txt')
    for c in get(e,pid)['components']:approve(r,pid,'components',c)
    b['source_component_hash']=e.get(f'/api/v1/projects/{pid}/foundation-horizontal-formwork-candidates',params={'component_id':s['id']}).json()['source_component_hash'];b['contact_ranges'][0]['plan_range']['evidence']['inner_x']=get(e,pid)['evidence'][-1]['id']
    if gate in ('review','calculate','book_review'):c=post(e,pid,b).json()
    if gate in ('calculate','book_review'):approve(r,pid,'components',c)
    if gate=='book_review':t=post(e,pid,{},'/takeoffs/calculate').json()[-1]
    doc=get(e,pid)['documents'][-1];(store.ROOT/'originals'/doc['hash']).write_bytes(b'corrupted-synthetic-inner-center-original')
    if gate=='preview':post(e,pid,b,'/components/preview',409)
    elif gate=='save':post(e,pid,b,expected=409)
    elif gate=='review':approve(r,pid,'components',c,expected=409)
    elif gate=='calculate':post(e,pid,{},'/takeoffs/calculate',409)
    else:approve(r,pid,'takeoffs',t,expected=409)


@pytest.mark.parametrize('fmt',['json','xlsx','pdf'])
@pytest.mark.parametrize('zero',[False,True])
def test_three_formats_strict_containment_and_old_quote_isolation(clients,fmt,zero):
    from backend.foundation_eccentric_window_reports import TITLES
    e,r,pid,_,q=priced_project(clients);e,r,pid,s,b,eid=setup(clients,pid,zero);c=post(e,pid,b).json();approve(r,pid,'components',c)
    for t in post(e,pid,{},'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
    post(e,pid,{},'/boq/generate');url=f'/api/v1/projects/{pid}/exports';res=e.get(url,params=dict(format=fmt,unpriced=True));assert res.status_code==200,res.text
    titles=set(TITLES[:4] if zero else TITLES[:3])
    if fmt=='json':assert titles<={t[0] for t in res.json()['tables']}
    elif fmt=='xlsx':assert titles<=set(load_workbook(io.BytesIO(res.content)).sheetnames)
    else:assert all(any(title in page.extract_text() for page in PdfReader(io.BytesIO(res.content)).pages) for title in titles)
    assert not any(t[0] in titles for t in e.get(url,params=dict(format='json',quote_id=q['id'])).json()['tables'])


@pytest.mark.parametrize('fault',['distance_squared','branch','window_side'])
def test_precise_two_center_overlap_certificate_tamper_rejected(clients,fault):
    e,r,pid,s,b,eid=setup(clients);d=copy.deepcopy(b['contact_ranges'][0]);d.update(name='合成内圆另项',plan_range=cc(eid,x='201',y='10',diameter='16'));b['contact_ranges'].append(d);c=post(e,pid,b).json()
    def forge(p):p['components'][-1]['preview']['horizontal_window_checks'][0][fault]={'distance_squared':'999','branch':'outer_separation','window_side':'right'}[fault]
    mutate(pid,forge);before=get(e,pid);approve(r,pid,'components',c,expected=409);assert get(e,pid)==before


def test_extremely_thin_positive_offset_net_is_not_rounded_to_geometric_zero():
    s=source();s['foundation_items'][0].pop('through_holes',None);r=calculate(body(s,x='0',y='0',outer='200',inner_x='0.0000000001',inner_y='0',inner='199.999999999'),s);d=r['contact_ranges'][0]
    assert Decimal(d['single_net_area'])>0 and r['quantity']=='0.000' and 'polygon_zero_contact' not in d;replay(r)


def test_selected_contact_entirely_inside_offset_window_exclusion_keeps_zero_without_false_hole_proof():
    s=end_source();r=calculate(body(s,face='hole_top',x='0',y='0',outer='600',inner_x='10',inner_y='-10',inner='500'),s);d=r['contact_ranges'][0]
    assert Decimal(d['single_gross_area'])==0 and 'polygon_zero_contact' not in d;replay(r)
