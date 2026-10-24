"""Internal hole-end windows, independently integrated cuts and immutable review."""
import copy,io,math,random
from decimal import Decimal,localcontext
import pytest
from openpyxl import load_workbook
from pypdf import PdfReader
from backend import store
from backend.engine import DomainError
from backend.foundation_horizontal_formwork import calculate
from backend.foundation_hole_end_windows import VERSION
from backend.foundation_horizontal_formwork_reports import HOLE_WINDOW_TITLES,tables
from tests.test_foundation_hole_end_formwork import source,body as full_body,contact,setup as end_setup
from tests.test_foundation_horizontal_windows import window,check as base_check
from tests.test_foundation_openings import hole
from tests.test_foundation_voids import rectangle
from tests.test_foundation_recesses import interval
from tests.test_post_foundation import equal_body,fixed_body,fixed_added,equal_added,fixed_adjusted,equal_adjusted,combined
from tests.test_foundation_round_step_windows import oracle
from tests.test_route_takeoff import evaluate
from tests.test_polygon_takeoff import post,mutate
from tests.test_workflow import get,approve,upload,create,priced_project


def body(s,eid='test',face='hole_top',index=1,bounds=('0','100','-100','100')):
    b=full_body(s,eid,face,index);b['contact_ranges'][0]['plan_range']=window(eid,*bounds);return b


def check(r):
    base_check(r)
    with localcontext() as ctx:
        ctx.prec=80
        for d in r['contact_ranges']:
            for m in [d.get('window_curved_intersection')]+[h.get('curved_intersection') for h in d['apertures']]:
                if m:
                    for q in m['terms']:
                        assert evaluate(q['formula_ast'],r['dimensions'])==Decimal(q['area'])
                        for k,a in q['formula_asts'].items():assert evaluate(a,r['dimensions'])==Decimal(q[k])
            for q in d.get('hole_end_window_coverage',{}).get('terms',[]):
                for k,a in q['formula_asts'].items():assert evaluate(a,r['dimensions'])==Decimal(q[k])
                for piece in q['pieces']:
                    for k,a in piece['formula_asts'].items():assert evaluate(a,r['dimensions'])==Decimal(piece[k])


@pytest.mark.parametrize('shape',['rectangular','stepped_rectangular','circular','stepped_circular'])
@pytest.mark.parametrize('face',['hole_top','hole_bottom'])
@pytest.mark.parametrize('make',[equal_body,fixed_body,fixed_added,equal_added,fixed_adjusted,equal_adjusted,combined,lambda s:combined(s,equal=True)])
def test_all_shapes_both_ends_eight_final_counts_half_window_immutable(shape,face,make):
    s=source(shape,make);b=body(s,face=face);before=copy.deepcopy((s,b));r=calculate(b,s);d=r['contact_ranges'][0];assert (s,b)==before and r['rule_version']==VERSION and Decimal(d['single_gross_area'])==Decimal('.003');assert Decimal(r['gross'])==Decimal('.003')*Decimal(d['total_count']);check(r)


@pytest.mark.parametrize('kind',['circle-circle','circle-rect','rect-circle','rect-rect'])
@pytest.mark.parametrize('face',['hole_top','hole_bottom'])
def test_shifted_adjacent_holes_half_window_independent_oracle(kind,face):
    a=interval('100','300') if face=='hole_top' else interval('300','500');b=interval('300','500') if face=='hole_top' else interval('100','300')
    hs=[hole(x='-10',y='5',diameter='40') if kind.startswith('circle') else rectangle(x='-10',y='5',length='40',width='60'),hole(name='相接H2',x='10',y='-5',diameter='40') if kind.endswith('circle') else rectangle(name='相接H2',x='10',y='-5',length='40',width='60')]
    s=source(holes=[{**h,'vertical_range':z} for h,z in zip(hs,[a,b])]);r=calculate(body(s,face=face,bounds=('-20','20','0','30')),s);d=r['contact_ranges'][0]
    box=[-.02,.02,0,.03];disks=[]
    for h in hs:
        p={k:float(v)/1000 for k,v in h['parameters'].items()}
        if 'diameter' in p:disks.append((p['x'],p['y'],p['diameter']/2))
        else:box=[max(box[0],p['x']-p['length']/2),min(box[1],p['x']+p['length']/2),max(box[2],p['y']-p['width']/2),min(box[3],p['y']+p['width']/2)]
    expected=oracle(disks,box) if disks else max(0,box[1]-box[0])*max(0,box[3]-box[2]);assert abs(float(d['apertures'][0]['single_area'])-expected)<3e-12;check(r)


@pytest.mark.parametrize('shape,kind',[('stepped_rectangular','circle'),('stepped_rectangular','rect'),('stepped_circular','circle'),('stepped_circular','rect')])
@pytest.mark.parametrize('side',['left','right'])
def test_step_roof_triple_intersection_with_opposite_layer_selected(shape,kind,side):
    h=hole(x='190') if kind=='circle' else rectangle(x='200');s=source(shape,holes=[{**h,'vertical_range':interval('0','200')}]);bounds=('170','195','0','35') if side=='left' else ('195','260','0','35');b=body(s,bounds=bounds);r=calculate(b,s);d=r['contact_ranges'][0];box=[float(v)/1000 for v in bounds];disks=[]
    if kind=='circle':disks.append((.19,0,.02))
    else:box=[max(box[0],.15),min(box[1],.25),max(box[2],-.03),min(box[3],.03)]
    if shape=='stepped_circular':disks.append((0,0,.2))
    else:box=[max(box[0],-.2),min(box[1],.2),max(box[2],-.15),min(box[3],.15)]
    expected=oracle(disks,box) if disks else max(0,box[1]-box[0])*max(0,box[3]-box[2]);assert abs(float(d['single_gross_area'])-expected)<3e-12 and d['level_index']==2 and d['concrete_side']=='above';check(r)


@pytest.mark.parametrize('face',['hole_top','hole_bottom'])
@pytest.mark.parametrize('bounds',[('-100','100','-100','100'),('-40','40','-20','20'),('0','0.000000001','-10','10')])
def test_two_disjoint_same_side_holes_complete_coverage_exact_zero_certificate(face,bounds):
    a=interval('100','300') if face=='hole_top' else interval('300','500');z=interval('300','500') if face=='hole_top' else interval('100','300')
    hs=[{**hole(x='0',diameter='80'),'vertical_range':a},{**rectangle(name='相接H2',x='-25',length='50',width='100'),'vertical_range':z},{**rectangle(name='相接H3',x='25',length='50',width='100'),'vertical_range':z}];s=source(holes=hs);r=calculate(body(s,face=face,bounds=bounds),s);d=r['contact_ranges'][0];assert Decimal(d['single_gross_area'])>0 and Decimal(d['single_net_area'])==0 and d['hole_end_window_coverage']['terms'];assert r['quantity']=='0.000';check(r)


def test_positive_gap_not_certified_and_wrong_side_not_deducted():
    hs=[{**hole(x='0',diameter='80'),'vertical_range':interval('100','300')},{**rectangle(name='相接H2',x='-25.000000001',length='50',width='100'),'vertical_range':interval('300','500')},{**rectangle(name='相接H3',x='25.000000001',length='50',width='100'),'vertical_range':interval('300','500')}];s=source(holes=hs);r=calculate(body(s,bounds=('-50','50','-50','50')),s);d=r['contact_ranges'][0];assert 'hole_end_window_coverage' not in d and Decimal(d['single_net_area'])>0;check(r)
    hs=[{**rectangle(x='-60'),'vertical_range':interval('100','300')},{**rectangle(x='60',name='同侧投影孔'),'vertical_range':interval('100','300')}];s=source(holes=hs);r=calculate(body(s,bounds=('-200','200','-100','100')),s);assert not r['contact_ranges'][0]['apertures'];check(r)


@pytest.mark.parametrize('bounds',[('200','300','0','100'),('50','60','-100','100'),('-100','100','100','200'),('0','0.000000001','0','0.000000001')])
def test_zero_tangent_outside_and_tiny_positive_window(bounds):
    s=source(holes=[{**hole(x='0',diameter='100'),'vertical_range':interval('100','300')}]);r=calculate(body(s,bounds=bounds),s);assert (Decimal(r['contact_ranges'][0]['single_gross_area'])>0)==(bounds[1]=='0.000000001');check(r)


@pytest.mark.parametrize('source_unit,source_factor',[('mm','1'),('cm','.1'),('dm','.01'),('m','.001'),('km','.000001')])
@pytest.mark.parametrize('window_unit,window_factor',[('mm','1'),('cm','.1'),('dm','.01'),('m','.001'),('km','.000001')])
def test_independent_source_and_window_units_low_context(source_unit,source_factor,window_unit,window_factor):
    s=source('stepped_circular',holes=[{**hole(x='190'),'vertical_range':interval('0','200')}]);expected=calculate(body(s,bounds=('170','205','0','35')),s);item=s['foundation_items'][0];item['length_unit']=source_unit
    for level in item['levels']:
        for k in level['parameters']:level['parameters'][k]=str(Decimal(level['parameters'][k])*Decimal(source_factor))
    b=body(s,bounds=('170','205','0','35'));c=b['contact_ranges'][0]['plan_range'];c['unit']=window_unit
    for k in c['parameters']:c['parameters'][k]=str(Decimal(c['parameters'][k])*Decimal(window_factor))
    with localcontext() as ctx:ctx.prec=6;r=calculate(b,s)
    assert r['quantity']==expected['quantity'] and r['gross']==expected['gross'];check(r)


def test_seeded_160_windows_four_shape_pairs_against_independent_quadrature():
    rng=random.Random(10601)
    for i in range(160):
        circle1=i%2==0;circle2=i%4<2;hs=[hole(x='-10',y='5',diameter='80') if circle1 else rectangle(x='-10',y='5',length='100',width='80'),hole(name='H2',x='10',y='-5',diameter='60') if circle2 else rectangle(name='H2',x='10',y='-5',length='80',width='100')];s=source(holes=[{**h,'vertical_range':interval('100','300') if j==0 else interval('300','500')} for j,h in enumerate(hs)]);x,y=rng.randint(-90,30),rng.randint(-80,30);bs=(str(x),str(x+rng.randint(1,90)),str(y),str(y+rng.randint(1,90)));r=calculate(body(s,bounds=bs),s);d=r['contact_ranges'][0];box=[float(v)/1000 for v in bs];disks=[]
        for j,h in enumerate(hs):
            p={k:float(v)/1000 for k,v in h['parameters'].items()}
            if 'diameter' in p:disks.append((p['x'],p['y'],p['diameter']/2))
            else:box=[max(box[0],p['x']-p['length']/2),min(box[1],p['x']+p['length']/2),max(box[2],p['y']-p['width']/2),min(box[3],p['y']+p['width']/2)]
        expected=0 if box[0]>=box[1] or box[2]>=box[3] else oracle(disks,box) if disks else (box[1]-box[0])*(box[3]-box[2]);assert abs(float(d['apertures'][0]['single_area'])-expected)<3e-12;(check(r) if i<12 else None)


def test_same_hole_end_claims_and_different_holes_or_ends_are_distinct():
    s=source();b=body(s);c=copy.deepcopy(b['contact_ranges'][0]);c['name']='另一半';c['plan_range']=window(x1='-100',x2='0',y1='-100',y2='100');b['contact_ranges'].append(c);check(calculate(b,s));c['plan_range']['parameters']['x_end']='0.000000001'
    with pytest.raises(DomainError,match='正重叠'):calculate(b,s)
    s=source(holes=[{**rectangle(x='0'),'vertical_range':interval('100','200')},{**rectangle(x='0',name='H2'),'vertical_range':interval('300','400')}]);b=body(s);c=copy.deepcopy(b['contact_ranges'][0]);c.update(name='另一孔',hole_index=2);b['contact_ranges'].append(c);r=calculate(b,s);assert not r['horizontal_window_checks'];check(r)


def setup(clients,pid=None,zero=False):
    e,r,pid,s,b,eid=end_setup(clients,pid);upload(e,pid,'仅合成孔端窗口软件验收，非真实工法。22圆台阶基础下Φ600高200上Φ400高400mm；H1圆孔Φ40中心190,0，Z0至200，H2圆孔Φ40中心175,0，Z200至400。H1孔顶取上侧L2，XY窗口170至205、0至35mm；共同原点为基础中心，X东Y北。孔端、窗口、相接孔及总面积修约3位须独立复核。','hole-end-window-qa.txt')
    for c in get(e,pid)['components']:approve(r,pid,'components',c)
    s=next(c for c in get(e,pid)['components'] if c['id']==s['id']);eid=get(e,pid)['evidence'][-1]['id'];s['object_hash']=store.digest(s);c=s['foundation_items'][0]['through_holes'];c['holes']=[{**hole(eid=eid,x='190'),'vertical_range':interval('0','200',eid)},{**hole(eid=eid,name='H2',x='175'),'vertical_range':interval('200','400',eid)}]
    if zero:c['holes']=[{**hole(eid=eid,x='0',diameter='80'),'vertical_range':interval('0','200',eid)},{**rectangle(eid=eid,name='H2',x='-25',length='50',width='100'),'vertical_range':interval('200','400',eid)},{**rectangle(eid=eid,name='H3',x='25',length='50',width='100'),'vertical_range':interval('200','400',eid)}]
    res=e.put(f'/api/v1/projects/{pid}/components/'+s['id'],json=s);assert res.status_code==200,res.text;approve(r,pid,'components',res.json());s=next(c for c in get(e,pid)['components'] if c['id']==s['id']);return e,r,pid,s,body(s,eid,bounds=('-50','50','-50','50') if zero else ('170','205','0','35')),eid


def test_preview_candidates_claims_readonly_and_cross_record_frame(clients):
    e,r,pid,s,b,eid=setup(clients);before=get(e,pid);assert e.get(f'/api/v1/projects/{pid}/foundation-horizontal-formwork-candidates',params=dict(component_id=s['id'])).status_code==200;post(e,pid,b,'/components/preview');assert get(e,pid)==before;c=post(e,pid,b).json();approve(r,pid,'components',c);post(e,pid,b,expected=409);other=copy.deepcopy(b);other['contact_ranges'][0]['plan_range']['parameters'].update(x_start='205',x_end='230');post(e,pid,other);bad=copy.deepcopy(other);bad['contact_ranges'][0]['plan_range']['coordinate_basis']='另一基准';post(e,pid,bad,expected=409);post(e,pid,full_body(s,eid,face='hole_top'),expected=409)


@pytest.mark.parametrize('fault',['gross','cut','direct','trace','edge','window','ast','frame','z','side','source','book','coverage'])
def test_saved_input_and_trace_tamper_rejected(clients,fault):
    e,r,pid,s,b,eid=setup(clients,zero=fault=='coverage');c=post(e,pid,b).json()
    if fault=='book':approve(r,pid,'components',c);post(e,pid,{},'/takeoffs/calculate')
    def forge(p):
        c=p['components'][-1];d=c['preview']['contact_ranges'][0]
        if fault=='gross':d['single_gross_area']='0'
        elif fault=='cut':d['apertures'][0]['single_area']='0'
        elif fault=='direct':d['apertures'][0]['window_direct_area']='0'
        elif fault=='trace':d['window_curved_intersection']['terms'][0]['area']='0'
        elif fault=='edge':d['window_curved_intersection']['terms'][0]['upper_boundary']='forged'
        elif fault=='window':c['contact_ranges'][0]['plan_range']['parameters']['x_end']='210'
        elif fault=='ast':d['formula_asts']['x_start']={'constant':'0'}
        elif fault=='frame':d['plan_frame_hash']='forged'
        elif fault=='z':d['z']='0.201'
        elif fault=='side':d['concrete_side']='below'
        elif fault=='source':c['source_geometry']['foundation_items'][0]['through_holes']['holes'][0]['parameters']['x']='180'
        elif fault=='coverage':d['hole_end_window_coverage']['terms'][0]['holes']=[]
        else:p['takeoffs'][-1]['quantity']='999'
    mutate(pid,forge);coll='takeoffs' if fault=='book' else 'components';approve(r,pid,coll,get(e,pid)[coll][-1],expected=409)


@pytest.mark.parametrize('gate',['preview','save','review','calculate','book_review'])
def test_original_bytes_five_gates(clients,gate):
    e,r,pid,s,b,eid=setup(clients)
    if gate in ('review','calculate','book_review'):c=post(e,pid,b).json()
    if gate in ('calculate','book_review'):approve(r,pid,'components',c)
    if gate=='book_review':t=post(e,pid,{},'/takeoffs/calculate').json()[-1]
    doc=get(e,pid)['documents'][-1];(store.ROOT/'originals'/doc['hash']).write_bytes(b'corrupt')
    if gate=='preview':post(e,pid,b,'/components/preview',409)
    elif gate=='save':post(e,pid,b,expected=409)
    elif gate=='review':approve(r,pid,'components',c,expected=409)
    elif gate=='calculate':post(e,pid,{},'/takeoffs/calculate',409)
    else:approve(r,pid,'takeoffs',t,expected=409)


def test_roles_tenants_foreign_proof_history_concurrency_source_stale(clients):
    e,r,pid,s,b,eid=setup(clients);post(r,pid,b,expected=403);assert clients['other'].get(f'/api/v1/projects/{pid}').status_code==404;foreign=create(e)['id'];upload(e,foreign,'异项目原件');bad=copy.deepcopy(b);bad['contact_ranges'][0]['plan_range']['coordinate_evidence_id']=get(e,foreign)['evidence'][-1]['id'];post(e,pid,bad,expected=404);c=post(e,pid,b).json();approve(r,pid,'components',c)
    for t in post(e,pid,{},'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
    post(e,pid,{},'/boq/generate');before=get(e,pid);b['object_hash']=store.digest(before['components'][-1]);c=copy.deepcopy(b['contact_ranges'][0]);c['name']='补充窗口';c['plan_range']=window(eid,'170','205','-35','0');b['contact_ranges'].append(c);url=f'/api/v1/projects/{pid}/components/'+before['components'][-1]['id'];res=e.put(url,json=b);assert res.status_code==200,res.text;after=get(e,pid);assert after['component_history'][-1]==before['components'][-1] and after['boq']==before['boq'] and e.put(url,json=b).status_code==409;approve(r,pid,'components',res.json());s['object_hash']=store.digest(s);s['foundation_items'][0]['levels'][1]['parameters']['diameter']='420';res=e.put(f'/api/v1/projects/{pid}/components/'+s['id'],json=s);assert res.status_code==200,res.text;approve(r,pid,'components',res.json());assert get(e,pid)['components'][-1]['state']=='stale'


@pytest.mark.parametrize('fmt',['json','xlsx','pdf'])
@pytest.mark.parametrize('zero',[False,True])
def test_three_formats_interval_coverage_and_old_quote_isolation(clients,fmt,zero):
    e,r,pid,_,q=priced_project(clients);e,r,pid,s,b,eid=setup(clients,pid,zero);c=post(e,pid,b).json();approve(r,pid,'components',c)
    for t in post(e,pid,{},'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
    post(e,pid,{},'/boq/generate');url=f'/api/v1/projects/{pid}/exports';res=e.get(url,params=dict(format=fmt,unpriced=True));assert res.status_code==200,res.text;titles=set(HOLE_WINDOW_TITLES if zero else HOLE_WINDOW_TITLES[:2])
    if fmt=='json':assert titles<={t[0] for t in res.json()['tables']}
    elif fmt=='xlsx':assert titles<=set(load_workbook(io.BytesIO(res.content)).sheetnames)
    else:assert titles<=set(t for p in PdfReader(io.BytesIO(res.content)).pages for t in titles if t in p.extract_text())
    assert not any(t[0] in set(HOLE_WINDOW_TITLES) for t in e.get(url,params=dict(format='json',quote_id=q['id'])).json()['tables']);assert tables(dict(takeoffs=[]))==[]
