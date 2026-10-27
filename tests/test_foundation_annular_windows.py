"""Synthetic concentric annuli: independent integration, exact radial topology and evidence gates."""
import copy,io,random,math
from decimal import Decimal,localcontext
from fractions import Fraction
import pytest
from openpyxl import load_workbook
from pypdf import PdfReader
from backend import store
from backend.compound_takeoff import PI
from backend.engine import DomainError
from backend.foundation_horizontal_formwork import calculate
from backend.foundation_annular_windows import VERSION,METHOD,validate,separation
from backend.foundation_horizontal_windows import checks
from tests.test_foundation_circular_windows import body as cb,config as cc,setup as old_setup,independent_area,replay
from tests.test_foundation_horizontal_windows import source,window,FRAME
from tests.test_foundation_hole_end_formwork import source as end_source
from tests.test_foundation_nested_windows import config as nc,block,box
from tests.test_foundation_multi_windows import config as mc,part
from tests.test_foundation_perforated_windows import config as pc
from tests.test_foundation_simple_windows import config as sc,C_SHAPE
from tests.test_foundation_openings import hole
from tests.test_foundation_voids import rectangle
from tests.test_foundation_recesses import interval
from tests.test_post_foundation import equal_body,fixed_body,fixed_added,equal_added,fixed_adjusted,equal_adjusted,combined
from tests.test_polygon_takeoff import post,mutate
from tests.test_workflow import get,approve,upload,create,priced_project


def config(eid='test',x='120',y='10',outer='240',inner='80',unit='mm'):
    c=cc(eid,x,y,outer,unit);c.update(shape='annular',parameters=dict(x=str(x),y=str(y),outer_diameter=str(outer),inner_diameter=str(inner)),exclusion_basis='合成内圆不支模，非实体孔；仅软件验证',exclusion_evidence_id=eid);c['evidence']={k:eid for k in c['parameters']};return c


def body(s,eid='test',face='bottom',**kw):
    b=cb(s,eid,face);b['contact_ranges'][0]['plan_range']=config(eid,**kw);return b


def ring(**kw):return dict(plan_range=config(**kw))
def disk(**kw):return dict(plan_range=cc(**kw))


@pytest.mark.parametrize('shape',['rectangular','circular','stepped_rectangular','stepped_circular'])
@pytest.mark.parametrize('face',['top','bottom','hole_top','hole_bottom'])
@pytest.mark.parametrize('make',[equal_body,fixed_body,fixed_added,equal_added,fixed_adjusted,equal_adjusted,combined,lambda s:combined(s,equal=True)])
def test_shapes_ends_counts_no_mutation_and_ast_replay(shape,face,make):
    s=end_source(shape,make) if face.startswith('hole') else source(shape,make);b=body(s,face=face);before=copy.deepcopy((s,b));r=calculate(b,s);d=r['contact_ranges'][0]
    assert before==(s,b) and r['rule_version']==VERSION and d['annular_window'] and not d.get('circular_window')
    with localcontext() as ctx:
        ctx.prec=80;assert Decimal(d['window_area'])==Decimal(PI)*Decimal('.0128')
        assert Decimal(d['total_net_area'])==Decimal(d['single_net_area'])*Decimal(d['total_count'])
    assert r['source_evidence']['horizontal_formwork.1.window_inner_basis']=='test';replay(r)


@pytest.mark.parametrize('shape',['rectangular','circular','stepped_rectangular','stepped_circular'])
@pytest.mark.parametrize('face',['top','bottom','hole_bottom','hole_top'])
def test_outer_disk_minus_inner_disk_matches_all_gross_cut_and_net(shape,face):
    s=end_source(shape) if face.startswith('hole') else source(shape);r=calculate(body(s,face=face),s);a=calculate(cb(s,face=face,x='120',y='10',diameter='240'),s);b=calculate(cb(s,face=face,x='120',y='10',diameter='80'),s)
    with localcontext() as ctx:
        ctx.prec=80
        for k in ['single_gross_area','single_opening_area','single_net_area','total_net_area']:
            assert abs(Decimal(r['contact_ranges'][0][k])-(Decimal(a['contact_ranges'][0][k])-Decimal(b['contact_ranges'][0][k])))<Decimal('1e-73')
    replay(r)


@pytest.mark.parametrize('shape',['rectangular','circular','stepped_rectangular','stepped_circular'])
@pytest.mark.parametrize('offset',[-350,-160,0,160,350])
def test_20_independent_numeric_ring_face_and_hole_intersections(shape,offset):
    s=source(shape);s['foundation_items'][0]['through_holes']['holes']=[{**rectangle(x='150',length='100',width='60'),'vertical_range':interval('0','150')}]
    r=calculate(body(s,x=offset,y='30',outer='360',inner='160'),s);d=r['contact_ranges'][0];disks=[(offset/1000,.03,.18)];boxes=[]
    if shape in ('rectangular','stepped_rectangular'):boxes=[(-.2,.2,-.25,.25)] if shape=='rectangular' else [(-.3,.3,-.25,.25)]
    else:disks.append((0.,0.,.3))
    expected=independent_area(disks,boxes,excluded_disks=[(offset/1000,.03,.08)],excluded_boxes=[(.1,.2,-.03,.03)])
    assert abs(float(d['single_net_area'])-expected)<3e-10;replay(r)


@pytest.mark.parametrize('shape',['stepped_rectangular','stepped_circular'])
@pytest.mark.parametrize('x',['0','170','280','700'])
def test_platform_actual_ring_and_upper_exclusion_independent_integral(shape,x):
    s=source(shape);s['foundation_items'][0].pop('through_holes',None);r=calculate(body(s,face='step',x=x,y='0',outer='220',inner='80'),s);d=r['contact_ranges'][0];ds=[(float(x)/1000,0.,.11)];inner=(float(x)/1000,0.,.04)
    expected=independent_area(ds,[(-.3,.3,-.25,.25)],excluded_disks=[inner],excluded_boxes=[(-.2,.2,-.15,.15)]) if shape=='stepped_rectangular' else independent_area(ds+[(0.,0.,.3)],excluded_disks=[inner,(0.,0.,.2)])
    assert abs(float(d['single_net_area'])-expected)<3e-10;replay(r)


@pytest.mark.parametrize('unit,factor',[('mm','1'),('cm','.1'),('dm','.01'),('m','.001'),('km','.000001')])
def test_common_units_low_precision_and_replay(unit,factor):
    s=source();b=body(s);c=b['contact_ranges'][0]['plan_range'];c['unit']=unit;c['parameters']={k:str(Decimal(v)*Decimal(factor)) for k,v in c['parameters'].items()}
    with localcontext() as ctx:ctx.prec=6;r=calculate(b,s)
    assert r['quantity']==calculate(body(s),s)['quantity'];replay(r)


@pytest.mark.parametrize('fault',['zero_inner','negative_inner','equal','larger','bool','nan','radius','eccentric','area','extra','evidence','exclude','exclude_proof','confirm','frame','unit'])
def test_schema_strict_inner_outer_and_exclusion_evidence(fault):
    c=config()
    if fault in ['zero_inner','negative_inner','equal','larger','bool','nan']:c['parameters']['inner_diameter']={'zero_inner':'0','negative_inner':'-1','equal':'240','larger':'241','bool':True,'nan':'NaN'}[fault]
    elif fault=='radius':c['parameters']['radius']=c['parameters'].pop('outer_diameter')
    elif fault=='eccentric':c['parameters']['inner_x']='0'
    elif fault=='area':c['area']='100'
    elif fault=='extra':c['holes']=[]
    elif fault=='evidence':c['evidence']['inner_diameter']=''
    elif fault=='exclude':c['exclusion_basis']=''
    elif fault=='exclude_proof':c['exclusion_evidence_id']=''
    elif fault=='confirm':c['scope_confirmed']=1
    elif fault=='frame':c['coordinate_basis']=''
    else:c['unit']='px'
    with pytest.raises(DomainError):validate(c)


@pytest.mark.parametrize('gap',['-0.000000001','0','0.000000001'])
def test_exact_outer_and_inner_tangencies_circle_and_ring(gap):
    a=ring(x='0',y='0',outer='200',inner='100');g=Decimal(gap)
    for other,expected in [(disk(x=str(Decimal(200)+g),y='0',diameter='200'),g<0),
                            (disk(x=str(Decimal(30)+g),y='0',diameter='40'),g>0),
                            (ring(x=str(Decimal(200)+g),y='0',outer='200',inner='50'),g<0),
                            (ring(x=str(Decimal(30)+g),y='0',outer='40',inner='20'),g>0)]:
        for left,right in [(a,other),(other,a)]:
            q=separation(left,right);assert q['overlap']==expected
            if not expected:assert q['state']==('boundary_touch' if gap=='0' else 'separated') and q['method']==METHOD


def test_concentric_radial_ranges_and_right_annulus_witness():
    a=ring(x='0',y='0',outer='200',inner='100')
    for other,expected in [(disk(x='0',y='0',diameter='100'),False),(disk(x='0',y='0',diameter='101'),True),(ring(x='0',y='0',outer='100',inner='40'),False),(ring(x='0',y='0',outer='300',inner='200'),False),(ring(x='0',y='0',outer='201',inner='199'),True),(a,True)]:
        assert separation(a,other)['overlap']==expected and separation(other,a)['overlap']==expected
    q=separation(disk(x='0',y='0',diameter='80'),a);assert q['annulus_side']=='right' and q['branch']=='other_in_inner'


def test_120_random_circle_ring_pairs_independent_euclidean_radial_conditions():
    rng=random.Random(11301)
    for _ in range(120):
        x,y,X,Y=[rng.randint(-200,200) for _ in range(4)];r,R=sorted(rng.sample(range(1,120),2));s,S=sorted(rng.sample(range(1,120),2));s=0 if rng.randrange(2) else s
        distance=math.hypot(x-X,y-Y);expected=distance<R+S and distance+S>r and distance+R>s
        a=ring(x=x,y=y,outer=2*R,inner=2*r);b=ring(x=X,y=Y,outer=2*S,inner=2*s) if s else disk(x=X,y=Y,diameter=2*S)
        assert separation(a,b)['overlap']==expected and separation(b,a)['overlap']==expected


def test_120_random_rectangles_independent_axis_extrema_and_proof_replay():
    rng=random.Random(11302)
    for _ in range(120):
        x,y=[Fraction(rng.randint(-150,150)) for _ in range(2)];r,R=sorted(rng.sample(range(1,120),2));L,B=rng.randint(-150,150),rng.randint(-150,150);X,Y=L+rng.randint(1,100),B+rng.randint(1,100)
        minimum=max(L-x,0,x-X)**2+max(B-y,0,y-Y)**2;maximum=max((L-x)**2,(X-x)**2)+max((B-y)**2,(Y-y)**2);expected=minimum<R*R and maximum>r*r
        q=separation(ring(x=x,y=y,outer=2*R,inner=2*r),dict(plan_range=window('test',str(L),str(X),str(B),str(Y))));assert q['overlap']==expected
        if not expected:
            p=q['components'][0];assert Fraction(p['minimum_squared'])*1000000==minimum and Fraction(p['maximum_squared'])*1000000==maximum
            for w in p['edge_witnesses']:
                cx,cy=Fraction(q['center_x']),Fraction(q['center_y']);nx,ny=Fraction(w['nearest_x']),Fraction(w['nearest_y']);assert Fraction(w['distance_squared'])==(nx-cx)**2+(ny-cy)**2


def test_connected_net_components_not_global_extrema_and_inner_holes():
    a=ring(x='0',y='0',outer='200',inner='100')
    # One block inside excluded inner disk; another beyond outer disk. Combining global extrema would falsely report overlap.
    c=mc([part(box(-10,10,-10,10),name='内圆内'),part(box(150,170,-10,10),name='外圆外')]);q=separation(a,dict(plan_range=c));assert not q['overlap'] and [p['branch'] for p in q['components']]==['inside_inner','outside_outer']
    c=pc(outer=box(-150,150,-150,150),inners=[box(-110,110,-110,110)]);assert not separation(a,dict(plan_range=c))['overlap']
    c=nc([block(box(-150,150,-150,150),holes=[box(-110,110,-110,110)]),block(box(-10,10,-10,10),name='内圆岛',parent=dict(block_index=1,hole_index=1))]);assert not separation(a,dict(plan_range=c))['overlap']
    c['parts'][1]['parameters']=part(box(60,80,-10,10))['parameters'];assert separation(a,dict(plan_range=c))['overlap']
    assert separation(a,dict(plan_range=sc(C_SHAPE)))['overlap']


@pytest.mark.parametrize('gap',['0','0.000000001'])
def test_physical_coverage_ring_proof_and_tiny_remaining_contact(gap):
    v=str(Decimal(25)+Decimal(gap));hs=[{**hole(x='0',diameter='80'),'vertical_range':interval('100','300')},{**rectangle(name='H2',x='-'+v,length='50',width='100'),'vertical_range':interval('300','500')},{**rectangle(name='H3',x=v,length='50',width='100'),'vertical_range':interval('300','500')}]
    s=end_source(holes=hs);r=calculate(body(s,face='hole_top',x='0',y='0',outer='60',inner='20'),s);d=r['contact_ranges'][0]
    assert ('polygon_zero_contact' in d)==(gap=='0') and (Decimal(d['single_net_area'])==0)==(gap=='0');replay(r)


def test_thin_positive_annulus_and_source_inside_excluded_inner_disk():
    s=source();s['foundation_items'][0].pop('through_holes',None);r=calculate(body(s,x='0',y='0',outer='200',inner='199.999999999'),s);d=r['contact_ranges'][0]
    assert Decimal(d['single_net_area'])>0 and r['quantity']=='0.000' and 'polygon_zero_contact' not in d;replay(r)
    s=end_source();d=calculate(body(s,face='hole_top',x='0',y='0',outer='600',inner='500'),s)['contact_ranges'][0];assert Decimal(d['single_gross_area'])==0 and 'polygon_zero_contact' not in d


def setup(clients,pid=None,zero=False):
    e,r,pid,s,b,eid=old_setup(clients,pid,zero);b['contact_ranges'][0]['plan_range']=config(eid,x='0' if zero else '196',y='0' if zero else '8',outer='40' if zero else '42',inner='12' if zero else '16');return e,r,pid,s,b,eid


def test_roles_foreign_exclusion_evidence_history_concurrency_and_source_stale(clients):
    e,r,pid,s,b,eid=setup(clients);before=get(e,pid);post(e,pid,b,'/components/preview');assert get(e,pid)==before
    post(r,pid,b,expected=403);assert clients['other'].get(f'/api/v1/projects/{pid}').status_code==404
    foreign=create(e)['id'];upload(e,foreign,'异项目合成');bad=copy.deepcopy(b);bad['contact_ranges'][0]['plan_range']['exclusion_evidence_id']=get(e,foreign)['evidence'][-1]['id'];post(e,pid,bad,expected=404)
    c=post(e,pid,b).json();approve(r,pid,'components',c);c=next(x for x in get(e,pid)['components'] if x['id']==c['id']);b['object_hash']=store.digest(c);b['contact_ranges'][0]['plan_range']['parameters']['inner_diameter']='18'
    res=e.put(f'/api/v1/projects/{pid}/components/'+c['id'],json=b);assert res.status_code==200,res.text;assert e.put(f'/api/v1/projects/{pid}/components/'+c['id'],json=b).status_code==409
    assert get(e,pid)['component_history'][-1]==c;approve(r,pid,'components',res.json())
    s['object_hash']=store.digest(s);s['foundation_items'][0]['levels'][1]['parameters']['diameter']='420';res=e.put(f'/api/v1/projects/{pid}/components/'+s['id'],json=s);assert res.status_code==200;approve(r,pid,'components',res.json());assert next(x for x in get(e,pid)['components'] if x['id']==c['id'])['state']=='stale'


def test_inner_disk_can_be_claimed_but_ring_overlap_cross_record_and_frame_refused(clients):
    e,r,pid,s,b,eid=setup(clients);post(e,pid,b)
    other=copy.deepcopy(b);other['name']='合成内圆空区另项';other['contact_ranges'][0]['plan_range']=cc(eid,x='196',y='8',diameter='16');post(e,pid,other)
    bad=copy.deepcopy(other);bad['contact_ranges'][0]['plan_range']['parameters']['diameter']='16.000000001';before=get(e,pid);post(e,pid,bad,expected=409);assert get(e,pid)==before
    bad=copy.deepcopy(other);bad['contact_ranges'][0]['plan_range']['coordinate_basis']='其他基准';post(e,pid,bad,expected=409)


@pytest.mark.parametrize('fault',['input','normalized','window_area','quantity','inner_basis','trace','proof','expression'])
def test_stored_annulus_or_evidence_tamper_rejected(clients,fault):
    e,r,pid,s,b,eid=setup(clients,zero=True);c=post(e,pid,b).json()
    def forge(p):
        c=p['components'][-1];d=c['preview']['contact_ranges'][0]
        if fault=='input':c['contact_ranges'][0]['plan_range']['parameters']['inner_diameter']='13'
        elif fault=='normalized':d['window_annulus']['inner_diameter']='999'
        elif fault=='window_area':d['window_exclusion_area']='999'
        elif fault=='quantity':d['single_net_area']='1'
        elif fault=='inner_basis':c['contact_ranges'][0]['plan_range']['exclusion_basis']='篡改内圆用途'
        elif fault=='trace':d['polygon_net_trace']['partitions']=[]
        elif fault=='proof':d['polygon_zero_contact']['net_trace']['partitions'][0]['remaining']=[{}]
        else:d['window_expression_definitions'][0]['value']='999'
    mutate(pid,forge);before=get(e,pid);approve(r,pid,'components',c,expected=409);assert get(e,pid)==before


@pytest.mark.parametrize('gate',['preview','save','review','calculate','book_review'])
def test_inner_exclusion_original_bytes_five_gates(clients,gate):
    e,r,pid,s,b,eid=setup(clients);upload(e,pid,'合成内圆不支模原件，非真实工程','annulus-inner-basis.txt')
    for c in get(e,pid)['components']:approve(r,pid,'components',c)
    b['source_component_hash']=e.get(f'/api/v1/projects/{pid}/foundation-horizontal-formwork-candidates',params={'component_id':s['id']}).json()['source_component_hash'];b['contact_ranges'][0]['plan_range']['exclusion_evidence_id']=get(e,pid)['evidence'][-1]['id']
    if gate in ('review','calculate','book_review'):c=post(e,pid,b).json()
    if gate in ('calculate','book_review'):approve(r,pid,'components',c)
    if gate=='book_review':t=post(e,pid,{},'/takeoffs/calculate').json()[-1]
    doc=get(e,pid)['documents'][-1];(store.ROOT/'originals'/doc['hash']).write_bytes(b'corrupted-synthetic-annulus-inner-original')
    if gate=='preview':post(e,pid,b,'/components/preview',409)
    elif gate=='save':post(e,pid,b,expected=409)
    elif gate=='review':approve(r,pid,'components',c,expected=409)
    elif gate=='calculate':post(e,pid,{},'/takeoffs/calculate',409)
    else:approve(r,pid,'takeoffs',t,expected=409)


@pytest.mark.parametrize('fmt',['json','xlsx','pdf'])
@pytest.mark.parametrize('zero',[False,True])
def test_three_formats_and_old_quote_isolation(clients,fmt,zero):
    from backend.foundation_annular_window_reports import TITLES
    e,r,pid,_,q=priced_project(clients);e,r,pid,s,b,eid=setup(clients,pid,zero);c=post(e,pid,b).json();approve(r,pid,'components',c)
    for t in post(e,pid,{},'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
    post(e,pid,{},'/boq/generate');url=f'/api/v1/projects/{pid}/exports';res=e.get(url,params=dict(format=fmt,unpriced=True));assert res.status_code==200,res.text
    titles=set(TITLES[:3] if zero else TITLES[:2])
    if fmt=='json':assert titles<={t[0] for t in res.json()['tables']}
    elif fmt=='xlsx':assert titles<=set(load_workbook(io.BytesIO(res.content)).sheetnames)
    else:assert all(any(title in page.extract_text() for page in PdfReader(io.BytesIO(res.content)).pages) for title in titles)
    assert not any(t[0] in titles for t in e.get(url,params=dict(format='json',quote_id=q['id'])).json()['tables'])

@pytest.mark.parametrize('fault',['distance_squared','branch','annulus_side'])
def test_exact_radial_certificate_tamper_rejected(clients,fault):
    e,r,pid,s,b,eid=setup(clients);other=copy.deepcopy(b['contact_ranges'][0]);other['name']='合成内圆另项';other['plan_range']=cc(eid,x='196',y='8',diameter='16');b['contact_ranges'].append(other);c=post(e,pid,b).json()
    def forge(p):
        q=p['components'][-1]['preview']['horizontal_window_checks'][0]
        q[fault]={'distance_squared':'999','branch':'outer_separation','annulus_side':'right'}[fault]
    mutate(pid,forge);before=get(e,pid);approve(r,pid,'components',c,expected=409);assert get(e,pid)==before
