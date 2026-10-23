"""Synthetic XY window clipping, independent exact-area oracles and workflow gates."""
import copy
import io
import json
import random
from decimal import Decimal, localcontext
from fractions import Fraction
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pypdf import PdfReader
from backend import store
from backend.main import app
from backend.engine import DomainError
from backend.surface_takeoff import boundary, partition, calculate, snapshot
from backend.polygon_cap_partitions import checked
from tests.test_polygon_takeoff import payload as polygon, ring
from tests.test_surface_takeoff import payload as surface
from tests.test_polygon_side_partitions import cuts as side_cuts
from tests.test_route_takeoff import evaluate
from tests.test_workflow import get, approve, upload, create, priced_project


def geometry(eid='test'):
    s=polygon(eid);s.update(id='cap-source',revision=1,multiplicity='2')
    s['polygon_rings']=[ring('凹形外轮廓',[(0,0),(6,0),(6,4),(4,4),(4,6),(0,6)],eid),
                        ring('跨窗口方孔',[(1,1),(3,1),(3,3),(1,3)],eid,'hole')]
    s['measurement_basis']='合成外32m²、贯通孔4m²，净28m²；厚200mm、重复2；XY窗口扣孔，非专业计量依据'
    return s


def groups(source,eid='test',top=True,x='2',y='2'):
    f=next(f for f in boundary(source)['faces'] if f['normal']==('z+' if top else 'z-'))
    return [{'face_id':f['face_id'],'u_cuts':[{'value':x,'unit':'m','evidence_id':eid}],
             'v_cuts':[{'value':y,'unit':'m','evidence_id':eid}]}]


def payload(source,eid='test',cuts=None):
    cuts=groups(source,eid) if cuts is None else cuts;view=partition(source,cuts)
    b=surface(source,eid,included=lambda f:False)
    b.update(name='合成凹形扣孔局部防水（非真实工程）',purpose='waterproof',face_subdivisions=cuts,
             measurement_basis='合成共同设计X/Y按2m分割，仅顶部X0至2、Y0至2的扣孔净区计入，其余排除；3m²×2=6m²',
             face_decisions=[{'face_id':f['face_id'],'included':bool(f.get('cap_clip_rule')) and Decimal(f['u_to'])<=2 and Decimal(f['v_to'])<=2,
                              'basis':'合成按XY窗口扣孔后的净区域核定，非矩形窗口面积','evidence_id':eid} for f in view['faces']])
    return b


def setup(clients,pid=None):
    e,r=clients['estimator'],clients['reviewer'];pid=pid or create(e)['id']
    upload(e,pid,'合成软件验收，非真实工程：同一水平设计平面及共同XY。外轮廓(0,0)(6,0)(6,4)(4,4)(4,6)(0,6)m，32m²；贯通孔(1,1)(3,1)(3,3)(1,3)m，4m²，净28m²；厚200mm，重复2。顶部按共同X=200cm、Y=2000mm分割，4个窗口扣孔净量分别3、7、7、11m²，仅左下净3m²计入，重复2=6m²；底部及全部侧面排除。','cap-window-basis.txt')
    eid=get(e,pid)['evidence'][-1]['id'];s=geometry(eid);s.pop('id');s.pop('revision');s['name']='合成凹形扣孔等厚实体（非真实工程）'
    response=e.post(f'/api/v1/projects/{pid}/components',json=s);assert response.status_code==200,response.text
    for c in get(e,pid)['components']:approve(r,pid,'components',c)
    source=next(c for c in get(e,pid)['components'] if c['id']==response.json()['id'])
    return e,r,pid,source,payload(source,eid),eid


def post(e,pid,path,body=None,expected=200):
    res=e.post(f'/api/v1/projects/{pid}'+path,json=body or {});assert res.status_code==expected,res.text;return res


def mutate(pid,action):
    with store.transaction() as c:
        p=json.loads(c.execute('select data from projects where id=?',(pid,)).fetchone()['data']);action(p);store.save(c,p)


@pytest.mark.parametrize('top',[False,True])
def test_concave_and_cross_window_hole_exact_net_areas_and_ast(top):
    s=geometry();g=groups(s,top=top);before=copy.deepcopy(s);r=partition(s,g)
    children=[f for f in r['faces'] if f.get('cap_clip_rule')]
    assert s==before and len(r['faces'])==15 and [Decimal(c['area']) for c in children]==[3,7,7,11]
    assert [Decimal(c['cap_clip_rings'][1]['area']) for c in children]==[1,1,1,1]
    assert r['surface_cap_partition_checks'][0]['exact_rational_tiling'] and Decimal(r['surface_cap_partition_checks'][0]['area_difference'])==0
    b=payload(s,cuts=g);out=calculate(b,s);assert Decimal(out['quantity'])==6 and out['rule_version']=='polygon-boundary-1.2'
    with localcontext() as ctx:
        ctx.prec=80
        assert evaluate(out['formula_ast'],out['dimensions'])==6
        for f in out['surface_faces']:assert evaluate(f['formula_ast'],out['dimensions'])==Decimal(f['area'])
    assert len({f['face_id'] for f in r['faces']})==len(r['faces'])


def test_nonterminating_intersections_keep_exact_rational_area_and_audit_precision():
    s=geometry();s['polygon_rings']=[ring('斜三角形',[(0,0),(3,0),(0,2)])]
    r=partition(s,groups(s,x='1',y='1'));children=[f for f in r['faces'] if f.get('cap_clip_rule')]
    expected=[Fraction(1),Fraction(2,3),Fraction(5,4),Fraction(1,12)]
    assert [Fraction(int(f['cap_rational_area']['numerator']),int(f['cap_rational_area']['denominator'])) for f in children]==expected
    assert sum(expected)==3
    with localcontext() as ctx:
        ctx.prec=80
        for f in children:assert evaluate(f['formula_ast'],r['dimensions'])==Decimal(f['area'])
    with localcontext() as ctx:
        ctx.prec=200;check=r['surface_cap_partition_checks'][0]
        assert Decimal(check['children_area'])==sum((Decimal(c['area']) for c in children),Decimal(0))
        assert Decimal(check['area_difference'])==Decimal(check['children_area'])-3


def test_disconnected_concave_intersection_counts_all_pieces_without_bridge_area():
    s=geometry();s['polygon_rings']=[ring('U形',[(0,0),(6,0),(6,6),(4,6),(4,2),(2,2),(2,6),(0,6)])]
    g=groups(s,x='3',y='3');g[0]['u_cuts']=[];r=partition(s,g)
    children=[f for f in r['faces'] if f.get('cap_clip_rule')];assert [Decimal(f['area']) for f in children]==[16,12]
    assert Decimal(r['surface_cap_partition_checks'][0]['parent_area'])==28


def test_empty_hole_only_outside_and_boundary_touch_windows_do_not_create_faces():
    s=geometry();s['polygon_rings']=[ring('三角形',[(0,0),(6,0),(0,6)]),ring('方孔',[(1,1),(2,1),(2,2),(1,2)],role='hole')]
    g=groups(s,x='2',y='2');g[0]['u_cuts']+=[{'value':v,'unit':'m','evidence_id':'test'} for v in ('1','3')];g[0]['v_cuts']+=[{'value':v,'unit':'m','evidence_id':'test'} for v in ('1','3')]
    r=partition(s,g);check=r['surface_cap_partition_checks'][0]
    assert check['zero_window_count']==2 and check['child_count']==14
    assert all(Decimal(f['area'])>0 for f in r['faces'] if f.get('cap_clip_rule'))


@pytest.mark.parametrize('mode',['reverse','translate','units','swap_axes'])
def test_orientation_negative_design_coordinates_and_unit_equivalence(mode):
    s=geometry();g=groups(s)
    for row in s['polygon_rings']:
        if mode=='reverse':row['points'].reverse()
        for p in row['points']:
            x,y=[Decimal(p['parameters'][k]) for k in ('x','y')]
            if mode=='translate':p['parameters']={'x':str(x-100),'y':str(y-200)}
            elif mode=='units':p.update(parameters={'x':str(x*1000),'y':str(y*1000)},unit='mm')
            elif mode=='swap_axes':p['parameters']={'x':str(y),'y':str(x)}
    g=groups(s,x='-98' if mode=='translate' else '2',y='-198' if mode=='translate' else '2')
    r=partition(s,g);assert sorted(Decimal(f['area']) for f in r['faces'] if f.get('cap_clip_rule'))==[3,7,7,11]


def test_cut_order_equivalent_units_and_unchanged_cells_preserve_identity():
    s=geometry();a=groups(s);a[0]['u_cuts'].append({'value':'4','unit':'m','evidence_id':'test'})
    b=copy.deepcopy(a);b[0]['u_cuts'].reverse();b[0]['v_cuts'][0].update(value='2000',unit='mm')
    def ids(g):return [(f['face_id'],f['area']) for f in partition(s,g)['faces']]
    assert ids(a)==ids(b)


def test_random_staircase_grid_oracle_uses_unit_cell_counts():
    rng=random.Random(6101)
    for _ in range(20):
        w=rng.randint(3,8);hs=[rng.randint(2,7) for _ in range(w)];coords=[(0,0),(w,0),(w,hs[-1])]
        for x in range(w-1,0,-1):
            if hs[x]!=hs[x-1]:coords.extend([(x,hs[x]),(x,hs[x-1])])
        coords.append((0,hs[0]));s=geometry();s['polygon_rings']=[ring('阶梯',coords)]
        xcut=rng.randint(1,w-1);ycut=rng.randint(1,max(hs)-1);r=partition(s,groups(s,x=str(xcut),y=str(ycut)))
        for f in r['faces']:
            if not f.get('cap_clip_rule'):continue
            expected=sum(Decimal(f['u_from'])<=x<Decimal(f['u_to']) and Decimal(f['v_from'])<=y<Decimal(f['v_to']) for x in range(w) for y in range(hs[x]))
            assert Decimal(f['area'])==expected


def test_cap_and_side_partitions_coexist_and_legacy_side_only_stays_identical():
    from tests.test_polygon_surface import geometry as side_geometry
    s=side_geometry();side=side_cuts(s);cap=groups(s)
    from backend.polygon_side_partitions import partition as legacy
    assert partition(s,side)==legacy(s,side,copy.deepcopy(boundary(s)))
    r=partition(s,side+cap);assert r['surface_partition_checks'] and r['surface_cap_partition_checks']
    assert len({c['parameter'] for c in r['surface_subdivisions']})==4
    b=payload(s,cuts=side+cap);b['face_decisions']=[{**d,'included':True} for d in b['face_decisions']]
    out=calculate(b,s)
    assert abs(Decimal(out['quantity'])-Decimal(396))<Decimal('1e-25')
    assert out['rule_version']=='polygon-boundary-1.2'


def test_no_partition_and_removed_partition_preserve_old_results():
    s=geometry();b=surface(s);original=calculate(b,s);assert calculate({**b,'face_subdivisions':[]},s)==original
    assert partition(s,[])==boundary(s) and original['rule_version']=='polygon-boundary-1.0'


@pytest.mark.parametrize('case',['unknown','duplicate_face','empty','fields','outside_x','outside_y','on_low','on_high','duplicate_unit','bad_unit','nan','precision','empty_proof','extra_cut','not_list','too_many'])
def test_invalid_design_windows_fail_closed(case):
    s=geometry();g=groups(s);c=g[0]['u_cuts'][0]
    if case=='unknown':g[0]['face_id']='old'
    elif case=='duplicate_face':g.append(copy.deepcopy(g[0]))
    elif case=='empty':g[0]['u_cuts']=[];g[0]['v_cuts']=[]
    elif case=='fields':g[0]['area']='100'
    elif case=='outside_x':c['value']='7'
    elif case=='outside_y':g[0]['v_cuts'][0]['value']='7'
    elif case=='on_low':c['value']='0'
    elif case=='on_high':c['value']='6'
    elif case=='duplicate_unit':g[0]['u_cuts'].append({'value':'2000','unit':'mm','evidence_id':'test'})
    elif case=='bad_unit':c['unit']='m²'
    elif case=='nan':c['value']='NaN'
    elif case=='precision':c['value']='.0000000000000000001'
    elif case=='empty_proof':c['evidence_id']=' '
    elif case=='extra_cut':c['x']='1'
    elif case=='not_list':g[0]['u_cuts']='bad'
    else:g[0]['u_cuts']*=9
    with pytest.raises(DomainError):partition(s,g)


def test_eight_lines_limit_zero_cells_and_fraction_complexity_guard():
    s=geometry();g=groups(s);g[0]['u_cuts']=[{'value':str(i/10),'unit':'m','evidence_id':'test'} for i in range(1,9)];g[0]['v_cuts']=copy.deepcopy(g[0]['u_cuts'])
    r=partition(s,g);assert len([f for f in r['faces'] if f.get('cap_clip_rule')])==81
    with pytest.raises(DomainError,match='8192'):checked(Fraction(1,1<<8193))


def test_candidate_preview_saved_geometry_review_books_and_boq(clients):
    e,r,pid,s,b,eid=setup(clients);before=get(e,pid)
    c=post(e,pid,'/surface-face-candidates/partition',b).json();assert c['base_faces'][1]['u_axis']=='x' and c['base_faces'][1]['u_to']=='6'
    view=post(e,pid,'/components/preview',b).json();assert Decimal(view['quantity'])==6 and get(e,pid)==before
    comp=post(e,pid,'/components',{**b,'source_geometry':{'fake':'ignored'}}).json();assert comp['source_geometry']==snapshot(s)
    post(e,pid,'/takeoffs/calculate',expected=409);approve(r,pid,'components',comp)
    for t in post(e,pid,'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
    assert any(v['unit']=='m²' and Decimal(v['quantity'])==6 for v in post(e,pid,'/boq/generate').json())


@pytest.mark.parametrize('fault',['bytes','missing','superseded','evidence','coordinate','preview','cut','rational','clip_ring','closure'])
def test_source_and_clip_proof_tampering_rejected_at_review(clients,fault):
    e,r,pid,s,b,eid=setup(clients);comp=post(e,pid,'/components',b).json()
    if fault in ('bytes','missing'):
        path=store.ROOT/'originals'/get(e,pid)['documents'][0]['hash']
        if fault=='bytes':path.write_bytes(b'corrupt')
        else:path.unlink()
    else:
        def damage(p):
            c=p['components'][-1]
            if fault=='superseded':p['documents'][0]['superseded']=True
            elif fault=='evidence':p['evidence'][0]['text']+=' changed'
            elif fault=='coordinate':p['components'][-2]['polygon_rings'][0]['points'][1]['parameters']['x']='7'
            elif fault=='preview':p['components'][-2]['preview']['quantity']='999'
            elif fault=='cut':c['face_subdivisions'][0]['u_cuts'][0]['value']='3'
            elif fault=='rational':next(f for f in c['preview']['surface_faces'] if f.get('cap_clip_rule'))['cap_rational_area']['numerator']='999'
            elif fault=='clip_ring':next(f for f in c['preview']['surface_faces'] if f.get('cap_clip_rule'))['cap_clip_rings'][1]['area']='999'
            else:c['preview']['surface_cap_partition_checks'][0]['exact_rational_tiling']=False
        mutate(pid,damage)
    current=get(e,pid)['components'][-1];res=r.post(f'/api/v1/projects/{pid}/reviews',json={'collection':'components','object_id':comp['id'],'object_hash':store.digest(current),'decision':'approve','reason':'合成校验'})
    assert res.status_code in (409,422),res.text


def test_book_clip_record_tampering_cannot_be_reviewed(clients):
    e,r,pid,s,b,eid=setup(clients);comp=post(e,pid,'/components',b).json();approve(r,pid,'components',comp);book=post(e,pid,'/takeoffs/calculate').json()[-1]
    mutate(pid,lambda p:p['takeoffs'][-1]['surface_cap_partition_checks'][0].update(numerator='999'))
    approve(r,pid,'takeoffs',book,expected=409)


def test_partition_revision_history_atomicity_and_new_boq(clients):
    e,r,pid,s,b,eid=setup(clients);comp=post(e,pid,'/components',b).json();approve(r,pid,'components',comp)
    for t in post(e,pid,'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
    post(e,pid,'/boq/generate');current=get(e,pid)['components'][-1]
    g=groups(s,eid,x='3');changed=payload(s,eid,g);view=partition(s,g)
    changed['face_decisions']=[{**d,'included':bool(f.get('cap_clip_rule')) and Decimal(f['u_to'])<=3 and Decimal(f['v_to'])<=2} for d,f in zip(changed['face_decisions'],view['faces'])]
    changed['object_hash']=store.digest(current);url=f'/api/v1/projects/{pid}/components/'+comp['id']
    res=e.put(url,json=changed);assert res.status_code==200,res.text;p=get(e,pid)
    assert Decimal(res.json()['preview']['quantity'])==8 and p['component_history'][-1]==current and all(t['state']=='stale' for t in p['takeoffs'])
    assert any(v['unit']=='m²' and Decimal(v['quantity'])==6 for v in p['boq']) and e.put(url,json=changed).status_code==409
    approve(r,pid,'components',res.json())
    for t in post(e,pid,'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
    assert any(v['unit']=='m²' and Decimal(v['quantity'])==8 for v in post(e,pid,'/boq/generate').json())


def test_source_revision_permissions_tenant_and_project_scope(clients):
    e,r,pid,s,b,eid=setup(clients)
    post(r,pid,'/surface-face-candidates/partition',b,403);post(clients['other'],pid,'/surface-face-candidates/partition',b,404);post(TestClient(app),pid,'/surface-face-candidates/partition',b,401)
    second=create(e)['id'];post(e,second,'/surface-face-candidates/partition',b,404)
    old=get(e,pid)['components'][0];res=e.put(f'/api/v1/projects/{pid}/components/'+s['id'],json={**copy.deepcopy(old),'object_hash':store.digest(old),'thickness':'300'});assert res.status_code==200,res.text
    approve(r,pid,'components',res.json());post(e,pid,'/surface-face-candidates/partition',b,409)
    from tests.test_project_access import restrict
    restrict(clients['admin'],pid,[]);post(e,pid,'/surface-face-candidates/partition',b,404)


@pytest.mark.parametrize('fmt',['json','xlsx','pdf'])
def test_export_clip_and_closed_area_tables_without_changing_old_quote(clients,fmt):
    e,r,pid,_,quote=priced_project(clients);e,r,pid,s,b,eid=setup(clients,pid)
    comp=post(e,pid,'/components',b).json();approve(r,pid,'components',comp)
    for t in post(e,pid,'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
    base=f'/api/v1/projects/{pid}/exports';res=e.get(base,params={'format':fmt,'unpriced':True});assert res.status_code==200,res.text
    titles=['净面窗口分区校验','净面窗口逐孔扣减','表面分区坐标']
    if fmt=='json':assert all(any(t[0]==title for t in res.json()['tables']) for title in titles)
    elif fmt=='xlsx':assert set(titles)<=set(load_workbook(io.BytesIO(res.content)).sheetnames)
    else:assert all(title in ''.join(p.extract_text() for p in PdfReader(io.BytesIO(res.content)).pages) for title in titles)
    old=e.get(base,params={'format':'json','quote_id':quote['id']}).json();assert old['quote']==quote and not any(t[0] in titles for t in old['tables'])


def test_notch_only_window_at_vertex_is_zero_area():
    s=geometry();r=partition(s,groups(s,x='4',y='4'));check=r['surface_cap_partition_checks'][0]
    assert check['zero_window_count']==1 and check['child_count']==3
    assert sorted(Decimal(f['area']) for f in r['faces'] if f.get('cap_clip_rule'))==[8,8,12]


def test_tiny_exact_window_at_large_origin_preserves_positive_area():
    s=geometry();s['multiplicity']='1'
    with localcontext() as ctx:
        ctx.prec=80;o=Decimal('1000000000');d=Decimal('.00000001')
        s['polygon_rings']=[ring('微小设计方形',[(o,o),(o+2*d,o),(o+2*d,o+2*d),(o,o+2*d)])]
        r=partition(s,groups(s,x=str(o+d),y=str(o+d)))
        assert all(Decimal(f['area'])==d*d for f in r['faces'] if f.get('cap_clip_rule'))
        assert r['surface_cap_partition_checks'][0]['exact_rational_tiling'] is True
