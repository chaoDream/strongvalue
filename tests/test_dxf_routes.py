import copy
import io
from decimal import Decimal
import pytest
from openpyxl import load_workbook
from pypdf import PdfReader
from backend import store
from backend.dxf_geometry import parse
from backend.engine import DomainError
from tests.test_workflow import create,get,upload,approve,priced_project


def polyline(points=None,handle='A',flags=8,vertex_flags=32):
    points=points or [(0,0,0),(3000,4000,12000),(3000,4000,17000)]
    out=[(0,'POLYLINE'),(5,handle),(8,'PIPE'),(70,str(flags))]
    for i,p in enumerate(points):
        out += [(0,'VERTEX'),(5,format(int(handle,16)*1000+i,'X')),(8,'PIPE'),(70,str(vertex_flags)),(10,str(p[0])),(20,str(p[1])),(30,str(p[2]))]
    return out+[(0,'SEQEND'),(5,handle+'FFFF')]


def line():
    return [(0,'LINE'),(5,'B'),(8,'LINE-TEST'),(10,'-1000'),(20,'-2000'),(30,'-3000'),(11,'2000'),(21,'2000'),(31,'9000')]


def dxf(entities=None,unit=4,blocks=None):
    pairs=[(0,'SECTION'),(2,'HEADER')]+([] if unit is None else [(9,'$INSUNITS'),(70,str(unit))])+[(0,'ENDSEC')]
    if blocks:pairs += [(0,'SECTION'),(2,'BLOCKS')]+blocks+[(0,'ENDSEC')]
    pairs += [(0,'SECTION'),(2,'ENTITIES')]+(entities if entities is not None else polyline())+[(0,'ENDSEC'),(0,'EOF')]
    return ''.join(str(c)+'\n'+str(v)+'\n' for c,v in pairs).encode()


def setup(clients,entities=None,unit=4,pid=None):
    c,r=clients['estimator'],clients['reviewer'];pid=pid or create(c)['id']
    upload(c,pid,'合成独立设计依据：共同WCS原点，XYZ方向与设计高程已确认；P1至P2的三维直线距离13m，容差0m；有序设计中心线，不含管件，重复2，总量小数3位。','cad-basis.txt')
    basis=get(c,pid)['evidence'][-1]['id']
    res=c.post(f'/api/v1/projects/{pid}/documents',data={'specialty':'建筑'},files={'file':('design.dxf',dxf(entities,unit),'application/dxf')})
    assert res.status_code==200,res.text
    did=res.json()['document']['id'];url=f'/api/v1/projects/{pid}/documents/{did}/dxf-route-candidates'
    return c,r,pid,did,basis,url


def read(c,url,query='',expected=200):
    r=c.get(url+query);assert r.status_code==expected,r.text
    return r.json()


def body(row,eid,**extra):
    return {'request_id':'synthetic-dxf-request-0001','object_hash':row['object_hash'],'name':'合成DXF设计管线',
            'coordinate_basis':'已核定共同WCS、XYZ方向和设计高程；仅合成测试','measurement_basis':'已核定设计中心线，不计管件和施工附加量；仅数学验收',
            'basis_evidence_id':eid,'calibration_evidence_id':eid,'calibration':{'from_point':1,'to_point':2,'distance':'13','unit':'m','tolerance':'0'},
            'multiplicity':'2','rounding_digits':3,'reason':'按原CAD对象与独立已知长度校核单位，确认设计中心线用途',
            'source_confirmed':True,**extra}


def preview(c,url,row,request,expected=200):
    r=c.post(url+'/'+row['id']+'/preview',json=request);assert r.status_code==expected,r.text
    return r.json()


def adopt(c,url,row,request,view,expected=200):
    r=c.post(url+'/'+row['id']+'/adopt',json={**request,'preview_hash':view['preview_hash']});assert r.status_code==expected,r.text
    return r.json()


def mutate(pid,fn):
    import json
    with store.transaction() as conn:
        p=json.loads(conn.execute('SELECT data FROM projects WHERE id=?',(pid,)).fetchone()['data']);fn(p);store.save(conn,p)


def test_parser_reads_actual_wcs_line_and_3d_polyline_only_from_entities():
    data=parse(dxf(polyline()+line(),blocks=line()))
    assert data['unit']=='mm' and len(data['objects'])==2
    assert all(r['eligible'] for r in data['objects'])
    assert Decimal(data['objects'][0]['single_length_m'])==18
    assert Decimal(data['objects'][1]['single_length_m'])==13
    assert data['objects'][0]['points'][1]=={'x':'3000','y':'4000','z':'12000'}


@pytest.mark.parametrize('entities',[
    polyline(flags=9),polyline(flags=10),polyline(flags=24),polyline(vertex_flags=0),
    polyline(points=[(0,0,0),(1,1,1),(0,0,0)]),polyline(points=[(0,0,0),(2,0,0),(1,0,0)]),
    line()+[(67,'1')],line()+[(410,'Layout1')],line()+[(210,'1')],line()+[(39,'1')],line()+[(60,'1')],
    [p for p in line() if p[0]!=31],line()+[(230,'NaN')],line()+[(230,'sNaN')],
    [(0,'LWPOLYLINE'),(5,'A'),(8,'PIPE')],[(0,'INSERT'),(5,'A'),(8,'PIPE')],[(0,'ARC'),(5,'A'),(8,'PIPE')]])
def test_unsupported_or_ambiguous_geometry_is_not_silently_used(entities):
    rows=parse(dxf(entities))['objects'];assert rows and all(not r['eligible'] and r['issue'] for r in rows)


@pytest.mark.parametrize('unit',[None,0,1,7])
def test_unknown_or_unsupported_insunits_does_not_guess_scale(unit):
    r=parse(dxf(unit=unit));assert r['unit'] is None and not r['objects'][0]['eligible']


@pytest.mark.parametrize('raw',[b'AutoCAD Binary DXF\r\n',b'0\nSECTION\n2\nENTITIES\n',b'0\nEOF\n99\ntrailing\n',dxf(line()+line())])
def test_malformed_or_duplicate_handles_are_rejected(raw):
    with pytest.raises(DomainError):parse(raw)


def test_read_preview_and_adoption_are_source_bound_and_independently_reviewed(clients):
    c,r,pid,did,eid,url=setup(clients);before=get(c,pid)
    data=read(c,url);assert data==read(r,url) and get(c,pid)==before
    row=data['candidates'][0];request=body(row,eid);view=preview(c,url,row,request)
    assert view['calibration']['passed'] and Decimal(view['calibration']['cad_distance_m'])==13
    assert view['result']['quantity']=='36.000' and get(c,pid)==before
    result=adopt(c,url,row,request,view);comp=result['component']
    assert comp['state']=='draft' and len(get(c,pid)['components'])==1 and not get(c,pid)['takeoffs']
    assert comp['dxf_trace']['candidate']['handle']=='A'
    assert comp['dxf_trace']['calibration']['known_distance_m']=='13'
    assert adopt(c,url,row,request,view)['duplicate']
    approve(c,pid,'components',comp,expected=403);approve(r,pid,'components',comp)
    result=c.post(f'/api/v1/projects/{pid}/takeoffs/calculate');assert result.status_code==200,result.text
    assert result.json()[0]['quantity']=='36.000' and result.json()[0]['state']=='draft'
    approve(r,pid,'takeoffs',result.json()[0]);boq=c.post(f'/api/v1/projects/{pid}/boq/generate').json()[0]
    assert boq['quantity']=='36.000' and boq['unit']=='m'


def test_wrong_scale_calibration_is_visible_and_never_auto_corrects(clients):
    c,r,pid,did,eid,url=setup(clients,unit=6);row=read(c,url)['candidates'][0];request=body(row,eid)
    view=preview(c,url,row,request);assert not view['calibration']['passed'] and Decimal(view['calibration']['cad_distance_m'])==13000
    before=get(c,pid);adopt(c,url,row,request,view,409);assert get(c,pid)==before


def test_calibration_is_chord_distance_not_polyline_length_and_inclusive_tolerance(clients):
    c,r,pid,did,eid,url=setup(clients);row=read(c,url)['candidates'][0]
    request=body(row,eid,calibration={'from_point':1,'to_point':2,'distance':'1300.01','unit':'cm','tolerance':'0.01'})
    view=preview(c,url,row,request);assert view['calibration']['passed'] and Decimal(view['calibration']['delta_m'])==Decimal('-.0001')
    request['calibration'].update(distance='18',unit='m',tolerance='0')
    view=preview(c,url,row,request);assert not view['calibration']['passed']


@pytest.mark.parametrize('extra',[
    {'source_confirmed':'true'},{'rounding_digits':True},{'multiplicity':'1.2'},
    {'reason':' '},{'object_hash':'bad'},{'parameters':{}},
    {'calibration':{'from_point':1,'to_point':1,'distance':'13','unit':'m','tolerance':'0'}},
    {'calibration':{'from_point':1,'to_point':4,'distance':'13','unit':'m','tolerance':'0'}},
    {'calibration':{'from_point':1,'to_point':2,'distance':'0','unit':'m','tolerance':'0'}},
    {'calibration':{'from_point':1,'to_point':2,'distance':'13','unit':'m','tolerance':'-1'}}])
def test_invalid_inputs_are_atomic(clients,extra):
    c,r,pid,did,eid,url=setup(clients);row=read(c,url)['candidates'][0];before=get(c,pid)
    preview(c,url,row,body(row,eid,**extra),422);assert get(c,pid)==before


def test_roles_source_confirmation_stale_preview_and_duplicate_source(clients):
    c,r,pid,did,eid,url=setup(clients);row=read(c,url)['candidates'][0];request=body(row,eid);view=preview(c,url,row,request)
    read(clients['other'],url,expected=404);adopt(r,url,row,request,view,403);adopt(clients['other'],url,row,request,view,404)
    adopt(c,url,row,{**request,'source_confirmed':False},view,422)
    adopt(c,url,row,{**request,'name':'changed'},view,409)
    adopt(c,url,row,request,view)
    request['request_id']='synthetic-second-cad-object-001';view=preview(c,url,row,request)
    adopt(c,url,row,request,view,409)


@pytest.mark.parametrize('change',['bytes','deleted','job','document_hash','source_evidence','superseded'])
def test_source_changes_block_adoption_and_saved_route_review(clients,change):
    c,r,pid,did,eid,url=setup(clients);row=read(c,url)['candidates'][0];request=body(row,eid);view=preview(c,url,row,request)
    comp=adopt(c,url,row,request,view)['component'];approve(r,pid,'components',comp)
    def action(p):
        if change=='job':next(j for j in p['jobs'] if j['document_id']==did)['state']='failed'
        elif change=='document_hash':next(d for d in p['documents'] if d['id']==did)['hash']='0'*64
        elif change=='source_evidence':next(e for e in p['evidence'] if e['id']==eid)['text']+='changed'
        elif change=='superseded':next(d for d in p['documents'] if d['id']==did)['superseded']=True
    if change in ('bytes','deleted'):
        path=store.ROOT/'originals'/row['document_hash']
        if change=='bytes':path.write_bytes(b'corrupt')
        else:path.unlink()
    else:mutate(pid,action)
    if change!='source_evidence':read(c,url,expected=409)
    approve(r,pid,'components',comp,expected=409)
    response=c.post(f'/api/v1/projects/{pid}/takeoffs/calculate');assert response.status_code==409


def test_explicit_manual_detachment_and_source_bound_revision(clients):
    c,r,pid,did,eid,url=setup(clients);row=read(c,url)['candidates'][0];request=body(row,eid);view=preview(c,url,row,request)
    comp=adopt(c,url,row,request,view)['component']
    changed=copy.deepcopy(comp);changed['object_hash']=store.digest(comp);changed['route_points'][1]['parameters']['x']='4000'
    assert c.put(f'/api/v1/projects/{pid}/components/'+comp['id'],json=changed).status_code==409
    revised=body(row,eid,request_id='synthetic-rebind-cad-route-001',component_id=comp['id'],component_hash=store.digest(comp))
    new=adopt(c,url,row,revised,preview(c,url,row,revised))['component'];assert new['id']==comp['id'] and new['revision']==2
    changed=copy.deepcopy(new);changed['object_hash']=store.digest(new);changed.pop('dxf_trace')
    changed['route_points'][1]['parameters']['x']='4000'
    assert c.put(f'/api/v1/projects/{pid}/components/'+new['id'],json=changed).status_code==409
    changed.update(detach_dxf=True,detach_reason='对照独立设计依据转为人工坐标修订，保留原CAD采用历史')
    for point in changed['route_points']:point['evidence']={k:eid for k in ('x','y','z')}
    result=c.put(f'/api/v1/projects/{pid}/components/'+new['id'],json=changed);assert result.status_code==200,result.text
    assert 'dxf_trace' not in result.json() and len(get(c,pid)['dxf_route_adoptions'])==2


def test_layer_pagination_and_original_quote_snapshot_exports(clients):
    c,r,pid,_,old=priced_project(clients);c,r,pid,did,eid,url=setup(clients,polyline()+line(),pid=pid)
    data=read(c,url,'?layer=LINE-TEST&limit=1');assert data['total']==1 and data['candidates'][0]['kind']=='LINE'
    assert len(read(c,url,'?offset=1&limit=1')['candidates'])==1
    for q in ['?offset=-1','?limit=0','?limit=201']:read(c,url,q,422)
    row=read(c,url)['candidates'][0];request=body(row,eid);adopt(c,url,row,request,preview(c,url,row,request))
    base=f'/api/v1/projects/{pid}'
    for fmt in ['json','xlsx','pdf']:
        res=c.get(base+'/exports',params={'format':fmt,'unpriced':True});assert res.status_code==200
        if fmt=='json':assert 'DXF坐标采用与单位校核' in str(res.json()['tables'])
        elif fmt=='xlsx':assert 'DXF坐标采用与单位校核' in load_workbook(io.BytesIO(res.content)).sheetnames
        else:assert 'DXF坐标采用与单位校核' in ''.join(p.extract_text() for p in PdfReader(io.BytesIO(res.content)).pages)
    historic=c.get(base+'/exports',params={'format':'json','quote_id':old['id']}).json()
    assert not any(t[0]=='DXF坐标采用与单位校核' for t in historic['tables'])
    assert historic['quote']==old


def test_duplicate_coordinate_groups_are_visible_and_ineligible():
    row=parse(dxf().replace(b'30\n0\n',b'30\n0\n30\n0\n',1))['objects'][0]
    assert not row['eligible'] and '重复' in row['issue']


@pytest.mark.parametrize('field',['trace','route','coordinate_evidence_id','multiplicity_evidence_id','rounding_evidence_id','parameters','clone'])
def test_generic_edit_cannot_rewrite_cad_binding_or_clone_source(clients,field):
    c,r,pid,did,eid,url=setup(clients);row=read(c,url)['candidates'][0];request=body(row,eid)
    comp=adopt(c,url,row,request,preview(c,url,row,request))['component'];changed=copy.deepcopy(comp)
    changed['object_hash']=store.digest(comp)
    if field=='trace':changed['dxf_trace']['reason']='rewritten'
    elif field=='route':changed['route_points'][1]['parameters']['x']='6000'
    elif field=='parameters':changed['parameters']={'extra':'1'}
    elif field!='clone':changed[field]=comp['dxf_trace']['evidence_id']
    before=get(c,pid)
    result=c.post(f'/api/v1/projects/{pid}/components',json=changed) if field=='clone' else c.put(f'/api/v1/projects/{pid}/components/'+comp['id'],json=changed)
    assert result.status_code==409,result.text
    assert get(c,pid)==before


@pytest.mark.parametrize('field',['candidate','calibration_input','proofs','adoption'])
def test_corrupt_saved_trace_blocks_review_without_server_error(clients,field):
    c,r,pid,did,eid,url=setup(clients);row=read(c,url)['candidates'][0];request=body(row,eid)
    comp=adopt(c,url,row,request,preview(c,url,row,request))['component']
    def action(p):
        if field=='adoption':p['dxf_route_adoptions']=[]
        else:
            p['components'][0]['dxf_trace'][field]=None
            p['dxf_route_adoptions'][0]['trace'][field]=None
    mutate(pid,action)
    approve(r,pid,'components',get(c,pid)['components'][0],expected=409)


def test_known_length_cannot_use_current_cad_as_its_own_independent_basis(clients):
    c,r,pid,did,eid,url=setup(clients);row=read(c,url)['candidates'][0]
    own=next(e['id'] for e in get(c,pid)['evidence'] if e['document_id']==did)
    preview(c,url,row,body(row,eid,calibration_evidence_id=own),422)
