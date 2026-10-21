import copy
import io
import json
import pytest
from openpyxl import load_workbook
from pypdf import PdfReader
from backend import store
from tests.test_workflow import priced_project,get,approve


def setup(clients):
    e,r,pid,comp,left=priced_project(clients);base=f'/api/v1/projects/{pid}';approve(r,pid,'quotes',left)
    p=get(e,pid);b=p['boq'][0];eid=p['evidence'][0]['id']
    change=e.post(base+'/changes',json={'boq_id':b['id'],'quantity':'20','reason':'合成数学签证测试，非真实合同变更','evidence_id':eid}).json()
    approve(r,pid,'changes',change)
    response=e.post(base+'/changes/'+change['id']+'/apply');assert response.status_code==200,response.text
    approve(r,pid,'boq',get(e,pid)['boq'][0]);response=e.post(base+'/quotes/calculate');assert response.status_code==200,response.text
    right=response.json();approve(r,pid,'quotes',right)
    preview=e.post(base+'/engineering-settlements/preview',json={'baseline_id':left['id'],'target_id':right['id']});assert preview.status_code==200,preview.text
    result=preview.json()
    body={'name':'合成工程结算差异记录','basis':'仅软件数学验收：同一范围的基线与变更后报价，不是实际合同结算',
          'baseline_id':left['id'],'target_id':right['id'],'preview_hash':result['preview_hash'],'previous_hash':'','request_id':store.uid(),'contract_evidence_id':eid,
          'explanations':[{'key':x['key'],'reason':'仅软件测试，数量或费用随合成签证变化，不代表价款权利认定','evidence_id':eid,'change_ids':[change['id']] if x['kind']=='line' else []} for x in result['rows']]}
    return e,r,pid,left,right,change,body,result


def submit(e,pid,body):
    response=e.post(f'/api/v1/projects/{pid}/engineering-settlements',json=body);assert response.status_code==200,response.text
    return response.json()


def test_settlement_review_math_exports_and_quote_immutability(clients):
    e,r,pid,left,right,change,body,preview=setup(clients);base=f'/api/v1/projects/{pid}';before=get(e,pid)
    assert preview['comparison']['total_delta']=='540.81' and preview['comparison']['attribution']['quantity']=='525.05' and preview['comparison']['fee_delta']=='15.76'
    item=submit(e,pid,body);p=get(e,pid)
    assert p['quotes']==before['quotes'] and p['boq']==before['boq'] and p['report_approval']==before['report_approval']
    assert item['project_name']==before['name'] and item['changes'][0]['id']==change['id'] and item['comparison']['closed']
    assert e.post(base+'/engineering-settlements',json=body).json()==item
    assert not e.get(base+'/engineering-settlements/status').json()['items'][0]['available']
    draft=e.get(base+'/engineering-settlements/'+item['id']+'/exports?format=json');assert '草稿' in draft.json()['status']
    approve(r,pid,'engineering_settlements',item)
    assert e.get(base+'/engineering-settlements/status').json()['items'][0]['available']
    exported=e.get(base+'/engineering-settlements/'+item['id']+'/exports?format=json');assert '已复核' in exported.json()['status']
    x=e.get(base+'/engineering-settlements/'+item['id']+'/exports?format=xlsx');wb=load_workbook(io.BytesIO(x.content));assert '结算差异说明' in wb.sheetnames
    pdf=e.get(base+'/engineering-settlements/'+item['id']+'/exports?format=pdf');assert pdf.content.startswith(b'%PDF')
    assert '540.81' in ''.join(page.extract_text() for page in PdfReader(io.BytesIO(pdf.content)).pages)
    artifacts=e.get(base+'/artifacts').json();artifact=next(a for a in artifacts if a.get('settlement_id')==item['id'] and a['format']=='json' and a['approved'])
    assert e.get(base+'/artifacts/'+artifact['id']+'/file').content==exported.content
    approve(r,pid,'engineering_settlements',get(e,pid)['engineering_settlements'][0],'reject')
    assert get(e,pid)['report_approval']==before['report_approval'] and get(e,pid)['quotes']==before['quotes']
    assert '草稿' in e.get(base+'/engineering-settlements/'+item['id']+'/exports?format=json').json()['status']
    assert e.get(base+'/artifacts/'+artifact['id']+'/file').content==exported.content


@pytest.mark.parametrize('change',[{'name':''},{'basis':''},{'contract_evidence_id':'foreign'},{'preview_hash':'stale'},{'previous_hash':'stale'},{'explanations':[]},
                                  {'explanations':[{'key':'foreign','reason':'test','evidence_id':'foreign'}]}])
def test_submission_validation_is_atomic(clients,change):
    e,r,pid,left,right,signed,body,preview=setup(clients);before=get(e,pid)
    response=e.post(f'/api/v1/projects/{pid}/engineering-settlements',json={**body,**change})
    assert response.status_code in (404,409,422),response.text
    assert get(e,pid)==before


def test_duplicate_explanation_foreign_proof_and_wrong_change_are_rejected(clients):
    e,r,pid,left,right,signed,body,preview=setup(clients);base=f'/api/v1/projects/{pid}/engineering-settlements'
    rows=body['explanations']
    for changes in [[rows[0],rows[0]],[{**rows[0],'evidence_id':'foreign'},rows[1]],[rows[0],{**rows[1],'change_ids':[signed['id']]}],[{**rows[0],'change_ids':['foreign']},rows[1]]]:
        response=e.post(base,json={**body,'explanations':changes});assert response.status_code in (404,409,422),response.text
    assert get(e,pid)['engineering_settlements']==[]


def test_roles_tenant_and_optimistic_review_guards(clients):
    e,r,pid,left,right,signed,body,preview=setup(clients);base=f'/api/v1/projects/{pid}'
    assert r.post(base+'/engineering-settlements',json=body).status_code==403
    assert clients['other'].post(base+'/engineering-settlements/preview',json=body).status_code==404
    item=submit(e,pid,body);approve(e,pid,'engineering_settlements',item,expected=403)
    assert clients['other'].get(base+'/engineering-settlements/status').status_code==404
    assert clients['other'].get(base+'/engineering-settlements/'+item['id']+'/exports').status_code==404
    response=r.post(base+'/reviews',json={'collection':'engineering_settlements','object_id':item['id'],'object_hash':'stale','decision':'approve','reason':'test'});assert response.status_code==409
    assert e.post(base+'/engineering-settlements',json={**body,'basis':'different'}).status_code==409


def test_latest_revision_requires_review_and_old_versions_remain_unchanged(clients):
    e,r,pid,left,right,signed,body,preview=setup(clients);base=f'/api/v1/projects/{pid}';first=submit(e,pid,body);approve(r,pid,'engineering_settlements',first)
    old=copy.deepcopy(get(e,pid)['engineering_settlements'][0]);previous=e.get(base+'/engineering-settlements/status').json()['previous_hash']
    second=submit(e,pid,{**body,'request_id':store.uid(),'previous_hash':previous,'name':'合成修订版'})
    assert second['version']==2 and get(e,pid)['engineering_settlements'][0]==old
    approve(r,pid,'engineering_settlements',old,expected=409);approve(r,pid,'engineering_settlements',second)
    listing=e.get(base+'/engineering-settlements/status').json()['items'];assert not listing[0]['available'] and listing[1]['available']


def test_quote_revocation_and_reapproval_do_not_reactivate_old_record(clients):
    e,r,pid,left,right,signed,body,preview=setup(clients);base=f'/api/v1/projects/{pid}';item=submit(e,pid,body);approve(r,pid,'engineering_settlements',item)
    approve(r,pid,'quotes',right,'reject');assert not e.get(base+'/engineering-settlements/status').json()['items'][0]['available']
    approve(r,pid,'quotes',right)
    approve(r,pid,'engineering_settlements',get(e,pid)['engineering_settlements'][0],expected=409)


def test_new_quote_and_source_edits_expire_old_preview_and_review(clients):
    e,r,pid,left,right,signed,body,preview=setup(clients);base=f'/api/v1/projects/{pid}';item=submit(e,pid,body)
    new=e.post(base+'/quotes/calculate');assert new.status_code==200
    approve(r,pid,'quotes',new.json())
    approve(r,pid,'engineering_settlements',item,expected=409)
    assert e.post(base+'/engineering-settlements',json={**body,'request_id':store.uid(),'previous_hash':store.digest(get(e,pid)['engineering_settlements'][0])}).status_code==409


@pytest.mark.parametrize('side',['baseline','target'])
def test_both_quotes_must_have_valid_approval(clients,side):
    e,r,pid,left,right,signed,body,preview=setup(clients)
    approve(r,pid,'quotes',left if side=='baseline' else right,'reject')
    response=e.post(f'/api/v1/projects/{pid}/engineering-settlements/preview',json=body);assert response.status_code==409


def test_same_or_reversed_quote_selection_is_rejected(clients):
    e,r,pid,left,right,signed,body,preview=setup(clients)
    for a,b in [(left,left),(right,left)]:
        assert e.post(f'/api/v1/projects/{pid}/engineering-settlements/preview',json={'baseline_id':a['id'],'target_id':b['id']}).status_code==409


def test_missing_and_tampered_evidence_never_passes_review(clients):
    e,r,pid,left,right,signed,body,preview=setup(clients);base=f'/api/v1/projects/{pid}';item=submit(e,pid,body)
    doc=get(e,pid)['documents'][0];path=store.ROOT/'originals'/doc['hash'];raw=path.read_bytes();path.write_bytes(b'tampered')
    approve(r,pid,'engineering_settlements',item,expected=409)
    draft=e.get(base+'/engineering-settlements/'+item['id']+'/exports?format=json').json();assert '草稿' in draft['status'] and draft['issue']
    path.write_bytes(raw);path.unlink();approve(r,pid,'engineering_settlements',item,expected=409)


def test_record_integrity_and_excel_formula_text_preservation(clients):
    e,r,pid,left,right,signed,body,preview=setup(clients);base=f'/api/v1/projects/{pid}';body['basis']='=HYPERLINK("https://example.invalid")'+'合成说明'*150
    item=submit(e,pid,body)
    data=e.get(base+'/engineering-settlements/'+item['id']+'/exports?format=xlsx').content;wb=load_workbook(io.BytesIO(data))
    assert all(c.data_type!='f' for s in wb for row in s for c in row)
    with store.transaction() as c:
        p=json.loads(c.execute('SELECT data FROM projects WHERE id=?',(pid,)).fetchone()['data']);p['engineering_settlements'][0]['basis']='tampered';store.save(c,p)
    assert e.get(base+'/engineering-settlements/'+item['id']+'/exports').status_code==409


def test_zero_difference_still_requires_contract_basis_and_review(clients):
    e,r,pid,left,right,signed,body,preview=setup(clients);base=f'/api/v1/projects/{pid}'
    current=e.post(base+'/quotes/calculate').json();approve(r,pid,'quotes',current)
    value=e.post(base+'/engineering-settlements/preview',json={'baseline_id':right['id'],'target_id':current['id']}).json()
    assert value['rows']==[] and value['comparison']['total_delta']=='0.00'
    item=submit(e,pid,{**body,'baseline_id':right['id'],'target_id':current['id'],'preview_hash':value['preview_hash'],'explanations':[]})
    approve(r,pid,'engineering_settlements',item)


def test_quantity_and_price_effects_remain_explicit_with_fee_closure(clients):
    e,r,pid,left,right,signed,body,preview=setup(clients);base=f'/api/v1/projects/{pid}';b=get(e,pid)['boq'][0];quotas=copy.deepcopy(b['quotas']);quotas[0]['resources'][0]['price']='200.005'
    value=e.put(base+'/boq/'+b['id'],json={'quotas':quotas}).json();approve(r,pid,'boq',value)
    target=e.post(base+'/quotes/calculate').json();approve(r,pid,'quotes',target)
    value=e.post(base+'/engineering-settlements/preview',json={'baseline_id':left['id'],'target_id':target['id']}).json()
    assert value['comparison']['attribution']['quantity']=='525.05'
    assert value['comparison']['attribution']['unit_price']=='2100.00'
    assert value['comparison']['fee_delta']=='78.76' and value['comparison']['total_delta']=='2703.81'
    assert value['comparison']['closed']


def test_different_professional_scope_is_not_a_comparable_contract_basis(clients):
    e,r,pid,left,right,signed,body,preview=setup(clients)
    with store.transaction() as c:
        p=json.loads(c.execute('SELECT data FROM projects WHERE id=?',(pid,)).fetchone()['data']);p['quotes'][0]['scope']=['建筑','电气']
        q=p['quotes'][0];q['hash']=store.digest({k:v for k,v in q.items() if k!='hash'});store.save(c,p)
    assert e.post(f'/api/v1/projects/{pid}/engineering-settlements/preview',json=body).status_code==409


def test_project_membership_applies_to_records_and_exports(clients):
    e,r,pid,left,right,signed,body,preview=setup(clients);base=f'/api/v1/projects/{pid}';item=submit(e,pid,body)
    admin=clients['admin'];target=admin.get(base+'/access').json()
    response=admin.put(base+'/access',json={'mode':'restricted','member_ids':[],'object_hash':target['object_hash'],'reason':'合成访问控制验证'});assert response.status_code==200
    assert e.get(base+'/engineering-settlements/status').status_code==404
    assert e.get(base+'/engineering-settlements/'+item['id']+'/exports').status_code==404
    assert admin.get(base+'/engineering-settlements/status').status_code==200
