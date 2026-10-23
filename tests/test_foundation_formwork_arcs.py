"""Synthetic cylindrical contact arcs; one evidenced local XY frame, degrees only."""
import copy,io,random
from decimal import Decimal,localcontext,ROUND_HALF_UP
import pytest
from openpyxl import load_workbook
from pypdf import PdfReader
from backend import store
from backend.compound_takeoff import PI
from backend.engine import DomainError
from backend.foundation_formwork import calculate
from tests.test_foundation_formwork import foundation,body as formwork_body,contact,setup as formwork_setup
from tests.test_post_foundation import equal_body,fixed_body,fixed_added,equal_added,fixed_adjusted,equal_adjusted,combined
from tests.test_route_takeoff import evaluate
from tests.test_polygon_takeoff import post,mutate
from tests.test_workflow import get,approve,upload,create,priced_project

TITLES={'基础模板圆弧角度','基础模板圆弧基准与原件','基础模板圆弧分离校验'}
FRAME='合成共同角度基准：同一基础局部XY圆心，+X对应合成图纸设计东向、+Y为北向；从+Z向原点俯视，0°为+X，从+X向+Y逆时针。仅数学软件验收，无真实工程定位或工法效力。'


def arc(eid='test',start='0',end='90'):
 return dict(parameters=dict(angle_start=start,angle_end=end),unit='°',evidence=dict(angle_start=eid,angle_end=eid),coordinate_basis=FRAME,coordinate_evidence_id=eid,scope_confirmed=True)


def arc_contact(eid='test',start='0',end='90',z_start='0',z_end='200',level=1,name='合成C1'):
 return {**contact(eid,'circumference',z_start,z_end,level,name),'arc_range':arc(eid,start,end),'measurement_basis':'仅合成圆柱指定角度及标高支模，未选范围不计，不含孔壁或水平面，继承全部基础个数，不再乘重复。'}


def body(s,eid='test'):
 b=formwork_body(s,eid);b.update(name='合成圆弧基础模板（非真实工程）',measurement_basis='合成22个圆台阶基础，下层Φ600mm，C1角0至90°、C2角270至360°，均Z0至200mm，每段.66πm²，合计1.32πm²→4.147m²；周期边界相接但无正面积交叠，不自动选小弧或补未选面。');b['contact_ranges']=[arc_contact(eid),arc_contact(eid,'270','360',name='合成C2')];return b


def setup(clients,pid=None,shape='stepped_circular'):
 e,r,pid,s,_,_=formwork_setup(clients,pid,shape)
 upload(e,pid,'合成圆弧模板软件验收，非真实工程或合同。已审22个圆台阶基础，下层实际直径600净高200mm，上层实际直径400净高400mm，同心紧接，材料79.200kg不变，基础计量2.330m3不变。下层整周外侧模板Z0至200mm为2.64πm2→8.294m2。C1角0至90°、C2角270至360°均Z0至200mm，分别.66πm2，共1.32πm2→4.147m2。0°和360°是同一周期边界，只有边界相接无正面积重叠。C1止角改120°后合计1.54πm2→4.838m2。C1止角改300°与C2在270至300°正重叠且标高正相交应拒；如两段标高仅相接可共存。统一使用本原件及共同角度基准说明：同一基础局部XY圆心，+X对应设计东向、+Y北向，从+Z向原点俯视0°为+X、向+Y逆时针为正。0≤起角<止角≤360°，跨0拆分、不求余、不倒向、不取小弧；弧度及无角度原件或未确认阻断。真实是否支模须专业原件核定，以上只数学软件验收。TEST-ARC仅合成模板清单映射码，非官方条目。','foundation-formwork-arcs-basis.txt')
 for c in get(e,pid)['components']:approve(r,pid,'components',c)
 s=get(e,pid)['components'][-1];eid=get(e,pid)['evidence'][-1]['id'];return e,r,pid,s,body(s,eid),eid


def check_asts(r):
 with localcontext() as ctx:
  ctx.prec=80;assert evaluate(r['formula_ast'],r['dimensions'])==Decimal(r['quantity'])
  for d in r['contact_ranges']:
   for k,a in d['formula_asts'].items():assert evaluate(a,r['dimensions'])==Decimal(d['normalized_parameters'][k] if k in ('z_start','z_end') else d[k]),k
  for c in r['arc_checks']:
   for k,a in c['formula_asts'].items():assert evaluate(a,r['dimensions'])==Decimal(c[k])


def test_two_arcs_periodic_seam_formula_balances_source_exact():
 s=foundation(shape='stepped_circular');b=body(s);before=copy.deepcopy((b,s));r=calculate(b,s);assert (b,s)==before and r['quantity']=='4.147' and r['rule_version']=='foundation-outer-formwork-1.1' and r['deduction']=='0';assert [d['total_count'] for d in r['contact_ranges']]==['22','22'];assert r['arc_checks'][0]['state']=='angular_tangent' and Decimal(r['arc_checks'][0]['angle_gap'])==0;check_asts(r)
 with localcontext() as ctx:ctx.prec=80;assert Decimal(r['gross'])==Decimal(PI)*Decimal('1.32')


@pytest.mark.parametrize('shape',['circular','stepped_circular'])
@pytest.mark.parametrize('make',[equal_body,fixed_body,fixed_added,equal_added,fixed_adjusted,equal_adjusted,combined,lambda s:combined(s,equal=True)])
def test_arcs_inherit_eight_final_count_modes_once(shape,make):
 s=foundation(make,shape);b=body(s);r=calculate(b,s)
 with localcontext() as ctx:ctx.prec=80;expected=Decimal(PI)*Decimal('.6')*Decimal('.2')/2*Decimal(r['contact_ranges'][0]['total_count']);assert Decimal(r['gross'])==expected
 check_asts(r)


@pytest.mark.parametrize('unit,factor',[('mm','1'),('cm','.1'),('dm','.01'),('m','.001'),('km','.000001')])
def test_z_units_independent_degree_unit_and_low_context(unit,factor):
 s=foundation(shape='stepped_circular');b=body(s)
 for d in b['contact_ranges']:
  d['unit']=unit
  for k in d['parameters']:d['parameters'][k]=str(Decimal(d['parameters'][k])*Decimal(factor))
 with localcontext() as ctx:ctx.prec=6;r=calculate(b,s)
 assert r==calculate(b,s) and r['quantity']=='4.147';check_asts(r)


@pytest.mark.parametrize('left_end,right_start,accepted,state',[('90','90',True,'angular_tangent'),('89.999999999','90',True,'angular_disjoint'),('90.000000001','90',False,None),('360','0',False,None)])
def test_exact_angle_boundary(left_end,right_start,accepted,state):
 s=foundation(shape='stepped_circular');b=body(s);b['contact_ranges']=[arc_contact(end=left_end),arc_contact(start=right_start,end='180',name='合成C2')]
 if accepted:r=calculate(b,s);assert r['arc_checks'][0]['state']==state;check_asts(r)
 else:
  with pytest.raises(DomainError,match='重叠'):calculate(b,s)


@pytest.mark.parametrize('start,end',[('0','360'),('0','180'),('180','360'),('270','360'),('0.000000000000001','0.000000000000002')])
def test_large_arcs_and_tiny_arcs_never_choose_smaller_arc_or_round_angles(start,end):
 s=foundation(shape='stepped_circular');b=body(s);b['rounding_digits']=12;b['contact_ranges']=[arc_contact(start=start,end=end)];r=calculate(b,s);check_asts(r)
 with localcontext() as ctx:ctx.prec=80;span=Decimal(end)-Decimal(start);expected=Decimal('.6')*Decimal(PI)*span/360*Decimal('.2')*22;assert Decimal(r['gross'])==expected and Decimal(r['quantity'])==expected.quantize(Decimal('1e-12'),rounding=ROUND_HALF_UP)


@pytest.mark.parametrize('z_start,state',[('100','vertical_tangent'),('100.000000001','vertically_disjoint')])
def test_angle_overlap_allowed_only_when_z_nonpositive_intersection(z_start,state):
 s=foundation(shape='stepped_circular');b=body(s);b['contact_ranges']=[arc_contact(end='300',z_end='100'),arc_contact(start='270',end='360',z_start=z_start,name='合成C2')];r=calculate(b,s);assert r['arc_checks'][0]['state']==state;check_asts(r)
 b['contact_ranges'][1]['parameters']['z_start']='99.999999999'
 with pytest.raises(DomainError,match='重叠'):calculate(b,s)


def test_full_circle_claim_and_adjacent_arc_subdivision_equal_full():
 s=foundation(shape='stepped_circular');b=body(s);full=contact(face='circumference',end='200');b['contact_ranges']=[full];old=calculate(b,s);b['contact_ranges']=[arc_contact(start=str(i*90),end=str((i+1)*90),name='合成C'+str(i)) for i in range(4)];r=calculate(b,s);assert r['quantity']==old['quantity']=='8.294';check_asts(r)
 b['contact_ranges'].append({**full,'name':'整周重复'})
 with pytest.raises(DomainError,match='重叠'):calculate(b,s)
 b['contact_ranges']=[arc_contact(start='0',end='360')];assert calculate(b,s)['quantity']==old['quantity']


def test_different_layers_same_angles_and_frame_but_independent_diameters():
 s=foundation(shape='stepped_circular');b=body(s);b['contact_ranges']=[arc_contact(),arc_contact(level=2,z_start='200',z_end='600',name='合成C2')];r=calculate(b,s);assert r['quantity']=='4.838' and not r['arc_checks'];check_asts(r)
 b['contact_ranges'][1]['arc_range']['coordinate_basis']='另一角度方向'
 with pytest.raises(DomainError,match='相同角度基准'):calculate(b,s)


def test_mixed_rectangular_full_circle_and_arcs_only_selected_ranges():
 s=foundation(shape='stepped_circular');m=s['source_geometry'];m['material_items'].append({**copy.deepcopy(m['material_items'][0]),'name':'合成第二立柱项'});s['source_component_hash']=store.digest({k:v for k,v in m.items() if k!='state'});d=foundation()['foundation_items'][0];d['material_index']=2;s['foundation_items'].append(d);b=body(s);b['contact_ranges'].append({**contact(name='合成矩形侧'),'material_index':2});r=calculate(b,s);assert r['quantity']=='10.549' and len(r['contact_ranges'])==3;check_asts(r)


def test_seeded_independent_area_oracle_and_rounding():
 rng=random.Random(9401)
 for _ in range(70):
  s=foundation(shape='stepped_circular');s['foundation_items'][0].pop('through_holes');s['foundation_items'][0]['levels'][0]['parameters']['diameter']=str(rng.randint(400,999));b=body(s);start=rng.randint(0,358);end=rng.randint(start+1,360);bottom=rng.randint(0,198);top=rng.randint(bottom+1,200);b['contact_ranges']=[arc_contact(start=str(start),end=str(end),z_start=str(bottom),z_end=str(top))];b['rounding_digits']=rng.randint(0,12);r=calculate(b,s)
  with localcontext() as ctx:ctx.prec=80;expected=Decimal(s['foundation_items'][0]['levels'][0]['parameters']['diameter'])/1000*Decimal(PI)*(end-start)/360*Decimal(top-bottom)/1000*22;assert Decimal(r['gross'])==expected and Decimal(r['quantity'])==expected.quantize(Decimal(1).scaleb(-b['rounding_digits']),rounding=ROUND_HALF_UP)
  check_asts(r)


@pytest.mark.parametrize('fault',['null','empty','extra','missing','params','proofs','start-proof','end-proof','frame','frame-proof','unconfirmed','truthy','unit','radian','negative','reverse','zero','over360','nan','inf','bool','precision','length-instead','rectangular','frame-different','frame-proof-different'])
def test_strict_arc_config(fault):
 s=foundation(shape='stepped_circular');b=body(s);c=b['contact_ranges'][0]['arc_range']
 if fault=='null':b['contact_ranges'][0]['arc_range']=None
 elif fault=='empty':b['contact_ranges'][0]['arc_range']={}
 elif fault=='extra':c['radius']='300'
 elif fault=='missing':c.pop('unit')
 elif fault=='params':c['parameters'].pop('angle_start')
 elif fault=='proofs':c['evidence'].pop('angle_start')
 elif fault in ('start-proof','end-proof'):c['evidence']['angle_start' if fault=='start-proof' else 'angle_end']=' '
 elif fault=='frame':c['coordinate_basis']=''
 elif fault=='frame-proof':c['coordinate_evidence_id']=''
 elif fault in ('unconfirmed','truthy'):c['scope_confirmed']=False if fault=='unconfirmed' else 1
 elif fault in ('unit','radian'):c['unit']='mm' if fault=='unit' else 'rad'
 elif fault in ('negative','reverse','zero','over360','nan','inf','bool','precision'):c['parameters']['angle_start' if fault=='negative' else 'angle_end']={'negative':'-1','reverse':'-10','zero':'0','over360':'360.000000001','nan':'NaN','inf':'Infinity','bool':True,'precision':'90.0000000000000000001'}[fault]
 elif fault=='length-instead':c['parameters']['arc_length']='600'
 elif fault=='rectangular':s=foundation();b['source_component_hash']=store.digest({k:v for k,v in s.items() if k!='state'});b['contact_ranges'][0]['face']='x_plus'
 elif fault=='frame-different':b['contact_ranges'][1]['arc_range']['coordinate_basis']='另一个角度基准'
 elif fault=='frame-proof-different':b['contact_ranges'][1]['arc_range']['coordinate_evidence_id']='other'
 with pytest.raises(DomainError):calculate(b,s)


def test_preview_claims_frames_across_records_and_forged_id(clients):
 e,r,pid,s,b,eid=setup(clients);before=get(e,pid);assert post(e,pid,b,'/components/preview').json()['quantity']=='4.147' and get(e,pid)==before
 first=copy.deepcopy(b);first['contact_ranges']=first['contact_ranges'][:1];c=post(e,pid,first).json();approve(r,pid,'components',c);second=copy.deepcopy(b);second['contact_ranges']=second['contact_ranges'][1:];d=post(e,pid,second).json();approve(r,pid,'components',d);post(e,pid,{**first,'id':c['id']},expected=409)
 second['contact_ranges'][0]['arc_range']['parameters']['angle_start']='89.999999999';post(e,pid,second,expected=409)
 second=copy.deepcopy(b);second['contact_ranges']=[arc_contact(eid,'90','180',name='合成C3')];second['contact_ranges'][0]['arc_range']['coordinate_basis']='方向基准变化';post(e,pid,second,expected=409)
 second['contact_ranges'][0]['arc_range']['coordinate_basis']=FRAME;assert post(e,pid,second).status_code==200


@pytest.mark.parametrize('fault',['angle','frame','proof','confirm','width','span','periodic-gap','formula','book'])
def test_tamper_rebuild(clients,fault):
 e,r,pid,s,b,eid=setup(clients);c=post(e,pid,b).json()
 if fault=='book':approve(r,pid,'components',c);post(e,pid,{},'/takeoffs/calculate')
 def forge(p):
  c=p['components'][-1]
  if fault=='angle':c['contact_ranges'][0]['arc_range']['parameters']['angle_end']='120'
  elif fault=='frame':
   for d in c['contact_ranges']:d['arc_range']['coordinate_basis']='改写共同基准'
  elif fault=='proof':c['contact_ranges'][0]['arc_range']['evidence']['angle_start']='forged'
  elif fault=='confirm':c['contact_ranges'][0]['arc_range']['scope_confirmed']=False
  elif fault=='width':c['preview']['contact_ranges'][0]['width']='1'
  elif fault=='span':c['preview']['contact_ranges'][0]['angle_span']='1'
  elif fault=='periodic-gap':c['preview']['arc_checks'][0]['angle_gap']='999'
  elif fault=='formula':c['preview']['contact_ranges'][0]['formula_asts']['angle_span']={'constant':'1'}
  else:p['takeoffs'][-1]['quantity']='999'
 mutate(pid,forge);coll='takeoffs' if fault=='book' else 'components';approve(r,pid,coll,get(e,pid)[coll][-1],expected=404 if fault=='proof' else 422 if fault=='confirm' else 409)


@pytest.mark.parametrize('gate',['preview','save','review','calculate','book_review'])
def test_angle_original_bytes_five_gates(clients,gate):
 e,r,pid,s,b,eid=setup(clients)
 if gate in ('review','calculate','book_review'):c=post(e,pid,b).json()
 if gate in ('calculate','book_review'):approve(r,pid,'components',c)
 if gate=='book_review':t=post(e,pid,{},'/takeoffs/calculate').json()[-1]
 doc=get(e,pid)['documents'][-1];(store.ROOT/'originals'/doc['hash']).write_bytes(b'corrupt')
 if gate=='preview':post(e,pid,b,'/components/preview',409)
 elif gate=='save':post(e,pid,b,expected=409)
 elif gate=='review':approve(r,pid,'components',c,expected=409)
 elif gate=='calculate':post(e,pid,{},'/takeoffs/calculate',409)
 else:approve(r,pid,'takeoffs',t,expected=409)


def test_revision_history_books_boq_concurrency_disable_and_source_stale(clients):
 e,r,pid,s,b,eid=setup(clients);c=post(e,pid,b).json();approve(r,pid,'components',c)
 for t in post(e,pid,{},'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
 post(e,pid,{},'/boq/generate');p=get(e,pid);old=p['components'][-1];new=copy.deepcopy(b);new['object_hash']=store.digest(old);new['contact_ranges'][0]['arc_range']['parameters']['angle_end']='120';url=f'/api/v1/projects/{pid}/components/'+old['id'];res=e.put(url,json=new);assert res.status_code==200 and res.json()['preview']['quantity']=='4.838';assert get(e,pid)['component_history'][-1]==old and get(e,pid)['boq']==p['boq'] and all(t['state']=='stale' for t in get(e,pid)['takeoffs']);assert e.put(url,json=new).status_code==409;approve(r,pid,'components',res.json())
 new['contact_ranges']=[contact(eid,face='circumference',end='200')];new['object_hash']=store.digest(get(e,pid)['components'][-1]);res=e.put(url,json=new);assert res.status_code==200 and res.json()['preview']['quantity']=='8.294' and res.json()['preview']['rule_version']=='foundation-outer-formwork-1.0';approve(r,pid,'components',res.json())
 s['object_hash']=store.digest(s);s['foundation_items'][0]['levels'][1]['parameters']['height']='500';res=e.put(f'/api/v1/projects/{pid}/components/'+s['id'],json=s);assert res.status_code==200;approve(r,pid,'components',res.json());assert get(e,pid)['components'][-1]['state']=='stale'


def test_roles_tenants_and_foreign_angle_evidence(clients):
 e,r,pid,s,b,eid=setup(clients);post(r,pid,b,expected=403);assert clients['other'].get(f'/api/v1/projects/{pid}').status_code==404;other=create(e)['id'];upload(e,other,'异项目圆弧角度依据');b['contact_ranges'][0]['arc_range']['evidence']['angle_start']=get(e,other)['evidence'][-1]['id'];post(e,pid,b,expected=404)


@pytest.mark.parametrize('shape',['circular','stepped_circular'])
@pytest.mark.parametrize('fmt',['json','xlsx','pdf'])
def test_three_exports_boq_and_old_quote(clients,shape,fmt):
 e,r,pid,_,q=priced_project(clients);e,r,pid,s,b,eid=setup(clients,pid,shape);c=post(e,pid,b).json();approve(r,pid,'components',c)
 for t in post(e,pid,{},'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
 boq=post(e,pid,{},'/boq/generate').json();assert any(d['quantity']=='4.147' and d['unit']=='m²' for d in boq)
 url=f'/api/v1/projects/{pid}/exports';res=e.get(url,params=dict(format=fmt,unpriced=True));assert res.status_code==200,res.text
 if fmt=='json':assert TITLES<={d[0] for d in res.json()['tables']}
 elif fmt=='xlsx':assert TITLES<=set(load_workbook(io.BytesIO(res.content)).sheetnames)
 else:assert all(v in ''.join(p.extract_text() for p in PdfReader(io.BytesIO(res.content)).pages) for v in TITLES)
 assert not any(d[0] in TITLES for d in e.get(url,params=dict(format='json',quote_id=q['id'])).json()['tables'])


def test_many_arcs_long_frame_and_proofs_pdf(clients):
 e,r,pid,s,b,eid=setup(clients);b['contact_ranges']=[arc_contact(eid,str(i*18),str((i+1)*18),name='合成C'+str(i)) for i in range(20)]
 for d in b['contact_ranges']:d['arc_range']['coordinate_basis']='合成'+('共同图纸角度基准及范围说明'*130)
 c=post(e,pid,b).json();assert c['preview']['quantity']=='8.294';approve(r,pid,'components',c)
 for t in post(e,pid,{},'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
 res=e.get(f'/api/v1/projects/{pid}/exports',params=dict(format='pdf',unpriced=True));assert res.status_code==200,res.text
