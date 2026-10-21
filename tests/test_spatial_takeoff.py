"""Synthetic 3D geometry and workflow verification; no professional measurement claim."""
import copy
import io
import json
import random
from decimal import Decimal, localcontext
import pytest
from openpyxl import load_workbook
from pypdf import PdfReader
from backend import store
from backend.engine import DomainError
from backend.spatial_takeoff import calculate
from tests.test_compound_takeoff import ast_value
from tests.test_workflow import create, get, upload, approve, priced_project


def region(name, x, y, z, length, width, height, operation='add', eid='test', unit='m'):
    fields = ['x', 'y', 'z', 'length', 'width', 'height']
    return {'name': name, 'operation': operation, 'parameters': dict(zip(fields, map(str, [x, y, z, length, width, height]))),
            'unit': unit, 'evidence': {k: eid for k in fields+['measurement_basis']}, 'measurement_basis': '仅合成长方体算例，非专业计量规则'}


def payload(eid='test'):
    return {'name': '合成错层构件（非真实工程）', 'kind': 'spatial_regions', 'parameters': {}, 'quantity_type': 'volume',
            'regions': [region('主体', 0, 0, 0, 4, 3, 2, eid=eid), region('翼块', 3, 0, 1, 2, 2, 2, eid=eid),
                        region('扣减A', 1, 1, '.5', 1, 1, 1, 'subtract', eid), region('扣减B', '1.5', 1, 1, 1, 1, 1, 'subtract', eid)],
            'multiplicity': '1', 'multiplicity_evidence_id': eid, 'coordinate_basis': '同一设计原点，X右Y上，Z为共同高程基准向上；仅合成测试',
            'coordinate_evidence_id': eid, 'measurement_basis': '长方体并集减扣减并集；相交部分仅扣一次；无专业效力', 'basis_evidence_id': eid}


def setup(clients):
    e, r = clients['estimator'], clients['reviewer']; pid = create(e)['id']
    upload(e, pid, '仅合成：共同原点X右Y上Z向上。主体XYZ(0,0,0)长宽高4×3×2m，翼块(3,0,1)2×2×2m；扣减A(1,1,0.5)1×1×1m，扣减B(1.5,1,1)1×1×1m；整组1。')
    return e, r, pid, payload(get(e, pid)['evidence'][0]['id'])


def test_different_elevations_and_overlapping_additions_and_deductions():
    body = payload(); before = copy.deepcopy(body); result = calculate(body); s = result['spatial_summary']
    assert body == before
    assert Decimal(s['naive_add_volume']) == 32 and Decimal(s['add_union_volume']) == 30 and Decimal(s['add_overlap_volume']) == 2
    assert Decimal(s['naive_deduct_volume']) == 2 and Decimal(s['deduct_union_volume']) == Decimal('1.75')
    assert Decimal(s['deduct_overlap_volume']) == Decimal('.25') and Decimal(result['quantity']) == Decimal('28.25')
    assert result['unit'] == 'm³' and len(result['spatial_intersections']) == 2 and len(result['spatial_layers']) == 5
    assert sum(Decimal(z['net_volume']) for z in result['spatial_layers']) == Decimal('28.25')
    assert [Decimal(z['net_volume']) for z in result['spatial_layers']] == list(map(Decimal, ['6','5.5','6.25','6.5','4']))
    with localcontext() as ctx:
        ctx.prec = 80
        assert ast_value(result['formula_ast'], result['dimensions']) == Decimal(result['quantity'])


def test_same_xy_at_separate_elevations_is_not_overlap():
    body = payload(); body['regions'] = [region('bottom',0,0,0,2,2,1), region('top',0,0,2,2,2,1)]
    result = calculate(body)
    assert not result['spatial_intersections'] and Decimal(result['quantity']) == 8
    assert len(result['spatial_layers']) == 2


def test_multiple_coverage_is_not_sum_of_pairwise_intersections():
    body = payload(); body['regions'] = [region(str(i),0,0,0,2,2,2) for i in range(3)]+[region('hole'+str(i),0,0,0,1,1,1,'subtract') for i in range(3)]
    r = calculate(body); s = r['spatial_summary']
    assert Decimal(s['add_overlap_volume']) == 16 and Decimal(s['deduct_overlap_volume']) == 2
    assert sum(Decimal(x['volume']) for x in r['spatial_intersections'] if x['operation']=='add') == 24
    assert Decimal(r['quantity']) == 7


def test_negative_xyz_mixed_units_translation_and_repeat():
    body = payload(); body['regions'] = [region('left',-2000,-1000,-2000,4000,3000,2000,unit='mm'),region('right',100,-100,-100,200,200,200,unit='cm')]; body['multiplicity']='2'
    r = calculate(body)
    assert Decimal(r['quantity']) == 60 and r['spatial_regions'][0]['normalized_parameters']['z'] == '-2'
    with localcontext() as ctx:
        ctx.prec=80
        assert ast_value(r['formula_ast'],r['dimensions']) == 60
    body['regions'] = [region('left',98,99,98,4,3,2),region('right',101,99,99,2,2,2)]
    assert Decimal(calculate(body)['quantity']) == 60


def test_touching_faces_and_full_deduction():
    body=payload(); body['regions']=[region('a',0,0,0,1,1,1),region('b',1,0,0,1,1,1),region('c',0,0,1,1,1,1)]
    r=calculate(body); assert not r['spatial_intersections'] and Decimal(r['quantity'])==3
    body['regions'] += [region('bottom',0,0,0,2,1,1,'subtract'),region('top',0,0,1,1,1,1,'subtract')]
    assert Decimal(calculate(body)['quantity'])==0


def test_deduction_can_cross_vertical_addition_joint():
    body=payload(); body['regions']=[region('a',0,0,0,2,2,1),region('b',0,0,1,2,2,1),region('ded',0,0,'.5',1,1,1,'subtract')]
    assert Decimal(calculate(body)['quantity'])==7


def test_matches_independent_integer_voxel_oracle():
    rng=random.Random(2601)
    for run in range(40):
        body=payload(); body['regions']=[]; cells=set(); cut=set()
        for i in range(rng.randint(1,7)):
            x,y,z=[rng.randint(-3,3) for _ in range(3)]; l,w,h=[rng.randint(1,3) for _ in range(3)]
            body['regions'].append(region('add'+str(i),x,y,z,l,w,h))
            cells |= {(a,b,c) for a in range(x,x+l) for b in range(y,y+w) for c in range(z,z+h)}
        for i in range(rng.randint(0,6)):
            x,y,z=rng.choice(sorted(cells)); body['regions'].append(region('ded'+str(i),x,y,z,1,1,1,'subtract')); cut.add((x,y,z))
        rng.shuffle(body['regions']); r=calculate(body)
        assert Decimal(r['quantity'])==len(cells-cut)
        with localcontext() as ctx:
            ctx.prec=80
            assert ast_value(r['formula_ast'],r['dimensions'])==len(cells-cut)


@pytest.mark.parametrize('case',['outside_xy','outside_z','z_gap','all_deduct','negative_height','zero_height','nan','unit','precision','duplicate','missing_z_proof','too_many','thickness','count','overflow','rotation','area_mode'])
def test_invalid_models_fail_explicitly(case):
    body=payload(); r=body['regions'][0]
    if case=='outside_xy':body['regions'][2]['parameters']['x']='20'
    elif case=='outside_z':body['regions'][2]['parameters']['z']='-.1'
    elif case=='z_gap':body['regions']=[region('a',0,0,0,1,1,1),region('b',0,0,2,1,1,1),region('ded',0,0,0,1,1,3,'subtract')]
    elif case=='all_deduct':
        for x in body['regions']:x['operation']='subtract'
    elif case=='negative_height':r['parameters']['height']='-1'
    elif case=='zero_height':r['parameters']['height']='0'
    elif case=='nan':r['parameters']['z']='NaN'
    elif case=='unit':r['unit']='m³'
    elif case=='precision':r['unit']='mm';r['parameters']['z']='.000000000000000001'
    elif case=='duplicate':body['regions'][1]['name']=r['name']
    elif case=='missing_z_proof':del r['evidence']['z']
    elif case=='too_many':body['regions']*=6
    elif case=='thickness':body['thickness']='1'
    elif case=='count':body['multiplicity']='1.5'
    elif case=='overflow':body['regions']=[region('huge',0,0,0,'1e12','1e12','1e12')]
    elif case=='area_mode':body['quantity_type']='area'
    else:r['parameters']['rotation']='30'
    with pytest.raises(DomainError):calculate(body)


def test_preview_approval_boq_and_all_exports(clients):
    e,r,pid,body=setup(clients);base='/api/v1/projects/'+pid;before=get(e,pid)
    response=e.post(base+'/components/preview',json=body)
    assert response.status_code==200,response.text
    assert get(e,pid)==before
    comp=e.post(base+'/components',json=body)
    assert comp.status_code==200,comp.text
    assert e.post(base+'/takeoffs/calculate').status_code==409
    approve(r,pid,'components',comp.json())
    response=e.post(base+'/takeoffs/calculate');assert response.status_code==200,response.text
    book=response.json()[0];approve(r,pid,'takeoffs',book)
    assert Decimal(e.post(base+'/boq/generate').json()[0]['quantity'])==Decimal('28.25')
    result=e.get(base+'/exports?unpriced=true&format=json').json()
    assert result['unpriced']['takeoffs'][0]['coordinate_basis']==body['coordinate_basis']
    exported=e.get(base+'/exports?unpriced=true&format=xlsx');assert exported.status_code==200,exported.text
    wb=load_workbook(io.BytesIO(exported.content),data_only=False)
    assert {'空间并集与扣减','长方体原始坐标','空间分层计算','空间分层条带','空间相交定位','算量原始单位换算'}<=set(wb.sheetnames)
    assert all(c.data_type!='f' for s in wb for row in s for c in row)
    pdf=e.get(base+'/exports?unpriced=true&format=pdf');assert pdf.status_code==200,pdf.text
    text='\n'.join(page.extract_text() for page in PdfReader(io.BytesIO(pdf.content)).pages)
    assert '空间并集与扣减' in text and '空间分层条带' in text


def test_concurrent_revision_atomicity_and_old_book_invalidation(clients):
    e,r,pid,body=setup(clients);base='/api/v1/projects/'+pid
    comp=e.post(base+'/components',json=body).json();approve(r,pid,'components',comp)
    book=e.post(base+'/takeoffs/calculate').json()[0];approve(r,pid,'takeoffs',book)
    before=get(e,pid);assert e.put(base+'/components/'+comp['id'],json=body).status_code==409
    body['object_hash']=e.get(base+'/hashes').json()['components'][comp['id']]
    bad=copy.deepcopy(body);bad['regions'][2]['parameters']['x']='100'
    assert e.put(base+'/components/'+comp['id'],json=bad).status_code==422 and get(e,pid)==before
    body['regions'][1]['parameters']['height']='3'
    revised=e.put(base+'/components/'+comp['id'],json=body)
    assert revised.status_code==200,revised.text
    assert revised.json()['revision']==2 and Decimal(revised.json()['preview']['quantity'])==Decimal('32.25')
    assert get(e,pid)['takeoffs'][0]['state']=='stale'
    assert e.put(base+'/components/'+comp['id'],json=body).status_code==409


@pytest.mark.parametrize('gate',['save','component_review','calculate','book_review'])
def test_corrupt_original_blocks_all_processing_gates(clients,gate):
    e,r,pid,body=setup(clients);base='/api/v1/projects/'+pid
    if gate!='save':comp=e.post(base+'/components',json=body).json()
    if gate in ('calculate','book_review'):approve(r,pid,'components',comp)
    if gate=='book_review':book=e.post(base+'/takeoffs/calculate').json()[0]
    doc=get(e,pid)['documents'][0];(store.ROOT/'originals'/doc['hash']).write_bytes(b'corrupt')
    if gate=='save':assert e.post(base+'/components',json=body).status_code==409
    elif gate=='component_review':approve(r,pid,'components',comp,expected=409)
    elif gate=='calculate':assert e.post(base+'/takeoffs/calculate').status_code==409
    else:approve(r,pid,'takeoffs',book,expected=409)


def test_permissions_foreign_proof_and_superseded_source(clients):
    e,r,pid,body=setup(clients);base='/api/v1/projects/'+pid
    assert r.post(base+'/components/preview',json=body).status_code==403
    assert clients['other'].post(base+'/components',json=body).status_code==404
    pid2=create(e)['id'];upload(e,pid2,'其他工程');eid=get(e,pid2)['evidence'][0]['id']
    foreign=copy.deepcopy(body);foreign['regions'][0]['evidence']['x']=eid
    assert e.post(base+'/components',json=foreign).status_code==404
    comp=e.post(base+'/components',json=body).json();approve(r,pid,'components',comp)
    upload(e,pid,'新版本坐标',name='fixture.txt')
    assert get(e,pid)['components'][0]['state']=='stale'
    approve(r,pid,'components',comp,expected=409)


def test_cannot_drop_structured_fields_or_mutate_calculation_book(clients):
    e,r,pid,body=setup(clients);base='/api/v1/projects/'+pid
    comp=e.post(base+'/components',json=body).json();approve(r,pid,'components',comp)
    e.post(base+'/takeoffs/calculate')
    with store.transaction() as c:
        p=json.loads(c.execute('SELECT data FROM projects WHERE id=?',(pid,)).fetchone()['data']);p['takeoffs'][0]['spatial_summary']['net_volume']='999';store.save(c,p)
    approve(r,pid,'takeoffs',get(e,pid)['takeoffs'][0],expected=409)
    changed={'kind':'slab','parameters':{},'object_hash':e.get(base+'/hashes').json()['components'][comp['id']]}
    assert e.put(base+'/components/'+comp['id'],json=changed).status_code==409


def test_new_spatial_geometry_never_rewrites_quote_snapshot(clients):
    e,r,pid,_,q=priced_project(clients);base='/api/v1/projects/'+pid;old=copy.deepcopy(q)
    eid=get(e,pid)['evidence'][0]['id'];response=e.post(base+'/components',json=payload(eid))
    assert response.status_code==200,response.text
    assert get(e,pid)['quotes'][0]==old and q['id'] in get(e,pid)['stale_quotes']
