"""Synthetic mixed-length packing, explicit proof status and source isolation."""
import copy,io,random
from decimal import Decimal,localcontext
from fractions import Fraction
from collections import Counter
import pytest
from openpyxl import load_workbook
from pypdf import PdfReader
from backend import store,mixed_stock_cutting as mixed
from backend.engine import DomainError
from backend.railing_material import calculate,snapshot
from tests.test_railing_material import geometry,payload,member,setup as material_setup
from tests.test_railing_posts import body as equal_body
from tests.test_fixed_posts import body as fixed_body
from tests.test_additional_posts import body as fixed_added
from tests.test_equal_additional_posts import body as equal_added
from tests.test_post_adjustments import body as fixed_adjusted
from tests.test_equal_post_adjustments import body as equal_adjusted
from tests.test_combined_post_adjustments import body as combined
from tests.test_post_spacing import body as local_body
from tests.test_path_post_spacing import body as path_body
from tests.test_stock_allowance import config as allowance
from tests.test_stock_cutting import config as cutting
from tests.test_panel_material import panel
from tests.test_route_takeoff import evaluate
from tests.test_railing_takeoff import interval
from tests.test_polygon_takeoff import post,mutate
from tests.test_workflow import get,approve,upload,create,priced_project

PARAMS=('stock_length','kerf_per_piece','trim_start','trim_end','reusable_min_length')
TITLES={'混合下料汇总','混合下料需求','混合下料原材明细','混合下料逐件位置','混合下料长度质量','混合下料依据与边界'}


def config(eid='test',members=None,stock='6000'):
 return dict(name='合成同规格方管混排',members=members or ['合成长料','合成短料'],stock_length=stock,kerf_per_piece='3',trim_start='10',trim_end='10',reusable_min_length='50',length_unit='mm',material_basis='合成同材质同截面方管A、每米3kg，可混排，非真实工程',measurement_basis='单件下料完整，每件一次锯缝含最后一件，两端切除未重复，不使用库存或增加净清单',compatibility_confirmed=True,cut_basis_confirmed=True,evidence={k:eid for k in (*PARAMS,'material_basis','measurement_basis')})


def body(s,eid='test'):
 b=payload(s,eid);b.update(name='合成不同长度混合下料（非真实工程）',measurement_basis='长料每组2件×3.5m×3kg/m＋短料每组2件×2.4m×3kg/m=35.4kg，来源重复2，净70.8kg；混排附表不增加净量或报价')
 b['material_items']=[{**member('合成长料','cut_pieces',eid),'quantity':'2','cut_length':'3500','measurement_basis':'合成每组2件，单件完整下料3500mm，每米3000g，同规格方管A'},{**member('合成短料','cut_pieces',eid),'quantity':'2','cut_length':'2400','measurement_basis':'合成每组2件，单件完整下料2400mm，每米3000g，同规格方管A'}];b['mixed_cutting']=[config(eid)];return b


def setup(clients,pid=None):
 e,r,pid,s,_,_=material_setup(clients,pid)
 upload(e,pid,'合成混排软件验收，非真实工程。长料每组2件3500mm、短料每组2件2400mm，两项均同材质同截面方管A、每米3000g，允许混排。来源重复2，长短各4件，共8件净23.6m70.8kg。单件下料完整，每件一次3mm锯缝含最后一件，原材每根两端各切除10mm，未重复；留用余料下限50mm。原材6000mm时4根，每根长短各1件，余74mm待留用，长度下界4，已证最少。购入24m72kg=净23.6＋锯缝.024＋切除.08＋待留用.296m，差额0。不自动入库或增加净清单。原材改5900mm，长短同根放不下，需6根；长度下界5但搜索穷尽证明6，购入35.4m106.2kg，净70.8kg不变须重新复核。','mixed-stock-cutting-basis.txt')
 for c in get(e,pid)['components']:approve(r,pid,'components',c)
 s=get(e,pid)['components'][-1];eid=get(e,pid)['evidence'][-1]['id'];return e,r,pid,s,body(s,eid),eid


def result(s,b):return calculate(b,s)['mixed_stock_cutting']['groups'][0]


def test_exact_mixed_plan_balance_count_asts_and_no_net_mutation():
 s=geometry();b=body(s);before=copy.deepcopy((s,b));old=copy.deepcopy(b);old.pop('mixed_cutting');base=calculate(old,s);v=calculate(b,s);g=v['mixed_stock_cutting']['groups'][0]
 assert (s,b)==before and v['quantity']==base['quantity']=='70.800' and v['formula_ast']==base['formula_ast'] and v['material_items']==base['material_items'] and v['rule_version']=='railing-material-mass-1.14'
 assert g['proof']['status']=='optimal_bound' and g['proof']['lower_bound']==g['stock_count']==4 and g['required_piece_count']==8 and g['balance_difference']=='0'
 assert g['total']==dict(net_length='23.6',kerf_length='0.024',trim_length='0.08',scrap_length='0',retained_length='0.296',loss_length='0.104',purchased_length='24',net_mass='70.8',kerf_mass='0.072',trim_mass='0.24',scrap_mass='0',retained_mass='0.888',loss_mass='0.312',purchased_mass='72')
 assert len([k for k in v['source_evidence'] if '.mixed_cutting.' in k])==7
 seen=[]
 for bar in g['bars']:
  assert bar['piece_count']==2 and bar['offcut_length']=='0.074' and bar['offcut_disposition']=='retain'
  a,z=bar['pieces'];assert (a['from_s'],a['to_s'],a['kerf_to_s'])==('0.01','3.51','3.513') and (z['from_s'],z['to_s'],z['kerf_to_s'])==('3.513','5.913','5.916');seen.extend((p['item_index'],p['piece_index']) for p in bar['pieces'])
  for k,ast in bar['formula_asts'].items():assert evaluate(ast,v['dimensions'])==Decimal(bar['values'][k])
 assert sorted(seen)==[(i,j) for i in (1,2) for j in range(1,5)]
 for d in g['demands']:assert evaluate(d['count_formula_ast'],v['dimensions'])==d['piece_count']
 for k,ast in g['total_formula_asts'].items():assert evaluate(ast,v['dimensions'])==Decimal(g['total'][k])
 for k,ast in v['mixed_stock_cutting']['total_formula_asts'].items():assert evaluate(ast,v['dimensions'])==Decimal(v['mixed_stock_cutting']['total'][k])


def test_search_exhaustion_proves_more_than_volume_lower_bound():
 s=geometry();b=body(s);b['mixed_cutting'][0]['stock_length']='5900';g=result(s,b);p=g['proof'];assert p['status']=='optimal_search' and p['search_exhausted'] and p['lower_bound']==5 and p['stock_count']==6 and p['visited_nodes']>0 and p['unproven_gap']==0
 assert g['total']['purchased_mass']=='106.2' and g['total']['retained_length']=='11.656'


def test_search_improves_first_fit_and_budget_status_is_explicit():
 sizes=[6,5,3,2,2,2];bins,p=mixed.pack(sizes,10);assert p['initial_stock_count']==3 and p['stock_count']==2 and p['status']=='optimal_bound' and p['visited_nodes']>0
 bins,p=mixed.pack(sizes,10,limit=0);assert p['stock_count']==3 and p['status']=='feasible_unproven' and not p['search_exhausted'] and p['unproven_gap']==1
 assert sorted(j for b in bins for j in b)==list(range(6)) and all(sum(sizes[j] for j in b)<=10 for b in bins)


def test_unproven_status_survives_complete_result_and_export(clients,monkeypatch):
 monkeypatch.setattr(mixed,'SEARCH_LIMIT',0);e,r,pid,s,b,eid=setup(clients);b['mixed_cutting'][0]['stock_length']='5900';c=post(e,pid,b).json();assert c['preview']['mixed_stock_cutting']['groups'][0]['proof']['status']=='feasible_unproven';approve(r,pid,'components',c)
 for t in post(e,pid,{},'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
 res=e.get(f'/api/v1/projects/{pid}/exports',params=dict(format='json',unpriced=True));assert res.status_code==200 and '可行但尚未证明最优' in str(res.json()['tables'])


@pytest.mark.parametrize('make',[equal_body,fixed_body,fixed_added,equal_added,fixed_adjusted,equal_adjusted,combined,lambda s:combined(s,equal=True),local_body,path_body])
def test_all_post_modes_mix_with_fixed_pieces_without_changing_layout(make):
 s=geometry();b=make(s);m=b['material_items'][0];m['name']='合成立柱';b['material_items'].append({**member('合成短料','cut_pieces'),'quantity':'2','cut_length':'2400'});old=calculate(copy.deepcopy(b),s);b['mixed_cutting']=[config(members=['合成立柱','合成短料'])];v=calculate(b,s);assert v['quantity']==old['quantity'] and v['material_items']==old['material_items'] and v['mixed_stock_cutting']['covered_item_count']==2


@pytest.mark.parametrize('unit,factor',[('m',Decimal('.001')),('mm',Decimal('1')),('cm',Decimal('.1')),('dm',Decimal('.01')),('km',Decimal('.000001'))])
def test_five_parameter_units(unit,factor):
 s=geometry();b=body(s);g=b['mixed_cutting'][0];g['length_unit']=unit
 for k in PARAMS:g[k]=str(Decimal(g[k])*factor)
 assert result(s,b)['total']['purchased_mass']=='72'


@pytest.mark.parametrize('weight,unit',[('3','kg/m'),('3000','g/m'),('.003','kg/mm'),('3','g/mm')])
def test_compatible_weight_units(weight,unit):
 s=geometry();b=body(s);b['material_items'][1].update(unit_weight=weight,weight_unit=unit);assert result(s,b)['normalized_unit_weight']=='3'


@pytest.mark.parametrize('threshold,status',[('74','retain'),('74.000000001','scrap'),('73.999999999','retain')])
def test_exact_offcut_threshold(threshold,status):
 s=geometry();b=body(s);b['mixed_cutting'][0]['reusable_min_length']=threshold;assert all(x['offcut_disposition']==status for x in result(s,b)['bars'])


def test_zero_offcut_and_identical_lengths_cardinality_bound():
 s=geometry();b=body(s)
 for m in b['material_items']:m['cut_length']='3000'
 c=b['mixed_cutting'][0];c.update(stock_length='6000',kerf_per_piece='0',trim_start='0',trim_end='0');g=result(s,b);assert g['stock_count']==4 and g['proof']['cardinality_lower_bound']==4 and all(x['offcut_disposition']=='none' for x in g['bars'])


@pytest.mark.parametrize('fault',['null','empty','dict','too-many-groups','missing','extra','proof','extra-proof','blank-proof','material-basis','basis','compatibility','cut-confirmation','truthy','members-type','one-member','many-members','duplicate-member','unknown','cross-group','duplicate-name','weight','single-cut','allowance','continuous','panel','negative','zero-stock','zero-threshold','unit','precision','nan','trim','overlong','group-count'])
def test_invalid_configuration_and_overlap_are_blocked(fault):
 s=geometry();b=body(s);c=b['mixed_cutting'][0]
 if fault=='null':b['mixed_cutting']=None
 elif fault=='empty':b['mixed_cutting']=[]
 elif fault=='dict':b['mixed_cutting']={}
 elif fault=='too-many-groups':b['mixed_cutting']=[c]*11
 elif fault=='missing':c.pop('stock_length')
 elif fault=='extra':c['force_optimal']=True
 elif fault=='proof':c['evidence'].pop('material_basis')
 elif fault=='extra-proof':c['evidence']['scope']='test'
 elif fault=='blank-proof':c['evidence']['stock_length']=''
 elif fault=='material-basis':c['material_basis']=''
 elif fault=='basis':c['measurement_basis']=''
 elif fault=='compatibility':c['compatibility_confirmed']=False
 elif fault=='cut-confirmation':c['cut_basis_confirmed']=False
 elif fault=='truthy':c['compatibility_confirmed']=1
 elif fault=='members-type':c['members']=[{},'合成短料']
 elif fault=='one-member':c['members']=['合成长料']
 elif fault=='many-members':c['members']=['x']*9
 elif fault=='duplicate-member':c['members']=['合成长料']*2
 elif fault=='unknown':c['members']=['不存在','合成短料']
 elif fault=='cross-group':b['mixed_cutting'].append({**copy.deepcopy(c),'name':'另一组'})
 elif fault=='duplicate-name':b['mixed_cutting'].append(copy.deepcopy(c))
 elif fault=='weight':b['material_items'][1]['unit_weight']='3001'
 elif fault=='single-cut':b['material_items'][0]['cutting_plan']=cutting()
 elif fault=='allowance':b['material_items'][0]['stock_allowance']=allowance()
 elif fault=='continuous':b['material_items'][0]=member('合成长料','continuous')
 elif fault=='panel':b['material_items'][0]=panel();c['members']=[b['material_items'][0]['name'],'合成短料']
 elif fault=='negative':c['kerf_per_piece']='-1'
 elif fault=='zero-stock':c['stock_length']='0'
 elif fault=='zero-threshold':c['reusable_min_length']='0'
 elif fault=='unit':c['length_unit']='m²'
 elif fault=='precision':c['stock_length']='6000.000000000000001'
 elif fault=='nan':c['trim_start']='NaN'
 elif fault=='trim':c['trim_start']='5990'
 elif fault=='overlong':c['stock_length']='3500'
 else:b['material_items'][0]['quantity']='101'
 with pytest.raises(DomainError):calculate(b,s)


def test_name_binding_is_stable_when_member_selection_order_changes():
 s=geometry();b=body(s);g=result(s,b);b['mixed_cutting'][0]['members'].reverse();h=result(s,b);assert g['bars']==h['bars'] and g['proof']==h['proof']
 b['material_items'].reverse();h=result(s,b);assert {d['name']:d['piece_count'] for d in h['demands']}=={'合成长料':4,'合成短料':4}


def test_max_200_pieces_and_table_400_limit():
 s=geometry();b=body(s)
 for m in b['material_items']:m['quantity']='50'
 assert result(s,b)['required_piece_count']==200
 b2=copy.deepcopy(b)
 for m in b2['material_items']:m['name']+='二'
 b['material_items']+=copy.deepcopy(b2['material_items']);b['mixed_cutting'].append(config(members=[m['name'] for m in b2['material_items']]));b['mixed_cutting'][1]['name']='第二组';assert calculate(b,s)['mixed_stock_cutting']['covered_item_count']==4
 for m in b2['material_items']:m['name']+='三'
 b['material_items']+=copy.deepcopy(b2['material_items']);b['mixed_cutting'].append(config(members=[m['name'] for m in b2['material_items']]));b['mixed_cutting'][2]['name']='第三组'
 with pytest.raises(DomainError,match='400'):calculate(b,s)


def test_seeded_independent_bruteforce_optimum_and_length_oracle():
 rng=random.Random(8101)
 def oracle(sizes,cap):
  best=len(sizes);bins=[]
  def visit(i):
   nonlocal best
   if i==len(sizes):best=min(best,len(bins));return
   seen=set()
   for j,used in enumerate(bins):
    if used in seen or used+sizes[i]>cap:continue
    seen.add(used);bins[j]+=sizes[i];visit(i+1);bins[j]-=sizes[i]
   if len(bins)+1<best:bins.append(sizes[i]);visit(i+1);bins.pop()
  visit(0);return best
 for _ in range(50):
  cap=rng.randint(6,20);sizes=sorted([rng.randint(1,cap) for _ in range(rng.randint(2,8))],reverse=True);bins,p=mixed.pack(sizes,cap);assert len(bins)==oracle(sizes,cap) and p['status']!='feasible_unproven';assert all(sum(sizes[j] for j in b)<=cap for b in bins) and sorted(j for b in bins for j in b)==list(range(len(sizes)))


def test_low_caller_precision_and_legacy_absence():
 s=geometry();b=body(s)
 with localcontext() as c:c.prec=6;v=calculate(b,s)
 assert v==calculate(b,s);b.pop('mixed_cutting');v=calculate(b,s);assert 'mixed_stock_cutting' not in v and v['rule_version']=='railing-material-mass-1.0'


@pytest.mark.parametrize('fmt',['json','xlsx','pdf'])
def test_full_lifecycle_three_exports_net_boq_old_quote(clients,fmt):
 e,r,pid,_,quote=priced_project(clients);e,r,pid,s,b,eid=setup(clients,pid);before=get(e,pid);post(e,pid,b,'/components/preview');assert get(e,pid)==before;c=post(e,pid,b).json();assert c['mixed_cutting']==b['mixed_cutting'];approve(r,pid,'components',c)
 for t in post(e,pid,{},'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
 assert any(v['quantity']=='70.800' for v in post(e,pid,{},'/boq/generate').json());url=f'/api/v1/projects/{pid}/exports';res=e.get(url,params=dict(format=fmt,unpriced=True));assert res.status_code==200,res.text
 if fmt=='json':assert TITLES<={t[0] for t in res.json()['tables']}
 elif fmt=='xlsx':assert TITLES<=set(load_workbook(io.BytesIO(res.content)).sheetnames)
 else:assert all(t in ''.join(p.extract_text() for p in PdfReader(io.BytesIO(res.content)).pages) for t in TITLES)
 old=e.get(url,params=dict(format='json',quote_id=quote['id'])).json();assert old['quote']==quote and not any(t[0] in TITLES for t in old['tables'])


@pytest.mark.parametrize('fault',['input','proof','status','assignment','position','count','ledger','ast','summary','book'])
def test_saved_raw_and_derived_mixed_records_rebuilt(clients,fault):
 e,r,pid,s,b,eid=setup(clients);c=post(e,pid,b).json()
 if fault=='book':approve(r,pid,'components',c);post(e,pid,{},'/takeoffs/calculate')
 def forge(p):
  c=p['components'][-1];g=c['preview']['mixed_stock_cutting']['groups'][0]
  if fault=='input':c['mixed_cutting'][0]['stock_length']='5900'
  elif fault=='proof':c['mixed_cutting'][0]['evidence']['material_basis']='forged'
  elif fault=='status':g['proof']['status']='feasible_unproven'
  elif fault=='assignment':g['bars'][0]['pieces'][0]['name']='伪造'
  elif fault=='position':g['bars'][0]['pieces'][0]['to_s']='3.5'
  elif fault=='count':g['demands'][0]['piece_count']=3
  elif fault=='ledger':g['total']['retained_length']='0'
  elif fault=='ast':g['total_formula_asts']['purchased_mass']={'constant':'70.8'}
  elif fault=='summary':c['preview']['mixed_stock_cutting']['stock_count']=1
  else:p['takeoffs'][-1]['mixed_stock_cutting']['groups'][0]['proof']['status']='forged'
 mutate(pid,forge);coll='takeoffs' if fault=='book' else 'components';approve(r,pid,coll,get(e,pid)[coll][-1],expected=404 if fault=='proof' else 409)


def test_same_net_revision_history_stale_books_removal_and_source(clients):
 e,r,pid,s,b,eid=setup(clients);c=post(e,pid,b).json();approve(r,pid,'components',c)
 for t in post(e,pid,{},'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
 post(e,pid,{},'/boq/generate');before=get(e,pid);old=before['components'][-1];new=copy.deepcopy(b);new['mixed_cutting'][0]['stock_length']='5900';new['object_hash']=store.digest(old);url=f'/api/v1/projects/{pid}/components/'+old['id'];res=e.put(url,json=new);assert res.status_code==200,res.text;p=get(e,pid);assert p['component_history'][-1]==old and p['boq']==before['boq'] and all(t['state']=='stale' for t in p['takeoffs']) and res.json()['preview']['quantity']=='70.800';assert e.put(url,json=new).status_code==409;approve(r,pid,'components',res.json())
 for t in post(e,pid,{},'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
 assert any(v['quantity']=='70.800' for v in post(e,pid,{},'/boq/generate').json())
 current=get(e,pid)['components'][-1];new['object_hash']=store.digest(current);new.pop('mixed_cutting');res=e.put(url,json=new);assert res.status_code==200 and 'mixed_cutting' not in res.json() and 'mixed_stock_cutting' not in res.json()['preview'];assert get(e,pid)['component_history'][-1]==current;approve(r,pid,'components',res.json())
 s['object_hash']=store.digest(s);s['edge_decisions'][0].update(mode='ranges',ranges=[interval(0,10,eid)]);res=e.put(f'/api/v1/projects/{pid}/components/'+s['id'],json=s);assert res.status_code==200;approve(r,pid,'components',res.json());assert get(e,pid)['components'][-1]['state']=='stale'


@pytest.mark.parametrize('gate',['preview','save','review','calculate','book_review'])
def test_mixed_original_bytes_gate(clients,gate):
 e,r,pid,s,b,eid=setup(clients)
 if gate in ('review','calculate','book_review'):c=post(e,pid,b).json()
 if gate in ('calculate','book_review'):approve(r,pid,'components',c)
 if gate=='book_review':t=post(e,pid,{},'/takeoffs/calculate').json()[-1]
 d=get(e,pid)['documents'][-1];(store.ROOT/'originals'/d['hash']).write_bytes(b'corrupt')
 if gate=='preview':post(e,pid,b,'/components/preview',409)
 elif gate=='save':post(e,pid,b,expected=409)
 elif gate=='review':approve(r,pid,'components',c,expected=409)
 elif gate=='calculate':post(e,pid,{},'/takeoffs/calculate',409)
 else:approve(r,pid,'takeoffs',t,expected=409)


def test_mixed_roles_tenant_foreign_evidence(clients):
 e,r,pid,s,b,eid=setup(clients);post(r,pid,b,expected=403);assert clients['other'].get(f'/api/v1/projects/{pid}').status_code==404
 other=create(e)['id'];upload(e,other,'异项目');b['mixed_cutting'][0]['evidence']['material_basis']=get(e,other)['evidence'][-1]['id'];post(e,pid,b,expected=404)


def test_long_names_pdf_and_material_overlap_export(clients):
 e,r,pid,s,b,eid=setup(clients);b['name']='合成'+('长名称验收'*39)+'末';b['mixed_cutting'][0]['name']='合成'+('长混排组名'*39)+'末';c=post(e,pid,b).json();approve(r,pid,'components',c)
 for t in post(e,pid,{},'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
 res=e.get(f'/api/v1/projects/{pid}/exports',params=dict(format='pdf',unpriced=True));assert res.status_code==200,res.text
