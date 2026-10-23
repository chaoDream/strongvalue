"""Synthetic connected-path spacing; independent geometry and lifecycle checks."""
import copy,io,random
from decimal import Decimal,localcontext
from fractions import Fraction
import pytest
from openpyxl import load_workbook
from pypdf import PdfReader
from backend import store
from backend.engine import DomainError
from backend.railing_material import calculate
from backend.railing_takeoff import snapshot
from backend.path_post_spacing import sign,normalize,expression
from tests.test_post_spacing import check as local_check,setup as local_setup,body as moved_body
from tests.test_railing_posts import geometry,body as equal_body
from tests.test_fixed_posts import body as fixed_body
from tests.test_additional_posts import body as fixed_added,line_source
from tests.test_equal_additional_posts import body as equal_added
from tests.test_post_adjustments import body as fixed_adjusted
from tests.test_equal_post_adjustments import body as equal_adjusted
from tests.test_combined_post_adjustments import body as combined
from tests.test_route_takeoff import point,evaluate
from tests.test_route_railing import payload as railing_payload
from tests.test_railing_takeoff import interval,geometry as polygon,payload as polygon_railing
from tests.test_polygon_takeoff import post,mutate
from tests.test_workflow import get,approve,priced_project,upload,create


def check(limit='2500',unit='mm',eid='test'):
 c=local_check(limit,unit,eid);c.update(scope='connected_path',measurement_basis='合成沿已审连续安装路径核定；跨转角、相接区间及闭环首尾，缺口断开，开端及无柱路径另核');return c


def body(s,eid='test'):
 b=moved_body(s,eid);b['name']='合成连续路径柱距（非真实工程）';b['material_items'][0]['spacing_check']=check(eid=eid);return b


def setup(clients,pid=None):
 e,r,pid,s,_,_=local_setup(clients,pid)
 upload(e,pid,'合成连续路径柱距验收，非真实工程。已审AB13m与BC5m连续连接，B端柱只归AB，沿三维来源路径核查11根10对，AB末柱到BC首柱5/3m，最大2.5m，上限2500mm。跨转角取沿路径距离，不取直线弦长；缺口断开，开端和无柱路径另核；各根不变，总79.2kg。范围变更及同质量上限修订仍须独立复核。','path-post-spacing-basis.txt')
 for c in get(e,pid)['components']:approve(r,pid,'components',c)
 s=get(e,pid)['components'][-1];eid=get(e,pid)['evidence'][-1]['id'];return e,r,pid,s,body(s,eid),eid


def result(s,b):return calculate(b,s)['material_items'][0]['post_spacing']


@pytest.mark.parametrize('make',[equal_body,fixed_body,fixed_added,equal_added,fixed_adjusted,equal_adjusted,combined,lambda s:combined(s,equal=True)])
def test_eight_modes_preserve_layout_and_mass_and_connect_shared_vertex(make):
 s=geometry();b=make(s);old=calculate(copy.deepcopy(b),s);b['material_items'][0]['spacing_check']=check();before=copy.deepcopy((s,b));v=calculate(b,s);g=v['material_items'][0]['post_spacing'];p=g['paths'][0]
 assert (s,b)==before and v['quantity']==old['quantity'] and v['formula_ast']==old['formula_ast'] and v['rule_version']=='railing-material-mass-1.11'
 assert v['material_items'][0]['post_layout']==old['material_items'][0]['post_layout'] and len(g['paths'])==1 and not p['closed'] and p['pair_count']==p['post_count']-1
 assert sum(a['cross_range'] for a in p['pairs'])==1 and p['unposted_start_distance']==p['unposted_end_distance']=='0'
 with localcontext() as c:
  c.prec=80
  for a in p['pairs']:assert abs(evaluate(a['distance_formula_ast'],v['dimensions'])-Decimal(a['distance']))<Decimal('1e-75')


def test_corner_with_unowned_vertex_uses_arc_not_chord_and_blocks_limit():
 s=geometry();b=equal_body(s);m=b['material_items'][0];m['post_ranges'][0]['end_post']=False;m['spacing_check']=check('4','m');g=result(s,b);a=next(a for a in g['paths'][0]['pairs'] if a['cross_range'])
 assert Fraction(74,21)-Fraction(a['distance'])<Fraction(1,10**75) and Decimal(a['distance'])>Decimal('3.5') and len(a['segments'])==2
 m['spacing_check']['max_spacing']='3.5'
 with pytest.raises(DomainError,match='连续路径.*超过'):calculate(b,s)


@pytest.mark.parametrize('basis,expected',[('spatial','3.5238095238'),('horizontal','3.3333333333')])
def test_source_selected_length_basis(basis,expected):
 s=geometry(basis=basis);b=equal_body(s);m=b['material_items'][0];m['post_ranges'][0]['end_post']=False;m['spacing_check']=check('4','m');a=next(a for a in result(s,b)['paths'][0]['pairs'] if a['cross_range']);assert a['distance'].startswith(expected)


@pytest.mark.parametrize('touch',[True,False])
def test_same_edge_touch_connects_and_gap_breaks(touch):
 s=line_source(12);s['edge_decisions'][0].update(mode='ranges',ranges=[interval(0,4),interval(4 if touch else 7,12)]);b=equal_body(s);m=b['material_items'][0];m['post_ranges'][0]['end_post']=False;m['post_ranges'][1]['start_post']=False;m['spacing_check']=check('4','m');g=result(s,b)
 assert len(g['paths'])==(1 if touch else 2)
 assert sum(a['cross_range'] for p in g['paths'] for a in p['pairs'])==int(touch)
 if touch:assert next(a for a in g['paths'][0]['pairs'] if a['cross_range'])['distance']=='4'
 else:assert all(not a['cross_range'] for p in g['paths'] for a in p['pairs'])


def closed_source():
 src=polygon();s=polygon_railing(src);s.update(id='polygon',revision=1,source_geometry=snapshot(src))
 for d in s['edge_decisions']:d.update(mode='full',ranges=[])
 return s


@pytest.mark.parametrize('starts',[False,True])
def test_polygon_and_hole_separate_closed_rings_wrap_exactly_once(starts):
 s=closed_source();b=equal_body(s);m=b['material_items'][0]
 for r in m['post_ranges']:r.update(start_post=starts,end_post=not starts)
 m['spacing_check']=check('2','m');g=result(s,b);assert len(g['paths'])==2
 for p in g['paths']:
  assert p['closed'] and p['pair_count']==p['post_count'] and p['unposted_start_distance'] is None and sum(a['wrap'] for a in p['pairs'])==1 and len(p['range_ids'])==3
 assert g['pair_count']==calculate(b,s)['material_items'][0]['post_layout']['post_count']


def test_closed_one_root_means_full_circuit_and_no_root_ring_is_disclosed():
 s=closed_source();b=equal_body(s);m=b['material_items'][0]
 for i,r in enumerate(m['post_ranges']):r.update(spacing='100',spacing_unit='m',start_post=i==0,end_post=False)
 m['spacing_check']=check('24','m');g=result(s,b);p=next(p for p in g['paths'] if p['post_count']);assert p['closed'] and p['pair_count']==1 and p['pairs'][0]['from_post']==p['pairs'][0]['to_post'] and p['pairs'][0]['distance']=='24' and len(p['pairs'][0]['segments'])==4
 assert any(p['fully_unposted'] and p['status']=='no_pairs' for p in g['paths'])
 m['spacing_check']['max_spacing']='23.999999999999'
 with pytest.raises(DomainError,match='闭环首尾.*超过'):calculate(b,s)


def test_excluded_edge_breaks_ring_but_selected_seam_stays_connected():
 s=closed_source();s['edge_decisions'][1].update(mode='exclude');b=equal_body(s);m=b['material_items'][0]
 for r in m['post_ranges']:r.update(start_post=True,end_post=False)
 m['spacing_check']=check('2','m');g=result(s,b);p=next(p for p in g['paths'] if not p['closed']);assert p['range_labels']==['R1 E3 整段','R1 E1 整段'] and not any(a['wrap'] for a in p['pairs']) and sum(a['cross_range'] for a in p['pairs'])==1


def test_empty_selected_range_inside_path_is_traversed_not_skipped():
 s=line_source(6);s['edge_decisions'][0].update(mode='ranges',ranges=[interval(0,2),interval(2,4),interval(4,6)]);b=equal_body(s);m=b['material_items'][0]
 for i,r in enumerate(m['post_ranges']):r.update(spacing='10',spacing_unit='m',start_post=i==0,end_post=i==2)
 m['spacing_check']=check('6','m');p=result(s,b)['paths'][0];assert p['post_count']==2 and p['pairs'][0]['distance']=='6' and len(p['pairs'][0]['segments'])==3
 m['spacing_check']['max_spacing']='5.999'
 with pytest.raises(DomainError,match='超过'):calculate(b,s)


def test_no_open_pairs_rejected_and_partial_boundary_gaps_are_separate():
 s=line_source(6);b=equal_body(s);m=b['material_items'][0];r=m['post_ranges'][0];r.update(spacing='10',spacing_unit='m',start_post=False);m['spacing_check']=check('10','m')
 with pytest.raises(DomainError,match='没有可核查'):calculate(b,s)
 r.update(spacing='2',start_post=False,end_post=False);p=result(s,b)['paths'][0];assert p['pair_count']==1 and p['unposted_start_distance']==p['unposted_end_distance']=='2'


@pytest.mark.parametrize('bad',['local','unknown','',None,{},True])
def test_explicit_unknown_scope_rejected(bad):
 s=geometry();b=body(s);b['material_items'][0]['spacing_check']['scope']=bad
 with pytest.raises(DomainError,match='连续路径'):calculate(b,s)


@pytest.mark.parametrize('fault',['missing','extra','no-proof','no-basis','zero','negative','unit'])
def test_strict_new_schema_and_evidence(fault):
 s=geometry();b=body(s);c=b['material_items'][0]['spacing_check']
 if fault=='missing':c.pop('max_spacing')
 elif fault=='extra':c['default']=True
 elif fault=='no-proof':c['evidence'].pop('measurement_basis')
 elif fault=='no-basis':c['measurement_basis']=''
 elif fault=='zero':c['max_spacing']='0'
 elif fault=='negative':c['max_spacing']='-1'
 else:c['spacing_unit']='cm2'
 with pytest.raises(DomainError):calculate(b,s)


@pytest.mark.parametrize('unit,value',[('m','2.5'),('mm','2500'),('cm','250'),('dm','25'),('km','.0025')])
def test_cap_units_equality_and_tiny_exceedance(unit,value):
 s=geometry();b=body(s);b['material_items'][0]['spacing_check']=check(value,unit);assert result(s,b)['normalized_limit']=='2.5'
 b['material_items'][0]['spacing_check']=check('2.49999999999999999','m')
 with pytest.raises(DomainError,match='超过'):calculate(b,s)


def test_multi_radical_cross_edge_certified_limit_and_equivalent_classes():
 s=line_source(4);src=s['source_geometry'];src['route_points']=[point('A',0,0,0),point('B',1,1,0),point('C',2,2,1)];s['source_component_hash']=store.digest(snapshot(src));s['edge_decisions']=railing_payload(src)['edge_decisions'];b=equal_body(s);m=b['material_items'][0]
 for i,r in enumerate(m['post_ranges']):r.update(spacing='10',spacing_unit='m',start_post=i==0,end_post=i==1)
 m['spacing_check']=check('3.14626436994197234','m')
 with pytest.raises(DomainError,match='超过'):calculate(b,s)
 m['spacing_check']['max_spacing']='3.14626436994197235';a=result(s,b)['paths'][0]['pairs'][0];assert len(a['exact_distance'])==2 and '√' in expression(a['exact_distance'])
 assert normalize([(Fraction(1),Fraction(8)),(Fraction(-2),Fraction(2))])==[]
 assert sign([(Fraction(1),Fraction(2)),(Fraction(-1),Fraction(3))])==-1
 assert sign([(Fraction(1),Fraction(2)),(Fraction(-1),Fraction(2)),(Fraction(1,10**180),Fraction(1))])==1


def test_seeded_independent_rational_path_oracle():
 rng=random.Random(7801)
 for _ in range(50):
  lengths=[rng.randint(3,20) for _ in range(rng.randint(2,6))];s=line_source(4);src=s['source_geometry'];xs=[0]
  for length in lengths:xs.append(xs[-1]+length)
  src['route_points']=[point(str(i),x,0,0) for i,x in enumerate(xs)];s['source_component_hash']=store.digest(snapshot(src));s['edge_decisions']=railing_payload(src)['edge_decisions'];b=equal_body(s);m=b['material_items'][0];coords=[]
  for i,(r,L) in enumerate(zip(m['post_ranges'],lengths)):
   step=rng.randint(1,5);n=(L+step-1)//step;start=i==0;end=rng.choice([True,False]);r.update(spacing=str(step),spacing_unit='m',start_post=start,end_post=end)
   coords.extend(Fraction(xs[i])+Fraction(L*j,n) for j in range(n+1) if (j!=0 or start) and (j!=n or end))
  m['spacing_check']=check('100','m');p=result(s,b)['paths'][0];expected=[b-a for a,b in zip(coords,coords[1:])];assert len(expected)==p['pair_count']
  for a,oracle in zip(p['pairs'],expected):
   assert len(a['exact_distance'])==1;t=a['exact_distance'][0];assert t['radicand']==dict(numerator='1',denominator='1') and Fraction(int(t['coefficient']['numerator']),int(t['coefficient']['denominator']))==oracle


@pytest.mark.parametrize('fmt',['json','xlsx','pdf'])
def test_workflow_exports_old_quote_and_read_only_preview(clients,fmt):
 e,r,pid,_,quote=priced_project(clients);e,r,pid,s,b,eid=setup(clients,pid);before=get(e,pid);post(e,pid,b,'/components/preview');assert get(e,pid)==before;c=post(e,pid,b).json();approve(r,pid,'components',c)
 for t in post(e,pid,{},'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
 assert any(v['quantity']=='79.200' for v in post(e,pid,{},'/boq/generate').json());url=f'/api/v1/projects/{pid}/exports';res=e.get(url,params=dict(format=fmt,unpriced=True));assert res.status_code==200,res.text;titles={'立柱连续路径柱距','立柱路径相邻柱距','立柱路径核查依据'}
 if fmt=='json':assert titles<={t[0] for t in res.json()['tables']}
 elif fmt=='xlsx':assert titles<=set(load_workbook(io.BytesIO(res.content)).sheetnames)
 else:assert all(t in ''.join(p.extract_text() for p in PdfReader(io.BytesIO(res.content)).pages) for t in titles)
 old=e.get(url,params=dict(format='json',quote_id=quote['id'])).json();assert old['quote']==quote and not any(t[0] in titles for t in old['tables'])


@pytest.mark.parametrize('fault',['scope','proof','pair','exact','path','ast','book'])
def test_saved_path_input_and_outputs_rebuilt(clients,fault):
 e,r,pid,s,b,eid=setup(clients);c=post(e,pid,b).json()
 if fault=='book':approve(r,pid,'components',c);post(e,pid,{},'/takeoffs/calculate')
 def forge(p):
  c=p['components'][-1];g=c['preview']['material_items'][0]['post_spacing']
  if fault=='scope':c['material_items'][0]['spacing_check'].pop('scope')
  elif fault=='proof':c['material_items'][0]['spacing_check']['evidence']['measurement_basis']='forged'
  elif fault=='pair':g['paths'][0]['pairs'][-1]['distance']='0'
  elif fault=='exact':g['paths'][0]['pairs'][0]['exact_distance'][0]['coefficient']['numerator']='99'
  elif fault=='path':g['paths'][0]['closed']=True
  elif fault=='ast':g['paths'][0]['pairs'][0]['distance_formula_ast']={'constant':'2.5'}
  else:p['takeoffs'][-1]['material_items'][0]['post_spacing']['paths'][0]['pair_count']=99
 mutate(pid,forge);coll='takeoffs' if fault=='book' else 'components';approve(r,pid,coll,get(e,pid)[coll][-1],expected=404 if fault=='proof' else 409)


def test_local_to_path_same_mass_revision_history_and_stale_boq(clients):
 e,r,pid,s,b,eid=setup(clients);oldbody=copy.deepcopy(b);oldbody['material_items'][0]['spacing_check'].pop('scope');old=post(e,pid,oldbody).json();approve(r,pid,'components',old)
 for t in post(e,pid,{},'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
 post(e,pid,{},'/boq/generate');before=get(e,pid);old=before['components'][-1];b['object_hash']=store.digest(old);url=f'/api/v1/projects/{pid}/components/'+old['id'];res=e.put(url,json=b);assert res.status_code==200,res.text;p=get(e,pid);assert p['component_history'][-1]==old and p['boq']==before['boq'] and all(t['state']=='stale' for t in p['takeoffs']) and res.json()['preview']['quantity']=='79.200';assert e.put(url,json=b).status_code==409;approve(r,pid,'components',res.json())
 for t in post(e,pid,{},'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
 assert get(e,pid)['takeoffs'][-1]['material_items'][0]['post_spacing']['pair_count']==10;assert any(v['quantity']=='79.200' for v in post(e,pid,{},'/boq/generate').json())
 s['object_hash']=store.digest(s);s['edge_decisions'][0].update(mode='ranges',ranges=[interval(0,10,eid)]);res=e.put(f'/api/v1/projects/{pid}/components/'+s['id'],json=s);assert res.status_code==200;approve(r,pid,'components',res.json());assert get(e,pid)['components'][-1]['state']=='stale';post(e,pid,b,'/components/preview',409)


@pytest.mark.parametrize('gate',['preview','save','review','calculate','book_review'])
def test_path_original_bytes_gate(clients,gate):
 e,r,pid,s,b,eid=setup(clients)
 if gate in ('review','calculate','book_review'):c=post(e,pid,b).json()
 if gate in ('calculate','book_review'):approve(r,pid,'components',c)
 if gate=='book_review':t=post(e,pid,{},'/takeoffs/calculate').json()[-1]
 doc=get(e,pid)['documents'][-1];(store.ROOT/'originals'/doc['hash']).write_bytes(b'corrupt')
 if gate=='preview':post(e,pid,b,'/components/preview',409)
 elif gate=='save':post(e,pid,b,expected=409)
 elif gate=='review':approve(r,pid,'components',c,expected=409)
 elif gate=='calculate':post(e,pid,{},'/takeoffs/calculate',409)
 else:approve(r,pid,'takeoffs',t,expected=409)


def test_path_roles_tenant_and_foreign_evidence(clients):
 e,r,pid,s,b,eid=setup(clients);post(r,pid,b,expected=403);assert clients['other'].get(f'/api/v1/projects/{pid}').status_code==404
 other=create(e)['id'];upload(e,other,'其他工程');b['material_items'][0]['spacing_check']['evidence']['measurement_basis']=get(e,other)['evidence'][-1]['id'];post(e,pid,b,expected=404)


@pytest.mark.parametrize('direction',['from_start','from_end'])
def test_partial_intervals_both_fixed_anchors_and_added_positions(direction):
 s=line_source(12);s['edge_decisions'][0].update(mode='ranges',ranges=[interval(2,10)]);b=fixed_added(s);m=b['material_items'][0];m['post_ranges'][0]['direction']=direction;m['post_ranges'][0]['additional_posts'][0]['s']='2500';m['spacing_check']=check('2','m');p=result(s,b)['paths'][0];assert p['post_count']==6 and p['pair_count']==5 and p['max_distance']=='2' and p['unposted_start_distance']==p['unposted_end_distance']=='0'


def test_reordered_input_ranges_cannot_change_source_order():
 s=line_source(12);s['edge_decisions'][0].update(mode='ranges',ranges=[interval(4,12),interval(0,4)]);b=equal_body(s);b['material_items'][0]['spacing_check']=check('4','m');g=result(s,b);b['material_items'][0]['post_ranges'].reverse();assert result(s,b)==g


def test_open_empty_component_is_disclosed_without_joining_over_gap():
 s=line_source(12);s['edge_decisions'][0].update(mode='ranges',ranges=[interval(0,2),interval(3,4),interval(5,12)]);b=equal_body(s);m=b['material_items'][0]
 m['post_ranges'][1].update(spacing='10',spacing_unit='m',start_post=False,end_post=False);m['spacing_check']=check('2','m');g=result(s,b);assert len(g['paths'])==3 and g['unpaired_path_count']==1;empty=next(p for p in g['paths'] if p['fully_unposted']);assert empty['max_distance'] is None and empty['unposted_start_distance'] is None and not empty['pairs']


def test_maximum_1000_posts_and_5000_total_still_apply():
 s=line_source(1998);b=fixed_body(s);m=b['material_items'][0];m['spacing_check']=check('2','m');assert result(s,b)['pair_count']==999
 b['material_items']=[{**copy.deepcopy(m),'name':str(i)} for i in range(5)];assert sum(v['post_spacing']['pair_count'] for v in calculate(b,s)['material_items'])==4995
 b['material_items'].append({**copy.deepcopy(m),'name':'6'})
 with pytest.raises(DomainError,match='5000'):calculate(b,s)


def test_mixed_local_and_path_material_exports_and_long_names(clients):
 e,r,pid,s,b,eid=setup(clients);m=copy.deepcopy(b['material_items'][0]);m['spacing_check'].pop('scope');b['material_items'].append(m);b['name']='合成'+('长名称验收'*39)+'末';b['material_items'][0]['name']='合成'+('长材料名'*39)+'末';c=post(e,pid,b).json();approve(r,pid,'components',c)
 for t in post(e,pid,{},'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
 res=e.get(f'/api/v1/projects/{pid}/exports',params=dict(format='pdf',unpriced=True));assert res.status_code==200,res.text;txt=''.join(p.extract_text() for p in PdfReader(io.BytesIO(res.content)).pages)
 assert all(t in txt for t in ('立柱连续路径柱距','立柱路径相邻柱距','立柱路径核查依据','立柱区间实际柱距','立柱相邻柱距明细','立柱柱距核查依据'))
