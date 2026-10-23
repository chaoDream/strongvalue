import copy
import io
from decimal import Decimal
import pytest
from openpyxl import load_workbook
from pypdf import PdfReader
from backend import store
from tests.test_dxf_routes import setup,polyline,read,body,preview,adopt,mutate
from tests.test_dxf_assembly import segment,selections,resolve,assembled
from tests.test_workflow import get,approve,upload,priced_project

POINTS=[(0,0,0),(3000,4000,12000),(3000,4000,17000),(6000,8000,17000)]


def prepared(clients,pid=None,extra=None):
    c,r,pid,did,eid,url=setup(clients,polyline(POINTS)+(extra or []),pid=pid)
    upload(c,pid,'合成独立设计依据：原P1至P2三维距离13m，原P2至P3及原P3至P4各5m，容差0m；共同WCS、XYZ方向和设计高程已核对；各段中心线范围明确，不计管件；重复2组，总量修约3位。','fragment-basis.txt')
    eid=get(c,pid)['evidence'][-1]['id'];source=read(c,url)['candidates'][0]
    return c,r,pid,did,eid,url,source


def part(row,start=2,end=4,reverse=True):
    return {'candidate_id':row['id'],'object_hash':row['object_hash'],'from_point':start,'to_point':end,'reverse':reverse}


def select(c,url,selection,expected=200):
    result=c.post(url.replace('dxf-route-candidates','dxf-route-fragment'),json=selection)
    assert result.status_code==expected,result.text
    return result.json().get('candidate',result.json())


def request(row,eid,distance='5',**extra):
    return body(row,eid,fragment=row['fragment'],calibration={'from_point':1,'to_point':2,'distance':distance,'unit':'m','tolerance':'0'},**extra)


def test_fragment_read_and_preview_keep_original_vertex_numbers_and_are_readonly(clients):
    c,r,pid,did,eid,url,source=prepared(clients);before=get(c,pid);selection=part(source)
    row=select(c,url,selection);assert row==select(r,url,selection) and get(c,pid)==before
    assert row['source_object']=={k:v for k,v in source.items() if k!='object_hash'}
    assert row['points']==list(reversed(source['points'][1:4])) and row['source_point_indices']==[4,3,2]
    assert row['fragment']==selection and Decimal(row['single_length_m'])==10
    view=preview(c,url,row,request(row,eid));assert view['calibration']['passed'] and view['result']['quantity']=='20.000'
    assert view['result']['route_points'][0]['parameters']==source['points'][3] and get(c,pid)==before


@pytest.mark.parametrize('changes',[
    {'from_point':0},{'to_point':102},{'from_point':True},{'to_point':False},{'from_point':'2'},
    {'from_point':2.0},{'from_point':4,'to_point':2},{'from_point':2,'to_point':2},{'to_point':5},
    {'reverse':'true'},{'reverse':1},{'pixels':[1,2]},{'scale':'1'},{'object_hash':'bad'},{'candidate_id':'bad'}])
def test_invalid_ranges_and_display_parameters_are_atomic(clients,changes):
    c,r,pid,did,eid,url,source=prepared(clients);before=get(c,pid)
    select(c,url,{**part(source),**changes},422);assert get(c,pid)==before


@pytest.mark.parametrize('missing',['from_point','to_point','reverse'])
def test_range_endpoints_and_direction_must_be_explicit(clients,missing):
    c,r,pid,did,eid,url,source=prepared(clients);selection=part(source);selection.pop(missing)
    select(c,url,selection,422)


def test_adjacent_disjoint_fragments_share_vertex_but_not_segments_and_form_correct_boq(clients):
    c,r,pid,did,eid,url,source=prepared(clients)
    first=select(c,url,part(source,1,2,False));second=select(c,url,part(source,2,4,True));comps=[]
    for i,(row,distance) in enumerate([(first,'13'),(second,'5')]):
        req=request(row,eid,distance,request_id='synthetic-disjoint-fragment-000'+str(i))
        comp=adopt(c,url,row,req,preview(c,url,row,req))['component'];comps.append(comp)
        assert comp['state']=='draft'
    assert len(get(c,pid)['components'])==2 and not get(c,pid)['takeoffs']
    for comp in comps:approve(r,pid,'components',comp)
    result=c.post(f'/api/v1/projects/{pid}/takeoffs/calculate');assert result.status_code==200,result.text
    assert sorted(t['quantity'] for t in result.json())==['20.000','26.000']
    for t in result.json():approve(r,pid,'takeoffs',t)
    result=c.post(f'/api/v1/projects/{pid}/boq/generate');assert result.status_code==200,result.text
    assert sum(Decimal(b['quantity']) for b in result.json())==46


@pytest.mark.parametrize('case',['same','reverse','overlap','whole_then_fragment','fragment_then_whole','assembly_then_fragment','fragment_then_assembly'])
def test_overlapping_original_segment_claims_are_blocked_across_all_modes(clients,case):
    extra=segment('B',POINTS[-1],(6000,8000,18000));c,r,pid,did,eid,url,source=prepared(clients,extra=extra)
    row=select(c,url,part(source));req=request(row,eid)
    whole_req=body(source,eid,request_id='synthetic-whole-source-00001')
    parts=selections(read(c,url)['candidates']);parts[1]['reverse']=False
    combined=resolve(c,url,parts);combined_req=body(combined,eid,segments=parts,request_id='synthetic-assembly-source-001')
    if case=='whole_then_fragment':old,old_req=source,whole_req
    elif case=='assembly_then_fragment':old,old_req=combined,combined_req
    else:old,old_req=row,req
    adopt(c,url,old,old_req,preview(c,url,old,old_req))
    if case=='fragment_then_whole':new,new_req=source,whole_req
    elif case=='fragment_then_assembly':new,new_req=combined,combined_req
    else:
        selection=part(source,2,4,case!='reverse')
        if case=='overlap':selection=part(source,1,3,False)
        new=select(c,url,selection);new_req=request(new,eid,'13' if case=='overlap' else '5',request_id='synthetic-overlap-attempt-001')
    before=get(c,pid);adopt(c,url,new,new_req,preview(c,url,new,new_req),409);assert get(c,pid)==before


def test_whole_source_can_be_revised_to_fragment_then_remaining_range_adopted(clients):
    c,r,pid,did,eid,url,source=prepared(clients);req=body(source,eid)
    old=adopt(c,url,source,req,preview(c,url,source,req))['component'];approve(r,pid,'components',old)
    t=c.post(f'/api/v1/projects/{pid}/takeoffs/calculate').json()[0];approve(r,pid,'takeoffs',t)
    first=select(c,url,part(source,1,2,False));req=request(first,eid,'13',request_id='synthetic-whole-to-fragment',component_id=old['id'],component_hash=store.digest(get(c,pid)['components'][0]))
    new=adopt(c,url,first,req,preview(c,url,first,req))['component'];assert new['revision']==2 and new['id']==old['id']
    p=get(c,pid);assert p['takeoffs'][0]['state']=='stale' and p['component_history'][0]['dxf_trace']['candidate']['kind']=='POLYLINE'
    rest=select(c,url,part(source));req=request(rest,eid,request_id='synthetic-remaining-fragment')
    adopt(c,url,rest,req,preview(c,url,rest,req));assert len(get(c,pid)['components'])==2


def test_revision_cannot_expand_into_another_fragments_claim(clients):
    c,r,pid,did,eid,url,source=prepared(clients);saved=[]
    for i,(selection,distance) in enumerate([(part(source,1,2,False),'13'),(part(source),'5')]):
        row=select(c,url,selection);req=request(row,eid,distance,request_id='synthetic-initial-split-000'+str(i))
        saved.append(adopt(c,url,row,req,preview(c,url,row,req))['component'])
    expanded=select(c,url,part(source,1,3,False));req=request(expanded,eid,'13',request_id='synthetic-expanded-split-001',component_id=saved[0]['id'],component_hash=store.digest(saved[0]))
    before=get(c,pid);adopt(c,url,expanded,req,preview(c,url,expanded,req),409);assert get(c,pid)==before


def test_stale_selection_preview_and_mixed_modes_are_rejected(clients):
    c,r,pid,did,eid,url,source=prepared(clients);row=select(c,url,part(source));req=request(row,eid);view=preview(c,url,row,req)
    bad=copy.deepcopy(req);bad['fragment']['reverse']=False;adopt(c,url,row,bad,view,409)
    bad=copy.deepcopy(req);bad['fragment']['object_hash']='0'*64;adopt(c,url,row,bad,view,409)
    bad=copy.deepcopy(req);bad['fragment']=None;adopt(c,url,row,bad,view,409)
    bad=copy.deepcopy(req);bad['segments']=[{'candidate_id':source['id'],'object_hash':source['object_hash'],'reverse':False}]*2
    preview(c,url,row,bad,422)
    select(c,url,{**part(source),'candidate_id':'0'*32},409)


@pytest.mark.parametrize('field',['fragment','source_object','source_point_indices','points'])
def test_saved_fragment_provenance_is_rebuilt_before_review(clients,field):
    c,r,pid,did,eid,url,source=prepared(clients);row=select(c,url,part(source));req=request(row,eid)
    comp=adopt(c,url,row,req,preview(c,url,row,req))['component']
    def corrupt(p):
        for trace in [p['components'][0]['dxf_trace'],p['dxf_route_adoptions'][0]['trace']]:trace['candidate'][field]=None
    mutate(pid,corrupt);approve(r,pid,'components',get(c,pid)['components'][0],expected=409)


@pytest.mark.parametrize('change',['bytes','job','evidence'])
def test_source_changes_invalidate_saved_fragment_and_prevent_new_adoption(clients,change):
    c,r,pid,did,eid,url,source=prepared(clients);row=select(c,url,part(source));req=request(row,eid);view=preview(c,url,row,req)
    comp=adopt(c,url,row,req,view)['component'];approve(r,pid,'components',comp)
    if change=='bytes':(store.ROOT/'originals'/source['document_hash']).write_bytes(b'changed')
    else:
        def action(p):
            if change=='job':next(j for j in p['jobs'] if j['document_id']==did)['state']='failed'
            else:next(e for e in p['evidence'] if e['id']==eid)['text']+='changed'
        mutate(pid,action)
    approve(r,pid,'components',get(c,pid)['components'][0],expected=409)
    assert c.post(f'/api/v1/projects/{pid}/takeoffs/calculate').status_code==409
    req['request_id']='synthetic-new-stale-fragment';adopt(c,url,row,req,view,409)


def test_roles_tenants_idempotency_and_old_assembly_request_hash(clients):
    c,r,pid,did,eid,url,source=prepared(clients);selection=part(source);row=select(c,url,selection);req=request(row,eid);view=preview(c,url,row,req)
    select(clients['other'],url,selection,404);adopt(r,url,row,req,view,403)
    result=adopt(c,url,row,req,view);assert adopt(c,url,row,req,view)['duplicate'] and len(get(c,pid)['components'])==1
    assert result['component']['dxf_trace']['candidate']['source_point_indices']==[4,3,2]
    c,r,pid,did,eid,url,objects,row,req=assembled(clients)
    view=preview(c,url,row,req);result=adopt(c,url,row,req,view)
    assert result['adoption']['request_hash']==store.digest({'document_id':did,'candidate_id':row['id'],**req,'component_id':None,'component_hash':None,'preview_hash':view['preview_hash']})


def test_fragment_ledger_exports_and_old_quote_isolation(clients):
    c,r,pid,_,old=priced_project(clients);c,r,pid,did,eid,url,source=prepared(clients,pid=pid)
    row=select(c,url,part(source));req=request(row,eid);adopt(c,url,row,req,preview(c,url,row,req))
    base=f'/api/v1/projects/{pid}/exports'
    for fmt in ['json','xlsx','pdf']:
        res=c.get(base,params={'format':fmt,'unpriced':True});assert res.status_code==200,res.text
        if fmt=='json':assert 'DXF原对象取段范围' in str(res.json()['tables']) and 'P1→原P4' in str(res.json()['tables'])
        elif fmt=='xlsx':assert 'DXF原对象取段范围' in load_workbook(io.BytesIO(res.content)).sheetnames
        else:assert 'DXF原对象取段范围' in ''.join(p.extract_text() for p in PdfReader(io.BytesIO(res.content)).pages)
    historic=c.get(base,params={'format':'json','quote_id':old['id']}).json()
    assert not any(t[0]=='DXF原对象取段范围' for t in historic['tables']) and historic['quote']==old
