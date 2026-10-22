"""Synthetic station math, workflow and source dependency checks; no professional claim."""
import copy
import io
import json
import random
from decimal import Decimal, ROUND_HALF_UP, localcontext
import pytest
from openpyxl import load_workbook
from pypdf import PdfReader
from backend import store
from backend.engine import DomainError
from backend.trench_takeoff import calculate, snapshot
from tests.test_route_takeoff import payload as route_payload, point
from tests.test_workflow import create, get, upload, approve, priced_project


def evaluate(node, dimensions):
    if 'dimension' in node:return Decimal(dimensions[node['dimension']])
    if 'constant' in node:return Decimal(node['constant'])
    values=[evaluate(a,dimensions) for a in node['args']]
    if node['op']=='sqrt':return values[0].sqrt()
    if node['op']=='round':return values[0].quantize(Decimal(1).scaleb(-node['digits']),rounding=ROUND_HALF_UP)
    if node['op']=='add':return sum(values,Decimal(0))
    if node['op']=='subtract':return values[0]-values[1]
    if node['op']=='divide':return values[0]/values[1]
    assert node['op']=='multiply'
    result=Decimal(1)
    for v in values:result*=v
    return result


def geometry(eid='test'):
    s=route_payload(eid);s.update(id='synthetic-route',revision=1)
    s['route_points']=s['route_points'][:2]
    return s


def section(i,b,t,h,eid='test',unit='m'):
    return {'point_index':i,'parameters':{'bottom_width':str(b),'top_width':str(t),'depth':str(h)},'unit':unit,
            'evidence':{k:eid for k in ['bottom_width','top_width','depth','measurement_basis']},'measurement_basis':'合成已核定梯形断面，非实际施工方案'}


def payload(source,eid='test'):
    return {'name':'合成变断面沟槽（非真实工程）','kind':'station_trench','parameters':{},'quantity_type':'volume',
            'source_component_id':source['id'],'source_component_hash':store.digest(snapshot(source)),
            'method':'linear_prismoidal','stations':[section(1,2,4,1,eid),section(2,4000,8000,3000,eid,'mm')],
            'measurement_basis':'仅合成：段内底宽、顶宽与挖深分别线性变化，水平长度5m，按棱台法，不含折点重叠；非专业签审','basis_evidence_id':eid,
            'geology_basis':'仅合成地质及场地假设，非真实地勘结论','geology_evidence_id':eid,
            'construction_basis':'仅合成已核定上下宽，无自动放坡或工作面系数','construction_evidence_id':eid,
            'rounding_digits':3,'rounding_evidence_id':eid}


def setup(clients):
    e,r=clients['estimator'],clients['reviewer'];pid=create(e)['id'];base='/api/v1/projects/'+pid
    upload(e,pid,'合成设计管线XYZ(0,0,0)m至(3000,4000,12000)mm，空间长13m，水平5m，共同XYZ基准，重复1，长度总量修约3位。')
    s=geometry(get(e,pid)['evidence'][0]['id']);s.pop('id');s.pop('revision')
    response=e.post(base+'/components',json=s);assert response.status_code==200,response.text
    approve(r,pid,'components',response.json());source=get(e,pid)['components'][0]
    upload(e,pid,'仅合成断面：起点底宽2顶宽4挖深1m，终点底宽4000顶宽8000挖深3000mm。上下宽与挖深分别线性变化；无转角交接重叠。地质场地/施工工作面均仅为假设，不自动取系数。体积总量修约3位。',name='sections.txt')
    # New source documents invalidate approvals; re-review the unchanged route after upload.
    approve(r,pid,'components',get(e,pid)['components'][0])
    source=get(e,pid)['components'][0];assert source['state']=='approved'
    eid=get(e,pid)['evidence'][-1]['id']
    return e,r,pid,source,payload(source,eid),eid


def test_sloping_route_uses_horizontal_length_and_exact_linear_integration():
    source=geometry();b=payload(source);before=copy.deepcopy(b);r=calculate(b,source)
    assert b==before and r['quantity']=='47.500' and r['unit']=='m³'
    s=r['trench_segments'][0]
    assert [Decimal(s[k]) for k in ['horizontal_length','spatial_length','start_area','end_area','mid_area','mean_area','volume']]==[5,13,3,18,9,Decimal('9.5'),Decimal('47.5')]
    assert r['trench_stations'][1]['normalized_parameters']=={'bottom_width':'4','top_width':'8','depth':'3'}
    assert len(r['unit_conversions'])==12
    with localcontext() as ctx:
        ctx.prec=80
        assert evaluate(r['formula_ast'],r['dimensions'])==Decimal(r['quantity'])
        assert evaluate(s['formula_ast'],r['dimensions'])==Decimal('47.5')


def test_average_end_area_is_distinct_explicit_method():
    source=geometry();b=payload(source);b['method']='average_end_area';r=calculate(b,source)
    assert r['quantity']=='52.500' and r['trench_summary']['method_label']=='平均端面积法'
    assert '近似' in r['notice']
    with localcontext() as ctx:
        ctx.prec=80
        assert evaluate(r['formula_ast'],r['dimensions'])==Decimal('52.5')


def test_same_depth_constant_sections_and_zero_depth_are_supported():
    s=geometry();b=payload(s);b['stations']=[section(1,2,4,1),section(2,2,4,1)]
    assert calculate(b,s)['quantity']=='15.000'
    b['method']='average_end_area';assert calculate(b,s)['quantity']=='15.000'
    for v in b['stations']:v['parameters']['depth']='0'
    assert calculate(b,s)['quantity']=='0.000'


def test_source_repetition_round_only_total_and_station_order_is_irrelevant():
    s=geometry();s['multiplicity']='2';b=payload(s);b['rounding_digits']=0
    assert calculate(b,s)['quantity']=='95'
    b['stations'].reverse();assert calculate(b,s)['quantity']=='95'


def test_seeded_linear_sections_match_independent_polynomial_integral():
    rng=random.Random(3001)
    for _ in range(30):
        s=geometry();s['route_points']=[point('a',-3,-4,-12),point('b',0,0,0)]
        b=payload(s);b0,b1,h0,h1=[Decimal(rng.randint(1,9)) for _ in range(4)];t0=b0+2;t1=b1+3
        b['stations']=[section(1,b0,t0,h0),section(2,b1,t1,h1)];b['rounding_digits']=12
        # Integrate ((b+t)/2) * h for independently linear width/depth over t in [0,1].
        with localcontext() as ctx:
            ctx.prec=80
            w0=(b0+t0)/2;dw=(b1+t1)/2-w0;dh=h1-h0
            expected=(w0*h0+(w0*dh+dw*h0)/2+dw*dh/3)*5
            r=calculate(b,s)
            assert Decimal(r['quantity'])==expected.quantize(Decimal('.000000000001'))
            assert evaluate(r['formula_ast'],r['dimensions'])==Decimal(r['quantity'])


@pytest.mark.parametrize('case',['missing','duplicate','unknown','non_int_index','unknown_field','missing_width','missing_proof','negative_depth','inverted','zero_width','nan','unit','precision','mode','parameters','method','method_type','rounding','missing_geology','missing_construction','blank_geology_proof','overflow','vertical'])
def test_invalid_sections_and_missing_basis_fail(case):
    s=geometry();b=payload(s);st=b['stations'][0]
    if case=='missing':b['stations'].pop()
    elif case=='duplicate':b['stations'][1]['point_index']=1
    elif case=='unknown':st['point_index']=99
    elif case=='non_int_index':st['point_index']=True
    elif case=='unknown_field':st['slope']='.5'
    elif case=='missing_width':del st['parameters']['top_width']
    elif case=='missing_proof':del st['evidence']['depth']
    elif case=='negative_depth':st['parameters']['depth']='-1'
    elif case=='inverted':st['parameters']['top_width']='1'
    elif case=='zero_width':st['parameters']['bottom_width']='0'
    elif case=='nan':st['parameters']['depth']='NaN'
    elif case=='unit':st['unit']='m³'
    elif case=='precision':st['unit']='mm';st['parameters']['depth']='.000000000000000001'
    elif case=='mode':b['quantity_type']='length'
    elif case=='parameters':b['parameters']={'volume':'1'}
    elif case=='method':b['method']='automatic_slope'
    elif case=='method_type':b['method']={}
    elif case=='rounding':b['rounding_digits']=True
    elif case=='missing_geology':del b['geology_basis']
    elif case=='missing_construction':b['construction_basis']=''
    elif case=='blank_geology_proof':b['geology_evidence_id']=''
    elif case=='overflow':
        for st in b['stations']:st['unit']='m';st['parameters']={k:'1e12' for k in ['bottom_width','top_width','depth']}
    elif case=='vertical':
        s['route_points']=[point('a',0,0,0),point('b',0,0,5)];b['source_component_hash']=store.digest(snapshot(s))
    with pytest.raises(DomainError):calculate(b,s)


def test_multi_segment_accumulation_and_ast():
    s=geometry();s['route_points'] += [point('c',9,12,13)];b=payload(s);b['stations'].append(section(3,4,8,3))
    r=calculate(b,s);assert r['quantity']=='227.500' and len(r['trench_segments'])==2
    with localcontext() as ctx:
        ctx.prec=80
        assert evaluate(r['formula_ast'],r['dimensions'])==Decimal('227.5')


def test_candidates_preview_save_review_boq_and_exports(clients):
    e,r,pid,source,b,eid=setup(clients);base='/api/v1/projects/'+pid;before=get(e,pid)
    c=e.get(base+'/trench-station-candidates',params={'component_id':source['id']});assert c.status_code==200,c.text
    assert c.json()['route_segments'][0]['horizontal_length']=='5' and get(e,pid)==before
    response=e.post(base+'/components/preview',json=b);assert response.status_code==200,response.text
    assert get(e,pid)==before
    response=e.post(base+'/components',json={**b,'source_geometry':{'forged':True}});assert response.status_code==200,response.text
    comp=response.json();assert comp['source_geometry']==snapshot(source)
    assert get(e,pid)['components'][0]['state']=='approved'
    assert e.post(base+'/takeoffs/calculate').status_code==409
    approve(r,pid,'components',comp);books=e.post(base+'/takeoffs/calculate').json()
    for book in books:approve(r,pid,'takeoffs',book)
    items=e.post(base+'/boq/generate').json();assert next(x for x in items if x['unit']=='m³')['quantity']=='47.500'
    j=e.get(base+'/exports?unpriced=true&format=json').json()['unpriced']['takeoffs'][1]
    assert j['trench_segments']==books[1]['trench_segments'] and j['source_geometry']==snapshot(source)
    response=e.get(base+'/exports?unpriced=true&format=xlsx');assert response.status_code==200,response.text
    wb=load_workbook(io.BytesIO(response.content));assert {'变断面沟槽汇总','沟槽站点断面','沟槽分段体积','沟槽来源坐标','算量原始单位换算'}<=set(wb.sheetnames)
    assert all(c.data_type!='f' for sheet in wb for row in sheet for c in row)
    response=e.get(base+'/exports?unpriced=true&format=pdf');assert response.status_code==200,response.text
    text='\n'.join(p.extract_text() for p in PdfReader(io.BytesIO(response.content)).pages)
    assert '变断面沟槽汇总' in text and '沟槽分段体积' in text and '47.500' in text


def test_revision_atomicity_source_approval_preserved_and_old_books_stale(clients):
    e,r,pid,source,b,eid=setup(clients);base='/api/v1/projects/'+pid
    comp=e.post(base+'/components',json=b).json();approve(r,pid,'components',comp)
    books=e.post(base+'/takeoffs/calculate').json()
    for book in books:approve(r,pid,'takeoffs',book)
    before=get(e,pid);assert e.put(base+'/components/'+comp['id'],json=b).status_code==409
    b['object_hash']=e.get(base+'/hashes').json()['components'][comp['id']]
    bad=copy.deepcopy(b);bad['stations'][0]['parameters']['top_width']='1'
    assert e.put(base+'/components/'+comp['id'],json=bad).status_code==422 and get(e,pid)==before
    b['method']='average_end_area';response=e.put(base+'/components/'+comp['id'],json=b)
    assert response.status_code==200,response.text
    assert response.json()['revision']==2 and response.json()['preview']['quantity']=='52.500'
    p=get(e,pid);assert p['components'][0]['state']=='approved' and all(t['state']=='stale' for t in p['takeoffs'])
    assert e.put(base+'/components/'+comp['id'],json=b).status_code==409


def test_source_revision_and_unapproval_block_old_station_hash(clients):
    e,r,pid,source,b,eid=setup(clients);base='/api/v1/projects/'+pid
    comp=e.post(base+'/components',json=b).json();approve(r,pid,'components',comp)
    route=copy.deepcopy(source);route['object_hash']=e.get(base+'/hashes').json()['components'][source['id']];route['route_points'][1]['parameters']['x']='6000'
    response=e.put(base+'/components/'+source['id'],json=route);assert response.status_code==200,response.text
    assert e.get(base+'/trench-station-candidates',params={'component_id':source['id']}).status_code==409
    approve(r,pid,'components',response.json())
    assert e.post(base+'/components/preview',json=b).status_code==409
    approve(r,pid,'components',get(e,pid)['components'][1],expected=409)
    fresh=e.get(base+'/trench-station-candidates',params={'component_id':source['id']});assert fresh.status_code==200
    assert fresh.json()['source_component_hash']!=b['source_component_hash']


@pytest.mark.parametrize('gate',['preview','save','component_review','calculate','book_review'])
def test_independent_section_original_corrupt_blocks_all_gates(clients,gate):
    e,r,pid,source,b,eid=setup(clients);base='/api/v1/projects/'+pid
    if gate not in ('preview','save'):comp=e.post(base+'/components',json=b).json()
    if gate in ('calculate','book_review'):approve(r,pid,'components',comp)
    if gate=='book_review':book=e.post(base+'/takeoffs/calculate').json()[1]
    doc=get(e,pid)['documents'][-1];(store.ROOT/'originals'/doc['hash']).write_bytes(b'corrupt')
    # A section-only corruption does not claim the independent route original is corrupt.
    assert e.get(base+'/trench-station-candidates',params={'component_id':source['id']}).status_code==200
    if gate=='preview':assert e.post(base+'/components/preview',json=b).status_code==409
    elif gate=='save':assert e.post(base+'/components',json=b).status_code==409
    elif gate=='component_review':approve(r,pid,'components',comp,expected=409)
    elif gate=='calculate':assert e.post(base+'/takeoffs/calculate').status_code==409
    else:approve(r,pid,'takeoffs',book,expected=409)


def test_permissions_foreign_proofs_wrong_source_and_type_guard(clients):
    e,r,pid,source,b,eid=setup(clients);base='/api/v1/projects/'+pid
    assert r.get(base+'/trench-station-candidates',params={'component_id':source['id']}).status_code==200
    assert r.post(base+'/components',json=b).status_code==403
    assert clients['other'].get(base+'/trench-station-candidates',params={'component_id':source['id']}).status_code==404
    pid2=create(e)['id'];upload(e,pid2,'其他项目')
    bad=copy.deepcopy(b);bad['geology_evidence_id']=get(e,pid2)['evidence'][0]['id']
    assert e.post(base+'/components',json=bad).status_code==404
    comp=e.post(base+'/components',json=b).json()
    assert e.get(base+'/trench-station-candidates',params={'component_id':comp['id']}).status_code==409
    changed={'kind':'trench','parameters':{},'object_hash':e.get(base+'/hashes').json()['components'][comp['id']]}
    assert e.put(base+'/components/'+comp['id'],json=changed).status_code==409


@pytest.mark.parametrize('target',['component','book'])
def test_tampered_calculation_is_rejected(clients,target):
    e,r,pid,source,b,eid=setup(clients);base='/api/v1/projects/'+pid
    comp=e.post(base+'/components',json=b).json()
    if target=='book':approve(r,pid,'components',comp);e.post(base+'/takeoffs/calculate')
    with store.transaction() as db:
        p=json.loads(db.execute('SELECT data FROM projects WHERE id=?',(pid,)).fetchone()['data'])
        if target=='component':p['components'][1]['preview']['trench_segments'][0]['volume']='999'
        else:p['takeoffs'][1]['trench_segments'][0]['volume']='999'
        store.save(db,p)
    col='components' if target=='component' else 'takeoffs'
    approve(r,pid,col,get(e,pid)[col][1],expected=409)


def test_adding_route_and_trench_preserves_quote_snapshot(clients):
    e,r,pid,_,q=priced_project(clients);base='/api/v1/projects/'+pid;old=copy.deepcopy(q);eid=get(e,pid)['evidence'][0]['id']
    s=geometry(eid);s.pop('id');s.pop('revision');source=e.post(base+'/components',json=s).json();approve(r,pid,'components',source)
    source=get(e,pid)['components'][-1];response=e.post(base+'/components',json=payload(source,eid));assert response.status_code==200,response.text
    p=get(e,pid);assert p['quotes'][0]==old and q['id'] in p['stale_quotes']
