"""Synthetic combined layout identities, exact controls and workflow regression."""
import copy,io,random
from decimal import Decimal,localcontext
from fractions import Fraction
import pytest
from openpyxl import load_workbook
from pypdf import PdfReader
from backend import store
from backend.engine import DomainError
from backend.railing_material import calculate
from backend.railing_takeoff import snapshot
from tests.test_additional_posts import body as fixed_body,addition,line_source
from tests.test_equal_additional_posts import body as equal_body
from tests.test_equal_post_adjustments import avoid,move as regular_move
from tests.test_railing_posts import geometry,setup as source_setup
from tests.test_fixed_posts import stations
from tests.test_route_takeoff import evaluate,point
from tests.test_route_railing import payload as railing_payload
from tests.test_railing_takeoff import interval
from tests.test_polygon_takeoff import post,mutate
from tests.test_workflow import get,approve,upload,priced_project,create


def move(kind='regular',node=1,s='2500',unit='mm',eid='test'):
 m=regular_move(node,s,unit,eid);m['node_type']=kind;m['evidence']['node_type']=eid;return m


def body(s,eid='test',equal=False):
 b=(equal_body if equal else fixed_body)(s,eid);b['name']='合成含加柱禁布移柱验收（非真实工程）';m=b['material_items'][0];m.update(mode='spaced_added_adjusted_posts' if equal else 'fixed_added_adjusted_posts',name='合成同规格含加柱移柱')
 for i,r in enumerate(m['post_ranges']):
  r.update(avoidance_ranges=[avoid('1400','1600',eid=eid,label='原加柱禁布'),avoid('1800','2200',eid=eid,label='原规则禁布')] if i==0 else [avoid('2400','2550',eid=eid,label='BC加柱禁布')],post_moves=[move(eid=eid),move('additional',1,'1700',eid=eid)] if i==0 else [move('additional',1,'2600',eid=eid),move('regular',2,'3500',eid=eid)])
 b['measurement_basis']='合成原规则11根＋原构造加柱2根，按类型和原编号移4根，根数规格保持13根，总93.6kg；仅验收软件';return b


def setup(clients,pid=None,equal=False):
 e,r,pid,s,_,_=source_setup(clients,pid)
 upload(e,pid,'合成含加柱移柱验收，非真实工程。AB三维13m，BC5m，共享B只归AB。定距AB起端2m、BC终端2m，或最大2m均分，原规则11根；原构造加柱AB S1500mm及BC S2500mm，各加柱输入序号1。禁布闭区间AB[1400,1600]、[1800,2200]mm及BC[2400,2550]mm。AB原规则j1移2500mm，原加柱1移1700mm；BC原规则j2移3500mm、原加柱1移2600mm。类型及编号分别有据，不移动端柱、不换序、不重合，原布置必须合法。每根1200mm、3000g/m，来源重复2，13根仍93.6kg；修订AB加柱目标1750mm后仍93.6kg。来源AB缩至S0到10m需重新核定。','combined-post-adjustments-basis.txt')
 for c in get(e,pid)['components']:approve(r,pid,'components',c)
 s=get(e,pid)['components'][-1];eid=get(e,pid)['evidence'][-1]['id'];return e,r,pid,s,body(s,eid,equal),eid


@pytest.mark.parametrize('equal',[False,True])
def test_complete_original_identity_count_mass_and_ast(equal):
 s=geometry();b=body(s,equal=equal);before=copy.deepcopy((b,s));v=calculate(b,s);assert before==(b,s)
 g=v['material_items'][0]['post_layout'];base=calculate((equal_body if equal else fixed_body)(s),s)['material_items'][0]['post_layout']
 assert v['quantity']=='93.600' and v['rule_version']=='railing-material-mass-1.9'
 assert (g['base_post_count'],g['additional_post_count'],g['post_count'],g['moved_post_count'],g['avoidance_range_count'])==(11,2,13,4,3)
 assert g['rule_version']==('railing-equal-added-adjusted-posts-1.0' if equal else 'railing-fixed-added-adjusted-posts-1.0')
 assert [p['original_s'] for p in g['post_stations']]==[p['s'] for p in base['post_stations']]
 assert [p['index'] for p in g['post_stations']]==list(range(1,14))
 moved=[p for p in g['post_stations'] if p['moved']];assert len(moved)==4
 assert [(p['original_node_type'],p['original_node_index'],p['s']) for p in moved]==[('additional',1,'1.7'),('regular',1,'2.5'),('additional',1,'2.6'),('regular',2,'3.5')]
 with localcontext() as c:
  c.prec=80;assert evaluate(v['formula_ast'],v['dimensions'])==Decimal('93.6')
  for p in g['post_stations']:
   assert evaluate(p['station_formula_ast'],v['dimensions'])==Decimal(p['s'])
   assert evaluate(p['original_station_formula_ast'],v['dimensions'])==Decimal(p['original_s'])
  for r in g['post_ranges']:assert evaluate(r['formula_ast'],v['dimensions'])==r['post_count'];assert all(a['final_post_count']==0 for a in r['avoidance_ranges'])
 assert len([k for k in g['source_evidence'] if k.endswith('.node_type')])==4


@pytest.mark.parametrize('equal',[False,True])
@pytest.mark.parametrize('fault',['missing_type','bad_type','list_type','missing_type_proof','blank_type_proof','extra_move','bool_node','string_node','missing_add','end_node','same','repeat','closed_boundary','cross_regular','cross_added','collision','unmoved_added','invalid_original','no_additions','no_controls','wrong_schema','missing_list','bad_count','extra_direction'])
def test_strict_originals_targets_and_combined_controls(equal,fault):
 s=geometry();b=body(s,equal=equal);item=b['material_items'][0];r=item['post_ranges'][0];m=r['post_moves'][1]
 if fault=='missing_type':m.pop('node_type')
 elif fault=='bad_type':m['node_type']='start'
 elif fault=='list_type':m['node_type']=[]
 elif fault=='missing_type_proof':m['evidence'].pop('node_type')
 elif fault=='blank_type_proof':m['evidence']['node_type']=' '
 elif fault=='extra_move':m['count']=1
 elif fault=='bool_node':m['node_index']=True
 elif fault=='string_node':m['node_index']='1'
 elif fault=='missing_add':m['node_index']=2
 elif fault=='end_node':r['post_moves'][0]['node_index']=7
 elif fault=='same':m['s']='1500'
 elif fault=='repeat':r['post_moves'].append(copy.deepcopy(m))
 elif fault=='closed_boundary':m['s']='1600'
 elif fault=='cross_regular':m['s']='2600'
 elif fault=='cross_added':r['post_moves'][0]['s']='1650'
 elif fault=='collision':m['s']='2500'
 elif fault=='unmoved_added':r['post_moves'].pop(1)
 elif fault=='invalid_original':r['additional_posts'][0]['s']='2000' if not equal else '13000'
 elif fault=='no_additions':
  for p in item['post_ranges']:p['additional_posts']=[]
 elif fault=='no_controls':
  for p in item['post_ranges']:p.update(post_moves=[],avoidance_ranges=[])
 elif fault=='wrong_schema':item['mode']='spaced_adjusted_posts' if equal else 'fixed_adjusted_posts'
 elif fault=='missing_list':r.pop('additional_posts')
 elif fault=='bad_count':item['quantity']='13'
 elif fault=='extra_direction':m['direction']='from_end'
 with pytest.raises(DomainError):calculate(b,s)


@pytest.mark.parametrize('equal',[False,True])
@pytest.mark.parametrize('basis',['spatial','horizontal'])
@pytest.mark.parametrize('unit,factor',[('m',Decimal(1)),('mm',Decimal(1000)),('cm',Decimal(100)),('dm',Decimal(10)),('km',Decimal('.001'))])
def test_units_basis_and_shared_endpoint(equal,basis,unit,factor):
 s=geometry(basis=basis);b=body(s,equal=equal)
 for r in b['material_items'][0]['post_ranges']:
  r.update(additional_posts=[addition(str(Decimal('1.5')*factor),unit)],avoidance_ranges=[avoid(str(Decimal('1.4')*factor),str(Decimal('1.6')*factor),unit)],post_moves=[move('additional',1,str(Decimal('1.62')*factor),unit)])
 assert calculate(b,s)['quantity']==('93.600' if basis=='spatial' else '64.800')


@pytest.mark.parametrize('equal',[False,True])
def test_unsorted_addition_input_identity_is_not_sorted_index(equal):
 s=line_source(8);b=body(s,equal=equal);r=b['material_items'][0]['post_ranges'][0]
 r.update(additional_posts=[addition('5000',label='输入第一'),addition('1000',label='输入第二')],avoidance_ranges=[avoid('900','1100')],post_moves=[move('additional',2,'1200')])
 g=calculate(b,s)['material_items'][0]['post_layout'];p=next(p for p in g['post_stations'] if p['moved'])
 assert p['additional_index']==2 and p['additional_label']=='输入第二' and p['original_s']=='1' and p['s']=='1.2'


@pytest.mark.parametrize('equal',[False,True])
def test_partial_and_multiple_moves_may_pass_original_locations_but_not_final_order(equal):
 s=line_source(12);s['edge_decisions'][0].update(mode='ranges',ranges=[interval(2,10)]);b=body(s,equal=equal);r=b['material_items'][0]['post_ranges'][0]
 r.update(additional_posts=[addition('3000')],avoidance_ranges=[],post_moves=[move('additional',1,'4500'),move('regular',1,'5000')])
 assert stations(calculate(b,s)['material_items'][0]['post_layout'],'S1 区间1')==[2,Decimal('4.5'),5,6,8,10]
 r['post_moves'][0]['s']='1500'
 with pytest.raises(DomainError):calculate(b,s)


@pytest.mark.parametrize('equal',[False,True])
def test_seeded_independent_rational_combined_oracle(equal):
 rng=random.Random(7601)
 for _ in range(30):
  L=rng.randint(9,35);s=line_source(L);b=body(s,equal=equal);r=b['material_items'][0]['post_ranges'][0]
  # Place and move an original added post in the first gap without changing its order.
  n=(L+1)//2;first=Fraction(L,n) if equal else Fraction(2);original=Fraction(1,2);target=Fraction(rng.randint(7,12),10)
  assert target<first
  r.update(additional_posts=[addition(str(float(original)),unit='m')],avoidance_ranges=[avoid('.4','.6','m')],post_moves=[move('additional',1,str(float(target)),'m')])
  g=calculate(b,s)['material_items'][0]['post_layout'];expected=[Fraction(0),target]+[Fraction(L*k,n) if equal else Fraction(2*k) for k in range(1,n)]+[Fraction(L)]
  with localcontext() as c:
   c.prec=80
   for p,e in zip(g['post_stations'],expected):assert abs(Decimal(p['s'])-Decimal(e.numerator)/Decimal(e.denominator))<Decimal('1e-75')


@pytest.mark.parametrize('equal',[False,True])
@pytest.mark.parametrize('fmt',['json','xlsx','pdf'])
def test_workflow_review_exports_and_old_quote(clients,equal,fmt):
 e,r,pid,_,quote=priced_project(clients);e,r,pid,s,b,eid=setup(clients,pid,equal);before=get(e,pid);post(e,pid,b,'/components/preview');assert get(e,pid)==before;c=post(e,pid,b).json();approve(r,pid,'components',c)
 for t in post(e,pid,{},'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
 assert any(t['quantity']=='93.600' for t in post(e,pid,{},'/boq/generate').json())
 url=f'/api/v1/projects/{pid}/exports';res=e.get(url,params=dict(format=fmt,unpriced=True));assert res.status_code==200,res.text
 titles={'构造加柱数量组成','构造加柱端距及依据','立柱禁布范围核查','立柱移前移后端距','含加柱移柱规则与范围'}
 if fmt=='json':assert titles<={v[0] for v in res.json()['tables']}
 elif fmt=='xlsx':assert titles<=set(load_workbook(io.BytesIO(res.content)).sheetnames)
 else:
  txt=''.join(p.extract_text() for p in PdfReader(io.BytesIO(res.content)).pages);assert all(t in txt for t in titles);assert '原加柱 #1' in txt and '原规则 j=1' in txt
 old=e.get(url,params=dict(format='json',quote_id=quote['id'])).json();assert old['quote']==quote and not any(t[0] in titles for t in old['tables'])


@pytest.mark.parametrize('fault',['type','node','original','position','count','addition','type_proof','book'])
def test_saved_inputs_and_record_tampering_blocked(clients,fault):
 e,r,pid,s,b,eid=setup(clients);c=post(e,pid,b).json()
 if fault=='book':approve(r,pid,'components',c);post(e,pid,{},'/takeoffs/calculate')
 def forge(p):
  c=p['components'][-1];g=c['preview']['material_items'][0]['post_layout'];m=c['material_items'][0]['post_ranges'][0]['post_moves'][1]
  if fault=='type':m['node_type']='regular'
  elif fault=='node':m['node_index']=2
  elif fault=='original':g['post_stations'][1]['original_node_type']='regular'
  elif fault=='position':g['post_stations'][1]['s']='1.75'
  elif fault=='count':g['additional_post_count']=9
  elif fault=='addition':c['material_items'][0]['post_ranges'][0]['additional_posts'][0]['s']='1400'
  elif fault=='type_proof':m['evidence']['node_type']='forged'
  else:p['takeoffs'][-1]['material_items'][0]['post_layout']['post_stations'][1]['original_node_index']=9
 mutate(pid,forge);coll='takeoffs' if fault=='book' else 'components';obj=get(e,pid)[coll][-1];res=r.post(f'/api/v1/projects/{pid}/reviews',json=dict(collection=coll,object_id=obj['id'],object_hash=r.get(f'/api/v1/projects/{pid}/hashes').json()[coll][obj['id']],decision='approve',reason='合成篡改应阻断'));assert res.status_code in (404,409,422),res.text


@pytest.mark.parametrize('equal',[False,True])
def test_same_mass_revision_exact_history_concurrency_and_source(clients,equal):
 e,r,pid,s,b,eid=setup(clients,equal=equal);c=post(e,pid,b).json();approve(r,pid,'components',c)
 for t in post(e,pid,{},'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
 post(e,pid,{},'/boq/generate');p=get(e,pid);old=p['components'][-1];old_boq=copy.deepcopy(p['boq']);new=copy.deepcopy(b);new['material_items'][0]['post_ranges'][0]['post_moves'][1]['s']='1750';new['object_hash']=store.digest(old);url=f'/api/v1/projects/{pid}/components/'+c['id'];res=e.put(url,json=new);assert res.status_code==200,res.text
 p=get(e,pid);assert p['component_history'][-1]==old and all(t['state']=='stale' for t in p['takeoffs']) and p['boq']==old_boq;assert res.json()['preview']['quantity']=='93.600';assert e.put(url,json=new).status_code==409;approve(r,pid,'components',res.json())
 for t in post(e,pid,{},'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
 g=get(e,pid)['takeoffs'][-1]['material_items'][0]['post_layout'];assert next(p for p in g['post_stations'] if p['range_label']=='S1 整段' and p['original_node_type']=='additional')['s']=='1.75'
 assert any(v['quantity']=='93.600' for v in post(e,pid,{},'/boq/generate').json())
 s['object_hash']=store.digest(s);s['edge_decisions'][0].update(mode='ranges',ranges=[interval(0,10,eid)]);res=e.put(f'/api/v1/projects/{pid}/components/'+s['id'],json=s);assert res.status_code==200;approve(r,pid,'components',res.json());assert get(e,pid)['components'][-1]['state']=='stale';post(e,pid,b,'/components/preview',409)


@pytest.mark.parametrize('gate',['preview','save','review','calculate','book_review'])
def test_original_bytes_every_gate(clients,gate):
 e,r,pid,s,b,eid=setup(clients)
 if gate in ('review','calculate','book_review'):c=post(e,pid,b).json()
 if gate in ('calculate','book_review'):approve(r,pid,'components',c)
 if gate=='book_review':t=post(e,pid,{},'/takeoffs/calculate').json()[-1]
 doc=get(e,pid)['documents'][-1];(store.ROOT/'originals'/doc['hash']).write_bytes(b'corrupt')
 if gate=='preview':post(e,pid,b,'/components/preview',409)
 elif gate=='save':post(e,pid,b,expected=409)
 elif gate=='review':approve(r,pid,'components',c,expected=409)
 elif gate=='calculate':post(e,pid,{},'/takeoffs/calculate',409)
 else:approve(r,pid,'takeoffs',t,expected=409)


def test_roles_tenant_foreign_type_proof(clients):
 e,r,pid,s,b,eid=setup(clients);post(r,pid,b,expected=403);assert clients['other'].get(f'/api/v1/projects/{pid}').status_code==404
 other=create(e)['id'];upload(e,other,'其他工程原件');b['material_items'][0]['post_ranges'][0]['post_moves'][1]['evidence']['node_type']=get(e,other)['evidence'][-1]['id'];post(e,pid,b,expected=404)


@pytest.mark.parametrize('equal',[False,True])
def test_irrational_original_boundary_and_added_neighbor_are_exact(equal):
 s=line_source(1);src=s['source_geometry'];src['route_points']=[point('A',0,0,0),point('B',1,1,0)];s['source_component_hash']=store.digest(snapshot(src));s['edge_decisions']=railing_payload(src)['edge_decisions'];b=body(s,equal=equal);r=b['material_items'][0]['post_ranges'][0]
 r.update(spacing='1',spacing_unit='m',additional_posts=[addition('.70710678118654752','m')],avoidance_ranges=[avoid('.70710678118654751','.70710678118654753','m')],post_moves=[move('additional',1,'.6','m')]+([move('regular',1,'.8','m')] if equal else []))
 with localcontext() as c:c.prec=6;g=calculate(b,s)['material_items'][0]['post_layout']
 assert g['post_ranges'][0]['avoidance_ranges'][0]['original_post_count']==(2 if equal else 1)
 r['post_moves'][0]['s']='.70710678118654753'
 with pytest.raises(DomainError):calculate(b,s)
 if equal:
  r.update(avoidance_ranges=[],post_moves=[move('additional',1,'.70710678118654753','m')])
  with pytest.raises(DomainError,match='重合或交叉'):calculate(b,s)


@pytest.mark.parametrize('equal',[False,True])
def test_addition_move_and_table_count_limits(equal):
 s=line_source(202);b=body(s,equal=equal);r=b['material_items'][0]['post_ranges'][0]
 r.update(additional_posts=[addition(str(2*i+Decimal('.5')),'m',label=str(i)) for i in range(100)],avoidance_ranges=[],post_moves=[move('additional',i+1,str(2*i+Decimal('.75')),'m') for i in range(100)])
 g=calculate(b,s)['material_items'][0]['post_layout'];assert g['additional_post_count']==100 and g['moved_post_count']==100
 r['post_moves'].append(move('regular',1,'2.2','m'))
 with pytest.raises(DomainError,match='100'):calculate(b,s)
 r['post_moves'].pop();r['additional_posts'].append(addition('200.5','m',label='101'))
 with pytest.raises(DomainError,match='100'):calculate(b,s)
 # 999 original regular roots plus one additional root meet the per-item cap.
 s=line_source(1996);b=body(s,equal=equal);r=b['material_items'][0]['post_ranges'][0];r.update(additional_posts=[addition('.5','m')],avoidance_ranges=[],post_moves=[move('additional',1,'.75','m')]);v=calculate(b,s);assert v['material_items'][0]['post_layout']['post_count']==1000
 b['material_items']=[{**copy.deepcopy(b['material_items'][0]),'name':str(i)} for i in range(5)];assert sum(i['post_layout']['post_count'] for i in calculate(b,s)['material_items'])==5000
 b['material_items'].append({**copy.deepcopy(b['material_items'][0]),'name':'6'})
 with pytest.raises(DomainError,match='5000'):calculate(b,s)
 r=b['material_items'][0]['post_ranges'][0];r['additional_posts'].append(addition('1','m',label='1001'));b['material_items']=b['material_items'][:1]
 with pytest.raises(DomainError,match='1000'):calculate(b,s)


@pytest.mark.parametrize('equal',[False,True])
@pytest.mark.parametrize('starts',[False,True])
def test_closed_polygon_hole_and_shared_ownership(equal,starts):
 from tests.test_railing_takeoff import geometry as pg,payload as pp
 src=pg();s=pp(src);s.update(id='polygon',revision=1,source_geometry=snapshot(src))
 for d in s['edge_decisions']:d.update(mode='full',ranges=[])
 b=body(s,equal=equal)
 for r in b['material_items'][0]['post_ranges']:
  r.update(start_post=starts,end_post=not starts,additional_posts=[addition('.5','m')],avoidance_ranges=[avoid('.4','.6','m')],post_moves=[move('additional',1,'.7','m')])
 g=calculate(b,s)['material_items'][0]['post_layout'];assert g['additional_post_count']==len(g['post_ranges']) and g['moved_post_count']==len(g['post_ranges'])
 b['material_items'][0]['post_ranges'][0]['end_post' if starts else 'start_post']=True
 with pytest.raises(DomainError,match='共享端点'):calculate(b,s)


def test_mixed_modes_preserve_each_rule_and_namespace():
 s=geometry();b=body(s);other=body(s,equal=True)['material_items'][0];other['name']='另一种均分';b['material_items'].append(other);v=calculate(b,s);assert v['quantity']=='187.200'
 with localcontext() as c:c.prec=80;assert evaluate(v['formula_ast'],v['dimensions'])==Decimal('187.2')
 assert {p['original_node_type'] for i in v['material_items'] for p in i['post_layout']['post_stations']}=={'regular','additional'}
