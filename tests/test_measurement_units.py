"""Exact source-unit arithmetic and approval regression, using synthetic evidence only."""
import copy
import io
from decimal import Decimal, localcontext
import pytest
from openpyxl import load_workbook
from pypdf import PdfReader
from backend.compound_takeoff import calculate
from backend.engine import DomainError
from backend.measurement_units import profiles
from tests.test_compound_takeoff import payload, ast_value, setup
from tests.test_workflow import get, approve, priced_project


def unit_payload(eid='test'):
    body=payload(eid)
    for part in body['measurement_parts']:
        part['parameter_units']={k:'m' for k in part['parameters']}
    part=body['measurement_parts'][0]
    part['parameters']={'length':'10000','width':'500','thickness':'300'}
    part['parameter_units']={'length':'mm','width':'cm','thickness':'mm'}
    return body


def test_mixed_source_units_preserve_raw_values_and_reproduce_ast():
    body=unit_payload();original=copy.deepcopy(body);result=calculate(body)
    assert body==original
    assert Decimal(result['quantity'])==Decimal('15.2') and result['rule_version']=='compound-geometry-1.1'
    row=result['measurement_parts'][0]
    assert row['parameters']['length']=='10000' and row['parameter_units']['length']=='mm'
    assert row['normalized_parameters']=={'length':'10','width':'5','thickness':'0.3'}
    assert row['unit_conversions'][0]['factor']=='0.001'
    assert result['dimensions']['1.length']=='10000'
    with localcontext() as ctx:
        ctx.prec=80
        assert ast_value(result['formula_ast'],result['dimensions'])==Decimal(result['quantity'])


@pytest.mark.parametrize('kind,params,units,expected',[
    ('pipe',{'length':'1.25'},{'length':'km'},'1250'),
    ('area',{'length':'200','width':'500'},{'length':'cm','width':'mm'},'1'),
    ('walls',{'length':'1000','width':'5000','thickness':'250','height':'300'},{'length':'cm','width':'mm','thickness':'mm','height':'cm'},'21.75'),
    ('masonry',{'length':'1000','height':'30','thickness':'200'},{'length':'mm','height':'dm','thickness':'mm'},'.6'),
    ('trench',{'length':'1000','top_width':'200','bottom_width':'1000','depth':'20'},{'length':'cm','top_width':'cm','bottom_width':'mm','depth':'dm'},'30'),
    ('backfill',{'excavation':'15000000000','occupied':'500000'},{'excavation':'mm³','occupied':'cm³'},'14.5'),
    ('rebar_schedule',{'length':'3450','count':'10','unit_weight':'617'},{'length':'mm','count':'1','unit_weight':'g/m'},'21.2865'),
    ('rebar_schedule',{'length':'2','count':'3','unit_weight':'.617'},{'length':'m','count':'1','unit_weight':'g/mm'},'3.702'),
])
def test_supported_dimensions_and_exact_factors(kind,params,units,expected):
    body=unit_payload();row=body['measurement_parts'][0]
    row.update(kind=kind,parameters=params,parameter_units=units,evidence={k:'test' for k in [*params,'multiplicity','measurement_basis']})
    body['measurement_parts']=[row]
    result=calculate(body)
    assert Decimal(result['quantity'])==Decimal(expected)
    with localcontext() as ctx:
        ctx.prec=80
        assert ast_value(result['formula_ast'],result['dimensions'])==Decimal(expected)


def test_ring_conversion_keeps_pi_precision():
    body=unit_payload();row=body['measurement_parts'][0]
    row.update(kind='ring',parameters={'outer_diameter':'3000','inner_diameter':'200','height':'40'},parameter_units={'outer_diameter':'mm','inner_diameter':'cm','height':'dm'},evidence={k:'test' for k in ['outer_diameter','inner_diameter','height','multiplicity','measurement_basis']})
    body['measurement_parts']=[row]
    result=calculate(body)
    with localcontext() as ctx:
        ctx.prec=80
        assert Decimal(result['quantity'])==Decimal('3.1415926535897932384626433832795028841971693993751')*5
        assert ast_value(result['formula_ast'],result['dimensions'])==Decimal(result['quantity'])


@pytest.mark.parametrize('case',['empty','missing','extra','wrong_dimension','unknown','not_string','underflow','overflow','nan'])
def test_invalid_source_units_and_converted_ranges(case):
    body=unit_payload();row=body['measurement_parts'][0]
    if case=='empty':row['parameter_units']={}
    elif case=='missing':del row['parameter_units']['width']
    elif case=='extra':row['parameter_units']['deduction']='m³'
    elif case=='wrong_dimension':row['parameter_units']['length']='cm³'
    elif case=='unknown':row['parameter_units']['length']='pixel'
    elif case=='not_string':row['parameter_units']['length']=['mm']
    elif case=='underflow':row['parameters']['length']='.000000000000000001'
    elif case=='overflow':row['parameters']['length']='1000000000000';row['parameter_units']['length']='km'
    else:row['parameters']['length']='NaN'
    with pytest.raises(DomainError):calculate(body)


def test_legacy_compounds_keep_original_calculation_schema(clients):
    e,r,pid,body=setup(clients);base='/api/v1/projects/'+pid
    legacy=calculate(body)
    assert legacy['rule_version']=='compound-geometry-1.0'
    assert all('unit_conversions' not in row and 'normalized_parameters' not in row for row in legacy['measurement_parts'])
    comp=e.post(base+'/components',json=body).json()
    approve(r,pid,'components',comp)
    result=e.post(base+'/takeoffs/calculate')
    assert result.status_code==200,result.text
    assert result.json()[0]['formula_ast']==legacy['formula_ast']


def test_authenticated_profiles_and_dimension_specific_units(clients):
    c=clients['estimator'];response=c.get('/api/v1/takeoff-unit-types')
    assert response.status_code==200
    assert response.json()==profiles()
    assert response.json()['slab']['length']['canonical']=='m'
    assert 'cm³' not in response.json()['slab']['length']['factors']
    assert response.json()['backfill']['excavation']['factors']['cm³']=='0.000001'
    assert set(response.json()['rebar_schedule']['count']['factors'])=={'1'}
    from fastapi.testclient import TestClient
    from backend.main import app
    with TestClient(app) as anonymous:assert anonymous.get('/api/v1/takeoff-unit-types').status_code==401


def test_source_units_review_boq_and_exported_conversion_book(clients):
    e,r,pid,_=setup(clients);base='/api/v1/projects/'+pid;eid=get(e,pid)['evidence'][0]['id']
    body=unit_payload(eid)
    preview=e.post(base+'/components/preview',json=body)
    assert preview.status_code==200,preview.text
    comp=e.post(base+'/components',json=body).json();approve(r,pid,'components',comp)
    book=e.post(base+'/takeoffs/calculate').json()[0];approve(r,pid,'takeoffs',book)
    assert Decimal(e.post(base+'/boq/generate').json()[0]['quantity'])==Decimal('15.2')
    response=e.get(base+'/exports?format=xlsx&unpriced=true')
    assert response.status_code==200,response.text
    wb=load_workbook(io.BytesIO(response.content),data_only=False)
    assert '算量原始单位换算' in wb.sheetnames
    rows=list(wb['算量原始单位换算'].values)
    assert any('10000' in str(row) and 'mm' in str(row) and '0.001' in str(row) for row in rows)
    assert all(cell.data_type!='f' for sheet in wb for row in sheet for cell in row)
    response=e.get(base+'/exports?format=json&unpriced=true')
    assert response.json()['unpriced']['takeoffs'][0]['dimensions']['1.length']=='10000'
    pdf=e.get(base+'/exports?format=pdf&unpriced=true')
    assert pdf.status_code==200,pdf.text
    text='\n'.join(page.extract_text() for page in PdfReader(io.BytesIO(pdf.content)).pages)
    assert '算量原始单位换算' in text and 'measurement-units-1.0' in text.replace('\n','')


def test_unit_revision_invalidates_book_and_preserves_old_quote(clients):
    e,r,pid,_,quote=priced_project(clients);base='/api/v1/projects/'+pid
    eid=get(e,pid)['evidence'][0]['id'];body=unit_payload(eid)
    comp=e.post(base+'/components',json=body).json()
    # The existing simple component was invalidated by adding the new one.
    for current in get(e,pid)['components']:approve(r,pid,'components',current)
    books=e.post(base+'/takeoffs/calculate').json()
    for book in books:approve(r,pid,'takeoffs',book)
    old=copy.deepcopy(get(e,pid)['quotes'][0])
    body['object_hash']=e.get(base+'/hashes').json()['components'][comp['id']]
    body['measurement_parts'][0]['parameter_units']['length']='cm'
    response=e.put(base+'/components/'+comp['id'],json=body)
    assert response.status_code==200,response.text
    assert response.json()['state']=='draft'
    assert Decimal(response.json()['preview']['quantity'])==Decimal('150.2')
    assert all(book['state']=='stale' for book in get(e,pid)['takeoffs'])
    assert get(e,pid)['quotes'][0]==old and old['id']==quote['id']
