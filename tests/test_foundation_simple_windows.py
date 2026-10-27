"""Synthetic concave windows: independent area/oracle and existing workflow gates."""
import copy
import io
import math
import random
from fractions import Fraction
from decimal import Decimal, localcontext
import pytest
from openpyxl import load_workbook
from pypdf import PdfReader
from backend import store
from backend.engine import DomainError
from backend.foundation_horizontal_formwork import calculate
from backend.foundation_simple_windows import VERSION, separation
from tests.test_foundation_polygon_windows import polygon, body as convex_body, replay, clip, setup as convex_setup
from tests.test_foundation_horizontal_windows import source, window
from tests.test_foundation_hole_end_formwork import source as end_source
from tests.test_foundation_openings import hole
from tests.test_foundation_voids import rectangle
from tests.test_foundation_recesses import interval
from tests.test_foundation_round_step_windows import adaptive
from tests.test_post_foundation import equal_body, fixed_body, fixed_added, equal_added, fixed_adjusted, equal_adjusted, combined
from tests.test_polygon_takeoff import post, mutate
from tests.test_workflow import get, approve, upload, create, priced_project

C_SHAPE=[(-100,-100),(100,-100),(100,-40),(-40,-40),(-40,40),(100,40),(100,100),(-100,100)]


def config(pts=C_SHAPE,eid='test',unit='mm'):
    c=polygon(pts,eid,unit);c['shape']='simple_polygon';return c


def body(s,eid='test',face='bottom',pts=C_SHAPE):
    b=convex_body(s,eid,face);b['contact_ranges'][0]['plan_range']=config(pts,eid);return b


@pytest.mark.parametrize('shape',['rectangular','circular','stepped_rectangular','stepped_circular'])
@pytest.mark.parametrize('face',['top','bottom','hole_top','hole_bottom'])
@pytest.mark.parametrize('make',[equal_body,fixed_body,fixed_added,equal_added,fixed_adjusted,equal_adjusted,combined,lambda s:combined(s,equal=True)])
def test_shapes_ends_counts_multi_span_replay_and_immutable(shape,face,make):
    s=end_source(shape,make) if face.startswith('hole') else source(shape,make)
    b=body(s,face=face);before=copy.deepcopy((s,b));r=calculate(b,s);d=r['contact_ranges'][0]
    assert before==(s,b) and r['rule_version']==VERSION and d['simple_polygon_window']
    assert Decimal(d['window_area'])==Decimal('.0288')
    assert any(q['window_region_index']==2 for q in d['polygon_net_trace']['partitions']) or face.startswith('hole')
    replay(r)


@pytest.mark.parametrize('shape,face',[(shape,face) for shape in ['rectangular','circular','stepped_rectangular','stepped_circular'] for face in ['top','bottom','hole_top','hole_bottom','step'] if face!='step' or shape.startswith('stepped_')])
def test_concave_window_equals_three_disjoint_rectangles(shape,face):
    s=end_source(shape);r=calculate(body(s,face=face),s);d=r['contact_ranges'][0]
    expected={k:Decimal(0) for k in ('single_gross_area','single_opening_area','single_net_area')}
    with localcontext() as ctx:
        ctx.prec=80
        for bounds in [('-100','-40','-100','100'),('-40','100','-100','-40'),('-40','100','40','100')]:
            b=convex_body(s,face=face);b['contact_ranges'][0]['plan_range']=window('test',*bounds)
            part=calculate(b,s)['contact_ranges'][0]
            for k in expected:expected[k]+=Decimal(part[k])
        assert all(abs(Decimal(d[k])-v)<Decimal('1e-75') for k,v in expected.items())
    replay(r)


def test_40_seeded_slanted_concave_shapes_independent_rectangle_clip():
    rng=random.Random(10801);s=source()
    for _ in range(40):
        sx,sy=rng.randint(-200,200),rng.randint(-200,200);shear=rng.choice([-2,-1,0,1,2])
        pts=[(x+shear*y+sx,y+sy) for x,y in C_SHAPE]
        r=calculate(body(s,pts=pts),s);d=r['contact_ranges'][0]
        expected=(clip(pts,(-200,200,-250,250))-clip(pts,(100,200,-30,30)))/1e6
        assert abs(float(d['single_net_area'])-expected)<3e-16
    replay(r)


def oracle(pts,include,exclude=()):
    xs=sorted({p[0] for p in pts}|{x+s*r for x,y,r in include+exclude for s in (-1,0,1)})
    def f(x):
        ys=sorted(a[1]+(x-a[0])*(b[1]-a[1])/(b[0]-a[0]) for a,b in zip(pts,pts[1:]+pts[:1]) if min(a[0],b[0])<x<max(a[0],b[0]))
        spans=list(zip(ys[::2],ys[1::2]))
        for cx,cy,r in include:
            if abs(x-cx)>=r:return 0.
            dy=math.sqrt(max(0,r*r-(x-cx)**2));spans=[(max(lo,cy-dy),min(hi,cy+dy)) for lo,hi in spans if max(lo,cy-dy)<min(hi,cy+dy)]
        for cx,cy,r in exclude:
            if abs(x-cx)>=r:continue
            dy=math.sqrt(max(0,r*r-(x-cx)**2));cut=(cy-dy,cy+dy);out=[]
            for lo,hi in spans:
                if cut[1]<=lo or cut[0]>=hi:out.append((lo,hi));continue
                if lo<cut[0]:out.append((lo,cut[0]))
                if cut[1]<hi:out.append((cut[1],hi))
            spans=out
        return sum(hi-lo for lo,hi in spans)
    return sum(adaptive(lambda x:f(math.nextafter(a,b) if x==a else math.nextafter(b,a) if x==b else x),a,b,1e-13) for a,b in zip(xs,xs[1:]))


@pytest.mark.parametrize('face',['top','bottom','step','hole_top'])
def test_slanted_multi_span_curves_independent_quadrature(face):
    hs=[{**hole(x='190',diameter='40'),'vertical_range':interval('0','200')},
        {**hole(x='175',name='相接孔',diameter='40'),'vertical_range':interval('200','400')}]
    s=end_source('stepped_circular',holes=hs)
    for shift in (-10,0,10):
        pts=[(190+x*.35+y*.1+shift,y*.35) for x,y in C_SHAPE]
        b=body(s,face=face,pts=pts);r=calculate(b,s);p=[(x/1000,y/1000) for x,y in pts]
        inc=[(0,0,.2 if face in ('top','hole_top') else .3)];exc=[]
        if face=='hole_top':inc.append((.19,0,.02));exc=[(.175,0,.02)]
        elif face=='bottom':exc=[(.19,0,.02)]
        elif face=='top':exc=[]
        else:exc=[(0,0,.2),(.19,0,.02)]
        assert abs(float(r['contact_ranges'][0]['single_net_area'])-oracle(p,inc,exc))<8e-11
    replay(r)


@pytest.mark.parametrize('unit,factor',[('mm','1'),('cm','.1'),('dm','.01'),('m','.001'),('km','.000001')])
@pytest.mark.parametrize('reverse',[False,True])
def test_units_direction_and_low_call_precision(unit,factor,reverse):
    s=source('circular');expected=calculate(body(s),s)
    pts=list(reversed(C_SHAPE)) if reverse else C_SHAPE;b=body(s,pts=pts);c=b['contact_ranges'][0]['plan_range'];c['unit']=unit
    c['parameters']={k:str(Decimal(v)*Decimal(factor)) for k,v in c['parameters'].items()}
    with localcontext() as ctx:ctx.prec=6;r=calculate(b,s)
    assert r['quantity']==expected['quantity'] and r['gross']==expected['gross'];replay(r)


@pytest.mark.parametrize('pts',[
    [(0,0),(1,0)],[(0,0),(1,0),(1,0)],[(0,0),(1,0),(2,0)],
    [(0,0),(100,100),(0,100),(100,0)],[(0,0),(100,0),(100,100),(50,0),(0,100)],
    [(0,0),(100,0),(100,100),(50,100),(0,100)],[(0,100),(59,-81),(-95,31),(95,31),(-59,-81)]])
def test_invalid_topology_rejected_without_repair(pts):
    s=source()
    with pytest.raises(DomainError):calculate(body(s,pts=pts),s)


@pytest.mark.parametrize('fault',['extra','missing','empty_proof','bool','nan','unit','confirm','too_many','nonconsecutive'])
def test_strict_schema_and_evidence(fault):
    s=source();b=body(s);c=b['contact_ranges'][0]['plan_range']
    if fault=='extra':c['holes']=[]
    elif fault=='missing':del c['parameters']['x1']
    elif fault=='empty_proof':c['evidence']['x1']=' '
    elif fault=='bool':c['parameters']['x1']=True
    elif fault=='nan':c['parameters']['x1']='NaN'
    elif fault=='unit':c['unit']='m²'
    elif fault=='confirm':c['scope_confirmed']=1
    elif fault=='too_many':c['parameters']={f'{a}{i}':str(i) for i in range(1,14) for a in ('x','y')}
    else:c['parameters']['x0']=c['parameters'].pop('x1')
    with pytest.raises(DomainError):calculate(b,s)


def test_twelve_vertices_three_vertical_pieces():
    pts=[(-100,-100),(100,-100),(100,-60),(20,-60),(20,-20),(80,-20),(80,20),(20,20),(20,60),(100,60),(100,100),(-100,100)]
    s=source('circular');r=calculate(body(s,pts=pts),s);d=r['contact_ranges'][0]
    assert Decimal(d['window_area'])==Decimal('.0328') and any(q['window_region_index']==3 for q in d['polygon_net_trace']['partitions'])
    replay(r)


@pytest.mark.parametrize('gap',['0','0.000000001'])
def test_all_lobes_zero_proof_or_tiny_positive_gap(gap):
    v=str(Decimal('25')+Decimal(gap));hs=[{**hole(x='0',diameter='80'),'vertical_range':interval('100','300')},
      {**rectangle(name='H2',x='-'+v,length='50',width='100'),'vertical_range':interval('300','500')},
      {**rectangle(name='H3',x=v,length='50',width='100'),'vertical_range':interval('300','500')}]
    s=end_source(holes=hs);pts=[(x/2,y/2) for x,y in C_SHAPE];r=calculate(body(s,face='hole_top',pts=pts),s);d=r['contact_ranges'][0]
    assert ('polygon_zero_contact' in d)==(gap=='0') and (Decimal(d['single_net_area'])==0)==(gap=='0')
    assert any(q['window_region_index']==2 for q in d['polygon_net_trace']['partitions']);replay(r)


def test_concave_bay_disjoint_boundary_touch_containment_and_tiny_overlap():
    a={'plan_range':config()}
    for other,expected in [(window('test','0','80','-20','20'),False),(window('test','-40','80','-40','40'),False),
        (window('test','-40.000000001','80','-40','40'),True),(config([(-200,-200),(200,-200),(200,200),(-200,200)]),True)]:
        r=separation(a,{'plan_range':other});assert r['overlap']==expected
    r=separation(a,{'plan_range':window('test','-40','80','-40','40')});assert r['state']=='boundary_touch'
    assert separation(a,{'plan_range':window('test','0','80','-20','20')})['state']=='separated'


def test_100_seeded_overlap_checks_against_independent_rectangle_clipping():
    rng=random.Random(10802);a={'plan_range':config()}
    for _ in range(100):
        x,y=rng.randint(-160,160),rng.randint(-160,160);X,Y=x+rng.randint(1,180),y+rng.randint(1,180)
        exact_points=[tuple(map(Fraction,p)) for p in C_SHAPE]
        assert separation(a,{'plan_range':window('test',str(x),str(X),str(y),str(Y))})['overlap']==(clip(exact_points,tuple(map(Fraction,(x,X,y,Y))))>0)


def test_window_check_witnesses_follow_named_left_and_right_records():
    from backend.foundation_horizontal_windows import checks
    left=dict(index=1,name='C形',material_index=1,face='top',plan_range=config())
    right=dict(index=2,name='凹口内矩形',material_index=1,face='top',plan_range=window('test','0','80','-20','20'))
    check=checks([left,right])[0]
    assert check['left_index']==1 and check['left']=='C形'
    assert check['right_index']==2 and check['right']=='凹口内矩形'
    assert check['partitions']==separation(left,right)['partitions']
    assert all(len(q['left'])==2 and len(q['right'])==1 for q in check['partitions'])


def setup(clients,pid=None,zero=False):
    e,r,pid,s,b,eid=convex_setup(clients,pid,zero)
    pts=[(x/2,y/2) for x,y in C_SHAPE] if zero else [(190+x*.35,y*.35) for x,y in C_SHAPE]
    b['contact_ranges'][0]['plan_range']=config(pts,eid);return e,r,pid,s,b,eid


def test_workflow_readonly_foreign_roles_history_claims_and_source_stale(clients):
    e,r,pid,s,b,eid=setup(clients);before=get(e,pid);post(e,pid,b,'/components/preview');assert get(e,pid)==before
    post(r,pid,b,expected=403);assert clients['other'].get(f'/api/v1/projects/{pid}').status_code==404
    foreign=create(e)['id'];upload(e,foreign,'异项目合成资料');bad=copy.deepcopy(b);bad['contact_ranges'][0]['plan_range']['evidence']['x1']=get(e,foreign)['evidence'][-1]['id'];post(e,pid,bad,expected=404)
    c=post(e,pid,b).json();approve(r,pid,'components',c);post(e,pid,b,expected=409)
    c=next(x for x in get(e,pid)['components'] if x['id']==c['id']);b['object_hash']=store.digest(c);b['name']='合成凹窗口V2'
    res=e.put(f'/api/v1/projects/{pid}/components/'+c['id'],json=b);assert res.status_code==200,res.text
    assert get(e,pid)['component_history'][-1]==c;approve(r,pid,'components',res.json())
    s['object_hash']=store.digest(s);s['foundation_items'][0]['levels'][1]['parameters']['diameter']='420'
    res=e.put(f'/api/v1/projects/{pid}/components/'+s['id'],json=s);assert res.status_code==200,res.text;approve(r,pid,'components',res.json())
    assert next(x for x in get(e,pid)['components'] if x['id']==c['id'])['state']=='stale'


def test_cross_record_notch_window_allowed_actual_overlap_and_frame_rejected(clients):
    e,r,pid,s,b,eid=setup(clients);post(e,pid,b)
    other=copy.deepcopy(b);other['name']='凹口内独立矩形';other['contact_ranges'][0]['plan_range']=window(eid,'180','200','-10','10')
    post(e,pid,other)
    bad=copy.deepcopy(other);bad['contact_ranges'][0]['plan_range']['parameters']['x_start']='175.999999999'
    before=get(e,pid);post(e,pid,bad,expected=409);assert get(e,pid)==before
    bad=copy.deepcopy(other);bad['contact_ranges'][0]['plan_range']['parameters'].update(x_start='230',x_end='250');bad['contact_ranges'][0]['plan_range']['coordinate_basis']='不同基准'
    post(e,pid,bad,expected=409)


@pytest.mark.parametrize('fault',['region','definition','trace','point','proof'])
def test_tamper_rejected_during_independent_review(clients,fault):
    e,r,pid,s,b,eid=setup(clients,zero=True);c=post(e,pid,b).json()
    def forge(p):
        c=p['components'][-1];d=c['preview']['contact_ranges'][0]
        if fault=='region':d['polygon_net_trace']['partitions'][0]['window_region_index']=99
        elif fault=='definition':d['window_expression_definitions'][0]['value']='999'
        elif fault=='trace':d['polygon_gross_trace']['terms'][0]['area']='999'
        elif fault=='point':c['contact_ranges'][0]['plan_range']['parameters']['x1']='-51'
        else:d['polygon_zero_contact']['net_trace']['partitions'][0]['remaining']=[{}]
    mutate(pid,forge);approve(r,pid,'components',c,expected=409)


@pytest.mark.parametrize('gate',['preview','save','review','calculate','book_review'])
def test_actual_original_bytes_all_five_gates(clients,gate):
    e,r,pid,s,b,eid=setup(clients)
    if gate in ('review','calculate','book_review'):c=post(e,pid,b).json()
    if gate in ('calculate','book_review'):approve(r,pid,'components',c)
    if gate=='book_review':t=post(e,pid,{},'/takeoffs/calculate').json()[-1]
    doc=get(e,pid)['documents'][-1];(store.ROOT/'originals'/doc['hash']).write_bytes(b'corrupted-synthetic-original')
    if gate=='preview':post(e,pid,b,'/components/preview',409)
    elif gate=='save':post(e,pid,b,expected=409)
    elif gate=='review':approve(r,pid,'components',c,expected=409)
    elif gate=='calculate':post(e,pid,{},'/takeoffs/calculate',409)
    else:approve(r,pid,'takeoffs',t,expected=409)


@pytest.mark.parametrize('fmt',['json','xlsx','pdf'])
@pytest.mark.parametrize('zero',[False,True])
def test_three_formats_new_tables_and_old_quote_isolation(clients,fmt,zero):
    from backend.foundation_horizontal_formwork_reports import SIMPLE_POLYGON_TITLES
    e,r,pid,_,q=priced_project(clients);e,r,pid,s,b,eid=setup(clients,pid,zero);c=post(e,pid,b).json();approve(r,pid,'components',c)
    for t in post(e,pid,{},'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
    post(e,pid,{},'/boq/generate');url=f'/api/v1/projects/{pid}/exports';res=e.get(url,params=dict(format=fmt,unpriced=True));assert res.status_code==200,res.text
    titles=set(SIMPLE_POLYGON_TITLES if zero else SIMPLE_POLYGON_TITLES[:2])
    if fmt=='json':assert titles<={t[0] for t in res.json()['tables']}
    elif fmt=='xlsx':assert titles<=set(load_workbook(io.BytesIO(res.content)).sheetnames)
    else:assert all(any(t in page.extract_text() for page in PdfReader(io.BytesIO(res.content)).pages) for t in titles)
    assert not any(t[0] in titles for t in e.get(url,params=dict(format='json',quote_id=q['id'])).json()['tables'])
