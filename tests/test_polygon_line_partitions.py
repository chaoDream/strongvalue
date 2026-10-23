"""Synthetic infinite-line net geometry; no real engineering approval implied."""
import copy
import io
import random
from decimal import Decimal, localcontext
from fractions import Fraction
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pypdf import PdfReader
from backend import store
from backend.main import app
from backend.engine import DomainError
from backend.surface_takeoff import boundary, partition, calculate
from tests.test_polygon_cap_partitions import geometry, groups as windows, setup as cap_setup, post, mutate
from tests.test_polygon_takeoff import ring
from tests.test_surface_takeoff import payload as surface
from tests.test_polygon_side_partitions import cuts as side_cuts
from tests.test_workflow import approve, get, create, priced_project, upload
from tests.test_route_takeoff import evaluate


def line(eid='test', a=(0,0), b=(6,6), unit='m'):
    return dict(zip(('x1','y1','x2','y2'),map(str,(*a,*b))),unit=unit,evidence={k:eid for k in ('x1','y1','x2','y2')},
                basis='合成斜线无限延伸，按规范侧核定不同防水做法；非专业依据',evidence_id=eid)


def groups(s,eid='test',lines=None,top=True,xy=False):
    g=windows(s,eid,top=top)
    if not xy:g[0].update(u_cuts=[],v_cuts=[])
    g[0]['line_cuts']=lines if lines is not None else [line(eid)]
    return g


def children(r):return [f for f in r['faces'] if f.get('cap_cut_sides')]


def fraction(f):return Fraction(int(f['cap_rational_area']['numerator']),int(f['cap_rational_area']['denominator']))


def payload(s,eid='test',g=None):
    g=groups(s,eid) if g is None else g;view=partition(s,g);b=surface(s,eid,included=lambda f:False)
    b.update(face_subdivisions=g,name='合成斜向防水范围（非真实工程）',purpose='waterproof',
             measurement_basis='合成只计入顶部斜线L1规范正侧，其他面排除，14m²×2=28m²',
             face_decisions=[dict(face_id=f['face_id'],included=bool(f.get('cap_cut_sides')) and f['cap_cut_sides'][0]['positive'],
                                  basis='合成逐区有据核定，非凸框面积、非专业签审',evidence_id=eid) for f in view['faces']])
    return b


def setup(clients,pid=None):
    e,r,pid,s,_,eid=cap_setup(clients,pid)
    upload(e,pid,'合成追加斜向做法依据（非真实工程）：顶部以共同设计点(0,0)m及(6,6)m定义无限直线，不是有限线段。外轮廓两侧均16m²，贯通方孔各扣2m²，净区各14m²；仅X−Y正侧计入，重复2，工程量28m²。修订为(0,1)至(6,7)m后两净区16.5、11.5m²，仅正侧计入，重复2为33m²。','line-partition-basis.txt')
    for c in get(e,pid)['components']:approve(r,pid,'components',c)
    s=next(c for c in get(e,pid)['components'] if c['id']==s['id'])
    eid=get(e,pid)['evidence'][-1]['id'];return e,r,pid,s,payload(s,eid),eid


@pytest.mark.parametrize('top',[False,True])
def test_diagonal_through_concavity_and_hole_exact_areas_ast_and_closure(top):
    s=geometry();g=groups(s,top=top);before=copy.deepcopy((s,g));r=partition(s,g)
    assert (s,g)==before and len(r['faces'])==13
    assert [fraction(f) for f in children(r)]==[14,14]
    assert all([Decimal(c['area']) for c in f['cap_clip_rings']]==[16,2] for f in children(r))
    out=calculate(payload(s,g=g),s);assert out['quantity']=='28' and out['rule_version']=='polygon-boundary-1.3'
    with localcontext() as ctx:
        ctx.prec=80
        assert evaluate(out['formula_ast'],out['dimensions'])==28
        for f in out['surface_faces']:assert evaluate(f['formula_ast'],out['dimensions'])==Decimal(f['area'])
    check=r['surface_cap_partition_checks'][0];assert check['exact_rational_tiling'] and check['area_difference']=='0'


def test_xy_windows_plus_diagonal_have_six_regions_and_count_once():
    s=geometry();g=groups(s,xy=True);r=partition(s,g)
    assert [fraction(f) for f in children(r)]==[Fraction('1.5'),Fraction('1.5'),7,7,Fraction('5.5'),Fraction('5.5')]
    assert sum(fraction(f) for f in children(r))==28
    assert Decimal(calculate(payload(s,g=g),s)['quantity'])==28


@pytest.mark.parametrize('mode',['reverse_endpoints','other_endpoints','units','order'])
def test_equivalent_lines_and_input_order_keep_geometric_identity(mode):
    s=geometry();g=groups(s,lines=[line(),line(a=(0,6),b=(6,0))]);other=copy.deepcopy(g)
    if mode=='reverse_endpoints':
        for c in other[0]['line_cuts']:c['x1'],c['x2'],c['y1'],c['y2']=c['x2'],c['x1'],c['y2'],c['y1']
    elif mode=='other_endpoints':other[0]['line_cuts'][0].update(x1='-5',y1='-5',x2='9',y2='9')
    elif mode=='units':
        for c in other[0]['line_cuts']:
            c['unit']='mm'
            for k in ('x1','y1','x2','y2'):c[k]=str(Decimal(c[k])*1000)
    else:other[0]['line_cuts'].reverse()
    def ids(g):return sorted((f['face_id'],fraction(f)) for f in children(partition(s,g)))
    assert ids(g)==ids(other)


def test_adding_a_cut_keeps_unchanged_domains_and_old_faces_are_not_reused():
    s=geometry();g=groups(s);a=partition(s,g);g[0]['line_cuts'].append(line(a=(0,-1),b=(6,5)));b=partition(s,g)
    first={f['face_id'] for f in children(a)};second={f['face_id'] for f in children(b)}
    assert len(first&second)==1 and len(second)==3
    assert not first & {f['face_id'] for f in boundary(s)['faces']}


def test_nonterminating_intersections_save_exact_convex_domain_and_area():
    s=geometry();s['polygon_rings']=[ring('方形',[(0,0),(2,0),(2,2),(0,2)])]
    r=partition(s,groups(s,lines=[line(a=(0,0),b=(3,2))]))
    assert sorted(fraction(f) for f in children(r))==[Fraction(4,3),Fraction(8,3)]
    exact_y=[Fraction(int(p['exact_y']['numerator']),int(p['exact_y']['denominator'])) for f in children(r) for p in f['cap_clip_domain']]
    assert Fraction(4,3) in exact_y
    with localcontext() as ctx:
        ctx.prec=200;c=r['surface_cap_partition_checks'][0]
        assert Decimal(c['children_area'])==sum((Decimal(f['area']) for f in children(r)),Decimal(0))
        assert Decimal(c['area_difference'])==Decimal(c['children_area'])-4


def below_segment(m,b,x0,x1,height):
    """Independent integral of clamped affine height over one staircase column."""
    knots={Fraction(x0),Fraction(x1)}
    if m:
        for y in (0,height):
            x=(y-b)/m
            if x0<x<x1:knots.add(x)
    knots=sorted(knots);total=Fraction(0)
    for a,z in zip(knots,knots[1:]):
        va=max(Fraction(0),min(Fraction(height),m*a+b));vz=max(Fraction(0),min(Fraction(height),m*z+b))
        total+=(z-a)*(va+vz)/2
    return total


def test_thirty_random_concave_staircases_match_independent_vertical_integral():
    rng=random.Random(6201)
    for _ in range(30):
        w=6;hs=[rng.randint(2,7) for _ in range(w)];pts=[(0,0),(w,0),(w,hs[-1])]
        for x in range(w-1,0,-1):
            if hs[x]!=hs[x-1]:pts.extend([(x,hs[x]),(x,hs[x-1])])
        pts.append((0,hs[0]));s=geometry();s['polygon_rings']=[ring('阶梯',pts)]
        # Keep input values finite decimal while intersections may be repeating.
        dy=rng.randint(2,9);m=Fraction(dy,w);b=Fraction(rng.randint(-1,1));cut=line(a=(0,b),b=(w,b+dy))
        r=partition(s,groups(s,lines=[cut]));expected=sum((below_segment(m,b,x,x+1,h) for x,h in enumerate(hs)),Fraction(0))
        actual=next(fraction(f) for f in children(r) if f['cap_cut_sides'][0]['positive'])
        assert actual==expected and sum(fraction(f) for f in children(r))==sum(hs)


@pytest.mark.parametrize('transform',['reverse_rings','translate','swap','negative_slope'])
def test_orientation_origin_axes_and_slope(transform):
    s=geometry();g=groups(s)
    for row in s['polygon_rings']:
        if transform=='reverse_rings':row['points'].reverse()
        for p in row['points']:
            x,y=(Decimal(p['parameters'][k]) for k in ('x','y'))
            if transform=='translate':x-=100;y-=200
            elif transform=='swap':x,y=y,x
            elif transform=='negative_slope':y=-y
            p['parameters']={'x':str(x),'y':str(y)}
    if transform=='translate':g=groups(s,lines=[line(a=(-100,-200),b=(-94,-194))])
    elif transform=='negative_slope':g=groups(s,lines=[line(a=(0,0),b=(6,-6))])
    else:g=groups(s)
    assert sorted(fraction(f) for f in children(partition(s,g)))==[14,14]


def test_u_shape_disconnected_portions_share_one_region_without_bridge_area():
    s=geometry();s['polygon_rings']=[ring('U形',[(0,0),(6,0),(6,6),(4,6),(4,2),(2,2),(2,6),(0,6)])]
    # Above y=.1x+3 contains separate arms in the same convex clipping domain.
    r=partition(s,groups(s,lines=[line(a=(0,3),b=(10,4))]))
    assert sorted(fraction(f) for f in children(r))==[Fraction('10.8'),Fraction('17.2')]


def test_four_crossing_or_parallel_lines_and_hole_only_regions():
    s=geometry();parallel=[line(a=(0,i),b=(6,6+i)) for i in (-1,0,1,2)]
    r=partition(s,groups(s,lines=parallel));assert len(children(r))==5 and sum(fraction(f) for f in children(r))==28
    # Four sides of a diamond wholly within the hole create a zero-net domain.
    diamond=[line(a=(0,'2.5'),b=('2.5',0)),line(a=(0,'5.5'),b=('5.5',0)),line(a=(0,'-.5'),b=(6,'5.5')),line(a=(0,'.5'),b=(6,'6.5'))]
    r=partition(s,groups(s,lines=diamond));assert r['surface_cap_partition_checks'][0]['zero_window_count']>0
    assert sum(fraction(f) for f in children(r))==28 and all(fraction(f)>0 for f in children(r))


def test_caps_sides_and_legacy_versions_coexist_without_rewriting_outputs():
    s=geometry();side=side_cuts(s);side[0]['u_cuts'][0]['value']='1';side[0]['u_cuts'][0]['unit']='m';side[0]['v_cuts'][0]['value']='100';g=groups(s)+side
    r=partition(s,g);assert r.get('surface_partition_checks') and r.get('surface_partition_lines')
    from backend.polygon_cap_partitions import partition as legacy
    assert partition(s,[])==boundary(s)
    for old in (windows(s),side):
        assert partition(s,old)==legacy(s,old,copy.deepcopy(boundary(s)))
    assert calculate(surface(s),s)['rule_version']=='polygon-boundary-1.0'
    assert calculate(payload(s,g=windows(s)),s)['rule_version']=='polygon-boundary-1.2'


@pytest.mark.parametrize('case',['same','horizontal','vertical','outside','touch','duplicate','duplicate_reversed','empty','five','nonlist','missing_coordinate','extra_field','bad_unit','bad_number','precision','coordinate_proof','basis','basis_proof','side','unknown','duplicate_group','notdict'])
def test_invalid_lines_fail_closed(case):
    s=geometry();g=groups(s);c=g[0]['line_cuts'][0]
    if case=='same':c.update(x2='0',y2='0')
    elif case=='horizontal':c['y2']='0'
    elif case=='vertical':c['x2']='0'
    elif case=='outside':c.update(y1='-10',y2='-4')
    elif case=='touch':c.update(y1='-6',y2='0')
    elif case=='duplicate':g[0]['line_cuts'].append(line(a=(1,1),b=(5,5)))
    elif case=='duplicate_reversed':g[0]['line_cuts'].append(line(a=(6,6),b=(0,0)))
    elif case=='empty':g[0]['line_cuts']=[]
    elif case=='five':g[0]['line_cuts']*=5
    elif case=='nonlist':g[0]['line_cuts']='bad'
    elif case=='missing_coordinate':c.pop('x1')
    elif case=='extra_field':c['area']='100'
    elif case=='bad_unit':c['unit']='m²'
    elif case=='bad_number':c['x1']='NaN'
    elif case=='precision':c['x1']='.0000000000000000001'
    elif case=='coordinate_proof':c['evidence']['x1']=' '
    elif case=='basis':c['basis']=' '
    elif case=='basis_proof':c['evidence_id']=' '
    elif case=='side':g[0]['face_id']=boundary(s)['faces'][2]['face_id']
    elif case=='unknown':g[0]['face_id']='old'
    elif case=='duplicate_group':g*=2
    else:g=[False]
    with pytest.raises(DomainError):partition(s,g)


def test_read_only_candidate_preview_save_review_book_and_boq(clients):
    e,r,pid,s,b,eid=setup(clients);before=get(e,pid)
    candidate=post(e,pid,'/surface-face-candidates/partition',b).json();assert len(candidate['surface_partition_lines'])==1
    view=post(e,pid,'/components/preview',b).json();assert view['quantity']=='28' and get(e,pid)==before
    saved=post(e,pid,'/components',{**b,'source_geometry':{'fake':True}}).json();assert saved['preview']=={k:v for k,v in view.items() if k!='source_stamp'}
    post(e,pid,'/takeoffs/calculate',expected=409);approve(r,pid,'components',saved)
    for t in post(e,pid,'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
    assert any(v['unit']=='m²' and Decimal(v['quantity'])==28 for v in post(e,pid,'/boq/generate').json())


@pytest.mark.parametrize('fault',['coordinate','proof','bytes','source','line_record','domain','exact_domain','net_fraction','closure'])
def test_line_geometry_and_proof_tamper_rejected_on_review(clients,fault):
    e,r,pid,s,b,eid=setup(clients);saved=post(e,pid,'/components',b).json()
    if fault=='bytes':(store.ROOT/'originals'/get(e,pid)['documents'][-1]['hash']).write_bytes(b'changed')
    else:
        def damage(p):
            c=p['components'][-1]
            if fault=='coordinate':c['face_subdivisions'][0]['line_cuts'][0]['y1']='1'
            elif fault=='proof':next(ev for ev in p['evidence'] if ev['id']==eid)['text']+=' changed'
            elif fault=='source':p['components'][0]['polygon_rings'][0]['points'][1]['parameters']['x']='7'
            elif fault=='line_record':c['preview']['surface_partition_lines'][0]['coefficients'][0]='999'
            elif fault=='closure':c['preview']['surface_cap_partition_checks'][0]['numerator']='999'
            else:
                f=next(f for f in c['preview']['surface_faces'] if f.get('cap_clip_domain'))
                if fault=='domain':f['cap_clip_domain'][0]['x']='999'
                elif fault=='exact_domain':f['cap_clip_domain'][0]['exact_x']['numerator']='999'
                else:f['cap_rational_area']['numerator']='999'
        mutate(pid,damage)
    current=get(e,pid)['components'][-1];response=r.post(f'/api/v1/projects/{pid}/reviews',json=dict(collection='components',object_id=saved['id'],object_hash=store.digest(current),decision='approve',reason='合成独立复核'))
    assert response.status_code in (409,422),response.text


def test_line_revision_and_removal_preserve_history_require_new_decisions_and_rebuild(clients):
    e,r,pid,s,b,eid=setup(clients);saved=post(e,pid,'/components',b).json();approve(r,pid,'components',saved)
    for t in post(e,pid,'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
    post(e,pid,'/boq/generate');old=get(e,pid)['components'][-1]
    new=payload(s,eid,groups(s,eid,lines=[line(eid,a=(0,1),b=(6,7))]));new['object_hash']=store.digest(old)
    url=f'/api/v1/projects/{pid}/components/'+saved['id']
    bad=copy.deepcopy(new);bad['face_decisions']=b['face_decisions'];assert e.put(url,json=bad).status_code==422
    res=e.put(url,json=new);assert res.status_code==200,res.text;assert Decimal(res.json()['preview']['quantity'])==33
    p=get(e,pid);assert p['component_history'][-1]==old and all(t['state']=='stale' for t in p['takeoffs'])
    assert any(v['unit']=='m²' and Decimal(v['quantity'])==28 for v in p['boq'])
    assert e.put(url,json=new).status_code==409;approve(r,pid,'components',res.json())
    for t in post(e,pid,'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
    assert any(v['unit']=='m²' and Decimal(v['quantity'])==33 for v in post(e,pid,'/boq/generate').json())
    removed=surface(s,eid,included=lambda f:f['normal']=='z+');removed.update(object_hash=store.digest(get(e,pid)['components'][-1]),face_subdivisions=[])
    response=e.put(url,json=removed);assert response.status_code==200,response.text
    assert response.json()['preview']['rule_version']=='polygon-boundary-1.0' and Decimal(response.json()['preview']['quantity'])==56


def test_reviewed_book_line_record_cannot_be_forged(clients):
    e,r,pid,s,b,eid=setup(clients);saved=post(e,pid,'/components',b).json();approve(r,pid,'components',saved)
    book=post(e,pid,'/takeoffs/calculate').json()[-1]
    mutate(pid,lambda p:p['takeoffs'][-1]['surface_partition_lines'][0]['normalized_coordinates'].update(x1='999'))
    approve(r,pid,'takeoffs',book,expected=409)


def test_line_endpoint_permissions_evidence_and_source_invalidation(clients):
    e,r,pid,s,b,eid=setup(clients)
    post(r,pid,'/surface-face-candidates/partition',b,403);post(clients['other'],pid,'/surface-face-candidates/partition',b,404)
    post(TestClient(app),pid,'/surface-face-candidates/partition',b,401)
    second=create(e)['id'];post(e,second,'/surface-face-candidates/partition',b,404)
    bad=copy.deepcopy(b);bad['face_subdivisions'][0]['line_cuts'][0]['evidence']['x1']='missing';post(e,pid,'/surface-face-candidates/partition',bad,404)
    old=get(e,pid)['components'][0];res=e.put(f'/api/v1/projects/{pid}/components/'+s['id'],json={**copy.deepcopy(old),'object_hash':store.digest(old),'thickness':'300'});assert res.status_code==200,res.text
    approve(r,pid,'components',res.json());post(e,pid,'/surface-face-candidates/partition',b,409)
    from tests.test_project_access import restrict
    restrict(clients['admin'],pid,[]);post(e,pid,'/surface-face-candidates/partition',b,404)


@pytest.mark.parametrize('fmt',['json','xlsx','pdf'])
def test_three_exports_and_old_quote_do_not_mix_new_line_data(clients,fmt):
    e,r,pid,_,quote=priced_project(clients);e,r,pid,s,b,eid=setup(clients,pid)
    c=post(e,pid,'/components',b).json();approve(r,pid,'components',c)
    for t in post(e,pid,'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
    url=f'/api/v1/projects/{pid}/exports';res=e.get(url,params={'format':fmt,'unpriced':True});assert res.status_code==200,res.text
    titles=['净面斜线分割依据','净面斜线裁剪范围']
    if fmt=='json':
        assert all(any(t[0]==title for t in res.json()['tables']) for title in titles)
        assert res.json()['unpriced']['takeoffs'][-1]['surface_faces'][1]['cap_clip_domain'][0]['exact_x']['denominator']=='1'
    elif fmt=='xlsx':assert set(titles)<=set(load_workbook(io.BytesIO(res.content)).sheetnames)
    else:assert all(title in ''.join(p.extract_text() for p in PdfReader(io.BytesIO(res.content)).pages) for title in titles)
    old=e.get(url,params={'format':'json','quote_id':quote['id']}).json()
    assert old['quote']==quote and not any(t[0] in titles for t in old['tables'])


def test_microgeometry_at_large_origin_and_low_calling_precision():
    s=geometry()
    with localcontext() as ctx:
        ctx.prec=80;o=Decimal('1000000000');d=Decimal('.00000001')
        s['polygon_rings']=[ring('微小方形',[(o,o),(o+2*d,o),(o+2*d,o+2*d),(o,o+2*d)])]
        g=groups(s,lines=[line(a=(o,o),b=(o+2*d,o+2*d))])
    with localcontext() as ctx:ctx.prec=10;r=partition(s,g)
    assert all(fraction(f)==Fraction(2,10**16) for f in children(r))


def test_combined_budget_blocks_before_silently_truncating_faces():
    s=geometry();s['polygon_rings']=[ring('U形',[(0,0),(6,0),(6,6),(4,6),(4,2),(2,2),(2,6),(0,6)]),ring('孔',[(.5,.5),(1.5,.5),(1.5,1.5),(.5,1.5)],role='hole')]
    base=boundary(s);g=[]
    for f in base['faces'][2:13]:
        g.append(dict(face_id=f['face_id'],u_cuts=[dict(value=str(Decimal(f['edge_length'])*i/10),unit='m',evidence_id='test') for i in range(1,9)],
                      v_cuts=[dict(value=str(Decimal('.2')*i/10),unit='m',evidence_id='test') for i in range(1,9)]))
    cap=groups(s,xy=True,lines=[line(a=(0,i),b=(6,6+i)) for i in map(Decimal,('-.2','0','.2','.4'))])[0]
    cap['u_cuts']=[dict(value=str(Decimal(i)/10),unit='m',evidence_id='test') for i in range(1,9)]
    cap['v_cuts']=copy.deepcopy(cap['u_cuts']);g.append(cap)
    last=base['faces'][13]
    g.append(dict(face_id=last['face_id'],u_cuts=[dict(value=str(Decimal(last['edge_length'])*i/4),unit='m',evidence_id='test') for i in (1,2,3)],v_cuts=[dict(value=str(Decimal('.05')*i),unit='m',evidence_id='test') for i in (1,2)]))
    with pytest.raises(DomainError,match='1000'):partition(s,g)


def test_original_numeric_inputs_are_accepted_and_export_records_are_strings():
    s=geometry();c=line();c.update(x1=0,y1=0,x2=6,y2=6)
    r=partition(s,groups(s,lines=[c]));assert r['surface_partition_lines'][0]['original']['x1']=='0'
