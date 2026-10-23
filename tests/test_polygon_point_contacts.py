"""Synthetic exact point-contact topology and review gates; no professional signoff."""
import copy
import io
import random
from decimal import Decimal, localcontext
from fractions import Fraction
from functools import cmp_to_key
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pypdf import PdfReader
from backend import store
from backend.main import app
from backend.engine import DomainError
from backend.surface_takeoff import boundary, partition, calculate
from backend.polygon_point_contacts import angle_compare, trace_loops
from backend.polygon_connected_partitions import exact_point
from backend.polygon_takeoff import inside
from tests.test_polygon_takeoff import ring
from tests.test_polygon_connected_partitions import geometry as connected_geometry, groups as connected_groups, pieces, frac
from tests.test_polygon_line_partitions import line, groups as line_groups
from tests.test_polygon_cap_partitions import post, mutate
from tests.test_surface_takeoff import payload as surface
from tests.test_route_takeoff import evaluate
from tests.test_workflow import create, get, upload, approve, priced_project


def geometry(eid='test'):
    s=connected_geometry(eid,holes=False)
    s['polygon_rings']=[ring('单点会合双翼外轮廓',[(0,0),(6,0),(6,4),(3,1),(0,4)],eid)]
    s['measurement_basis']='合成外15m²、无孔，厚200mm、重复2。Y=1m以上两翼仅在(3,1)接点相会，各4.5m²，按内部连通分开；非专业依据'
    return s


def groups(s,eid='test',y='1',top=True):
    g=connected_groups(s,eid,y=y,top=top)
    g[0]['point_contact']=dict(mode='interior',basis='合成按净区内部连通处理，零宽接点不合并两块面积，不证明结构连接或做法连续',evidence_id=eid)
    return g


def payload(s,eid='test',g=None):
    g=groups(s,eid) if g is None else g;v=partition(s,g);b=surface(s,eid,included=lambda f:False)
    b.update(name='合成单点会合左翼防水（非真实工程）',purpose='waterproof',face_subdivisions=g,
             measurement_basis='合成顶部Y=1m以上两翼仅一点会合，各4.5m²，按内部连通分开，仅左翼计入，重复2=9m²；其余明确排除',
             face_decisions=[dict(face_id=f['face_id'],included=bool(f.get('cap_component_rings')) and Decimal(f['v_from'])>=1 and f['cap_component_index']==1,
                                  basis='合成逐块核定，零宽接点不证明实际施工接触',evidence_id=eid) for f in v['faces']])
    return b


def setup(clients,pid=None):
    e,r=clients['estimator'],clients['reviewer'];pid=pid or create(e)['id']
    upload(e,pid,'合成软件验收（非真实工程）：共同水平设计XY。外轮廓(0,0)(6,0)(6,4)(3,1)(0,4)m，无孔，面积15m²，厚200mm、重复2。顶部Y=100cm分割，上区两翼在(3,1)m仅一点会合，各三角净4.5m²。明确按净区内部连通口径，零宽接点不合并面积，只计左翼9m²，其余排除，不代表结构连接或做法贯通。合成软件修订Y=200cm时两翼分离，各净2m²，仅左翼计入4m²；须重新核定。','point-contact-basis.txt')
    eid=get(e,pid)['evidence'][-1]['id'];s=geometry(eid);s.pop('id');s.pop('revision');s['name']='合成单点会合双翼等厚实体（非真实工程）'
    saved=post(e,pid,'/components',s).json()
    for c in get(e,pid)['components']:approve(r,pid,'components',c)
    s=next(c for c in get(e,pid)['components'] if c['id']==saved['id']);return e,r,pid,s,payload(s,eid),eid


@pytest.mark.parametrize('top',[False,True])
def test_point_only_connection_keeps_positive_width_interiors_as_separate_pieces(top):
    s=geometry();g=groups(s,top=top);before=copy.deepcopy((s,g));v=partition(s,g);fs=pieces(v)
    assert (s,g)==before and [frac(f) for f in fs]==[6,Fraction('4.5'),Fraction('4.5')] and len(v['faces'])==9
    checks=v['surface_point_contact_checks'];assert [c['contact_count'] for c in checks]==[0,1] and checks[1]['component_count']==2
    contact=checks[1]['contacts'][0];assert exact_point(contact['point'])==(3,1) and contact['degree']==2 and len(contact['pairings'])==2
    assert all(f['cap_contact_markers'] for f in fs[1:]) and fs[0]['cap_contact_markers']==[]
    assert all(inside(exact_point(f['cap_label_point']),[exact_point(p) for p in f['cap_component_rings'][0]['points']])==1 for f in fs)
    out=calculate(payload(s,g=g),s);assert Decimal(out['quantity'])==9 and out['rule_version']=='polygon-boundary-1.5'
    assert out['source_evidence']['cap_contact.1.basis']=='test'
    with localcontext() as ctx:
        ctx.prec=80
        assert evaluate(out['formula_ast'],out['dimensions'])==9
        for f in fs:assert evaluate(f['formula_ast'],v['dimensions'])==Decimal(f['area'])


def test_tangent_hole_creates_weak_outer_walk_keeps_void_and_exact_original_deductions():
    s=geometry();s['polygon_rings']=[ring('方形',[(0,0),(6,0),(6,6),(0,6)]),ring('菱形孔',[(3,1),(4,2),(3,3),(2,2)],role='hole')]
    v=partition(s,groups(s));assert [frac(f) for f in pieces(v)]==[6,28]
    c=v['surface_point_contact_checks'][1];assert c['contact_count']==1 and c['repeated_vertex_walks']==1
    upper=pieces(v)[1];ring_points=[exact_point(p) for p in upper['cap_component_rings'][0]['points']]
    assert len(ring_points)==9 and ring_points.count((3,1))==2
    assert inside((3,2),ring_points)==-1 and inside((1,2),ring_points)==1
    assert Decimal(v['surface_cap_clip_sources'][1]['cap_clip_rings'][1]['area'])==2
    assert len(upper['cap_component_rings'])==1 and frac(upper)==28


@pytest.mark.parametrize('change',['reverse_ring','rotate_start','reverse_cut','cut_units','translate','swap_axes','shear'])
def test_contact_topology_transforms_and_equivalent_input_identity(change):
    s=geometry();g=groups(s)
    if change=='reverse_ring':s['polygon_rings'][0]['points'].reverse()
    elif change=='rotate_start':s['polygon_rings'][0]['points']=s['polygon_rings'][0]['points'][2:]+s['polygon_rings'][0]['points'][:2]
    elif change=='cut_units':
        original=partition(s,g);g[0]['v_cuts'][0].update(value='1000',unit='mm');assert [f['face_id'] for f in pieces(original)]==[f['face_id'] for f in pieces(partition(s,g))]
    if change in ('translate','swap_axes','shear'):
        for p in s['polygon_rings'][0]['points']:
            x,y=(Decimal(p['parameters'][k]) for k in ('x','y'))
            if change=='translate':x-=100;y-=200
            elif change=='swap_axes':x,y=y,x
            else:y+=2*x
            p['parameters']={'x':str(x),'y':str(y)}
        if change=='translate':g=groups(s,y='-199')
        elif change=='swap_axes':g=groups(s);g[0]['u_cuts']=g[0]['v_cuts'];g[0]['v_cuts']=[]
        else:g=line_groups(s,lines=[line(a=(0,1),b=(6,13))]);g[0].update(separate_regions=True,point_contact=groups(s)[0]['point_contact'])
    if change=='reverse_cut':
        g=line_groups(s,lines=[line(a=(0,1),b=(6,7))]);g[0].update(separate_regions=True,point_contact=groups(s)[0]['point_contact'])
        a=partition(s,g);g[0]['line_cuts']=[line(a=(6,7),b=(0,1))];b=partition(s,g)
        assert [f['face_id'] for f in pieces(a)]==[f['face_id'] for f in pieces(b)];return
    if change in ('reverse_ring','rotate_start'):g=groups(s)
    v=partition(s,g);assert sorted(frac(f) for f in pieces(v))==[Fraction('4.5'),Fraction('4.5'),6]


def test_twenty_random_sawtooth_lobes_match_independent_triangle_areas_and_contacts():
    rng=random.Random(6401)
    for _ in range(20):
        n=rng.randint(2,8);hs=[rng.randint(2,8) for _ in range(n)];w=2*n;coords=[(0,0),(w,0),(w,hs[-1])]
        for i in range(n-1,0,-1):coords.extend([(2*i,1),(0 if i==1 else 2*i-1,hs[i-1])])
        s=geometry();s['polygon_rings']=[ring('锯齿',coords)];v=partition(s,groups(s))
        assert sorted(frac(f) for f in pieces(v))==sorted([Fraction(w)]+[Fraction(h-1) for h in hs])
        assert sum(c['contact_count'] for c in v['surface_point_contact_checks'])==n-1
        assert all(c['exact_rational_tiling'] for c in v['surface_cap_component_checks'])


def test_multiple_tangent_holes_and_remaining_closed_hole_are_not_double_deducted():
    s=geometry();s['polygon_rings']=[ring('outer',[(0,0),(10,0),(10,6),(0,6)]),
        ring('tangent1',[(2,1),(3,2),(2,3),(1,2)],role='hole'),ring('tangent2',[(6,1),(7,2),(6,3),(5,2)],role='hole'),
        ring('closed',[(8,3),(9,3),(9,4),(8,4)],role='hole')]
    v=partition(s,groups(s));assert sorted(frac(f) for f in pieces(v))==[10,45]
    upper=pieces(v)[1];assert len(upper['cap_component_rings'])==2 and upper['cap_component_rings'][1]['role']=='hole'
    assert v['surface_point_contact_checks'][1]['contact_count']==2 and v['surface_point_contact_checks'][1]['repeated_vertex_walks']==1


def test_default_rejection_and_all_older_versions_are_preserved():
    s=geometry();g=groups(s);g[0].pop('point_contact')
    with pytest.raises(DomainError,match='点接触|分支'):partition(s,g)
    s=connected_geometry();g=connected_groups(s)
    from backend.polygon_connected_partitions import partition as previous
    old=previous(s,g,copy.deepcopy(boundary(s)));assert partition(s,g)==old and not old.get('surface_point_contact_checks')
    b=surface(s);assert calculate(b,s)['rule_version']=='polygon-boundary-1.0'
    new=copy.deepcopy(g);new[0]['point_contact']=groups(s)[0]['point_contact'];v=partition(s,new)
    assert all(c['contact_count']==0 for c in v['surface_point_contact_checks']) and sorted(frac(f) for f in pieces(v))==[5,5,16]


@pytest.mark.parametrize('case',['false','string','wrong_mode','missing_basis','blank_basis','missing_proof','extra','no_separate','side','source_touch'])
def test_point_contact_choice_requires_explicit_scope_and_evidence(case):
    s=geometry();g=groups(s);c=g[0]['point_contact']
    if case=='false':g[0]['point_contact']=False
    elif case=='string':g[0]['point_contact']='interior'
    elif case=='wrong_mode':c['mode']='snap'
    elif case=='missing_basis':c.pop('basis')
    elif case=='blank_basis':c['basis']=' '
    elif case=='missing_proof':c['evidence_id']=''
    elif case=='extra':c['epsilon']='.01'
    elif case=='no_separate':g[0].pop('separate_regions')
    elif case=='side':g[0]['face_id']=boundary(s)['faces'][2]['face_id']
    else:s['polygon_rings'].append(ring('源孔触外边',[(0,1),(1,1),(1,2),(0,2)],role='hole'))
    with pytest.raises(DomainError):partition(s,g)


def test_angle_order_exact_axes_and_clockwise_wrap_with_huge_and_tiny_rays():
    rays=[(1,0),(1,1),(0,1),(-1,1),(-1,0),(-1,-1),(0,-1),(1,-1)]
    assert sorted(reversed(rays),key=cmp_to_key(angle_compare))==rays
    a=(Fraction(10**100),Fraction(1,10**100));b=(Fraction(10**100),Fraction(2,10**100))
    assert angle_compare(a,b)<0 and angle_compare(b,a)>0 and angle_compare((1,1),(2,2))==0


def test_invalid_embedding_cannot_be_resolved_by_arbitrary_pairing():
    from collections import defaultdict
    a,b,c=(Fraction(0),Fraction(0)),(Fraction(1),Fraction(0)),(Fraction(2),Fraction(0))
    remaining={(a,b),(b,a)};out=defaultdict(list,{a:[b],b:[a]});inc=defaultdict(list,{a:[b],b:[a]})
    with pytest.raises(DomainError,match='折返'):trace_loops(remaining,out,inc,[])
    out=defaultdict(list,{a:[b,c]});inc=defaultdict(list,{a:[b]})
    with pytest.raises(DomainError,match='平衡'):trace_loops({(a,b)},out,inc,[])


def test_read_only_rejection_preview_save_independent_review_and_boq(clients):
    e,r,pid,s,b,eid=setup(clients);before=get(e,pid);old=copy.deepcopy(b);old['face_subdivisions'][0].pop('point_contact')
    post(e,pid,'/surface-face-candidates/partition',old,422)
    candidate=post(e,pid,'/surface-face-candidates/partition',b).json();assert candidate['surface_point_contact_checks'][1]['contact_count']==1
    preview=post(e,pid,'/components/preview',b).json();assert Decimal(preview['quantity'])==9 and get(e,pid)==before
    c=post(e,pid,'/components',b).json();post(e,pid,'/takeoffs/calculate',expected=409);approve(r,pid,'components',c)
    for t in post(e,pid,'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
    assert any(v['unit']=='m²' and Decimal(v['quantity'])==9 for v in post(e,pid,'/boq/generate').json())


@pytest.mark.parametrize('fault',['mode','basis','evidence','bytes','coordinate','pair','count','marker','weak_boundary','source'])
def test_contact_records_and_proofs_are_rebuilt_before_review(clients,fault):
    e,r,pid,s,b,eid=setup(clients);comp=post(e,pid,'/components',b).json()
    if fault=='bytes':(store.ROOT/'originals'/get(e,pid)['documents'][-1]['hash']).write_bytes(b'changed')
    else:
        def damage(p):
            c=p['components'][-1];check=c['preview']['surface_point_contact_checks'][1]
            if fault=='mode':c['face_subdivisions'][0]['point_contact']['mode']='snap'
            elif fault=='basis':c['face_subdivisions'][0]['point_contact']['basis']+=' changed'
            elif fault=='evidence':p['evidence'][0]['text']+=' changed'
            elif fault=='coordinate':check['contacts'][0]['point']['exact_x']['numerator']='999'
            elif fault=='pair':check['contacts'][0]['pairings'][0]['outgoing_to']['x']='999'
            elif fault=='count':check['contact_count']=0
            elif fault=='marker':next(f for f in c['preview']['surface_faces'] if f.get('cap_contact_markers'))['cap_contact_markers'][0]['x']='999'
            elif fault=='weak_boundary':check['repeated_vertex_walks']=999
            else:p['components'][0]['polygon_rings'][0]['points'][1]['parameters']['x']='7'
        mutate(pid,damage)
    current=get(e,pid)['components'][-1];res=r.post(f'/api/v1/projects/{pid}/reviews',json=dict(collection='components',object_id=comp['id'],object_hash=store.digest(current),decision='approve',reason='合成核查'))
    assert res.status_code in (409,422),res.text


def test_calculation_book_contact_proof_cannot_be_forged(clients):
    e,r,pid,s,b,eid=setup(clients);c=post(e,pid,'/components',b).json();approve(r,pid,'components',c)
    book=next(t for t in post(e,pid,'/takeoffs/calculate').json() if t['component_id']==c['id'])
    mutate(pid,lambda p:next(t for t in p['takeoffs'] if t['id']==book['id'])['surface_point_contact_checks'][1]['setting'].update(basis='forged'))
    approve(r,pid,'takeoffs',book,expected=409)


def test_revision_can_remove_geometric_contact_but_keeps_explicit_rule_and_history(clients):
    e,r,pid,s,b,eid=setup(clients);c=post(e,pid,'/components',b).json();approve(r,pid,'components',c)
    for t in post(e,pid,'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
    post(e,pid,'/boq/generate');old=get(e,pid)['components'][-1]
    new=payload(s,eid,groups(s,eid,y='2'));new['object_hash']=store.digest(old);url=f'/api/v1/projects/{pid}/components/'+c['id']
    bad={**new,'face_decisions':b['face_decisions']};assert e.put(url,json=bad).status_code==422
    res=e.put(url,json=new);assert res.status_code==200,res.text
    assert Decimal(res.json()['preview']['quantity'])==4 and res.json()['preview']['rule_version']=='polygon-boundary-1.5'
    assert all(c['contact_count']==0 for c in res.json()['preview']['surface_point_contact_checks'])
    p=get(e,pid);assert p['component_history'][-1]==old and all(t['state']=='stale' for t in p['takeoffs'])
    assert any(v['unit']=='m²' and Decimal(v['quantity'])==9 for v in p['boq']) and e.put(url,json=new).status_code==409
    approve(r,pid,'components',res.json())
    for t in post(e,pid,'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
    assert any(v['unit']=='m²' and Decimal(v['quantity'])==4 for v in post(e,pid,'/boq/generate').json())
    # Once geometry has no contact, explicitly removing the new choice restores1.4.
    g=groups(s,eid,y='2');g[0].pop('point_contact');normal=payload(s,eid,g);assert calculate(normal,s)['rule_version']=='polygon-boundary-1.4'


def test_permissions_foreign_evidence_and_source_revision_are_enforced(clients):
    e,r,pid,s,b,eid=setup(clients)
    post(r,pid,'/surface-face-candidates/partition',b,403);post(clients['other'],pid,'/surface-face-candidates/partition',b,404);post(TestClient(app),pid,'/surface-face-candidates/partition',b,401)
    second=create(e)['id'];post(e,second,'/surface-face-candidates/partition',b,404)
    upload(e,second,'另一个合成项目的口径证明');foreign=get(e,second)['evidence'][-1]['id'];bad=copy.deepcopy(b);bad['face_subdivisions'][0]['point_contact']['evidence_id']=foreign;post(e,pid,'/surface-face-candidates/partition',bad,404)
    old=get(e,pid)['components'][0];res=e.put(f'/api/v1/projects/{pid}/components/'+s['id'],json={**copy.deepcopy(old),'object_hash':store.digest(old),'thickness':'300'});assert res.status_code==200,res.text
    approve(r,pid,'components',res.json());post(e,pid,'/surface-face-candidates/partition',b,409)
    from tests.test_project_access import restrict
    restrict(clients['admin'],pid,[]);post(e,pid,'/surface-face-candidates/partition',b,404)


@pytest.mark.parametrize('fmt',['json','xlsx','pdf'])
def test_contact_tables_and_exact_proofs_export_without_changing_old_quote(clients,fmt):
    e,r,pid,_,quote=priced_project(clients);e,r,pid,s,b,eid=setup(clients,pid);c=post(e,pid,'/components',b).json();approve(r,pid,'components',c)
    for t in post(e,pid,'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
    url=f'/api/v1/projects/{pid}/exports';res=e.get(url,params={'format':fmt,'unpriced':True});assert res.status_code==200,res.text
    titles=['点接触处理依据','接点边界配对记录']
    if fmt=='json':assert all(any(t[0]==title for t in res.json()['tables']) for title in titles) and any(t.get('surface_point_contact_checks') for t in res.json()['unpriced']['takeoffs'])
    elif fmt=='xlsx':assert set(titles)<=set(load_workbook(io.BytesIO(res.content)).sheetnames)
    else:assert all(title in ''.join(p.extract_text() for p in PdfReader(io.BytesIO(res.content)).pages) for title in titles)
    old=e.get(url,params={'format':'json','quote_id':quote['id']}).json();assert old['quote']==quote and not any(t[0] in titles for t in old['tables'])


def test_tiny_point_contact_at_large_origin_uses_exact_order_and_positive_areas():
    s=geometry()
    with localcontext() as ctx:
        ctx.prec=80;o=Decimal('1000000000');scale=Decimal('.00000001')
        for p in s['polygon_rings'][0]['points']:p['parameters']={k:str(o+Decimal(v)*scale) for k,v in p['parameters'].items()}
        g=groups(s,y=str(o+scale))
    with localcontext() as ctx:ctx.prec=10;v=partition(s,g)
    assert sorted(frac(f) for f in pieces(v))==[Fraction(45,10**17),Fraction(45,10**17),Fraction(6,10**16)]
    assert sum(c['contact_count'] for c in v['surface_point_contact_checks'])==1


def test_tangent_hole_contact_with_repeating_fraction_intersections_and_ast():
    s=geometry();s['polygon_rings']=[ring('方形',[(0,0),(10,0),(10,10),(0,10)]),ring('三角孔',[(7,4),(8,6),(6,6)],role='hole')]
    g=line_groups(s,lines=[line(a=(0,1),b=(7,4))]);g[0].update(separate_regions=True,point_contact=groups(s)[0]['point_contact'])
    v=partition(s,g);assert sorted(frac(f) for f in pieces(v))==[Fraction(220,7),Fraction(466,7)]
    assert sum(c['contact_count'] for c in v['surface_point_contact_checks'])==1
    assert any(p['exact_y']['denominator']=='7' for f in pieces(v) for r in f['cap_component_rings'] for p in r['points'])
    with localcontext() as ctx:
        ctx.prec=80
        for f in pieces(v):assert evaluate(f['formula_ast'],v['dimensions'])==Decimal(f['area'])


def test_new_contact_mode_respects_existing_complexity_budget(monkeypatch):
    import backend.polygon_connected_partitions as connected
    s=geometry();g=groups(s)
    # Exhaust the shared work budget with a real contact geometry.
    monkeypatch.setattr(connected,'MAX_EDGE_TESTS',50)
    with pytest.raises(DomainError,match='2000000'):partition(s,g)
