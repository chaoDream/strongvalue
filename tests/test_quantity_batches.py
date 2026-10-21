import copy
import io
import json
import zipfile
import uuid
import pytest
from openpyxl import load_workbook
from backend import store
from backend.workflows import allocation_diagnostics
from tests.test_workflow import get, approve, priced_project, upload


def setup(clients):
    e,r,pid,comp,q=priced_project(clients);base=f'/api/v1/projects/{pid}'
    b=get(e,pid)['boq'][0]
    response=e.post(base+'/boq/'+b['id']+'/split',json={'parts':[{'name':'合成底板甲','quantity':'7'},{'name':'合成底板乙','quantity':'8'}]})
    assert response.status_code==200,response.text
    for b in response.json():approve(r,pid,'boq',b)
    p=get(e,pid);eid=p['evidence'][0]['id']
    body={'name':'合成批量调量单','reason':'测试合同变更，不代表真实签证','request_id':str(uuid.uuid4()),
          'entries':[{'boq_id':b['id'],'object_hash':store.digest(b),'quantity':str(int(b['quantity'])+1),
                      'basis':'合成逐项尺寸展开依据','evidence_id':eid} for b in p['boq']]}
    return e,r,pid,q,body


def create(e,pid,body):
    res=e.post(f'/api/v1/projects/{pid}/quantity-change-batches',json=body)
    assert res.status_code==200,res.text
    return res.json()['batch']


def review_all(r,e,pid,batch):
    p=get(e,pid)
    for cid in batch['change_ids']:approve(r,pid,'changes',next(c for c in p['changes'] if c['id']==cid))


def test_atomic_apply_shared_source_preview_and_snapshot_exports(clients):
    e,r,pid,old,body=setup(clients);base=f'/api/v1/projects/{pid}'
    before=get(e,pid);batch=create(e,pid,body)
    assert batch['impact']['before_total']=='1622.40'
    assert batch['impact']['after_total']=='1838.73'
    assert batch['impact']['delta']=='216.33'
    assert get(e,pid)['boq']==before['boq']
    review_all(r,e,pid,batch)
    res=e.post(base+'/quantity-change-batches/'+batch['id']+'/apply');assert res.status_code==200,res.text
    p=get(e,pid)
    assert [b['quantity'] for b in p['boq']]==['8','9']
    assert all(b['state']=='draft' for b in p['boq'])
    assert len(p['takeoffs'])==2 and len(p['takeoff_history'])==1
    assert allocation_diagnostics(p)['passed']
    for b in p['boq']:approve(r,pid,'boq',b)
    quote=e.post(base+'/quotes/calculate').json()
    assert quote['total']==batch['impact']['after_total']
    assert quote['snapshot']['quantity_batches'][-1]['state']=='applied'
    assert p['quotes'][0]==old
    workbook=e.get(base+'/exports',params={'quote_id':quote['id'],'format':'xlsx'})
    w=load_workbook(io.BytesIO(workbook.content));assert '批量调量台账' in w.sheetnames
    assert '变更签证台账' in w.sheetnames
    assert e.get(base+'/exports',params={'quote_id':quote['id'],'format':'pdf'}).status_code==200
    snapshot=get(e,pid)
    assert e.post(base+'/quantity-change-batches/'+batch['id']+'/apply').json()['duplicate']
    assert get(e,pid)==snapshot
    assert e.post(base+'/quantity-change-batches',json=body).json()['duplicate']
    assert get(e,pid)['quotes'][-1]==quote
    from backend import delivery
    with store.connect() as c:
        payload,_,_=delivery.export_bundle(c,get(e,pid),None)
    z=zipfile.ZipFile(io.BytesIO(payload))
    raw=json.loads(z.read('unpriced-snapshot.json'))
    assert raw['quantity_batches'][0]['id']==batch['id']
    assert set(raw['applied_changes'])==set(batch['change_ids'])
    assert len(raw['changes'])==2


def test_partial_review_and_child_endpoint_do_not_partially_apply(clients):
    e,r,pid,q,body=setup(clients);base=f'/api/v1/projects/{pid}';batch=create(e,pid,body)
    p=get(e,pid);approve(r,pid,'changes',p['changes'][0])
    before=get(e,pid)
    assert e.post(base+'/changes/'+batch['change_ids'][0]+'/apply').status_code==409
    assert e.post(base+'/quantity-change-batches/'+batch['id']+'/apply').status_code==409
    assert get(e,pid)==before
    assert r.post(base+'/quantity-change-batches/'+batch['id']+'/apply').status_code==403


def test_mid_apply_failure_rolls_back_quantities_history_and_audit(clients,monkeypatch):
    from backend import quantity_batches
    from fastapi import HTTPException
    e,r,pid,q,body=setup(clients);base=f'/api/v1/projects/{pid}';batch=create(e,pid,body);review_all(r,e,pid,batch)
    before=get(e,pid);audit=e.get(base+'/audit').json();original=quantity_batches.apply_change_data;calls=[]
    def failing(p,c):
        calls.append(c['id'])
        if len(calls)==2:raise HTTPException(409,'测试第二项失败')
        return original(p,c)
    monkeypatch.setattr(quantity_batches,'apply_change_data',failing)
    assert e.post(base+'/quantity-change-batches/'+batch['id']+'/apply').status_code==409
    assert len(calls)==2 and get(e,pid)==before
    assert e.get(base+'/audit').json()==audit


@pytest.mark.parametrize('mutation',[
    lambda d:d.update(entries=[]),lambda d:d['entries'].append(copy.deepcopy(d['entries'][0])),
    lambda d:d['entries'][0].update(quantity='-1'),lambda d:d['entries'][0].update(quantity='NaN'),
    lambda d:d['entries'][0].update(basis=''),lambda d:d.update(reason=''),lambda d:d.update(request_id='bad'),
])
def test_invalid_batch_has_no_partial_changes(clients,mutation):
    e,r,pid,q,body=setup(clients);before=get(e,pid);mutation(body)
    res=e.post(f'/api/v1/projects/{pid}/quantity-change-batches',json=body)
    assert res.status_code==422,res.text
    assert get(e,pid)==before


def test_zero_and_unchanged_quantities_are_distinguished(clients):
    e,r,pid,q,body=setup(clients);base=f'/api/v1/projects/{pid}'
    body['entries'][0]['quantity']='7'
    assert e.post(base+'/quantity-change-batches',json=body).status_code==422
    body['entries'][0]['quantity']='0';batch=create(e,pid,body);review_all(r,e,pid,batch)
    assert e.post(base+'/quantity-change-batches/'+batch['id']+'/apply').status_code==200
    assert get(e,pid)['boq'][0]['quantity']=='0'
    assert allocation_diagnostics(get(e,pid))['passed']


def test_duplicate_request_payload_and_cross_tenant_role_boundaries(clients):
    e,r,pid,q,body=setup(clients);base=f'/api/v1/projects/{pid}'
    assert r.post(base+'/quantity-change-batches',json=body).status_code==403
    assert clients['other'].post(base+'/quantity-change-batches',json=body).status_code==404
    batch=create(e,pid,body);before=get(e,pid)
    assert e.post(base+'/quantity-change-batches',json=body).json()['duplicate']
    body['entries'][0]['quantity']='11'
    assert e.post(base+'/quantity-change-batches',json=body).status_code==409
    assert get(e,pid)==before
    assert clients['other'].get(base+'/quantity-change-batches/'+batch['id']).status_code==404
    assert clients['other'].post(base+'/quantity-change-batches/'+batch['id']+'/apply').status_code==404


@pytest.mark.parametrize('change',['rules','date','source','corrupt'])
def test_changed_context_or_source_rejects_entire_application(clients,change):
    e,r,pid,q,body=setup(clients);base=f'/api/v1/projects/{pid}';batch=create(e,pid,body);review_all(r,e,pid,batch)
    if change=='rules':
        rules=get(e,pid)['rules'][0];rules['name']='新测试规则';e.put(base+'/rules',json=rules)
    elif change=='date':e.patch(base+'/settings',json={'pricing_date':'2020-01-01','reason':'测试适用日期变化'})
    elif change=='source':upload(e,pid,'新的合成数量依据')
    else:
        p=get(e,pid);(store.ROOT/'originals'/p['documents'][0]['hash']).write_bytes(b'corrupt')
    before=get(e,pid)
    assert e.post(base+'/quantity-change-batches/'+batch['id']+'/apply').status_code==409
    assert get(e,pid)==before
    assert not before['applied_changes']


def test_stale_line_hash_blocks_submission(clients):
    e,r,pid,q,body=setup(clients);body['entries'][-1]['object_hash']='outdated'
    before=get(e,pid);res=e.post(f'/api/v1/projects/{pid}/quantity-change-batches',json=body)
    assert res.status_code==409 and get(e,pid)==before


def test_cancel_keeps_reviews_but_cannot_be_applied_or_reapproved(clients):
    e,r,pid,q,body=setup(clients);base=f'/api/v1/projects/{pid}';batch=create(e,pid,body);review_all(r,e,pid,batch)
    before=get(e,pid);batch=before['quantity_batches'][0]
    url=base+'/quantity-change-batches/'+batch['id']+'/cancel'
    assert e.post(url,json={'reason':'测试取消','object_hash':'wrong'}).status_code==409
    assert e.post(url,json={'reason':'需补充测量依据，取消本单','object_hash':store.digest(batch)}).status_code==200
    p=get(e,pid);assert p['boq']==before['boq'] and p['reviews']==before['reviews']
    assert all(c['state']=='cancelled' for c in p['changes'])
    assert not any('变更签证' in msg for msg in p['blockers'])
    assert e.post(base+'/quantity-change-batches/'+batch['id']+'/apply').status_code==409
    approve(r,pid,'changes',p['changes'][0],expected=409)
    assert r.get(base+'/review-batch-targets?collection=changes').json()['total']==0
    assert e.post(base+'/quantity-change-batches',json=body).json()['batch']['state']=='cancelled'


def test_unverified_pricing_does_not_fabricate_money_impact(clients):
    e,r,pid,q,body=setup(clients);base=f'/api/v1/projects/{pid}'
    rule=get(e,pid)['rules'][0];rule['name']='未审规则';assert e.put(base+'/rules',json=rule).status_code==200
    batch=create(e,pid,body)
    assert not batch['impact']['available']
    assert 'after_total' not in batch['impact'] and 'delta' not in batch['impact']


def test_read_targets_bind_displayed_quantities_and_hashes(clients):
    e,r,pid,q,body=setup(clients);base=f'/api/v1/projects/{pid}'
    targets=e.get(base+'/quantity-change-targets').json();p=get(e,pid)
    assert len(targets)==2
    for t in targets:
        b=next(b for b in p['boq'] if b['id']==t['id'])
        assert t['quantity']==b['quantity'] and t['object_hash']==store.digest(b)
    assert clients['other'].get(base+'/quantity-change-targets').status_code==404
    batch=create(e,pid,body);target=e.get(base+'/quantity-change-batches/'+batch['id']).json()
    assert target['batch']==batch and target['object_hash']==store.digest(batch)
