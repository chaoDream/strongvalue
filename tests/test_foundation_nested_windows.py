"""Synthetic hierarchy, independent parity/clipping and unchanged evidence gates."""
import copy
import io
import itertools
import random
from decimal import Decimal,localcontext
from fractions import Fraction
import pytest
from openpyxl import load_workbook
from pypdf import PdfReader
from backend import store
from backend.engine import DomainError
from backend.foundation_horizontal_formwork import calculate
from backend.foundation_nested_windows import VERSION,METHOD,validate,separation
from backend.foundation_horizontal_windows import checks
from tests.test_foundation_multi_windows import part,config as mc,body as mb,replay_blocks,polygon_vertices,setup as old_setup
from tests.test_foundation_perforated_windows import excluded
from tests.test_foundation_polygon_windows import clip,quadrature
from tests.test_foundation_horizontal_windows import source,window
from tests.test_foundation_hole_end_formwork import source as end_source
from tests.test_foundation_openings import hole
from tests.test_foundation_voids import rectangle
from tests.test_foundation_recesses import interval
from tests.test_post_foundation import equal_body,fixed_body,fixed_added,equal_added,fixed_adjusted,equal_adjusted,combined
from tests.test_polygon_takeoff import post,mutate
from tests.test_workflow import get,approve,upload,create,priced_project


def box(x,X,y,Y):return [(x,y),(X,y),(X,Y),(x,Y)]


def block(pts,eid='test',name='合成根块',holes=(),parent=None):
    return {**part(pts,eid,name,holes),'parent':parent}


def config(parts=None,eid='test',unit='mm'):
    c=mc(parts if parts is not None else [block(box(-100,100,-100,100),eid,holes=[box(-60,60,-60,60)]),
        block(box(-20,20,-20,20),eid,'合成岛区',parent={'block_index':1,'hole_index':1})],eid,unit)
    c['shape']='nested_polygon';return c


def body(s,eid='test',face='bottom',parts=None):
    b=mb(s,eid,face);b['contact_ranges'][0]['plan_range']=config(parts,eid);return b


@pytest.mark.parametrize('shape',['rectangular','circular','stepped_rectangular','stepped_circular'])
@pytest.mark.parametrize('face',['top','bottom','hole_top','hole_bottom'])
@pytest.mark.parametrize('make',[equal_body,fixed_body,fixed_added,equal_added,fixed_adjusted,equal_adjusted,combined,lambda s:combined(s,equal=True)])
def test_shapes_ends_final_counts_hierarchy_immutable_and_replay(shape,face,make):
    s=end_source(shape,make) if face.startswith('hole') else source(shape,make);b=body(s,face=face);before=copy.deepcopy((s,b));r=calculate(b,s);d=r['contact_ranges'][0]
    assert before==(s,b) and r['rule_version']==VERSION and Decimal(d['window_area'])==Decimal('.0272')
    assert d['window_nesting']==[dict(block_index=1,parent=None,depth=0),dict(block_index=2,parent=dict(block_index=1,hole_index=1),depth=1)]
    assert all(q['nesting_depth']==q['block_index']-1 for q in d['polygon_net_trace']['partitions']);replay_blocks(r)


@pytest.mark.parametrize('shape,face',[(s,f) for s in ['rectangular','circular','stepped_rectangular','stepped_circular'] for f in ['top','bottom','hole_bottom','hole_top']]+[('stepped_rectangular','step'),('stepped_circular','step')])
def test_parent_exclusion_plus_independent_island_matches_combined_without_double_count(shape,face):
    from backend.foundation_multi_windows import part_config
    s=end_source(shape) if face.startswith('hole') else source(shape)
    parts=None
    if face=='step':parts=[block(box(140,270,-100,100),holes=[box(160,230,-60,60)]),block(box(180,210,-20,20),name='岛区',parent=dict(block_index=1,hole_index=1))]
    b=body(s,face=face,parts=parts);r=calculate(b,s);c=b['contact_ranges'][0]['plan_range'];children=[]
    for p in c['parts']:
        child=copy.deepcopy(b);child['contact_ranges'][0]['plan_range']=part_config(c,p);children.append(calculate(child,s)['contact_ranges'][0])
    with localcontext() as ctx:
        ctx.prec=80
        for k in ['single_gross_area','single_opening_area','single_net_area','total_net_area']:
            assert abs(Decimal(r['contact_ranges'][0][k])-sum((Decimal(d[k]) for d in children),Decimal(0)))<Decimal('1e-73')
    replay_blocks(r)


def chain():
    return [block(box(-120,120,-120,120),name='根块',holes=[box(-100,100,-100,100)]),
      block(box(-80,80,-80,80),name='一级岛',holes=[box(-60,60,-60,60)],parent=dict(block_index=1,hole_index=1)),
      block(box(-40,40,-40,40),name='二级岛',holes=[box(-20,20,-20,20)],parent=dict(block_index=2,hole_index=1)),
      block(box(-10,10,-10,10),name='三级岛',parent=dict(block_index=3,hole_index=1))]


def test_four_levels_and_all_24_input_orders_with_explicit_reference_remap():
    s=source();parts=chain();r=calculate(body(s,parts=parts),s);assert [p['depth'] for p in r['contact_ranges'][0]['window_nesting']]==[0,1,2,3]
    original=r['contact_ranges'][0]
    for order in itertools.permutations(range(4)):
        mapping={old+1:new+1 for new,old in enumerate(order)};changed=copy.deepcopy([parts[i] for i in order])
        for p in changed:
            if p['parent'] is not None:p['parent']['block_index']=mapping[p['parent']['block_index']]
        current=calculate(body(s,parts=changed),s);assert current['quantity']==r['quantity']
        with localcontext() as ctx:
            ctx.prec=80;assert abs(Decimal(current['contact_ranges'][0]['single_net_area'])-Decimal(original['single_net_area']))<Decimal('1e-75')
    replay_blocks(r)


def test_sibling_islands_and_separate_root_forest():
    parts=[block(box(-120,120,-120,120),holes=[box(-100,100,-100,100)]),
      block(box(-80,-20,-60,60),name='左岛',parent=dict(block_index=1,hole_index=1)),
      block(box(20,80,-60,60),name='右岛',parent=dict(block_index=1,hole_index=1)),
      block(box(300,400,-100,100),name='其他根块')]
    s=source();r=calculate(body(s,parts=parts),s);d=r['contact_ranges'][0]
    assert [v['depth'] for v in d['window_nesting']]==[0,1,1,0] and any(c['relation']=='strictly_disjoint' for c in d['window_part_checks']);replay_blocks(r)


def test_40_seeded_islands_independent_exact_rectangle_clipping():
    rng=random.Random(11101);s=source()
    for _ in range(40):
        x,y=rng.randint(-140,100),rng.randint(-140,100);outer=box(x-100,x+100,y-100,y+100);h=box(x-70,x+70,y-70,y+70)
        rx,ry=rng.randint(5,40),rng.randint(5,40);island=[(x-rx,y),(x,y-ry),(x+rx,y),(x,y+ry)]
        parts=[block(outer,holes=[h]),block(island,name='菱形岛区',parent=dict(block_index=1,hole_index=1))]
        r=calculate(body(s,parts=parts),s)
        def area(pts):
            pts=[tuple(map(Fraction,p)) for p in pts];return clip(pts,tuple(map(Fraction,(-200,200,-250,250))))-clip(pts,tuple(map(Fraction,(100,200,-30,30))))
        expected=(area(outer)-area(h)+area(island))/1000000
        assert abs(float(r['contact_ranges'][0]['single_net_area'])-float(expected))<2e-16
    replay_blocks(r)


@pytest.mark.parametrize('offset',[-300,-160,-40,0,40,160,300])
def test_curved_root_and_island_independent_adaptive_meter_integral(offset):
    s=source('circular');s['foundation_items'][0].pop('through_holes',None)
    outer=box(offset-160,offset+160,-160,160);h=box(offset-100,offset+100,-100,100);island=[(offset-60,0),(offset,-60),(offset+60,0),(offset,60)]
    r=calculate(body(s,face='top',parts=[block(outer,holes=[h]),block(island,name='岛区',parent=dict(block_index=1,hole_index=1))]),s)
    meters=lambda pts:[(x/1000,y/1000) for x,y in pts]
    expected=quadrature(meters(outer),[(0,0,.3)])-quadrature(meters(h),[(0,0,.3)])+quadrature(meters(island),[(0,0,.3)])
    assert abs(float(r['contact_ranges'][0]['single_net_area'])-expected)<3e-11;replay_blocks(r)


@pytest.mark.parametrize('unit,factor',[('mm','1'),('cm','.1'),('dm','.01'),('m','.001'),('km','.000001')])
@pytest.mark.parametrize('reverse',[False,True])
def test_common_units_orientation_and_low_call_precision(unit,factor,reverse):
    s=source();b=body(s);c=b['contact_ranges'][0]['plan_range'];c['unit']=unit
    for p in c['parts']:
        for ring in [p]+p['holes']:
            pts=polygon_vertices(ring)
            if reverse:pts.reverse()
            ring['parameters']={a+str(i):str(Decimal(str(v))*Decimal(factor)) for i,xy in enumerate(pts,1) for a,v in zip(('x','y'),xy)}
    with localcontext() as ctx:ctx.prec=6;r=calculate(b,s)
    assert Decimal(r['contact_ranges'][0]['window_area'])==Decimal('.0272');replay_blocks(r)


def test_48_point_and_eight_hole_limits():
    poly=[(-20,-10),(-16,-16),(-10,-20),(10,-20),(16,-16),(20,-10),(20,10),(16,16),(10,20),(-10,20),(-16,16),(-20,10)]
    scale=lambda n,dx=0:[(x*n+dx,y*n) for x,y in poly]
    parts=[block(scale(10),holes=[scale(8)]),block(scale(5),name='岛区',parent=dict(block_index=1,hole_index=1)),block(scale(5,500),name='另一根块')]
    s=source();r=calculate(body(s,parts=parts),s);replay_blocks(r)
    parts[1]['holes']=[excluded([(-2,-2),(2,-2),(0,2)])]
    with pytest.raises(DomainError,match='48'):validate(config(parts))
    tri=lambda x,y:[(x-4,y-4),(x+4,y-4),(x,y+4)]
    parts=[block(box(-200,200,-200,200),holes=[box(-100,100,-100,100),tri(-150,-150),tri(-150,0),tri(-150,150)]),
      block(box(-60,60,-60,60),name='岛区',holes=[tri(x,y) for x,y in [(-30,-30),(30,-30),(-30,30),(30,30)]],parent=dict(block_index=1,hole_index=1)),block(box(400,500,400,500),name='远根块')]
    validate(config(parts));parts[2]['holes']=[excluded(tri(450,450))]
    with pytest.raises(DomainError,match='8'):validate(config(parts))


@pytest.mark.parametrize('fault',['root_claim','wrong_hole','missing_hole','missing_block','self','cycle','boolean','string','extra_parent','part_extra','one','five','no_nesting','missing_proof','confirm','unit','part_unit','hole_cross','outer_cross'])
def test_explicit_parent_schema_and_invalid_topology(fault):
    c=config();p=c['parts'][1]
    if fault=='root_claim':p['parent']=None
    elif fault=='wrong_hole':c['parts'][0]['holes'].append(excluded(box(-90,-70,-90,-70),name='另一孔'));p['parent']['hole_index']=2
    elif fault=='missing_hole':p['parent']['hole_index']=4
    elif fault=='missing_block':p['parent']['block_index']=3
    elif fault=='self':p['parent']['block_index']=2
    elif fault=='cycle':p['holes']=[excluded(box(-10,10,-10,10))];c['parts'][0]['parent']=dict(block_index=2,hole_index=1)
    elif fault=='boolean':p['parent']['block_index']=True
    elif fault=='string':p['parent']['hole_index']='1'
    elif fault=='extra_parent':p['parent']['depth']=1
    elif fault=='part_extra':p['depth']=1
    elif fault=='one':c['parts']=c['parts'][:1]
    elif fault=='five':c['parts']*=3
    elif fault=='no_nesting':p['parameters']=part(box(200,300,-100,100))['parameters'];p['parent']=None
    elif fault=='missing_proof':p['evidence_id']=''
    elif fault=='confirm':c['scope_confirmed']=1
    elif fault=='unit':c['unit']='px'
    elif fault=='part_unit':p['unit']='mm'
    elif fault=='hole_cross':p['parameters']=part(box(-70,20,-20,20))['parameters']
    else:p['parameters']=part(box(-110,20,-20,20))['parameters']
    with pytest.raises(DomainError):validate(c)


def test_nearest_parent_cannot_skip_level_and_concave_hole_chord_cannot_escape():
    c=config(chain());c['parts'][2]['parent']=dict(block_index=1,hole_index=1)
    with pytest.raises(DomainError,match='最近'):validate(c)
    from tests.test_foundation_simple_windows import C_SHAPE
    c=config([block(box(-150,150,-150,150),holes=[C_SHAPE]),block([(-70,-70),(70,-70),(-70,70)],name='越凹口岛',parent=dict(block_index=1,hole_index=1))])
    with pytest.raises(DomainError):validate(c)


@pytest.mark.parametrize('gap',['0','0.000000001'])
def test_island_strict_boundary_gap_without_tolerance(gap):
    x=str(Decimal(-60)+Decimal(gap));c=config();c['parts'][1]['parameters']=part([(x,-20),(20,-20),(20,20),(x,20)])['parameters']
    if gap=='0':
        with pytest.raises(DomainError):validate(c)
    else:validate(c)


def test_100_seeded_net_claims_against_independent_fraction_clip():
    rng=random.Random(11102);a={'plan_range':config()}
    for _ in range(100):
        x,y=rng.randint(-120,120),rng.randint(-120,120);b=tuple(map(Fraction,(x,x+rng.randint(1,100),y,y+rng.randint(1,100))))
        def area(p):return clip([tuple(map(Fraction,p)) for p in polygon_vertices(p)],b)
        expected=sum(area(p)-sum(area(h) for h in p['holes']) for p in a['plan_range']['parts'])>0
        assert separation(a,{'plan_range':window('test',*[str(v) for v in b])})['overlap']==expected
    a.update(index=1,name='嵌套区',material_index=1,face='top');b=dict(index=2,name='残余空区',material_index=1,face='top',plan_range=window('test','-50','-30','-10','10'))
    q=checks([a,b])[0];assert q['method']==METHOD and q['left_index']==1 and q['partitions']==separation(a,b)['partitions']
    assert separation(a,{'plan_range':window('test','-10','10','-10','10')})['overlap']


@pytest.mark.parametrize('gap',['0','0.000000001'])
def test_root_zero_but_island_gap_prevents_whole_zero(gap):
    v=str(Decimal(25)+Decimal(gap));hs=[{**hole(x='0',diameter='80'),'vertical_range':interval('100','300')},
      {**rectangle(name='H2',x='-'+v,length='50',width='100'),'vertical_range':interval('300','500')},
      {**rectangle(name='H3',x=v,length='50',width='100'),'vertical_range':interval('300','500')}]
    s=end_source(holes=hs);r=calculate(body(s,face='hole_top'),s);d=r['contact_ranges'][0]
    assert ('polygon_zero_contact' in d)==(gap=='0') and (Decimal(d['single_net_area'])==0)==(gap=='0')
    assert Decimal(d['window_parts'][0]['single_gross_area'])==0 and Decimal(d['window_parts'][1]['single_gross_area'])>0
    replay_blocks(r)


def setup(clients,pid=None,zero=False):
    e,r,pid,s,b,eid=old_setup(clients,pid,zero)
    outer=box(-50,50,-50,50) if zero else box(170,210,-30,30)
    h=box(-30,30,-30,30) if zero else box(180,200,-20,20)
    island=box(-10,10,-10,10) if zero else box(192,198,5,15)
    b['contact_ranges'][0]['plan_range']=config([block(outer,eid,holes=[h]),block(island,eid,'合成内孔岛区',parent=dict(block_index=1,hole_index=1))],eid)
    return e,r,pid,s,b,eid


def test_workflow_readonly_roles_foreign_island_basis_history_concurrency_source_stale(clients):
    e,r,pid,s,b,eid=setup(clients);before=get(e,pid);post(e,pid,b,'/components/preview');assert get(e,pid)==before
    post(r,pid,b,expected=403);assert clients['other'].get(f'/api/v1/projects/{pid}').status_code==404
    foreign=create(e)['id'];upload(e,foreign,'异项目合成');bad=copy.deepcopy(b);bad['contact_ranges'][0]['plan_range']['parts'][1]['evidence_id']=get(e,foreign)['evidence'][-1]['id'];post(e,pid,bad,expected=404)
    c=post(e,pid,b).json();approve(r,pid,'components',c);c=next(x for x in get(e,pid)['components'] if x['id']==c['id']);b['object_hash']=store.digest(c);b['contact_ranges'][0]['plan_range']['parts'][1]['parameters']['x1']='192.5'
    res=e.put(f'/api/v1/projects/{pid}/components/'+c['id'],json=b);assert res.status_code==200,res.text;assert e.put(f'/api/v1/projects/{pid}/components/'+c['id'],json=b).status_code==409
    assert get(e,pid)['component_history'][-1]==c;approve(r,pid,'components',res.json())
    s['object_hash']=store.digest(s);s['foundation_items'][0]['levels'][1]['parameters']['diameter']='420'
    res=e.put(f'/api/v1/projects/{pid}/components/'+s['id'],json=s);assert res.status_code==200,res.text;approve(r,pid,'components',res.json())
    assert next(x for x in get(e,pid)['components'] if x['id']==c['id'])['state']=='stale'


def test_residual_parent_hole_can_be_claimed_but_island_overlap_cannot(clients):
    e,r,pid,s,b,eid=setup(clients);post(e,pid,b)
    other=copy.deepcopy(b);other['name']='父孔残余空区';other['contact_ranges'][0]['plan_range']=window(eid,'181','190','5','15');post(e,pid,other)
    bad=copy.deepcopy(other);bad['contact_ranges'][0]['plan_range']=window(eid,'191','192.000000001','5','15');before=get(e,pid);post(e,pid,bad,expected=409);assert get(e,pid)==before
    bad=copy.deepcopy(other);bad['contact_ranges'][0]['plan_range']['coordinate_basis']='其他XY正向';post(e,pid,bad,expected=409)


@pytest.mark.parametrize('fault',['parent_input','parent_trace','depth','whole_hierarchy','relation','island_quantity','proof_depth','child_proof','whole_proof','expression'])
def test_hierarchy_or_quantity_tamper_rejected(clients,fault):
    e,r,pid,s,b,eid=setup(clients,zero=True);c=post(e,pid,b).json()
    def forge(p):
        c=p['components'][-1];d=c['preview']['contact_ranges'][0]
        if fault=='parent_input':c['contact_ranges'][0]['plan_range']['parts'][1]['parent']=None
        elif fault=='parent_trace':d['window_parts'][1]['parent']=None
        elif fault=='depth':d['window_parts'][1]['nesting_depth']=0
        elif fault=='whole_hierarchy':d['window_nesting']=[]
        elif fault=='relation':d['window_part_checks']=[]
        elif fault=='island_quantity':d['window_parts'][1]['single_gross_area']='999'
        elif fault=='proof_depth':d['polygon_net_trace']['partitions'][0]['nesting_depth']=99
        elif fault=='child_proof':d['window_parts'][1]['polygon_zero_contact']['net_trace']['partitions'][0]['remaining']=[{}]
        elif fault=='whole_proof':d['polygon_zero_contact']['net_trace']['partitions'][0]['remaining']=[{}]
        else:d['window_expression_definitions'][0]['value']='999'
    mutate(pid,forge);before=get(e,pid);approve(r,pid,'components',c,expected=422 if fault=='parent_input' else 409);assert get(e,pid)==before


@pytest.mark.parametrize('gate',['preview','save','review','calculate','book_review'])
def test_island_relation_original_bytes_all_five_gates(clients,gate):
    e,r,pid,s,b,eid=setup(clients);upload(e,pid,'合成岛区位于B1 W1内部的支模及坐标原件','island-parent-basis.txt')
    for c in get(e,pid)['components']:approve(r,pid,'components',c)
    b['source_component_hash']=e.get(f'/api/v1/projects/{pid}/foundation-horizontal-formwork-candidates',params={'component_id':s['id']}).json()['source_component_hash'];block=b['contact_ranges'][0]['plan_range']['parts'][1];block['evidence_id']=get(e,pid)['evidence'][-1]['id'];block['evidence']={k:block['evidence_id'] for k in block['evidence']}
    if gate in ('review','calculate','book_review'):c=post(e,pid,b).json()
    if gate in ('calculate','book_review'):approve(r,pid,'components',c)
    if gate=='book_review':t=post(e,pid,{},'/takeoffs/calculate').json()[-1]
    doc=get(e,pid)['documents'][-1];(store.ROOT/'originals'/doc['hash']).write_bytes(b'corrupted-synthetic-island-parent-original')
    if gate=='preview':post(e,pid,b,'/components/preview',409)
    elif gate=='save':post(e,pid,b,expected=409)
    elif gate=='review':approve(r,pid,'components',c,expected=409)
    elif gate=='calculate':post(e,pid,{},'/takeoffs/calculate',409)
    else:approve(r,pid,'takeoffs',t,expected=409)


@pytest.mark.parametrize('fmt',['json','xlsx','pdf'])
@pytest.mark.parametrize('zero',[False,True])
def test_three_formats_hierarchy_and_old_quote_isolation(clients,fmt,zero):
    from backend.foundation_multi_window_reports import NESTED_TITLES,NESTING_TITLE
    e,r,pid,_,q=priced_project(clients);e,r,pid,s,b,eid=setup(clients,pid,zero);c=post(e,pid,b).json();approve(r,pid,'components',c)
    for t in post(e,pid,{},'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
    post(e,pid,{},'/boq/generate');url=f'/api/v1/projects/{pid}/exports';res=e.get(url,params=dict(format=fmt,unpriced=True));assert res.status_code==200,res.text
    titles=set(NESTED_TITLES[:5] if zero else NESTED_TITLES[:4])|{NESTING_TITLE}
    if fmt=='json':assert titles<={t[0] for t in res.json()['tables']}
    elif fmt=='xlsx':assert titles<=set(load_workbook(io.BytesIO(res.content)).sheetnames)
    else:assert all(any(title in page.extract_text() for page in PdfReader(io.BytesIO(res.content)).pages) for title in titles)
    assert not any(t[0] in titles for t in e.get(url,params=dict(format='json',quote_id=q['id'])).json()['tables'])
