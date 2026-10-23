"""Synthetic fixed-step post layouts, not a structural design acceptance."""
import copy,io,random
from decimal import Decimal,localcontext
import pytest
from openpyxl import load_workbook
from pypdf import PdfReader
from backend.engine import DomainError
from backend import store
from backend.railing_material import calculate
from backend.railing_takeoff import snapshot as source_snapshot
from tests.test_railing_posts import member as equal_member,body as equal_body,geometry,material_setup
from tests.test_route_railing import payload as railing_payload
from tests.test_route_takeoff import point,evaluate
from tests.test_railing_takeoff import interval
from tests.test_polygon_takeoff import post,mutate
from tests.test_workflow import get,approve,upload,priced_project,create


def member(s,eid='test'):
 v=equal_member(s,eid);v.update(name='合成定距立柱',mode='fixed_posts',measurement_basis='固定2m，从明确锚端逐次定距，余跨留另一端，不均分；每根1200mm、3000g/m；共享端点一次')
 for i,r in enumerate(v['post_ranges']):
  r.update(direction='from_start' if i==0 else 'from_end',measurement_basis='固定2000mm，AB从起端、BC从终端定距，余跨分别留终端与起端；共享B只归AB终端')
  r['evidence']['direction']=eid
 return v


def body(s,eid='test',mixed=False):
 b=equal_body(s,eid,mixed);b['material_items'][-1]=member(s,eid);b['measurement_basis']='合成定距留余跨及独立下料质量验收，无真实工程效力';return b


def setup(clients,pid=None):
 e,r,pid,s,_,eid=material_setup(clients,pid)
 upload(e,pid,'合成定距留余跨验收，非真实工程：AB三维13m、BC5m，来源重复2。固定间距各2000mm；AB从起端定距S0,2,4,6,8,10,12,13m，终端余跨1m，共8根；BC从终端定距，起端不计（B只归AB），S1,3,5m，共3根，起端余跨1m。每根下料1200mm、3000g/m，单组11根39.6kg，总79.2kg，修约3位。AB改3000mm时S0,3,6,9,12,13m，共6根、终端余跨1m；BC仍3根，总64.8kg。改锚向仅改变位置，不改变根数。','fixed-posts-basis.txt')
 for c in get(e,pid)['components']:approve(r,pid,'components',c)
 s=get(e,pid)['components'][-1];eid=get(e,pid)['evidence'][-1]['id'];return e,r,pid,s,body(s,eid),eid


def stations(g,label):return [Decimal(p['s']) for p in g['post_stations'] if p['range_label']==label]


def test_fixed_anchor_directions_remainders_mass_and_asts():
 s=geometry();b=body(s);before=copy.deepcopy((s,b));v=calculate(b,s);assert (s,b)==before and v['quantity']=='79.200' and v['rule_version']=='railing-material-mass-1.4'
 g=v['material_items'][0]['post_layout'];assert g['rule_version']=='railing-fixed-step-posts-1.0' and g['post_count']==11
 assert stations(g,'S1 整段')==[0,2,4,6,8,10,12,13] and stations(g,'S2 整段')==[1,3,5]
 assert [(r['regular_interval_count'],Decimal(r['residual_spacing']),r['residual_side']) for r in g['post_ranges']]==[(6,1,'end'),(2,1,'start')]
 assert all('actual_spacing' not in r for r in g['post_ranges'])
 with localcontext() as c:
  c.prec=80;assert evaluate(v['formula_ast'],v['dimensions'])==Decimal('79.2')
  for p in g['post_stations']:assert evaluate(p['station_formula_ast'],v['dimensions'])==Decimal(p['s'])
  for r in g['post_ranges']:assert evaluate(r['residual_formula_ast'],v['dimensions'])==Decimal(r['residual_spacing'])
 old=calculate(equal_body(s),s);assert old['rule_version']=='railing-material-mass-1.3' and stations(old['material_items'][0]['post_layout'],'S1 整段')[1]!=2
 assert Decimal(calculate(body(s,mixed=True),s)['quantity'])==Decimal('806.4')


@pytest.mark.parametrize('direction',['from_start','from_end'])
@pytest.mark.parametrize('length,spacing,regular,residual,expected',[(6,2,3,0,[0,2,4,6]),(1,2,0,1,[0,1]),(5,2,2,1,None)])
def test_exact_multiple_shorter_than_step_and_both_anchors(direction,length,spacing,regular,residual,expected):
 s=geometry();src=s['source_geometry'];src['route_points']=[point('A',0,0,0),point('B',length,0,0)];s['source_component_hash']=store.digest(source_snapshot(src));s['edge_decisions']=railing_payload(src)['edge_decisions'];b=body(s);r=b['material_items'][0]['post_ranges'][0];r.update(direction=direction,spacing=str(spacing),spacing_unit='m');g=calculate(b,s)['material_items'][0]['post_layout'];z=g['post_ranges'][0]
 assert z['regular_interval_count']==regular and Decimal(z['residual_spacing'])==residual and (z['residual_side']=='none')==(residual==0)
 assert stations(g,'S1 整段')==(expected if expected else [0,2,4,5] if direction=='from_start' else [0,1,3,5])


@pytest.mark.parametrize('basis',['spatial','horizontal'])
@pytest.mark.parametrize('unit,raw',[('m','2'),('mm','2000'),('cm','200'),('dm','20'),('km','.002')])
def test_spacing_units_and_source_basis(basis,unit,raw):
 s=geometry(basis=basis);b=body(s)
 for r in b['material_items'][0]['post_ranges']:r.update(spacing=raw,spacing_unit=unit)
 v=calculate(b,s);assert v['quantity']==('79.200' if basis=='spatial' else '50.400')


def test_reverse_anchor_and_endpoint_ownership_change_positions_independently():
 s=geometry();b=body(s);b['material_items'][0]['post_ranges'][0]['direction']='from_end';g=calculate(b,s)['material_items'][0]['post_layout'];assert stations(g,'S1 整段')==[0,1,3,5,7,9,11,13] and g['post_count']==11
 b['material_items'][0]['post_ranges'][1]['start_post']=True
 with pytest.raises(DomainError,match='共享端点'):calculate(b,s)
 b['material_items'][0]['post_ranges'][0]['end_post']=False;assert calculate(b,s)['quantity']=='79.200'


def test_local_ranges_keep_source_s_origin_and_do_not_bridge_gaps():
 s=geometry();s['edge_decisions'][0].update(mode='ranges',ranges=[interval(1,6),interval(8,12)]);b=body(s)
 for r in b['material_items'][0]['post_ranges']:r.update(start_post=True,end_post=True)
 g=calculate(b,s)['material_items'][0]['post_layout'];assert stations(g,'S1 区间1')==[1,3,5,6] and stations(g,'S1 区间2')==[8,10,12]


@pytest.mark.parametrize('spacing,n',[('1.41421356237309505',1),('1.41421356237309504',2)])
def test_irrational_boundary_not_classified_as_exact_multiple(spacing,n):
 s=geometry();src=s['source_geometry'];src['route_points']=[point('A',0,0,0),point('B',1,1,0)];s['source_component_hash']=store.digest(source_snapshot(src));s['edge_decisions']=railing_payload(src)['edge_decisions'];b=body(s);r=b['material_items'][0]['post_ranges'][0];r.update(direction='from_end',spacing=spacing,spacing_unit='m')
 with localcontext() as c:c.prec=9;g=calculate(b,s)['material_items'][0]['post_layout']
 z=g['post_ranges'][0];assert z['interval_count']==n and z['residual_side']=='start' and Decimal(z['residual_spacing'])>0 and z['regular_interval_count']==n-1
 if n==2:assert 0<stations(g,'S1 整段')[1]<Decimal('1e-16')


@pytest.mark.parametrize('fault',['missing_direction','blank_direction','bad_direction','bool_direction','missing_direction_proof','blank_direction_proof','extra_direction_proof','equal_mode_direction','old_range','duplicate','manual_count','zero_spacing','nan','missing_end','string_end'])
def test_strict_fixed_schema(fault):
 s=geometry();b=body(s);v=b['material_items'][0];r=v['post_ranges'][0]
 if fault=='missing_direction':r.pop('direction')
 elif fault=='blank_direction':r['direction']=''
 elif fault=='bad_direction':r['direction']='reverse_s'
 elif fault=='bool_direction':r['direction']=True
 elif fault=='missing_direction_proof':r['evidence'].pop('direction')
 elif fault=='blank_direction_proof':r['evidence']['direction']=''
 elif fault=='extra_direction_proof':v['evidence']['direction']='test'
 elif fault=='equal_mode_direction':v['mode']='spaced_posts'
 elif fault=='old_range':r['range_id']='old'
 elif fault=='duplicate':v['post_ranges'][1]['range_id']=r['range_id']
 elif fault=='manual_count':v['quantity']='11'
 elif fault=='zero_spacing':r['spacing']='0'
 elif fault=='nan':r['spacing']='NaN'
 elif fault=='missing_end':r.pop('end_post')
 elif fault=='string_end':r['end_post']='true'
 with pytest.raises(DomainError):calculate(b,s)


def test_seeded_independent_integer_station_oracle():
 rng=random.Random(7101)
 for _ in range(40):
  s=geometry();src=s['source_geometry'];L=rng.randint(1,500);d=rng.randint(1,50);src['route_points']=[point('A',0,0,0),point('B',L,0,0)];s['source_component_hash']=store.digest(source_snapshot(src));s['edge_decisions']=railing_payload(src)['edge_decisions'];b=body(s);r=b['material_items'][0]['post_ranges'][0];r.update(spacing=str(d),spacing_unit='m',direction=rng.choice(['from_start','from_end']))
  expected=sorted(set([0,L]+(list(range(d,L,d)) if r['direction']=='from_start' else list(range(L-d,0,-d)))))
  g=calculate(b,s)['material_items'][0]['post_layout'];assert stations(g,'S1 整段')==expected and g['post_count']==len(expected)


@pytest.mark.parametrize('fmt',['json','xlsx','pdf'])
def test_workflow_exports_and_old_quote_isolation(clients,fmt):
 e,r,pid,_,quote=priced_project(clients);e,r,pid,s,b,eid=setup(clients,pid);before=get(e,pid);post(e,pid,b,'/components/preview');assert get(e,pid)==before
 c=post(e,pid,b).json();approve(r,pid,'components',c)
 for t in post(e,pid,{},'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
 assert any(v['quantity']=='79.200' for v in post(e,pid,{},'/boq/generate').json())
 url=f'/api/v1/projects/{pid}/exports';res=e.get(url,params=dict(format=fmt,unpriced=True));assert res.status_code==200,res.text
 titles=['定距立柱质量汇总','定距立柱锚向余跨与归属','定距立柱跨数精确校验','定距立柱逐根端距']
 if fmt=='json':assert set(titles)<={t[0] for t in res.json()['tables']}
 elif fmt=='xlsx':assert set(titles)<=set(load_workbook(io.BytesIO(res.content)).sheetnames)
 else:assert all(v in ''.join(p.extract_text() for p in PdfReader(io.BytesIO(res.content)).pages) for v in titles)
 old=e.get(url,params=dict(format='json',quote_id=quote['id'])).json();assert old['quote']==quote and not any(t[0] in titles for t in old['tables'])


@pytest.mark.parametrize('fault',['direction','station','residual','side','formula','count','source','book'])
def test_saved_and_book_records_are_rebuilt(clients,fault):
 e,r,pid,s,b,eid=setup(clients);c=post(e,pid,b).json()
 if fault=='book':approve(r,pid,'components',c);post(e,pid,{},'/takeoffs/calculate')
 def forge(p):
  c=p['components'][-1];v=c['preview']['material_items'][0];g=v['post_layout']
  if fault=='direction':c['material_items'][0]['post_ranges'][0]['direction']='from_end'
  elif fault=='station':g['post_stations'][1]['s']='1'
  elif fault=='residual':g['post_ranges'][0]['residual_spacing']='2'
  elif fault=='side':g['post_ranges'][0]['residual_side']='start'
  elif fault=='formula':g['post_ranges'][0]['residual_formula_ast']={'constant':'1'}
  elif fault=='count':g['post_count']=999
  elif fault=='source':c['source_geometry']['edge_decisions'][0]['mode']='exclude'
  else:p['takeoffs'][-1]['material_items'][0]['post_layout']['post_stations'][1]['s']='1'
 mutate(pid,forge);coll='takeoffs' if fault=='book' else 'components';approve(r,pid,coll,get(e,pid)[coll][-1],expected=409)


def test_revision_history_concurrency_and_source_invalidation(clients):
 e,r,pid,s,b,eid=setup(clients);c=post(e,pid,b).json();approve(r,pid,'components',c)
 for t in post(e,pid,{},'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
 post(e,pid,{},'/boq/generate');old=get(e,pid)['components'][-1];new=copy.deepcopy(b);new['material_items'][0]['post_ranges'][0]['spacing']='3000';new['object_hash']=store.digest(old);url=f'/api/v1/projects/{pid}/components/'+c['id'];res=e.put(url,json=new);assert res.status_code==200,res.text
 p=get(e,pid);assert p['component_history'][-1]==old and all(t['state']=='stale' for t in p['takeoffs']) and any(v['quantity']=='79.200' for v in p['boq']);assert res.json()['preview']['quantity']=='64.800' and e.put(url,json=new).status_code==409;approve(r,pid,'components',res.json())
 for t in post(e,pid,{},'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
 assert any(v['quantity']=='64.800' for v in post(e,pid,{},'/boq/generate').json())
 s['object_hash']=store.digest(s);s['edge_decisions'][0].update(mode='ranges',ranges=[interval(0,10,eid)]);res=e.put(f'/api/v1/projects/{pid}/components/'+s['id'],json=s);assert res.status_code==200,res.text;approve(r,pid,'components',res.json());assert get(e,pid)['components'][-1]['state']=='stale';post(e,pid,b,'/components/preview',409)


@pytest.mark.parametrize('gate',['preview','save','review','calculate','book_review'])
def test_original_bytes_gates(clients,gate):
 e,r,pid,s,b,eid=setup(clients)
 if gate in ('review','calculate','book_review'):c=post(e,pid,b).json()
 if gate in ('calculate','book_review'):approve(r,pid,'components',c)
 if gate=='book_review':t=post(e,pid,{},'/takeoffs/calculate').json()[-1]
 doc=get(e,pid)['documents'][-1];(store.ROOT/'originals'/doc['hash']).write_bytes(b'corrupt')
 if gate=='preview':post(e,pid,b,'/components/preview',409)
 elif gate=='save':post(e,pid,b,expected=409)
 elif gate=='review':approve(r,pid,'components',c,expected=409)
 elif gate=='calculate':post(e,pid,{},'/takeoffs/calculate',409)
 else:approve(r,pid,'takeoffs',t,expected=409)


def test_roles_tenant_and_foreign_direction_evidence(clients):
 e,r,pid,s,b,eid=setup(clients);post(r,pid,b,expected=403);assert clients['other'].get(f'/api/v1/projects/{pid}').status_code==404
 other=create(e)['id'];upload(e,other,'其他工程原件');b['material_items'][0]['post_ranges'][0]['evidence']['direction']=get(e,other)['evidence'][-1]['id'];post(e,pid,b,expected=404)


def test_fixed_and_equal_rows_keep_separate_dimension_namespaces():
 s=geometry();b=body(s);b['material_items'].append(equal_member(s));v=calculate(b,s)
 assert v['quantity']=='158.400' and v['material_items'][0]['post_layout']['rule_version']=='railing-fixed-step-posts-1.0' and v['material_items'][1]['post_layout']['rule_version']=='railing-equal-bay-posts-1.0'
 with localcontext() as c:
  c.prec=80;assert evaluate(v['formula_ast'],v['dimensions'])==Decimal('158.4')
  for r in v['material_items']:assert evaluate(r['formula_ast'],v['dimensions'])==Decimal('39.6')


@pytest.mark.parametrize('direction',['from_start','from_end'])
def test_closed_polygon_and_hole_fixed_endpoint_ownership(direction):
 from tests.test_railing_takeoff import geometry as pg,payload as pp
 src=pg();s=pp(src);s.update(id='polygon',revision=1,source_geometry=source_snapshot(src))
 for d in s['edge_decisions']:d.update(mode='full',ranges=[])
 b=body(s)
 for r in b['material_items'][0]['post_ranges']:r.update(direction=direction,start_post=True,end_post=False)
 v=calculate(b,s);assert v['material_items'][0]['post_layout']['post_count']==15
 b['material_items'][0]['post_ranges'][0]['end_post']=True
 with pytest.raises(DomainError,match='共享端点'):calculate(b,s)
