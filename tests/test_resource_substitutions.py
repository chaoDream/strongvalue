"""Synthetic substitution workflow, not a real material mix or performance verification."""
import copy
import io
import uuid
import pytest
from openpyxl import load_workbook
from backend import store, main as core
from tests.test_workflow import priced_project, get, approve
from tests.test_quota_conversion import conversion
from tests.test_unpriced_resources import pending_resource


def setup(clients, split=False):
    e,r,pid,comp,old=priced_project(clients);base=f'/api/v1/projects/{pid}'
    row=e.get(base+'/material-substitution-targets').json()[0]
    source=row['resource'];new={**source,'pricing_status':'priced','source_type':'供应商报价','valid_from':'2026-01-01',
                               'name':'替换材料甲','resource_code':'TEST-NEW-A','specification':'合成规格A',
                               'fee_allocation':'合成材料费','consumption':'1.2','price':'120'}
    values=[new]
    if split:values.append({**new,'name':'配合材料乙','resource_code':'TEST-NEW-B','specification':'合成规格B','consumption':'.2','price':'30'})
    body={**{k:row[k] for k in ['key','boq_hash','resource_hash']},'request_id':str(uuid.uuid4()),'name':'纯数学材料替换单',
          'reason':'仅合成材料替换及消耗测试，不验证性能','source':'合成替换依据','license':'仅数学测试','evidence_id':source['evidence_id'],'after':values}
    return e,r,pid,old,row,body


def propose(e,pid,body):
    response=e.post(f'/api/v1/projects/{pid}/material-substitutions',json=body)
    assert response.status_code==200,response.text
    return response.json()['substitution']


def test_split_material_workflow_impact_reapproval_export_and_immutable_quote(clients):
    e,r,pid,old,row,body=setup(clients,True);base=f'/api/v1/projects/{pid}'
    before=get(e,pid);s=propose(e,pid,body)
    assert get(e,pid)['boq']==before['boq'], 'proposal cannot change current pricing'
    assert s['impact']['available'] and s['impact']['after_total']=='2317.50'
    assert s['impact']['delta']=='695.10'
    assert e.post(base+'/material-substitutions/'+s['id']+'/apply').status_code==409
    assert approve(e,pid,'resource_substitutions',s,expected=403).status_code==403
    approve(r,pid,'resource_substitutions',s)
    applied=e.post(base+'/material-substitutions/'+s['id']+'/apply');assert applied.status_code==200,applied.text
    assert e.post(base+'/material-substitutions/'+s['id']+'/apply').json()['duplicate']
    p=get(e,pid);b=p['boq'][0]
    assert b['state']=='draft' and b['quantity']==before['boq'][0]['quantity']
    assert len(b['quotas'][0]['resources'])==2
    assert all(x['substitution_id']==s['id'] for x in b['quotas'][0]['resources'])
    assert old['id'] in p['stale_quotes'] and e.post(base+'/quotes/calculate').status_code==409
    approve(r,pid,'boq',b)
    new=e.post(base+'/quotes/calculate');assert new.status_code==200,new.text
    new=new.json();assert new['total']==s['impact']['after_total']
    assert next(q for q in get(e,pid)['quotes'] if q['id']==old['id'])==old
    assert new['snapshot']['applied_resource_substitutions']==[s['id']]
    for fmt in ['xlsx','pdf','json']:
        response=e.get(base+'/exports',params={'format':fmt,'quote_id':new['id']});assert response.status_code==200
        if fmt=='xlsx':assert '材料替换台账' in load_workbook(io.BytesIO(response.content)).sheetnames
        if fmt=='json':assert '材料替换台账' in [t[0] for t in response.json()['tables']]
    approve(r,pid,'resource_substitutions',get(e,pid)['resource_substitutions'][0],decision='reject')
    assert get(e,pid)['boq'][0]['state']=='stale'
    approve(r,pid,'boq',get(e,pid)['boq'][0],expected=409)
    assert e.post(base+'/quotes/calculate').status_code==409
    assert any('材料替换审批已失效' in x for x in get(e,pid)['blockers'])
    approve(r,pid,'resource_substitutions',get(e,pid)['resource_substitutions'][0],expected=409)


def test_pending_replacement_applies_without_fabricated_amounts(clients):
    e,r,pid,old,row,body=setup(clients);base=f'/api/v1/projects/{pid}'
    body['after']=[pending_resource(body['evidence_id'])]
    s=propose(e,pid,body);assert not s['impact']['available']
    approve(r,pid,'resource_substitutions',s)
    assert e.post(base+'/material-substitutions/'+s['id']+'/apply').status_code==200
    assert len(e.get(base+'/pricing-tasks').json())==1
    approve(r,pid,'boq',get(e,pid)['boq'][0])
    response=e.post(base+'/quotes/calculate');assert response.status_code==409 and '待询价' in response.text
    assert e.get(base+'/exports?format=xlsx&unpriced=true').status_code==200
    raw=e.get(base+'/exports?format=json&unpriced=true').json()['unpriced']
    assert raw['applied_resource_substitutions']==[s['id']]
    assert raw['resource_substitutions'][0]['id']==s['id']
    import json,zipfile
    bundle=e.get(base+'/exports?format=zip&unpriced=true');assert bundle.status_code==200
    archive=zipfile.ZipFile(io.BytesIO(bundle.content))
    assert json.loads(archive.read('unpriced-snapshot.json'))['applied_resource_substitutions']==[s['id']]


@pytest.mark.parametrize('change', ['boq','rules','date'])
def test_changed_baseline_blocks_review_and_apply(clients,change):
    e,r,pid,old,row,body=setup(clients);base=f'/api/v1/projects/{pid}'
    s=propose(e,pid,body);approve(r,pid,'resource_substitutions',s)
    with store.transaction() as c:
        p=core.task_project(c,pid,{'tenant':'local'})
        if change=='boq':p['boq'][0]['features']+=' changed'
        if change=='rules':p['rules'][0]['fees'][0]['rate']='.04'
        if change=='date':p['pricing_date']='2026-12-01'
        store.save(c,p)
    before=get(e,pid)
    assert e.post(base+'/material-substitutions/'+s['id']+'/apply').status_code==409
    assert get(e,pid)==before
    approve(r,pid,'resource_substitutions',before['resource_substitutions'][0],expected=409)


@pytest.mark.parametrize('mutation', ['foreign_evidence','superseded','missing_original','bad_price','bad_consumption','unknown_field','equipment','empty_spec','historical_price','too_many'])
def test_invalid_proposals_do_not_write(clients,mutation):
    e,r,pid,old,row,body=setup(clients);base=f'/api/v1/projects/{pid}'
    if mutation=='foreign_evidence':body['evidence_id']='foreign'
    if mutation=='bad_price':body['after'][0]['price']='NaN'
    if mutation=='bad_consumption':body['after'][0]['consumption']='-1'
    if mutation=='unknown_field':body['after'][0]['price_change_id']='forged'
    if mutation=='equipment':body['after'][0]['type']='设备'
    if mutation=='empty_spec':body['after'][0]['specification']=' '
    if mutation=='historical_price':body['after'][0]['source_type']='历史审核价'
    if mutation=='too_many':body['after']*=21
    if mutation in ('superseded','missing_original'):
        p=get(e,pid)
        if mutation=='missing_original':(store.ROOT/'originals'/p['documents'][0]['hash']).unlink()
        else:
            with store.transaction() as c:
                p=core.task_project(c,pid,{'tenant':'local'});p['documents'][0]['superseded']=True;store.save(c,p)
    before=get(e,pid)
    result=e.post(base+'/material-substitutions',json=body)
    assert result.status_code in (404,409,422),result.text
    assert get(e,pid)==before


def test_duplicate_request_access_and_payload_mismatch(clients):
    e,r,pid,old,row,body=setup(clients);base=f'/api/v1/projects/{pid}'
    s=propose(e,pid,body)
    response=e.post(base+'/material-substitutions',json=body);assert response.json()['duplicate']
    body['reason']+=' changed'
    assert e.post(base+'/material-substitutions',json=body).status_code==409
    assert len(get(e,pid)['resource_substitutions'])==1
    assert clients['other'].get(base+'/material-substitution-targets').status_code==404
    assert clients['other'].post(base+'/material-substitutions/'+s['id']+'/apply').status_code==404
    assert r.post(base+'/material-substitutions',json=body).status_code==403


def test_identity_tampering_and_duplicate_marker_are_blocked(clients):
    e,r,pid,old,row,body=setup(clients);base=f'/api/v1/projects/{pid}'
    s=propose(e,pid,body);approve(r,pid,'resource_substitutions',s);assert e.post(base+'/material-substitutions/'+s['id']+'/apply').status_code==200
    b=get(e,pid)['boq'][0];qs=copy.deepcopy(b['quotas']);qs[0]['resources'][0]['consumption']='9'
    changed=e.put(base+'/boq/'+b['id'],json={'quotas':qs});assert changed.status_code==200
    approve(r,pid,'boq',changed.json(),expected=409)
    assert e.post(base+'/quotes/calculate').status_code==409
    qs=copy.deepcopy(b['quotas']);qs[0]['resources']*=2
    changed=e.put(base+'/boq/'+b['id'],json={'quotas':qs});assert changed.status_code==200
    assert any('替换材料与已应用记录不一致' in x for x in get(e,pid)['blockers'])


def test_quota_conversion_preserved_and_applied_exactly_once(clients):
    e,r,pid,old,row,body=setup(clients);base=f'/api/v1/projects/{pid}'
    b=get(e,pid)['boq'][0];qs=copy.deepcopy(b['quotas']);qs[0]['conversion']=conversion(body['evidence_id'])
    saved=e.put(base+'/boq/'+b['id'],json={'quotas':qs});approve(r,pid,'boq',saved.json())
    row=e.get(base+'/material-substitution-targets').json()[0]
    body.update({k:row[k] for k in ['key','boq_hash','resource_hash']})
    s=propose(e,pid,body);approve(r,pid,'resource_substitutions',s)
    assert e.post(base+'/material-substitutions/'+s['id']+'/apply').status_code==200
    b=get(e,pid)['boq'][0];assert b['quotas'][0]['conversion']==qs[0]['conversion']
    approve(r,pid,'boq',b)
    new=e.post(base+'/quotes/calculate');assert new.status_code==200,new.text
    assert new.json()['lines'][0]['unit_price']=='15.84'
    assert new.json()['total']==s['impact']['after_total']


@pytest.mark.parametrize('kind',['superseded','missing_bytes'])
def test_original_rechecked_after_review(clients,kind):
    e,r,pid,old,row,body=setup(clients);base=f'/api/v1/projects/{pid}'
    s=propose(e,pid,body);approve(r,pid,'resource_substitutions',s)
    p=get(e,pid)
    if kind=='missing_bytes':(store.ROOT/'originals'/p['documents'][0]['hash']).unlink()
    else:
        with store.transaction() as c:
            p=core.task_project(c,pid,{'tenant':'local'});p['documents'][0]['superseded']=True;store.save(c,p)
    before=get(e,pid)
    assert e.post(base+'/material-substitutions/'+s['id']+'/apply').status_code==409
    assert get(e,pid)==before


def test_pending_price_fill_preserves_substitution_authorization(clients):
    e,r,pid,old,row,body=setup(clients);base=f'/api/v1/projects/{pid}'
    after=pending_resource(body['evidence_id']);body['after']=[after]
    s=propose(e,pid,body);approve(r,pid,'resource_substitutions',s)
    assert e.post(base+'/material-substitutions/'+s['id']+'/apply').status_code==200
    b=get(e,pid)['boq'][0];qs=copy.deepcopy(b['quotas']);entry=qs[0]['resources'][0]
    entry.update({k:v for k,v in row['resource'].items() if k not in ['name','unit','consumption','type']})
    entry.update(price='120',pricing_status='priced',source_type='供应商报价')
    updated=e.put(base+'/boq/'+b['id'],json={'quotas':qs});assert updated.status_code==200,updated.text
    approve(r,pid,'boq',updated.json())
    quote=e.post(base+'/quotes/calculate');assert quote.status_code==200,quote.text
    assert quote.json()['lines'][0]['unit_price']=='126.00'


def test_reviewed_substitution_can_be_saved_as_source_linked_catalog_candidate(clients):
    e,r,pid,old,row,body=setup(clients);base=f'/api/v1/projects/{pid}'
    s=propose(e,pid,body);approve(r,pid,'resource_substitutions',s)
    assert e.post(base+'/material-substitutions/'+s['id']+'/apply').status_code==200
    b=get(e,pid)['boq'][0];approve(r,pid,'boq',b)
    saved=e.post(base+'/catalog',json={'boq_id':b['id']});assert saved.status_code==200,saved.text
    candidate=saved.json();approve(r,pid,'catalog',candidate)
    assert candidate['quotas'][0]['resources'][0]['substitution_id']==s['id']
