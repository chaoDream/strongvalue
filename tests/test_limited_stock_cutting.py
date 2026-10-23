"""Finite supply synthetic packing, independent oracle and workflow isolation."""
import copy,io,random
from decimal import Decimal,localcontext
from fractions import Fraction
from collections import Counter
import pytest
from pypdf import PdfReader
from openpyxl import load_workbook
from backend import store,limited_stock_cutting as limited,mixed_stock_cutting as mixed
from backend.engine import DomainError
from backend.railing_material import calculate
from tests.test_multi_stock_cutting import body as multi_body,setup as multi_setup,geometry,member,equal_body,fixed_body,fixed_added,equal_added,fixed_adjusted,equal_adjusted,combined,local_body,path_body
from tests.test_mixed_stock_cutting import config as mixed_config
from tests.test_railing_takeoff import interval
from tests.test_route_takeoff import evaluate
from tests.test_polygon_takeoff import post,mutate
from tests.test_workflow import get,approve,upload,create,priced_project

TITLE='有限供应选用台账'


def limits(eid='test',counts=('2','2','2')):
 return dict(quantities=[dict(quantity=n,evidence_id=eid) for n in counts],scope_confirmed=True,measurement_basis='合成仅本组本次方案各定尺供应上限，须据各根数原件，非库存占用或跨组共享额度，不自动采购或入库，非真实工程',evidence_id=eid)


def body(s,eid='test'):
 b=multi_body(s,eid);b['name']='合成有限供应选材（非真实工程）';c=b['mixed_cutting'][0]['stock_choices'];c.update(supply_confirmed=False,measurement_basis='合成同规格方管A，6000/3523/2423mm统一加工；先最小购长，同长度再最少根数，供应量按独立有限供应依据',supply_limits=limits(eid));return b


def setup(clients,pid=None):
 e,r,pid,s,b,eid=multi_setup(clients,pid)
 upload(e,pid,'合成有限供应软件验收，非真实工程。同规格方管A每米3000g，长短每组各2件3500/2400mm，来源重复2共8件净70.8kg。定尺6000/3523/2423mm，每件3mm锯缝含最后一件，每根两端各10mm、余料下限50mm。仅本混排组本次供应上限各2根，不表示库存预留或跨组共享；选S1/S2/S3各2根，购入23.892m71.676kg，净23.6m、锯缝.024m、端切.12m、待留用.148m，平衡0。S1只1根时无法满足。修订上限S1为4、S2/S3为0时仅选四根6000mm，购入24m72kg，净量不变需重新复核。目标先最小购长同长度再最少根数，不计采购价、不自动入库或消耗实际库存。','limited-stock-basis.txt')
 for c in get(e,pid)['components']:approve(r,pid,'components',c)
 s=get(e,pid)['components'][-1];eid=get(e,pid)['evidence'][-1]['id'];return e,r,pid,s,body(s,eid),eid


def group(s,b):return calculate(b,s)['mixed_stock_cutting']['groups'][0]


def test_exact_finite_plan_counts_ast_balance_and_net_unchanged():
 s=geometry();b=body(s);before=copy.deepcopy((s,b));v=calculate(b,s);g=v['mixed_stock_cutting']['groups'][0];base=calculate(multi_body(s),s)
 assert (s,b)==before and v['quantity']==base['quantity']=='70.800' and v['material_items']==base['material_items'] and v['formula_ast']==base['formula_ast'] and v['rule_version']=='railing-material-mass-1.16'
 assert g['proof']['supply_model']=='finite_limits' and g['proof']['status']=='length_optimal_search' and g['stock_count']==6 and g['total']['purchased_mass']=='71.676'
 assert g['total']['retained_length']=='0.148' and g['total']['trim_length']=='0.12' and g['balance_difference']=='0'
 for o in g['stock_selection']['options']:
  assert o['supply_limit']==o['stock_count']==2 and o['unselected_supply']==0
  assert evaluate(o['supply_limit_formula_ast'],v['dimensions'])==2 and evaluate(o['unselected_supply_formula_ast'],v['dimensions'])==0
 assert len([k for k in v['source_evidence'] if '.mixed_cutting.' in k])==14
 for bar in g['bars']:
  for k,ast in bar['formula_asts'].items():assert evaluate(ast,v['dimensions'])==Decimal(bar['values'][k])
 for k,ast in g['total_formula_asts'].items():assert evaluate(ast,v['dimensions'])==Decimal(g['total'][k])


def test_zero_supply_and_unused_quantities():
 s=geometry();b=body(s);c=b['mixed_cutting'][0]['stock_choices'];c['supply_limits']=limits(counts=('5','0','0'));g=group(s,b)
 assert g['stock_count']==4 and g['total']['purchased_mass']=='72'
 assert [(o['stock_count'],o['unselected_supply']) for o in g['stock_selection']['options']]==[(4,1),(0,0),(0,0)]


@pytest.mark.parametrize('counts',[('0','0','0'),('1','2','2'),('0','2','2'),('0','4','0')])
def test_insufficient_supply_proven_by_capacity_or_search(counts):
 s=geometry();b=body(s);b['mixed_cutting'][0]['stock_choices']['supply_limits']=limits(counts=counts)
 with pytest.raises(DomainError,match='已由容量约束或穷尽搜索证明'):group(s,b)


def test_feasible_search_after_all_initial_heuristics_fail_and_unknown_budget():
 sizes=[10,9,8,8,5,5,3,3];stocks=[7,8,10];counts=[1,2,3];plan,p=limited.pack(sizes,stocks,0,counts)
 assert p['initial_purchased_length'] is None and p['initial_stock_count'] is None and (p['purchased_length'],p['stock_count'])==(53,6) and p['status']=='length_optimal_search'
 with pytest.raises(DomainError,match='不能据此判断供应不足'):limited.pack(sizes,stocks,0,counts,limit=0)
 # Existing feasible seed plus exhausted budget is a usable unproven plan.
 plan,p=limited.pack([6,5,3,2,2,2],[10,11],0,[2,2],limit=0);assert p['status']=='length_feasible_unproven' and not p['search_exhausted']


def test_lower_bound_and_lexicographic_tie():
 plan,p=limited.pack([5,5],[5,10],0,[2,1]);assert (p['purchased_length'],p['stock_count'])==(10,1) and p['status']=='length_optimal_bound'


@pytest.mark.parametrize('fault',['null','list','empty','extra','missing','unconfirmed','truthy','both-supply','neither-supply','no-basis','no-proof','bad-proof','quantities-null','quantities-dict','missing-row','extra-row','row-null','row-list','row-extra','row-missing','row-no-proof','row-bad-proof','negative','fraction','too-large','nan','infinite','bool','precision'])
def test_strict_supply_input_and_mode_conflicts(fault):
 s=geometry();b=body(s);c=b['mixed_cutting'][0]['stock_choices'];l=c['supply_limits'];q=l['quantities'][0]
 if fault=='null':c['supply_limits']=None
 elif fault=='list':c['supply_limits']=[]
 elif fault=='empty':c['supply_limits']={}
 elif fault=='extra':l['inventory_reserved']=True
 elif fault=='missing':l.pop('scope_confirmed')
 elif fault=='unconfirmed':l['scope_confirmed']=False
 elif fault=='truthy':l['scope_confirmed']=1
 elif fault=='both-supply':c['supply_confirmed']=True
 elif fault=='neither-supply':c.pop('supply_limits')
 elif fault=='no-basis':l['measurement_basis']=''
 elif fault=='no-proof':l['evidence_id']=''
 elif fault=='bad-proof':l['evidence_id']=False
 elif fault=='quantities-null':l['quantities']=None
 elif fault=='quantities-dict':l['quantities']={}
 elif fault=='missing-row':l['quantities'].pop()
 elif fault=='extra-row':l['quantities'].append(q)
 elif fault=='row-null':l['quantities'][0]=None
 elif fault=='row-list':l['quantities'][0]=[]
 elif fault=='row-extra':q['used']=2
 elif fault=='row-missing':q.pop('quantity')
 elif fault=='row-no-proof':q['evidence_id']=''
 elif fault=='row-bad-proof':q['evidence_id']=True
 elif fault=='negative':q['quantity']='-1'
 elif fault=='fraction':q['quantity']='1.5'
 elif fault=='too-large':q['quantity']='201'
 elif fault=='nan':q['quantity']='NaN'
 elif fault=='infinite':q['quantity']='Infinity'
 elif fault=='bool':q['quantity']=True
 else:q['quantity']='2.000000000000000001'
 with pytest.raises(DomainError):calculate(b,s)


@pytest.mark.parametrize('unit,factor',[('m',Decimal('.001')),('mm',Decimal('1')),('cm',Decimal('.1')),('dm',Decimal('.01')),('km',Decimal('.000001'))])
def test_units_and_precision(unit,factor):
 s=geometry();b=body(s);g=b['mixed_cutting'][0];g['length_unit']=unit
 for k in ('stock_length','kerf_per_piece','trim_start','trim_end','reusable_min_length'):g[k]=str(Decimal(g[k])*factor)
 for o in g['stock_choices']['alternatives']:o['stock_length']=str(Decimal(o['stock_length'])*factor)
 with localcontext() as ctx:ctx.prec=6;v=calculate(b,s)
 assert v==calculate(b,s) and v['mixed_stock_cutting']['total']['purchased_mass']=='71.676'


@pytest.mark.parametrize('make',[equal_body,fixed_body,fixed_added,equal_added,fixed_adjusted,equal_adjusted,combined,lambda s:combined(s,equal=True),local_body,path_body])
def test_all_post_modes_and_spacing_preserved(make):
 s=geometry();b=make(s);b['material_items'][0]['name']='合成立柱';b['material_items'].append({**member('合成短料','cut_pieces'),'quantity':'2','cut_length':'2400'});old=calculate(copy.deepcopy(b),s);g=mixed_config(members=['合成立柱','合成短料']);g['stock_choices']=body(s)['mixed_cutting'][0]['stock_choices'];g['stock_choices']['supply_limits']=limits(counts=('200','200','200'));b['mixed_cutting']=[g];v=calculate(b,s);assert v['quantity']==old['quantity'] and v['material_items']==old['material_items']


def test_independent_exhaustive_finite_capacity_oracle():
 rng=random.Random(8302)
 def oracle(sizes,stocks,trim,counts):
  best=None;bins=[];used=[0]*len(stocks)
  def walk(i):
   nonlocal best
   if i==len(sizes):
    value=(sum(stocks[j] for j,_ in bins),len(bins));best=value if best is None else min(best,value);return
   for k,(j,load) in enumerate(bins):
    if load+sizes[i]<=stocks[j]-trim:bins[k]=(j,load+sizes[i]);walk(i+1);bins[k]=(j,load)
   for j,stock in enumerate(stocks):
    if used[j]<counts[j] and sizes[i]<=stock-trim:used[j]+=1;bins.append((j,sizes[i]));walk(i+1);bins.pop();used[j]-=1
  walk(0);return best
 for _ in range(80):
  trim=rng.randint(0,2);stocks=sorted(rng.sample(range(4,14),3));counts=[rng.randint(0,3) for _ in stocks];sizes=sorted([rng.randint(1,max(stocks)-trim) for _ in range(rng.randint(2,6))],reverse=True);expected=oracle(sizes,stocks,trim,counts)
  if expected is None:
   with pytest.raises(DomainError,match='已由容量约束或穷尽搜索证明'):limited.pack(sizes,stocks,trim,counts)
  else:
   plan,p=limited.pack(sizes,stocks,trim,counts);assert (p['purchased_length'],p['stock_count'])==expected and p['status']!='length_feasible_unproven'
   used=Counter(j for j,_ in plan);assert all(used[j]<=counts[j] for j in range(3)) and sorted(i for _,ps in plan for i in ps)==list(range(len(sizes)))


def test_200_pieces_and_five_candidates():
 s=geometry();b=body(s)
 for m in b['material_items']:m['quantity']='50'
 c=b['mixed_cutting'][0]['stock_choices'];c['alternatives']+=[dict(stock_length=x,evidence_id='test') for x in ('5926','8000')];c['supply_limits']=limits(counts=('200',)*5);g=group(s,b);assert g['required_piece_count']==200 and all(o['stock_count']<=200 for o in g['stock_selection']['options'])


@pytest.mark.parametrize('fmt',['json','xlsx','pdf'])
def test_lifecycle_exports_and_old_quote_isolation(clients,fmt):
 e,r,pid,_,quote=priced_project(clients);e,r,pid,s,b,eid=setup(clients,pid);before=get(e,pid);post(e,pid,b,'/components/preview');assert get(e,pid)==before;c=post(e,pid,b).json();approve(r,pid,'components',c)
 for t in post(e,pid,{},'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
 assert any(t['quantity']=='70.800' for t in post(e,pid,{},'/boq/generate').json());url=f'/api/v1/projects/{pid}/exports';res=e.get(url,params=dict(format=fmt,unpriced=True));assert res.status_code==200,res.text
 if fmt=='json':assert TITLE in str(res.json()['tables']) and '有限供应范围' in str(res.json()['tables'])
 elif fmt=='xlsx':assert TITLE in load_workbook(io.BytesIO(res.content)).sheetnames
 else:assert TITLE in ''.join(p.extract_text() for p in PdfReader(io.BytesIO(res.content)).pages)
 assert TITLE not in str(e.get(url,params=dict(format='json',quote_id=quote['id'])).json()['tables'])


@pytest.mark.parametrize('fault',['raw','proof','scope','limit','used','remaining','ast','status','book'])
def test_saved_rebuild_and_review_hashes(clients,fault):
 e,r,pid,s,b,eid=setup(clients);c=post(e,pid,b).json()
 if fault=='book':approve(r,pid,'components',c);post(e,pid,{},'/takeoffs/calculate')
 def forge(p):
  c=p['components'][-1];g=c['preview']['mixed_stock_cutting']['groups'][0]
  if fault=='raw':c['mixed_cutting'][0]['stock_choices']['supply_limits']['quantities'][0]['quantity']='3'
  elif fault=='proof':c['mixed_cutting'][0]['stock_choices']['supply_limits']['quantities'][0]['evidence_id']='forged'
  elif fault=='scope':c['mixed_cutting'][0]['stock_choices']['supply_limits']['scope_confirmed']=False
  elif fault=='limit':g['stock_selection']['options'][0]['supply_limit']=3
  elif fault=='used':g['stock_selection']['options'][0]['stock_count']=1
  elif fault=='remaining':g['stock_selection']['options'][0]['unselected_supply']=1
  elif fault=='ast':g['stock_selection']['options'][0]['unselected_supply_formula_ast']={'constant':'1'}
  elif fault=='status':g['proof']['supply_model']='unlimited'
  else:p['takeoffs'][-1]['mixed_stock_cutting']['groups'][0]['stock_selection']['options'][0]['supply_limit']=3
 mutate(pid,forge);coll='takeoffs' if fault=='book' else 'components';approve(r,pid,coll,get(e,pid)[coll][-1],expected=404 if fault=='proof' else 422 if fault=='scope' else 409)


def test_same_net_revision_history_stale_books_removal_and_source(clients):
 e,r,pid,s,b,eid=setup(clients);c=post(e,pid,b).json();approve(r,pid,'components',c)
 for t in post(e,pid,{},'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
 post(e,pid,{},'/boq/generate');before=get(e,pid);old=before['components'][-1];new=copy.deepcopy(b);new['mixed_cutting'][0]['stock_choices']['supply_limits']=limits(eid,('4','0','0'));new['object_hash']=store.digest(old);url=f'/api/v1/projects/{pid}/components/'+old['id'];res=e.put(url,json=new);assert res.status_code==200,res.text
 p=get(e,pid);assert p['component_history'][-1]==old and p['boq']==before['boq'] and all(t['state']=='stale' for t in p['takeoffs']) and res.json()['preview']['quantity']=='70.800' and res.json()['preview']['mixed_stock_cutting']['total']['purchased_mass']=='72';assert e.put(url,json=new).status_code==409;approve(r,pid,'components',res.json())
 current=get(e,pid)['components'][-1];new['object_hash']=store.digest(current);new['mixed_cutting'][0]['stock_choices'].pop('supply_limits');new['mixed_cutting'][0]['stock_choices']['supply_confirmed']=True;res=e.put(url,json=new);assert res.status_code==200 and res.json()['preview']['rule_version']=='railing-material-mass-1.15' and not res.json()['preview']['mixed_stock_cutting']['groups'][0]['stock_selection'].get('supply_limited');approve(r,pid,'components',res.json())
 s['object_hash']=store.digest(s);s['edge_decisions'][0].update(mode='ranges',ranges=[interval(0,10,eid)]);res=e.put(f'/api/v1/projects/{pid}/components/'+s['id'],json=s);assert res.status_code==200;approve(r,pid,'components',res.json());assert get(e,pid)['components'][-1]['state']=='stale'


@pytest.mark.parametrize('gate',['preview','save','review','calculate','book_review'])
def test_limit_original_bytes_gates(clients,gate):
 e,r,pid,s,b,eid=setup(clients)
 if gate in ('review','calculate','book_review'):c=post(e,pid,b).json()
 if gate in ('calculate','book_review'):approve(r,pid,'components',c)
 if gate=='book_review':t=post(e,pid,{},'/takeoffs/calculate').json()[-1]
 d=get(e,pid)['documents'][-1];(store.ROOT/'originals'/d['hash']).write_bytes(b'corrupt')
 if gate=='preview':post(e,pid,b,'/components/preview',409)
 elif gate=='save':post(e,pid,b,expected=409)
 elif gate=='review':approve(r,pid,'components',c,expected=409)
 elif gate=='calculate':post(e,pid,{},'/takeoffs/calculate',409)
 else:approve(r,pid,'takeoffs',t,expected=409)


def test_roles_tenant_and_foreign_limit_evidence(clients):
 e,r,pid,s,b,eid=setup(clients);post(r,pid,b,expected=403);assert clients['other'].get(f'/api/v1/projects/{pid}').status_code==404
 other=create(e)['id'];upload(e,other,'异项目');fid=get(e,other)['evidence'][-1]['id'];original=copy.deepcopy(b);b['mixed_cutting'][0]['stock_choices']['supply_limits']['quantities'][0]['evidence_id']=fid;post(e,pid,b,expected=404)
 b=original;b['mixed_cutting'][0]['stock_choices']['supply_limits']['evidence_id']=fid;post(e,pid,b,expected=404)


def test_infeasible_and_unknown_previews_do_not_mutate_project(clients,monkeypatch):
 e,r,pid,s,b,eid=setup(clients);before=get(e,pid);b['mixed_cutting'][0]['stock_choices']['supply_limits']=limits(eid,('1','2','2'));res=post(e,pid,b,'/components/preview',422);assert '已由容量约束或穷尽搜索证明' in res.text and get(e,pid)==before
 monkeypatch.setattr(mixed,'SEARCH_LIMIT',0);c=b['mixed_cutting'][0];c.update(stock_length='10000',kerf_per_piece='0',trim_start='0',trim_end='0');c['stock_choices']['alternatives']=[dict(stock_length=x,evidence_id=eid) for x in ('7000','8000')];c['stock_choices']['supply_limits']=limits(eid,('6','2','4'))
 b['material_items']=[{**member('合成'+str(x),'cut_pieces',eid),'cut_length':str(x*1000),'quantity':str(n)} for x,n in ((10,1),(9,1),(8,2),(5,2),(3,2))];c['members']=[m['name'] for m in b['material_items']];res=post(e,pid,b,'/components/preview',422);assert '不能据此判断供应不足' in res.text and get(e,pid)==before


def test_no_initial_solution_survives_export_and_long_names(clients):
 e,r,pid,s,b,eid=setup(clients);c=b['mixed_cutting'][0];c.update(stock_length='10000',kerf_per_piece='0',trim_start='0',trim_end='0');c['stock_choices']['alternatives']=[dict(stock_length=x,evidence_id=eid) for x in ('7000','8000')];c['stock_choices']['supply_limits']=limits(eid,('6','2','4'))
 b['material_items']=[{**member('合成'+str(x),'cut_pieces',eid),'cut_length':str(x*1000),'quantity':str(n)} for x,n in ((10,1),(9,1),(8,2),(5,2),(3,2))];c['members']=[m['name'] for m in b['material_items']];b['name']='合成'+('长名称验收'*39)+'末';c['name']='合成'+('长混排组名'*39)+'末';obj=post(e,pid,b).json();assert obj['preview']['mixed_stock_cutting']['groups'][0]['proof']['initial_purchased_length'] is None;approve(r,pid,'components',obj)
 for t in post(e,pid,{},'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
 for fmt in ('json','pdf'):
  res=e.get(f'/api/v1/projects/{pid}/exports',params=dict(format=fmt,unpriced=True));assert res.status_code==200,res.text


def test_existing_feasible_but_unproven_export(clients,monkeypatch):
 monkeypatch.setattr(mixed,'SEARCH_LIMIT',0);e,r,pid,s,b,eid=setup(clients);obj=post(e,pid,b).json();assert obj['preview']['mixed_stock_cutting']['groups'][0]['proof']['status']=='length_feasible_unproven';approve(r,pid,'components',obj)
 for t in post(e,pid,{},'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
 res=e.get(f'/api/v1/projects/{pid}/exports',params=dict(format='json',unpriced=True));assert res.status_code==200 and '尚未证明双目标最优' in str(res.json()['tables'])
