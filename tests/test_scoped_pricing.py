"""Synthetic scoped fees: ordered scopes, exact closure, evidence binding and compatible previews."""
import copy
import io
import uuid
from decimal import Decimal
import pytest
from openpyxl import load_workbook
from backend import store
from backend.engine import calculate,DomainError
from tests.test_workflow import get,approve
from tests.test_engineering_structure import setup as hierarchy_setup,create as create_structure,quote


def fee(name,rate,basis='subtotal',eid='synthetic'):
    return {'name':name,'rate':rate,'basis':basis,'source':'仅数学费用规则','evidence_id':eid}


def math_input():
    resource={'name':'合成材料','type':'材料','consumption':'1.05','price':'100.005'}
    items=[{'id':key,'quantity':qty,'quotas':[{'code':'TEST','resources':[resource]}]} for key,qty in [('a','7'),('b','8')]]
    structure={'nodes':[{'id':'s1','kind':'single','parent_id':None,'name':'单项甲'}, {'id':'u1','kind':'unit','parent_id':'s1','name':'单位甲'},
                        {'id':'s2','kind':'single','parent_id':None,'name':'单项乙'}, {'id':'u2','kind':'unit','parent_id':'s2','name':'单位乙'}],
               'assignments':{'a':'u1','b':'u2'}}
    rules={'fees':[fee('项目费','.03','children_total')],'node_fees':[{'node_id':'u1','fees':[fee('单位费','.1')]},{'node_id':'s1','fees':[fee('单项费','.05','children_total')]}]}
    return items,rules,structure


def setup(clients):
    e,r,pid,old,body=hierarchy_setup(clients);base=f'/api/v1/projects/{pid}'
    body['assignments']={b['id']:'u1' if Decimal(b['quantity'])==7 else 'u2' for b in get(e,pid)['boq']}
    s=create_structure(e,pid,body);approve(r,pid,'engineering_structures',s)
    p=get(e,pid);rule=copy.deepcopy(p['rules'][0]);eid=p['evidence'][0]['id']
    rule.update(node_fees=[{'node_id':'u1','fees':[fee('合成单位费','.1',eid=eid)]},{'node_id':'s1','fees':[fee('合成单项费','.05','children_total',eid)]}],
                structure_hash=p['engineering_structure_hash'],fees=[fee('合成项目费','.03','children_total',eid)])
    return e,r,pid,old,rule


def save(e,pid,rule):
    result=e.put(f'/api/v1/projects/{pid}/rules',json=rule);assert result.status_code==200,result.text
    return result.json()


def test_math_bottom_up_no_double_count_and_local_scope_keys():
    items,rules,s=math_input();result=calculate(items,rules,s)
    assert result['subtotal']=='1575.15' and result['total']=='1739.76'
    assert [f['amount'] for f in result['fees']]==['73.51','40.43','50.67']
    assert [f['scope_kind'] for f in result['fees']]==['unit','single','project']
    assert result['fees'][1]['base_amount']=='808.58'
    assert result['fees'][2]['base_amount']=='1689.09'
    assert sum(Decimal(f['amount']) for f in result['fees'])+Decimal(result['subtotal'])==Decimal(result['total'])
    assert items==math_input()[0]
    # Reusing the same key in distinct scopes is intentional and must not merge fees.
    for group in rules['node_fees']:
        group['fees'][0]['key']='same'
    rules['fees'][0]['key']='same'
    assert calculate(items,rules,s)['total']=='1739.76'


def test_lower_fee_running_and_within_scope_fee_references():
    items,rules,s=math_input()
    rules['node_fees'][1]['fees']=[fee('下级附加','.1','lower_fees'),fee('本级附加','.5','fee:下级附加')]
    rules['fees']=[fee('项目累计','.01','running')]
    result=calculate(items,rules,s)
    assert [f['amount'] for f in result['fees']]==['73.51','7.35','3.68','16.60']
    assert result['total']=='1676.29'


@pytest.mark.parametrize('mutation',['missing_structure','missing_assignment','unknown_node','duplicate_group','missing_evidence','missing_rate','forward_reference','cross_scope','negative_rate','bad_rounding','cycle','duplicate_node'])
def test_invalid_scope_rules_are_not_silently_ignored(mutation):
    items,rules,s=math_input()
    if mutation=='missing_structure':s=None
    if mutation=='missing_assignment':s['assignments'].pop('b')
    if mutation=='unknown_node':rules['node_fees'][0]['node_id']='missing'
    if mutation=='duplicate_group':rules['node_fees'].append(copy.deepcopy(rules['node_fees'][0]))
    if mutation=='missing_evidence':rules['node_fees'][0]['fees'][0].pop('evidence_id')
    if mutation=='missing_rate':rules['node_fees'][0]['fees'][0].pop('rate')
    if mutation=='forward_reference':rules['node_fees'][0]['fees'][0]['basis']='fee:later'
    if mutation=='cross_scope':rules['node_fees'][1]['fees'][0]['basis']='fee:单位费'
    if mutation=='negative_rate':rules['node_fees'][0]['fees'][0]['rate']='-.1'
    if mutation=='bad_rounding':rules['rounding_config']={'mode':'unexpected'}
    if mutation=='cycle':s['nodes'][1]['parent_id']='u1'
    if mutation=='duplicate_node':s['nodes'].append(copy.deepcopy(s['nodes'][0]))
    with pytest.raises(DomainError):calculate(items,rules,s)


def test_scope_review_quote_export_and_immutable_legacy(clients):
    e,r,pid,old,rule=setup(clients);base=f'/api/v1/projects/{pid}'
    original_hash=get(e,pid)['rule_hash']
    saved=save(e,pid,rule)
    assert e.put(base+'/rules',json={**rule,'object_hash':original_hash}).status_code==409
    assert e.post(base+'/quotes/calculate').status_code==409
    approve(r,pid,'rules',saved);q=quote(e,pid)
    assert q['engine']=='pricing-3.0' and q['total']=='1739.76'
    summary=q['engineering_summary'];assert summary['scoped_fees'] and summary['project_fees']=='50.67'
    assert summary['descendant_fees']=='113.94'
    assert next(n for n in summary['nodes'] if n['id']=='s1')['total']=='849.01'
    assert next(x for x in get(e,pid)['quotes'] if x['id']==old['id'])==old
    for fmt in ['json','xlsx','pdf','zip']:
        exported=e.get(base+'/exports',params={'format':fmt,'quote_id':q['id']});assert exported.status_code==200
        if fmt=='xlsx':assert '分级取费规则' in load_workbook(io.BytesIO(exported.content)).sheetnames
    approve(r,pid,'rules',saved,decision='reject')
    assert e.post(base+'/quotes/calculate').status_code==409


def test_rebound_structure_required_and_stale_submission_rejected(clients):
    e,r,pid,old,rule=setup(clients);base=f'/api/v1/projects/{pid}'
    saved=save(e,pid,rule);approve(r,pid,'rules',saved)
    target=e.get(base+'/engineering-structure-targets').json()
    s=create_structure(e,pid,{**target,'name':'新结构版本','reason':'合成重核','evidence_id':get(e,pid)['evidence'][0]['id'],'request_id':str(uuid.uuid4())})
    approve(r,pid,'engineering_structures',s)
    assert e.post(base+'/quotes/calculate').status_code==409
    assert e.put(base+'/rules',json=rule).status_code==409
    rule['structure_hash']=get(e,pid)['engineering_structure_hash'];approved=save(e,pid,rule);approve(r,pid,'rules',approved)
    assert quote(e,pid)['total']=='1739.76'


@pytest.mark.parametrize('kind',['foreign','superseded','missing_bytes'])
def test_node_fee_evidence_is_checked_at_save_review_and_quote(clients,kind):
    from tests.test_workflow import upload
    e,r,pid,old,rule=setup(clients);base=f'/api/v1/projects/{pid}'
    upload(e,pid,'合成费用依据，非真实费率','node-fee.txt');p=get(e,pid);eid=p['evidence'][-1]['id']
    rule['node_fees'][0]['fees'][0]['evidence_id']='foreign' if kind=='foreign' else eid
    if kind=='foreign':assert e.put(base+'/rules',json=rule).status_code==404;return
    saved=save(e,pid,rule);approve(r,pid,'rules',saved)
    doc=next(d for d in p['documents'] if d['id']==p['evidence'][-1]['document_id'])
    if kind=='missing_bytes':(store.ROOT/'originals'/doc['hash']).unlink()
    else:
        from backend import main as core
        with store.transaction() as c:
            p=core.task_project(c,pid,{'tenant':'local'});next(d for d in p['documents'] if d['id']==doc['id'])['superseded']=True;store.save(c,p)
    assert e.post(base+'/quotes/calculate').status_code==409


def test_quantity_batch_preview_matches_new_scoped_quote(clients):
    e,r,pid,old,rule=setup(clients);base=f'/api/v1/projects/{pid}'
    saved=save(e,pid,rule);approve(r,pid,'rules',saved)
    p=get(e,pid);eid=p['evidence'][0]['id']
    body={'name':'合成分级费调量','reason':'仅数学合同变更测试','request_id':str(uuid.uuid4()),
          'entries':[{'boq_id':b['id'],'object_hash':store.digest(b),'quantity':str(Decimal(b['quantity'])+1),'basis':'合成逐项计量','evidence_id':eid} for b in p['boq']]}
    response=e.post(base+'/quantity-change-batches',json=body);assert response.status_code==200,response.text
    batch=response.json()['batch'];assert batch['impact']['available'] and batch['impact']['after_total']=='1972.84'
    p=get(e,pid)
    for cid in batch['change_ids']:approve(r,pid,'changes',next(c for c in p['changes'] if c['id']==cid))
    assert e.post(base+'/quantity-change-batches/'+batch['id']+'/apply').status_code==200
    for b in get(e,pid)['boq']:approve(r,pid,'boq',b)
    assert quote(e,pid)['total']==batch['impact']['after_total']


def test_material_replacement_preview_includes_all_scoped_fees(clients):
    e,r,pid,old,rule=setup(clients);base=f'/api/v1/projects/{pid}'
    saved=save(e,pid,rule);approve(r,pid,'rules',saved)
    row=next(v for v in e.get(base+'/material-substitution-targets').json() if v['boq_name']=='合成底板甲');resource=row['resource']
    after={**resource,'name':'合成替换材','resource_code':'TEST-REPLACE','specification':'合成规格','fee_allocation':'材料费','consumption':'1.2','price':'120',
           'pricing_status':'priced','source_type':'供应商报价','valid_from':'2026-01-01'}
    body={**{k:row[k] for k in ['key','boq_hash','resource_hash']},'name':'合成分级费材料替换','reason':'仅数学测试','source':'合成依据','license':'仅测试','evidence_id':resource['evidence_id'],
          'after':[after],'request_id':str(uuid.uuid4())}
    response=e.post(base+'/material-substitutions',json=body);assert response.status_code==200,response.text
    s=response.json()['substitution'];assert s['impact']['available'] and s['impact']['after_total']=='2064.45'
    approve(r,pid,'resource_substitutions',s);assert e.post(base+'/material-substitutions/'+s['id']+'/apply').status_code==200
    for b in get(e,pid)['boq']:approve(r,pid,'boq',b)
    assert quote(e,pid)['total']==s['impact']['after_total']


def test_sensitivity_uses_selected_snapshot_scopes_and_fee_rounding(clients):
    e,r,pid,old,rule=setup(clients);base=f'/api/v1/projects/{pid}'
    saved=save(e,pid,rule);approve(r,pid,'rules',saved);q=quote(e,pid)
    response=e.post(base+'/scenarios',json={'quote_id':q['id'],'material_factor':'1','quantity_factor':'1'})
    assert response.status_code==200,response.text
    assert response.json()['total']==q['total']
    assert response.json()['result']['scoped_pricing']['project']['own_fees']=='50.67'
    items,rules,s=math_input();rules['rounding_config']={'mode':'DOWN','unit_price_digits':2}
    assert calculate(items,rules,s)['total']=='1739.58'
