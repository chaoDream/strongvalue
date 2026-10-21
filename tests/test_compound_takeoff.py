"""Synthetic arithmetic and workflow checks; no professional measurement authority."""
import copy
import io
import json
from decimal import Decimal, localcontext
import pytest
from openpyxl import load_workbook
from backend import store
from backend.compound_takeoff import calculate
from backend.engine import DomainError, takeoff
from tests.test_workflow import create, get, upload, approve, priced_project


def payload(eid='test'):
    def part(name, operation, length, width, thickness, count):
        return {'name': name, 'operation': operation, 'kind': 'slab',
                'parameters': {'length': length, 'width': width, 'thickness': thickness},
                'multiplicity': count, 'measurement_basis': '合成数学依据，非专业规范',
                'evidence': {k: eid for k in ['length', 'width', 'thickness', 'multiplicity', 'measurement_basis']}}
    return {'name': '合成底板分项（非真实工程）', 'kind': 'compound', 'parameters': {},
            'basis_evidence_id': eid, 'measurement_basis': '已核对各实体不重叠；洞口按列示尺寸扣除，仅用于数学验收',
            'measurement_parts': [part('主底板', 'add', '10', '5', '.3', '1'),
                                  part('独立加厚区', 'add', '2', '1', '.2', '2'),
                                  part('穿板洞口', 'subtract', '1', '1', '.3', '2')]}


def setup(clients):
    e, r = clients['estimator'], clients['reviewer']
    pid = create(e)['id']
    upload(e, pid, '仅合成：底板10×5×0.3，加厚2×1×0.2共2处，洞口1×1×0.3共2处；加厚区不重叠。')
    eid = get(e, pid)['evidence'][0]['id']
    return e, r, pid, payload(eid)


def ast_value(ast, dimensions):
    if 'constant' in ast: return Decimal(ast['constant'])
    if 'dimension' in ast: return Decimal(dimensions[ast['dimension']])
    values = [ast_value(x, dimensions) for x in ast['args']]
    if ast['op'] == 'add': return sum(values, Decimal(0))
    result = values[0]
    for value in values[1:]:
        if ast['op'] == 'multiply': result *= value
        elif ast['op'] == 'subtract': result -= value
        else: result /= value
    return result


def test_arithmetic_breakdown_and_reproducible_ast():
    result = calculate(payload())
    assert Decimal(result['gross']) == Decimal('15.8')
    assert Decimal(result['deduction']) == Decimal('.6')
    assert Decimal(result['quantity']) == Decimal('15.2')
    with localcontext() as ctx:
        ctx.prec = 80
        assert ast_value(result['formula_ast'], result['dimensions']) == Decimal(result['quantity'])
    assert len(result['source_evidence']) == 16
    assert result['measurement_parts'][2]['quantity'] == '0.6'


@pytest.mark.parametrize('kind,parameters', [
    ('pipe', {'length':'3.25'}), ('area', {'length':'2','width':'3'}),
    ('masonry', {'length':'4','height':'2','thickness':'.2'}),
    ('walls', {'length':'10','width':'5','height':'3','thickness':'.25'}),
    ('trench', {'length':'10','bottom_width':'1','top_width':'2','depth':'2'}),
    ('backfill', {'excavation':'20','occupied':'3.5'}),
    ('ring', {'outer_diameter':'3','inner_diameter':'2','height':'4'}),
    ('rebar_schedule', {'length':'2','count':'10','unit_weight':'.617'}),
])
def test_each_supported_geometry_ast(kind, parameters):
    body = payload()
    body['measurement_parts'] = [{**body['measurement_parts'][0], 'kind':kind, 'parameters':parameters,
                                 'multiplicity':'3', 'evidence':{k:'test' for k in [*parameters,'multiplicity','measurement_basis']}}]
    result = calculate(body)
    with localcontext() as ctx:
        ctx.prec = 80
        assert Decimal(result['quantity']) == Decimal(takeoff(kind,{**parameters,'deduction':'0'})['quantity'])*3
        assert ast_value(result['formula_ast'], result['dimensions']) == Decimal(result['quantity'])


@pytest.mark.parametrize('case', ['negative','fraction_count','huge_count','missing_proof','mixed_units','unknown_dimension','all_deduct','over_deduct','duplicate','too_many','missing_basis','unknown_row'])
def test_reject_invalid_measurement(case):
    body = payload()
    row = body['measurement_parts'][0]
    if case == 'negative': row['parameters']['length'] = '-1'
    elif case == 'fraction_count': row['multiplicity'] = '1.5'
    elif case == 'huge_count': row['multiplicity'] = '1000001'
    elif case == 'missing_proof': del row['evidence']['multiplicity']
    elif case == 'mixed_units':
        row.update(kind='area', parameters={'length':'10','width':'5'}, evidence={k:'test' for k in ['length','width','multiplicity','measurement_basis']})
    elif case == 'unknown_dimension': row['parameters']['deduction'] = '0'
    elif case == 'all_deduct':
        for part in body['measurement_parts']: part['operation'] = 'subtract'
    elif case == 'over_deduct': body['measurement_parts'][2]['multiplicity'] = '100'
    elif case == 'duplicate': body['measurement_parts'][1]['name'] = row['name']
    elif case == 'too_many': body['measurement_parts'] *= 17
    elif case == 'missing_basis': body['basis_evidence_id'] = ''
    elif case == 'unknown_row': row['script'] = 'invalid'
    with pytest.raises(DomainError): calculate(body)


def test_preview_review_takeoff_and_boq(clients):
    e,r,pid,body = setup(clients)
    base = '/api/v1/projects/'+pid
    before = get(e,pid)
    preview = e.post(base+'/components/preview',json=body)
    assert preview.status_code == 200, preview.text
    assert get(e,pid) == before
    comp = e.post(base+'/components',json=body)
    assert comp.status_code == 200, comp.text
    comp = comp.json()
    assert Decimal(comp['preview']['quantity']) == Decimal('15.2')
    assert e.post(base+'/takeoffs/calculate').status_code == 409
    approve(r,pid,'components',comp)
    response = e.post(base+'/takeoffs/calculate')
    assert response.status_code == 200, response.text
    book = response.json()[0]
    assert book['measurement_parts'] == comp['preview']['measurement_parts']
    approve(r,pid,'takeoffs',book)
    boq = e.post(base+'/boq/generate').json()[0]
    assert boq['quantity'] == book['quantity']
    assert boq['unit'] == 'm³'


def test_revision_guard_atomicity_and_downstream_invalidation(clients):
    e,r,pid,body = setup(clients)
    base = '/api/v1/projects/'+pid
    comp = e.post(base+'/components',json=body).json()
    approve(r,pid,'components',comp)
    book = e.post(base+'/takeoffs/calculate').json()[0]
    approve(r,pid,'takeoffs',book)
    before = get(e,pid)
    assert e.put(base+'/components/'+comp['id'],json=body).status_code == 409
    body['object_hash'] = e.get(base+'/hashes').json()['components'][comp['id']]
    bad = copy.deepcopy(body);bad['measurement_parts'][0]['parameters']['length'] = '-2'
    assert e.put(base+'/components/'+comp['id'],json=bad).status_code == 422
    assert get(e,pid) == before
    body['measurement_parts'][0]['parameters']['length'] = '12'
    revised = e.put(base+'/components/'+comp['id'],json=body)
    assert revised.status_code == 200, revised.text
    assert revised.json()['revision'] == 2
    assert revised.json()['state'] == 'draft'
    assert get(e,pid)['takeoffs'][0]['state'] == 'stale'
    assert e.put(base+'/components/'+comp['id'],json=body).status_code == 409
    assert get(e,pid)['component_history'][0]['preview']['quantity'] == comp['preview']['quantity']


def test_permissions_foreign_evidence_and_steel_specialty(clients):
    e,r,pid,body = setup(clients)
    base = '/api/v1/projects/'+pid
    assert r.post(base+'/components',json=body).status_code == 403
    assert r.post(base+'/components/preview',json=body).status_code == 403
    assert clients['other'].post(base+'/components',json=body).status_code == 404
    pid2 = create(e)['id'];upload(e,pid2,'其他项目原件')
    body['measurement_parts'][0]['evidence']['length'] = get(e,pid2)['evidence'][0]['id']
    assert e.post(base+'/components',json=body).status_code == 404
    eid = body['basis_evidence_id']
    body['measurement_parts'] = [{'name':'钢筋','operation':'add','kind':'rebar_schedule','parameters':{'length':'2','count':'2','unit_weight':'.617'},'multiplicity':'1','measurement_basis':'测试','evidence':{k:eid for k in ['length','count','unit_weight','multiplicity','measurement_basis']}}]
    assert e.post(base+'/components',json=body).status_code == 409
    upload(e,pid,'合成钢筋资料',name='steel.txt',specialty='结构')
    eid2 = get(e,pid)['evidence'][-1]['id']
    body['measurement_parts'][0]['evidence'] = {k:eid2 for k in body['measurement_parts'][0]['evidence']}
    assert e.post(base+'/components',json=body).status_code == 200


@pytest.mark.parametrize('stage',['save','component_review','calculate','book_review'])
def test_missing_or_tampered_original_blocks_each_gate(clients, stage):
    e,r,pid,body = setup(clients)
    base = '/api/v1/projects/'+pid
    if stage != 'save': comp = e.post(base+'/components',json=body).json()
    if stage in ('calculate','book_review'): approve(r,pid,'components',comp)
    if stage == 'book_review': book = e.post(base+'/takeoffs/calculate').json()[0]
    doc = get(e,pid)['documents'][0]
    original = store.ROOT/'originals'/doc['hash']
    original.write_bytes(b'changed')
    if stage == 'save': assert e.post(base+'/components',json=body).status_code == 409
    elif stage == 'component_review': approve(r,pid,'components',comp,expected=409)
    elif stage == 'calculate': assert e.post(base+'/takeoffs/calculate').status_code == 409
    else: approve(r,pid,'takeoffs',book,expected=409)
    assert any('原件' in x for x in get(e,pid)['blockers']) if stage != 'save' else True


def test_superseded_evidence_and_snapshot_exports(clients):
    e,r,pid,body = setup(clients)
    base = '/api/v1/projects/'+pid
    comp = e.post(base+'/components',json=body).json()
    approve(r,pid,'components',comp)
    book = e.post(base+'/takeoffs/calculate').json()[0]
    approve(r,pid,'takeoffs',book)
    e.post(base+'/boq/generate')
    response = e.get(base+'/exports?format=xlsx&unpriced=true')
    assert response.status_code == 200, response.text
    wb = load_workbook(io.BytesIO(response.content),data_only=True)
    assert '分项算量与扣减明细' in wb.sheetnames
    assert wb['分项算量与扣减明细'].max_row >= 4
    assert '计价规则包未审定' in str(list(wb['缺项与警示'].values))
    assert '计价规则包未审定' not in str(list(wb['分项规则与精度'].values))
    assert '穿板洞口' in str(list(wb['分项算量与扣减明细'].values))
    data = e.get(base+'/exports?format=json&unpriced=true').json()
    assert 'compound-geometry-1.0' in json.dumps(data)
    pdf = e.get(base+'/exports?format=pdf&unpriced=true')
    assert pdf.status_code == 200, pdf.text
    from pypdf import PdfReader
    text = '\n'.join(page.extract_text() for page in PdfReader(io.BytesIO(pdf.content)).pages)
    assert '穿板洞口' in text and '分项算量与扣减明细' in text
    upload(e,pid,'改版尺寸',name='fixture.txt')
    assert get(e,pid)['components'][0]['state'] == 'stale'
    approve(r,pid,'components',comp,expected=409)
    assert e.post(base+'/components/preview',json=body).status_code == 409


def test_existing_quote_snapshot_survives_new_compound(clients):
    e,r,pid,_,quote = priced_project(clients)
    base = '/api/v1/projects/'+pid
    saved = copy.deepcopy(quote['snapshot'])
    eid = get(e,pid)['evidence'][0]['id']
    assert e.post(base+'/components',json=payload(eid)).status_code == 200
    p = get(e,pid)
    assert p['quotes'][0]['snapshot'] == saved
    assert quote['id'] in p['stale_quotes']
