"""Synthetic backfill ledger checks. Containment claims require independent professional review."""
import copy
import io
import json
import random
from decimal import Decimal, localcontext
import pytest
from fastapi import HTTPException
from openpyxl import load_workbook
from pypdf import PdfReader
from backend import store
from backend.engine import DomainError, takeoff
from backend.volume_balance import calculate, reference, snapshot
from tests.test_trench_takeoff import setup as trench_setup, geometry as route_geometry, payload as trench_payload, evaluate
from backend.trench_takeoff import calculate as trench_quantity
from tests.test_workflow import create, get, upload, approve, priced_project


def primitive(cid, kind, params, eid='test'):
    return {'id':cid,'revision':1,'name':'合成'+cid,'kind':kind,'parameters':{**params,'deduction':'0'},
            'measurement_basis':'仅合成已核定体积，非真实工程','evidence':{k:eid for k in list(params)+['deduction']},
            'preview':takeoff(kind,{**params,'deduction':'0'})}


def sources():
    return {'excavation':primitive('excavation','trench',{'length':'5','bottom_width':'2','top_width':'4','depth':'3'}),
            'occupations':[primitive('solid-a','slab',{'length':'2','width':'2','thickness':'1'}),primitive('solid-b','slab',{'length':'1','width':'2','thickness':'1'})]}


def payload(src,eid='test'):
    er=reference(src['excavation'])
    return {'name':'合成关联回填（非真实工程）','kind':'volume_balance','parameters':{},'quantity_type':'volume',
            'excavation_component_id':er['component_id'],'excavation_component_hash':er['component_hash'],
            'occupations':[{'component_id':s['id'],'component_hash':reference(s)['component_hash'],'factor':'1' if i==0 else '.5',
                            'basis':'仅合成：已核定同域不重复，按明确比例扣除，无真实包含关系认定','evidence_id':eid} for i,s in enumerate(src['occupations'])],
            'boundaries_confirmed':True,'measurement_basis':'仅合成：同一计量范围，扣减不重复，来源既定体积及重复数量不再调整，非专业验收','basis_evidence_id':eid,
            'rounding_digits':3,'rounding_evidence_id':eid}


def setup(clients):
    e,r,pid,route,trench,eid=trench_setup(clients);base='/api/v1/projects/'+pid
    response=e.post(base+'/components',json=trench);assert response.status_code==200,response.text
    trench_id=response.json()['id']
    ids=[]
    for cid,length in [('solid-a','2'),('solid-b','1')]:
        s=primitive(cid,'slab',{'length':length,'width':'2','thickness':'1'},eid)
        for k in ['id','revision','preview']:s.pop(k)
        response=e.post(base+'/components',json=s);assert response.status_code==200,response.text
        ids.append(response.json()['id'])
    upload(e,pid,'仅合成回填平衡：已核定同一范围且不重叠；开挖47.500m³，实体A为4m³扣1，实体B为2m³扣0.5，回填42.500m³。来源量含既定重复数量与修约，不再乘数量。净回填总量修约3位；无真实专业效力。',name='balance.txt')
    for comp in get(e,pid)['components']:approve(r,pid,'components',comp)
    p=get(e,pid);src={'excavation':snapshot(next(c for c in p['components'] if c['id']==trench_id)),
                     'occupations':[snapshot(next(c for c in p['components'] if c['id']==cid)) for cid in ids]}
    beid=p['evidence'][-1]['id']
    return e,r,pid,src,payload(src,beid),beid


def test_source_quantities_fractional_deductions_and_ast():
    src=sources();b=payload(src);before=copy.deepcopy((src,b));r=calculate(b,src)
    assert (src,b)==before and r['quantity']=='40.000' and Decimal(r['gross'])==45 and Decimal(r['deduction'])==5
    assert [Decimal(v['deduction']) for v in r['balance_deductions']]==[4,1]
    assert r['dimensions']['excavation.quantity']=='45' and r['dimensions']['occupied.2.factor']=='.5'
    with localcontext() as ctx:
        ctx.prec=80
        assert evaluate(r['formula_ast'],r['dimensions'])==40
        for d in r['balance_deductions']:assert evaluate(d['formula_ast'],r['dimensions'])==Decimal(d['deduction'])


def test_source_existing_rounding_is_respected_not_reversed():
    src=sources();route=route_geometry();trench=trench_payload(route);trench.update(id='rounded-trench',revision=1,source_geometry=route,rounding_digits=0)
    trench['preview']=trench_quantity(trench);trench['evidence']=trench['preview']['source_evidence'];src['excavation']=trench
    b=payload(src);r=calculate(b,src)
    assert trench['preview']['gross']=='47.5' and r['gross']=='48' and r['quantity']=='43.000'


def test_no_deductions_full_deduction_and_rounding_after_sum():
    src=sources();src['occupations']=[];b=payload(src);assert calculate(b,src)['quantity']=='45.000'
    src=sources();src['excavation']['parameters']['length']='0.555555555555555556';src['excavation']['preview']=takeoff('trench',src['excavation']['parameters'])
    b=payload(src);b['rounding_digits']=2;r=calculate(b,src)
    assert r['quantity']=='0.00' and Decimal(r['balance_summary']['unrounded_quantity'])>0
    src=sources();src['excavation']=primitive('ex','trench',{'length':'5','bottom_width':'1','top_width':'1','depth':'1'});b=payload(src)
    assert calculate(b,src)['quantity']=='0.000'


def test_seeded_balances_match_independent_decimal_ledger():
    rng=random.Random(3101)
    for _ in range(30):
        src=sources();src['occupations']=[]
        for i in range(5):src['occupations'].append(primitive(str(i),'slab',{'length':str(rng.randint(1,5)),'width':'1','thickness':'1'}))
        b=payload(src);expected=Decimal(45)
        for row,s in zip(b['occupations'],src['occupations']):
            row['factor']=str(Decimal(rng.randint(1,10))/10)
            expected-=Decimal(s['parameters']['length'])*Decimal(row['factor'])
        r=calculate(b,src);assert Decimal(r['quantity'])==expected
        with localcontext() as ctx:
            ctx.prec=80
            assert evaluate(r['formula_ast'],r['dimensions'])==expected


@pytest.mark.parametrize('case',['unconfirmed','string_confirmation','missing','duplicate','self','unknown','extra','missing_factor','factor_zero','factor_negative','factor_over','nan','precision','blank_basis','blank_evidence','missing_rounding','rounding_bool','mode','parameters','too_many','over_deduct'])
def test_invalid_balance_inputs_fail(case):
    src=sources();b=payload(src);d=b['occupations'][0]
    if case=='unconfirmed':b['boundaries_confirmed']=False
    elif case=='string_confirmation':b['boundaries_confirmed']='true'
    elif case=='missing':b['occupations'].pop()
    elif case=='duplicate':b['occupations'][1]=copy.deepcopy(d)
    elif case=='self':d['component_id']=b['excavation_component_id']
    elif case=='unknown':d['component_id']='unknown'
    elif case=='extra':d['quantity']='999'
    elif case=='missing_factor':del d['factor']
    elif case=='factor_zero':d['factor']='0'
    elif case=='factor_negative':d['factor']='-1'
    elif case=='factor_over':d['factor']='1.1'
    elif case=='nan':d['factor']='NaN'
    elif case=='precision':d['factor']='.1234567890123456789'
    elif case=='blank_basis':d['basis']=''
    elif case=='blank_evidence':d['evidence_id']=''
    elif case=='missing_rounding':del b['rounding_digits']
    elif case=='rounding_bool':b['rounding_digits']=True
    elif case=='mode':b['quantity_type']='area'
    elif case=='parameters':b['parameters']={'volume':'1'}
    elif case=='too_many':b['occupations']*=26;src['occupations']*=26
    elif case=='over_deduct':
        src['excavation']=primitive('small','trench',{'length':'1','bottom_width':'1','top_width':'1','depth':'1'});b=payload(src)
    with pytest.raises(DomainError):calculate(b,src)


@pytest.mark.parametrize('target',['excavation','occupation'])
def test_old_source_version_hash_rejected(target):
    src=sources();b=payload(src)
    s=src['excavation'] if target=='excavation' else src['occupations'][0];s['revision']=2
    with pytest.raises(HTTPException) as error:calculate(b,src)
    assert error.value.status_code==409


def test_candidates_preview_review_books_boq_and_exports(clients):
    e,r,pid,src,b,eid=setup(clients);base='/api/v1/projects/'+pid;before=get(e,pid)
    selector={'excavation_component_id':b['excavation_component_id'],'occupation_component_ids':[v['component_id'] for v in b['occupations']]}
    c=e.post(base+'/volume-balance-candidates',json=selector);assert c.status_code==200,c.text
    assert c.json()['excavation']['quantity']=='47.500' and get(e,pid)==before
    response=e.post(base+'/components/preview',json=b);assert response.status_code==200,response.text
    assert response.json()['quantity']=='42.500' and get(e,pid)==before
    response=e.post(base+'/components',json={**b,'source_geometry':{'forged':True}});assert response.status_code==200,response.text
    comp=response.json();assert comp['source_geometry']==src
    assert all(s['state']=='approved' for s in get(e,pid)['components'][:-1])
    assert e.post(base+'/takeoffs/calculate').status_code==409
    approve(r,pid,'components',comp)
    response=e.post(base+'/takeoffs/calculate');assert response.status_code==200,response.text
    for book in response.json():approve(r,pid,'takeoffs',book)
    items=e.post(base+'/boq/generate').json();assert next(x for x in items if x['name']==b['name'])['quantity']=='42.500'
    j=e.get(base+'/exports?unpriced=true&format=json').json()['unpriced']['takeoffs'][-1]
    assert j['balance_deductions']==comp['preview']['balance_deductions'] and j['source_geometry']==src
    response=e.get(base+'/exports?unpriced=true&format=xlsx');assert response.status_code==200,response.text
    wb=load_workbook(io.BytesIO(response.content));assert {'关联回填平衡','回填逐项扣减','回填来源计算索引','回填范围与修约'}<=set(wb.sheetnames)
    assert all(c.data_type!='f' for s in wb for row in s for c in row)
    response=e.get(base+'/exports?unpriced=true&format=pdf');assert response.status_code==200,response.text
    text='\n'.join(p.extract_text() for p in PdfReader(io.BytesIO(response.content)).pages)
    assert '关联回填平衡' in text and '回填逐项扣减' in text and '42.500' in text


def test_concurrent_edit_atomic_failure_and_source_approvals_preserved(clients):
    e,r,pid,src,b,eid=setup(clients);base='/api/v1/projects/'+pid
    comp=e.post(base+'/components',json=b).json();approve(r,pid,'components',comp)
    books=e.post(base+'/takeoffs/calculate').json()
    for book in books:approve(r,pid,'takeoffs',book)
    before=get(e,pid);assert e.put(base+'/components/'+comp['id'],json=b).status_code==409
    b['object_hash']=e.get(base+'/hashes').json()['components'][comp['id']]
    bad=copy.deepcopy(b);bad['occupations'][0]['factor']='2'
    assert e.put(base+'/components/'+comp['id'],json=bad).status_code==422 and get(e,pid)==before
    b['occupations'][1]['factor']='1';response=e.put(base+'/components/'+comp['id'],json=b);assert response.status_code==200,response.text
    assert response.json()['revision']==2 and response.json()['preview']['quantity']=='41.500'
    p=get(e,pid);assert all(s['state']=='approved' for s in p['components'][:-1]) and all(t['state']=='stale' for t in p['takeoffs'])
    assert e.put(base+'/components/'+comp['id'],json=b).status_code==409


@pytest.mark.parametrize('gate',['preview','save','component_review','calculate','book_review'])
def test_allocation_original_corruption_blocks_processing(clients,gate):
    e,r,pid,src,b,eid=setup(clients);base='/api/v1/projects/'+pid
    if gate not in ('preview','save'):comp=e.post(base+'/components',json=b).json()
    if gate in ('calculate','book_review'):approve(r,pid,'components',comp)
    if gate=='book_review':book=e.post(base+'/takeoffs/calculate').json()[-1]
    doc=get(e,pid)['documents'][-1];(store.ROOT/'originals'/doc['hash']).write_bytes(b'corrupt')
    selector={'excavation_component_id':b['excavation_component_id'],'occupation_component_ids':[v['component_id'] for v in b['occupations']]}
    assert e.post(base+'/volume-balance-candidates',json=selector).status_code==200
    if gate=='preview':assert e.post(base+'/components/preview',json=b).status_code==409
    elif gate=='save':assert e.post(base+'/components',json=b).status_code==409
    elif gate=='component_review':approve(r,pid,'components',comp,expected=409)
    elif gate=='calculate':assert e.post(base+'/takeoffs/calculate').status_code==409
    else:approve(r,pid,'takeoffs',book,expected=409)


def test_source_change_unapproval_and_forged_preview_rejected(clients):
    e,r,pid,src,b,eid=setup(clients);base='/api/v1/projects/'+pid
    comp=e.post(base+'/components',json=b).json();approve(r,pid,'components',comp)
    source=next(c for c in get(e,pid)['components'] if c['id']==b['occupations'][0]['component_id'])
    source['object_hash']=e.get(base+'/hashes').json()['components'][source['id']];source['parameters']['length']='3'
    response=e.put(base+'/components/'+source['id'],json=source);assert response.status_code==200,response.text
    assert e.post(base+'/components/preview',json=b).status_code==409
    for c in get(e,pid)['components'][:-1]:approve(r,pid,'components',c)
    assert e.post(base+'/components/preview',json=b).status_code==409
    approve(r,pid,'components',get(e,pid)['components'][-1],expected=409)
    # Current source preview is independently recomputed even when a client knows its hash.
    with store.transaction() as db:
        p=json.loads(db.execute('SELECT data FROM projects WHERE id=?',(pid,)).fetchone()['data']);next(c for c in p['components'] if c['id']==source['id'])['preview']['quantity']='999';store.save(db,p)
    selector={'excavation_component_id':b['excavation_component_id'],'occupation_component_ids':[source['id']]}
    assert e.post(base+'/volume-balance-candidates',json=selector).status_code==409


def test_permissions_cross_project_wrong_unit_duplicate_and_recursive_sources(clients):
    e,r,pid,src,b,eid=setup(clients);base='/api/v1/projects/'+pid
    selector={'excavation_component_id':b['excavation_component_id'],'occupation_component_ids':[v['component_id'] for v in b['occupations']]}
    assert r.post(base+'/volume-balance-candidates',json=selector).status_code==200
    assert r.post(base+'/components',json=b).status_code==403
    assert clients['other'].post(base+'/volume-balance-candidates',json=selector).status_code==404
    assert e.post(base+'/volume-balance-candidates',json={**selector,'occupation_component_ids':[selector['occupation_component_ids'][0]]*2}).status_code==422
    route=get(e,pid)['components'][0]
    assert e.post(base+'/volume-balance-candidates',json={**selector,'occupation_component_ids':[route['id']]}).status_code==409
    pid2=create(e)['id'];upload(e,pid2,'其他项目');bad=copy.deepcopy(b);bad['occupations'][0]['evidence_id']=get(e,pid2)['evidence'][0]['id']
    assert e.post(base+'/components',json=bad).status_code==404
    comp=e.post(base+'/components',json=b).json();approve(r,pid,'components',comp)
    assert e.post(base+'/volume-balance-candidates',json={**selector,'occupation_component_ids':[comp['id']]}).status_code==409
    changed={'kind':'backfill','parameters':{},'object_hash':e.get(base+'/hashes').json()['components'][comp['id']]}
    assert e.put(base+'/components/'+comp['id'],json=changed).status_code==409


@pytest.mark.parametrize('target',['component','book'])
def test_saved_balance_tampering_cannot_pass_review(clients,target):
    e,r,pid,src,b,eid=setup(clients);base='/api/v1/projects/'+pid;comp=e.post(base+'/components',json=b).json()
    if target=='book':approve(r,pid,'components',comp);e.post(base+'/takeoffs/calculate')
    with store.transaction() as db:
        p=json.loads(db.execute('SELECT data FROM projects WHERE id=?',(pid,)).fetchone()['data'])
        if target=='component':p['components'][-1]['preview']['balance_deductions'][0]['deduction']='999'
        else:p['takeoffs'][-1]['balance_deductions'][0]['deduction']='999'
        store.save(db,p)
    col='components' if target=='component' else 'takeoffs';approve(r,pid,col,get(e,pid)[col][-1],expected=409)


def test_new_balance_never_rewrites_historical_quote(clients):
    e,r,pid,_,q=priced_project(clients);old=copy.deepcopy(q);base='/api/v1/projects/'+pid;eid=get(e,pid)['evidence'][0]['id']
    s=primitive('ex','trench',{'length':'5','bottom_width':'2','top_width':'4','depth':'3'},eid)
    for k in ['id','revision','preview']:s.pop(k)
    response=e.post(base+'/components',json=s);assert response.status_code==200,response.text;approve(r,pid,'components',response.json())
    src={'excavation':snapshot(get(e,pid)['components'][-1]),'occupations':[]};response=e.post(base+'/components',json=payload(src,eid));assert response.status_code==200,response.text
    assert get(e,pid)['quotes'][0]==old and q['id'] in get(e,pid)['stale_quotes']
