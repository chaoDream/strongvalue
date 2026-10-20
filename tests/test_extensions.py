import copy
import io
from decimal import Decimal

import pytest
from reportlab.pdfgen import canvas
from backend.engine import takeoff, calculate, DomainError
from backend.parsers import dimension_candidates, dxf_objects
from backend import gateway
from tests.test_workflow import create, get, upload, approve, priced_project, setup_boq


@pytest.mark.parametrize('kind,params,expected', [
    ('area', {'length': '10', 'width': '5'}, '50'),
    ('masonry', {'length': '10', 'height': '3', 'thickness': '.24', 'deduction': '.2'}, '7.00'),
    ('trench', {'length': '10', 'bottom_width': '1', 'top_width': '3', 'depth': '2'}, '40'),
    ('backfill', {'excavation': '40', 'occupied': '5', 'deduction': '1'}, '34'),
    ('rebar_schedule', {'length': '6', 'count': '12', 'unit_weight': '1.58'}, '113.76'),
])
def test_extended_geometry(kind, params, expected):
    assert Decimal(takeoff(kind, params)['quantity']) == Decimal(expected)


def test_ring_and_professional_boundaries():
    r = takeoff('ring', {'outer_diameter': '2', 'inner_diameter': '1', 'height': '3'})
    assert abs(Decimal(r['quantity'])-Decimal('7.068583470577034')) < Decimal('1e-14')
    result = calculate([{'quantity': r['quantity'], 'quotas': [{'code': 'test', 'resources': [{'consumption': '1', 'price': '100'}]}]}], {'fees': []})
    assert result['total'] == '706.86'
    with pytest.raises(DomainError): takeoff('rebar_schedule', {'length': '6', 'count': '1.5', 'unit_weight': '1.58'})
    with pytest.raises(DomainError): takeoff('backfill', {'excavation': '1', 'occupied': '2'})
    with pytest.raises(DomainError): takeoff('trench', {'length': '1', 'bottom_width': '3', 'top_width': '1', 'depth': '2'})


def test_rebar_requires_structural_evidence(clients):
    c=clients['estimator'];pid=create(c)['id'];upload(c,pid)
    eid=get(c,pid)['evidence'][0]['id']
    body={'kind':'rebar_schedule','name':'假设钢筋表', 'parameters':{'length':'6','count':'12','unit_weight':'1.58','deduction':'0'},
          'evidence':{k:eid for k in ['length','count','unit_weight','deduction']},'measurement_basis':'测试'}
    assert c.post(f'/api/v1/projects/{pid}/components',json=body).status_code==409
    upload(c,pid,'结构钢筋表：6米、12根、1.58kg/m，无扣减',name='steel.txt',specialty='结构')
    eid=get(c,pid)['evidence'][-1]['id'];body['evidence']={k:eid for k in body['parameters']}
    assert c.post(f'/api/v1/projects/{pid}/components',json=body).status_code==200


def test_fee_expression_rounding_and_references():
    items=[{'quantity':'3','quotas':[{'code':'t','resources':[{'type':'人工','consumption':'1','price':'1.005'},{'type':'机械','consumption':'2','price':'2'}]}]}]
    rules={'rounding_config':{'unit_price_digits':3,'mode':'HALF_UP'},'fees':[
        {'name':'管理','rate':'.1','basis':{'op':'add','args':['labor','machinery']}},
        {'name':'利润','rate':'.1','basis':'fee:管理'}]}
    q=calculate(items,rules)
    assert q['lines'][0]['unit_price']=='5.005'
    assert q['subtotal']=='15.02' and q['fees'][0]['amount']=='1.50' and q['total']=='16.67'
    for expr in ['fee:future',{'op':'eval','args':['labor','machinery']},{'op':'divide','args':['labor',{'constant':'0'}]}]:
        with pytest.raises(DomainError): calculate(items,{'fees':[{'name':'bad','rate':'.1','basis':expr}]})


def test_import_preview_commit_preserves_lineage_and_is_idempotent(clients):
    e,r=clients['estimator'],clients['reviewer'];pid=create(e,['合同计价'])['id'];base=f'/api/v1/projects/{pid}'
    raw='清单编码,项目名称,单位,工程量,项目特征,工程层级,综合单价\n001,底板,m³,15,测试清单,A / 池,100.01\n'
    res=e.post(base+'/boq/import-preview',files={'file':('fixture.csv',raw.encode())})
    assert res.status_code==200,res.text
    preview=res.json();assert preview['records'][0]['code']=='001'
    assert get(e,pid)['boq']==[]
    res=e.post(base+'/boq/imports/'+preview['id']+'/commit',json={});assert res.status_code==200
    assert e.post(base+'/boq/imports/'+preview['id']+'/commit',json={}).json()['duplicate']
    p=get(e,pid);b=p['boq'][0];t=p['takeoffs'][0]
    assert b['reference_unit_price']=='100.01' and b['quotas']==[]
    approve(r,pid,'boq',b,expected=409)
    approve(r,pid,'takeoffs',t);approve(r,pid,'boq',b)
    assert p['evidence'][0]['row']==2
    assert e.post(base+'/quotes/calculate').status_code==409
    assert clients['other'].post(base+'/boq/imports/'+preview['id']+'/commit',json={}).status_code==404


def test_invalid_import_and_duplicate_acknowledgment(clients):
    e=clients['estimator'];pid=create(e)['id'];base=f'/api/v1/projects/{pid}'
    raw='清单编码,项目名称,单位,工程量\nA,test,m,-1\n'
    p=e.post(base+'/boq/import-preview',files={'file':('bad.csv',raw.encode())}).json()
    assert p['errors'] and e.post(base+'/boq/imports/'+p['id']+'/commit',json={}).status_code==409
    raw='清单编码,项目名称,单位,工程量\nA,test,m,1\nA,test2,m,2\n'
    p=e.post(base+'/boq/import-preview',files={'file':('duplicate.csv',raw.encode())}).json()
    assert p['warnings']
    assert e.post(base+'/boq/imports/'+p['id']+'/commit',json={}).status_code==409
    assert e.post(base+'/boq/imports/'+p['id']+'/commit',json={'acknowledge_duplicates':True}).status_code==200


def test_split_merge_conservation_and_no_repeat_generation(clients):
    e,r,pid,comp,t,b,eid=setup_boq(clients);base=f'/api/v1/projects/{pid}'
    assert e.post(base+'/boq/'+b['id']+'/split',json={'parts':[{'quantity':'8'},{'quantity':'8'}]}).status_code==422
    parts=e.post(base+'/boq/'+b['id']+'/split',json={'parts':[{'quantity':'8'},{'quantity':'7'}]}).json()
    assert sum(Decimal(x['quantity']) for x in parts)==15
    merged=e.post(base+'/boq/merge',json={'ids':[x['id'] for x in parts]}).json()
    assert len(merged['quantity_links'])==2 and Decimal(merged['quantity'])==15
    assert len(e.post(base+'/boq/generate').json())==1
    assert len(get(e,pid)['boq_history'])==3


def test_change_approval_apply_compare_and_scenario_independence(clients):
    e,r,pid,comp,q=priced_project(clients);base=f'/api/v1/projects/{pid}';p=get(e,pid);b=p['boq'][0];eid=p['evidence'][0]['id']
    approve(r,pid,'quotes',q)
    original=copy.deepcopy(get(e,pid))
    res=e.post(base+'/scenarios',json={'quote_id':q['id'],'name':'测试测算','material_factor':'1.1'});assert res.status_code==200,res.text
    after=get(e,pid);assert after['quotes']==original['quotes'] and after['report_approval']==original['report_approval']
    assert Decimal(res.json()['total'])>Decimal(q['total'])
    ch=e.post(base+'/changes',json={'boq_id':b['id'],'quantity':'16','evidence_id':eid,'reason':'合成变更依据'}).json()
    assert e.post(base+'/changes/'+ch['id']+'/apply').status_code==409
    approve(r,pid,'quotes',q,expected=409)
    approve(r,pid,'changes',ch)
    applied=e.post(base+'/changes/'+ch['id']+'/apply');assert applied.status_code==200,applied.text
    assert e.post(base+'/changes/'+ch['id']+'/apply').json()['duplicate']
    assert e.post(base+'/boq/generate').status_code==409
    approve(r,pid,'boq',applied.json())
    q2=e.post(base+'/quotes/calculate').json()
    diff=e.get(base+f'/quotes/compare?left={q["id"]}&right={q2["id"]}').json()
    assert len(diff['lines'])==1 and diff['lines'][0]['before']['quantity']=='15.0'
    assert diff['lines'][0]['after']['quantity']=='16'
    assert Decimal(diff['total_delta'])==Decimal(q2['total'])-Decimal(q['total'])
    assert get(e,pid)['quotes'][0]==q
    raw=e.get(base+'/exports?format=xlsx&quote_id='+q2['id'])
    from openpyxl import load_workbook
    assert '变更签证台账' in load_workbook(io.BytesIO(raw.content)).sheetnames


def test_catalog_filters_and_review(clients):
    e,r,pid,comp,q=priced_project(clients);base=f'/api/v1/projects/{pid}';b=get(e,pid)['boq'][0]
    item=e.post(base+'/catalog',json={'boq_id':b['id']}).json()
    assert e.get(base+'/boq/'+b['id']+'/quota-candidates').json()==[]
    approve(r,pid,'catalog',item)
    candidates=e.get(base+'/boq/'+b['id']+'/quota-candidates').json()
    assert len(candidates)==1 and candidates[0]['quotas']==b['quotas']
    assert clients['other'].get(base+'/boq/'+b['id']+'/quota-candidates').status_code==404


def test_explicit_dimension_units_and_cross_drawing_conflict(clients):
    dims=dimension_candidates('构件：底板\n长度10000mm，宽度500cm，厚度0.3m')
    assert [Decimal(d['value']) for d in dims]==[Decimal(10),Decimal(5),Decimal('.3')]
    e,r=clients['estimator'],clients['reviewer'];pid=create(e)['id'];base=f'/api/v1/projects/{pid}'
    upload(e,pid,'构件：底板\n长度10m',name='plan.txt');upload(e,pid,'构件：底板\n长度11m',name='section.txt')
    p=get(e,pid);a,b=p['dimension_candidates'];assert a['state']==b['state']=='conflict'
    approve(r,pid,'dimension_candidates',a,expected=409)
    approve(r,pid,'dimension_candidates',b,'reject');approve(r,pid,'dimension_candidates',a)


def test_pdf_anchors_and_ascii_dxf_are_not_pixel_quantities(clients):
    c=clients['estimator'];pid=create(c)['id'];base=f'/api/v1/projects/{pid}'
    out=io.BytesIO();pdf=canvas.Canvas(out);pdf.drawString(100,700,'L=10m');pdf.save()
    res=c.post(base+'/documents',data={'specialty':'建筑'},files={'file':('drawing.pdf',out.getvalue())});assert res.status_code==200
    p=get(c,pid);objects=p['documents'][0]['vector_objects'];assert objects and objects[0]['method']=='pdf_text_anchor'
    assert p['takeoffs']==[] and p['dimension_candidates'][0]['state']=='draft'
    raw=b'0\nSECTION\n9\n$INSUNITS\n70\n4\n0\nTEXT\n8\nDIM\n10\n10\n20\n20\n1\nL=10000mm\n0\nDIMENSION\n42\n10000\n0\nENDSEC\n0\nEOF\n'
    objs,units=dxf_objects(raw);assert units=='mm' and objs[-1]['measured_value']=='10000'


def test_historical_pricing_date_invalidates_approval(clients):
    e,r,pid,comp,q=priced_project(clients);base=f'/api/v1/projects/{pid}';approve(r,pid,'quotes',q)
    res=e.patch(base+'/settings',json={'pricing_date':'2000-01-01','reason':'测试历史日期'})
    assert res.status_code==200 and res.json()['report_approval'] is None
    assert any('有效期' in text for text in res.json()['blockers'])
    assert e.post(base+'/quotes/calculate').status_code==409


def test_gateway_requires_consent_and_rejects_ungrounded_candidates(clients,monkeypatch):
    e=clients['estimator'];pid=create(e)['id'];upload(e,pid,'规模500m³/日');p=get(e,pid);did=p['documents'][0]['id'];base=f'/api/v1/projects/{pid}'
    assert e.post(base+'/documents/'+did+'/analyze',json={'consent':True}).status_code==409
    monkeypatch.setenv('QIANG_AI_GATEWAY_URL','http://localhost:9999/extract');monkeypatch.setenv('QIANG_AI_MODEL','test-model')
    calls=[]
    def mocked(doc,raw,evidence):
        calls.append({'raw':raw,'methods':[e['method'] for e in evidence]})
        return {'model':'test-model','warnings':[], 'candidates':[
            {'kind':'fact','key':'capacity','label':'规模','value':'500','unit':'m³/日','component_name':'','page':1,'quote':'规模500m³/日','bbox':None},
            {'kind':'dimension','key':'length','label':'长度','value':'999','unit':'m','component_name':'虚构','page':1,'quote':'不存在的来源','bbox':None}]}
    monkeypatch.setattr(gateway,'request_gateway',mocked)
    url=base+'/documents/'+did+'/analyze'
    assert e.post(url,json={'consent':False,'destination':'localhost:9999'}).status_code==422
    assert calls==[]
    result=e.post(url,json={'consent':True,'destination':'localhost:9999'});assert result.status_code==200
    p=get(e,pid);assert len(p['analyses'])==1 and len(p['analyses'][0]['rejected_candidates'])==1
    assert all(f['state']=='draft' for f in p['facts'])
    assert len(calls)==1
    assert e.post(url,json={'consent':True,'destination':'localhost:9999'}).json()['duplicate']
    assert len(calls)==1
    e.post(base+'/evidence',json={'document_id':did,'page':1,'text':'人工补充的来源标注'})
    assert not e.post(url,json={'consent':True,'destination':'localhost:9999'}).json()['duplicate']
    assert len(calls)==2 and all(not method.startswith('model_') for call in calls for method in call['methods'])


def test_preview_returns_actual_immutable_pdf_inline(clients):
    c=clients['estimator'];pid=create(c)['id'];base=f'/api/v1/projects/{pid}'
    raw=c.get(base+'/exports?format=pdf').content
    artifact=c.get(base+'/artifacts').json()[0]
    response=c.get(base+'/artifacts/'+artifact['id']+'/file?inline=true')
    assert response.content==raw and response.headers['content-disposition'].startswith('inline')


def test_responses_provider_payload_metadata_and_no_automatic_tools(monkeypatch):
    import json
    import httpx
    monkeypatch.setenv('QIANG_AI_PROTOCOL','openai_responses')
    monkeypatch.setenv('QIANG_AI_GATEWAY_URL','https://model.example.invalid/v1/responses')
    monkeypatch.setenv('QIANG_AI_MODEL','test-model')
    monkeypatch.setenv('QIANG_AI_GATEWAY_TOKEN','test-only-not-real')
    original_client=httpx.Client
    requests=[]
    candidate={'kind':'fact','key':'capacity','label':'规模','value':'500','unit':'m³/日','component_name':'','page':1,'quote':'规模500m³/日','bbox':None}
    def handler(request):
        body=json.loads(request.content);requests.append(body)
        assert body['store'] is False and body['input'][0]['content'][0]['file_data'].startswith('data:application/pdf;base64,')
        assert body['text']['format']['strict'] is True and 'tools' not in body
        schema=body['text']['format']['schema']
        assert set(schema['required'])==set(schema['properties'])
        assert set(schema['$defs']['Candidate']['required'])==set(schema['$defs']['Candidate']['properties'])
        return httpx.Response(200,json={'id':'test-response','model':'test-model-snapshot','status':'completed','output':[{'type':'message','content':[{'type':'output_text','text':json.dumps({'model':'invented-self-report','candidates':[candidate],'warnings':[]})}]}]})
    monkeypatch.setattr(gateway.httpx,'Client',lambda **kwargs:original_client(transport=httpx.MockTransport(handler),**kwargs))
    result=gateway.request_gateway({'name':'drawing.pdf','hash':'testhash'},b'%PDF-test',[])
    assert result['model']=='test-model' and result['provider_model']=='test-model-snapshot' and len(requests)==1
    for body in [{'status':'incomplete'},{'status':'completed','output':[{'content':[{'type':'refusal','refusal':'test'}]}]}]:
        with pytest.raises(DomainError): gateway.unpack_response(body,'test-model')


def test_member_admin_boundary_tenant_scope_and_session_revocation(clients):
    admin=clients['admin'];other=clients['other'];reviewer=clients['reviewer']
    body={'username':'new-reviewer','password':'A-Test-Only-Password-2026','role':'reviewer','tenant':'other'}
    assert reviewer.post('/api/v1/organization/members',json=body).status_code==403
    result=admin.post('/api/v1/organization/members',json=body)
    assert result.status_code==200 and 'password' not in result.json()
    member=result.json()
    assert not any(m['username']=='new-reviewer' for m in other.get('/api/v1/organization/members').json())
    assert other.patch('/api/v1/organization/members/'+member['id'],json={'role':'admin','active':True,'reason':'cross tenant'}).status_code==404
    from fastapi.testclient import TestClient
    from backend.main import app
    login=TestClient(app)
    assert login.post('/api/v1/auth/login',json={'username':body['username'],'password':body['password']}).status_code==200
    assert login.get('/api/v1/auth/me').json()['tenant']=='local'
    res=admin.patch('/api/v1/organization/members/'+member['id'],json={'role':'reviewer','active':False,'reason':'测试停用'})
    assert res.status_code==200 and login.get('/api/v1/auth/me').status_code==401
    assert login.post('/api/v1/auth/login',json={'username':body['username'],'password':body['password']}).status_code==401
    own=next(m for m in admin.get('/api/v1/organization/members').json() if m['username']=='admin')
    assert admin.patch('/api/v1/organization/members/'+own['id'],json={'role':'reviewer','active':False,'reason':'self'}).status_code==409
    events=admin.get('/api/v1/organization/audit').json()
    assert len(events)==2 and not any('password' in str(e['detail']) for e in events)


def test_golden_reconciliation_gate_and_stale_after_recalculation(clients):
    e,r,pid,comp,q=priced_project(clients);base=f'/api/v1/projects/{pid}'
    e.patch(base+'/settings',json={'require_golden_baseline':True,'reason':'启用数学测试对账门禁'})
    assert any('黄金底稿' in message for message in get(e,pid)['blockers'])
    raw='构件名称,单位,基准工程量,绝对容差,相对容差\n测试底板,m³,15,0,0\n'
    res=e.post(base+'/takeoffs/reconcile',data={'basis':'独立数学底稿，测试非专业签审'},files={'file':('baseline.csv',raw.encode())})
    assert res.status_code==200,res.text
    recon=res.json();assert recon['passed'] and recon['rows'][0]['delta']=='0.0'
    approve(r,pid,'reconciliations',recon)
    assert not any('黄金底稿' in message for message in get(e,pid)['blockers'])
    q2=e.post(base+'/quotes/calculate').json();approve(r,pid,'quotes',q2)
    comp['parameters']['thickness']='.4'
    comp=e.put(base+'/components/'+comp['id'],json=comp).json()
    approve(r,pid,'components',comp)
    assert e.post(base+'/takeoffs/calculate').status_code==200
    approve(r,pid,'reconciliations',recon,expected=409)
    assert any('黄金底稿' in message for message in get(e,pid)['blockers'])


def test_failed_baseline_cannot_be_approved_and_units_do_not_match(clients):
    e,r,pid,comp,q=priced_project(clients);base=f'/api/v1/projects/{pid}'
    for unit in ['m³','元']:
        raw=f'构件名称,单位,基准工程量,绝对容差,相对容差\n测试底板,{unit},2982485.80,0,0\n'
        record=e.post(base+'/takeoffs/reconcile',data={'basis':'错误底稿测试'},files={'file':('wrong.csv',raw.encode())}).json()
        assert not record['passed']
        approve(r,pid,'reconciliations',record,expected=409)


def test_source_version_replacement_does_not_make_false_current_conflict(clients):
    e=clients['estimator'];pid=create(e)['id']
    upload(e,pid,'构件：底板\n长度10m\n规模500m³/日')
    upload(e,pid,'构件：底板\n长度11m\n规模600m³/日')
    p=get(e,pid)
    assert [x['state'] for x in p['dimension_candidates']]==['superseded','draft']
    assert [x['state'] for x in p['facts']]==['superseded','draft']


def test_revoked_applied_change_requires_new_approved_correction(clients):
    e,r,pid,comp,q=priced_project(clients);base=f'/api/v1/projects/{pid}';p=get(e,pid);b=p['boq'][0];eid=p['evidence'][0]['id']
    ch=e.post(base+'/changes',json={'boq_id':b['id'],'quantity':'16','reason':'测试变更','evidence_id':eid}).json()
    approve(r,pid,'changes',ch);e.post(base+'/changes/'+ch['id']+'/apply')
    approve(r,pid,'changes',get(e,pid)['changes'][0],'reject')
    t=get(e,pid)['takeoffs'][0]
    approve(r,pid,'takeoffs',t,expected=409)
    correction=e.post(base+'/changes',json={'boq_id':b['id'],'quantity':'15','reason':'新签证修正回原数量','evidence_id':eid})
    assert correction.status_code==200,correction.text
    approve(r,pid,'changes',correction.json())
    res=e.post(base+'/changes/'+correction.json()['id']+'/apply');assert res.status_code==200,res.text
    assert all(t['state']=='approved' for t in get(e,pid)['takeoffs'])
    approve(r,pid,'boq',res.json())


def test_review_snapshot_and_hash_are_returned_atomically(clients):
    from backend import store
    e,r,pid,comp,t,b,eid=setup_boq(clients);base=f'/api/v1/projects/{pid}'
    snapshot=r.get(base+'/review-target',params={'collection':'boq','object_id':b['id']}).json()
    assert snapshot['object_hash']==store.digest(snapshot['object'])
    e.put(base+'/boq/'+b['id'],json={'features':'review opened before this edit'})
    response=r.post(base+'/reviews',json={'collection':'boq','object_id':b['id'],'object_hash':snapshot['object_hash'],'decision':'approve','reason':'stale UI must be rejected'})
    assert response.status_code==409
    assert clients['other'].get(base+'/review-target',params={'collection':'boq','object_id':b['id']}).status_code==404
