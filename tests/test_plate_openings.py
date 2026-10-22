"""Synthetic offset/disjoint circular openings. Geometry checks do not certify structural safety."""
import copy
import io
import json
import random
from decimal import Decimal, localcontext
import pytest
from openpyxl import load_workbook
from pypdf import PdfReader
from backend import store
from backend.compound_takeoff import PI
from backend.chamber_plate import calculate
from backend.engine import DomainError
from backend.volume_balance import snapshot
from tests.test_chamber_plate import parent, seal, payload as legacy_payload, setup as legacy_setup
from tests.test_trench_takeoff import evaluate
from tests.test_volume_balance import payload as balance_payload
from tests.test_workflow import get, upload, approve, create


def hole(name,x,y,diameter,eid='test',unit='m'):
    return {'name':name,'parameters':{'x':str(x),'y':str(y),'diameter':str(diameter)},'unit':unit,
            'evidence':{k:eid for k in ['x','y','diameter','measurement_basis']},'measurement_basis':'仅合成圆形贯通孔，中心相对于本板中心，局部坐标有据核定'}


def payload(source,eid='test'):
    b=legacy_payload(source,eid);b.update(name='合成偏心多孔井板（非真实工程）',opening_mode='offset_circles',measurement_basis='仅合成同层圆板与偏心多圆孔，逐孔局部位置及互不重叠范围有据核定，继承井数一次，回填范围另审')
    d=b['plate_sections'][0];d['parameters'].pop('opening_diameter');d['evidence'].pop('opening_diameter')
    d.update(coordinate_basis='板中心为原点，水平X向右Y向上，位置仅合成，非施工定位',openings=[hole('A','-.6',0,'.4',eid),hole('B',600,0,400,eid,'mm')])
    d['evidence']['coordinate_basis']=eid;d['measurement_basis']='仅合成水平圆板、两个板内且互不重叠的圆形贯通孔，范围及独立标高已核定'
    return b


def setup(clients,save=False):
    e,r,pid,source,trench,_,eid=legacy_setup(clients);base='/api/v1/projects/'+pid
    upload(e,pid,'合成圆板直径2.4m、厚0.2m、板底标高-0.2m；局部原点在板中心，X向右Y向上。贯通孔A中心(-0.6,0)m、直径0.4m；B中心(600,0)mm、直径400mm，完全位于板内且互不重叠。每井本层一块板，井数继承一次，净量总修约6位；无真实施工效力。',name='hole-layout.txt')
    for c in get(e,pid)['components']:approve(r,pid,'components',c)
    source=get(e,pid)['components'][-1];eid=get(e,pid)['evidence'][-1]['id'];b=payload(source,eid)
    if save:
        response=e.post(base+'/components',json=b);assert response.status_code==200,response.text
    return e,r,pid,source,trench,b,eid


def test_offset_holes_mixed_units_and_check_asts():
    s=parent();b=payload(s);before=copy.deepcopy((s,b));r=calculate(b,s);assert (s,b)==before
    assert r['quantity']=='0.854513' and r['plate_summary']['hole_count']==2 and r['rule_version']=='circular-plate-offset-openings-1.0'
    row=r['plate_sections'][0];assert row['openings'][1]['normalized_parameters']=={'x':'0.6','y':'0','diameter':'0.4'}
    assert row['opening_checks'][0]['state']=='disjoint'
    with localcontext() as ctx:
        ctx.prec=80
        assert Decimal(r['deduction'])==Decimal(PI)*Decimal('.016')
        assert evaluate(r['formula_ast'],r['dimensions'])==Decimal(r['quantity'])
        for h in row['openings']:
            assert evaluate(h['formula_ast'],r['dimensions'])==Decimal(h['single_volume'])
            assert evaluate(h['containment']['formula_ast'],r['dimensions'])==Decimal(h['containment']['margin'])
        c=row['opening_checks'][0];assert evaluate(c['formula_ast'],r['dimensions'])==Decimal(c['margin'])


def test_tangency_at_board_edge_and_between_holes_is_exact():
    s=parent();b=payload(s);b['plate_sections'][0]['openings']=[hole('edge',1,0,'.4')]
    r=calculate(b,s);assert r['plate_sections'][0]['openings'][0]['containment']['state']=='tangent'
    b['plate_sections'][0]['openings']=[hole('a',0,0,'.4'),hole('b','.4',0,'.4')]
    r=calculate(b,s);assert r['plate_sections'][0]['opening_checks'][0]['state']=='tangent'
    b['plate_sections'][0]['openings'][1]['parameters']['x']='.399999999999999999'
    with pytest.raises(DomainError,match='重叠'):calculate(b,s)
    b['plate_sections'][0]['openings']=[hole('edge','1.00000000000000001',0,'.4')]
    with pytest.raises(DomainError,match='越出'):calculate(b,s)


def test_empty_hole_table_and_concentric_layout_match_legacy_quantities():
    s=parent();b=payload(s);b['plate_sections'][0]['openings']=[];r=calculate(b,s)
    assert r['quantity']=='0.904779' and r['deduction']=='0' and not r['plate_sections'][0]['opening_checks']
    legacy=legacy_payload(s);old=calculate(legacy,s);before=store.digest((s,legacy,old))
    b['plate_sections'][0]['openings']=[hole('central',0,0,'.6')];r=calculate(b,s)
    assert all(r[k]==old[k] for k in ['quantity','gross','deduction'])
    assert calculate(legacy,s)==old and store.digest((s,legacy,old))==before
    assert old['rule_version']=='circular-chamber-plate-1.0' and 'opening_mode' not in old


def test_seeded_disjoint_holes_match_independent_area_difference():
    rng=random.Random(3601);locations=[('-.7','-.7'),('.7','-.7'),('-.7','.7'),('.7','.7'),('0','0')]
    for _ in range(30):
        s=parent();count=rng.randint(1,8);s['chamber_wells'][0]['multiplicity']=str(count);seal(s);b=payload(s);b['rounding_digits']=12
        d=b['plate_sections'][0];d['parameters']['diameter']='4';thickness=Decimal(rng.randint(1,20))/100;d['parameters']['thickness']=str(thickness)
        ds=[Decimal(rng.randint(1,60))/100 for _ in range(rng.randint(1,5))]
        d['openings']=[hole(str(i),x,y,ds[i]) for i,(x,y) in enumerate(locations[:len(ds)])]
        with localcontext() as ctx:
            ctx.prec=80;expected=Decimal(PI)/4*(16-sum((v*v for v in ds),Decimal(0)))*thickness*count;r=calculate(b,s)
            assert Decimal(r['quantity'])==expected.quantize(Decimal('.000000000001'))
            assert evaluate(r['formula_ast'],r['dimensions'])==Decimal(r['quantity'])
            for h in r['plate_sections'][0]['openings']:assert Decimal(h['containment']['margin'])>=0
            for c in r['plate_sections'][0]['opening_checks']:assert Decimal(c['margin'])>=0


@pytest.mark.parametrize('case',['overlap','contained','same_center','outside','whole_board','negative','zero','nan','unit','precision','too_many','not_list','duplicate_name','missing_x','missing_proof','extra_field','missing_coordinates','missing_coordinate_proof','legacy_parameters','unknown_mode','null_mode','role','mode','unconfirmed','quantity_override','missing_holes'])
def test_invalid_layouts_fail(case):
    s=parent();b=payload(s);d=b['plate_sections'][0];h=d['openings'][0]
    if case=='overlap':d['openings'][1]=hole('B','-.5',0,'.4')
    elif case=='contained':d['openings']=[hole('a',0,0,1),hole('b',0,0,'.2')]
    elif case=='same_center':d['openings'][1]=hole('B','-.6',0,'.4')
    elif case=='outside':h['parameters']['x']='1.1'
    elif case=='whole_board':h['parameters'].update(x='0',diameter='2.4')
    elif case=='negative':h['parameters']['diameter']='-.4'
    elif case=='zero':h['parameters']['diameter']='0'
    elif case=='nan':h['parameters']['y']='NaN'
    elif case=='unit':h['unit']='m²'
    elif case=='precision':h['unit']='mm';h['parameters']['x']='.000000000000000001'
    elif case=='too_many':d['openings']*=11
    elif case=='not_list':d['openings']={}
    elif case=='duplicate_name':d['openings'][1]['name']='A'
    elif case=='missing_x':del h['parameters']['x']
    elif case=='missing_proof':del h['evidence']['y']
    elif case=='extra_field':h['net_area']='.1'
    elif case=='missing_coordinates':d['coordinate_basis']=''
    elif case=='missing_coordinate_proof':del d['evidence']['coordinate_basis']
    elif case=='legacy_parameters':d['parameters']['opening_diameter']='.6'
    elif case=='unknown_mode':b['opening_mode']='union_all_circles'
    elif case=='null_mode':b['opening_mode']=None
    elif case=='role':b['plate_role']='unknown'
    elif case=='mode':b['quantity_type']='area'
    elif case=='unconfirmed':b['scope_confirmed']=False
    elif case=='quantity_override':b['multiplicity']='2'
    elif case=='missing_holes':del d['openings']
    with pytest.raises(DomainError):calculate(b,s)


def test_workflow_review_books_boq_and_exports_and_legacy_coexistence(clients):
    e,r,pid,source,trench,b,eid=setup(clients);base='/api/v1/projects/'+pid;before=get(e,pid)
    response=e.post(base+'/components/preview',json=b);assert response.status_code==200,response.text
    assert response.json()['quantity']=='0.854513' and get(e,pid)==before
    response=e.post(base+'/components',json={**b,'source_geometry':{'forged':True}});assert response.status_code==200,response.text
    new=response.json();assert new['opening_mode']=='offset_circles' and new['source_geometry']==snapshot(source)
    assert e.post(base+'/takeoffs/calculate').status_code==409;approve(r,pid,'components',new)
    old=legacy_payload(source,eid);old['name']='旧版同心孔井板（合成）';old=e.post(base+'/components',json=old).json();approve(r,pid,'components',old)
    response=e.post(base+'/takeoffs/calculate');assert response.status_code==200,response.text
    books=response.json();assert books[-2]['rule_version']=='circular-plate-offset-openings-1.0' and books[-1]['rule_version']=='circular-chamber-plate-1.0'
    for book in books:approve(r,pid,'takeoffs',book)
    items=e.post(base+'/boq/generate').json();assert [x['quantity'] for x in items[-2:]]==['0.854513','0.848230']
    js=e.get(base+'/exports?unpriced=true&format=json').json()['unpriced']['takeoffs'][-2]
    assert js['plate_sections'][0]['openings']==new['preview']['plate_sections'][0]['openings']
    response=e.get(base+'/exports?unpriced=true&format=xlsx');assert response.status_code==200,response.text
    wb=load_workbook(io.BytesIO(response.content));assert {'井板逐孔定位与扣减','井板孔包含与分离校验'}<=set(wb.sheetnames)
    assert all(c.data_type!='f' for s in wb for row in s for c in row)
    response=e.get(base+'/exports?unpriced=true&format=pdf');assert response.status_code==200,response.text
    text='\n'.join(p.extract_text() for p in PdfReader(io.BytesIO(response.content)).pages)
    assert '井板逐孔定位与扣减' in text and '井板孔包含与分离校验' in text and '0.854513' in text


def test_relocated_hole_same_quantity_still_invalidates_backfill(clients):
    e,r,pid,source,trench,b,eid=setup(clients,save=True);base='/api/v1/projects/'+pid;comp=get(e,pid)['components'][-1];approve(r,pid,'components',comp)
    t=e.post(base+'/components',json=trench);assert t.status_code==200,t.text;approve(r,pid,'components',t.json());p=get(e,pid)
    src={'excavation':snapshot(p['components'][-1]),'occupations':[snapshot(c) for c in p['components'][1:3]]};balance=balance_payload(src,eid);balance['rounding_digits']=6
    for d in balance['occupations']:d['factor']='1'
    response=e.post(base+'/components',json=balance);assert response.status_code==200,response.text
    assert response.json()['preview']['quantity']=='37.220709';approve(r,pid,'components',response.json())
    books=e.post(base+'/takeoffs/calculate');assert books.status_code==200,books.text
    for book in books.json():approve(r,pid,'takeoffs',book)
    before=get(e,pid);assert e.put(base+'/components/'+comp['id'],json=b).status_code==409 and get(e,pid)==before
    b['object_hash']=e.get(base+'/hashes').json()['components'][comp['id']];bad=copy.deepcopy(b);bad['plate_sections'][0]['openings'][0]['parameters']['x']='.6'
    assert e.put(base+'/components/'+comp['id'],json=bad).status_code==422 and get(e,pid)==before
    b['plate_sections'][0]['openings'][0]['parameters']['x']='-.7';response=e.put(base+'/components/'+comp['id'],json=b);assert response.status_code==200,response.text
    assert response.json()['preview']['quantity']=='0.854513' and response.json()['revision']==2
    p=get(e,pid);assert [c['state'] for c in p['components']]==['approved','approved','draft','approved','stale']
    assert all(t['state']=='stale' for t in p['takeoffs']);approve(r,pid,'components',p['components'][2]);approve(r,pid,'components',p['components'][-1],expected=409)


@pytest.mark.parametrize('gate',['preview','save','component_review','calculate','book_review'])
def test_hole_position_original_corruption_blocks_all_gates(clients,gate):
    e,r,pid,source,trench,b,eid=setup(clients);base='/api/v1/projects/'+pid
    if gate not in ('preview','save'):comp=e.post(base+'/components',json=b).json()
    if gate in ('calculate','book_review'):approve(r,pid,'components',comp)
    if gate=='book_review':book=e.post(base+'/takeoffs/calculate').json()[-1]
    doc=get(e,pid)['documents'][-1];(store.ROOT/'originals'/doc['hash']).write_bytes(b'corrupt')
    if gate=='preview':assert e.post(base+'/components/preview',json=b).status_code==409
    elif gate=='save':assert e.post(base+'/components',json=b).status_code==409
    elif gate=='component_review':approve(r,pid,'components',comp,expected=409)
    elif gate=='calculate':assert e.post(base+'/takeoffs/calculate').status_code==409
    else:approve(r,pid,'takeoffs',book,expected=409)


def test_foreign_hole_proof_and_permissions(clients):
    e,r,pid,source,trench,b,eid=setup(clients);base='/api/v1/projects/'+pid
    assert r.post(base+'/components',json=b).status_code==403 and clients['other'].post(base+'/components/preview',json=b).status_code==404
    other=create(e)['id'];upload(e,other,'其他项目孔位');bad=copy.deepcopy(b);bad['plate_sections'][0]['openings'][0]['evidence']['x']=get(e,other)['evidence'][0]['id']
    assert e.post(base+'/components',json=bad).status_code==404


@pytest.mark.parametrize('target',['component','book'])
def test_saved_check_tampering_cannot_pass_review(clients,target):
    e,r,pid,source,trench,b,eid=setup(clients,save=True);base='/api/v1/projects/'+pid
    if target=='book':approve(r,pid,'components',get(e,pid)['components'][-1]);e.post(base+'/takeoffs/calculate')
    with store.transaction() as db:
        p=json.loads(db.execute('SELECT data FROM projects WHERE id=?',(pid,)).fetchone()['data'])
        row=p['components'][-1]['preview']['plate_sections'][0] if target=='component' else p['takeoffs'][-1]['plate_sections'][0]
        row['openings'][0]['containment']['margin']='999';store.save(db,p)
    col='components' if target=='component' else 'takeoffs';approve(r,pid,col,get(e,pid)[col][-1],expected=409)
