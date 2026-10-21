"""Missing prices are explicit unresolved work, never synthetic zero-price quotes."""
import copy
import io
import pytest
from openpyxl import load_workbook
from backend import store
from backend.engine import calculate, DomainError
from tests.test_workflow import priced_project, setup_boq, get, approve


def pending_resource(eid):
    return {'name':'合成主材（非真实材料）','type':'材料','unit':'m³','consumption':'1.05','price':None,
            'pricing_status':'unpriced','resource_code':'TEST-MAT','specification':'仅数学测试规格',
            'fee_allocation':'合成材料费，不含安装费','evidence_id':eid}


def add_pending(e,pid):
    b=get(e,pid)['boq'][0];qs=copy.deepcopy(b['quotas']);qs[0]['resources']=[pending_resource(qs[0]['evidence_id'])]
    response=e.put(f'/api/v1/projects/{pid}/boq/'+b['id'],json={'quotas':qs,'object_hash':store.digest(b)})
    assert response.status_code==200,response.text
    return response.json()


def test_pending_work_survives_review_export_then_closes_only_after_price_saved(clients):
    e,r,pid,comp,old=priced_project(clients);base=f'/api/v1/projects/{pid}'
    b=add_pending(e,pid)
    tasks=e.get(base+'/pricing-tasks').json();assert len(tasks)==1
    assert tasks[0]['resource']['price'] is None and tasks[0]['boq_hash']==store.digest(b)
    assert tasks[0]['kind']=='price_inquiry' and get(e,pid)['pricing_tasks']==tasks
    approve(r,pid,'boq',b)
    quote=e.post(base+'/quotes/calculate');assert quote.status_code==409 and '待询价' in quote.text
    # Current draft exports must show unresolved work even when an older priced quote exists.
    for fmt in ['xlsx','pdf','json']:
        exported=e.get(base+'/exports',params={'format':fmt,'unpriced':'true'})
        assert exported.status_code==200,exported.text[:150] if fmt=='json' else exported.status_code
        if fmt=='xlsx':
            wb=load_workbook(io.BytesIO(exported.content));assert '待询价与匹配任务' in wb.sheetnames
            assert any('待询价' in str(c.value) for row in wb['资源价格表'] for c in row)
        if fmt=='json':
            assert exported.json()['quote'] is None
            assert '待询价与匹配任务' in [section[0] for section in exported.json()['tables']]
    current=get(e,pid)['boq'][0];qs=copy.deepcopy(current['quotas'])
    qs[0]['resources'][0]={**old['lines'][0]['quotas'][0]['resources'][0],**pending_resource(qs[0]['evidence_id']),
                          'pricing_status':'priced','source_type':'供应商报价','price':'120'}
    saved=e.put(base+'/boq/'+b['id'],json={'quotas':qs,'object_hash':store.digest(current)})
    assert saved.status_code==200,saved.text
    assert e.get(base+'/pricing-tasks').json()==[]
    assert e.post(base+'/quotes/calculate').status_code==409, 'saved price still requires review'
    approve(r,pid,'boq',saved.json())
    new=e.post(base+'/quotes/calculate');assert new.status_code==200,new.text
    assert new.json()['lines'][0]['unit_price']=='126.00'
    assert next(q for q in get(e,pid)['quotes'] if q['id']==old['id'])==old


@pytest.mark.parametrize('field,value', [('price','0'),('price','100'),('resource_code',''),('specification',' '),
    ('unit',None),('fee_allocation',''),('consumption','NaN'),('pricing_status','pending'),('type','未知'),
    ('source_type','模型猜测'),('name',False)])
def test_pending_validation_is_atomic(clients,field,value):
    e,r,pid,comp,old=priced_project(clients);base=f'/api/v1/projects/{pid}'
    b=get(e,pid)['boq'][0];qs=copy.deepcopy(b['quotas']);resource=pending_resource(qs[0]['evidence_id']);resource[field]=value;qs[0]['resources']=[resource]
    before=get(e,pid)
    response=e.put(base+'/boq/'+b['id'],json={'quotas':qs})
    assert response.status_code==422,response.text
    assert get(e,pid)==before


def test_pending_pure_engine_never_returns_partial_amounts():
    item={'quantity':'1','quotas':[{'code':'TEST','resources':[pending_resource('synthetic')]}]}
    with pytest.raises(DomainError,match='待询价'):
        calculate([item],{'fees':[]})


def test_missing_quota_task_and_tenant_isolation(clients):
    e,r,pid,comp,t,b,eid=setup_boq(clients);base=f'/api/v1/projects/{pid}'
    tasks=e.get(base+'/pricing-tasks').json()
    assert len(tasks)==1 and tasks[0]['kind']=='quota_match'
    assert clients['other'].get(base+'/pricing-tasks').status_code==404


def test_pending_evidence_and_optimistic_save(clients):
    e,r,pid,comp,old=priced_project(clients);base=f'/api/v1/projects/{pid}'
    b=get(e,pid)['boq'][0];qs=copy.deepcopy(b['quotas']);qs[0]['resources']=[pending_resource('foreign')]
    assert e.put(base+'/boq/'+b['id'],json={'quotas':qs}).status_code==404
    added=add_pending(e,pid)
    assert e.put(base+'/boq/'+b['id'],json={'quotas':b['quotas'],'object_hash':store.digest(b)}).status_code==409
    assert get(e,pid)['boq'][0]==added


def test_historical_price_cannot_silently_become_current_quote(clients):
    e,r,pid,comp,old=priced_project(clients);base=f'/api/v1/projects/{pid}'
    b=get(e,pid)['boq'][0];qs=copy.deepcopy(b['quotas']);qs[0]['resources'][0].update(pricing_status='priced',source_type='历史审核价')
    response=e.put(base+'/boq/'+b['id'],json={'quotas':qs});assert response.status_code==200
    approve(r,pid,'boq',response.json())
    blocked=e.post(base+'/quotes/calculate');assert blocked.status_code==409 and '历史审核价' in blocked.text


def test_explicit_priced_state_requires_a_source_type(clients):
    e,r,pid,comp,old=priced_project(clients);base=f'/api/v1/projects/{pid}'
    b=get(e,pid)['boq'][0];qs=copy.deepcopy(b['quotas']);qs[0]['resources'][0]['pricing_status']='priced'
    response=e.put(base+'/boq/'+b['id'],json={'quotas':qs})
    assert response.status_code==422 and '来源类型' in response.text


def test_current_unpriced_bundle_with_existing_quote_and_selection_conflicts(clients):
    import json,zipfile
    e,r,pid,comp,old=priced_project(clients);base=f'/api/v1/projects/{pid}'
    add_pending(e,pid)
    response=e.get(base+'/exports',params={'format':'zip','unpriced':True})
    assert response.status_code==200,response.text[:100]
    archive=zipfile.ZipFile(io.BytesIO(response.content))
    assert 'unpriced-snapshot.json' in archive.namelist() and 'quote.json' not in archive.namelist()
    draft=json.loads(archive.read('unpriced-snapshot.json'))
    assert draft['boq'][0]['quotas'][0]['resources'][0]['price'] is None
    check=e.get(base+'/delivery-check?unpriced=true').json()
    assert check['valid'] and check['quote_id'] is None and not check['approved']
    for suffix in ['/exports?format=json','/delivery-check?']:
        assert e.get(base+suffix+'&unpriced=true&quote_id='+old['id']).status_code==422
    historical=e.get(base+'/exports',params={'format':'json','quote_id':old['id']}).json()
    assert historical['quote']==old
    assert '待询价与匹配任务' not in [row[0] for row in historical['tables']]


def test_restricted_project_tasks_are_inaccessible(clients):
    from tests.test_project_access import restrict
    e,r,pid,comp,old=priced_project(clients)
    add_pending(e,pid)
    restrict(clients['admin'],pid,[])
    assert e.get(f'/api/v1/projects/{pid}/pricing-tasks').status_code==404
