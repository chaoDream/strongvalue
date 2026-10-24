"""Local horizontal XY windows; independent area, hole clipping and claim checks."""
import copy,io,math,random
from decimal import Decimal,localcontext,ROUND_HALF_UP
import pytest
from openpyxl import load_workbook
from pypdf import PdfReader
from backend import store
from backend.engine import DomainError
from backend.foundation_horizontal_formwork import calculate
from backend.foundation_horizontal_windows import VERSION,FIELDS
from backend.foundation_horizontal_formwork_reports import WINDOW_TITLES,WINDOW_CURVED_TITLE,tables
from backend.railing_material import snapshot
from tests.test_foundation_horizontal_formwork import body as full_body,contact,check as base_check,setup as full_setup
from tests.test_foundation_formwork import foundation
from tests.test_foundation_openings import hole
from tests.test_foundation_voids import rectangle
from tests.test_foundation_recesses import interval
from tests.test_post_foundation import equal_body,fixed_body,fixed_added,equal_added,fixed_adjusted,equal_adjusted,combined
from tests.test_route_takeoff import evaluate
from tests.test_polygon_takeoff import post,mutate
from tests.test_workflow import get,approve,upload,create,priced_project

FRAME='合成基础中心XY，X向设计东Y向北，与已审孔位共用，不随俯仰视方向翻转；仅软件验收。'


def window(eid='test',x1='0',x2='200',y1='-250',y2='250'):
    return dict(parameters=dict(zip(FIELDS,(x1,x2,y1,y2))),unit='mm',evidence={k:eid for k in FIELDS},coordinate_basis=FRAME,coordinate_evidence_id=eid,scope_confirmed=True)


def body(s,eid='test',face='bottom',bounds=None):
    return {**full_body(s,eid),'name':'合成水平局部窗口模板（非真实工程）','contact_ranges':[{**contact(eid,face),'plan_range':window(eid,*(bounds or ('0','200','-250','250')))}]}


def source(shape='rectangular',make=equal_body):return foundation(make,shape)


def check(r):
    base_check(r)
    with localcontext() as ctx:
        ctx.prec=80
        for d in r['contact_ranges']:
            for h in d['apertures']:
                for k,a in h.get('formula_asts',{}).items():assert evaluate(a,r['dimensions'])==Decimal(h[k])
            for m in [d.get('window_curved_intersection')]+[h.get('curved_intersection') for h in d['apertures']]:
                if m:
                    for q in m['terms']:assert evaluate(q['formula_ast'],r['dimensions'])==Decimal(q['area'])
        for c in r.get('horizontal_window_checks',[]):
            for k,a in c['formula_asts'].items():assert evaluate(a,r['dimensions'])==Decimal(c[k])


def test_partial_actual_hole_threshold_and_no_input_mutation():
    s=source();b=body(s);before=copy.deepcopy((s,b));r=calculate(b,s);d=r['contact_ranges'][0];assert (s,b)==before and r['quantity']=='2.068' and r['rule_version']==VERSION;assert Decimal(d['single_gross_area'])==Decimal('.1') and Decimal(d['single_opening_area'])==Decimal('.006');check(r)


@pytest.mark.parametrize('shape',['rectangular','stepped_rectangular','circular','stepped_circular'])
@pytest.mark.parametrize('face',['top','bottom'])
@pytest.mark.parametrize('make',[equal_body,fixed_body,fixed_added,equal_added,fixed_adjusted,equal_adjusted,combined,lambda s:combined(s,equal=True)])
def test_shapes_end_layers_and_eight_final_counts_partition_matches_whole(shape,face,make):
    s=source(shape,make);b=body(s,face=face,bounds=('-1000','0','-1000','1000'));b['contact_ranges'].append({**contact(face=face,name='合成右半端面'),'plan_range':window(x1='0',x2='1000',y1='-1000',y2='1000')});r=calculate(b,s);full=full_body(s);full['contact_ranges']=[contact(face=face)];whole=calculate(full,s);assert r['quantity']==whole['quantity'];assert r['horizontal_window_checks'][0]['state']=='boundary_touch';check(r)


@pytest.mark.parametrize('x1,x2,y1,y2,expected',[('-500','500','-500','500','.194'),('200','300','0','100','0'),('-300','-200','0','100','0'),('0','200','-30','0','.003'),('0','200','30','100','.014'),('100','200','-30','30','0')])
def test_rectangular_window_partial_hole_empty_tangent_and_outside(x1,x2,y1,y2,expected):
    s=source();r=calculate(body(s,bounds=(x1,x2,y1,y2)),s);assert Decimal(r['contact_ranges'][0]['single_net_area'])==Decimal(expected);check(r)


@pytest.mark.parametrize('unit,factor',[('mm','1'),('cm','.1'),('dm','.01'),('m','.001'),('km','.000001')])
def test_five_window_units_low_calling_precision(unit,factor):
    s=source();b=body(s);c=b['contact_ranges'][0]['plan_range'];c['unit']=unit
    for k in FIELDS:c['parameters'][k]=str(Decimal(c['parameters'][k])*Decimal(factor))
    with localcontext() as ctx:ctx.prec=6;r=calculate(b,s)
    assert r==calculate(b,s) and r['quantity']=='2.068';check(r)


def test_circle_segment_and_partial_circle_hole_independent_math():
    s=source('circular');c=s['foundation_items'][0]['through_holes'];c['holes']=[{**hole(x='80'),'vertical_range':interval('0','200')}];b=body(s,bounds=('80','400','-400','400'));r=calculate(b,s);d=r['contact_ranges'][0];R=.3;a=.08;expected=R*R*math.acos(a/R)-a*math.sqrt(R*R-a*a)-math.pi*.02*.02/2;assert float(d['single_net_area'])==pytest.approx(expected,abs=2e-16);assert d['window_curved_intersection'] and d['apertures'][0]['curved_intersection'];check(r)


def test_160_seeded_rectangle_windows_independent_integer_clip_and_rounding():
    rng=random.Random(10301)
    for _ in range(160):
        x1,x2=sorted(rng.sample(range(-400,401),2));y1,y2=sorted(rng.sample(range(-400,401),2));s=source();b=body(s,bounds=tuple(map(str,(x1,x2,y1,y2))));b['rounding_digits']=rng.randrange(13);r=calculate(b,s)
        with localcontext() as ctx:
            ctx.prec=80;gross=max(0,min(x2,200)-max(x1,-200))*max(0,min(y2,250)-max(y1,-250));cut=max(0,min(x2,200)-max(x1,100))*max(0,min(y2,30)-max(y1,-30));expected=Decimal(gross-cut)/1000000*22;assert Decimal(r['quantity'])==expected.quantize(Decimal(1).scaleb(-b['rounding_digits']),rounding=ROUND_HALF_UP)
        check(r)


@pytest.mark.parametrize('fault',['null','empty','extra','missing','no-evidence','blank-proof','basis','frame-proof','unconfirmed','truthy','zero-x','reverse-y','unit','non-number','manual-area','step','hole-end','same-face','full-face','different-frame','different-frame-proof'])
def test_strict_window_schema_and_claims(fault):
    s=source('stepped_rectangular');b=body(s);d=b['contact_ranges'][0];c=d['plan_range']
    if fault=='null':d['plan_range']=None
    elif fault=='empty':d['plan_range']={}
    elif fault=='extra':c['area']='1'
    elif fault=='missing':c['parameters'].pop('x_start')
    elif fault=='no-evidence':c['evidence'].pop('y_end')
    elif fault=='blank-proof':c['evidence']['x_start']=' '
    elif fault=='basis':c['coordinate_basis']=' '
    elif fault=='frame-proof':c['coordinate_evidence_id']=' '
    elif fault=='unconfirmed':c['scope_confirmed']=False
    elif fault=='truthy':c['scope_confirmed']=1
    elif fault=='zero-x':c['parameters']['x_end']='0'
    elif fault=='reverse-y':c['parameters']['y_end']='-300'
    elif fault=='unit':c['unit']='m²'
    elif fault=='non-number':c['parameters']['x_start']=True
    elif fault=='manual-area':d['area']='1'
    elif fault=='step':s=source('stepped_circular');b['source_component_hash']=store.digest(snapshot(s));d.update(face='step',level_index=1)
    elif fault=='hole-end':d.update(face='hole_top',hole_index=1)
    elif fault=='same-face':b['contact_ranges'].append({**copy.deepcopy(d),'name':'重叠'})
    elif fault=='full-face':b['contact_ranges'].append(contact(face='bottom',name='整面重叠'))
    else:
        other=copy.deepcopy(d);other['name']='另半窗口';other['plan_range']['parameters'].update(x_start='-300',x_end='0');other['plan_range']['coordinate_basis' if fault=='different-frame' else 'coordinate_evidence_id']='另一基准';b['contact_ranges'].append(other)
    with pytest.raises(DomainError):calculate(b,s)


def test_exact_separation_sub_unit_overlap_and_whole_face_reverse_order():
    s=source();b=body(s,bounds=('-200','0','-250','250'));other={**contact(face='bottom',name='右半'),'plan_range':window()};b['contact_ranges'].append(other);check(calculate(b,s));other['plan_range']['parameters']['x_start']='-0.000000001'
    with pytest.raises(DomainError,match='重叠'):calculate(b,s)
    b['contact_ranges']=[contact(face='bottom'),{**other,'name':'窗口'}]
    with pytest.raises(DomainError,match='重叠'):calculate(b,s)


def setup(clients,pid=None):
    e,r,pid,s,_,_=full_setup(clients,pid,'circular')
    upload(e,pid,FRAME+'仅合成22个Φ600×600mm圆柱基础，H1圆孔Φ40中心(80,0)Z0至200mm，H2矩形100×60中心(-80,0)Z0至150mm。体积阈值100m²不扣计量体积，物理孔仍参与模板。整体底面窗口X80至300、Y负300至300mm：毛面积为R.3m弓形，扣H1半圆.0002πm²，H2不交。V2保留原窗口，并加X0至80同Y范围，两窗口共享边界、合成右半圆扣H1全孔，净.0446π×22m²。支模与修约3位仅合成，TEST-WINDOW非官方编码。','foundation-horizontal-window-basis.txt')
    for c in get(e,pid)['components']:approve(r,pid,'components',c)
    s=next(c for c in get(e,pid)['components'] if c['id']==s['id']);eid=get(e,pid)['evidence'][-1]['id'];s['object_hash']=store.digest(s);c=s['foundation_items'][0]['through_holes'];c['deduction_rule']['threshold']='100';c['holes']=[{**hole(eid=eid,x='80'),'vertical_range':interval('0','200',eid)},{**rectangle(eid=eid,name='合成H2',x='-80'),'vertical_range':interval('0','150',eid)}];res=e.put(f'/api/v1/projects/{pid}/components/'+s['id'],json=s);assert res.status_code==200,res.text;approve(r,pid,'components',res.json());s=next(c for c in get(e,pid)['components'] if c['id']==s['id']);return e,r,pid,s,body(s,eid,bounds=('80','300','-300','300')),eid


def test_candidates_preview_cross_record_claims_same_frame_and_whole_face(clients):
    e,r,pid,s,b,eid=setup(clients);before=get(e,pid);post(e,pid,b,'/components/preview');assert get(e,pid)==before;c=post(e,pid,b).json();approve(r,pid,'components',c);post(e,pid,{**b,'id':c['id']},expected=409)
    second=body(s,eid,bounds=('0','80','-300','300'));post(e,pid,second);bad=copy.deepcopy(second);bad['contact_ranges'][0]['plan_range']['coordinate_basis']='不同基准';post(e,pid,bad,expected=409);post(e,pid,full_body(s,eid),expected=409)


@pytest.mark.parametrize('fault',['bounds','full','gross','hole','rule','trace','frame','ast','source','book'])
def test_saved_window_tamper_reconstruction_gates(clients,fault):
    e,r,pid,s,b,eid=setup(clients);c=post(e,pid,b).json()
    if fault=='book':approve(r,pid,'components',c);post(e,pid,{},'/takeoffs/calculate')
    def forge(p):
        c=p['components'][-1];d=c['preview']['contact_ranges'][0]
        if fault=='bounds':d['x_start']='0'
        elif fault=='full':d['full_face_area']='0'
        elif fault=='gross':d['single_gross_area']='0'
        elif fault=='hole':d['apertures'][0]['single_area']='0'
        elif fault=='rule':d['window_intersection_rule']='forged'
        elif fault=='trace':d['window_curved_intersection']['terms'][0]['area']='0'
        elif fault=='frame':d['plan_frame_hash']='forged'
        elif fault=='ast':d['formula_asts']['x_start']={'constant':'0'}
        elif fault=='source':c['source_geometry']['foundation_items'][0]['parameters']['diameter']='700'
        else:p['takeoffs'][-1]['quantity']='999'
    mutate(pid,forge);coll='takeoffs' if fault=='book' else 'components';approve(r,pid,coll,get(e,pid)[coll][-1],expected=409)


@pytest.mark.parametrize('gate',['preview','save','review','calculate','book_review'])
def test_window_original_integrity_five_gates(clients,gate):
    e,r,pid,s,b,eid=setup(clients)
    if gate in ('review','calculate','book_review'):c=post(e,pid,b).json()
    if gate in ('calculate','book_review'):approve(r,pid,'components',c)
    if gate=='book_review':t=post(e,pid,{},'/takeoffs/calculate').json()[-1]
    doc=get(e,pid)['documents'][-1];(store.ROOT/'originals'/doc['hash']).write_bytes(b'corrupt')
    if gate=='preview':post(e,pid,b,'/components/preview',409)
    elif gate=='save':post(e,pid,b,expected=409)
    elif gate=='review':approve(r,pid,'components',c,expected=409)
    elif gate=='calculate':post(e,pid,{},'/takeoffs/calculate',409)
    else:approve(r,pid,'takeoffs',t,expected=409)


def test_window_revision_history_concurrency_source_stale_permissions(clients):
    e,r,pid,s,b,eid=setup(clients);post(r,pid,b,expected=403);assert clients['other'].get(f'/api/v1/projects/{pid}/foundation-horizontal-formwork-candidates',params=dict(component_id=s['id'])).status_code==404
    foreign=create(e)['id'];upload(e,foreign,'异项目窗口依据');bad=copy.deepcopy(b);bad['contact_ranges'][0]['plan_range']['evidence']['x_start']=get(e,foreign)['evidence'][-1]['id'];post(e,pid,bad,expected=404)
    c=post(e,pid,b).json();approve(r,pid,'components',c)
    for t in post(e,pid,{},'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
    post(e,pid,{},'/boq/generate');before=get(e,pid);old=before['components'][-1];b['object_hash']=store.digest(old);b['contact_ranges'].append({**contact(eid,'bottom','合成补充窗口'),'plan_range':window(eid,'0','80','-300','300')});url=f'/api/v1/projects/{pid}/components/'+old['id'];res=e.put(url,json=b);assert res.status_code==200,res.text;after=get(e,pid);assert after['component_history'][-1]==old and after['boq']==before['boq'] and e.put(url,json=b).status_code==409;approve(r,pid,'components',res.json());s['object_hash']=store.digest(s);s['foundation_items'][0]['parameters']['diameter']='700';res=e.put(f'/api/v1/projects/{pid}/components/'+s['id'],json=s);assert res.status_code==200,res.text;approve(r,pid,'components',res.json());assert get(e,pid)['components'][-1]['state']=='stale'


@pytest.mark.parametrize('fmt',['json','xlsx','pdf'])
def test_three_formats_window_tables_full_precision_old_quote_excluded(clients,fmt):
    e,r,pid,_,q=priced_project(clients);e,r,pid,s,b,eid=setup(clients,pid);c=post(e,pid,b).json();approve(r,pid,'components',c)
    for t in post(e,pid,{},'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
    post(e,pid,{},'/boq/generate');url=f'/api/v1/projects/{pid}/exports';res=e.get(url,params=dict(format=fmt,unpriced=True));assert res.status_code==200,res.text;titles=set(WINDOW_TITLES)|{WINDOW_CURVED_TITLE}
    if fmt=='json':assert titles<={t[0] for t in res.json()['tables']}
    elif fmt=='xlsx':assert titles<=set(load_workbook(io.BytesIO(res.content)).sheetnames)
    else:assert all(t in ''.join(p.extract_text() for p in PdfReader(io.BytesIO(res.content)).pages) for t in titles)
    assert not any(t[0] in titles for t in e.get(url,params=dict(format='json',quote_id=q['id'])).json()['tables']);assert tables(dict(takeoffs=[]))==[]
