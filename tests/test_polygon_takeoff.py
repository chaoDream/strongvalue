"""Synthetic geometry and workflow checks, not real engineering acceptance."""
import copy
import io
import json
import random
from decimal import Decimal, localcontext
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pypdf import PdfReader
from backend import store
from backend.main import app
from backend.engine import DomainError
from backend.polygon_takeoff import calculate
from tests.test_route_takeoff import evaluate
from tests.test_workflow import create, get, upload, approve, priced_project


OUTER=[(0,0),(6,0),(6,3),(4,3),(4,5),(0,5)]
TRIANGLE=[(1,1),(2,1),(1,2)]
DIAMOND=[(3,'2.5'),('3.5',3),(3,'3.5'),('2.5',3)]


def ring(name, coords, eid='test', role='outer', unit='m'):
    return {'name':name,'role':role,'measurement_basis':'合成核定直边范围与逐孔扣减，无专业效力','evidence_id':eid,
            'points':[{'name':'P'+str(i+1),'parameters':{'x':str(x),'y':str(y)},'unit':unit,'evidence':{'x':eid,'y':eid}} for i,(x,y) in enumerate(coords)]}


def payload(eid='test',mode='volume'):
    b={'name':'合成斜孔凹形底板（非真实工程）','kind':'polygon_region','parameters':{},'quantity_type':mode,
       'polygon_rings':[ring('凹形外轮廓',OUTER,eid),ring('三角孔',TRIANGLE,eid,'hole'),ring('斜边菱形孔',DIAMOND,eid,'hole')],
       'coordinate_basis':'已核对同一水平设计平面、共同原点与XY方向及高程基准，仅合成软件验收','coordinate_evidence_id':eid,
       'measurement_basis':'外轮廓26m²，扣三角孔0.5m²、菱形孔0.5m²，净25m²；逐孔全扣，重复2，厚200mm，总量3位；仅合成',
       'basis_evidence_id':eid,'multiplicity':'2','multiplicity_evidence_id':eid,'rounding_digits':3,'rounding_evidence_id':eid,'boundaries_confirmed':True}
    if mode=='volume':b.update(thickness='200',thickness_unit='mm',thickness_evidence_id=eid)
    return b


def setup(clients,pid=None):
    e,r=clients['estimator'],clients['reviewer'];pid=pid or create(e)['id']
    upload(e,pid,'合成设计依据（非真实工程）：同一水平平面、共同原点XY方向及高程；外轮廓(0,0)(6,0)(6,3)(4,3)(4,5)(0,5)m，26m²；三角孔(1,1)(2,1)(1,2)m，0.5m²；菱形孔(3,2.5)(3.5,3)(3,3.5)(2.5,3)m，0.5m²。逐孔全扣，净25m²，统一竖向厚200mm，重复2，总量四舍五入3位。','polygon-basis.txt')
    return e,r,pid,payload(get(e,pid)['evidence'][-1]['id'])


def post(e,pid,b,endpoint='/components',expected=200):
    res=e.post('/api/v1/projects/'+pid+endpoint,json=b);assert res.status_code==expected,res.text;return res


def mutate(pid,action):
    with store.transaction() as c:
        p=json.loads(c.execute('select data from projects where id=?',(pid,)).fetchone()['data']);action(p);store.save(c,p)


def test_concave_outer_and_sloped_holes_exact_area_volume_and_ast():
    b=payload();before=copy.deepcopy(b);r=calculate(b);s=r['polygon_summary']
    assert b==before and r['quantity']=='10.000' and r['unit']=='m³'
    assert (Decimal(s['outer_area']),Decimal(s['hole_area']),Decimal(s['net_area']))==(26,1,25)
    assert [Decimal(v['area']) for v in r['polygon_rings']]==[26,Decimal('.5'),Decimal('.5')]
    assert len(r['polygon_checks'])==3 and s['hole_count']==2 and s['point_count']==13
    assert r['unit_conversions'][-1]['original_value']=='200' and s['thickness']=='0.2'
    with localcontext() as ctx:
        ctx.prec=80
        assert evaluate(r['formula_ast'],r['dimensions'])==Decimal(r['quantity'])
        for row in r['polygon_rings']:
            assert evaluate(row['formula_ast'],r['dimensions'])==Decimal(row['area'])
            for edge in row['edges']:assert evaluate(edge['formula_ast'],r['dimensions'])==Decimal(edge['cross_product'])


@pytest.mark.parametrize('mode',['area','volume'])
def test_reversal_cyclic_start_translation_and_affine_rotation_preserve_geometry(mode):
    b=payload(mode=mode);expected=Decimal(calculate(b)['quantity'])
    for r in b['polygon_rings']:
        r['points'].reverse();r['points']=r['points'][2:]+r['points'][:2]
        for p in r['points']:
            x,y=(Decimal(p['parameters'][k]) for k in ('x','y'))
            p['parameters']={'x':str((2*x+y-100)*1000),'y':str((-x+3*y-50)*1000)};p['unit']='mm'
    result=calculate(b);assert Decimal(result['quantity'])==expected*7
    with localcontext() as ctx:ctx.prec=80;assert evaluate(result['formula_ast'],result['dimensions'])==expected*7


def test_concave_negative_edge_contributions_and_collinear_forward_points_are_preserved():
    b=payload(mode='area');b['polygon_rings']=[ring('U形',[(0,0),(3,0),(6,0),(6,6),(4,6),(4,2),(2,2),(2,6),(0,6)])]
    r=calculate(b);assert r['quantity']=='56.000' and any(Decimal(e['cross_product'])<0 for e in r['polygon_rings'][0]['edges'])


def test_rounding_happens_only_after_holes_thickness_and_repeat():
    b=payload(mode='area');b['polygon_rings']=[ring('tiny',[(0,0),(1,0),(0,'.005')])];b['rounding_digits']=2
    r=calculate(b);assert r['quantity']=='0.01' and Decimal(r['polygon_summary']['net_area'])==Decimal('.0025')
    assert Decimal(r['polygon_summary']['unrounded_quantity'])==Decimal('.005')
    b['rounding_digits']=0;assert calculate(b)['quantity']=='0'


def test_tiny_design_shape_at_large_common_origin_uses_exact_decimal():
    b=payload(mode='area');b['multiplicity']='1';b['rounding_digits']=12
    with localcontext() as ctx:
        ctx.prec=80;o=Decimal('1000000000');d=Decimal('.00000001')
        b['polygon_rings']=[ring('small',[(o,o),(o+d,o),(o+d,o+d),(o,o+d)])]
    with localcontext() as ctx:ctx.prec=10;r=calculate(b)
    assert Decimal(r['polygon_summary']['net_area'])==Decimal('1e-16') and r['quantity']=='0.000000000000'


def test_seeded_step_polygons_match_independent_cell_count_and_affine_area_oracle():
    rng=random.Random(5701)
    for _ in range(35):
        w=rng.randint(2,8);heights=[rng.randint(2,8) for _ in range(w)];coords=[(0,0),(w,0),(w,heights[-1])]
        for x in range(w-1,0,-1):coords.extend([(x,heights[x]),(x,heights[x-1])])
        coords.append((0,heights[0]));coords=[p for i,p in enumerate(coords) if i==0 or coords[i-1]!=p]
        cells=[(x,y) for x in range(w) for y in range(heights[x])];chosen=rng.sample(cells,min(rng.randint(0,4),len(cells)))
        b=payload(mode='area');b['polygon_rings']=[ring('outer',coords)]
        for i,(x,y) in enumerate(chosen):
            x,y=Decimal(x),Decimal(y);a=Decimal('.25');c=Decimal('.75')
            b['polygon_rings'].append(ring(str(i),[(x+a,y+a),(x+c,y+a),(x+c,y+c),(x+a,y+c)],role='hole'))
        if rng.choice([True,False]):
            for ring_row in b['polygon_rings']:ring_row['points'].reverse()
        result=calculate(b);expected=(Decimal(sum(heights))-Decimal('.25')*len(chosen))*2
        assert Decimal(result['quantity'])==expected
        with localcontext() as ctx:ctx.prec=80;assert evaluate(result['formula_ast'],result['dimensions'])==expected


@pytest.mark.parametrize('coords',[
    [(0,0),(2,2),(0,2),(2,0)],[(0,0),(3,0),(1,0),(1,2)],[(0,0),(1,0),(2,0)],
    [(0,0),(2,0),(2,2),(1,0),(0,2)],[(0,0),(1,0),(1,1),(0,0)],[(0,0),(2,0),(2,2),(0,2),(0,1),(1,1),(1,0)],
])
def test_self_crossing_backtracking_touching_zero_area_and_repeated_closure_fail(coords):
    b=payload(mode='area');b['polygon_rings']=[ring('invalid',coords)]
    with pytest.raises(DomainError):calculate(b)


@pytest.mark.parametrize('case',['outside','boundary_vertex','boundary_edge','across_concavity','overlapping_holes','touching_holes','nested_holes','identical_holes'])
def test_invalid_hole_relations_never_silently_clip_or_double_deduct(case):
    b=payload(mode='area')
    coords={'outside':[(7,1),(8,1),(7,2)],'boundary_vertex':[(0,1),(1,1),(1,2)],'boundary_edge':[(0,1),(1,1),(1,2),(0,2)],
            'across_concavity':[(3,4),(5,2),(3,2)],'overlapping_holes':[(1,1),(2,1),(2,2),(1,2)],
            'touching_holes':[(2,1),(3,1),(3,2)],'nested_holes':[(1.1,1.1),(1.2,1.1),(1.1,1.2)],'identical_holes':TRIANGLE}[case]
    b['polygon_rings'][2]=ring('bad',coords,role='hole')
    with pytest.raises(DomainError):calculate(b)


@pytest.mark.parametrize('case',['mode','parameters','no_confirmation','confirmation_number','missing_coordinate_basis','missing_basis','missing_repeat_evidence','missing_rounding_evidence',
    'extra_point','missing_xy','missing_point_evidence','empty_ring_evidence','invalid_unit','nan','too_precise','duplicate_names','duplicate_rings','second_outer','first_hole','no_rings',
    'too_few','too_many_rings','too_many_points','total_points','fraction_count','zero_count','huge_count','bool_round','invalid_round','area_thickness','zero_thickness','overflow'])
def test_invalid_inputs_fail_explicitly(case):
    b=payload();r=b['polygon_rings'][0];p=r['points'][0]
    if case=='mode':b['quantity_type']='length'
    elif case=='parameters':b['parameters']={'length':'1'}
    elif case=='no_confirmation':b['boundaries_confirmed']=False
    elif case=='confirmation_number':b['boundaries_confirmed']=1
    elif case=='missing_coordinate_basis':b['coordinate_basis']=''
    elif case=='missing_basis':b['basis_evidence_id']=''
    elif case=='missing_repeat_evidence':b['multiplicity_evidence_id']=''
    elif case=='missing_rounding_evidence':b['rounding_evidence_id']=''
    elif case=='extra_point':p['parameters']['z']='0'
    elif case=='missing_xy':p['parameters'].pop('y')
    elif case=='missing_point_evidence':p['evidence'].pop('x')
    elif case=='empty_ring_evidence':r['evidence_id']=''
    elif case=='invalid_unit':p['unit']='m²'
    elif case=='nan':p['parameters']['x']='NaN'
    elif case=='too_precise':p['unit']='mm';p['parameters']['x']='0.000000000000000001'
    elif case=='duplicate_names':r['points'][1]['name']=p['name']
    elif case=='duplicate_rings':b['polygon_rings'][1]['name']=r['name']
    elif case=='second_outer':b['polygon_rings'][1]['role']='outer'
    elif case=='first_hole':r['role']='hole'
    elif case=='no_rings':b['polygon_rings']=[]
    elif case=='too_few':r['points']=r['points'][:2]
    elif case=='too_many_rings':b['polygon_rings']*=8
    elif case=='too_many_points':r['points']*=17
    elif case=='total_points':b['polygon_rings']=[copy.deepcopy(r) for _ in range(21)];[x.update(points=x['points']*2) for x in b['polygon_rings']]
    elif case=='fraction_count':b['multiplicity']='1.5'
    elif case=='zero_count':b['multiplicity']='0'
    elif case=='huge_count':b['multiplicity']='1000001'
    elif case=='bool_round':b['rounding_digits']=True
    elif case=='invalid_round':b['rounding_digits']=13
    elif case=='area_thickness':b['quantity_type']='area'
    elif case=='zero_thickness':b['thickness']='0'
    else:b['polygon_rings']=[ring('huge',[(0,0),('1e12',0),(0,'1e12')])]
    with pytest.raises(DomainError):calculate(b)


def test_unit_equivalent_duplicate_points_are_rejected():
    b=payload();r=b['polygon_rings'][0];r['points'][1].update(parameters={'x':'0','y':'0'},unit='mm')
    with pytest.raises(DomainError):calculate(b)


def test_exact_limits_allow_100_outer_points_20_holes_and_200_total_points():
    outer=[(x,0) for x in range(25)]+[(25,y) for y in range(25)]+[(x,25) for x in range(25,0,-1)]+[(0,y) for y in range(25,0,-1)]
    b=payload();b['polygon_rings']=[ring('外轮廓',outer)]
    for i in range(20):
        x,y=1+(i%5)*4,1+(i//5)*5
        b['polygon_rings'].append(ring('孔'+str(i),[(x,y),(x+2,y),(x+2,y+1),(x+1,y+2),(x,y+1)],role='hole'))
    r=calculate(b);assert r['polygon_summary']['point_count']==200 and r['polygon_summary']['hole_count']==20
    assert r['quantity']=='226.000' and len(r['polygon_checks'])==210


def test_preview_and_full_workflow_revision_keep_exact_history_and_invalidate_downstream(clients):
    e,r,pid,b=setup(clients);before=get(e,pid);view=post(e,pid,b,'/components/preview').json();assert get(e,pid)==before and view['quantity']=='10.000'
    comp=post(e,pid,b).json();post(e,pid,{},'/takeoffs/calculate',409);approve(r,pid,'components',comp)
    book=post(e,pid,{},'/takeoffs/calculate').json()[0];approve(r,pid,'takeoffs',book)
    assert post(e,pid,{},'/boq/generate').json()[0]['quantity']=='10.000'
    old=get(e,pid)['components'][0];changed={**copy.deepcopy(old),'object_hash':store.digest(old),'multiplicity':'3'}
    res=e.put(f'/api/v1/projects/{pid}/components/'+old['id'],json=changed);assert res.status_code==200,res.text
    new=res.json();assert new['revision']==2 and new['state']=='draft' and new['preview']['quantity']=='15.000'
    p=get(e,pid);assert p['component_history'][-1]==old and p['takeoffs'][0]['state']=='stale' and p['boq'][0]['quantity']=='10.000'
    assert e.put(f'/api/v1/projects/{pid}/components/'+old['id'],json=changed).status_code==409
    post(e,pid,{},'/takeoffs/calculate',409);approve(r,pid,'components',new)
    book=post(e,pid,{},'/takeoffs/calculate').json()[0];approve(r,pid,'takeoffs',book)
    assert post(e,pid,{},'/boq/generate').json()[0]['quantity']=='15.000'


@pytest.mark.parametrize('fault',['missing','bytes','superseded','evidence','document_hash','preview','coordinate','stamp'])
def test_current_source_and_saved_geometry_are_rechecked_at_review_and_recalculation(clients,fault):
    e,r,pid,b=setup(clients);comp=post(e,pid,b).json();p=get(e,pid);doc=p['documents'][0]
    if fault in ('missing','bytes'):
        path=store.ROOT/'originals'/doc['hash']
        if fault=='missing':path.unlink()
        else:path.write_bytes(b'corrupt')
    else:
        def damage(p):
            if fault=='superseded':p['documents'][0]['superseded']=True
            elif fault=='evidence':p['evidence'][0]['text']+=' changed'
            elif fault=='document_hash':p['evidence'][0]['document_hash']='0'*64
            elif fault=='preview':p['components'][0]['preview']['polygon_rings'][0]['area']='999'
            elif fault=='coordinate':p['components'][0]['polygon_rings'][0]['points'][1]['parameters']['x']='7'
            else:p['components'][0]['source_stamp']='0'*64
        mutate(pid,damage)
    current=get(e,pid)['components'][0];approve(r,pid,'components',current,expected=409)
    mutate(pid,lambda p:p['components'][0].update(state='approved'));post(e,pid,{},'/takeoffs/calculate',409)


def test_tampered_calculation_book_cannot_be_reviewed(clients):
    e,r,pid,b=setup(clients);comp=post(e,pid,b).json();approve(r,pid,'components',comp)
    t=post(e,pid,{},'/takeoffs/calculate').json()[0]
    mutate(pid,lambda p:p['takeoffs'][0]['polygon_checks'][0].update(relation='fake'))
    approve(r,pid,'takeoffs',t,expected=409)


def test_permission_tenant_membership_and_cross_project_evidence(clients):
    e,r,pid,b=setup(clients)
    post(r,pid,b,expected=403);post(r,pid,b,'/components/preview',403);post(clients['other'],pid,b,expected=404)
    c=TestClient(app);post(c,pid,b,expected=401)
    second=create(e)['id'];post(e,second,b,expected=404)
    from tests.test_project_access import restrict
    restrict(clients['admin'],pid,[]);post(e,pid,b,expected=404)


@pytest.mark.parametrize('fmt',['json','xlsx','pdf'])
def test_three_formats_include_geometry_evidence_and_keep_old_quote_snapshot(clients,fmt):
    e,r,pid,_,quote=priced_project(clients);e,r,pid,b=setup(clients,pid)
    comp=post(e,pid,b).json()
    for c in get(e,pid)['components']:approve(r,pid,'components',c)
    for t in post(e,pid,{},'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
    base=f'/api/v1/projects/{pid}/exports';res=e.get(base,params={'format':fmt,'unpriced':True});assert res.status_code==200,res.text
    names=['多边形净量汇总','多边形外轮廓与孔','多边形原始顶点','多边形逐边计算','多边形扣孔边界校验']
    if fmt=='json':
        tables=res.json()['tables'];assert all(any(t[0]==n for t in tables) for n in names);assert '10.000' in str(tables)
    elif fmt=='xlsx':assert all(n in load_workbook(io.BytesIO(res.content)).sheetnames for n in names)
    else:assert all(n in ''.join(p.extract_text() for p in PdfReader(io.BytesIO(res.content)).pages) for n in names)
    old=e.get(base,params={'format':'json','quote_id':quote['id']}).json();assert old['quote']==quote and not any(t[0] in names for t in old['tables'])
