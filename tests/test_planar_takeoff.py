"""Synthetic coordinate geometry checks, not professional measurement approval."""
import copy
import io
import json
import random
from decimal import Decimal,localcontext
import pytest
from openpyxl import load_workbook
from pypdf import PdfReader
from backend import store
from backend.engine import DomainError
from backend.planar_takeoff import calculate
from tests.test_compound_takeoff import ast_value
from tests.test_workflow import create,get,upload,approve,priced_project


def region(name,x,y,length,width,operation='add',eid='test',unit='m'):
    return {'name':name,'operation':operation,'parameters':dict(zip(['x','y','length','width'],map(str,[x,y,length,width]))),
            'unit':unit,'evidence':{k:eid for k in ['x','y','length','width','measurement_basis']},'measurement_basis':'仅合成坐标算例，非专业计量规则'}


def payload(eid='test',mode='volume'):
    body={'name':'合成L形底板（非真实工程）','kind':'planar_regions','parameters':{},'quantity_type':mode,
          'regions':[region('主体',0,0,4,3,eid=eid),region('翼板',3,0,2,2,eid=eid),region('扣减A',1,1,1,1,'subtract',eid),region('扣减B','1.5',1,1,1,'subtract',eid)],
          'multiplicity':'1','multiplicity_evidence_id':eid,'coordinate_basis':'同一设计原点，X向右、Y向上；坐标仅用于合成测试',
          'coordinate_evidence_id':eid,'measurement_basis':'按实际矩形并集计量，扣减重叠部分仅扣一次；无专业效力','basis_evidence_id':eid}
    if mode=='volume':body.update(thickness='200',thickness_unit='mm',thickness_evidence_id=eid)
    return body


def setup(clients):
    e,r=clients['estimator'],clients['reviewer'];pid=create(e)['id']
    upload(e,pid,'仅合成：共同原点X右Y上；主体(0,0)4m×3m，翼板(3,0)2m×2m；扣减A(1,1)1m×1m，扣减B(1.5,1)1m×1m；统一厚200mm；1组。')
    return e,r,pid,payload(get(e,pid)['evidence'][0]['id'])


def test_overlapping_additions_and_deductions_are_unioned_once():
    body=payload();before=copy.deepcopy(body);result=calculate(body);s=result['planar_summary']
    assert body==before
    assert Decimal(s['naive_add_area'])==16 and Decimal(s['add_union_area'])==14 and Decimal(s['add_overlap_area'])==2
    assert Decimal(s['naive_deduct_area'])==2 and Decimal(s['deduct_union_area'])==Decimal('1.5') and Decimal(s['deduct_overlap_area'])==Decimal('.5')
    assert Decimal(s['net_area'])==Decimal('12.5') and Decimal(result['quantity'])==Decimal('2.5')
    assert result['unit']=='m³' and len(result['planar_intersections'])==2
    assert sum(Decimal(x['net_area']) for x in result['planar_strips'])==Decimal('12.5')
    with localcontext() as ctx:
        ctx.prec=80
        assert ast_value(result['formula_ast'],result['dimensions'])==Decimal(result['quantity'])


def test_multiple_coverage_is_not_the_sum_of_pairwise_intersections():
    body=payload(mode='area');body['regions']=[region(str(i),0,0,2,2) for i in range(3)]+[region('hole'+str(i),0,0,1,1,'subtract') for i in range(3)]
    r=calculate(body);s=r['planar_summary']
    assert Decimal(s['add_overlap_area'])==8 and Decimal(s['deduct_overlap_area'])==2
    assert sum(Decimal(x['area']) for x in r['planar_intersections'] if x['operation']=='add')==12
    assert Decimal(r['quantity'])==3


def test_negative_coordinates_mixed_units_translation_and_repeat_count():
    body=payload(mode='area');body['regions']=[region('left',-2000,-1000,4000,3000,unit='mm'),region('right',100,-100,200,200,unit='cm')];body['multiplicity']='2'
    r=calculate(body)
    assert Decimal(r['quantity'])==28
    assert r['planar_regions'][0]['normalized_parameters']['x']=='-2'
    assert Decimal(r['planar_summary']['add_union_area'])==14
    body['regions']=[region('left',98,99,4,3),region('right',101,99,2,2)]
    assert Decimal(calculate(body)['quantity'])==28


def test_touching_edges_are_not_overlap_and_full_deduction_can_be_zero():
    body=payload(mode='area');body['regions']=[region('a',0,0,1,1),region('b',1,0,1,1),region('c',0,1,1,1)]
    r=calculate(body);assert not r['planar_intersections'] and Decimal(r['quantity'])==3
    body['regions'].append(region('deduct-all',0,0,2,1,'subtract'));assert Decimal(calculate(body)['quantity'])==1
    body['regions'].append(region('deduct-top',0,1,1,1,'subtract'));assert Decimal(calculate(body)['quantity'])==0


def test_deduction_can_cross_the_joint_of_two_additive_regions():
    body=payload(mode='area');body['regions']=[region('a',0,0,2,2),region('b',2,0,2,2),region('ded',1,0,2,2,'subtract')]
    assert Decimal(calculate(body)['quantity'])==4


def test_union_matches_independent_integer_cell_oracle():
    rng=random.Random(2501)
    for run in range(25):
        body=payload(mode='area');body['regions']=[];cells=set();cut=set()
        for i in range(rng.randint(1,8)):
            x,y=rng.randint(-4,4),rng.randint(-4,4);l,w=rng.randint(1,4),rng.randint(1,4)
            body['regions'].append(region('add'+str(i),x,y,l,w))
            cells|={(a,b) for a in range(x,x+l) for b in range(y,y+w)}
        for i in range(rng.randint(0,6)):
            x,y=rng.choice(sorted(cells));body['regions'].append(region('ded'+str(i),x,y,1,1,'subtract'));cut.add((x,y))
        rng.shuffle(body['regions']);r=calculate(body)
        assert Decimal(r['quantity'])==len(cells-cut)
        with localcontext() as ctx:
            ctx.prec=80
            assert ast_value(r['formula_ast'],r['dimensions'])==len(cells-cut)


@pytest.mark.parametrize('case',['outside','gap','all_deduct','negative_length','zero_width','nan','unit','precision','duplicate','missing_proof','too_many','area_thickness','count','overflow','unknown_field'])
def test_invalid_models_fail_explicitly(case):
    body=payload();r=body['regions'][0]
    if case=='outside':body['regions'][2]['parameters']['x']='20'
    elif case=='gap':body['regions']=[region('a',0,0,1,1),region('b',2,0,1,1),region('ded',0,0,3,1,'subtract')]
    elif case=='all_deduct':
        for x in body['regions']:x['operation']='subtract'
    elif case=='negative_length':r['parameters']['length']='-1'
    elif case=='zero_width':r['parameters']['width']='0'
    elif case=='nan':r['parameters']['x']='NaN'
    elif case=='unit':r['unit']='m²'
    elif case=='precision':r['unit']='mm';r['parameters']['x']='.000000000000000001'
    elif case=='duplicate':body['regions'][1]['name']=r['name']
    elif case=='missing_proof':del r['evidence']['x']
    elif case=='too_many':body['regions']*=8
    elif case=='area_thickness':body['quantity_type']='area'
    elif case=='count':body['multiplicity']='1.5'
    elif case=='overflow':body['regions']=[region('huge',0,0,'1e12','1e12')]
    else:r['parameters']['rotation']='30'
    with pytest.raises(DomainError):calculate(body)


def test_preview_approval_boq_and_all_exports(clients):
    e,r,pid,body=setup(clients);base='/api/v1/projects/'+pid;before=get(e,pid)
    response=e.post(base+'/components/preview',json=body)
    assert response.status_code==200,response.text
    assert get(e,pid)==before
    comp=e.post(base+'/components',json=body)
    assert comp.status_code==200,comp.text
    assert e.post(base+'/takeoffs/calculate').status_code==409
    approve(r,pid,'components',comp.json())
    response=e.post(base+'/takeoffs/calculate');assert response.status_code==200,response.text
    book=response.json()[0];approve(r,pid,'takeoffs',book)
    assert Decimal(e.post(base+'/boq/generate').json()[0]['quantity'])==Decimal('2.5')
    result=e.get(base+'/exports?unpriced=true&format=json').json()
    assert result['unpriced']['takeoffs'][0]['coordinate_basis']==body['coordinate_basis']
    exported=e.get(base+'/exports?unpriced=true&format=xlsx');assert exported.status_code==200,exported.text
    wb=load_workbook(io.BytesIO(exported.content),data_only=False)
    assert {'区域并集与扣减','平面区域与坐标','区域条带计算过程','区域相交定位','算量原始单位换算'}<=set(wb.sheetnames)
    assert all(c.data_type!='f' for s in wb for row in s for c in row)
    pdf=e.get(base+'/exports?unpriced=true&format=pdf');assert pdf.status_code==200,pdf.text
    text='\n'.join(page.extract_text() for page in PdfReader(io.BytesIO(pdf.content)).pages)
    assert '区域并集与扣减' in text and '区域条带计算过程' in text


def test_concurrent_revision_atomicity_and_old_book_invalidation(clients):
    e,r,pid,body=setup(clients);base='/api/v1/projects/'+pid
    comp=e.post(base+'/components',json=body).json();approve(r,pid,'components',comp)
    book=e.post(base+'/takeoffs/calculate').json()[0];approve(r,pid,'takeoffs',book)
    before=get(e,pid);assert e.put(base+'/components/'+comp['id'],json=body).status_code==409
    body['object_hash']=e.get(base+'/hashes').json()['components'][comp['id']]
    bad=copy.deepcopy(body);bad['regions'][2]['parameters']['x']='100'
    assert e.put(base+'/components/'+comp['id'],json=bad).status_code==422 and get(e,pid)==before
    body['thickness']='300'
    revised=e.put(base+'/components/'+comp['id'],json=body)
    assert revised.status_code==200,revised.text
    assert revised.json()['revision']==2 and Decimal(revised.json()['preview']['quantity'])==Decimal('3.75')
    assert get(e,pid)['takeoffs'][0]['state']=='stale'
    assert e.put(base+'/components/'+comp['id'],json=body).status_code==409


@pytest.mark.parametrize('gate',['save','component_review','calculate','book_review'])
def test_corrupt_original_blocks_all_processing_gates(clients,gate):
    e,r,pid,body=setup(clients);base='/api/v1/projects/'+pid
    if gate!='save':comp=e.post(base+'/components',json=body).json()
    if gate in ('calculate','book_review'):approve(r,pid,'components',comp)
    if gate=='book_review':book=e.post(base+'/takeoffs/calculate').json()[0]
    doc=get(e,pid)['documents'][0];(store.ROOT/'originals'/doc['hash']).write_bytes(b'corrupt')
    if gate=='save':assert e.post(base+'/components',json=body).status_code==409
    elif gate=='component_review':approve(r,pid,'components',comp,expected=409)
    elif gate=='calculate':assert e.post(base+'/takeoffs/calculate').status_code==409
    else:approve(r,pid,'takeoffs',book,expected=409)


def test_permissions_foreign_proof_and_superseded_source(clients):
    e,r,pid,body=setup(clients);base='/api/v1/projects/'+pid
    assert r.post(base+'/components/preview',json=body).status_code==403
    assert clients['other'].post(base+'/components',json=body).status_code==404
    pid2=create(e)['id'];upload(e,pid2,'其他工程');eid=get(e,pid2)['evidence'][0]['id']
    foreign=copy.deepcopy(body);foreign['regions'][0]['evidence']['x']=eid
    assert e.post(base+'/components',json=foreign).status_code==404
    comp=e.post(base+'/components',json=body).json();approve(r,pid,'components',comp)
    upload(e,pid,'新版本坐标',name='fixture.txt')
    assert get(e,pid)['components'][0]['state']=='stale'
    approve(r,pid,'components',comp,expected=409)


def test_cannot_drop_structured_fields_or_mutate_calculation_book(clients):
    e,r,pid,body=setup(clients);base='/api/v1/projects/'+pid
    comp=e.post(base+'/components',json=body).json();approve(r,pid,'components',comp)
    e.post(base+'/takeoffs/calculate')
    with store.transaction() as c:
        p=json.loads(c.execute('SELECT data FROM projects WHERE id=?',(pid,)).fetchone()['data']);p['takeoffs'][0]['planar_summary']['net_area']='999';store.save(c,p)
    approve(r,pid,'takeoffs',get(e,pid)['takeoffs'][0],expected=409)
    changed={'kind':'slab','parameters':{},'object_hash':e.get(base+'/hashes').json()['components'][comp['id']]}
    assert e.put(base+'/components/'+comp['id'],json=changed).status_code==409


def test_new_planar_geometry_never_rewrites_quote_snapshot(clients):
    e,r,pid,_,q=priced_project(clients);base='/api/v1/projects/'+pid;old=copy.deepcopy(q)
    eid=get(e,pid)['evidence'][0]['id'];response=e.post(base+'/components',json=payload(eid))
    assert response.status_code==200,response.text
    assert get(e,pid)['quotes'][0]==old and q['id'] in get(e,pid)['stale_quotes']
