"""Synthetic multi-circle exclusions: independent integrals, exact topology and workflow."""
import copy,io,random
from decimal import Decimal,localcontext
from fractions import Fraction
import pytest
from openpyxl import load_workbook
from pypdf import PdfReader
from backend import store
from backend.engine import DomainError
from backend.compound_takeoff import PI
from backend.foundation_horizontal_formwork import calculate
from backend.foundation_circular_holes_windows import VERSION,METHOD,validate,separation,topology
from tests.test_foundation_circular_windows import body as cb,config as cc,independent_area,replay as old_replay
from tests.test_foundation_annular_windows import config as ac,setup as old_setup
from tests.test_foundation_eccentric_windows import config as ec
from tests.test_foundation_horizontal_windows import source,window
from tests.test_foundation_hole_end_formwork import source as end_source
from tests.test_foundation_multi_windows import config as mc,part
from tests.test_foundation_nested_windows import config as nc,block,box
from tests.test_foundation_perforated_windows import config as pc
from tests.test_foundation_openings import hole
from tests.test_foundation_voids import rectangle
from tests.test_foundation_recesses import interval
from tests.test_post_foundation import equal_body,fixed_body,fixed_added,equal_added,fixed_adjusted,equal_adjusted,combined
from tests.test_polygon_takeoff import post,mutate
from tests.test_workflow import get,approve,upload,create,priced_project
from tests.test_route_takeoff import evaluate


def inner(eid,x,y,d,name):
    p=dict(x=str(x),y=str(y),diameter=str(d));return dict(name=name,parameters=p,evidence={k:eid for k in p},measurement_basis='合成不支模圆区，非真实工程',evidence_id=eid)


def config(eid='test',x='120',y='10',diameter='240',holes=None,unit='mm'):
    c=cc(eid,x,y,diameter,unit);c.update(shape='circular_with_holes',holes=holes if holes is not None else [inner(eid,'95','10','50','左内圆'),inner(eid,'155','20','40','右内圆')]);return c


def body(s,eid='test',face='bottom',**kw):
    b=cb(s,eid,face);b['contact_ranges'][0]['plan_range']=config(eid,**kw);return b


def punctured(**kw):return dict(plan_range=config(**kw))


def replay(r):
    old_replay(r)
    for d in r['contact_ranges']:
        if not d.get('circular_holes_window'):continue
        for h in d['window_circular_holes']:
            assert evaluate(h['formula_asts']['area'],r['dimensions'])==Decimal(h['area'])
            q=h['containment']
            for k,n in q['formula_asts'].items():assert Fraction(evaluate(n,r['dimensions']))==Fraction(q[k])
            assert Fraction(q['center_distance_squared'])<Fraction(q['radius_gap_squared'])
        for q in d['window_hole_checks']:
            for k,n in q['formula_asts'].items():assert Fraction(evaluate(n,r['dimensions']))==Fraction(q[k])
            assert Fraction(q['center_distance_squared'])>Fraction(q['radius_sum_squared'])


@pytest.mark.parametrize('shape',['rectangular','circular','stepped_rectangular','stepped_circular'])
@pytest.mark.parametrize('face',['top','bottom','hole_top','hole_bottom'])
@pytest.mark.parametrize('make',[equal_body,fixed_body,fixed_added,equal_added,fixed_adjusted,equal_adjusted,combined,lambda s:combined(s,equal=True)])
def test_shapes_ends_counts_immutable_and_all_ast_replay(shape,face,make):
    s=end_source(shape,make) if face.startswith('hole') else source(shape,make);b=body(s,face=face);before=copy.deepcopy((s,b));r=calculate(b,s);d=r['contact_ranges'][0]
    assert before==(s,b) and r['rule_version']==VERSION and d['circular_holes_window'] and len(d['window_circular_holes'])==2
    with localcontext() as ctx:
        ctx.prec=80;assert Decimal(d['window_area'])==Decimal(PI)*Decimal('.013375');assert Decimal(d['total_net_area'])==Decimal(d['single_net_area'])*Decimal(d['total_count'])
    assert r['source_evidence']['horizontal_formwork.1.window_holes.2.basis']=='test';replay(r)


@pytest.mark.parametrize('shape',['rectangular','circular','stepped_rectangular','stepped_circular'])
@pytest.mark.parametrize('face',['top','bottom','hole_bottom','hole_top'])
def test_outer_minus_each_actual_disk_full_gross_cut_net(shape,face):
    s=end_source(shape) if face.startswith('hole') else source(shape);r=calculate(body(s,face=face),s);outer=calculate(cb(s,face=face),s);holes=[calculate(cb(s,face=face,**h['parameters']),s) for h in config()['holes']]
    with localcontext() as ctx:
        ctx.prec=80
        for k in ['single_gross_area','single_opening_area','single_net_area','total_net_area']:
            assert abs(Decimal(r['contact_ranges'][0][k])-(Decimal(outer['contact_ranges'][0][k])-sum(Decimal(h['contact_ranges'][0][k]) for h in holes)))<Decimal('1e-73')
    replay(r)


@pytest.mark.parametrize('shape',['rectangular','circular','stepped_rectangular','stepped_circular'])
@pytest.mark.parametrize('offset',[-350,-160,0,160,350])
def test_independent_numeric_physical_hole_intersections(shape,offset):
    s=source(shape);s['foundation_items'][0]['through_holes']['holes']=[{**rectangle(x='150',length='100',width='60'),'vertical_range':interval('0','150')}]
    hs=[inner('test',offset-50,10,60,'W1'),inner('test',offset+55,45,70,'W2')];r=calculate(body(s,x=offset,y='30',diameter='360',holes=hs),s);d=r['contact_ranges'][0];disks=[(offset/1000,.03,.18)];boxes=[]
    if shape in ('rectangular','stepped_rectangular'):boxes=[(-.2,.2,-.25,.25)] if shape=='rectangular' else [(-.3,.3,-.25,.25)]
    else:disks.append((0.,0.,.3))
    expected=independent_area(disks,boxes,excluded_disks=[((offset-50)/1000,.01,.03),((offset+55)/1000,.045,.035)],excluded_boxes=[(.1,.2,-.03,.03)])
    assert abs(float(d['single_net_area'])-expected)<3e-10;replay(r)


@pytest.mark.parametrize('shape',['stepped_rectangular','stepped_circular'])
@pytest.mark.parametrize('x',['0','170','280','700'])
def test_platform_independent_integral(shape,x):
    s=source(shape);s['foundation_items'][0].pop('through_holes',None);hs=[inner('test',Decimal(x)-30,-10,40,'W1'),inner('test',Decimal(x)+30,15,40,'W2')];r=calculate(body(s,face='step',x=x,y='0',diameter='220',holes=hs),s)
    ds=[(float(x)/1000,0.,.11)];ex=[((float(x)-30)/1000,-.01,.02),((float(x)+30)/1000,.015,.02)]
    expected=independent_area(ds,[(-.3,.3,-.25,.25)],excluded_disks=ex,excluded_boxes=[(-.2,.2,-.15,.15)]) if shape=='stepped_rectangular' else independent_area(ds+[(0.,0.,.3)],excluded_disks=ex+[(0.,0.,.2)])
    assert abs(float(r['contact_ranges'][0]['single_net_area'])-expected)<3e-10;replay(r)


@pytest.mark.parametrize('unit,factor',[('mm','1'),('cm','.1'),('dm','.01'),('m','.001'),('km','.000001')])
def test_shared_units_and_low_caller_precision(unit,factor):
    s=source();b=body(s);c=b['contact_ranges'][0]['plan_range'];c['unit']=unit
    for obj in [c]+c['holes']:obj['parameters']={k:str(Decimal(v)*Decimal(factor)) for k,v in obj['parameters'].items()}
    with localcontext() as ctx:ctx.prec=6;r=calculate(b,s)
    assert r['quantity']==calculate(body(s),s)['quantity'];replay(r)


@pytest.mark.parametrize('gap',['-0.000000001','0','0.000000001'])
def test_exact_strict_internal_and_pair_tangency(gap):
    g=Decimal(gap);c=config(x='0',y='0',diameter='200',holes=[inner('test',Decimal(80)+g,0,40,'W1'),inner('test',-40,0,20,'W2')])
    if g<0:validate(c)
    else:
        with pytest.raises(DomainError,match='严格'):validate(c)
    c=config(x='0',y='0',diameter='200',holes=[inner('test',0,0,40,'W1'),inner('test',Decimal(40)+g,0,40,'W2')])
    if g>0:validate(c)
    else:
        with pytest.raises(DomainError,match='严格'):validate(c)


@pytest.mark.parametrize('fault',['one','five','overlap','contain','touch','name','empty_name','zero','bool','nan','missing','extra','offset','evidence','basis','basis_proof','unit','count','confirm'])
def test_schema_each_hole_own_original_and_limits(fault):
    c=config();h=c['holes'][1]
    if fault=='one':c['holes']=c['holes'][:1]
    elif fault=='five':c['holes']*=3
    elif fault in ('overlap','contain','touch'):h['parameters'].update(x={'overlap':'100','contain':'95','touch':'140'}[fault],y='10')
    elif fault=='name':h['name']=c['holes'][0]['name']
    elif fault=='empty_name':h['name']=''
    elif fault in ('zero','bool','nan'):h['parameters']['diameter']={'zero':'0','bool':True,'nan':'NaN'}[fault]
    elif fault=='missing':del h['parameters']['y']
    elif fault=='extra':h['area']='123'
    elif fault=='offset':h['parameters']['offset_x']='20'
    elif fault=='evidence':h['evidence']['y']=''
    elif fault=='basis':h['measurement_basis']=''
    elif fault=='basis_proof':h['evidence_id']=''
    elif fault=='unit':h['unit']='mm'
    elif fault=='count':h['count']='2'
    else:c['scope_confirmed']=1
    with pytest.raises(DomainError):validate(c)


def test_four_holes_order_does_not_change_quantity():
    hs=[inner('test',x,y,20,str(i)) for i,(x,y) in enumerate([(-40,-40),(40,-40),(40,40),(-40,40)])];s=source();b=body(s,x='0',y='0',diameter='200',holes=hs);a=calculate(b,s);b['contact_ranges'][0]['plan_range']['holes'].reverse();z=calculate(b,s)
    assert a['quantity']==z['quantity'] and a['contact_ranges'][0]['single_net_area']==z['contact_ranges'][0]['single_net_area'] and len(a['contact_ranges'][0]['window_hole_checks'])==6;replay(a);replay(z)


def test_100_seeded_curved_pairs_independent_numerical_intersections():
    rng=random.Random(11501)
    for _ in range(100):
        x,y,X,Y=[rng.randint(-140,140) for _ in range(4)];R,S=rng.randint(30,100),rng.randint(30,100)
        hs=[inner('test',x-R//3,y,R//4,'W1'),inner('test',x+R//3,y,R//4,'W2')];a=punctured(x=x,y=y,diameter=2*R,holes=hs);kind=rng.randrange(4);ex=[(float(h['parameters']['x'])/1000,y/1000,R//4/2000) for h in hs]
        if kind==0:b=dict(plan_range=cc(x=X,y=Y,diameter=2*S))
        elif kind==1:b=dict(plan_range=ac(x=X,y=Y,outer=2*S,inner=S//2));ex.append((X/1000,Y/1000,S//2/2000))
        elif kind==2:b=dict(plan_range=ec(x=X,y=Y,outer=2*S,inner_x=X+1,inner_y=Y-1,inner=S//2));ex.append(((X+1)/1000,(Y-1)/1000,S//2/2000))
        else:
            b=punctured(x=X,y=Y,diameter=2*S,holes=[inner('test',X-S//3,Y,S//4,'W1'),inner('test',X+S//3,Y,S//4,'W2')]);ex.extend([(float(h['parameters']['x'])/1000,Y/1000,S//4/2000) for h in b['plan_range']['holes']])
        expected=independent_area([(x/1000,y/1000,R/1000),(X/1000,Y/1000,S/1000)],excluded_disks=ex)>1e-11
        assert separation(a,b)['overlap']==expected and separation(b,a)['overlap']==expected


def test_120_rectangles_independent_per_hole_extrema_and_edge_witness():
    rng=random.Random(11502)
    for _ in range(120):
        L,B=rng.randint(-180,180),rng.randint(-180,180);X,Y=L+rng.randint(1,80),B+rng.randint(1,80);minimum=max(L,0,-X)**2+max(B,0,-Y)**2
        maxima=[max((L-x)**2,(X-x)**2)+max(B*B,Y*Y) for x in (-35,35)];expected=minimum<10000 and all(m>400 for m in maxima)
        a=punctured(x=0,y=0,diameter=200,holes=[inner('test',-35,0,40,'W1'),inner('test',35,0,40,'W2')]);q=separation(a,dict(plan_range=window('test',str(L),str(X),str(B),str(Y))));assert q['overlap']==expected
        if not expected:
            p=q['components'][0];assert Fraction(p['minimum_outer_squared'])*1000000==minimum
            assert [Fraction(m['maximum_squared'])*1000000 for m in p['inner_maxima']]==maxima
            for w in p['edge_witnesses']:
                nx,ny=Fraction(w['nearest_x']),Fraction(w['nearest_y']);assert Fraction(w['distance_squared'])==nx*nx+ny*ny


def test_disconnected_blocks_cannot_mix_holes_or_extrema_and_nested_island():
    a=punctured(x=0,y=0,diameter=200,holes=[inner('test',-35,0,40,'W1'),inner('test',35,0,40,'W2')])
    c=mc([part(box(-40,-30,-5,5),name='W1内'),part(box(30,40,-5,5),name='W2内'),part(box(150,170,-10,10),name='外圆外')]);q=separation(a,dict(plan_range=c));assert not q['overlap'] and [p['hole_index'] for p in q['components']]==[1,2,None]
    assert separation(a,dict(plan_range=window('test','-40','40','-5','5')))['overlap']
    c=pc(outer=box(-150,150,-150,150),inners=[box(-110,110,-110,110)]);assert not separation(a,dict(plan_range=c))['overlap']
    c=nc([block(box(-150,150,-150,150),holes=[box(-110,110,-110,110)]),block(box(30,40,-5,5),name='W2内岛',parent=dict(block_index=1,hole_index=1))]);assert not separation(a,dict(plan_range=c))['overlap']
    c['parts'][1]['parameters']=part(box(-5,5,-5,5))['parameters'];assert separation(a,dict(plan_range=c))['overlap']


@pytest.mark.parametrize('gap',['0','0.000000001'])
def test_zero_proof_all_candidates_or_positive_tiny_gap(gap):
    v=str(Decimal(25)+Decimal(gap));hs=[{**hole(x='0',diameter='80'),'vertical_range':interval('100','300')},{**rectangle(name='H2',x='-'+v,length='50',width='100'),'vertical_range':interval('300','500')},{**rectangle(name='H3',x=v,length='50',width='100'),'vertical_range':interval('300','500')}]
    s=end_source(holes=hs);r=calculate(body(s,face='hole_top',x=0,y=0,diameter=60,holes=[inner('test',-10,0,8,'W1'),inner('test',10,0,8,'W2')]),s);d=r['contact_ranges'][0]
    assert ('polygon_zero_contact' in d)==(gap=='0') and (Decimal(d['single_net_area'])==0)==(gap=='0');replay(r)
    if gap=='0':assert all(not p['remaining'] for p in d['polygon_zero_contact']['net_trace']['partitions'])


def setup(clients,pid=None,zero=False):
    e,r,pid,s,b,eid=old_setup(clients,pid,zero);b['contact_ranges'][0]['plan_range']=config(eid,x=0 if zero else 196,y=0 if zero else 8,diameter=40 if zero else 42,holes=[inner(eid,-8,0,6,'W1'),inner(eid,8,2,6,'W2')] if zero else [inner(eid,190,6,8,'W1'),inner(eid,201,10,10,'W2')]);return e,r,pid,s,b,eid


def test_workflow_roles_foreign_w2_history_concurrency_source_stale(clients):
    e,r,pid,s,b,eid=setup(clients);before=get(e,pid);post(e,pid,b,'/components/preview');assert get(e,pid)==before;post(r,pid,b,expected=403)
    foreign=create(e)['id'];upload(e,foreign,'异工程合成');bad=copy.deepcopy(b);bad['contact_ranges'][0]['plan_range']['holes'][1]['evidence']['x']=get(e,foreign)['evidence'][-1]['id'];post(e,pid,bad,expected=404)
    c=post(e,pid,b).json();approve(r,pid,'components',c);c=next(x for x in get(e,pid)['components'] if x['id']==c['id']);b['object_hash']=store.digest(c);b['contact_ranges'][0]['plan_range']['holes'][1]['parameters']['y']='11';res=e.put(f'/api/v1/projects/{pid}/components/'+c['id'],json=b);assert res.status_code==200,res.text;assert e.put(f'/api/v1/projects/{pid}/components/'+c['id'],json=b).status_code==409
    assert get(e,pid)['component_history'][-1]==c;approve(r,pid,'components',res.json());s['object_hash']=store.digest(s);s['foundation_items'][0]['levels'][1]['parameters']['diameter']='420';res=e.put(f'/api/v1/projects/{pid}/components/'+s['id'],json=s);assert res.status_code==200;approve(r,pid,'components',res.json());assert next(x for x in get(e,pid)['components'] if x['id']==c['id'])['state']=='stale'


def test_w2_claim_allowed_but_tiny_overlap_or_frame_mismatch_refused(clients):
    e,r,pid,s,b,eid=setup(clients);post(e,pid,b);other=copy.deepcopy(b);other['name']='合成W2另项';other['contact_ranges'][0]['plan_range']=cc(eid,x=201,y=10,diameter=10);post(e,pid,other)
    bad=copy.deepcopy(other);bad['contact_ranges'][0]['plan_range']['parameters']['x']='201.000000001';before=get(e,pid);post(e,pid,bad,expected=409);assert get(e,pid)==before
    bad=copy.deepcopy(other);bad['contact_ranges'][0]['plan_range']['coordinate_basis']='其他基准';post(e,pid,bad,expected=409)


@pytest.mark.parametrize('fault',['input','normalized','area','containment','pair','pair_ast','basis','zero','expression'])
def test_saved_inputs_topology_or_calculation_tamper_rejected(clients,fault):
    e,r,pid,s,b,eid=setup(clients,zero=True);c=post(e,pid,b).json()
    def forge(p):
        c=p['components'][-1];d=c['preview']['contact_ranges'][0];h=d['window_circular_holes'][1]
        if fault=='input':c['contact_ranges'][0]['plan_range']['holes'][1]['parameters']['x']='9'
        elif fault=='normalized':h['normalized_parameters']['x']='999'
        elif fault=='area':h['area']='1'
        elif fault=='containment':h['containment']['center_distance_squared']='999'
        elif fault=='pair':d['window_hole_checks'][0]['center_distance_squared']='999'
        elif fault=='pair_ast':d['window_hole_checks'][0]['formula_asts']['radius_sum_squared']={'constant':'0'}
        elif fault=='basis':c['contact_ranges'][0]['plan_range']['holes'][1]['measurement_basis']='篡改用途'
        elif fault=='zero':d['polygon_zero_contact']['net_trace']['partitions'][0]['remaining']=[{}]
        else:d['window_expression_definitions'][0]['value']='999'
    mutate(pid,forge);before=get(e,pid);approve(r,pid,'components',c,expected=409);assert get(e,pid)==before


@pytest.mark.parametrize('gate',['preview','save','review','calculate','book_review'])
def test_w2_basis_original_bytes_five_gates(clients,gate):
    e,r,pid,s,b,eid=setup(clients);upload(e,pid,'合成W2用途及坐标原件，非真实工程','w2-original.txt')
    for c in get(e,pid)['components']:approve(r,pid,'components',c)
    b['source_component_hash']=e.get(f'/api/v1/projects/{pid}/foundation-horizontal-formwork-candidates',params={'component_id':s['id']}).json()['source_component_hash'];b['contact_ranges'][0]['plan_range']['holes'][1]['evidence_id']=get(e,pid)['evidence'][-1]['id']
    if gate in ('review','calculate','book_review'):c=post(e,pid,b).json()
    if gate in ('calculate','book_review'):approve(r,pid,'components',c)
    if gate=='book_review':t=post(e,pid,{},'/takeoffs/calculate').json()[-1]
    doc=get(e,pid)['documents'][-1];(store.ROOT/'originals'/doc['hash']).write_bytes(b'corrupted-synthetic-w2-original')
    if gate=='preview':post(e,pid,b,'/components/preview',409)
    elif gate=='save':post(e,pid,b,expected=409)
    elif gate=='review':approve(r,pid,'components',c,expected=409)
    elif gate=='calculate':post(e,pid,{},'/takeoffs/calculate',409)
    else:approve(r,pid,'takeoffs',t,expected=409)


@pytest.mark.parametrize('fmt',['json','xlsx','pdf'])
@pytest.mark.parametrize('zero',[False,True])
def test_three_formats_topology_and_old_quote_isolation(clients,fmt,zero):
    from backend.foundation_circular_holes_reports import TITLES
    e,r,pid,_,q=priced_project(clients);e,r,pid,s,b,eid=setup(clients,pid,zero);c=post(e,pid,b).json();approve(r,pid,'components',c)
    for t in post(e,pid,{},'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
    post(e,pid,{},'/boq/generate');url=f'/api/v1/projects/{pid}/exports';res=e.get(url,params=dict(format=fmt,unpriced=True));assert res.status_code==200,res.text; titles=set(TITLES[:4] if zero else TITLES[:3])
    if fmt=='json':assert titles<={t[0] for t in res.json()['tables']}
    elif fmt=='xlsx':assert titles<=set(load_workbook(io.BytesIO(res.content)).sheetnames)
    else:assert all(any(title in p.extract_text() for p in PdfReader(io.BytesIO(res.content)).pages) for title in titles)
    assert not any(t[0] in titles for t in e.get(url,params=dict(format='json',quote_id=q['id'])).json()['tables'])


@pytest.mark.parametrize('fault',['distance_squared','hole_index','window_side'])
def test_full_window_certificate_tamper_rejected(clients,fault):
    e,r,pid,s,b,eid=setup(clients);d=copy.deepcopy(b['contact_ranges'][0]);d.update(name='W2另项',plan_range=cc(eid,x=201,y=10,diameter=10));b['contact_ranges'].append(d);c=post(e,pid,b).json()
    def forge(p):p['components'][-1]['preview']['horizontal_window_checks'][0][fault]={'distance_squared':'999','hole_index':1,'window_side':'right'}[fault]
    mutate(pid,forge);before=get(e,pid);approve(r,pid,'components',c,expected=409);assert get(e,pid)==before


def test_source_entirely_inside_one_exclusion_no_false_physical_hole_proof():
    s=end_source();r=calculate(body(s,face='hole_top',x=0,y=0,diameter=1000,holes=[inner('test',0,0,500,'W1'),inner('test',350,0,100,'W2')]),s);d=r['contact_ranges'][0];assert Decimal(d['single_gross_area'])==0 and 'polygon_zero_contact' not in d;replay(r)
