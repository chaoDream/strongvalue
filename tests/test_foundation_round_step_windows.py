"""Circular ledge windows: analytic disk/disk/window clipping and review gates."""
import copy,io,math,random
from decimal import Decimal,localcontext,ROUND_HALF_UP
import pytest
from openpyxl import load_workbook
from pypdf import PdfReader
from backend import store
from backend.engine import DomainError
from backend.railing_material import snapshot
from backend.foundation_horizontal_formwork import calculate
from backend.foundation_round_step_windows import VERSION
from backend.disk_window_intersection import VERSION as SWEEP_VERSION
from backend.foundation_horizontal_formwork_reports import ROUND_WINDOW_TITLE,STEP_WINDOW_TITLES,WINDOW_TITLES,WINDOW_CURVED_TITLE,tables
from tests.test_foundation_step_windows import body,check as rectangular_check
from tests.test_foundation_horizontal_windows import window,FRAME
from tests.test_foundation_step_formwork import source,contact,body as step_body,setup as step_setup
from tests.test_foundation_openings import hole
from tests.test_foundation_voids import rectangle
from tests.test_foundation_recesses import interval
from tests.test_post_foundation import equal_body,fixed_body,fixed_added,equal_added,fixed_adjusted,equal_adjusted,combined
from tests.test_route_takeoff import evaluate
from tests.test_polygon_takeoff import post,mutate
from tests.test_workflow import get,approve,upload,create,priced_project


def circle_source(make=equal_body,x='190',y='0',diameter='40'):
    return source('stepped_circular',make,holes=[{**hole(x=x,y=y,diameter=diameter),'vertical_range':interval('0','200')}])


def check(r):
    rectangular_check(r)
    with localcontext() as ctx:
        ctx.prec=80
        for d in r['contact_ranges']:
            for h in d['apertures']:
                for q in h.get('exposed_coverage',{}).get('terms',[]):
                    for k,a in q['formula_asts'].items():assert evaluate(a,r['dimensions'])==Decimal(q[k])
            for m in [d.get('window_curved_intersection'),d.get('window_upper_curved_intersection')]+[h.get('curved_intersection') for h in d['apertures']]:
                if m:
                    for q in m['terms']:
                        assert evaluate(q['formula_ast'],r['dimensions'])==Decimal(q['area'])
                        for k,a in q.get('formula_asts',{}).items():assert evaluate(a,r['dimensions'])==Decimal(q[k])


def adaptive(f,a,b,eps=1e-13):
    m=(a+b)/2;fa,fm,fb=f(a),f(m),f(b);v=(b-a)*(fa+4*fm+fb)/6
    def run(a,b,fa,fm,fb,v,eps,n):
        m=(a+b)/2;fl,fr=f((a+m)/2),f((m+b)/2);l=(m-a)*(fa+4*fl+fm)/6;r=(b-m)*(fm+4*fr+fb)/6;t=l+r
        if n<=0 or abs(t-v)<=15*eps:return t+(t-v)/15
        return run(a,m,fa,fl,fm,l,eps/2,n-1)+run(m,b,fm,fr,fb,r,eps/2,n-1)
    return run(a,b,fa,fm,fb,v,eps,28)


def oracle(disks,box):
    x1,x2,y1,y2=box;lo=max([x1]+[x-r for x,y,r in disks]);hi=min([x2]+[x+r for x,y,r in disks])
    if lo>=hi:return 0.
    cuts={lo,hi}
    for x,y,r in disks:
        cuts.add(x)
        for z in (y1,y2):
            if abs(z-y)<r:
                a=math.sqrt(r*r-(z-y)**2);cuts.update((x-a,x+a))
    if len(disks)==2:
        (x,y,r),(X,Y,R)=disks;dx,dy=X-x,Y-y;d=math.hypot(dx,dy)
        if abs(R-r)<d<R+r:
            a=(r*r-R*R+d*d)/(2*d);h=math.sqrt(max(0.,r*r-a*a));cuts.update((x+a*dx/d+h*dy/d,x+a*dx/d-h*dy/d))
    cuts=sorted(t for t in cuts if lo<=t<=hi)
    def f(t):
        widths=[math.sqrt(max(0.,r*r-(t-x)**2)) for x,y,r in disks];top=min([y2]+[y+w for (x,y,r),w in zip(disks,widths)]);bottom=max([y1]+[y-w for (x,y,r),w in zip(disks,widths)]);return max(0.,top-bottom)
    return sum(adaptive(f,a,b) for a,b in zip(cuts,cuts[1:]))


@pytest.mark.parametrize('make',[equal_body,fixed_body,fixed_added,equal_added,fixed_adjusted,equal_adjusted,combined,lambda s:combined(s,equal=True)])
@pytest.mark.parametrize('axis',['x','y'])
def test_eight_final_counts_two_partitions_match_whole(make,axis):
    s=circle_source(make);b=body(s,bounds=('-1000','0','-1000','1000') if axis=='x' else ('-1000','1000','-1000','0'));other=copy.deepcopy(b['contact_ranges'][0]);other['name']='另一半圆平台';other['plan_range']=window(x1='0' if axis=='x' else '-1000',x2='1000',y1='-1000' if axis=='x' else '0',y2='1000');b['contact_ranges'].append(other);before=copy.deepcopy((s,b));r=calculate(b,s);assert (s,b)==before and r['quantity']==calculate(step_body(s),s)['quantity'] and r['rule_version']==VERSION;check(r)


@pytest.mark.parametrize('x,y,bounds',[
    ('190','0',('170','220','0','30')),('-190','0',('-220','-170','0','30')),
    ('0','190',('0','30','170','220')),('0','-190',('0','30','-220','-170')),
    ('114','152',('100','140','130','180')),('-114','152',('-140','-100','130','180')),
    ('114','-152',('100','140','-180','-130')),('-114','-152',('-140','-100','-180','-130'))])
def test_both_arc_directions_and_y_offset_independent_integration(x,y,bounds):
    s=circle_source(x=x,y=y);r=calculate(body(s,bounds=bounds),s);h=r['contact_ranges'][0]['apertures'][0];box=tuple(float(v)/1000 for v in bounds);expected=oracle([(float(x)/1000,float(y)/1000,.02),(0,0,.2)],box);assert abs(float(h['covered_area'])-expected)<3e-12;check(r)


def test_half_lens_symmetry_and_new_interval_trace():
    s=circle_source();b=body(s,bounds=('170','220','0','30'));r=calculate(b,s);h=r['contact_ranges'][0]['apertures'][0];whole=calculate(step_body(s),s)['contact_ranges'][0]['apertures'][0];assert abs(float(h['covered_area'])*2-float(whole['covered_area']))<2e-18;m=h['curved_intersection'];assert m['rule_version']==SWEEP_VERSION and len(m['terms'])>=2
    with localcontext() as ctx:
        ctx.prec=80
        assert abs(sum(Decimal(t['area']) for t in m['terms'])-Decimal(h['covered_area']))<Decimal('1e-80')
        assert all(Decimal(t['x_from'])<Decimal(t['x_to']) and Decimal(t['area'])>0 and t['upper_boundary'] and t['lower_boundary'] for t in m['terms'])
    check(r)


@pytest.mark.parametrize('x,y,diameter',[('0','0','40'),('180','0','40'),('220','0','40'),('250','0','40'),('0','0','500'),('0','0','600'),('190','0','40')])
@pytest.mark.parametrize('bounds',[('-400','400','-400','400'),('170','300','0','200'),('-100','100','-100','100')])
def test_contained_tangent_disjoint_cover_zero_and_full_window(x,y,diameter,bounds):
    s=circle_source(x=x,y=y,diameter=diameter)
    if diameter=='600':
        with pytest.raises(DomainError,match='圆孔实际直径'):calculate(body(s,bounds=bounds),s)
        return
    r=calculate(body(s,bounds=bounds),s);box=tuple(float(v)/1000 for v in bounds);h=r['contact_ranges'][0]['apertures'][0];expected=oracle([(float(x)/1000,float(y)/1000,float(diameter)/2000),(0,0,.2)],box);assert abs(float(h['covered_area'])-expected)<3e-12
    check(r)


@pytest.mark.parametrize('kind',['rect','circle'])
def test_160_random_windows_independent_numerical_oracle(kind):
    rng=random.Random(10501);maximum=0.
    for i in range(80):
        x,y=[(190,0),(114,152),(0,190),(-114,152),(-190,0),(114,-152),(210,0),(160,120)][i%8];rad=[10,20,35][i%3];bounds=(*sorted(rng.sample(range(x-70,x+71),2)),*sorted(rng.sample(range(y-70,y+71),2)));box=tuple(v/1000 for v in bounds)
        h=hole(x=str(x),y=str(y),diameter=str(rad*2)) if kind=='circle' else rectangle(x=str(x),y=str(y),length='80',width='50');s=source('stepped_circular',holes=[{**h,'vertical_range':interval('0','200')}]);b=body(s,bounds=tuple(map(str,bounds)));r=calculate(b,s);d=r['contact_ranges'][0];actual=d['apertures'][0]
        if kind=='circle':common=oracle([(x/1000,y/1000,rad/1000),(0,0,.2)],box);inside=oracle([(x/1000,y/1000,rad/1000)],box)
        else:
            clipped=(max(box[0],x/1000-.04),min(box[1],x/1000+.04),max(box[2],y/1000-.025),min(box[3],y/1000+.025));inside=max(0.,clipped[1]-clipped[0])*max(0.,clipped[3]-clipped[2]);common=oracle([(0,0,.2)],clipped) if inside else 0.
        gross=oracle([(0,0,.3)],box)-oracle([(0,0,.2)],box);expected=(gross-inside+common)*22;err=abs(float(r['horizontal_formwork_summary']['unrounded_quantity'])-expected);maximum=max(maximum,err);assert err<3e-10,(i,kind,bounds,err);assert abs(float(actual['covered_area'])-common)<3e-12;check(r)


@pytest.mark.parametrize('unit,factor',[('mm','1'),('cm','.1'),('dm','.01'),('m','.001'),('km','.000001')])
@pytest.mark.parametrize('source_unit,source_factor',[('mm','1'),('cm','.1'),('dm','.01'),('m','.001'),('km','.000001')])
def test_25_unit_combinations_and_low_calling_precision(unit,factor,source_unit,source_factor):
    s=circle_source();item=s['foundation_items'][0];item['length_unit']=source_unit
    for t in item['levels']:
        for k in t['parameters']:t['parameters'][k]=str(Decimal(t['parameters'][k])*Decimal(source_factor))
    b=body(s,bounds=('170','220','0','30'));c=b['contact_ranges'][0]['plan_range'];c['unit']=unit
    for k in c['parameters']:c['parameters'][k]=str(Decimal(c['parameters'][k])*Decimal(factor))
    with localcontext() as ctx:ctx.prec=6;r=calculate(b,s)
    assert r==calculate(b,s) and r['quantity']=='0.011';check(r)


@pytest.mark.parametrize('delta',['0.00000001','0.000000000001','0.000000000000001'])
def test_near_outer_tangent_no_geometric_snap(delta):
    # A tiny lens is cut at y=0; its positive half must remain measurable.
    s=circle_source(x=str(Decimal('220')-Decimal(delta)));b=body(s,bounds=('190','250','0','30'));r=calculate(b,s);h=r['contact_ranges'][0]['apertures'][0];assert Decimal(h['covered_area'])>0;whole=calculate(step_body(s),s)['contact_ranges'][0]['apertures'][0];assert abs(Decimal(h['covered_area'])/Decimal(whole['covered_area'])-Decimal('.5'))<Decimal('1e-20');check(r)


def test_three_circle_levels_same_xy_and_whole_claim_guards():
    s=circle_source();third=copy.deepcopy(s['foundation_items'][0]['levels'][1]);third['parameters'].update(diameter='200',height='100');s['foundation_items'][0]['levels'].append(third);b=body(s,bounds=('0','300','0','300'));other=body(s,bounds=('0','300','0','300'),level=2)['contact_ranges'][0];other['name']='合成上平台';b['contact_ranges'].append(other);r=calculate(b,s);assert r['horizontal_window_checks']==[];check(r)
    bad=copy.deepcopy(b);bad['contact_ranges'].append({**contact(),'name':'整平台重复'})
    with pytest.raises(DomainError):calculate(bad,s)


def setup(clients,pid=None):
    e,r,pid,s,_,_=step_setup(clients,pid,shape='stepped_circular')
    upload(e,pid,FRAME+'仅合成22个同心圆台阶基础，下Φ600高200mm，上Φ400高400mm。H1圆孔Φ40中心(190,0)mm，H2矩形100×60中心(-200,0)mm，两孔Z0至200mm。体积阈值100m²不扣计量混凝土，物理孔仍扣模板，材料79.200kg保持。V1平台窗口X170至300、Y0至200mm，毛=下层圆与窗口交集减上层圆与窗口交集；H1仅扣窗口内外露圆孔，上层共同覆盖为半个两圆透镜。V2保留右窗，加X负300至负170、同Y范围左窗，H2先与窗口裁切再扣上层圆共同覆盖。支模和总量修约3位仅合成，TEST-ROUND-WINDOW非官方编码。','foundation-round-step-window-basis.txt')
    for c in get(e,pid)['components']:approve(r,pid,'components',c)
    p=get(e,pid);s=next(c for c in p['components'] if c['id']==s['id']);eid=p['evidence'][-1]['id'];s['object_hash']=store.digest(s);item=s['foundation_items'][0];item['through_holes']['deduction_rule'].update(threshold='100',evidence_id=eid);item['through_holes']['holes']=[{**hole(eid=eid,x='190'),'vertical_range':interval('0','200',eid)},{**rectangle(eid=eid,x='-200',name='合成H2'),'vertical_range':interval('0','200',eid)}];s['basis_evidence_id']=eid;res=e.put(f'/api/v1/projects/{pid}/components/'+s['id'],json=s);assert res.status_code==200,res.text;approve(r,pid,'components',res.json());s=next(c for c in get(e,pid)['components'] if c['id']==s['id']);return e,r,pid,s,body(s,eid,bounds=('170','300','0','200')),eid


def test_api_preview_readonly_claims_and_frame_rules(clients):
    e,r,pid,s,b,eid=setup(clients);before=get(e,pid);post(e,pid,b,'/components/preview');assert get(e,pid)==before;c=post(e,pid,b).json();approve(r,pid,'components',c);post(e,pid,{**b,'id':c['id']},expected=409);other=body(s,eid,bounds=('-300','-170','0','200'));post(e,pid,other);post(e,pid,step_body(s,eid),expected=409);bad=copy.deepcopy(other);bad['contact_ranges'][0]['plan_range']['coordinate_basis']='不同坐标';post(e,pid,bad,expected=409)


@pytest.mark.parametrize('fault',['upper','lower','covered','net','bounds','trace-area','trace-start','trace-boundary','trace-ast','rule','book'])
def test_stored_trace_and_geometry_tamper_rebuilt(clients,fault):
    e,r,pid,s,b,eid=setup(clients);c=post(e,pid,b).json()
    if fault=='book':approve(r,pid,'components',c);post(e,pid,{},'/takeoffs/calculate')
    def forge(p):
        d=p['components'][-1]['preview']['contact_ranges'][0];h=d['apertures'][0];m=h['curved_intersection']
        if fault in ('upper','lower'):d['window_'+fault+'_area']='0'
        elif fault in ('covered','net'):h['covered_area' if fault=='covered' else 'single_area']='0'
        elif fault=='bounds':d['plan_bounds']['x_start']='0'
        elif fault=='rule':m['integration_rule']='forged'
        elif fault=='book':p['takeoffs'][-1]['quantity']='99'
        else:m['terms'][0][{'trace-area':'area','trace-start':'x_from','trace-boundary':'upper_boundary','trace-ast':'formula_ast'}[fault]]={'constant':'0'} if fault=='trace-ast' else 'forged'
    mutate(pid,forge);coll='takeoffs' if fault=='book' else 'components';approve(r,pid,coll,get(e,pid)[coll][-1],expected=409)


@pytest.mark.parametrize('gate',['preview','save','review','calculate','book_review'])
def test_evidence_bytes_at_five_gates(clients,gate):
    e,r,pid,s,b,eid=setup(clients)
    if gate in ('review','calculate','book_review'):c=post(e,pid,b).json()
    if gate in ('calculate','book_review'):approve(r,pid,'components',c)
    if gate=='book_review':t=post(e,pid,{},'/takeoffs/calculate').json()[-1]
    doc=get(e,pid)['documents'][-1];(store.ROOT/'originals'/doc['hash']).write_bytes(b'corrupt')
    if gate=='preview':post(e,pid,b,'/components/preview',409)
    elif gate=='save':post(e,pid,b,expected=409)
    elif gate=='review':approve(r,pid,'components',c,expected=409)
    elif gate=='calculate':post(e,pid,{},'/takeoffs/calculate',409)
    else:approve(r,pid,'takeoffs',t,expected=409)


def test_history_roles_foreign_evidence_concurrency_and_source_change(clients):
    e,r,pid,s,b,eid=setup(clients);post(r,pid,b,expected=403);assert clients['other'].get(f'/api/v1/projects/{pid}').status_code==404;foreign=create(e)['id'];upload(e,foreign,'异项目圆平台窗口依据');bad=copy.deepcopy(b);bad['contact_ranges'][0]['plan_range']['evidence']['x_start']=get(e,foreign)['evidence'][-1]['id'];post(e,pid,bad,expected=404);c=post(e,pid,b).json();approve(r,pid,'components',c)
    for t in post(e,pid,{},'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
    post(e,pid,{},'/boq/generate');before=get(e,pid);old=before['components'][-1];b['object_hash']=store.digest(old);other=body(s,eid,bounds=('-300','-170','0','200'))['contact_ranges'][0];other['name']='左窗';b['contact_ranges'].append(other);url=f'/api/v1/projects/{pid}/components/'+old['id'];res=e.put(url,json=b);assert res.status_code==200,res.text;after=get(e,pid);assert after['component_history'][-1]==old and after['boq']==before['boq'] and e.put(url,json=b).status_code==409;approve(r,pid,'components',res.json());s['object_hash']=store.digest(s);s['foundation_items'][0]['levels'][1]['parameters']['diameter']='420';res=e.put(f'/api/v1/projects/{pid}/components/'+s['id'],json=s);assert res.status_code==200,res.text;approve(r,pid,'components',res.json());assert get(e,pid)['components'][-1]['state']=='stale'


@pytest.mark.parametrize('fmt',['json','xlsx','pdf'])
def test_three_formats_interval_table_and_old_quote_isolation(clients,fmt):
    e,r,pid,_,q=priced_project(clients);e,r,pid,s,b,eid=setup(clients,pid);c=post(e,pid,b).json();approve(r,pid,'components',c)
    for t in post(e,pid,{},'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
    post(e,pid,{},'/boq/generate');url=f'/api/v1/projects/{pid}/exports';res=e.get(url,params=dict(format=fmt,unpriced=True));assert res.status_code==200,res.text;titles=set(STEP_WINDOW_TITLES)|set(WINDOW_TITLES)|{WINDOW_CURVED_TITLE,ROUND_WINDOW_TITLE}
    if fmt=='json':
        exported={t[0]:t for t in res.json()['tables']};assert titles<=set(exported)
        assert exported[WINDOW_CURVED_TITLE][1][1]=='解析分项'
        from backend.foundation_horizontal_formwork_reports import CURVED_TITLE
        cells=[row[-1] for row in exported[CURVED_TITLE][2] if '两圆与矩形窗口分段圆弧解析积分' in row[-1]]
        assert cells and all('象限分项' not in cell for cell in cells)
        assert any('按圆支撑端点' in cell for cell in cells)
    elif fmt=='xlsx':assert titles<=set(load_workbook(io.BytesIO(res.content)).sheetnames)
    else:assert all(t in ''.join(p.extract_text() for p in PdfReader(io.BytesIO(res.content)).pages) for t in titles)
    assert not any(t[0] in titles for t in e.get(url,params=dict(format='json',quote_id=q['id'])).json()['tables']);assert tables(dict(takeoffs=[]))==[]


@pytest.mark.parametrize('kind',['rect','circle'])
def test_exact_complete_void_certificate_resolves_zero_without_tolerance(kind):
    if kind=='rect':
        hs={**rectangle(x='-190',length='80',width='50'),'vertical_range':interval('0','200')};s=source('stepped_circular',holes=[hs]);bs=('-210','-135','20','22')
    else:s=circle_source(x='114',y='-152',diameter='70');bs=('100','122','-165','-85')
    r=calculate(body(s,bounds=bs),s);d=r['contact_ranges'][0];h=d['apertures'][0];assert h['exposed_coverage']['terms'] and Decimal(d['single_gross_area'])>0 and Decimal(d['single_net_area'])==0 and r['quantity']=='0.000';assert h['formula_ast']==d['formula_asts']['single_gross_area'];check(r)


@pytest.mark.parametrize('kind',['rect','circle'])
def test_complete_void_certificate_does_not_hide_uncovered_positive_area(kind):
    if kind=='rect':
        hs={**rectangle(x='-190',length='80',width='50'),'vertical_range':interval('0','200')};s=source('stepped_circular',holes=[hs]);bs=('-210','220','20','22')
    else:s=circle_source(x='114',y='-152',diameter='70');bs=('100','122','-230','-85')
    r=calculate(body(s,bounds=bs),s);d=r['contact_ranges'][0];assert 'exposed_coverage' not in d['apertures'][0] and Decimal(d['single_net_area'])>0;check(r)


@pytest.mark.parametrize('fault',['margin','boundary'])
def test_complete_void_certificate_tamper_is_rebuilt(clients,fault):
    e,r,pid,s,b,eid=setup(clients);b['contact_ranges'][0]['plan_range']=window(eid,'-210','-135','20','22');c=post(e,pid,b).json();assert c['preview']['quantity']=='0.000'
    def forge(p):
        q=p['components'][-1]['preview']['contact_ranges'][0]['apertures'][1]['exposed_coverage']['terms'][0];q['lower_margin' if fault=='margin' else 'upper_boundary']='forged'
    mutate(pid,forge);approve(r,pid,'components',get(e,pid)['components'][-1],expected=409)


@pytest.mark.parametrize('fmt',['json','xlsx','pdf'])
def test_zero_contact_certificate_in_three_formats(clients,fmt):
    from backend.foundation_horizontal_formwork_reports import ROUND_COVERAGE_TITLE
    e,r,pid,s,b,eid=setup(clients);b['contact_ranges'][0]['plan_range']=window(eid,'-210','-135','20','22');c=post(e,pid,b).json();approve(r,pid,'components',c)
    for t in post(e,pid,{},'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
    post(e,pid,{},'/boq/generate');res=e.get(f'/api/v1/projects/{pid}/exports',params=dict(format=fmt,unpriced=True));assert res.status_code==200,res.text
    if fmt=='json':assert ROUND_COVERAGE_TITLE in {t[0] for t in res.json()['tables']}
    elif fmt=='xlsx':assert ROUND_COVERAGE_TITLE in load_workbook(io.BytesIO(res.content)).sheetnames
    else:assert ROUND_COVERAGE_TITLE in ''.join(p.extract_text() for p in PdfReader(io.BytesIO(res.content)).pages)
