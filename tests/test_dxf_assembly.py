import copy
import io
import pytest
from openpyxl import load_workbook
from pypdf import PdfReader
from backend import store
from tests.test_dxf_routes import setup,read,body,preview,adopt,polyline,mutate
from tests.test_workflow import get,approve,priced_project


def segment(handle,a,b,layer='PIPE'):
    return [(0,'LINE'),(5,handle),(8,layer)]+[(code,str(v)) for code,v in zip((10,20,30,11,21,31),(*a,*b))]


def sample():
    return segment('A',(0,0,0),(3000,4000,12000))+segment('B',(3000,4000,17000),(3000,4000,12000),'VERTICAL')


def selections(rows):
    return [{'candidate_id':r['id'],'object_hash':r['object_hash'],'reverse':i==1} for i,r in enumerate(rows)]


def resolve(c,url,parts,expected=200):
    r=c.post(url.replace('dxf-route-candidates','dxf-route-assembly'),json={'segments':parts})
    assert r.status_code==expected,r.text
    return r.json().get('candidate',r.json())


def assembled(clients,pid=None):
    c,r,pid,did,eid,url=setup(clients,sample(),pid=pid);objects=read(c,url)['candidates'];parts=selections(objects)
    row=resolve(c,url,parts);request=body(row,eid,segments=parts)
    return c,r,pid,did,eid,url,objects,row,request


def test_resolve_and_preview_are_readonly_and_keep_originals_and_direction(clients):
    c,r,pid,did,eid,url=setup(clients,sample());objects=read(c,url)['candidates'];parts=selections(objects);before=get(c,pid)
    row=resolve(c,url,parts);assert row==resolve(r,url,parts) and get(c,pid)==before
    assert row['kind']=='ASSEMBLY' and row['single_length_m']=='18' and len(row['points'])==3
    assert row['points'][2]=={'x':'3000','y':'4000','z':'17000'}
    assert row['source_objects'][1]['points'][0]==row['points'][2] and row['source_ranges'][1]['reverse']
    assert row['source_ranges'][0]['start_point']==1 and row['source_ranges'][0]['end_point']==2
    assert row['source_ranges'][1]['start_point']==2 and row['source_ranges'][1]['end_point']==3
    assert row['joins'][0]['incoming_end']==row['joins'][0]['outgoing_start']
    view=preview(c,url,row,body(row,eid,segments=parts));assert view['result']['quantity']=='36.000'
    assert get(c,pid)==before


def test_assembly_adoption_independent_review_takeoff_and_boq(clients):
    c,r,pid,did,eid,url,objects,row,request=assembled(clients)
    view=preview(c,url,row,request);result=adopt(c,url,row,request,view);comp=result['component']
    assert comp['state']=='draft' and not get(c,pid)['takeoffs'] and len(get(c,pid)['evidence'])==3
    assert adopt(c,url,row,request,view)['duplicate']
    approve(c,pid,'components',comp,expected=403);approve(r,pid,'components',comp)
    res=c.post(f'/api/v1/projects/{pid}/takeoffs/calculate');assert res.status_code==200,res.text
    assert res.json()[0]['quantity']=='36.000';approve(r,pid,'takeoffs',res.json()[0])
    assert c.post(f'/api/v1/projects/{pid}/boq/generate').json()[0]['quantity']=='36.000'


@pytest.mark.parametrize('fault',['empty','one','many','repeat','string_reverse','missing_reverse','extra','hash','id','missing','wrong_direction','order'])
def test_invalid_selection_is_atomic(clients,fault):
    c,r,pid,did,eid,url=setup(clients,sample());parts=selections(read(c,url)['candidates']);expected=422
    if fault=='empty':parts=[]
    elif fault=='one':parts=parts[:1]
    elif fault=='many':parts=parts*11
    elif fault=='repeat':parts[1]=copy.deepcopy(parts[0])
    elif fault=='string_reverse':parts[1]['reverse']='true'
    elif fault=='missing_reverse':parts[1].pop('reverse')
    elif fault=='extra':parts[1]['snap']=True
    elif fault=='hash':parts[1]['object_hash']='bad'
    elif fault=='id':parts[1]['candidate_id']='bad'
    elif fault=='missing':parts[1]['candidate_id']='0'*32;expected=409
    elif fault=='wrong_direction':parts[1]['reverse']=False
    else:parts.reverse()
    before=get(c,pid);resolve(c,url,parts,expected);assert get(c,pid)==before


@pytest.mark.parametrize('gap',['1','0.0000000001'])
def test_xy_touching_with_z_gap_does_not_snap(clients,gap):
    from decimal import Decimal
    z=Decimal('12000')+Decimal(gap)
    entities=segment('A',(0,0,0),(3000,4000,12000))+segment('B',(3000,4000,z),(3000,4000,17000))
    c,r,pid,did,eid,url=setup(clients,entities);parts=selections(read(c,url)['candidates']);parts[1]['reverse']=False
    error=resolve(c,url,parts,422);assert 'XYZ端点不完全一致' in error['detail']


def test_numerically_equal_join_keeps_both_raw_representations(clients):
    entities=segment('A',(0,0,0),(3000,4000,12000))+segment('B',('3000.0','4E3','12000.00'),(3000,4000,17000))
    c,r,pid,did,eid,url=setup(clients,entities);parts=selections(read(c,url)['candidates']);parts[1]['reverse']=False
    row=resolve(c,url,parts)
    assert row['points'][1]['x']=='3000' and row['joins'][0]['outgoing_start']['x']=='3000.0'
    assert row['joins'][0]['outgoing_start']['y']=='4E3'


@pytest.mark.parametrize('entities',[
    segment('A',(0,0,0),(3000,4000,12000))+segment('B',(3000,4000,12000),(0,0,0)),
    segment('A',(0,0,0),(3000,0,0))+segment('B',(3000,0,0),(1000,0,0)),
    segment('A',(0,0,0),(3000,4000,12000))+segment('B',(3000,4000,12000),(3000,4000,17000))+[(67,'1')]])
def test_closed_overlapping_or_unsupported_objects_are_not_assembled(clients,entities):
    c,r,pid,did,eid,url=setup(clients,entities);parts=selections(read(c,url)['candidates']);parts[1]['reverse']=False
    resolve(c,url,parts,422)


def test_mixed_polyline_and_line_join_preserves_all_vertices(clients):
    c,r,pid,did,eid,url=setup(clients,polyline()+segment('B',(3000,4000,18000),(3000,4000,17000)))
    row=resolve(c,url,selections(read(c,url)['candidates']));assert len(row['points'])==4 and row['single_length_m']=='19'


@pytest.mark.parametrize('count',[51,52])
def test_combined_vertex_bound_does_not_truncate(clients,count):
    entities=polyline([(i,0,0) for i in range(51)])+polyline([(i,0,0) for i in range(50,50+count)],handle='B')
    c,r,pid,did,eid,url=setup(clients,entities);parts=selections(read(c,url)['candidates']);parts[1]['reverse']=False
    row=resolve(c,url,parts,200 if count==51 else 422)
    if count==51:assert len(row['points'])==101 and row['points'][-1]['x']=='100'


@pytest.mark.parametrize('first',['single','assembly'])
def test_single_and_assembled_adoptions_cannot_reuse_same_source_object(clients,first):
    c,r,pid,did,eid,url,objects,row,request=assembled(clients)
    if first=='single':
        req=body(objects[0],eid);adopt(c,url,objects[0],req,preview(c,url,objects[0],req))
        request['request_id']='synthetic-assembly-after-single';adopt(c,url,row,request,preview(c,url,row,request),409)
    else:
        adopt(c,url,row,request,preview(c,url,row,request))
        req=body(objects[0],eid,request_id='synthetic-single-after-assembly');adopt(c,url,objects[0],req,preview(c,url,objects[0],req),409)


def test_existing_single_can_be_explicitly_revised_into_assembly(clients):
    c,r,pid,did,eid,url,objects,row,request=assembled(clients)
    req=body(objects[0],eid);old=adopt(c,url,objects[0],req,preview(c,url,objects[0],req))['component'];approve(r,pid,'components',old)
    t=c.post(f'/api/v1/projects/{pid}/takeoffs/calculate').json()[0];approve(r,pid,'takeoffs',t)
    request.update(request_id='synthetic-single-to-assembled',component_id=old['id'],component_hash=store.digest(get(c,pid)['components'][0]))
    revised=adopt(c,url,row,request,preview(c,url,row,request))['component']
    p=get(c,pid);assert revised['id']==old['id'] and revised['revision']==2 and revised['state']=='draft'
    assert p['takeoffs'][0]['state']=='stale' and len(p['components'])==1 and len(p['component_history'])==1
    assert p['dxf_route_adoptions'][0]['trace']['candidate']['kind']=='LINE'


@pytest.mark.parametrize('field',['selection','source_objects','source_ranges','joins'])
def test_stored_assembly_metadata_is_reconstructed_before_review(clients,field):
    c,r,pid,did,eid,url,objects,row,request=assembled(clients)
    comp=adopt(c,url,row,request,preview(c,url,row,request))['component']
    def change(p):
        for trace in [p['components'][0]['dxf_trace'],p['dxf_route_adoptions'][0]['trace']]:trace['candidate'][field]=None
    mutate(pid,change);approve(r,pid,'components',get(c,pid)['components'][0],expected=409)


def test_changed_selected_object_hash_and_stale_preview_are_blocked(clients):
    c,r,pid,did,eid,url,objects,row,request=assembled(clients)
    view=preview(c,url,row,request);bad=copy.deepcopy(request);bad['segments'][1]['object_hash']='0'*64
    adopt(c,url,row,bad,view,409)
    bad=copy.deepcopy(request);bad['segments'][1]['reverse']=False
    adopt(c,url,row,bad,view,422)
    mutate(pid,lambda p:next(j for j in p['jobs'] if j['document_id']==did).update(state='failed'))
    adopt(c,url,row,request,view,409)


def test_tenant_role_and_original_single_request_hash_compatibility(clients):
    c,r,pid,did,eid,url,objects,row,request=assembled(clients)
    resolve(clients['other'],url,request['segments'],404)
    view=preview(c,url,row,request);adopt(r,url,row,request,view,403)
    req=body(objects[0],eid);v=preview(c,url,objects[0],req);result=adopt(c,url,objects[0],req,v)
    legacy={**req,'component_id':None,'component_hash':None,'preview_hash':v['preview_hash']}
    assert result['adoption']['request_hash']==store.digest({'document_id':did,'candidate_id':objects[0]['id'],**legacy})


def test_three_export_formats_keep_assembly_and_old_quote_separate(clients):
    c,r,pid,_,old=priced_project(clients)
    c,r,pid,did,eid,url,objects,row,request=assembled(clients,pid=pid)
    adopt(c,url,row,request,preview(c,url,row,request));base=f'/api/v1/projects/{pid}/exports'
    for fmt in ['json','xlsx','pdf']:
        res=c.get(base,params={'format':fmt,'unpriced':True});assert res.status_code==200,res.text
        if fmt=='json':assert 'DXF多对象组装来源' in str(res.json()['tables'])
        elif fmt=='xlsx':assert 'DXF多对象组装来源' in load_workbook(io.BytesIO(res.content)).sheetnames
        else:assert 'DXF多对象组装来源' in ''.join(p.extract_text() for p in PdfReader(io.BytesIO(res.content)).pages)
    historic=c.get(base,params={'format':'json','quote_id':old['id']}).json()
    assert not any(t[0]=='DXF多对象组装来源' for t in historic['tables']) and historic['quote']==old
