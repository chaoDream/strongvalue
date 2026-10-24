"""Synthetic hole-wall strips: offset coordinates, physical cuts and review gates."""
import copy,io,random
from decimal import Decimal,localcontext,ROUND_HALF_UP
import pytest
from openpyxl import load_workbook
from pypdf import PdfReader
from backend import store
from backend.engine import DomainError
from backend.foundation_void_formwork import calculate
from backend.foundation_void_formwork_reports import TITLES,STRIP_TITLES
from tests.test_foundation_void_formwork import foundation,body,wall,setup,check_asts
from tests.test_foundation_voids import rectangle
from tests.test_foundation_recesses import interval
from tests.test_post_foundation import equal_body,fixed_body,fixed_added,equal_added,fixed_adjusted,equal_adjusted,combined
from tests.test_route_takeoff import evaluate
from tests.test_polygon_takeoff import post,mutate
from tests.test_workflow import get,approve,upload,create,priced_project

FRAME='合成基础中心为0，X向图右、Y向图上，孔位沿已审基础共同局部XY；面向不同孔侧不翻转，非真实工程。'


def strip(start='-30',end='0',unit='mm',eid='test'):
 return dict(parameters=dict(lateral_start=start,lateral_end=end),unit=unit,evidence=dict(lateral_start=eid,lateral_end=eid),coordinate_basis=FRAME,coordinate_evidence_id=eid,scope_confirmed=True)


def selected(eid='test',start='-30',end='0',**kwargs):return {**wall(eid,**kwargs),'lateral_range':strip(start,end,eid=eid)}


def strip_body(s,eid='test'):
 b=body(s,eid);b['contact_ranges']=[selected(eid,end='-10',name='合成下带'),selected(eid,start='-10',end='10',name='合成中带'),selected(eid,start='10',end='30',name='合成上带')];return b


def check(r):
 check_asts(r)
 with localcontext() as ctx:
  ctx.prec=80
  for c in r.get('void_lateral_checks',[]):
   for k,a in c['formula_asts'].items():assert evaluate(a,r['dimensions'])==Decimal(c[k])


def test_partition_equals_full_wall_actual_cuts_and_no_mutation():
 s=foundation();b=strip_body(s);before=copy.deepcopy((s,b));r=calculate(b,s);assert (s,b)==before;whole=body(s);whole['contact_ranges']=whole['contact_ranges'][:1];w=calculate(whole,s);assert r['quantity']==w['quantity']=='0.154' and r['gross']==w['gross'] and r['deduction']==w['deduction'];assert [Decimal(d['single_deduction_area']) for d in r['contact_ranges']]==list(map(Decimal,['.0005','.001','.0005']));assert r['rule_version']=='foundation-void-formwork-1.1' and w['rule_version']=='foundation-void-formwork-1.0';assert [d['state'] for d in r['void_lateral_checks']]==['lateral_tangent','laterally_disjoint','lateral_tangent'];check(r)


@pytest.mark.parametrize('side',['x_minus','x_plus','y_minus','y_plus'])
@pytest.mark.parametrize('make',[equal_body,fixed_body,fixed_added,equal_added,fixed_adjusted,equal_adjusted,combined,lambda s:combined(s,equal=True)])
def test_offset_holes_four_axes_and_eight_final_counts(side,make):
 s=foundation(make);h={**rectangle(x='40',y='60'),'vertical_range':interval('0','150')};axis='x' if side.startswith('x_') else 'y';other='y' if axis=='x' else 'x';sign=-1 if side.endswith('minus') else 1;center=40 if axis=='x' else 60;half=50 if axis=='x' else 30;n={**rectangle(name='合成相邻孔',x='40',y='60',length='40',width='40'),'vertical_range':interval('50','100')};n['parameters'][axis]=str(center+sign*(half+20));n['parameters'][other]=str((60 if axis=='x' else 40)+10);s['foundation_items'][0]['through_holes']['holes']=[h,n];b=body(s);b['contact_ranges']=[selected(start='40' if axis=='x' else '20',end='80' if axis=='x' else '60',face=side)];r=calculate(b,s);d=r['contact_ranges'][0];assert Decimal(d['width'])==Decimal('.04') and Decimal(d['single_deduction_area'])==Decimal('.0015');assert Decimal(r['void_formwork_summary']['unrounded_quantity'])==Decimal('.0045')*Decimal(d['total_count']);assert d['lateral_axis']==('Y' if axis=='x' else 'X');assert Decimal(d['face_lateral_start'])==Decimal('.03' if axis=='x' else '-.01');check(r)


@pytest.mark.parametrize('shape',['rectangular','stepped_rectangular','circular','stepped_circular'])
def test_rectangular_hole_strip_in_four_foundation_shapes(shape):
 s=foundation(shape=shape);r=calculate(strip_body(s),s);assert r['quantity']=='0.154';check(r)


@pytest.mark.parametrize('unit,factor',[('mm','1'),('cm','.1'),('dm','.01'),('m','.001'),('km','.000001')])
def test_independent_lateral_units_and_low_calling_precision(unit,factor):
 s=foundation();b=strip_body(s)
 for d in b['contact_ranges']:
  d['lateral_range']['unit']=unit
  for k,v in d['lateral_range']['parameters'].items():d['lateral_range']['parameters'][k]=str(Decimal(v)*Decimal(factor))
 with localcontext() as ctx:ctx.prec=6;r=calculate(b,s)
 assert r==calculate(b,s) and r['quantity']=='0.154';assert len(r['unit_conversions'])==12;check(r)


@pytest.mark.parametrize('side,x,y,start,end',[('x_plus','150','0','-15','15'),('x_minus','-150','0','-15','15'),('y_plus','0','220','-20','20'),('y_minus','0','-220','-20','20')])
def test_partial_exterior_contact_is_entirely_cut(side,x,y,start,end):
 s=foundation();s['foundation_items'][0]['through_holes']['holes']=[{**rectangle(x=x,y=y),'vertical_range':interval('0','150')}];b=body(s);b['contact_ranges']=[selected(start=start,end=end,face=side)];r=calculate(b,s);d=r['contact_ranges'][0];assert r['quantity']=='0.000' and d['single_gross_area']==d['single_deduction_area'];assert Decimal(d['deductions'][0]['lateral_start'])==Decimal(start)/1000;check(r)


def test_160_independent_strip_neighbor_intersection_oracles():
 rng=random.Random(9701)
 for _ in range(160):
  s=foundation();nl,nh=sorted(rng.sample(range(-30,31),2));zl,zh=sorted(rng.sample(range(151),2));a,b=sorted(rng.sample(range(-30,31),2));lo,hi=sorted(rng.sample(range(151),2));c=s['foundation_items'][0]['through_holes'];n={**rectangle(name='合成相邻',x='70',y=str(Decimal(nl+nh)/2),length='60',width=str(nh-nl)),'vertical_range':interval(str(zl),str(zh))};c['holes']=[c['holes'][0],n];v=body(s);v['contact_ranges']=[selected(start=str(a),end=str(b))];v['contact_ranges'][0]['parameters']=dict(z_start=str(lo),z_end=str(hi));v['rounding_digits']=rng.randrange(13);r=calculate(v,s)
  with localcontext() as ctx:ctx.prec=80;cut=max(0,min(b,nh)-max(a,nl))*max(0,min(hi,zh)-max(lo,zl));expected=Decimal((b-a)*(hi-lo)-cut)/1000000*22;assert Decimal(r['quantity'])==expected.quantize(Decimal(1).scaleb(-v['rounding_digits']),rounding=ROUND_HALF_UP)
  check(r)


@pytest.mark.parametrize('fault',['null','list','missing','extra','parameters','evidence','start-proof','end-proof','basis','basis-proof','confirm','truthy','unit','reverse','equal','left','right','nan','bool','precision','circular','arc','mixed-frame','mixed-proof','same-strip','full-overlap','tiny-overlap','hole-origin'])
def test_strict_lateral_input_and_no_snap(fault):
 s=foundation();b=strip_body(s);d=b['contact_ranges'][0];c=d['lateral_range']
 if fault=='null':d['lateral_range']=None
 elif fault=='list':d['lateral_range']=[]
 elif fault=='missing':c.pop('coordinate_basis')
 elif fault=='extra':c['width']='10'
 elif fault=='parameters':c['parameters']['width']='10'
 elif fault=='evidence':c['evidence'].pop('lateral_end')
 elif fault in ('start-proof','end-proof'):c['evidence']['lateral_start' if fault=='start-proof' else 'lateral_end']=' '
 elif fault=='basis':c['coordinate_basis']=' '
 elif fault=='basis-proof':c['coordinate_evidence_id']=' '
 elif fault in ('confirm','truthy'):c['scope_confirmed']=False if fault=='confirm' else 1
 elif fault=='unit':c['unit']='m²'
 elif fault in ('reverse','equal','left','right','nan','bool','precision'):c['parameters']['lateral_start' if fault=='left' else 'lateral_end']={'reverse':'-31','equal':'-30','left':'-30.000000000001','right':'30.000000000001','nan':'NaN','bool':True,'precision':'0.1234567890123456789'}[fault]
 elif fault=='circular':d.update(hole_index=2,face='circumference',parameters=dict(z_start='300',z_end='600'))
 elif fault=='arc':d['arc_range']={}
 elif fault in ('mixed-frame','mixed-proof'):b['contact_ranges'][1]['lateral_range']['coordinate_basis' if fault=='mixed-frame' else 'coordinate_evidence_id']='other'
 elif fault=='same-strip':b['contact_ranges'][1]['lateral_range']=copy.deepcopy(c)
 elif fault=='full-overlap':b['contact_ranges'][1].pop('lateral_range')
 elif fault=='tiny-overlap':b['contact_ranges'][1]['lateral_range']['parameters']['lateral_start']='-10.000000000001'
 elif fault=='hole-origin':d['face']='y_plus';c['parameters']=dict(lateral_start='-50',lateral_end='50')
 with pytest.raises(DomainError):calculate(b,s)


def test_shared_boundary_and_z_tangent_full_strip_coexist():
 s=foundation();b=strip_body(s);b['contact_ranges']=b['contact_ranges'][:1]+[wall(start='150',end='150',name='zero')]
 b['contact_ranges'][0]['parameters']['z_end']='50';b['contact_ranges'][1]['parameters']=dict(z_start='50',z_end='150');r=calculate(b,s);assert r['void_lateral_checks'][0]['state']=='vertical_tangent';check(r)


def test_partial_cut_fully_empty_interior_and_neighbor_z_tangent():
 s=foundation();b=strip_body(s);b['contact_ranges']=[selected(start='-20',end='20')];b['contact_ranges'][0]['parameters']=dict(z_start='50',z_end='100');r=calculate(b,s);assert r['quantity']=='0.000';b['contact_ranges'][0]['parameters']=dict(z_start='100',z_end='150');r=calculate(b,s);assert Decimal(r['deduction'])==0 and r['quantity']=='0.044';check(r)


def test_candidate_offset_bounds_preview_and_no_project_mutation(clients):
 e,r,pid,s,_,eid=setup(clients);before=get(e,pid);c=e.get(f'/api/v1/projects/{pid}/foundation-void-formwork-candidates',params=dict(component_id=s['id'])).json();h=c['foundation_items'][0]['holes'][0];assert h['lateral_bounds']['y_plus']==dict(start='0.10',end='0.20',axis='X');assert 'lateral_bounds' not in c['foundation_items'][0]['holes'][1];post(e,pid,strip_body(s,eid),'/components/preview');assert get(e,pid)==before


def test_cross_record_units_actual_id_full_strip_and_stale_claims(clients):
 e,r,pid,s,_,eid=setup(clients);b=strip_body(s,eid);b['contact_ranges']=b['contact_ranges'][:1];c=post(e,pid,b).json();approve(r,pid,'components',c);post(e,pid,{**b,'id':c['id']},expected=409);second=strip_body(s,eid);second['contact_ranges']=second['contact_ranges'][1:];
 for d in second['contact_ranges']:
  d['lateral_range']['unit']='m'
  d['lateral_range']['parameters']={k:str(Decimal(v)/1000) for k,v in d['lateral_range']['parameters'].items()}
 post(e,pid,second);second['contact_ranges'][0]['lateral_range']['parameters']['lateral_start']='-.010000000000001';post(e,pid,second,expected=409);whole=body(s,eid);whole['contact_ranges']=whole['contact_ranges'][:1];post(e,pid,whole,expected=409);mutate(pid,lambda p:p['components'][-2].update(state='stale'));post(e,pid,b,expected=409)


@pytest.mark.parametrize('fault',['basis','proof'])
def test_frame_shared_across_records_and_holes_even_disjoint(fault,clients):
 e,r,pid,s,_,eid=setup(clients);b=strip_body(s,eid);b['contact_ranges']=b['contact_ranges'][:1];post(e,pid,b);b['contact_ranges']=[selected(eid,start='-20',end='20',hole_index=3,name='另一孔')];b['contact_ranges'][0]['parameters']=dict(z_start='50',z_end='100');b['contact_ranges'][0]['lateral_range']['coordinate_basis' if fault=='basis' else 'coordinate_evidence_id']='other';post(e,pid,b,expected=409)


@pytest.mark.parametrize('fault',['raw','unit','frame','proof','width','face-bounds','cut','check','formula'])
def test_saved_strip_tamper_blocks_independent_review(clients,fault):
 e,r,pid,s,_,eid=setup(clients);c=post(e,pid,strip_body(s,eid)).json()
 def forge(p):
  c=p['components'][-1];d=c['contact_ranges'][0]['lateral_range'];v=c['preview']['contact_ranges'][0]
  if fault=='raw':d['parameters']['lateral_start']='-29'
  elif fault=='unit':d['unit']='cm'
  elif fault=='frame':d['coordinate_basis']='forged'
  elif fault=='proof':d['evidence']['lateral_start']='forged'
  elif fault=='width':v['width']='999'
  elif fault=='face-bounds':v['face_lateral_start']='-999'
  elif fault=='cut':v['deductions'][0]['lateral_end']='999'
  elif fault=='check':c['preview']['void_lateral_checks'][0]['lateral_gap']='999'
  else:v['formula_asts']['lateral_start']={'constant':'0'}
 mutate(pid,forge);approve(r,pid,'components',get(e,pid)['components'][-1],expected=404 if fault=='proof' else 422 if fault in ('unit','frame') else 409)


@pytest.mark.parametrize('gate',['preview','save','review','calculate','book_review'])
def test_strip_original_bytes_five_gates(clients,gate):
 e,r,pid,s,_,eid=setup(clients);b=strip_body(s,eid)
 if gate in ('review','calculate','book_review'):c=post(e,pid,b).json()
 if gate in ('calculate','book_review'):approve(r,pid,'components',c)
 if gate=='book_review':t=post(e,pid,{},'/takeoffs/calculate').json()[-1]
 doc=get(e,pid)['documents'][-1];(store.ROOT/'originals'/doc['hash']).write_bytes(b'corrupt')
 if gate=='preview':post(e,pid,b,'/components/preview',409)
 elif gate=='save':post(e,pid,b,expected=409)
 elif gate=='review':approve(r,pid,'components',c,expected=409)
 elif gate=='calculate':post(e,pid,{},'/takeoffs/calculate',409)
 else:approve(r,pid,'takeoffs',t,expected=409)


def test_full_to_strip_revision_history_concurrency_and_source_stale(clients):
 e,r,pid,s,b,eid=setup(clients);b['contact_ranges']=b['contact_ranges'][:1];old=post(e,pid,b).json();approve(r,pid,'components',old)
 for t in post(e,pid,{},'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
 post(e,pid,{},'/boq/generate');p=get(e,pid);old=p['components'][-1];b=strip_body(s,eid);b['contact_ranges']=b['contact_ranges'][:2];b['object_hash']=store.digest(old);url=f'/api/v1/projects/{pid}/components/'+old['id'];res=e.put(url,json=b);assert res.status_code==200 and res.json()['preview']['quantity']=='0.099';p2=get(e,pid);assert p2['component_history'][-1]==old and p2['boq']==p['boq'] and all(t['state']=='stale' for t in p2['takeoffs']);assert e.put(url,json=b).status_code==409;approve(r,pid,'components',res.json());s['object_hash']=store.digest(s);s['foundation_items'][0]['parameters']['height']='700';res=e.put(f'/api/v1/projects/{pid}/components/'+s['id'],json=s);assert res.status_code==200;approve(r,pid,'components',res.json());assert get(e,pid)['components'][-1]['state']=='stale'


def test_strip_roles_tenants_foreign_proof(clients):
 e,r,pid,s,_,eid=setup(clients);b=strip_body(s,eid);post(r,pid,b,expected=403);assert clients['other'].get(f'/api/v1/projects/{pid}').status_code==404;other=create(e)['id'];upload(e,other,'异项目原件');b['contact_ranges'][0]['lateral_range']['evidence']['lateral_start']=get(e,other)['evidence'][-1]['id'];post(e,pid,b,expected=404)


@pytest.mark.parametrize('fmt',['json','xlsx','pdf'])
def test_three_formats_and_old_quote_isolation(clients,fmt):
 e,r,pid,_,q=priced_project(clients);e,r,pid,s,b,eid=setup(clients,pid);c=post(e,pid,strip_body(s,eid)).json();approve(r,pid,'components',c)
 for t in post(e,pid,{},'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
 boq=post(e,pid,{},'/boq/generate').json();assert any(d['quantity']=='0.154' and d['unit']=='m²' for d in boq);url=f'/api/v1/projects/{pid}/exports';res=e.get(url,params=dict(format=fmt,unpriced=True));assert res.status_code==200,res.text
 if fmt=='json':assert set(TITLES+STRIP_TITLES)<={d[0] for d in res.json()['tables']}
 elif fmt=='xlsx':assert set(TITLES+STRIP_TITLES)<=set(load_workbook(io.BytesIO(res.content)).sheetnames)
 else:assert all(v in ''.join(p.extract_text() for p in PdfReader(io.BytesIO(res.content)).pages) for v in TITLES+STRIP_TITLES)
 assert not any(d[0] in TITLES+STRIP_TITLES for d in e.get(url,params=dict(format='json',quote_id=q['id'])).json()['tables'])


def test_multiple_adjacent_voids_clip_only_selected_transverse_parts():
 s=foundation();c=s['foundation_items'][0]['through_holes'];c['holes']=[c['holes'][0]]+[{**rectangle(name='合成相邻'+str(y),x='70',y=str(y),length='60',width='20'),'vertical_range':interval('0','150')} for y in (-20,20)];b=body(s);b['contact_ranges']=[selected(start='-15',end='15')];r=calculate(b,s);d=r['contact_ranges'][0];assert len(d['deductions'])==2 and r['quantity']=='0.066';assert all(Decimal(n['width'])==Decimal('.005') for n in d['deductions']);check(r)


def test_volume_threshold_never_restores_strip_contact():
 results=[]
 for threshold in ('0','100'):
  s=foundation();s['foundation_items'][0]['through_holes']['deduction_rule']['threshold']=threshold;r=calculate(strip_body(s),s);results.append((r['quantity'],r['gross'],r['deduction']));check(r)
 assert results[0]==results[1]
