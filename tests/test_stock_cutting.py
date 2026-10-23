"""Synthetic homogeneous cutting, exact capacity certificates and net isolation."""
import copy,io,random
from decimal import Decimal,localcontext
from fractions import Fraction
import pytest
from openpyxl import load_workbook
from pypdf import PdfReader
from backend import store
from backend.engine import DomainError
from backend.railing_material import calculate
from backend.railing_takeoff import snapshot
from tests.test_railing_material import geometry,payload,setup as material_setup
from tests.test_railing_posts import body as equal_body
from tests.test_fixed_posts import body as fixed_body
from tests.test_additional_posts import body as fixed_added
from tests.test_equal_additional_posts import body as equal_added
from tests.test_post_adjustments import body as fixed_adjusted
from tests.test_equal_post_adjustments import body as equal_adjusted
from tests.test_combined_post_adjustments import body as combined
from tests.test_path_post_spacing import body as path_body
from tests.test_post_spacing import body as local_body
from tests.test_stock_allowance import config as allowance
from tests.test_polygon_material import panel as polygon_panel
from tests.test_panel_material import panel as rect_panel
from tests.test_route_takeoff import evaluate
from tests.test_railing_takeoff import interval
from tests.test_polygon_takeoff import post,mutate
from tests.test_workflow import get,approve,upload,create,priced_project

PARAMS=('stock_length','kerf_per_piece','trim_start','trim_end','reusable_min_length')


def config(eid='test',stock='6000',unit='mm',kerf='3',start='10',end='10',minimum='1500'):
 return dict(stock_length=stock,kerf_per_piece=kerf,trim_start=start,trim_end=end,reusable_min_length=minimum,length_unit=unit,cut_basis_confirmed=True,measurement_basis='合成每件实际下料完整，每件计一次锯缝含本根最后一件，两端切除未重复；余料按核定下限分类，只列附表不增加净工程量或报价',evidence={k:eid for k in (*PARAMS,'measurement_basis')})


def body(s,eid='test'):
 b=payload(s,eid);b['name']='合成定尺下料与余料（非真实工程）';b['measurement_basis']='单组18m×2道×2kg/m+5件×1.2m×3kg/m=90kg，来源重复2=180kg；定尺下料附表不增加净量或报价';m=b['material_items'][1];m['quantity']='5';m['measurement_basis']='合成每组独立5件，每件实际下料1200mm，每米3000g，来源重复2共10件；下料已完整含本件构造尺寸，定尺另计锯缝及两端切除';m['cutting_plan']=config(eid);return b


def setup(clients,pid=None):
 e,r,pid,s,_,_=material_setup(clients,pid)
 upload(e,pid,'合成定尺软件验收，非真实工程。连续扶手两道不排料，净144kg。独立立柱每组5件实际下料1200mm、3000g/m，来源重复2共10件净36kg，材料表总净180kg。原材定尺6000mm，每件锯缝3mm（含最后一件），每根两端各切除10mm，未与锯缝重复；可留用余料下限1500mm。容量floor(5980/1203)=4，最少3根；前2根各4件、每根余1168mm报废，尾根2件、余3574mm待留用。购入18m=净12+锯缝.03+切除.06+报废2.336+待留用3.574，质量54kg；附表不增加180kg净清单或报价。改原材5000mm仍容量4/3根，购入15m45kg，净量不变须重新复核。阈值1168mm时等值余料可留用。','stock-cutting-basis.txt')
 for c in get(e,pid)['components']:approve(r,pid,'components',c)
 s=get(e,pid)['components'][-1];eid=get(e,pid)['evidence'][-1]['id'];return e,r,pid,s,body(s,eid),eid


def getplan(s,b):return calculate(b,s)['material_items'][-1]['stock_cutting']


def test_exact_capacity_patterns_balances_asts_and_net_quantity():
 s=geometry();b=body(s);before=copy.deepcopy((b,s));old=copy.deepcopy(b);old['material_items'][1].pop('cutting_plan');base=calculate(old,s);v=calculate(b,s);g=v['material_items'][1]['stock_cutting']
 assert (b,s)==before and v['quantity']==base['quantity']=='180.000' and v['formula_ast']==base['formula_ast'] and v['rule_version']=='railing-material-mass-1.13'
 assert (g['required_piece_count'],g['capacity'],g['stock_count'])==(10,'4',3)
 assert [(p['stock_count'],p['pieces_per_stock'],p['offcut_length'],p['offcut_disposition']) for p in g['patterns']]==[(2,4,'1.168','scrap'),(1,2,'3.574','retain')]
 assert [(p['first_stock'],p['last_stock']) for p in g['patterns']]==[(1,2),(3,3)]
 assert g['total']==dict(net_length='12',kerf_length='0.03',trim_length='0.06',scrap_length='2.336',retained_length='3.574',loss_length='2.426',purchased_length='18',net_mass='36',kerf_mass='0.09',trim_mass='0.18',scrap_mass='7.008',retained_mass='10.722',loss_mass='7.278',purchased_mass='54')
 assert g['balance']==dict(purchased_length='18',net_plus_loss_plus_retained='18',difference='0')
 assert v['cutting_summary']['item_count']==1 and v['cutting_summary']['unplanned_item_count']==1 and len([k for k in v['source_evidence'] if '.cutting_plan.' in k])==6
 with localcontext() as c:
  c.prec=80
  assert evaluate(g['required_count_formula_ast'],v['dimensions'])==10
  for k,ast in g['total_formula_asts'].items():assert evaluate(ast,v['dimensions'])==Decimal(g['total'][k])
  for k,ast in v['cutting_summary']['total_formula_asts'].items():assert evaluate(ast,v['dimensions'])==Decimal(v['cutting_summary']['total'][k])
  for p in g['patterns']:
   for k,ast in p['per_stock_formula_asts'].items():assert evaluate(ast,v['dimensions'])==Decimal(p['per_stock'][k])
 cert=g['capacity_certificate'];assert Decimal(cert['used_at_capacity'])<=Decimal(cert['usable_length'])<Decimal(cert['used_at_next'])
 cert=g['stock_count_certificate'];assert int(cert['previous_capacity'])<cert['required']<=int(cert['current_capacity'])


@pytest.mark.parametrize('make',[equal_body,fixed_body,fixed_added,equal_added,fixed_adjusted,equal_adjusted,combined,lambda s:combined(s,equal=True),local_body,path_body])
def test_eight_post_modes_and_both_spacing_scopes_unchanged(make):
 s=geometry();b=make(s);old=calculate(copy.deepcopy(b),s);b['material_items'][0]['cutting_plan']=config();v=calculate(b,s);r=v['material_items'][0];assert v['quantity']==old['quantity'] and r['post_layout']==old['material_items'][0]['post_layout'] and r.get('post_spacing')==old['material_items'][0].get('post_spacing')
 assert r['stock_cutting']['required_piece_count']==r['post_layout']['post_count']*2 and Decimal(r['stock_cutting']['total']['net_mass'])==Decimal(v['quantity'])


@pytest.mark.parametrize('unit,factor',[('m',Decimal('.001')),('mm',Decimal('1')),('cm',Decimal('.1')),('dm',Decimal('.01')),('km',Decimal('.000001'))])
def test_five_explicit_length_units(unit,factor):
 s=geometry();b=body(s);c=b['material_items'][1]['cutting_plan'];c['length_unit']=unit
 for k in PARAMS:c[k]=str(Decimal(c[k])*factor)
 assert getplan(s,b)['total']['purchased_mass']=='54' and getplan(s,b)['capacity']=='4'


@pytest.mark.parametrize('weight,unit',[('3','kg/m'),('3000','g/m'),('.003','kg/mm'),('3','g/mm')])
def test_weight_units(weight,unit):
 s=geometry();b=body(s);b['material_items'][1].update(unit_weight=weight,weight_unit=unit);assert getplan(s,b)['total']['purchased_mass']=='54'


@pytest.mark.parametrize('minimum,disposition',[('1168','retain'),('1168.000000001','scrap'),('1167.999999999','retain')])
def test_reusable_threshold_exact_equality(minimum,disposition):
 s=geometry();b=body(s);b['material_items'][1]['cutting_plan']['reusable_min_length']=minimum;assert getplan(s,b)['patterns'][0]['offcut_disposition']==disposition


@pytest.mark.parametrize('stock,capacity',[('4832','4'),('4831.999999999','3'),('4832.000000001','4')])
def test_capacity_boundary_and_one_kerf_for_every_piece(stock,capacity):
 s=geometry();b=body(s);b['material_items'][1]['cutting_plan']['stock_length']=stock;assert getplan(s,b)['capacity']==capacity


def test_zero_offcut_and_zero_kerf_trim_do_not_create_inventory():
 s=geometry();b=body(s);m=b['material_items'][1];m['quantity']='6';m['cutting_plan']=config(stock='4800',kerf='0',start='0',end='0');g=getplan(s,b);assert g['stock_count']==3 and len(g['patterns'])==1 and g['patterns'][0]['offcut_disposition']=='none' and g['total']['retained_length']==g['total']['loss_length']=='0'


def test_last_only_pattern_and_exact_large_capacity_string():
 s=geometry();b=body(s);m=b['material_items'][1];m.update(quantity='1',cut_length='1e-15',length_unit='m');m['cutting_plan']=config(stock='1000000000000',unit='m',kerf='0',start='0',end='0',minimum='1');g=getplan(s,b);assert g['capacity']=='1000000000000000000000000000' and g['stock_count']==1 and g['patterns'][0]['pieces_per_stock']==2


@pytest.mark.parametrize('fault',['null','list','missing','extra','proof','extra-proof','blank-proof','basis','unconfirmed','truthy','negative','zero-stock','zero-threshold','nan','infinite','precision','unit','trim-too-long','no-capacity','fractional-demand','too-many-pieces','conflict','continuous','panel','zero-cut-unit'])
def test_strict_input_limits_and_mutually_exclusive_plans(fault):
 s=geometry();b=body(s);m=b['material_items'][1];c=m['cutting_plan']
 if fault=='null':m['cutting_plan']=None
 elif fault=='list':m['cutting_plan']=[]
 elif fault=='missing':c.pop('trim_start')
 elif fault=='extra':c['stock_count']=3
 elif fault=='proof':c['evidence'].pop('trim_end')
 elif fault=='extra-proof':c['evidence']['unit']='test'
 elif fault=='blank-proof':c['evidence']['stock_length']=''
 elif fault=='basis':c['measurement_basis']=''
 elif fault=='unconfirmed':c['cut_basis_confirmed']=False
 elif fault=='truthy':c['cut_basis_confirmed']=1
 elif fault=='negative':c['kerf_per_piece']='-1'
 elif fault=='zero-stock':c['stock_length']='0'
 elif fault=='zero-threshold':c['reusable_min_length']='0'
 elif fault=='nan':c['trim_start']='NaN'
 elif fault=='infinite':c['stock_length']='Infinity'
 elif fault=='precision':c['stock_length']='.1234567890123456789'
 elif fault=='unit':c['length_unit']='m²'
 elif fault=='trim-too-long':c['trim_start']='5990'
 elif fault=='no-capacity':c['stock_length']='1200'
 elif fault=='fractional-demand':s['source_geometry']['multiplicity']='.5';s['source_component_hash']=store.digest(snapshot(s['source_geometry']));b['source_component_hash']=store.digest({k:v for k,v in s.items() if k!='state'})
 elif fault=='too-many-pieces':m['quantity']='1000000'
 elif fault=='conflict':m['stock_allowance']=allowance()
 elif fault=='continuous':b['material_items'][0]['cutting_plan']=c;m.pop('cutting_plan')
 elif fault=='panel':b['material_items']=[rect_panel()];b['material_items'][0]['cutting_plan']=c
 else:m['cut_length']='.000000000000000001';m['length_unit']='mm'
 with pytest.raises(DomainError):calculate(b,s)


def test_fractional_derived_demand_is_rejected_before_cutting():
 from backend.stock_cutting import calculate as cut_quantity
 s=geometry();b=body(s);row=calculate(b,s)['material_items'][1]
 with pytest.raises(DomainError,match='整数'):cut_quantity(config(),row,Decimal('.5'),'test.cutting')


def test_million_pieces_use_compact_patterns_without_expansion():
 s=geometry();b=body(s);b['material_items'][1]['quantity']='500000';g=getplan(s,b);assert g['required_piece_count']==1000000 and g['stock_count']==250000 and len(g['patterns'])==1 and g['patterns'][0]['last_stock']==250000


def test_two_items_may_use_different_procurement_methods_without_mixing():
 s=geometry();b=body(s);b['material_items'][0]['stock_allowance']=allowance();v=calculate(b,s);assert v['quantity']=='180.000' and v['stock_summary']['item_count']==v['cutting_summary']['item_count']==1 and v['cutting_summary']['total']['purchased_mass']=='54'


def test_seeded_independent_integer_capacity_search_and_ledger_oracle():
 rng=random.Random(8001)
 for _ in range(60):
  cut=rng.randint(50,1500);kerf=rng.randint(0,5);start=rng.randint(0,20);end=rng.randint(0,20);stock=rng.randint(cut+kerf+start+end,8000);count=rng.randint(1,100);minimum=rng.randint(1,500)
  s=geometry();b=body(s);m=b['material_items'][1];m.update(quantity=str(count),cut_length=str(cut));m['cutting_plan']=config(stock=str(stock),kerf=str(kerf),start=str(start),end=str(end),minimum=str(minimum));g=getplan(s,b)
  capacity=0
  while (capacity+1)*(cut+kerf)+start+end<=stock:capacity+=1
  required=count*2;bars=0
  while bars*capacity<required:bars+=1
  assert (int(g['capacity']),g['stock_count'])==(capacity,bars) and sum(p['produced_count'] for p in g['patterns'])==required
  assert Fraction(g['total']['purchased_length'])==Fraction(bars*stock,1000) and Fraction(g['total']['net_length'])==Fraction(required*cut,1000)
  assert Fraction(g['total']['purchased_length'])==sum((Fraction(g['total'][k+'_length']) for k in ('net','kerf','trim','scrap','retained')),Fraction(0))


def test_low_caller_precision_does_not_change_exact_capacity():
 s=geometry();b=body(s)
 with localcontext() as c:c.prec=6;v=calculate(b,s)
 assert v==calculate(b,s)


@pytest.mark.parametrize('fmt',['json','xlsx','pdf'])
def test_full_workflow_exports_net_boq_and_old_quote_isolation(clients,fmt):
 e,r,pid,_,quote=priced_project(clients);e,r,pid,s,b,eid=setup(clients,pid);before=get(e,pid);post(e,pid,b,'/components/preview');assert get(e,pid)==before;c=post(e,pid,b).json();approve(r,pid,'components',c)
 for t in post(e,pid,{},'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
 boq=post(e,pid,{},'/boq/generate').json();assert any(v['quantity']=='180.000' and v['unit']=='kg' for v in boq)
 url=f'/api/v1/projects/{pid}/exports';res=e.get(url,params=dict(format=fmt,unpriced=True));assert res.status_code==200,res.text;titles={'定尺下料汇总','定尺下料模式','定尺长度质量台账','定尺取整与原件'}
 if fmt=='json':assert titles<={x[0] for x in res.json()['tables']}
 elif fmt=='xlsx':assert titles<=set(load_workbook(io.BytesIO(res.content)).sheetnames)
 else:assert all(t in ''.join(p.extract_text() for p in PdfReader(io.BytesIO(res.content)).pages) for t in titles)
 old=e.get(url,params=dict(format='json',quote_id=quote['id'])).json();assert old['quote']==quote and not any(t[0] in titles for t in old['tables'])


@pytest.mark.parametrize('fault',['input','proof','capacity','count','pattern','offcut','certificate','balance','ast','summary','book'])
def test_saved_calculation_rebuild_rejects_tampering(clients,fault):
 e,r,pid,s,b,eid=setup(clients);c=post(e,pid,b).json()
 if fault=='book':approve(r,pid,'components',c);post(e,pid,{},'/takeoffs/calculate')
 def forge(p):
  c=p['components'][-1];v=c['preview'];g=v['material_items'][1]['stock_cutting']
  if fault=='input':c['material_items'][1]['cutting_plan']['stock_length']='5000'
  elif fault=='proof':c['material_items'][1]['cutting_plan']['evidence']['stock_length']='forged'
  elif fault=='capacity':g['capacity']='5'
  elif fault=='count':g['stock_count']=2
  elif fault=='pattern':g['patterns'][0]['produced_count']=10
  elif fault=='offcut':g['patterns'][0]['offcut_disposition']='retain'
  elif fault=='certificate':g['capacity_certificate']['used_at_next']='6'
  elif fault=='balance':g['balance']['difference']='1'
  elif fault=='ast':g['total_formula_asts']['purchased_length']={'constant':'12'}
  elif fault=='summary':v['cutting_summary']['total']['purchased_mass']='36'
  else:p['takeoffs'][-1]['material_items'][1]['stock_cutting']['stock_count']=2
 mutate(pid,forge);coll='takeoffs' if fault=='book' else 'components';approve(r,pid,coll,get(e,pid)[coll][-1],expected=404 if fault=='proof' else 409)


def test_same_net_revision_removal_history_concurrency_source_invalidation(clients):
 e,r,pid,s,b,eid=setup(clients);c=post(e,pid,b).json();approve(r,pid,'components',c)
 for t in post(e,pid,{},'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
 post(e,pid,{},'/boq/generate');before=get(e,pid);old=before['components'][-1];new=copy.deepcopy(b);new['material_items'][1]['cutting_plan']['stock_length']='5000';new['object_hash']=store.digest(old);url=f'/api/v1/projects/{pid}/components/'+old['id'];res=e.put(url,json=new);assert res.status_code==200,res.text;p=get(e,pid);assert p['component_history'][-1]==old and p['boq']==before['boq'] and all(t['state']=='stale' for t in p['takeoffs']) and res.json()['preview']['quantity']=='180.000' and res.json()['preview']['cutting_summary']['total']['purchased_mass']=='45';assert e.put(url,json=new).status_code==409;approve(r,pid,'components',res.json())
 for t in post(e,pid,{},'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
 assert any(v['quantity']=='180.000' for v in post(e,pid,{},'/boq/generate').json())
 current=get(e,pid)['components'][-1];new['object_hash']=store.digest(current);new['material_items'][1].pop('cutting_plan');res=e.put(url,json=new);assert res.status_code==200 and 'cutting_summary' not in res.json()['preview'];assert get(e,pid)['component_history'][-1]==current;approve(r,pid,'components',res.json())
 s['object_hash']=store.digest(s);s['edge_decisions'][0].update(mode='ranges',ranges=[interval(0,10,eid)]);res=e.put(f'/api/v1/projects/{pid}/components/'+s['id'],json=s);assert res.status_code==200;approve(r,pid,'components',res.json());assert get(e,pid)['components'][-1]['state']=='stale'


@pytest.mark.parametrize('gate',['preview','save','review','calculate','book_review'])
def test_cutting_original_bytes_gate(clients,gate):
 e,r,pid,s,b,eid=setup(clients)
 if gate in ('review','calculate','book_review'):c=post(e,pid,b).json()
 if gate in ('calculate','book_review'):approve(r,pid,'components',c)
 if gate=='book_review':t=post(e,pid,{},'/takeoffs/calculate').json()[-1]
 d=get(e,pid)['documents'][-1];(store.ROOT/'originals'/d['hash']).write_bytes(b'corrupt')
 if gate=='preview':post(e,pid,b,'/components/preview',409)
 elif gate=='save':post(e,pid,b,expected=409)
 elif gate=='review':approve(r,pid,'components',c,expected=409)
 elif gate=='calculate':post(e,pid,{},'/takeoffs/calculate',409)
 else:approve(r,pid,'takeoffs',t,expected=409)


def test_cutting_roles_tenant_foreign_evidence(clients):
 e,r,pid,s,b,eid=setup(clients);post(r,pid,b,expected=403);assert clients['other'].get(f'/api/v1/projects/{pid}').status_code==404
 other=create(e)['id'];upload(e,other,'异项目');b['material_items'][1]['cutting_plan']['evidence']['kerf_per_piece']=get(e,other)['evidence'][-1]['id'];post(e,pid,b,expected=404)


def test_long_names_pdf_remains_exportable(clients):
 e,r,pid,s,b,eid=setup(clients);b['name']='合成'+('长名称验收'*39)+'末';b['material_items'][1]['name']='合成'+('长材料名'*39)+'末';c=post(e,pid,b).json();approve(r,pid,'components',c)
 for t in post(e,pid,{},'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
 res=e.get(f'/api/v1/projects/{pid}/exports',params=dict(format='pdf',unpriced=True));assert res.status_code==200,res.text
