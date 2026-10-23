"""Synthetic tier stacks, source counts, precision and review lifecycle."""
import copy,io,random
from decimal import Decimal,localcontext,ROUND_HALF_UP
import pytest
from openpyxl import load_workbook
from pypdf import PdfReader
from backend import store
from backend.engine import DomainError
from backend.post_foundation import calculate
from tests.test_post_foundation import body as rectangular_body,source,setup as rectangular_setup,row as rect_row,equal_body,fixed_body,fixed_added,equal_added,fixed_adjusted,equal_adjusted,combined
from tests.test_polygon_takeoff import post,mutate
from tests.test_route_takeoff import evaluate
from tests.test_workflow import get,approve,upload,create,priced_project

TITLES={'台阶基础逐层体积','台阶基础逐层原件'}


def level(length='600',width='500',height='200',eid='test'):
 return dict(parameters=dict(length=length,width=width,height=height),evidence={k:eid for k in ('length','width','height')})


def stepped(eid='test',index=1):
 return dict(material_index=index,shape='stepped_rectangular',levels=[level(eid=eid),level('400','300','400',eid)],length_unit='mm',measurement_basis='合成每根立柱一个同心同轴矩形台阶基础，下层600×500×净高200mm，上层400×300×净高400mm，自下而上紧接无空隙、不外伸，层高非累计，不重复计基础数量，无实际工程效力',evidence=dict(measurement_basis=eid),stacking_confirmed=True)


def body(s,eid='test'):
 b=rectangular_body(s,eid);b['name']='合成台阶立柱基础（非真实工程）';b['foundation_items']=[stepped(eid)];b['measurement_basis']='合成一柱一台阶基础，各层同心垂直接触无空隙、上层不外伸，层高为本层净高；同一材料且无重复清单，不含未列孔洞或交叠，来源重复一次';return b


def setup(clients,pid=None):
 e,r,pid,m,b,old_eid=rectangular_setup(clients,pid)
 upload(e,pid,'合成台阶基础软件验收，非真实工程。已审最终立柱单组11根、重复2共22个基础，材料79.2kg保持。每个基础两层，自下而上同心同轴垂直紧接、无空隙且不外伸，下层600×500×净层高200mm体积.06m3，上层400×300×净层高400mm体积.048m3，底面相对Z=0、层界.2、顶.6m，每基础.108m3，全部2.376m3，总量修约3位，层数不乘基础个数。同一合成C30材料、无锚栓孔及预埋扣减，不含土方垫层钢筋模板，与既有基础清单不重复。上层净高改500mm，顶.7m，每基础.12m3，全部2.640m3，柱数和材料质量不变须重审。上层长改601mm会外伸须拒。原矩形基础400×500×600mm每个.12m3，全部2.64m3，本台阶修订替换原基础而非追加。','foundation-steps-basis.txt')
 for c in get(e,pid)['components']:approve(r,pid,'components',c)
 m=get(e,pid)['components'][-1];eid=get(e,pid)['evidence'][-1]['id'];return e,r,pid,m,body(m,eid),eid


def test_tier_geometry_counts_asts_and_no_mutation():
 s=source();b=body(s);before=copy.deepcopy((s,b));r=calculate(b,s);d=r['foundation_items'][0];assert (s,b)==before and r['quantity']=='2.376' and r['rule_version']=='railing-post-foundation-1.1' and r['foundation_summary']['total_count']=='22';assert d['single_foundation_volume']=='0.108' and d['single_group_volume']=='1.188'
 assert [tuple(Decimal(v[k]) for k in ('from_z','to_z','single_level_volume','total_volume')) for v in d['foundation_levels']]==[tuple(map(Decimal,vs)) for vs in [('0','.2','.06','1.32'),('.2','.6','.048','1.056')]]
 assert evaluate(r['formula_ast'],r['dimensions'])==Decimal(r['quantity'])
 for v in d['foundation_levels']:
  for k,a in v['formula_asts'].items():assert evaluate(a,r['dimensions'])==Decimal(v[k])
 for k,a in d['formula_asts'].items():assert evaluate(a,r['dimensions'])==Decimal(d[{'single_volume':'single_foundation_volume'}.get(k,k)])


@pytest.mark.parametrize('make',[equal_body,fixed_body,fixed_added,equal_added,fixed_adjusted,equal_adjusted,combined,lambda s:combined(s,equal=True)])
def test_all_eight_final_post_modes(make):
 s=source(make);r=calculate(body(s),s);assert Decimal(r['gross'])==Decimal('.108')*r['foundation_summary']['single_group_count']*2


@pytest.mark.parametrize('unit,factor',[('m',Decimal('.001')),('mm',Decimal('1')),('cm',Decimal('.1')),('dm',Decimal('.01')),('km',Decimal('.000001'))])
def test_units_and_low_context(unit,factor):
 s=source();b=body(s);d=b['foundation_items'][0];d['length_unit']=unit
 for v in d['levels']:
  for k in v['parameters']:v['parameters'][k]=str(Decimal(v['parameters'][k])*factor)
 with localcontext() as ctx:ctx.prec=6;r=calculate(b,s)
 assert r==calculate(b,s) and r['quantity']=='2.376'


@pytest.mark.parametrize('length,accepted',[('600',True),('599.999999999',True),('600.000000001',False)])
def test_exact_non_overhang_boundary(length,accepted):
 s=source();b=body(s);b['foundation_items'][0]['levels'][1]['parameters']['length']=length
 if accepted:assert calculate(b,s)['foundation_summary']['total_count']=='22'
 else:
  with pytest.raises(DomainError,match='上层台阶'):calculate(b,s)


def test_maximum_ten_tiers_equal_footprints_and_partial_mixed_selection():
 s=source();s['material_items']=[{**copy.deepcopy(s['material_items'][0]),'name':'合成柱'+str(i)} for i in range(3)];b=body(s);b['foundation_items'][0]['levels']=[level(height='100') for _ in range(10)];b['foundation_items'].append(rect_row(2));r=calculate(b,s);assert r['foundation_summary']['selected_item_count']==2 and r['foundation_summary']['unselected_post_item_count']==1 and r['quantity']=='9.240';assert Decimal(r['foundation_items'][0]['foundation_levels'][-1]['to_z'])==1;b['foundation_items'].reverse();assert calculate(b,s)==r


def test_seeded_independent_layer_sum_and_final_rounding():
 rng=random.Random(8701)
 for _ in range(60):
  s=source();b=body(s);d=b['foundation_items'][0];a,z=1000,900;expected=Decimal(0);levels=[]
  with localcontext() as ctx:
   ctx.prec=80
   for _ in range(rng.randint(2,10)):
    a=rng.randint(1,a);z=rng.randint(1,z);h=rng.randint(1,999);levels.append(level(str(a),str(z),str(h)));expected+=Decimal(a)*z*h/10**9
   d['levels']=levels;b['rounding_digits']=rng.randint(0,12);r=calculate(b,s);expected*=22;assert Decimal(r['gross'])==expected and Decimal(r['quantity'])==expected.quantize(Decimal(1).scaleb(-b['rounding_digits']),rounding=ROUND_HALF_UP)


@pytest.mark.parametrize('fault',['shape-null','shape-other','shape-list','missing-shape','extra-row','old-parameters','old-proofs','missing-row','unconfirmed','truthy','null-levels','dict-levels','empty','one','eleven','level-null','level-list','level-extra','level-missing','parameters-null','parameter-missing','parameter-extra','proofs-null','proof-missing','proof-extra','proof-blank','proof-bool','negative','zero','bool','nan','infinite','precision','unit','overhang-length','overhang-width','reverse','scope-blank','scope-proof-blank','scope-proof-bool'])
def test_strict_tier_inputs(fault):
 s=source();b=body(s);d=b['foundation_items'][0];v=d['levels'][0]
 if fault=='shape-null':d['shape']=None
 elif fault=='shape-other':d['shape']='circular'
 elif fault=='shape-list':d['shape']=[]
 elif fault=='missing-shape':d.pop('shape')
 elif fault=='extra-row':d['quantity']='22'
 elif fault=='old-parameters':d['parameters']=dict(length='1',width='1',height='1')
 elif fault=='old-proofs':d['evidence']['height']='test'
 elif fault=='missing-row':d.pop('stacking_confirmed')
 elif fault=='unconfirmed':d['stacking_confirmed']=False
 elif fault=='truthy':d['stacking_confirmed']=1
 elif fault=='null-levels':d['levels']=None
 elif fault=='dict-levels':d['levels']={}
 elif fault=='empty':d['levels']=[]
 elif fault=='one':d['levels']=[v]
 elif fault=='eleven':d['levels']=[v]*11
 elif fault=='level-null':d['levels'][0]=None
 elif fault=='level-list':d['levels'][0]=[]
 elif fault=='level-extra':v['z']='0'
 elif fault=='level-missing':v.pop('evidence')
 elif fault=='parameters-null':v['parameters']=None
 elif fault=='parameter-missing':v['parameters'].pop('height')
 elif fault=='parameter-extra':v['parameters']['cumulative_height']='200'
 elif fault=='proofs-null':v['evidence']=None
 elif fault=='proof-missing':v['evidence'].pop('height')
 elif fault=='proof-extra':v['evidence']['z']='test'
 elif fault=='proof-blank':v['evidence']['height']=''
 elif fault=='proof-bool':v['evidence']['height']=True
 elif fault in ('negative','zero','bool','nan','infinite','precision'):v['parameters']['height']={'negative':'-1','zero':'0','bool':True,'nan':'NaN','infinite':'Infinity','precision':'.1234567890123456789'}[fault]
 elif fault=='unit':d['length_unit']='m³'
 elif fault=='overhang-length':d['levels'][1]['parameters']['length']='601'
 elif fault=='overhang-width':d['levels'][1]['parameters']['width']='501'
 elif fault=='reverse':d['levels'].reverse()
 elif fault=='scope-blank':d['measurement_basis']=''
 elif fault=='scope-proof-blank':d['evidence']['measurement_basis']=''
 else:d['evidence']['measurement_basis']=True
 with pytest.raises(DomainError):calculate(b,s)


@pytest.mark.parametrize('fmt',['json','xlsx','pdf'])
def test_workflow_export_boq_and_old_quote(clients,fmt):
 e,r,pid,_,quote=priced_project(clients)
 if fmt=='json':upload(e,pid,'仅合成编码依据：TEST-PROFILE是软件测试编码，非官方条目，无真实计价效力。','mapping-code.txt',specialty='合同计价')
 e,r,pid,s,b,eid=setup(clients,pid);before=get(e,pid);res=post(e,pid,b,'/components/preview');assert res.json()['quantity']=='2.376' and get(e,pid)==before;c=post(e,pid,b).json();approve(r,pid,'components',c)
 for t in post(e,pid,{},'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
 boq=post(e,pid,{},'/boq/generate').json();assert any(v['quantity']=='2.376' and v['unit']=='m³' for v in boq) and any(v['quantity']=='79.200' for v in boq);url=f'/api/v1/projects/{pid}/exports';res=e.get(url,params=dict(format=fmt,unpriced=True));assert res.status_code==200,res.text
 if fmt=='json':assert TITLES<={t[0] for t in res.json()['tables']}
 elif fmt=='xlsx':assert TITLES<=set(load_workbook(io.BytesIO(res.content)).sheetnames)
 else:assert all(v in ''.join(p.extract_text() for p in PdfReader(io.BytesIO(res.content)).pages) for v in TITLES)
 assert not any(t[0] in TITLES for t in e.get(url,params=dict(format='json',quote_id=quote['id'])).json()['tables'])
 if fmt=='json':
  from tests.test_mapping_profiles import rule_body,save_rule,preview,selection
  p=get(e,pid);rule=rule_body(p,'post_foundation','volume',next(v['id'] for v in p['evidence'] if v['document_id']==next(d['id'] for d in p['documents'] if d['name']=='mapping-code.txt')));assert e.patch(f'/api/v1/projects/{pid}/settings',json=dict(boq_standard=rule['standard'],boq_edition=rule['edition'],reason='仅合成台阶基础清单口径验收')).status_code==200;save_rule(e,r,pid,rule);bid=next(v['id'] for v in p['boq'] if v['takeoff_id']==p['takeoffs'][-1]['id']);chosen=preview(e,pid,bid);assert chosen['state']=='unique',chosen;assert e.post(f'/api/v1/projects/{pid}/boq/apply-mappings',json=dict(selections=[selection(chosen)])).status_code==200;current=next(v for v in get(e,pid)['boq'] if v['id']==bid);assert current['quantity']=='2.376' and current['unit']=='m³';approve(r,pid,'boq',current)



@pytest.mark.parametrize('fault',['raw','proof','confirmation','shape','z','volume','ast','count','book'])
def test_tampering_rebuilt(clients,fault):
 e,r,pid,s,b,eid=setup(clients);c=post(e,pid,b).json()
 if fault=='book':approve(r,pid,'components',c);post(e,pid,{},'/takeoffs/calculate')
 def forge(p):
  c=p['components'][-1];d=c['preview']['foundation_items'][0]
  if fault=='raw':c['foundation_items'][0]['levels'][1]['parameters']['height']='500'
  elif fault=='proof':c['foundation_items'][0]['levels'][1]['evidence']['height']='forged'
  elif fault=='confirmation':c['foundation_items'][0]['stacking_confirmed']=False
  elif fault=='shape':d['shape']='other'
  elif fault=='z':d['foundation_levels'][1]['from_z']='0'
  elif fault=='volume':d['foundation_levels'][1]['total_volume']='2'
  elif fault=='ast':d['foundation_levels'][1]['formula_asts']['from_z']={'constant':'0'}
  elif fault=='count':d['total_count']='44'
  else:p['takeoffs'][-1]['foundation_items'][0]['foundation_levels'][1]['to_z']='.7'
 mutate(pid,forge);coll='takeoffs' if fault=='book' else 'components';approve(r,pid,coll,get(e,pid)[coll][-1],expected=404 if fault=='proof' else 422 if fault=='confirmation' else 409)


def test_revision_history_claims_switch_back_and_source(clients):
 e,r,pid,s,b,eid=setup(clients);c=post(e,pid,b).json();approve(r,pid,'components',c)
 for t in post(e,pid,{},'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
 post(e,pid,{},'/boq/generate');p=get(e,pid);old=p['components'][-1];new=copy.deepcopy(b);new['foundation_items'][0]['levels'][1]['parameters']['height']='500';new['object_hash']=store.digest(old);url=f'/api/v1/projects/{pid}/components/'+old['id'];res=e.put(url,json=new);assert res.status_code==200 and res.json()['preview']['quantity']=='2.640';assert get(e,pid)['component_history'][-1]==old and get(e,pid)['boq']==p['boq'] and all(t['state']=='stale' for t in get(e,pid)['takeoffs']);assert e.put(url,json=new).status_code==409;post(e,pid,{**b,'id':old['id']},expected=409);approve(r,pid,'components',res.json())
 current=get(e,pid)['components'][-1];rect=rectangular_body(s,eid);rect.update(id=old['id'],object_hash=store.digest(current));res=e.put(url,json=rect);assert res.status_code==200 and res.json()['preview']['rule_version']=='railing-post-foundation-1.0';approve(r,pid,'components',res.json())
 s['object_hash']=store.digest(s);s['material_items'][0]['post_ranges'][0]['spacing']='3000';res=e.put(f'/api/v1/projects/{pid}/components/'+s['id'],json=s);assert res.status_code==200;approve(r,pid,'components',res.json());assert get(e,pid)['components'][-1]['state']=='stale'


@pytest.mark.parametrize('gate',['preview','save','review','calculate','book_review'])
def test_original_bytes_gates(clients,gate):
 e,r,pid,s,b,eid=setup(clients)
 if gate in ('review','calculate','book_review'):c=post(e,pid,b).json()
 if gate in ('calculate','book_review'):approve(r,pid,'components',c)
 if gate=='book_review':t=post(e,pid,{},'/takeoffs/calculate').json()[-1]
 d=get(e,pid)['documents'][-1];(store.ROOT/'originals'/d['hash']).write_bytes(b'corrupt')
 if gate=='preview':post(e,pid,b,'/components/preview',409)
 elif gate=='save':post(e,pid,b,expected=409)
 elif gate=='review':approve(r,pid,'components',c,expected=409)
 elif gate=='calculate':post(e,pid,{},'/takeoffs/calculate',409)
 else:approve(r,pid,'takeoffs',t,expected=409)


def test_roles_tenant_and_foreign_layer_evidence(clients):
 e,r,pid,s,b,eid=setup(clients);post(r,pid,b,expected=403);assert clients['other'].get(f'/api/v1/projects/{pid}').status_code==404;oid=create(e)['id'];upload(e,oid,'异项目原件');fid=get(e,oid)['evidence'][-1]['id'];b['foundation_items'][0]['levels'][1]['evidence']['height']=fid;post(e,pid,b,expected=404)


def test_long_names_pdf(clients):
 e,r,pid,s,b,eid=setup(clients);b['name']='合成'+('台阶长名称'*39)+'末';b['foundation_items'][0]['levels']=[level(height='100',eid=eid) for _ in range(10)];c=post(e,pid,b).json();approve(r,pid,'components',c)
 for t in post(e,pid,{},'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
 res=e.get(f'/api/v1/projects/{pid}/exports',params=dict(format='pdf',unpriced=True));assert res.status_code==200,res.text
