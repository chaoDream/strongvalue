"""Synthetic equal-bay layout and mass. Not a structural design acceptance."""
import copy,io,random
from fractions import Fraction
from decimal import Decimal,localcontext,ROUND_HALF_UP
import pytest
from openpyxl import load_workbook
from pypdf import PdfReader
from backend.engine import DomainError
from backend import store
from backend.railing_material import calculate
from backend.railing_posts import ranges
from backend.railing_takeoff import calculate as railing_quantity,snapshot as source_snapshot
from tests.test_railing_material import geometry,payload,setup as material_setup
from tests.test_polygon_material import panel as polygon_panel
from tests.test_panel_material import panel as rectangle_panel
from tests.test_railing_takeoff import interval
from tests.test_route_takeoff import point,evaluate
from tests.test_route_railing import payload as railing_payload
from tests.test_polygon_takeoff import post,mutate
from tests.test_workflow import get,approve,upload,priced_project,create


def member(s,eid='test'):
 candidates=ranges(railing_quantity(s));plans=[]
 for i,r in enumerate(candidates):
  plans.append(dict(range_id=r['range_id'],spacing='2000',spacing_unit='mm',start_post=i==0,end_post=True,measurement_basis='核定最大间距2m均分跨间；公共转折点归前段终端，本段起端不另计，首尾端有柱；按来源S口径，非固定2m余跨',evidence={k:eid for k in ('spacing','start_post','end_post','measurement_basis')}))
 return dict(name='合成均分立柱',mode='spaced_posts',cut_length='1200',length_unit='mm',unit_weight='3000',weight_unit='g/m',post_ranges=plans,measurement_basis='每根实际下料1200mm、3000g/m；各区间均分跨间、两端归属如逐项依据，不含构造加柱及损耗',evidence={k:eid for k in ('cut_length','unit_weight','measurement_basis')})


def body(s,eid='test',mixed=False):
 b=payload(s,eid);b['material_items']=(b['material_items']+[rectangle_panel(eid),polygon_panel(eid)] if mixed else [])+[member(s,eid)];b['measurement_basis']='合成均分布柱及独立下料质量；共享端点一次，来源重复一次，无实际工程效力';return b


def setup(clients,pid=None):
 e,r,pid,s,_,eid=material_setup(clients,pid)
 upload(e,pid,'合成软件布柱验收，非真实工程：栏杆AB三维13m、BC5m，来源重复2。各区间核定最大间距2000mm，均分跨间，首端与终端计柱，共享B仅归AB终端、BC起端不计；AB7跨、6内部柱+两端共8根，BC3跨、2内部柱+终端共3根，单组11根。每根实际下料1200mm，每米质量3000g/m，单组39.6kg、总79.2kg，总质量修约3位。AB最大间距改3000mm时5跨6根，BC仍3根，合计9根、64.8kg。相邻局部区间、闭合末点需按原件明确归属，不推断交叉加柱。','equal-bay-posts-basis.txt')
 for c in get(e,pid)['components']:approve(r,pid,'components',c)
 s=get(e,pid)['components'][-1];eid=get(e,pid)['evidence'][-1]['id'];return e,r,pid,s,body(s,eid),eid


def test_equal_bay_layout_mass_count_once_and_asts():
 s=geometry();b=body(s);before=copy.deepcopy((s,b));v=calculate(b,s);assert (s,b)==before and v['quantity']=='79.200' and v['rule_version']=='railing-material-mass-1.3'
 row=v['material_items'][0];g=row['post_layout'];assert g['post_count']==11 and [r['interval_count'] for r in g['post_ranges']]==[7,3] and [r['post_count'] for r in g['post_ranges']]==[8,3]
 assert len(g['post_stations'])==11 and 'material.1.quantity' not in v['dimensions']
 with localcontext() as c:
  c.prec=80;assert evaluate(v['formula_ast'],v['dimensions'])==Decimal('79.2');assert evaluate(g['formula_ast'],v['dimensions'])==11
  for r in g['post_ranges']:
   cert=r['ceil_certificate'];a,z,bound=[Fraction(int(cert[k]['numerator']),int(cert[k]['denominator'])) for k in ('lower_spacing_square','range_length_square','upper_spacing_square')];assert a<z<=bound
  for p in g['post_stations']:assert abs(evaluate(p['station_formula_ast'],v['dimensions'])-Decimal(p['s']))<Decimal('1e-75')
 assert Decimal(calculate(body(s,mixed=True),s)['quantity'])==Decimal('806.4')


def test_shared_vertex_rejected_until_explicitly_owned_once():
 s=geometry();b=body(s);b['material_items'][0]['post_ranges'][1]['start_post']=True
 with pytest.raises(DomainError,match='共享端点'):calculate(b,s)
 b['material_items'][0]['post_ranges'][0]['end_post']=False;assert calculate(b,s)['quantity']=='79.200'


@pytest.mark.parametrize('unit,value',[('m','2'),('cm','200'),('mm','2000'),('dm','20'),('km','.002')])
def test_spacing_units(unit,value):
 s=geometry();b=body(s)
 for r in b['material_items'][0]['post_ranges']:r.update(spacing=value,spacing_unit=unit)
 assert calculate(b,s)['quantity']=='79.200'


def test_horizontal_and_spatial_spacing_bases_are_distinct():
 s=geometry(basis='horizontal');v=calculate(body(s),s);assert v['material_items'][0]['post_layout']['post_count']==7 and v['quantity']=='50.400'


def test_exact_square_ceil_above_and_below_irrational_boundary():
 s=geometry();src=s['source_geometry'];src['route_points']=[point('A',0,0,0),point('B',1,1,0)];s['source_component_hash']=store.digest(source_snapshot(src));s['edge_decisions']=railing_payload(src)['edge_decisions'];b=body(s);r=b['material_items'][0]['post_ranges'][0];r.update(spacing='1.41421356237309505',spacing_unit='m')
 assert calculate(b,s)['material_items'][0]['post_layout']['post_ranges'][0]['interval_count']==1
 r['spacing']='1.41421356237309504';assert calculate(b,s)['material_items'][0]['post_layout']['post_ranges'][0]['interval_count']==2


def test_partial_ranges_gaps_and_touching_interval_ownership():
 s=geometry();s['edge_decisions'][0].update(mode='ranges',ranges=[interval(0,2),interval(4,10)]);b=body(s)
 # Each isolated range can own its start; the following BC start is a distinct vertex.
 for r in b['material_items'][0]['post_ranges']:r['start_post']=True
 v=calculate(b,s);g=v['material_items'][0]['post_layout'];assert [r['post_count'] for r in g['post_ranges']]==[2,4,4] and g['post_count']==10
 s['edge_decisions'][0]['ranges']=[interval(0,2),interval(2,4)];b=body(s)
 for r in b['material_items'][0]['post_ranges']:r['start_post']=True
 with pytest.raises(DomainError,match='共享端点'):calculate(b,s)
 b['material_items'][0]['post_ranges'][1]['start_post']=False;assert calculate(b,s)['material_items'][0]['post_layout']['post_count']==7


def test_closed_polygon_vertex_ownership_and_holes():
 from tests.test_railing_takeoff import geometry as pg,payload as pp
 src=pg();s=pp(src);s.update(id='polygon',revision=1,source_geometry=source_snapshot(src))
 for d in s['edge_decisions']:d.update(mode='full',ranges=[])
 b=body(s)
 for r in b['material_items'][0]['post_ranges']:r.update(start_post=True,end_post=False)
 v=calculate(b,s);assert v['material_items'][0]['post_layout']['post_count']>0
 b['material_items'][0]['post_ranges'][0]['end_post']=True
 with pytest.raises(DomainError,match='共享端点'):calculate(b,s)


@pytest.mark.parametrize('fault',['missing','extra','unknown','duplicate','not_list','missing_range_field','extra_range_field','blank_basis','missing_proof','extra_proof','blank_proof','zero_spacing','negative_spacing','bool_spacing','unit','nonfinite','zero_length','manual_count','count_proof','string_start','missing_end','top_scope','end_type','too_dense'])
def test_strict_layout_inputs(fault):
 s=geometry();b=body(s);v=b['material_items'][0];r=v['post_ranges'][0]
 if fault=='missing':v['post_ranges'].pop()
 elif fault=='extra':v['post_ranges'].append(copy.deepcopy(r))
 elif fault=='unknown':r['range_id']='old'
 elif fault=='duplicate':v['post_ranges'][1]['range_id']=r['range_id']
 elif fault=='not_list':v['post_ranges']={}
 elif fault=='missing_range_field':r.pop('spacing')
 elif fault=='extra_range_field':r['quantity']='9'
 elif fault=='blank_basis':r['measurement_basis']=''
 elif fault=='missing_proof':r['evidence'].pop('spacing')
 elif fault=='extra_proof':r['evidence']['quantity']='test'
 elif fault=='blank_proof':r['evidence']['spacing']=''
 elif fault=='zero_spacing':r['spacing']='0'
 elif fault=='negative_spacing':r['spacing']='-1'
 elif fault=='bool_spacing':r['spacing']=True
 elif fault=='unit':r['spacing_unit']='m²'
 elif fault=='nonfinite':r['spacing']='NaN'
 elif fault=='zero_length':v['cut_length']='0'
 elif fault=='manual_count':v['quantity']='11'
 elif fault=='count_proof':v['evidence']['quantity']='test'
 elif fault=='string_start':r['start_post']='false'
 elif fault=='missing_end':r.pop('end_post')
 elif fault=='top_scope':b['scope_confirmed']=False
 elif fault=='end_type':r['end_post']=0
 elif fault=='too_dense':r.update(spacing='.001',spacing_unit='m')
 with pytest.raises(DomainError):calculate(b,s)


def test_no_generated_posts_rejected_and_limits():
 s=geometry();b=body(s)
 for r in b['material_items'][0]['post_ranges']:r.update(spacing='100',spacing_unit='m',start_post=False,end_post=False)
 with pytest.raises(DomainError,match='任何立柱'):calculate(b,s)
 src=s['source_geometry'];src['route_points']=[point('A',0,0,0),point('B',999,0,0)];s['source_component_hash']=store.digest(source_snapshot(src));s['edge_decisions']=railing_payload(src)['edge_decisions'];b=body(s);b['material_items'][0]['post_ranges'][0].update(spacing='1',spacing_unit='m');assert calculate(b,s)['material_items'][0]['post_layout']['post_count']==1000
 b['material_items']=[{**copy.deepcopy(b['material_items'][0]),'name':str(i)} for i in range(5)];assert Decimal(calculate(b,s)['quantity'])==36000
 b['material_items'].append({**copy.deepcopy(b['material_items'][0]),'name':'sixth'})
 with pytest.raises(DomainError,match='5000'):calculate(b,s)
 b['material_items']=b['material_items'][:1];b['material_items'][0]['post_ranges'][0]['spacing']='.999'
 with pytest.raises(DomainError,match='1000根'):calculate(b,s)


def test_seeded_independent_integer_ceiling_oracle():
 rng=random.Random(7001)
 for _ in range(35):
  s=geometry();src=s['source_geometry'];L=rng.randint(1,1000);src['route_points']=[point('A',0,0,0),point('B',L,0,0)];s['source_component_hash']=store.digest(source_snapshot(src));s['edge_decisions']=railing_payload(src)['edge_decisions'];b=body(s);v=b['material_items'][0];r=v['post_ranges'][0];d=rng.randint(1,50);r.update(spacing=str(d),spacing_unit='m');n=(L+d-1)//d
  assert calculate(b,s)['material_items'][0]['post_layout']['post_count']==n+1


@pytest.mark.parametrize('fmt',['json','xlsx','pdf'])
def test_full_workflow_exports_old_quote_and_readonly(clients,fmt):
 e,r,pid,_,quote=priced_project(clients);e,r,pid,s,b,eid=setup(clients,pid);before=get(e,pid);candidate=e.get(f'/api/v1/projects/{pid}/railing-material-candidates',params=dict(component_id=s['id']));assert candidate.status_code==200 and len(candidate.json()['post_ranges'])==2;post(e,pid,b,'/components/preview');assert get(e,pid)==before
 c=post(e,pid,b).json();approve(r,pid,'components',c)
 for t in post(e,pid,{},'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
 assert any(v['quantity']=='79.200' for v in post(e,pid,{},'/boq/generate').json())
 url=f'/api/v1/projects/{pid}/exports';res=e.get(url,params=dict(format=fmt,unpriced=True));assert res.status_code==200,res.text
 titles=['立柱均分布置质量','立柱逐区间间距与归属','立柱跨间数精确校验','立柱逐根端距']
 if fmt=='json':assert set(titles)<={t[0] for t in res.json()['tables']}
 elif fmt=='xlsx':assert set(titles)<=set(load_workbook(io.BytesIO(res.content)).sheetnames)
 else:assert all(v in ''.join(p.extract_text() for p in PdfReader(io.BytesIO(res.content)).pages) for v in titles)
 old=e.get(url,params=dict(format='json',quote_id=quote['id'])).json();assert old['quote']==quote and not any(t[0] in titles for t in old['tables'])


@pytest.mark.parametrize('fault',['count','station','certificate','formula','spacing','owner','source','book'])
def test_saved_and_book_tampering_rebuilt(clients,fault):
 e,r,pid,s,b,eid=setup(clients);c=post(e,pid,b).json()
 if fault=='book':approve(r,pid,'components',c);post(e,pid,{},'/takeoffs/calculate')
 def forge(p):
  c=p['components'][-1];v=c['preview']['material_items'][0];g=v['post_layout']
  if fault=='count':g['post_count']=999
  elif fault=='station':g['post_stations'][1]['s']='2'
  elif fault=='certificate':g['post_ranges'][0]['ceil_certificate']['range_length_square']['numerator']='170'
  elif fault=='formula':g['formula_ast']={'constant':'11'}
  elif fault=='spacing':c['material_items'][0]['post_ranges'][0]['spacing']='3000'
  elif fault=='owner':c['material_items'][0]['post_ranges'][0]['start_post']=False
  elif fault=='source':c['source_geometry']['edge_decisions'][0]['mode']='exclude'
  else:p['takeoffs'][-1]['material_items'][0]['post_layout']['post_count']=999
 mutate(pid,forge);coll='takeoffs' if fault=='book' else 'components';approve(r,pid,coll,get(e,pid)[coll][-1],expected=409)


def test_revision_history_concurrency_source_and_boq(clients):
 e,r,pid,s,b,eid=setup(clients);c=post(e,pid,b).json();approve(r,pid,'components',c)
 for t in post(e,pid,{},'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
 post(e,pid,{},'/boq/generate');old=get(e,pid)['components'][-1];new=copy.deepcopy(b);new['material_items'][0]['post_ranges'][0]['spacing']='3000';new['object_hash']=store.digest(old);url=f'/api/v1/projects/{pid}/components/'+c['id'];res=e.put(url,json=new);assert res.status_code==200,res.text
 p=get(e,pid);assert p['component_history'][-1]==old and all(t['state']=='stale' for t in p['takeoffs']) and any(v['quantity']=='79.200' for v in p['boq']);assert res.json()['preview']['quantity']=='64.800' and e.put(url,json=new).status_code==409;approve(r,pid,'components',res.json())
 for t in post(e,pid,{},'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
 assert any(v['quantity']=='64.800' for v in post(e,pid,{},'/boq/generate').json())
 s['object_hash']=store.digest(s);s['edge_decisions'][0].update(mode='ranges',ranges=[interval(0,10,eid)]);res=e.put(f'/api/v1/projects/{pid}/components/'+s['id'],json=s);assert res.status_code==200,res.text;approve(r,pid,'components',res.json());assert get(e,pid)['components'][-1]['state']=='stale';post(e,pid,b,'/components/preview',409)


@pytest.mark.parametrize('gate',['preview','save','review','calculate','book_review'])
def test_original_bytes_gates(clients,gate):
 e,r,pid,s,b,eid=setup(clients)
 if gate in ('review','calculate','book_review'):c=post(e,pid,b).json()
 if gate in ('calculate','book_review'):approve(r,pid,'components',c)
 if gate=='book_review':t=post(e,pid,{},'/takeoffs/calculate').json()[-1]
 doc=get(e,pid)['documents'][-1];(store.ROOT/'originals'/doc['hash']).write_bytes(b'corrupt')
 if gate=='preview':post(e,pid,b,'/components/preview',409)
 elif gate=='save':post(e,pid,b,expected=409)
 elif gate=='review':approve(r,pid,'components',c,expected=409)
 elif gate=='calculate':post(e,pid,{},'/takeoffs/calculate',409)
 else:approve(r,pid,'takeoffs',t,expected=409)


def test_roles_tenant_and_foreign_layout_evidence(clients):
 e,r,pid,s,b,eid=setup(clients);post(r,pid,b,expected=403);assert clients['other'].get(f'/api/v1/projects/{pid}').status_code==404
 other=create(e)['id'];upload(e,other,'其他工程合成原件');foreign=get(e,other)['evidence'][-1]['id'];b['material_items'][0]['post_ranges'][0]['evidence']['start_post']=foreign;post(e,pid,b,expected=404)
