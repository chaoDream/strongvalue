"""Synthetic exact boundary extraction and independently enumerated grid islands."""
import copy
import io
import json
import random
from collections import deque
from decimal import Decimal, localcontext
from fractions import Fraction
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pypdf import PdfReader
from backend import store
from backend.engine import DomainError
from backend.main import app
from backend.surface_takeoff import boundary, partition, calculate
from backend.polygon_connected_partitions import separate, MAX_EDGE_TESTS, exact_point
from backend.polygon_takeoff import inside
from tests.test_polygon_takeoff import ring
from tests.test_polygon_cap_partitions import geometry as cap_geometry, groups as cap_groups, post, mutate
from tests.test_polygon_line_partitions import groups as line_groups, line
from tests.test_surface_takeoff import payload as surface
from tests.test_workflow import create, get, upload, approve, priced_project
from tests.test_route_takeoff import evaluate


U=[(0,0),(6,0),(6,6),(4,6),(4,2),(2,2),(2,6),(0,6)]


def geometry(eid='test',holes=True):
    s=cap_geometry(eid);s['polygon_rings']=[ring('U形外轮廓',U,eid)]
    if holes:
        s['polygon_rings'] += [ring('左臂孔',[(.5,4),(1.5,4),(1.5,5),(.5,5)],eid,'hole'),ring('右臂孔',[(4.5,4),(5.5,4),(5.5,5),(4.5,5)],eid,'hole')]
    s['measurement_basis']='合成U外28m²，两臂各孔1m²，净26m²，厚200mm、重复2；不是专业计量依据'
    return s


def groups(s,eid='test',y='3',top=True):
    g=cap_groups(s,eid,top=top,y=y);g[0]['u_cuts']=[];g[0]['separate_regions']=True;return g


def pieces(result):return [f for f in result['faces'] if f.get('cap_component_rings')]


def frac(f):return Fraction(int(f['cap_rational_area']['numerator']),int(f['cap_rational_area']['denominator']))


def payload(s,eid='test',g=None):
    g=groups(s,eid) if g is None else g;v=partition(s,g);b=surface(s,eid,included=lambda f:False)
    def included(f):
        return bool(f.get('cap_component_rings')) and Decimal(f['v_from'])>=3 and max(Decimal(p['x']) for p in f['cap_component_rings'][0]['points'])<=2
    b.update(name='合成U形左臂独立防水（非真实工程）',purpose='waterproof',face_subdivisions=g,
             measurement_basis='合成按Y=3m裁剪，上部两独立臂各外6扣孔1=5m²，仅左臂计入，重复2=10m²；其他面明确排除',
             face_decisions=[dict(face_id=f['face_id'],included=included(f),basis='合成按独立块逐项核定做法，非专业签审',evidence_id=eid) for f in v['faces']])
    return b


def setup(clients,pid=None):
    e,r=clients['estimator'],clients['reviewer'];pid=pid or create(e)['id']
    upload(e,pid,'合成原始依据，仅软件验收：共同水平设计XY。U形外(0,0)(6,0)(6,6)(4,6)(4,2)(2,2)(2,6)(0,6)m面积28m²；左孔(.5,4)(1.5,4)(1.5,5)(.5,5)m面积1m²，右孔(4.5,4)(5.5,4)(5.5,5)(4.5,5)m面积1m²，净26m²。厚200mm、重复2。顶部Y=300cm分区，上区左右分离净块各外6−孔1=5m²，仅左臂计入10m²，其他面排除。停用分块后上区两臂合计10m²×2=20m²须另行明确核定。','connected-cap-basis.txt')
    eid=get(e,pid)['evidence'][-1]['id'];s=geometry(eid);s.pop('id');s.pop('revision');s['name']='合成U形双孔等厚实体（非真实工程）'
    comp=post(e,pid,'/components',s).json()
    for c in get(e,pid)['components']:approve(r,pid,'components',c)
    s=next(c for c in get(e,pid)['components'] if c['id']==comp['id'])
    return e,r,pid,s,payload(s,eid),eid


@pytest.mark.parametrize('top',[False,True])
def test_u_shape_arms_separate_with_each_hole_retained_and_exact_ast(top):
    s=geometry();g=groups(s,top=top);before=copy.deepcopy((s,g));v=partition(s,g);fs=pieces(v)
    assert (s,g)==before and [frac(f) for f in fs]==[16,5,5] and len(v['faces'])==20
    assert [[r['role'] for r in f['cap_component_rings']] for f in fs]==[['outer'],['outer','hole'],['outer','hole']]
    assert [f['cap_component_count'] for f in fs]==[1,2,2]
    for f in fs:
        p=exact_point(f['cap_label_point']);rings=[[exact_point(pt) for pt in r['points']] for r in f['cap_component_rings']]
        assert inside(p,rings[0])==1 and all(inside(p,h)==-1 for h in rings[1:])
    assert [c['component_count'] for c in v['surface_cap_component_checks']]==[1,2]
    assert all(c['exact_rational_tiling'] and Decimal(c['area_difference'])==0 for c in v['surface_cap_component_checks'])
    assert len({f['face_id'] for f in v['faces']})==len(v['faces'])
    with localcontext() as ctx:
        ctx.prec=80
        for f in fs:assert evaluate(f['formula_ast'],v['dimensions'])==Decimal(f['area'])
    b=payload(s,g=g);out=calculate(b,s);assert out['rule_version']=='polygon-boundary-1.4' and Decimal(out['quantity'])==10
    assert sum(Decimal(c['area']) for c in out['surface_cap_clip_sources'])==26
    if not top:assert all(f['normal']=='z-' for f in fs)


def test_opened_holes_become_notches_while_original_deduction_proof_is_preserved():
    s=geometry();s['polygon_rings']=[ring('外方形',[(0,0),(6,0),(6,6),(0,6)]),ring('长孔',[(2,1),(4,1),(4,5),(2,5)],role='hole')]
    v=partition(s,groups(s));assert [frac(f) for f in pieces(v)]==[14,14]
    assert all(len(f['cap_component_rings'])==1 and len(f['cap_component_rings'][0]['points'])==8 for f in pieces(v))
    assert all(Decimal(c['cap_clip_rings'][1]['area'])==4 for c in v['surface_cap_clip_sources'])


def test_hole_cut_across_both_sides_creates_two_independent_strip_pieces():
    s=geometry();s['polygon_rings']=[ring('方形',[(0,0),(6,0),(6,6),(0,6)]),ring('长孔',[(2,1),(4,1),(4,5),(2,5)],role='hole')]
    g=groups(s,y='2');g[0]['v_cuts'].append(dict(value='4',unit='m',evidence_id='test'));v=partition(s,g)
    assert sorted(frac(f) for f in pieces(v))==[4,4,10,10]
    middle=[c for c in v['surface_cap_component_checks'] if c['clip_index']==2][0];assert middle['component_count']==2 and Decimal(middle['clip_area'])==8


def test_slope_cut_split_arms_and_nonterminating_coordinates():
    s=geometry(holes=False);g=line_groups(s,lines=[line(a=(0,3),b=(10,4))]);g[0]['separate_regions']=True
    v=partition(s,g);assert sorted(frac(f) for f in pieces(v))==[Fraction(5),Fraction('5.8'),Fraction('17.2')]
    s['polygon_rings']=[ring('square',[(0,0),(2,0),(2,2),(0,2)])]
    g=line_groups(s,lines=[line(a=(0,0),b=(3,2))]);g[0]['separate_regions']=True;v=partition(s,g)
    assert sorted(frac(f) for f in pieces(v))==[Fraction(4,3),Fraction(8,3)]
    assert any(p['exact_y']['denominator']=='3' for f in pieces(v) for r in f['cap_component_rings'] for p in r['points'])


@pytest.mark.parametrize('change',['reverse_rings','rotate_start','collinear','swap_axes','negative_xy','units'])
def test_boundary_orientation_start_and_coordinate_transforms(change):
    s=geometry();g=groups(s)
    for r in s['polygon_rings']:
        if change=='reverse_rings':r['points'].reverse()
        elif change=='rotate_start':r['points']=r['points'][1:]+r['points'][:1]
        elif change=='collinear':
            a,b=r['points'][:2];p=copy.deepcopy(a);p['name']='额外共线点';p['parameters']={k:str((Decimal(a['parameters'][k])+Decimal(b['parameters'][k]))/2) for k in ('x','y')};r['points'].insert(1,p)
        for p in r['points']:
            x,y=(Decimal(p['parameters'][k]) for k in ('x','y'))
            if change=='swap_axes':x,y=y,x
            elif change=='negative_xy':x-=100;y-=200
            elif change=='units':x*=1000;y*=1000;p['unit']='mm'
            p['parameters']={'x':str(x),'y':str(y)}
    g=groups(s,y='-197' if change=='negative_xy' else '3')
    if change=='swap_axes':g[0]['u_cuts']=g[0]['v_cuts'];g[0]['v_cuts']=[]
    assert sorted(frac(f) for f in pieces(partition(s,g)))==[5,5,16]


def islands(cells):
    cells=set(cells);result=[]
    while cells:
        start=min(cells);cells.remove(start);q=deque([start]);island={start}
        while q:
            x,y=q.popleft()
            for p in [(x-1,y),(x+1,y),(x,y-1),(x,y+1)]:
                if p in cells:cells.remove(p);island.add(p);q.append(p)
        result.append(island)
    return result


def test_thirty_random_staircases_and_holes_match_independent_grid_flood_fill():
    rng=random.Random(6301)
    for _ in range(30):
        w=rng.randint(4,8);hs=[rng.randint(2,7) for _ in range(w)];coords=[(0,0),(w,0),(w,hs[-1])]
        for x in range(w-1,0,-1):
            if hs[x]!=hs[x-1]:coords.extend([(x,hs[x]),(x,hs[x-1])])
        coords.append((0,hs[0]));s=geometry(holes=False);s['polygon_rings']=[ring('阶梯',coords)]
        occupied={(x,y) for x in range(w) for y in range(hs[x])};hole_cells=rng.sample(sorted(occupied),3)
        for i,(x,y) in enumerate(hole_cells):s['polygon_rings'].append(ring('孔'+str(i),[(x+.25,y+.25),(x+.75,y+.25),(x+.75,y+.75),(x+.25,y+.75)],role='hole'))
        xcut=rng.randint(1,w-1);ycut=rng.randint(1,max(hs)-1);g=cap_groups(s,x=str(xcut),y=str(ycut));g[0]['separate_regions']=True;v=partition(s,g)
        expected=[]
        for x0,x1 in [(0,xcut),(xcut,w)]:
            for y0,y1 in [(0,ycut),(ycut,max(hs))]:
                for cellset in islands({p for p in occupied if x0<=p[0]<x1 and y0<=p[1]<y1}):
                    expected.append(Fraction(len(cellset))-Fraction(sum(p in cellset for p in hole_cells),4))
        assert sorted(frac(f) for f in pieces(v))==sorted(expected)
        for f in pieces(v):
            p=exact_point(f['cap_label_point']);rr=[[exact_point(pt) for pt in r['points']] for r in f['cap_component_rings']]
            assert inside(p,rr[0])==1 and all(inside(p,h)==-1 for h in rr[1:])


def test_equivalent_cut_units_and_order_and_unchanged_piece_identity():
    s=geometry();g=groups(s);v=partition(s,g);other=copy.deepcopy(g);other[0]['v_cuts'][0].update(value='3000',unit='mm')
    assert [f['face_id'] for f in pieces(v)]==[f['face_id'] for f in pieces(partition(s,other))]
    # A new vertical line through the empty gap leaves both arm boundaries intact.
    other[0]['u_cuts']=[dict(value='3',unit='m',evidence_id='test')]
    old_arms={f['face_id'] for f in pieces(v) if f['cap_component_count']==2}
    now={f['face_id'] for f in pieces(partition(s,other))};assert old_arms<=now


def test_mode_is_optional_and_all_legacy_outputs_remain_unchanged():
    s=geometry();g=groups(s);g[0].pop('separate_regions')
    from backend.polygon_line_partitions import partition as legacy
    assert partition(s,g)==legacy(s,g,copy.deepcopy(boundary(s)))
    assert calculate(surface(s),s)['rule_version']=='polygon-boundary-1.0'
    g=groups(s)+groups(s,top=False);v=partition(s,g)
    assert len(pieces(v))==6 and sum(frac(f) for f in pieces(v))==52


def test_cut_along_notch_zero_width_bridge_is_cancelled():
    s=geometry(holes=False);v=partition(s,groups(s,y='2'));assert sorted(frac(f) for f in pieces(v))==[8,8,12]


def test_point_touching_hole_branch_refuses_independent_assignment_but_legacy_still_works():
    s=geometry();s['polygon_rings']=[ring('外方形',[(0,0),(6,0),(6,6),(0,6)]),ring('菱形孔',[(3,1),(4,2),(3,3),(2,2)],role='hole')]
    g=groups(s,y='1');old=copy.deepcopy(g);old[0].pop('separate_regions');assert len(partition(s,old)['faces'])==11
    with pytest.raises(DomainError,match='点接触|分支'):partition(s,g)


@pytest.mark.parametrize('value',[False,'true',1,None])
def test_separation_flag_cannot_be_coerced(value):
    s=geometry();g=groups(s);g[0]['separate_regions']=value
    with pytest.raises(DomainError):partition(s,g)


def test_side_flag_missing_cuts_and_graph_budget_are_blocked():
    s=geometry();g=groups(s);g[0]['face_id']=boundary(s)['faces'][2]['face_id']
    with pytest.raises(DomainError):partition(s,g)
    g=groups(s);g[0]['v_cuts']=[]
    with pytest.raises(DomainError):partition(s,g)
    with pytest.raises(DomainError,match='2000000'):
        separate([[(Fraction(x),Fraction(y)) for x,y in U]],[(Fraction(0),Fraction(0)),(Fraction(6),Fraction(0)),(Fraction(6),Fraction(6)),(Fraction(0),Fraction(6))],[MAX_EDGE_TESTS])


def test_candidate_preview_independent_review_takeoffs_and_boq(clients):
    e,r,pid,s,b,eid=setup(clients);before=get(e,pid)
    c=post(e,pid,'/surface-face-candidates/partition',b).json();assert len(pieces(c))==3
    preview=post(e,pid,'/components/preview',b).json();assert Decimal(preview['quantity'])==10 and get(e,pid)==before
    comp=post(e,pid,'/components',b).json();post(e,pid,'/takeoffs/calculate',expected=409);approve(r,pid,'components',comp)
    for t in post(e,pid,'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
    assert any(v['unit']=='m²' and Decimal(v['quantity'])==10 for v in post(e,pid,'/boq/generate').json())


@pytest.mark.parametrize('fault',['mode','boundary','rational_point','clip_source','closure','piece_area','evidence','bytes','source'])
def test_derived_boundaries_and_original_clip_proofs_cannot_be_forged(clients,fault):
    e,r,pid,s,b,eid=setup(clients);comp=post(e,pid,'/components',b).json()
    if fault=='bytes':(store.ROOT/'originals'/get(e,pid)['documents'][-1]['hash']).write_bytes(b'bad')
    else:
        def damage(p):
            c=p['components'][-1];f=next(f for f in c['preview']['surface_faces'] if f.get('cap_component_rings'))
            if fault=='mode':c['face_subdivisions'][0].pop('separate_regions')
            elif fault=='boundary':f['cap_component_rings'][0]['points'][0]['x']='999'
            elif fault=='rational_point':f['cap_component_rings'][0]['points'][0]['exact_y']['numerator']='999'
            elif fault=='clip_source':c['preview']['surface_cap_clip_sources'][0]['cap_clip_rings'][1]['area']='999'
            elif fault=='closure':c['preview']['surface_cap_component_checks'][0]['component_count']=999
            elif fault=='piece_area':f['cap_rational_area']['numerator']='999'
            elif fault=='evidence':p['evidence'][-1]['text']+=' changed'
            else:p['components'][0]['polygon_rings'][0]['points'][1]['parameters']['x']='7'
        mutate(pid,damage)
    current=get(e,pid)['components'][-1];res=r.post(f'/api/v1/projects/{pid}/reviews',json=dict(collection='components',object_id=comp['id'],object_hash=store.digest(current),decision='approve',reason='合成核查'))
    assert res.status_code in (409,422),res.text


def test_saved_book_connected_geometry_proof_cannot_be_forged(clients):
    e,r,pid,s,b,eid=setup(clients);comp=post(e,pid,'/components',b).json();approve(r,pid,'components',comp);books=post(e,pid,'/takeoffs/calculate').json()
    book=next(t for t in books if t['component_id']==comp['id']);mutate(pid,lambda p:next(t for t in p['takeoffs'] if t['id']==book['id'])['surface_cap_component_checks'][0].update(numerator='999'))
    approve(r,pid,'takeoffs',book,expected=409)


def test_disabling_separation_requires_new_decisions_preserves_v1_and_rebuilds_boq(clients):
    e,r,pid,s,b,eid=setup(clients);comp=post(e,pid,'/components',b).json();approve(r,pid,'components',comp)
    for t in post(e,pid,'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
    post(e,pid,'/boq/generate');old=get(e,pid)['components'][-1];g=groups(s,eid);g[0].pop('separate_regions');v=partition(s,g)
    new=payload(s,eid,g);new.update(object_hash=store.digest(old),measurement_basis='合成停用分离，两臂做法共同计入，净10×2=20m²')
    new['face_decisions']=[dict(face_id=f['face_id'],included=bool(f.get('parent_face_id')) and Decimal(f['v_from'])>=3,basis='合成停用分离后明确核定合并做法',evidence_id=eid) for f in v['faces']]
    url=f'/api/v1/projects/{pid}/components/'+comp['id'];bad={**new,'face_decisions':old['face_decisions']};assert e.put(url,json=bad).status_code==422
    res=e.put(url,json=new);assert res.status_code==200,res.text;assert Decimal(res.json()['preview']['quantity'])==20 and res.json()['preview']['rule_version']=='polygon-boundary-1.2'
    p=get(e,pid);assert p['component_history'][-1]==old and all(t['state']=='stale' for t in p['takeoffs'])
    assert any(v['unit']=='m²' and Decimal(v['quantity'])==10 for v in p['boq']) and e.put(url,json=new).status_code==409
    approve(r,pid,'components',res.json())
    for t in post(e,pid,'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
    assert any(v['unit']=='m²' and Decimal(v['quantity'])==20 for v in post(e,pid,'/boq/generate').json())


def test_source_revision_scope_permissions_and_cross_project_proofs(clients):
    e,r,pid,s,b,eid=setup(clients)
    post(r,pid,'/surface-face-candidates/partition',b,403);post(clients['other'],pid,'/surface-face-candidates/partition',b,404);post(TestClient(app),pid,'/surface-face-candidates/partition',b,401)
    second=create(e)['id'];post(e,second,'/surface-face-candidates/partition',b,404)
    upload(e,second,'合成另一项目的证据');foreign=get(e,second)['evidence'][-1]['id'];bad=copy.deepcopy(b);bad['face_subdivisions'][0]['v_cuts'][0]['evidence_id']=foreign;post(e,pid,'/surface-face-candidates/partition',bad,404)
    old=get(e,pid)['components'][0];res=e.put(f'/api/v1/projects/{pid}/components/'+s['id'],json={**copy.deepcopy(old),'object_hash':store.digest(old),'thickness':'300'});assert res.status_code==200,res.text
    approve(r,pid,'components',res.json());post(e,pid,'/surface-face-candidates/partition',b,409)
    from tests.test_project_access import restrict
    restrict(clients['admin'],pid,[]);post(e,pid,'/surface-face-candidates/partition',b,404)


@pytest.mark.parametrize('fmt',['json','xlsx','pdf'])
def test_component_tables_in_three_exports_and_old_quote_snapshot_isolation(clients,fmt):
    e,r,pid,_,quote=priced_project(clients);e,r,pid,s,b,eid=setup(clients,pid);comp=post(e,pid,'/components',b).json();approve(r,pid,'components',comp)
    for t in post(e,pid,'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
    url=f'/api/v1/projects/{pid}/exports';res=e.get(url,params={'format':fmt,'unpriced':True});assert res.status_code==200,res.text
    titles=['净块分离与闭合校验','独立净块派生边界','分离前原区逐孔扣减']
    if fmt=='json':assert all(any(t[0]==title for t in res.json()['tables']) for title in titles) and any(t.get('surface_cap_component_checks') for t in res.json()['unpriced']['takeoffs'])
    elif fmt=='xlsx':assert set(titles)<=set(load_workbook(io.BytesIO(res.content)).sheetnames)
    else:assert all(title in ''.join(p.extract_text() for p in PdfReader(io.BytesIO(res.content)).pages) for title in titles)
    old=e.get(url,params={'format':'json','quote_id':quote['id']}).json();assert old['quote']==quote and not any(t[0] in titles for t in old['tables'])


def test_tiny_separated_arms_at_large_origin_preserve_exact_positive_areas():
    s=geometry(holes=False)
    with localcontext() as ctx:
        ctx.prec=80;o=Decimal('1000000000');scale=Decimal('.00000001')
        for p in s['polygon_rings'][0]['points']:p['parameters']={k:str(o+Decimal(v)*scale) for k,v in p['parameters'].items()}
        g=groups(s,y=str(o+3*scale))
    with localcontext() as ctx:ctx.prec=10;v=partition(s,g)
    assert sorted(frac(f) for f in pieces(v))==[Fraction(6,10**16),Fraction(6,10**16),Fraction(16,10**16)]


def test_connected_blocks_push_exactly_one_thousand_legacy_faces_over_budget():
    s=geometry();base=boundary(s);g=groups(s);side=base['faces'][2:]
    def cuts(f,nu,nv):
        return dict(face_id=f['face_id'],u_cuts=[dict(value=str(Decimal(f['edge_length'])*i/10),unit='m',evidence_id='test') for i in range(1,nu+1)],
                    v_cuts=[dict(value=str(Decimal(f['height'])*i/10),unit='m',evidence_id='test') for i in range(1,nv+1)])
    g.extend(cuts(f,8,8) for f in side[:12]);g.append(cuts(side[12],6,2));g.append(cuts(side[13],1,0))
    old=copy.deepcopy(g);old[0].pop('separate_regions');assert len(partition(s,old)['faces'])==1000
    with pytest.raises(DomainError,match='1000'):partition(s,g)


def test_axis_and_oblique_mixed_with_side_partition_and_partial_mode_selection():
    s=geometry(holes=False);g=groups(s)+cap_groups(s,top=False,x='3',y='3')
    side=boundary(s)['faces'][2];g.append(dict(face_id=side['face_id'],u_cuts=[dict(value='1',unit='m',evidence_id='test')],v_cuts=[dict(value='.1',unit='m',evidence_id='test')]))
    v=partition(s,g);assert v['surface_partition_checks'] and v['surface_cap_component_checks']
    assert all(f['normal']=='z+' for f in pieces(v))
    assert any(f.get('cap_clip_rule') and not f.get('cap_component_rings') and f['normal']=='z-' for f in v['faces'])
    assert Decimal(v['normal_totals']['z+'])==28 and Decimal(v['normal_totals']['z-'])==28


def test_cut_revision_recomputes_derived_boundaries_and_requires_new_review(clients):
    e,r,pid,s,b,eid=setup(clients);comp=post(e,pid,'/components',b).json();approve(r,pid,'components',comp)
    old=get(e,pid)['components'][-1];new=payload(s,eid,groups(s,eid,y='3.5'));new['object_hash']=store.digest(old)
    assert Decimal(post(e,pid,'/components/preview',new).json()['quantity'])==8
    res=e.put(f'/api/v1/projects/{pid}/components/'+comp['id'],json=new);assert res.status_code==200,res.text
    assert res.json()['state']=='draft' and Decimal(res.json()['preview']['quantity'])==8 and get(e,pid)['component_history'][-1]==old
    assert not {f['face_id'] for f in old['preview']['surface_faces'] if f.get('cap_component_rings')} & {f['face_id'] for f in res.json()['preview']['surface_faces'] if f.get('cap_component_rings')}
    approve(r,pid,'components',res.json())
    for t in post(e,pid,'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
    assert any(v['unit']=='m²' and Decimal(v['quantity'])==8 for v in post(e,pid,'/boq/generate').json())
