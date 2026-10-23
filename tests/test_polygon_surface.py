"""Synthetic extruded-polygon boundary math and audit gates, no professional signoff."""
import copy
import io
import json
import random
from decimal import Decimal, localcontext
import pytest
from fastapi.testclient import TestClient
from backend.main import app
from openpyxl import load_workbook
from pypdf import PdfReader
from backend import store
from backend.engine import DomainError
from backend.surface_takeoff import boundary, calculate, snapshot, partition
from tests.test_polygon_takeoff import payload as polygon, ring
from tests.test_route_takeoff import evaluate
from tests.test_surface_takeoff import payload as surface
from tests.test_workflow import get, approve, create, priced_project, upload


def geometry(eid='test', holes=True):
    b=polygon(eid);b.update(id='synthetic-prism',revision=1,thickness='2',thickness_unit='m',multiplicity='3')
    # 6x6 outer square, internal 3-4-5 triangle; cap30 and outer/inner sides48+24 per group.
    b['measurement_basis']='合成方形外轮廓36m²；三四五贯通孔6m²，净30m²；等厚2m、重复3；边界按逐面决定，非专业依据'
    b['polygon_rings']=[ring('方形外轮廓',[(0,0),(6,0),(6,6),(0,6)],eid)]
    if holes:b['polygon_rings'].append(ring('三四五三角孔',[(1,1),(4,1),(1,5)],eid,'hole'))
    return b


def setup(clients,pid=None):
    e,r=clients['estimator'],clients['reviewer'];pid=pid or create(e)['id']
    upload(e,pid,'合成软件验收：同一水平设计平面、共同XY原点方向。6×6m方形外轮廓，内孔(1,1)(4,1)(1,5)m，孔边3、5、4m。外36−孔6＝净30m²；统一竖向厚2m，重复3。单组上下各30m²、外侧48m²、孔内侧24m²。仅明确竖向全部计入模板、水平上下排除，单组72、总216m²。局部底面0不是绝对标高，不推断相邻实体接触。','polygon-surface-basis.txt')
    eid=get(e,pid)['evidence'][-1]['id'];b=geometry(eid)
    b.pop('id');b.pop('revision');b['name']='合成多边形等厚实体（非真实工程）'
    res=e.post(f'/api/v1/projects/{pid}/components',json=b);assert res.status_code==200,res.text
    for c in get(e,pid)['components']:approve(r,pid,'components',c)
    source=next(c for c in get(e,pid)['components'] if c['id']==res.json()['id'])
    return e,r,pid,source,surface(source,eid),eid


def post(e,pid,path,body=None,expected=200):
    res=e.post(f'/api/v1/projects/{pid}'+path,json=body or {});assert res.status_code==expected,res.text;return res


def mutate(pid,action):
    with store.transaction() as c:
        p=json.loads(c.execute('select data from projects where id=?',(pid,)).fetchone()['data']);action(p);store.save(c,p)


@pytest.mark.parametrize('purpose',['formwork','waterproof'])
def test_exact_caps_hole_sides_edge_identity_and_aggregate_ast(purpose):
    s=geometry();before=copy.deepcopy(s);b=boundary(s);p=surface(s);p['purpose']=purpose;r=calculate(p,s)
    assert s==before and len(b['faces'])==9 and len({f['face_id'] for f in b['faces']})==9
    assert {k:Decimal(v) for k,v in b['normal_totals'].items()}=={'z-':30,'z+':30,'outer':48,'hole':24}
    assert Decimal(b['total_area'])==132 and Decimal(r['quantity'])==216 and Decimal(r['deduction'])==180
    assert r['rule_version']=='polygon-boundary-1.0' and r['source_geometry']==snapshot(s)
    assert [Decimal(f['edge_length']) for f in b['faces'] if f['normal']=='hole']==[3,5,4]
    assert b['faces'][1]['local_z']=='2' and len(b['faces'][0]['boundary_rings'])==2
    with localcontext() as ctx:
        ctx.prec=80
        for f in b['faces']:assert evaluate(f['formula_ast'],b['dimensions'])==Decimal(f['area'])
        assert evaluate(r['formula_ast'],r['dimensions'])==216
    assert '局部' in r['notice'] and '未修约' in r['rounding_rule']


@pytest.mark.parametrize('select,expected',[(lambda f:True,396),(lambda f:False,0),(lambda f:f['normal']=='z+',90),(lambda f:f['normal']=='hole',72)])
def test_all_excluded_top_only_and_hole_only_are_explicit(select,expected):
    s=geometry();r=calculate(surface(s,included=select),s);assert Decimal(r['quantity'])==expected
    with localcontext() as ctx:ctx.prec=80;assert evaluate(r['formula_ast'],r['dimensions'])==expected


@pytest.mark.parametrize('mode',['reverse','translate','rotate','mixed_units','collinear','no_hole'])
def test_independent_perimeter_oracle_under_design_transformations(mode):
    s=geometry(holes=mode!='no_hole')
    for row in s['polygon_rings']:
        if mode=='reverse':row['points'].reverse()
        for pt in row['points']:
            x,y=[Decimal(pt['parameters'][k]) for k in ('x','y')]
            if mode=='translate':pt['parameters']={'x':str(x-99999),'y':str(y+99999)}
            if mode=='rotate':pt['parameters']={'x':str(-y),'y':str(x)}
            if mode=='mixed_units':pt.update(parameters={'x':str(x*1000),'y':str(y*1000)},unit='mm')
    if mode=='collinear':s['polygon_rings'][0]=ring('带共线点',[(0,0),(3,0),(6,0),(6,6),(0,6)])
    b=boundary(s);expected=120 if mode=='no_hole' else 132
    assert Decimal(b['total_area'])==expected
    r=calculate(surface(s,included=lambda f:True),s);assert Decimal(r['quantity'])==expected*3


def test_concave_slanted_and_irrational_edges_use_unrounded_geometry_not_volume():
    s=polygon();s.update(id='s',revision=1,rounding_digits=0,thickness='.123',thickness_unit='m')
    b=boundary(s);assert len(b['faces'])==15
    with localcontext() as ctx:
        ctx.prec=80
        perimeter=Decimal(22)+2+Decimal(2).sqrt()+4*Decimal('.5').sqrt()
        expected=50+perimeter*Decimal('.123')
        assert abs(Decimal(b['total_area'])-expected)<Decimal('1e-75')
        assert [Decimal(f['area']) for f in b['faces'][:2]]==[25,25]
        for f in b['faces']:assert evaluate(f['formula_ast'],b['dimensions'])==Decimal(f['area'])
    s['rounding_digits']=12;assert boundary(s)==b  # reviewed volume display precision never changes geometry


def test_large_origin_tiny_geometry_and_source_count_applied_once():
    s=geometry();s.update(thickness='1',multiplicity='7')
    with localcontext() as ctx:
        ctx.prec=80;o=Decimal('1000000000');d=Decimal('.00000001')
        s['polygon_rings']=[ring('微小方形',[(o,o),(o+d,o),(o+d,o+d),(o,o+d)])]
        expected=2*d*d+4*d
        b=boundary(s);assert Decimal(b['total_area'])==expected
        r=calculate(surface(s,included=lambda f:True),s);assert Decimal(r['quantity'])==expected*7


def test_area_mode_and_nonempty_polygon_subdivisions_rejected():
    s=geometry();s['quantity_type']='area'
    with pytest.raises(DomainError):boundary(s)
    s=geometry();assert partition(s,[])==boundary(s)
    with pytest.raises(DomainError,match='局部分区'):partition(s,[{'face_id':boundary(s)['faces'][0]['face_id'],'u_cuts':[],'v_cuts':[]}])


def test_candidate_preview_review_calculate_boq_and_client_snapshot_ignored(clients):
    e,r,pid,s,b,eid=setup(clients);before=get(e,pid);url=f'/api/v1/projects/{pid}/surface-face-candidates'
    candidate=e.get(url,params={'component_id':s['id']});assert candidate.status_code==200,candidate.text
    assert candidate.json()==r.get(url,params={'component_id':s['id']}).json()
    assert Decimal(candidate.json()['total_area'])==132
    view=post(e,pid,'/components/preview',b).json();assert get(e,pid)==before and Decimal(view['quantity'])==216
    comp=post(e,pid,'/components',{**b,'source_geometry':{'fake':'ignored'}}).json();assert comp['source_geometry']==snapshot(s)
    assert get(e,pid)['components'][-2]['state']=='approved';post(e,pid,'/takeoffs/calculate',expected=409)
    approve(r,pid,'components',comp)
    for t in post(e,pid,'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
    items=post(e,pid,'/boq/generate').json();assert next(v for v in items if v['unit']=='m²')['quantity']=='216'


@pytest.mark.parametrize('fault',['bytes','missing','superseded','proof','coordinate','preview','mapping','stamp'])
def test_source_integrity_checked_before_candidate_preview_and_save(clients,fault):
    e,r,pid,s,b,eid=setup(clients)
    if fault in ('bytes','missing'):
        path=store.ROOT/'originals'/get(e,pid)['documents'][0]['hash']
        if fault=='bytes':path.write_bytes(b'corrupt')
        else:path.unlink()
    else:
        def damage(p):
            c=p['components'][-1]
            if fault=='superseded':p['documents'][0]['superseded']=True
            elif fault=='proof':p['evidence'][0]['text']+=' changed'
            elif fault=='coordinate':c['polygon_rings'][0]['points'][1]['parameters']['x']='7'
            elif fault=='preview':c['preview']['quantity']='999'
            elif fault=='mapping':c['preview']['polygon_rings'][0]['edges'][0]['cross_product']='999'
            else:c['source_stamp']='0'*64
        mutate(pid,damage)
    before=get(e,pid)
    assert e.get(f'/api/v1/projects/{pid}/surface-face-candidates',params={'component_id':s['id']}).status_code==409
    post(e,pid,'/components/preview',b,409);post(e,pid,'/components',b,409);assert get(e,pid)==before


@pytest.mark.parametrize('gate',['review','calculate','book_review','snapshot'])
def test_downstream_rechecks_and_saved_source_snapshots(clients,gate):
    e,r,pid,s,b,eid=setup(clients);comp=post(e,pid,'/components',b).json()
    if gate in ('calculate','book_review'):approve(r,pid,'components',comp)
    if gate=='book_review':book=post(e,pid,'/takeoffs/calculate').json()[-1]
    def damage(p):
        if gate=='snapshot':p['components'][-1]['source_geometry']['polygon_rings'][0]['points'][0]['parameters']['x']='9'
        elif gate=='book_review':p['takeoffs'][-1]['surface_faces'][2]['edge_length']='999'
        else:p['components'][-2]['polygon_rings'][0]['points'][1]['parameters']['x']='7'
    mutate(pid,damage)
    if gate=='calculate':post(e,pid,'/takeoffs/calculate',expected=409)
    elif gate=='book_review':approve(r,pid,'takeoffs',book,expected=409)
    else:approve(r,pid,'components',comp,expected=409)


def test_source_revision_resets_hashes_and_preserves_exact_history_and_old_boq(clients):
    e,r,pid,s,b,eid=setup(clients);comp=post(e,pid,'/components',b).json();approve(r,pid,'components',comp)
    for t in post(e,pid,'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
    post(e,pid,'/boq/generate');base=f'/api/v1/projects/{pid}'
    old=next(c for c in get(e,pid)['components'] if c['id']==s['id'])
    revised=e.put(base+'/components/'+s['id'],json={**copy.deepcopy(old),'object_hash':store.digest(old),'thickness':'3'});assert revised.status_code==200,revised.text
    p=get(e,pid);current=p['components'][-1];assert current['state']=='stale' and all(t['state']=='stale' for t in p['takeoffs'])
    assert next(v for v in p['boq'] if v['unit']=='m²')['quantity']=='216'
    approve(r,pid,'components',revised.json());post(e,pid,'/components/preview',b,409)
    fresh=next(c for c in get(e,pid)['components'] if c['id']==s['id']);new=surface(fresh,eid);new['object_hash']=store.digest(current)
    res=e.put(base+'/components/'+comp['id'],json=new);assert res.status_code==200,res.text
    assert res.json()['revision']==2 and Decimal(res.json()['preview']['quantity'])==324
    assert get(e,pid)['component_history'][-1]==current
    approve(r,pid,'components',res.json())
    for t in post(e,pid,'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
    assert next(v for v in post(e,pid,'/boq/generate').json() if v['unit']=='m²')['quantity']=='324'


def test_role_tenant_cross_project_area_and_partition_gate(clients):
    e,r,pid,s,b,eid=setup(clients);base=f'/api/v1/projects/{pid}'
    post(r,pid,'/components',b,403);post(clients['other'],pid,'/components',b,404);post(TestClient(app),pid,'/components',b,401)
    second=create(e)['id'];assert e.get(f'/api/v1/projects/{second}/surface-face-candidates',params={'component_id':s['id']}).status_code==404
    post(e,pid,'/surface-face-candidates/partition',{**b,'face_subdivisions':[{'face_id':'x','u_cuts':[],'v_cuts':[]}]},422)
    area=polygon(eid,'area');a=post(e,pid,'/components',area).json();approve(r,pid,'components',a)
    assert e.get(base+'/surface-face-candidates',params={'component_id':a['id']}).status_code==409
    from tests.test_project_access import restrict
    restrict(clients['admin'],pid,[]);assert e.get(base+'/surface-face-candidates',params={'component_id':s['id']}).status_code==404


@pytest.mark.parametrize('fmt',['json','xlsx','pdf'])
def test_three_exports_and_old_quote_isolation(clients,fmt):
    e,r,pid,_,quote=priced_project(clients);e,r,pid,s,b,eid=setup(clients,pid)
    comp=post(e,pid,'/components',b).json();approve(r,pid,'components',comp)
    for t in post(e,pid,'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
    base=f'/api/v1/projects/{pid}/exports';res=e.get(base,params={'format':fmt,'unpriced':True});assert res.status_code==200,res.text
    titles=['逐面计量决定','多边形表面原始坐标','多边形表面来源口径']
    if fmt=='json':
        data=res.json();assert all(any(t[0]==n for t in data['tables']) for n in titles)
        item=next(t for t in data['unpriced']['takeoffs'] if t.get('surface_summary'));assert item['source_geometry']==snapshot(s) and Decimal(item['quantity'])==216
    elif fmt=='xlsx':assert set(titles)<=set(load_workbook(io.BytesIO(res.content)).sheetnames)
    else:assert all(n in ''.join(p.extract_text() for p in PdfReader(io.BytesIO(res.content)).pages) for n in titles)
    old=e.get(base,params={'format':'json','quote_id':quote['id']}).json();assert old['quote']==quote and not any(t[0] in titles for t in old['tables'])


def test_random_concave_staircases_match_independent_area_perimeter_oracle():
    rng=random.Random(5901)
    for _ in range(25):
        w=rng.randint(2,10);heights=[rng.randint(1,9) for _ in range(w)]
        coords=[(0,0),(w,0),(w,heights[-1])]
        for x in range(w-1,0,-1):
            if heights[x]!=heights[x-1]:coords.extend([(x,heights[x]),(x,heights[x-1])])
        coords.append((0,heights[0]));s=geometry();s.update(thickness='.125',multiplicity='2')
        s['polygon_rings']=[ring('阶梯凹形',coords)]
        perimeter=2*w+heights[0]+heights[-1]+sum(abs(a-b) for a,b in zip(heights,heights[1:]))
        expected=2*Decimal(sum(heights))+Decimal(perimeter)*Decimal('.125')
        result=calculate(surface(s,included=lambda f:True),s)
        assert Decimal(result['surface_summary']['total_area'])==expected and Decimal(result['quantity'])==expected*2


@pytest.mark.parametrize('fault',['missing','duplicate','unknown','non_bool','extra','blank_basis','blank_evidence'])
def test_polygon_face_decisions_cannot_omit_or_override_geometry(fault):
    s=geometry();b=surface(s);d=b['face_decisions']
    if fault=='missing':d.pop()
    elif fault=='duplicate':d[1]=d[0]
    elif fault=='unknown':d[0]['face_id']='forged'
    elif fault=='non_bool':d[0]['included']=1
    elif fault=='extra':d[0]['area']='999'
    elif fault=='blank_basis':d[0]['basis']=' '
    else:d[0]['evidence_id']=' '
    with pytest.raises(DomainError):calculate(b,s)
