"""Synthetic contact surfaces; no inference that soil faces require formwork."""
import copy,io,random
from decimal import Decimal,localcontext,ROUND_HALF_UP
import pytest
from fastapi import HTTPException
from openpyxl import load_workbook
from pypdf import PdfReader
from backend import store
from backend.compound_takeoff import PI
from backend.engine import DomainError
from backend.foundation_formwork import calculate,layers
from backend.post_foundation import calculate as foundation_quantity
from backend.railing_material import snapshot
from tests.test_foundation_deduction_rules import setup as threshold_setup,body as threshold_body
from tests.test_foundation_voids import body as geometry_body,rectangle,SHAPES
from tests.test_foundation_recesses import interval
from tests.test_post_foundation import source,equal_body,fixed_body,fixed_added,equal_added,fixed_adjusted,equal_adjusted,combined
from tests.test_route_takeoff import evaluate
from tests.test_polygon_takeoff import post,mutate
from tests.test_workflow import get,approve,upload,create,priced_project

TITLES={'基础模板接触面积汇总','基础模板逐面标高与净量','基础模板侧开口实际扣减','基础模板原始范围与依据','基础模板做法与计算规则'}


def foundation(make=equal_body,shape='rectangular',eid='test',edge=True):
 s=source(make);b=threshold_body(s,eid,shape);d=b['foundation_items'][0]
 if edge and shape=='rectangular':
  d['through_holes']['holes']=[{**rectangle(x='150',eid=eid),'vertical_range':interval('0','150',eid)}]
  d['through_holes']['deduction_rule']['threshold']='100'
 return {**b,'id':'synthetic-foundation','revision':1,'source_geometry':s}


def contact(eid='test',face='x_plus',start='0',end='600',level=1,name='合成C1'):
 return dict(name=name,material_index=1,level_index=level,face=face,parameters=dict(z_start=start,z_end=end),unit='mm',evidence={k:eid for k in ('z_start','z_end','measurement_basis')},measurement_basis='仅软件合成接触范围，所选侧面实际支模，按真实开口扣面积，不由体积扣减阈值恢复侧面，未选面不计，非真实工程工法。')


def body(s,eid='test'):
 return dict(name='合成基础外侧模板（非真实工程）',kind='foundation_formwork',parameters={},quantity_type='area',source_component_id=s['id'],source_component_hash=store.digest(snapshot(s)),contact_ranges=[contact(eid)],scope_confirmed=True,material='合成模板材料及支模做法，仅软件验收，不作施工依据',material_evidence_id=eid,measurement_basis='合成每个已审基础选入X正侧Z0至600mm，来源全部22基础，单面毛.3m²，实际侧开口宽60×高150mm扣.009m²，净.291m²；全部6.402m²，混凝土阈值100m²不扣孔体积但不能填实侧面；不额外乘重复，无模板清单重复。',basis_evidence_id=eid,rounding_digits=3,rounding_evidence_id=eid)


def setup(clients,pid=None,shape='rectangular'):
 e,r,pid,s,b,eid=threshold_setup(clients,pid,shape)
 upload(e,pid,'合成模板软件验收，非真实工程或合同规范。来源单组11根、重复2共22基础，材料79.200kg。矩形基础400×500×600mm，孔心X150,Y0mm，矩形孔长100宽60mm，Z0至150mm，接触X正侧边缘；混凝土单孔面积阈值100m2全部不扣，计量混凝土2.640m3，孔仍真实存在。模板X正侧Z0至600mm，单面毛.3m2、侧开口.009m2、净.291m2；乘22基础模板6.402m2。模板起Z改100mm则净5.434m2；同面0至100及100至600相接可分项，正高度重叠拒绝，异面可同高。圆柱或圆台阶仅整周按πD乘接触高。未选面不计，土侧不默认支模，不含内部孔壁、顶底、台阶水平面及支撑。TEST-FORMWORK仅合成模板映射编码，非官方条目。','foundation-formwork-basis.txt')
 for c in get(e,pid)['components']:approve(r,pid,'components',c)
 s=get(e,pid)['components'][-1];b=threshold_body(s,eid,shape)
 if shape=='rectangular':b['foundation_items'][0]['through_holes']=foundation(eid=eid)['foundation_items'][0]['through_holes']
 f=post(e,pid,b).json();approve(r,pid,'components',f);f=get(e,pid)['components'][-1];eid=get(e,pid)['evidence'][-1]['id'];b=body(f,eid)
 if shape!='rectangular':
  b['contact_ranges']=[contact(eid,face='circumference' if 'circular' in shape else 'x_plus',end='200' if 'stepped' in shape else '600')]
 return e,r,pid,f,b,eid


def test_real_side_opening_despite_volume_threshold_and_all_asts():
 s=foundation();b=body(s);before=copy.deepcopy((s,b));r=calculate(b,s);assert (s,b)==before and r['quantity']=='6.402' and r['unit']=='m²' and foundation_quantity(s)['quantity']=='2.640' and Decimal(foundation_quantity(s)['deduction'])==0
 assert Decimal(r['gross'])==Decimal('6.6') and Decimal(r['deduction'])==Decimal('.198') and r['contact_ranges'][0]['total_count']=='22'
 with localcontext() as ctx:
  ctx.prec=80;assert evaluate(r['formula_ast'],r['dimensions'])==Decimal(r['quantity'])
  for d in r['contact_ranges']:
   for k,a in d['formula_asts'].items():assert evaluate(a,r['dimensions'])==Decimal(d['normalized_parameters'][k] if k in ('z_start','z_end') else d[k])
   for h in d['apertures']:
    for k,a in h['formula_asts'].items():assert evaluate(a,r['dimensions'])==Decimal(h[k])


@pytest.mark.parametrize('shape',SHAPES)
@pytest.mark.parametrize('make',[equal_body,fixed_body,fixed_added,equal_added,fixed_adjusted,equal_adjusted,combined,lambda s:combined(s,equal=True)])
def test_shapes_eight_final_counts_and_source_repetition_once(shape,make):
 s=foundation(make,shape);f=foundation_quantity(s);b=body(s);b['contact_ranges']=[contact(face='circumference' if 'circular' in shape else 'y_minus',end='200' if 'stepped' in shape else '600')];r=calculate(b,s);d=r['contact_ranges'][0]
 with localcontext() as ctx:
  ctx.prec=80;p=layers(f['foundation_items'][0])[0]['normalized_parameters'];width=Decimal(p['diameter'])*Decimal(PI) if 'circular' in shape else Decimal(p['length']);area=width*(Decimal('.2') if 'stepped' in shape else Decimal('.6'));assert Decimal(d['single_gross_area'])==area and Decimal(r['gross'])==area*Decimal(f['foundation_items'][0]['total_count']) and r['deduction']=='0'


@pytest.mark.parametrize('unit,factor',[('mm','1'),('cm','.1'),('dm','.01'),('m','.001'),('km','.000001')])
def test_units_low_context_and_no_per_face_rounding(unit,factor):
 s=foundation();b=body(s);d=b['contact_ranges'][0];d['unit']=unit
 for k in d['parameters']:d['parameters'][k]=str(Decimal(d['parameters'][k])*Decimal(factor))
 with localcontext() as ctx:ctx.prec=6;r=calculate(b,s)
 assert r==calculate(b,s) and r['quantity']=='6.402'


@pytest.mark.parametrize('face,x,y,length,width,expected',[('x_minus','-150','0','100','60','.009'),('x_plus','150','0','100','60','.009'),('y_minus','0','-220','100','60','.015'),('y_plus','0','220','100','60','.015')])
def test_all_four_sides_and_corner_opening(face,x,y,length,width,expected):
 s=foundation();s['foundation_items'][0]['through_holes']['holes'][0]['parameters']=dict(x=x,y=y,length=length,width=width);b=body(s);b['contact_ranges'][0]['face']=face;r=calculate(b,s);assert Decimal(r['contact_ranges'][0]['single_opening_area'])==Decimal(expected)
 # A corner opening removes a strip on both distinct faces, without double counting either face.
 s['foundation_items'][0]['through_holes']['holes'][0]['parameters']=dict(x='150',y='220',length='100',width='60');b=body(s);b['contact_ranges']=[contact(face='x_plus'),contact(face='y_plus',name='合成C2')];assert Decimal(calculate(b,s)['deduction'])==Decimal('.024')*22


@pytest.mark.parametrize('start,end,expected',[('0','100','0.968'),('100','600','5.434'),('150','600','4.950'),('150.000000001','600','4.950'),('0','150','1.452'),('0','600','6.402')])
def test_exact_clipping_to_actual_side_aperture(start,end,expected):
 s=foundation();b=body(s);b['contact_ranges']=[contact(start=start,end=end)];assert calculate(b,s)['quantity']==expected


def test_adjacent_split_equal_whole_and_duplicate_overlap_rejected():
 s=foundation();b=body(s);whole=calculate(b,s);b['contact_ranges']=[contact(end='100'),contact(start='100',name='合成C2')];assert calculate(b,s)['quantity']==whole['quantity'];b['contact_ranges'][1]['parameters']['z_start']='99.999999999'
 with pytest.raises(DomainError,match='重叠'):calculate(b,s)


def test_circle_tangency_and_internal_rectangles_do_not_remove_side_area():
 for shape in SHAPES:
  s=foundation(shape=shape,edge=False);b=body(s);b['contact_ranges']=[contact(face='circumference' if 'circular' in shape else 'x_plus',end='200' if 'stepped' in shape else '600')];assert calculate(b,s)['deduction']=='0'
 s=foundation();c=s['foundation_items'][0]['through_holes'];h=c['holes'][0];h.pop('shape');h['parameters']=dict(x='180',y='0',diameter='40');h['evidence']={k:'test' for k in ('x','y','diameter','measurement_basis')};r=calculate(body(s),s);assert r['deduction']=='0'


def test_stepped_boundary_and_two_layers_not_twice_counted():
 s=foundation(shape='stepped_rectangular');b=body(s);b['contact_ranges']=[contact(end='200'),contact(start='200',end='600',level=2,name='合成C2')];r=calculate(b,s);assert r['quantity']=='4.840' and [d['total_count'] for d in r['contact_ranges']]==['22','22']
 b['contact_ranges'][0]['parameters']['z_end']='200.000000001'
 with pytest.raises(DomainError,match='本层'):calculate(b,s)


def test_seeded_independent_surface_oracle_and_final_rounding():
 rng=random.Random(9301)
 for _ in range(70):
  s=foundation(edge=False);d=s['foundation_items'][0];d.pop('through_holes');d['parameters']={k:str(rng.randint(1,999)) for k in ('length','width','height')};b=body(s);b['rounding_digits']=rng.randint(0,12);b['contact_ranges']=[contact(face=f,end=d['parameters']['height'],name=f) for f in ('x_minus','x_plus','y_minus','y_plus')];r=calculate(b,s)
  with localcontext() as ctx:ctx.prec=80;expected=2*(Decimal(d['parameters']['length'])+Decimal(d['parameters']['width']))*Decimal(d['parameters']['height'])/1000000*22;assert Decimal(r['gross'])==expected and Decimal(r['quantity'])==expected.quantize(Decimal(1).scaleb(-b['rounding_digits']),rounding=ROUND_HALF_UP)


@pytest.mark.parametrize('fault',['empty','null-row','extra','missing','bool-index','index','level','bool-level','face','face-type','circular-face','negative','reverse','zero','outside','precision','nan','bool','unit','parameters','evidence','blank-proof','basis','name','duplicate-name','duplicate-range','count','confirm','type','top-parameters','rounding','material','too-many'])
def test_strict_invalid_inputs(fault):
 s=foundation();b=body(s);d=b['contact_ranges'][0]
 if fault=='empty':b['contact_ranges']=[]
 elif fault=='null-row':b['contact_ranges']=[None]
 elif fault=='extra':d['count']='2'
 elif fault=='missing':d.pop('face')
 elif fault in ('index','bool-index'):d['material_index']=2 if fault=='index' else True
 elif fault in ('level','bool-level'):d['level_index']=2 if fault=='level' else True
 elif fault=='face':d['face']='top'
 elif fault=='face-type':d['face']=[]
 elif fault=='circular-face':d['face']='circumference'
 elif fault in ('negative','reverse','zero','outside','precision','nan','bool'):d['parameters']['z_start' if fault=='negative' else 'z_end']={'negative':'-1','reverse':'-1','zero':'0','outside':'601','precision':'600.0000000000000000001','nan':'NaN','bool':True}[fault]
 elif fault=='unit':d['unit']='m²'
 elif fault=='parameters':d['parameters']['height']='600'
 elif fault=='evidence':d['evidence'].pop('z_start')
 elif fault=='blank-proof':d['evidence']['z_start']=' '
 elif fault=='basis':d['measurement_basis']=''
 elif fault=='name':d['name']=''
 elif fault in ('duplicate-name','duplicate-range'):b['contact_ranges'].append(copy.deepcopy(d));b['contact_ranges'][-1]['name']=d['name'] if fault=='duplicate-name' else 'other'
 elif fault=='count':b['multiplicity']='2'
 elif fault=='confirm':b['scope_confirmed']=1
 elif fault=='type':b['quantity_type']='volume'
 elif fault=='top-parameters':b['parameters']={'area':'1'}
 elif fault=='rounding':b['rounding_digits']=True
 elif fault=='material':b['material']=''
 elif fault=='too-many':b['contact_ranges']=[copy.deepcopy(d)]*201
 with pytest.raises(DomainError):calculate(b,s)


def test_candidates_preview_no_mutation_claims_cross_records_and_revision(clients):
 e,r,pid,s,b,eid=setup(clients);p=get(e,pid);url=f'/api/v1/projects/{pid}';res=e.get(url+'/foundation-formwork-candidates',params=dict(component_id=s['id']));assert res.status_code==200 and res.json()['foundation_items'][0]['total_count']=='22';assert res.json()['foundation_items'][0]['levels'][0]['faces']==['x_minus','x_plus','y_minus','y_plus']
 assert post(e,pid,b,'/components/preview').json()['quantity']=='6.402' and get(e,pid)==p
 c=post(e,pid,b).json();approve(r,pid,'components',c);post(e,pid,b,expected=409);post(e,pid,{**b,'id':c['id']},expected=409)
 second=copy.deepcopy(b);second['contact_ranges'][0]['face']='x_minus';post(e,pid,second)
 c=get(e,pid)['components'][-2];new=copy.deepcopy(b);new['contact_ranges'][0]['parameters']['z_start']='100';new['object_hash']=store.digest(c);res=e.put(url+'/components/'+c['id'],json=new);assert res.status_code==200 and res.json()['preview']['quantity']=='5.434';assert get(e,pid)['component_history'][-1]==c;assert e.put(url+'/components/'+c['id'],json=new).status_code==409


@pytest.mark.parametrize('fault',['source','preview','range','evidence','geometry','formula','book'])
def test_tamper_reconstruction(clients,fault):
 e,r,pid,s,b,eid=setup(clients);c=post(e,pid,b).json()
 if fault=='book':approve(r,pid,'components',c);post(e,pid,{},'/takeoffs/calculate')
 def forge(p):
  c=p['components'][-1]
  if fault=='source':c['source_component_hash']='bad'
  elif fault=='preview':c['preview']['quantity']='99'
  elif fault=='range':c['contact_ranges'][0]['parameters']['z_start']='100'
  elif fault=='evidence':c['evidence']['formwork.basis']='bad'
  elif fault=='geometry':c['source_geometry']['foundation_items'][0]['parameters']['height']='100'
  elif fault=='formula':c['preview']['contact_ranges'][0]['formula_asts']['single_net_area']={'constant':'0'}
  else:p['takeoffs'][-1]['quantity']='99'
 mutate(pid,forge);collection='takeoffs' if fault=='book' else 'components';approve(r,pid,collection,get(e,pid)[collection][-1],expected=409)


@pytest.mark.parametrize('gate',['preview','save','review','calculate','book_review'])
def test_original_bytes_all_five_gates(clients,gate):
 e,r,pid,s,b,eid=setup(clients)
 if gate in ('review','calculate','book_review'):c=post(e,pid,b).json()
 if gate in ('calculate','book_review'):approve(r,pid,'components',c)
 if gate=='book_review':t=post(e,pid,{},'/takeoffs/calculate').json()[-1]
 doc=get(e,pid)['documents'][-1];(store.ROOT/'originals'/doc['hash']).write_bytes(b'corrupt')
 if gate=='preview':post(e,pid,b,'/components/preview',409)
 elif gate=='save':post(e,pid,b,expected=409)
 elif gate=='review':approve(r,pid,'components',c,expected=409)
 elif gate=='calculate':post(e,pid,{},'/takeoffs/calculate',409)
 else:approve(r,pid,'takeoffs',t,expected=409)


def test_roles_foreign_evidence_and_source_revision_transitive_stale(clients):
 e,r,pid,s,b,eid=setup(clients);post(r,pid,b,expected=403);assert clients['other'].get(f'/api/v1/projects/{pid}/foundation-formwork-candidates',params=dict(component_id=s['id'])).status_code==404
 other=create(e)['id'];upload(e,other,'异项目模板原件');bad=copy.deepcopy(b);bad['material_evidence_id']=get(e,other)['evidence'][-1]['id'];post(e,pid,bad,expected=404)
 c=post(e,pid,b).json();approve(r,pid,'components',c)
 s['object_hash']=store.digest(s);s['foundation_items'][0]['parameters']['height']='700';res=e.put(f'/api/v1/projects/{pid}/components/'+s['id'],json=s);assert res.status_code==200;approve(r,pid,'components',res.json());assert get(e,pid)['components'][-1]['state']=='stale'
 with pytest.raises(HTTPException):calculate(b,get(e,pid)['components'][-2])


@pytest.mark.parametrize('shape',SHAPES)
@pytest.mark.parametrize('fmt',['json','xlsx','pdf'])
def test_three_exports_boq_profile_and_old_quote_isolation(clients,shape,fmt):
 e,r,pid,_,q=priced_project(clients);e,r,pid,s,b,eid=setup(clients,pid,shape);c=post(e,pid,b).json();approve(r,pid,'components',c)
 for t in post(e,pid,{},'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
 boq=post(e,pid,{},'/boq/generate').json();assert any(v['unit']=='m²' and v['quantity']==c['preview']['quantity'] for v in boq)
 url=f'/api/v1/projects/{pid}/exports';res=e.get(url,params=dict(format=fmt,unpriced=True));assert res.status_code==200,res.text
 if fmt=='json':assert TITLES<={v[0] for v in res.json()['tables']}
 elif fmt=='xlsx':assert TITLES<=set(load_workbook(io.BytesIO(res.content)).sheetnames)
 else:assert all(v in ''.join(p.extract_text() for p in PdfReader(io.BytesIO(res.content)).pages) for v in TITLES)
 assert not any(v[0] in TITLES for v in e.get(url,params=dict(format='json',quote_id=q['id'])).json()['tables'])


def test_adjacent_claim_ranges_different_units_and_same_face_different_level(clients):
 e,r,pid,s,b,eid=setup(clients);b['contact_ranges'][0]['parameters']['z_end']='100';c=post(e,pid,b).json();approve(r,pid,'components',c);b['contact_ranges'][0]['parameters']=dict(z_start='.1',z_end='.6');b['contact_ranges'][0]['unit']='m';assert post(e,pid,b).json()['preview']['quantity']=='5.434';b['contact_ranges'][0]['parameters']['z_start']='.099999999999';post(e,pid,b,expected=409)


@pytest.mark.parametrize('shape',SHAPES)
def test_every_recorded_formula_and_conversion_has_current_source_proof(shape):
 s=foundation(shape=shape);b=body(s);f=foundation_quantity(s);b['contact_ranges']=[contact(face='circumference' if 'circular' in shape else 'x_plus',end='200' if 'stepped' in shape else '600')];r=calculate(b,s)
 with localcontext() as ctx:
  ctx.prec=80
  for d in r['contact_ranges']:
   for k,node in d['formula_asts'].items():assert evaluate(node,r['dimensions'])==Decimal(d['normalized_parameters'][k] if k in ('z_start','z_end') else d[k])
 for c in r['unit_conversions']:assert c['parameter'] in r['source_evidence'] and c['parameter'] in r['dimensions']
 assert f==foundation_quantity(s)


def test_multiple_apertures_touching_z_and_full_face_opening_zero_net():
 s=foundation();c=s['foundation_items'][0]['through_holes'];h=c['holes'][0];second=copy.deepcopy(h);second['name']='合成H2';second['vertical_range']=interval('150','300');c['holes'].append(second);r=calculate(body(s),s);assert len(r['contact_ranges'][0]['apertures'])==2 and r['quantity']=='6.204'
 h['parameters']['width']='500';c['holes']=[h];b=body(s);b['contact_ranges'][0]['parameters']['z_end']='150';r=calculate(b,s);assert r['quantity']=='0.000' and Decimal(r['gross'])==Decimal(r['deduction'])>0


def test_multiple_source_items_only_selected_contact_counts():
 s=foundation();m=s['source_geometry'];m['material_items'].append({**copy.deepcopy(m['material_items'][0]),'name':'合成第二立柱项'});s['source_component_hash']=store.digest(snapshot(m));second=copy.deepcopy(s['foundation_items'][0]);second['material_index']=2;s['foundation_items'].append(second);b=body(s);r=calculate(b,s);assert r['quantity']=='6.402' and r['contact_ranges'][0]['total_count']=='22'
 second=copy.deepcopy(b['contact_ranges'][0]);second.update(name='合成C2',material_index=2);b['contact_ranges'].append(second);assert calculate(b,s)['quantity']=='12.804'


def test_revision_preserves_history_old_boq_and_stales_books_then_ancestor_stales_template(clients):
 e,r,pid,s,b,eid=setup(clients);c=post(e,pid,b).json();approve(r,pid,'components',c)
 for t in post(e,pid,{},'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
 post(e,pid,{},'/boq/generate');p=get(e,pid);c=p['components'][-1];new=copy.deepcopy(b);new['object_hash']=store.digest(c);new['contact_ranges'][0]['parameters']['z_start']='100';res=e.put(f'/api/v1/projects/{pid}/components/'+c['id'],json=new);assert res.status_code==200 and res.json()['preview']['quantity']=='5.434';p2=get(e,pid);assert p2['component_history'][-1]==c and p2['boq']==p['boq'] and all(t['state']=='stale' for t in p2['takeoffs']);approve(r,pid,'components',res.json())
 m=p2['components'][-3];assert m['kind']=='railing_material';m['object_hash']=store.digest(m);m['material_items'][0]['post_ranges'][0]['spacing']='3000';res=e.put(f'/api/v1/projects/{pid}/components/'+m['id'],json=m);assert res.status_code==200;p3=get(e,pid);assert p3['components'][-2]['state']==p3['components'][-1]['state']=='stale'


def test_maximum_ranges_long_proofs_export_and_area_mapping(clients):
 e,r,pid,s,b,eid=setup(clients);b['contact_ranges']=[contact(eid,start=str(i*3),end=str((i+1)*3),name='合成C'+str(i)) for i in range(200)];b['contact_ranges'][-1]['measurement_basis']='合成'+('实际接触范围及扣减依据'*140);b['material']='合成'+('模板做法依据'*140);c=post(e,pid,b).json();assert c['preview']['quantity']=='6.402';approve(r,pid,'components',c)
 for t in post(e,pid,{},'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
 post(e,pid,{},'/boq/generate');res=e.get(f'/api/v1/projects/{pid}/exports',params=dict(format='pdf',unpriced=True));assert res.status_code==200,res.text
