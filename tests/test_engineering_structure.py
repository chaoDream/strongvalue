"""Synthetic engineering hierarchy and immutable, unallocated project-fee rollup contracts."""
import copy
import io
import uuid
from decimal import Decimal
import pytest
from openpyxl import load_workbook
from backend import store, main as core
from backend.quote_comparison import compare_quotes
from tests.test_workflow import get, approve, priced_project
from tests.test_quantity_batches import setup as split_project


def setup(clients):
    e,r,pid,old,change=split_project(clients);base=f'/api/v1/projects/{pid}'
    target=e.get(base+'/engineering-structure-targets').json();boq=target['boq']
    body={'request_id':str(uuid.uuid4()),'name':'合成两单项三级结构','reason':'纯数学归属测试，不代表真实工程划分',
          'source_hash':target['source_hash'],'base_hash':target['base_hash'],'evidence_id':get(e,pid)['evidence'][0]['id'],
          'nodes':[{'id':'s1','name':'合成单项甲','kind':'single','parent_id':None},
                   {'id':'u1','name':'合成单位甲','kind':'unit','parent_id':'s1'},
                   {'id':'s2','name':'合成单项乙','kind':'single','parent_id':None},
                   {'id':'u2','name':'合成单位乙','kind':'unit','parent_id':'s2'}],
          'assignments':{boq[0]['id']:'u1',boq[1]['id']:'u2'}}
    return e,r,pid,old,body


def create(e,pid,body):
    response=e.post(f'/api/v1/projects/{pid}/engineering-structures',json=body)
    assert response.status_code==200,response.text
    return response.json()['structure']


def quote(e,pid):
    response=e.post(f'/api/v1/projects/{pid}/quotes/calculate');assert response.status_code==200,response.text
    return response.json()


def test_reviewed_rollup_exact_closure_snapshots_exports_and_history(clients):
    e,r,pid,old,body=setup(clients);base=f'/api/v1/projects/{pid}'
    s=create(e,pid,body)
    assert e.post(base+'/quotes/calculate').status_code==409
    approve(e,pid,'engineering_structures',s,expected=403)
    approve(r,pid,'engineering_structures',s)
    q=quote(e,pid);summary=q['engineering_summary']
    assert summary['closed'] and summary['subtotal']==q['subtotal']
    assert summary['total']==q['total']
    assert Decimal(summary['subtotal'])+Decimal(summary['project_fees'])==Decimal(summary['total'])
    assert summary['project_fees']=='47.25'
    units=[n for n in summary['nodes'] if n['kind']=='unit'];singles=[n for n in summary['nodes'] if n['kind']=='single']
    assert sum(Decimal(n['subtotal']) for n in units)==Decimal(q['subtotal'])
    assert sum(Decimal(n['subtotal']) for n in singles)==Decimal(q['subtotal'])
    assert [n['line_count'] for n in units]==[1,1]
    assert next(v for v in get(e,pid)['quotes'] if v['id']==old['id'])==old
    for fmt in ['xlsx','pdf','json','zip']:
        result=e.get(base+'/exports',params={'format':fmt,'quote_id':q['id']});assert result.status_code==200
        if fmt=='xlsx':
            wb=load_workbook(io.BytesIO(result.content));assert '工程结构与归属' in wb.sheetnames
            rows=list(wb['分级汇总'].values);assert any('建设项目' in row for row in rows)
        if fmt=='json':assert result.json()['quote']['engineering_summary']==summary
    old_export=e.get(base+'/exports',params={'format':'json','quote_id':old['id']}).json()
    assert '工程结构与归属' not in [section[0] for section in old_export['tables']]
    p=get(e,pid);comparison=compare_quotes(p,old['id'],q['id'])
    assert comparison['engineering_structure_changed'] and comparison['engineering_assignment_changes']
    approve(r,pid,'quotes',q)
    approve(r,pid,'engineering_structures',s,decision='reject')
    changed=get(e,pid)
    assert q['id'] in changed['stale_quotes'] and changed['report_approval'] is None
    assert e.post(base+'/quotes/calculate').status_code==409
    draft=e.get(base+'/exports?format=json&unpriced=true').json()['unpriced']
    assert draft['engineering_structures'][-1]['state']=='rejected'


@pytest.mark.parametrize('kind',['duplicate_id','cycle','orphan','single_parent','missing_unit','same_sibling_name','missing_assignment','unknown_assignment','single_assignment','unknown_field','blank_name','too_many'])
def test_invalid_structure_is_atomic(clients,kind):
    e,r,pid,old,body=setup(clients)
    if kind=='duplicate_id':body['nodes'][1]['id']='s1'
    if kind=='cycle':body['nodes'][1]['parent_id']='u1'
    if kind=='orphan':body['nodes'][1]['parent_id']='missing'
    if kind=='single_parent':body['nodes'][0]['parent_id']='s2'
    if kind=='missing_unit':body['nodes'].pop()
    if kind=='same_sibling_name':body['nodes'][2]['name']=body['nodes'][0]['name']
    if kind=='missing_assignment':body['assignments'].pop(next(iter(body['assignments'])))
    if kind=='unknown_assignment':body['assignments']['foreign']='u1'
    if kind=='single_assignment':body['assignments'][next(iter(body['assignments']))]='s1'
    if kind=='unknown_field':body['nodes'][0]['amount']='0'
    if kind=='blank_name':body['nodes'][0]['name']=' '
    if kind=='too_many':body['nodes']*=60
    before=get(e,pid)
    response=e.post(f'/api/v1/projects/{pid}/engineering-structures',json=body)
    assert response.status_code in (409,422),response.text
    assert get(e,pid)==before


def test_structure_revision_reassignment_and_historical_review_guard(clients):
    e,r,pid,old,body=setup(clients);base=f'/api/v1/projects/{pid}'
    s=create(e,pid,body);approve(r,pid,'engineering_structures',s);left=quote(e,pid)
    target=e.get(base+'/engineering-structure-targets').json();body.update(source_hash=target['source_hash'],base_hash=target['base_hash'],request_id=str(uuid.uuid4()))
    body['assignments']={bid:'u2' for bid in body['assignments']}
    second=create(e,pid,body)
    approve(r,pid,'engineering_structures',get(e,pid)['engineering_structures'][0],decision='reject',expected=409)
    approve(r,pid,'engineering_structures',second);right=quote(e,pid)
    assert left['total']==right['total']
    p=get(e,pid);comparison=compare_quotes(p,left['id'],right['id'])
    assert comparison['closed'] and comparison['total_delta']=='0.00'
    assert comparison['engineering_structure_changed'] and len(comparison['engineering_assignment_changes'])==1
    assert next(n for n in right['engineering_summary']['nodes'] if n['id']=='u1')['line_count']==0
    assert next(v for v in p['quotes'] if v['id']==left['id'])==left


@pytest.mark.parametrize('kind',['name','scope'])
def test_boq_identity_changes_require_a_new_structure(clients,kind):
    e,r,pid,old,body=setup(clients);base=f'/api/v1/projects/{pid}'
    s=create(e,pid,body);approve(r,pid,'engineering_structures',s)
    b=get(e,pid)['boq'][0]
    if kind=='name':
        updated=e.put(base+'/boq/'+b['id'],json={'name':b['name']+' changed'}).json();approve(r,pid,'boq',updated)
    else:
        split=e.post(base+'/boq/'+b['id']+'/split',json={'parts':[{'name':'新甲','quantity':'3'},{'name':'新乙','quantity':'4'}]})
        assert split.status_code==200,split.text
        for entry in split.json():approve(r,pid,'boq',entry)
    assert e.post(base+'/quotes/calculate').status_code==409
    assert any('工程层级' in v for v in get(e,pid)['blockers'])


def test_price_changes_do_not_invalidate_declared_assignment(clients):
    e,r,pid,old,body=setup(clients);base=f'/api/v1/projects/{pid}'
    s=create(e,pid,body);approve(r,pid,'engineering_structures',s)
    b=get(e,pid)['boq'][0];qs=copy.deepcopy(b['quotas']);qs[0]['resources'][0]['price']='120'
    updated=e.put(base+'/boq/'+b['id'],json={'quotas':qs}).json()
    for entry in get(e,pid)['boq']:approve(r,pid,'boq',entry)
    q=quote(e,pid);assert q['engineering_summary']['structure_id']==s['id']


@pytest.mark.parametrize('kind',['foreign','superseded','missing_bytes'])
def test_evidence_rechecked_and_access_scoped(clients,kind):
    e,r,pid,old,body=setup(clients);base=f'/api/v1/projects/{pid}'
    assert clients['other'].get(base+'/engineering-structure-targets').status_code==404
    if kind=='foreign':
        body['evidence_id']='foreign';assert e.post(base+'/engineering-structures',json=body).status_code==404;return
    s=create(e,pid,body);approve(r,pid,'engineering_structures',s)
    p=get(e,pid)
    if kind=='missing_bytes':
        q=quote(e,pid)
        path=store.ROOT/'originals'/p['documents'][0]['hash']
        path.write_bytes(b'corrupted synthetic original')
        approve(r,pid,'quotes',q,expected=409)
        path.unlink()
    else:
        with store.transaction() as c:
            p=core.task_project(c,pid,{'tenant':'local'});p['documents'][0]['superseded']=True;store.save(c,p)
    assert e.post(base+'/quotes/calculate').status_code==409


def test_duplicate_and_stale_submission_contract(clients):
    e,r,pid,old,body=setup(clients);base=f'/api/v1/projects/{pid}'
    s=create(e,pid,body)
    assert e.post(base+'/engineering-structures',json=body).json()['duplicate']
    body['reason']+=' changed';assert e.post(base+'/engineering-structures',json=body).status_code==409
    body['request_id']=str(uuid.uuid4());assert e.post(base+'/engineering-structures',json=body).status_code==409
    assert len(get(e,pid)['engineering_structures'])==1
