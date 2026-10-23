import copy
import io
from decimal import Decimal
import pytest
from openpyxl import load_workbook
from pypdf import PdfReader
from backend import store
from tests.test_dxf_routes import setup,polyline,read,body,preview,adopt,mutate
from tests.test_dxf_fragments import prepared,part,select,request
from tests.test_dxf_assembly import resolve
from tests.test_workflow import get,approve,upload,priced_project

B=[(9000,12000,19000),(9000,12000,17000),(6000,8000,17000),(6000,8000,18000)]


def ranged(clients,pid=None,b=None):
    c,r,pid,did,eid,url,a=prepared(clients,pid=pid,extra=polyline(b or B,handle='B'))
    upload(c,pid,'合成组装独立依据：A原P2至P3和B原P2至P3三维直线距离各5m，容差0m；A原P2至P4顺接B原P3至P2为设计中心线，共同WCS及高程已核定；重复2组，总量修约3位。','assembly-fragment-basis.txt')
    eid=get(c,pid)['evidence'][-1]['id'];objects=read(c,url)['candidates']
    parts=[part(objects[0],2,4,False),part(objects[1],2,3,True)]
    row=resolve(c,url,parts)
    req=body(row,eid,segments=parts,calibration={'from_point':1,'to_point':2,'distance':'5','unit':'m','tolerance':'0'})
    return c,r,pid,did,eid,url,objects,parts,row,req


def test_fragment_assembly_is_readonly_and_keeps_full_sources_indices_and_direction(clients):
    c,r,pid,did,eid,url,objects,parts,row,req=ranged(clients);before=get(c,pid)
    assert row==resolve(r,url,parts) and get(c,pid)==before
    assert row['version']=='dxf-wcs-assembly-1.1' and row['single_length_m']=='15' and len(row['points'])==4
    assert row['points']==objects[0]['points'][1:4]+[objects[1]['points'][1]]
    assert row['selection']==parts
    for i,indices in enumerate([[2,3,4],[3,2]]):
        src=row['source_objects'][i]
        assert src['kind']=='FRAGMENT' and src['source_point_indices']==indices
        assert src['source_object']=={k:v for k,v in objects[i].items() if k!='object_hash'}
    assert [(p['start_point'],p['end_point'],p['reverse']) for p in row['source_ranges']]==[(1,3,False),(3,4,True)]
    assert row['joins'][0]['point_index']==3
    assert preview(c,url,row,req)['result']['quantity']=='30.000' and get(c,pid)==before


def test_fragment_assembly_adoption_review_takeoff_boq_roles_and_idempotency(clients):
    c,r,pid,did,eid,url,objects,parts,row,req=ranged(clients)
    resolve(clients['other'],url,parts,404);view=preview(c,url,row,req);adopt(r,url,row,req,view,403)
    saved=adopt(c,url,row,req,view);comp=saved['component']
    assert comp['state']=='draft' and not get(c,pid)['takeoffs']
    assert adopt(c,url,row,req,view)['duplicate'] and len(get(c,pid)['components'])==1
    approve(r,pid,'components',comp)
    t=c.post(f'/api/v1/projects/{pid}/takeoffs/calculate').json()[0];assert t['quantity']=='30.000'
    approve(r,pid,'takeoffs',t)
    assert c.post(f'/api/v1/projects/{pid}/boq/generate').json()[0]['quantity']=='30.000'


def test_whole_object_and_fragment_can_be_mixed_without_reversing_twice(clients):
    c,r,pid,did,eid,url,objects,parts,row,req=ranged(clients)
    parts[0]={k:parts[0][k] for k in ['candidate_id','object_hash','reverse']}
    row=resolve(c,url,parts);assert row['single_length_m']=='28'
    assert row['source_objects'][0]['kind']=='POLYLINE' and row['source_objects'][1]['kind']=='FRAGMENT'
    assert row['points'][-1]==objects[1]['points'][1]
    req=body(row,eid,segments=parts);assert preview(c,url,row,req)['result']['quantity']=='56.000'


def test_adjacent_fragments_of_same_raw_object_form_one_path_and_hold_full_claim(clients):
    c,r,pid,did,eid,url,source=prepared(clients)
    parts=[part(source,1,2,False),part(source,2,4,False)];row=resolve(c,url,parts)
    assert row['points']==source['points'] and row['single_length_m']=='23'
    req=body(row,eid,segments=parts);adopt(c,url,row,req,preview(c,url,row,req))
    other=select(c,url,part(source,3,4,True));req=request(other,eid,request_id='synthetic-already-assembled-part')
    adopt(c,url,other,req,preview(c,url,other,req),409)


@pytest.mark.parametrize('case',['same','reverse','overlapping','whole_and_fragment'])
def test_overlapping_raw_ranges_are_rejected_inside_one_sequence(clients,case):
    c,r,pid,did,eid,url,source=prepared(clients)
    parts=[part(source,1,3,False),part(source,1,3,case=='reverse')]
    if case=='overlapping':parts[1]=part(source,2,4,False)
    if case=='whole_and_fragment':parts[0]={k:parts[0][k] for k in ['candidate_id','object_hash','reverse']}
    before=get(c,pid);err=resolve(c,url,parts,422)
    assert '重复使用了CAD原始管段' in err['detail'] and get(c,pid)==before


@pytest.mark.parametrize('changes',[
    {'from_point':None},{'to_point':None},{'from_point':True},{'to_point':3.0},{'from_point':'2'},
    {'from_point':0},{'to_point':102},{'to_point':5},{'from_point':3,'to_point':2},
    {'from_point':2,'to_point':2},{'reverse':'true'},{'pixels':[1,2]},
    {'fragment':{'from_point':2,'to_point':3}},{'snap_tolerance':'1'}])
def test_invalid_fragment_selection_in_assembly_is_atomic(clients,changes):
    c,r,pid,did,eid,url,objects,parts,row,req=ranged(clients);before=get(c,pid)
    parts[1].update(changes);resolve(c,url,parts,422);assert get(c,pid)==before


@pytest.mark.parametrize('missing',['from_point','to_point','reverse'])
def test_partial_range_or_implicit_direction_cannot_be_used(clients,missing):
    c,r,pid,did,eid,url,objects,parts,row,req=ranged(clients)
    parts[1].pop(missing);resolve(c,url,parts,422)


def test_fragment_candidate_identity_cannot_replace_raw_source_reference(clients):
    c,r,pid,did,eid,url,objects,parts,row,req=ranged(clients)
    frag=select(c,url,parts[1]);parts[1].update(candidate_id=frag['id'],object_hash=frag['object_hash'])
    resolve(c,url,parts,409)


@pytest.mark.parametrize('first',['assembly','fragment','whole','legacy_assembly'])
def test_fragment_assembly_claims_conflict_with_existing_modes(clients,first):
    c,r,pid,did,eid,url,objects,parts,row,req=ranged(clients)
    if first=='assembly':
        adopt(c,url,row,req,preview(c,url,row,req));new=objects[1];newreq=body(new,eid,request_id='synthetic-whole-after-range-assembly',calibration={'from_point':2,'to_point':3,'distance':'5','unit':'m','tolerance':'0'})
    else:
        if first=='fragment':
            old=select(c,url,parts[0]);oldreq=request(old,eid,request_id='synthetic-part-before-assembly')
        elif first=='whole':old=objects[0];oldreq=body(old,eid)
        else:
            oldparts=[{'candidate_id':objects[0]['id'],'object_hash':objects[0]['object_hash'],'reverse':False},part(objects[1],3,4,False)]
            old=resolve(c,url,oldparts);oldreq=body(old,eid,segments=oldparts)
        adopt(c,url,old,oldreq,preview(c,url,old,oldreq));new=row;newreq=req
    before=get(c,pid);adopt(c,url,new,newreq,preview(c,url,new,newreq),409);assert get(c,pid)==before


def test_remaining_range_can_be_adopted_without_overclaiming_whole_source(clients):
    c,r,pid,did,eid,url,objects,parts,row,req=ranged(clients)
    comp=adopt(c,url,row,req,preview(c,url,row,req))['component']
    remaining=select(c,url,part(objects[0],1,2,False));extra=request(remaining,eid,'13',request_id='synthetic-remaining-range-after-assembly')
    second=adopt(c,url,remaining,extra,preview(c,url,remaining,extra))['component']
    for v in [comp,second]:approve(r,pid,'components',v)
    ts=c.post(f'/api/v1/projects/{pid}/takeoffs/calculate').json();assert sorted(t['quantity'] for t in ts)==['26.000','30.000']
    for t in ts:approve(r,pid,'takeoffs',t)
    assert sum(Decimal(v['quantity']) for v in c.post(f'/api/v1/projects/{pid}/boq/generate').json())==56


def test_revision_of_range_assembly_releases_previous_claims_and_stales_downstream(clients):
    c,r,pid,did,eid,url,objects,parts,row,req=ranged(clients)
    old=adopt(c,url,row,req,preview(c,url,row,req))['component'];approve(r,pid,'components',old)
    t=c.post(f'/api/v1/projects/{pid}/takeoffs/calculate').json()[0];approve(r,pid,'takeoffs',t)
    fragment=select(c,url,part(objects[0],2,3,False));req=request(fragment,eid,request_id='synthetic-assembly-revised-to-part',component_id=old['id'],component_hash=store.digest(get(c,pid)['components'][0]))
    new=adopt(c,url,fragment,req,preview(c,url,fragment,req))['component']
    assert new['revision']==2 and get(c,pid)['takeoffs'][0]['state']=='stale'
    assert get(c,pid)['component_history'][0]['dxf_trace']['candidate']=={k:v for k,v in row.items() if k!='object_hash'}
    released=select(c,url,parts[1]);req=request(released,eid,request_id='synthetic-released-assembly-source')
    adopt(c,url,released,req,preview(c,url,released,req))


@pytest.mark.parametrize('field',['fragment','source_object','source_point_indices','selection','source_ranges','joins'])
def test_saved_nested_provenance_is_rebuilt_before_review(clients,field):
    c,r,pid,did,eid,url,objects,parts,row,req=ranged(clients)
    comp=adopt(c,url,row,req,preview(c,url,row,req))['component']
    def corrupt(p):
        for trace in [p['components'][0]['dxf_trace'],p['dxf_route_adoptions'][0]['trace']]:
            target=trace['candidate']['source_objects'][1] if field in ['fragment','source_object','source_point_indices'] else trace['candidate']
            target[field]=None
    mutate(pid,corrupt);approve(r,pid,'components',get(c,pid)['components'][0],expected=409)


@pytest.mark.parametrize('change',['bytes','job','evidence','selection'])
def test_stale_sources_and_changed_selection_cannot_use_prior_preview(clients,change):
    c,r,pid,did,eid,url,objects,parts,row,req=ranged(clients);view=preview(c,url,row,req)
    if change=='bytes':(store.ROOT/'originals'/objects[0]['document_hash']).write_bytes(b'changed')
    elif change=='selection':req['segments'][1]['from_point']=1
    else:
        def mutate_source(p):
            if change=='job':next(j for j in p['jobs'] if j['document_id']==did)['state']='failed'
            else:next(e for e in p['evidence'] if e['id']==eid)['text']+='changed'
        mutate(pid,mutate_source)
    before=get(c,pid);adopt(c,url,row,req,view,409);assert get(c,pid)==before


def test_fragment_join_keeps_exact_raw_string_from_both_sides(clients):
    b=copy.deepcopy(B);b[2]=('6000.00','8E3','17000.0')
    c,r,pid,did,eid,url,objects,parts,row,req=ranged(clients,b=b)
    assert row['points'][2]=={'x':'6000','y':'8000','z':'17000'}
    assert row['joins'][0]['outgoing_start']=={'x':'6000.00','y':'8E3','z':'17000.0'}
    assert row['source_objects'][1]['points'][0]==row['joins'][0]['outgoing_start']


@pytest.mark.parametrize('count',[52,53])
def test_fragment_assembly_vertex_limit_does_not_truncate(clients,count):
    c,r,pid,did,eid,url=setup(clients,polyline([(i,0,0) for i in range(101)])+polyline([(i,0,0) for i in range(50,151)],handle='B'))
    objects=read(c,url)['candidates'];parts=[part(objects[0],2,51,False),part(objects[1],1,count,False)]
    row=resolve(c,url,parts,200 if count==52 else 422)
    if count==52:assert len(row['points'])==101 and row['points'][-1]['x']=='101'


def test_unsupported_whole_source_cannot_be_made_eligible_by_selecting_a_part(clients):
    c,r,pid,did,eid,url=setup(clients,polyline(flags=9)+polyline(handle='B'))
    objects=read(c,url)['candidates'];resolve(c,url,[part(objects[0],1,2,False),part(objects[1],2,3,False)],422)


@pytest.mark.parametrize('fmt',['json','xlsx','pdf'])
def test_mapping_export_uses_snapshot_and_preserves_old_quotes(clients,fmt):
    c,r,pid,_,old=priced_project(clients);c,r,pid,did,eid,url,objects,parts,row,req=ranged(clients,pid=pid)
    adopt(c,url,row,req,preview(c,url,row,req));base=f'/api/v1/projects/{pid}/exports'
    res=c.get(base,params={'format':fmt,'unpriced':True});assert res.status_code==200,res.text
    if fmt=='json':
        table=next(t for t in res.json()['tables'] if t[0]=='DXF组装取段顶点映射')
        assert 'P3→原P3；P4→原P2' in str(table) and '原P2—原P3 / 反向' in str(table)
    elif fmt=='xlsx':assert 'DXF组装取段顶点映射' in load_workbook(io.BytesIO(res.content)).sheetnames
    else:assert 'DXF组装取段顶点映射' in ''.join(p.extract_text() for p in PdfReader(io.BytesIO(res.content)).pages)
    historic=c.get(base,params={'format':'json','quote_id':old['id']}).json()
    assert not any(t[0]=='DXF组装取段顶点映射' for t in historic['tables']) and historic['quote']==old
