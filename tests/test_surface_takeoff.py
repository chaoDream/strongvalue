"""Synthetic boundary-face math and evidence workflow; no professional approval."""
import copy
import io
import json
import random
from decimal import Decimal, localcontext
import pytest
from openpyxl import load_workbook
from pypdf import PdfReader
from backend import store
from backend.engine import DomainError
from backend.surface_takeoff import boundary, calculate, snapshot
from tests.test_spatial_takeoff import payload as spatial_payload, region
from tests.test_compound_takeoff import ast_value
from tests.test_workflow import create, get, upload, approve, priced_project


def geometry(eid='test'):
    s=spatial_payload(eid);s.update(id='synthetic-solid',revision=1,name='合成开口池体（非真实工程）')
    s['regions']=[region('外块',0,0,0,4,3,2,eid=eid),region('净空',1,1,'.5',2,1,'1.5','subtract',eid)]
    return s


def payload(source,eid='test',included=None):
    faces=boundary(source)['faces']
    return {'name':'合成竖向模板接触面（非真实工程）','kind':'surface_faces','parameters':{},'quantity_type':'area','purpose':'formwork',
            'source_component_id':source['id'],'source_component_hash':store.digest(snapshot(source)),
            'measurement_basis':'仅合成：竖向边界全部计入，水平面全部排除，无专业计量效力','basis_evidence_id':eid,
            'face_decisions':[{'face_id':f['face_id'],'included':included(f) if included else not f['normal'].startswith('z'),
                               'basis':'合成逐面核定：计入竖向面，排除水平面','evidence_id':eid} for f in faces]}


def setup(clients):
    e,r=clients['estimator'],clients['reviewer'];pid=create(e)['id']
    upload(e,pid,'仅合成XYZ同基准：外块(0,0,0)4×3×2m；净空(1,1,0.5)2×1×1.5m，顶部敞口。单组1；竖向面计入模板，水平面排除。数学测试无专业效力。')
    eid=get(e,pid)['evidence'][0]['id'];s=geometry(eid);s.pop('id');s.pop('revision')
    response=e.post('/api/v1/projects/'+pid+'/components',json=s);assert response.status_code==200,response.text
    source=response.json();approve(r,pid,'components',source)
    source=get(e,pid)['components'][0]
    return e,r,pid,source,payload(source,eid),eid


def test_open_pool_outer_and_inner_boundary_and_face_selection():
    s=geometry();b=boundary(s);r=calculate(payload(s),s)
    assert Decimal(b['total_area'])==61 and Decimal(r['quantity'])==37
    assert Decimal(r['surface_summary']['excluded_area'])==24
    assert sum(Decimal(f['area']) for f in b['faces'] if f['normal'].startswith('z'))==24
    with localcontext() as ctx:
        ctx.prec=80
        assert ast_value(r['formula_ast'],r['dimensions'])==37
        for f in b['faces']:assert ast_value(f['formula_ast'],b['dimensions'])==Decimal(f['area'])


def test_shared_contact_faces_disappear_and_identical_solids_do_not_duplicate():
    s=geometry();s['regions']=[region('a',0,0,0,2,2,2),region('b',2,0,0,1,2,2)]
    b=boundary(s);assert Decimal(b['total_area'])==32 and len(b['faces'])==6
    s['regions'].append(region('duplicate',0,0,0,2,2,2))
    duplicate=boundary(s)
    assert duplicate['faces']==b['faces'] and duplicate['total_area']==b['total_area']


def test_internal_cavity_and_through_opening_add_their_boundary():
    s=geometry();s['regions']=[region('outer',0,0,0,3,3,3),region('void',1,1,1,1,1,1,'subtract')]
    assert Decimal(boundary(s)['total_area'])==60  # 54 outside + 6 cavity
    s['regions']=[region('outer',0,0,0,4,3,2),region('hole',1,1,0,1,1,2,'subtract')]
    assert Decimal(boundary(s)['total_area'])==58  # 52 outside - 2 + 8 opening


def test_negative_xyz_units_repetition_and_all_excluded():
    s=geometry();s['regions']=[region('a',-1000,-2000,-3000,2000,3000,4000,unit='mm')];s['multiplicity']='3'
    r=calculate(payload(s,included=lambda f:True),s)
    assert Decimal(r['surface_summary']['total_area'])==52 and Decimal(r['quantity'])==156
    assert r['unit_conversions'][0]['source_unit']=='mm'
    zero=calculate(payload(s,included=lambda f:False),s);assert Decimal(zero['quantity'])==0
    with localcontext() as ctx:
        ctx.prec=80
        assert ast_value(r['formula_ast'],r['dimensions'])==156
        assert ast_value(zero['formula_ast'],zero['dimensions'])==0


def test_surface_matches_independent_voxel_exposed_face_oracle():
    rng=random.Random(2701)
    offsets=[(1,0,0),(-1,0,0),(0,1,0),(0,-1,0),(0,0,1),(0,0,-1)]
    for run in range(35):
        s=geometry();s['regions']=[];cells=set();cuts=set()
        for i in range(rng.randint(1,6)):
            x,y,z=[rng.randint(-3,3) for _ in range(3)];l,w,h=[rng.randint(1,3) for _ in range(3)]
            s['regions'].append(region('a'+str(i),x,y,z,l,w,h))
            cells|={(a,b,c) for a in range(x,x+l) for b in range(y,y+w) for c in range(z,z+h)}
        for i in range(rng.randint(0,4)):
            x,y,z=rng.choice(sorted(cells));s['regions'].append(region('cut'+str(i),x,y,z,1,1,1,'subtract'));cuts.add((x,y,z))
        cells-=cuts
        if not cells:continue
        expected=sum((x+dx,y+dy,z+dz) not in cells for x,y,z in cells for dx,dy,dz in offsets)
        b=boundary(s);assert Decimal(b['total_area'])==expected
        assert len({f['face_id'] for f in b['faces']})==len(b['faces'])
        r=calculate(payload(s,included=lambda f:True),s)
        with localcontext() as ctx:
            ctx.prec=80
            assert ast_value(r['formula_ast'],r['dimensions'])==expected


@pytest.mark.parametrize('case',['missing','duplicate','unknown','non_bool','extra','blank_basis','blank_evidence','bad_purpose','bad_unit','bad_parameters'])
def test_invalid_face_decisions_fail(case):
    s=geometry();p=payload(s)
    if case=='missing':p['face_decisions'].pop()
    elif case=='duplicate':p['face_decisions'][1]=p['face_decisions'][0]
    elif case=='unknown':p['face_decisions'][0]['face_id']='invented'
    elif case=='non_bool':p['face_decisions'][0]['included']='false'
    elif case=='extra':p['face_decisions'][0]['area']='999'
    elif case=='blank_basis':p['face_decisions'][0]['basis']=' '
    elif case=='blank_evidence':p['face_decisions'][0]['evidence_id']=''
    elif case=='bad_purpose':p['purpose']='automatic'
    elif case=='bad_unit':p['quantity_type']='volume'
    else:p['parameters']={'area':'999'}
    with pytest.raises(DomainError):calculate(p,s)


def test_empty_solid_has_no_selectable_boundary():
    s=geometry();s['regions']=[region('a',0,0,0,1,1,1),region('cut',0,0,0,1,1,1,'subtract')]
    with pytest.raises(DomainError):boundary(s)


def test_candidate_preview_save_review_and_all_exports(clients):
    e,r,pid,source,body,eid=setup(clients);base='/api/v1/projects/'+pid;before=get(e,pid)
    candidates=e.get(base+'/surface-face-candidates',params={'component_id':source['id']});assert candidates.status_code==200,candidates.text
    assert Decimal(candidates.json()['total_area'])==61 and get(e,pid)==before
    preview=e.post(base+'/components/preview',json=body);assert preview.status_code==200,preview.text
    assert Decimal(preview.json()['quantity'])==37 and get(e,pid)==before
    response=e.post(base+'/components',json={**body,'source_geometry':{'malicious':'ignored'}});assert response.status_code==200,response.text
    comp=response.json();assert comp['source_geometry']==snapshot(source)
    assert get(e,pid)['components'][0]['state']=='approved'  # derived-face editing preserves its unchanged source approval
    assert e.post(base+'/takeoffs/calculate').status_code==409
    approve(r,pid,'components',comp)
    books=e.post(base+'/takeoffs/calculate');assert books.status_code==200,books.text
    for book in books.json():approve(r,pid,'takeoffs',book)
    boq=e.post(base+'/boq/generate').json();assert {b['unit']:Decimal(b['quantity']) for b in boq}=={'m³':Decimal(21),'m²':Decimal(37)}
    exported=e.get(base+'/exports?unpriced=true&format=json').json()
    assert exported['unpriced']['takeoffs'][1]['surface_faces']==preview.json()['surface_faces']
    x=e.get(base+'/exports?unpriced=true&format=xlsx');assert x.status_code==200,x.text
    wb=load_workbook(io.BytesIO(x.content),data_only=False)
    assert {'表面计量汇总','逐面计量决定','表面几何来源'}<=set(wb.sheetnames)
    assert all(c.data_type!='f' for ws in wb for row in ws for c in row)
    pdf=e.get(base+'/exports?unpriced=true&format=pdf');assert pdf.status_code==200,pdf.text
    texts='\n'.join(p.extract_text() or '' for p in PdfReader(io.BytesIO(pdf.content)).pages)
    assert '逐面计量决定' in texts and '表面几何来源' in texts


def test_revision_atomicity_current_source_and_approval_staleness(clients):
    e,r,pid,source,body,eid=setup(clients);base='/api/v1/projects/'+pid
    comp=e.post(base+'/components',json=body).json();approve(r,pid,'components',comp)
    for t in e.post(base+'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
    before=get(e,pid);assert e.put(base+'/components/'+comp['id'],json=body).status_code==409
    body['object_hash']=e.get(base+'/hashes').json()['components'][comp['id']]
    bad=copy.deepcopy(body);bad['face_decisions'].pop()
    assert e.put(base+'/components/'+comp['id'],json=bad).status_code==422 and get(e,pid)==before
    body['purpose']='waterproof';body['name']='合成防水表面';revised=e.put(base+'/components/'+comp['id'],json=body)
    assert revised.status_code==200,revised.text
    assert revised.json()['revision']==2 and get(e,pid)['components'][0]['state']=='approved'
    assert all(t['state']=='stale' for t in get(e,pid)['takeoffs'])
    source['object_hash']=e.get(base+'/hashes').json()['components'][source['id']];source['regions'][0]['parameters']['height']='3'
    response=e.put(base+'/components/'+source['id'],json=source);assert response.status_code==200,response.text
    approve(r,pid,'components',response.json())
    assert e.post(base+'/components/preview',json=body).status_code==409
    approve(r,pid,'components',revised.json(),expected=409)


@pytest.mark.parametrize('gate',['save','component_review','calculate','book_review'])
def test_corrupt_original_blocks_all_gates(clients,gate):
    e,r,pid,source,body,eid=setup(clients);base='/api/v1/projects/'+pid
    if gate!='save':comp=e.post(base+'/components',json=body).json()
    if gate in ('calculate','book_review'):approve(r,pid,'components',comp)
    if gate=='book_review':book=e.post(base+'/takeoffs/calculate').json()[1]
    doc=get(e,pid)['documents'][0];(store.ROOT/'originals'/doc['hash']).write_bytes(b'corrupt')
    if gate=='save':assert e.post(base+'/components',json=body).status_code==409
    elif gate=='component_review':approve(r,pid,'components',comp,expected=409)
    elif gate=='calculate':assert e.post(base+'/takeoffs/calculate').status_code==409
    else:approve(r,pid,'takeoffs',book,expected=409)


def test_permissions_foreign_source_and_face_proofs(clients):
    e,r,pid,source,body,eid=setup(clients);base='/api/v1/projects/'+pid
    assert r.post(base+'/components',json=body).status_code==403
    assert clients['other'].get(base+'/surface-face-candidates',params={'component_id':source['id']}).status_code==404
    second=create(e)['id'];upload(e,second,'其他工程');foreign=get(e,second)['evidence'][0]['id']
    bad=copy.deepcopy(body);bad['face_decisions'][0]['evidence_id']=foreign
    assert e.post(base+'/components',json=bad).status_code==404
    assert e.get('/api/v1/projects/'+second+'/surface-face-candidates',params={'component_id':source['id']}).status_code==404
    approve(r,pid,'components',source,decision='reject')
    assert e.post(base+'/components/preview',json=body).status_code==409


def test_snapshot_and_book_tampering_and_type_downgrade_blocked(clients):
    e,r,pid,source,body,eid=setup(clients);base='/api/v1/projects/'+pid
    comp=e.post(base+'/components',json=body).json();approve(r,pid,'components',comp)
    books=e.post(base+'/takeoffs/calculate').json()
    with store.transaction() as c:
        p=json.loads(c.execute('SELECT data FROM projects WHERE id=?',(pid,)).fetchone()['data']);p['takeoffs'][1]['surface_summary']['included_area']='999';store.save(c,p)
    approve(r,pid,'takeoffs',books[1],expected=409)
    changed={'kind':'area','parameters':{},'object_hash':e.get(base+'/hashes').json()['components'][comp['id']]}
    assert e.put(base+'/components/'+comp['id'],json=changed).status_code==409
    with store.transaction() as c:
        p=json.loads(c.execute('SELECT data FROM projects WHERE id=?',(pid,)).fetchone()['data']);p['components'][1]['source_geometry']['regions'][0]['parameters']['height']='99';store.save(c,p)
    approve(r,pid,'components',comp,expected=409)


def test_new_surface_never_rewrites_historical_quote(clients):
    e,r,pid,_,quote=priced_project(clients);base='/api/v1/projects/'+pid;old=copy.deepcopy(quote);eid=get(e,pid)['evidence'][0]['id']
    s=geometry(eid);s.pop('id');s.pop('revision');source=e.post(base+'/components',json=s).json();approve(r,pid,'components',source)
    source=next(x for x in get(e,pid)['components'] if x['id']==source['id'])
    response=e.post(base+'/components',json=payload(source,eid));assert response.status_code==200,response.text
    assert get(e,pid)['quotes'][0]==old and quote['id'] in get(e,pid)['stale_quotes']


def test_decimal_boundary_at_large_origin_retains_small_dimensions():
    s=geometry();s['regions']=[region('thin','1e12',0,0,'1e-18',1,1)]
    with localcontext() as ctx:
        ctx.prec=80
        b=boundary(s);assert Decimal(b['total_area'])==Decimal('2.000000000000000004')
        r=calculate(payload(s,included=lambda f:True),s)
        assert ast_value(r['formula_ast'],r['dimensions'])==Decimal('2.000000000000000004')


def test_multiple_derived_purposes_keep_geometry_and_each_other_approval(clients):
    e,r,pid,source,body,eid=setup(clients);base='/api/v1/projects/'+pid
    first=e.post(base+'/components',json=body).json();approve(r,pid,'components',first)
    second_body=copy.deepcopy(body);second_body.update(name='合成防水',purpose='waterproof')
    second=e.post(base+'/components',json=second_body);assert second.status_code==200,second.text
    approve(r,pid,'components',second.json())
    assert all(c['state']=='approved' for c in get(e,pid)['components'])
    body['object_hash']=e.get(base+'/hashes').json()['components'][first['id']];body['measurement_basis']='合成新口径，仍需逐面核对'
    revised=e.put(base+'/components/'+first['id'],json=body);assert revised.status_code==200,revised.text
    states={c['id']:c['state'] for c in get(e,pid)['components']}
    assert states[source['id']]=='approved' and states[second.json()['id']]=='approved' and states[first['id']]=='draft'


def test_superseded_source_and_unapproved_candidates_are_blocked(clients):
    e,r,pid,source,body,eid=setup(clients);base='/api/v1/projects/'+pid
    comp=e.post(base+'/components',json=body).json()
    upload(e,pid,'新版本来源及做法',name='fixture.txt')
    assert e.get(base+'/surface-face-candidates',params={'component_id':source['id']}).status_code==409
    assert e.post(base+'/components',json=body).status_code==409
    approve(r,pid,'components',comp,expected=409)
