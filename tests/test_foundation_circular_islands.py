"""Synthetic circular islands: explicit hierarchy, independent integrals and gates."""
import copy,io,itertools,random
from decimal import Decimal,localcontext
from fractions import Fraction
import pytest
from openpyxl import load_workbook
from pypdf import PdfReader
from backend import store
from backend.engine import DomainError
from backend.compound_takeoff import PI
from backend.foundation_horizontal_formwork import calculate
from backend.foundation_circular_islands import VERSION,METHOD,validate,separation,part_config,topology
from tests.test_foundation_circular_holes_windows import inner,config as hc,replay as replay_circles,setup as old_setup
from tests.test_foundation_circular_windows import body as cb,config as cc,independent_area
from tests.test_foundation_eccentric_windows import config as ec
from tests.test_foundation_annular_windows import config as ac
from tests.test_foundation_horizontal_windows import source,window
from tests.test_foundation_hole_end_formwork import source as end_source
from tests.test_foundation_multi_windows import replay_blocks,config as mc,part
from tests.test_foundation_nested_windows import config as nc,block as poly_block,box
from tests.test_foundation_perforated_windows import config as pc
from tests.test_foundation_openings import hole
from tests.test_foundation_voids import rectangle
from tests.test_foundation_recesses import interval
from tests.test_post_foundation import equal_body,fixed_body,fixed_added,equal_added,fixed_adjusted,equal_adjusted,combined
from tests.test_polygon_takeoff import post,mutate
from tests.test_workflow import get,approve,upload,create,priced_project
from tests.test_route_takeoff import evaluate


def block(eid='test',x=0,y=0,diameter=200,name='合成根块',holes=None,parent=None):return inner(eid,x,y,diameter,name)|dict(measurement_basis='合成本块实际支模及最近父关系说明，仅软件验证',holes=holes or [],parent=parent)
def parent(b=1,h=1):return dict(block_index=b,hole_index=h)
def config(eid='test',parts=None,unit='mm'):
    c=hc(eid,unit=unit);return {k:v for k,v in c.items() if k not in ('parameters','evidence','holes')}|dict(shape='nested_circular',parts=parts if parts is not None else [block(eid,holes=[inner(eid,10,0,100,'父内圆')]),block(eid,20,0,40,'合成岛区',parent=parent())])
def body(s,eid='test',face='bottom',parts=None):
    b=cb(s,eid,face);b['contact_ranges'][0]['plan_range']=config(eid,parts);return b
def wrapped(c):return dict(plan_range=c)


def replay(r):
    replay_blocks(r)
    for d in r['contact_ranges']:
        for q in d['window_part_checks']:
            for k,n in q['formula_asts'].items():assert Fraction(evaluate(n,r['dimensions']))==Fraction(q[k])
        child={**r,'contact_ranges':[{**d,**p} for p in d['window_parts']]};replay_circles(child)


@pytest.mark.parametrize('shape',['rectangular','circular','stepped_rectangular','stepped_circular'])
@pytest.mark.parametrize('face',['top','bottom','hole_top','hole_bottom'])
@pytest.mark.parametrize('make',[equal_body,fixed_body,fixed_added,equal_added,fixed_adjusted,equal_adjusted,combined,lambda s:combined(s,equal=True)])
def test_shapes_ends_counts_hierarchy_immutable_and_full_ast_replay(shape,face,make):
    s=end_source(shape,make) if face.startswith('hole') else source(shape,make);b=body(s,face=face);before=copy.deepcopy((s,b));r=calculate(b,s);d=r['contact_ranges'][0]
    assert (s,b)==before and r['rule_version']==VERSION and d['nested_circular_window'] and not d.get('multi_polygon_window')
    assert d['window_nesting']==[dict(block_index=1,parent=None,depth=0),dict(block_index=2,parent=parent(),depth=1)]
    with localcontext() as ctx:ctx.prec=80;assert abs(Decimal(d['window_area'])-Decimal(PI)*Decimal('.0079'))<Decimal('1e-78')
    assert r['source_evidence']['horizontal_formwork.1.parts.2.basis']=='test';replay(r)


@pytest.mark.parametrize('shape,face',[(s,f) for s in ['rectangular','circular','stepped_rectangular','stepped_circular'] for f in ['top','bottom','hole_bottom','hole_top']]+[('stepped_rectangular','step'),('stepped_circular','step')])
def test_parent_disk_minus_hole_plus_island_independent_gross_cut_net(shape,face):
    s=end_source(shape) if face.startswith('hole') else source(shape);b=body(s,face=face);r=calculate(b,s);children=[]
    for x,y,d in [(0,0,200),(10,0,100),(20,0,40)]:children.append(calculate(cb(s,face=face,x=x,y=y,diameter=d),s)['contact_ranges'][0])
    with localcontext() as ctx:
        ctx.prec=80
        for k in ['single_gross_area','single_opening_area','single_net_area','total_net_area']:
            assert abs(Decimal(r['contact_ranges'][0][k])-(Decimal(children[0][k])-Decimal(children[1][k])+Decimal(children[2][k])))<Decimal('1e-73')
    replay(r)


@pytest.mark.parametrize('shape',['rectangular','circular','stepped_rectangular','stepped_circular'])
@pytest.mark.parametrize('x',[-300,-140,0,140,300])
def test_20_independent_numerical_integrals_with_physical_hole(shape,x):
    s=source(shape);s['foundation_items'][0]['through_holes']['holes']=[{**rectangle(x='150',length='100',width='60'),'vertical_range':interval('0','150')}]
    ps=[block(x=x,y=30,diameter=260,holes=[inner('test',x+15,20,140,'W1')]),block(x=x+25,y=25,diameter=60,name='岛',parent=parent())];r=calculate(body(s,parts=ps),s);ds=[];bs=[]
    if shape in ('rectangular','stepped_rectangular'):bs=[(-.2,.2,-.25,.25)] if shape=='rectangular' else [(-.3,.3,-.25,.25)]
    else:ds=[(0.,0.,.3)]
    a=independent_area(ds+[(x/1000,.03,.13)],bs,excluded_disks=[((x+15)/1000,.02,.07)],excluded_boxes=[(.1,.2,-.03,.03)])
    b=independent_area(ds+[((x+25)/1000,.025,.03)],bs,excluded_boxes=[(.1,.2,-.03,.03)])
    assert abs(float(r['contact_ranges'][0]['single_net_area'])-a-b)<3e-10;replay(r)


def chain(eid='test'):
    return [block(eid,0,0,200,'B1',[inner(eid,0,0,160,'W1')]),block(eid,0,0,120,'B2',[inner(eid,0,0,80,'W2')],parent()),block(eid,0,0,60,'B3',[inner(eid,0,0,40,'W3')],parent(2)),block(eid,0,0,20,'B4',parent=parent(3))]


def test_four_levels_all_24_orderings_and_immediate_parent_remapping():
    s=source();base=chain();expected=None
    for order in itertools.permutations(range(4)):
        ps=copy.deepcopy([base[i] for i in order]);mapping={old+1:new+1 for new,old in enumerate(order)}
        for p in ps:
            if p['parent']:p['parent']['block_index']=mapping[p['parent']['block_index']]
        b=body(s,parts=ps);r=calculate(b,s);d=r['contact_ranges'][0];assert [h['depth'] for h in d['window_nesting']]==list(order)
        if expected is None:expected=r['quantity']
        assert r['quantity']==expected;replay(r)


@pytest.mark.parametrize('gap',['-0.000000001','0','0.000000001'])
def test_strict_parent_hole_tangency_no_tolerance(gap):
    c=config();c['parts'][1]['parameters']['x']=str(Decimal(40)+Decimal(gap))
    if Decimal(gap)<0:validate(c)
    else:
        with pytest.raises(DomainError):validate(c)


@pytest.mark.parametrize('fault',['missing_parent','wrong_hole','root','self','bool','cycle','skip','no_island','touch_roots','overlap_roots','unit','one','five','name','missing_evidence','basis','extra','hole_schema','five_holes','nine_holes'])
def test_explicit_hierarchy_strict_schema_and_limits(fault):
    c=config();p=c['parts'][1]
    if fault=='missing_parent':del p['parent']
    elif fault=='wrong_hole':p['parent']['hole_index']=2
    elif fault=='root':p['parent']=None
    elif fault=='self':p['parent']=parent(2)
    elif fault=='bool':p['parent']['block_index']=True
    elif fault=='cycle':c['parts'][0]['parent']=parent(2);p['holes']=[inner('test',20,0,10,'内孔')]
    elif fault=='skip':c['parts']=chain();c['parts'][2]['parent']=parent(1)
    elif fault=='no_island':c['parts'][1]=block(x=300,name='分离块')
    elif fault in ('touch_roots','overlap_roots'):c['parts'].append(block(x=200 if fault=='touch_roots' else 199,name='相邻根块'))
    elif fault=='unit':p['unit']='mm'
    elif fault=='one':c['parts']=c['parts'][:1]
    elif fault=='five':c['parts']*=3
    elif fault=='name':p['name']=c['parts'][0]['name']
    elif fault=='missing_evidence':p['evidence']['x']=''
    elif fault=='basis':p['measurement_basis']=''
    elif fault=='extra':c['count']='2'
    elif fault=='hole_schema':c['parts'][0]['holes'][0]['area']='1'
    elif fault=='five_holes':c['parts'][0]['holes']*=5
    else:
        c['parts']=chain()
        for p in c['parts'][:3]:p['holes']=[inner('test',x,y,10,str(i)) for i,(x,y) in enumerate([(-10,-10),(10,-10),(10,10)])]
    with pytest.raises(DomainError):validate(c)


@pytest.mark.parametrize('unit,factor',[('mm','1'),('cm','.1'),('dm','.01'),('m','.001'),('km','.000001')])
def test_all_circle_units_and_low_precision(unit,factor):
    s=source();b=body(s);c=b['contact_ranges'][0]['plan_range'];c['unit']=unit
    for p in c['parts']:
        for h in [p]+p['holes']:h['parameters']={k:str(Decimal(v)*Decimal(factor)) for k,v in h['parameters'].items()}
    with localcontext() as ctx:ctx.prec=6;r=calculate(b,s)
    assert r['quantity']==calculate(body(s),s)['quantity'];replay(r)


def test_multiple_roots_and_two_children_in_different_holes_strictly_separate():
    ps=[block(diameter=300,holes=[inner('test',-60,0,80,'左W'),inner('test',60,0,80,'右W')]),block(x=-60,diameter=30,name='左岛',parent=parent()),block(x=60,diameter=30,name='右岛',parent=parent(1,2)),block(x=400,diameter=100,name='分离根')]
    c=config(parts=ps);validate(c);r=calculate(body(source(),parts=ps),source());assert [h['depth'] for h in r['contact_ranges'][0]['window_nesting']]==[0,1,1,0];replay(r)


def test_complete_net_pair_checks_islands_and_unused_hole_space_with_mixed_shapes():
    a=wrapped(config());cases=[(cc(x=20,y=0,diameter=40),True),(cc(x=-20,y=0,diameter=10),False),(ac(x=10,y=0,outer=100,inner=80),False),(ec(x=0,y=0,outer=300,inner_x=0,inner_y=0,inner=220),False),(hc(x=0,y=0,diameter=200),True),(window('test','10','30','-5','5'),True),(window('test','-25','-15','-5','5'),False),(mc([part(box(-25,-15,-5,5),name='孔空区'),part(box(150,160,-5,5),name='外区')]),False),(pc(outer=box(-150,150,-150,150),inners=[box(-110,110,-110,110)]),False),(nc([poly_block(box(-150,150,-150,150),holes=[box(-110,110,-110,110)]),poly_block(box(15,25,-5,5),name='内岛',parent=parent())]),True)]
    for c,overlap in cases:
        for left,right in [(a,wrapped(c)),(wrapped(c),a)]:
            q=separation(left,right);assert q['overlap']==overlap
            if not overlap:assert q['method']==METHOD and len(q['pairs'])==2
    b=wrapped(config(parts=[block(x=300,name='远根',holes=[inner('test',300,0,100,'父W')]),block(x=310,diameter=40,name='远岛',parent=parent())]));q=separation(a,b);assert not q['overlap'] and len(q['pairs'])==4


def test_80_seeded_disk_pairs_against_independent_boolean_integral():
    rng=random.Random(11601);a=wrapped(config())
    for _ in range(80):
        x,y,d=rng.randint(-150,150),rng.randint(-150,150),rng.randint(2,80);c=cc(x=x,y=y,diameter=d)
        base=[(x/1000,y/1000,d/2000)];area=independent_area(base+[(0.,0.,.1)],excluded_disks=[(.01,0.,.05)])+independent_area(base+[(.02,0.,.02)])
        assert separation(a,wrapped(c))['overlap']==(area>1e-11)


@pytest.mark.parametrize('gap',['0','0.000000001'])
def test_every_root_and_island_candidate_covered_or_tiny_positive_gap(gap):
    v=str(Decimal(25)+Decimal(gap));hs=[{**hole(x=0,diameter=80),'vertical_range':interval('100','300')},{**rectangle(name='H2',x='-'+v,length=50,width=100),'vertical_range':interval('300','500')},{**rectangle(name='H3',x=v,length=50,width=100),'vertical_range':interval('300','500')}]
    s=end_source(holes=hs);ps=[block(diameter=60,holes=[inner('test',0,0,40,'W')]),block(diameter=20,name='岛',parent=parent())];r=calculate(body(s,face='hole_top',parts=ps),s);d=r['contact_ranges'][0];assert ('polygon_zero_contact' in d)==(gap=='0') and (Decimal(d['single_net_area'])==0)==(gap=='0');replay(r)
    if gap=='0':assert {q['block_index'] for q in d['polygon_zero_contact']['net_trace']['partitions']}=={1,2}


def setup(clients,pid=None,zero=False):
    e,r,pid,s,b,eid=old_setup(clients,pid,zero)
    ps=[block(eid,0,0,40,holes=[inner(eid,0,0,26,'父排除W')]),block(eid,2,0,12,'圆岛',parent=parent())] if zero else [block(eid,196,8,42,holes=[inner(eid,201,10,16,'父排除W')]),block(eid,201,10,10,'圆岛',parent=parent())]
    b['contact_ranges'][0]['plan_range']=config(eid,ps);return e,r,pid,s,b,eid


def test_roles_foreign_island_original_history_concurrency_and_source_stale(clients):
    e,r,pid,s,b,eid=setup(clients);before=get(e,pid);post(e,pid,b,'/components/preview');assert get(e,pid)==before;post(r,pid,b,expected=403)
    foreign=create(e)['id'];upload(e,foreign,'异工程合成');bad=copy.deepcopy(b);bad['contact_ranges'][0]['plan_range']['parts'][1]['evidence_id']=get(e,foreign)['evidence'][-1]['id'];post(e,pid,bad,expected=404)
    c=post(e,pid,b).json();approve(r,pid,'components',c);c=next(x for x in get(e,pid)['components'] if x['id']==c['id']);b['object_hash']=store.digest(c);b['contact_ranges'][0]['plan_range']['parts'][1]['parameters']['y']='11';res=e.put(f'/api/v1/projects/{pid}/components/'+c['id'],json=b);assert res.status_code==200,res.text;assert e.put(f'/api/v1/projects/{pid}/components/'+c['id'],json=b).status_code==409;assert get(e,pid)['component_history'][-1]==c;approve(r,pid,'components',res.json())
    s['object_hash']=store.digest(s);s['foundation_items'][0]['levels'][1]['parameters']['diameter']='420';res=e.put(f'/api/v1/projects/{pid}/components/'+s['id'],json=s);assert res.status_code==200;approve(r,pid,'components',res.json());assert next(x for x in get(e,pid)['components'] if x['id']==c['id'])['state']=='stale'


def test_unused_parent_hole_claim_allowed_but_island_overlap_or_frame_refused(clients):
    e,r,pid,s,b,eid=setup(clients);post(e,pid,b);other=copy.deepcopy(b);other['name']='父孔剩余环带另项';other['contact_ranges'][0]['plan_range']=ec(eid,x=201,y=10,outer=16,inner_x=201,inner_y=10,inner=10);post(e,pid,other)
    bad=copy.deepcopy(other);bad['contact_ranges'][0]['plan_range']['parameters']['inner_x']='201.000000001';before=get(e,pid);post(e,pid,bad,expected=409);assert get(e,pid)==before
    bad=copy.deepcopy(other);bad['contact_ranges'][0]['plan_range']['coordinate_basis']='其他基准';post(e,pid,bad,expected=409)


@pytest.mark.parametrize('fault',['input_parent','island_source','hierarchy','topology','topology_ast','island_area','zero','expression','hole_area'])
def test_saved_parent_or_geometry_tamper_rejected(clients,fault):
    e,r,pid,s,b,eid=setup(clients,zero=True);c=post(e,pid,b).json()
    def forge(p):
        c=p['components'][-1];d=c['preview']['contact_ranges'][0]
        if fault=='input_parent':c['contact_ranges'][0]['plan_range']['parts'][1]['parent']=None
        elif fault=='island_source':c['contact_ranges'][0]['plan_range']['parts'][1]['evidence_id']='fake'
        elif fault=='hierarchy':d['window_nesting'][1]['depth']=0
        elif fault=='topology':d['window_part_checks'][0]['limit_squared']='999'
        elif fault=='topology_ast':d['window_part_checks'][0]['formula_asts']['limit_squared']={'constant':'0'}
        elif fault=='island_area':d['window_parts'][1]['single_net_area']='1'
        elif fault=='zero':d['polygon_zero_contact']['net_trace']['partitions'][0]['remaining']=[{}]
        elif fault=='expression':d['window_expression_definitions'][0]['value']='999'
        else:d['window_parts'][0]['window_circular_holes'][0]['area']='999'
    mutate(pid,forge);before=get(e,pid);approve(r,pid,'components',c,expected=422 if fault=='input_parent' else 404 if fault=='island_source' else 409);assert get(e,pid)==before


@pytest.mark.parametrize('gate',['preview','save','review','calculate','book_review'])
def test_island_parent_relation_original_bytes_five_gates(clients,gate):
    e,r,pid,s,b,eid=setup(clients);upload(e,pid,'仅合成：圆岛须在最近父块W内并实际支模','island-relation.txt')
    for c in get(e,pid)['components']:approve(r,pid,'components',c)
    b['source_component_hash']=e.get(f'/api/v1/projects/{pid}/foundation-horizontal-formwork-candidates',params={'component_id':s['id']}).json()['source_component_hash'];b['contact_ranges'][0]['plan_range']['parts'][1]['evidence_id']=get(e,pid)['evidence'][-1]['id']
    if gate in ('review','calculate','book_review'):c=post(e,pid,b).json()
    if gate in ('calculate','book_review'):approve(r,pid,'components',c)
    if gate=='book_review':t=post(e,pid,{},'/takeoffs/calculate').json()[-1]
    doc=get(e,pid)['documents'][-1];(store.ROOT/'originals'/doc['hash']).write_bytes(b'corrupted-synthetic-island-original')
    if gate=='preview':post(e,pid,b,'/components/preview',409)
    elif gate=='save':post(e,pid,b,expected=409)
    elif gate=='review':approve(r,pid,'components',c,expected=409)
    elif gate=='calculate':post(e,pid,{},'/takeoffs/calculate',409)
    else:approve(r,pid,'takeoffs',t,expected=409)


@pytest.mark.parametrize('fmt',['json','xlsx','pdf'])
@pytest.mark.parametrize('zero',[False,True])
def test_three_formats_parent_and_source_records_old_quote_isolation(clients,fmt,zero):
    e,r,pid,_,q=priced_project(clients);e,r,pid,s,b,eid=setup(clients,pid,zero);c=post(e,pid,b).json();approve(r,pid,'components',c)
    for t in post(e,pid,{},'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
    post(e,pid,{},'/boq/generate');url=f'/api/v1/projects/{pid}/exports';res=e.get(url,params=dict(format=fmt,unpriced=True));assert res.status_code==200,res.text
    titles={'基础圆岛净块与父关系','基础圆岛严格父子拓扑','基础圆岛窗口参数','基础圆岛严格拓扑','基础圆岛分段积分'}|({'基础圆岛整体零接触','基础圆岛零接触校验'} if zero else set())
    if fmt=='json':assert titles<={t[0] for t in res.json()['tables']}
    elif fmt=='xlsx':assert titles<=set(load_workbook(io.BytesIO(res.content)).sheetnames)
    else:assert all(any(title in p.extract_text() for p in PdfReader(io.BytesIO(res.content)).pages) for title in titles)
    assert not any(t[0] in titles for t in e.get(url,params=dict(format='json',quote_id=q['id'])).json()['tables'])


@pytest.mark.parametrize('fault',['block_index','hole_index','distance'])
def test_full_window_block_pair_certificate_tamper(clients,fault):
    e,r,pid,s,b,eid=setup(clients);d=copy.deepcopy(b['contact_ranges'][0]);d.update(name='父孔剩余',plan_range=ec(eid,x=201,y=10,outer=16,inner_x=201,inner_y=10,inner=10));b['contact_ranges'].append(d);c=post(e,pid,b).json()
    def forge(p):
        q=p['components'][-1]['preview']['horizontal_window_checks'][0]['pairs'][0]
        if fault=='block_index':q['left_block_index']=2
        elif fault=='hole_index':q['proof']['hole_index']=2
        else:q['proof']['distance_squared']='999'
    mutate(pid,forge);before=get(e,pid);approve(r,pid,'components',c,expected=409);assert get(e,pid)==before


def test_repeated_w_indices_keep_distinct_block_identifiers_in_report_tables():
    from backend.foundation_circular_island_reports import tables
    s=source();b=body(s,parts=chain());r=calculate(b,s);ts=tables({'takeoffs':[{**b,**r}]});strict=next(rows for title,_,rows in ts if title=='基础圆岛严格拓扑')
    for i in [1,2,3]:assert any(' / B'+str(i)+' B'+str(i)+' / C1 / W1' in row[0] for row in strict)


def test_four_blocks_eight_internal_circles_limit_is_usable():
    ps=[]
    for i,x in enumerate([0,400]):
        hs=[inner('test',x+dx,dy,40,str(j)) for j,(dx,dy) in enumerate([(-40,-40),(40,-40),(40,40),(-40,40)])]
        ps.append(block(x=x,diameter=200,name='root'+str(i),holes=hs))
    ps.extend([block(x=-40,y=-40,diameter=20,name='island0',parent=parent(1)),block(x=360,y=-40,diameter=20,name='island1',parent=parent(2))]);validate(config(parts=ps));s=source();r=calculate(body(s,parts=ps),s);replay(r)
