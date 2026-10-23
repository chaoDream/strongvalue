"""Circular foundation inputs, independently computed quantities and lifecycle."""
import copy,io,random
from decimal import Decimal,localcontext,ROUND_HALF_UP
import pytest
from openpyxl import load_workbook
from pypdf import PdfReader
from backend import store
from backend.compound_takeoff import PI
from backend.engine import DomainError
from backend.post_foundation import calculate
from backend.railing_material import calculate as material_quantity
from tests.test_foundation_steps import stepped as rect_steps
from tests.test_post_foundation import body as rectangular_body,source,setup as rectangular_setup,row as rect_row,equal_body,fixed_body,fixed_added,equal_added,fixed_adjusted,equal_adjusted,combined
from tests.test_polygon_takeoff import post,mutate
from tests.test_route_takeoff import evaluate
from tests.test_workflow import get,approve,upload,create,priced_project

ROUND_TABLE='圆形基础圆周率与口径'
SHAPES=('circular','stepped_circular')


def level(diameter='600',height='200',eid='test'):
 return dict(parameters=dict(diameter=diameter,height=height),evidence={k:eid for k in ('diameter','height')})


def row(eid='test',shape='stepped_circular',index=1):
 b=dict(material_index=index,shape=shape,length_unit='mm',measurement_basis='合成每根已审立柱一个竖直实心圆形基础，填写实际直径而非半径，各层同心同轴自下而上紧接无空隙、不外伸，层高为本层净高，层数不乘基础个数；无孔洞及交叠，非真实工程',evidence=dict(measurement_basis=eid))
 if shape=='circular':b.update(parameters=dict(diameter='600',height='600'));b['evidence'].update({k:eid for k in ('diameter','height')})
 else:b.update(levels=[level(eid=eid),level('400','400',eid)],stacking_confirmed=True)
 return b


def body(s,eid='test',shape='stepped_circular'):
 b=rectangular_body(s,eid);b.update(name='合成圆形立柱基础（非真实工程）',material='合成C30实心圆形基础，同项同尺寸、各层同一材料，非真实工程',foundation_items=[row(eid,shape)],measurement_basis='合成一柱一实心圆柱或圆台阶基础，实际直径逐项据原件核定，各层同心同轴垂直接触无空隙、不外伸，层高非累计，来源重复一次且未与既有清单重复，不含孔洞扣减或土方钢筋模板');return b


def setup(clients,pid=None,shape='stepped_circular'):
 e,r,pid,m,_,old_eid=rectangular_setup(clients,pid)
 upload(e,pid,'合成圆形基础软件验收，非真实工程。已审最终立柱单组11根重复2共22个基础，原材料79.2kg保持。竖直实心圆柱实际直径600mm、净高600mm，半径为300mm不可当直径，π/4×.6²×.6×22=3.732212072464674m3，修约3位3.732。圆形台阶自下而上同心同轴紧接无空隙、不外伸，下层实际直径600mm净层高200mm、上层实际直径400mm净层高400mm，相对底面Z=0、层界.2、顶.6m，每个.034πm3，总.748πm3=2.349911304885165m3，修约2.350。上层净高改500mm则顶.7m、总.836πm3=2.626371458401067m3，修约2.626，基础数仍22；上层直径600.000000001mm外伸须拒。各层同一合成C30，圆柱实心无洞，各基础独立不交叠，不重复既有基础清单，不含锚栓孔预埋、土方垫层钢筋模板及承载设计。TEST-ROUND是仅合成软件映射编码，非官方条目，不适用实际工程。','foundation-round-basis.txt')
 for c in get(e,pid)['components']:approve(r,pid,'components',c)
 m=get(e,pid)['components'][-1];eid=get(e,pid)['evidence'][-1]['id'];return e,r,pid,m,body(m,eid,shape),eid


@pytest.mark.parametrize('shape',SHAPES)
def test_counts_pi_full_formula_and_no_mutation(shape):
 s=source();b=body(s,shape=shape);before=copy.deepcopy((s,b));r=calculate(b,s);d=r['foundation_items'][0];assert (s,b)==before and r['quantity']==('3.732' if shape=='circular' else '2.350') and r['rule_version']=='railing-post-foundation-1.2' and r['foundation_summary']['total_count']=='22' and r['foundation_pi_constant']==PI and material_quantity(s)['quantity']=='79.200'
 with localcontext() as ctx:
  ctx.prec=80;assert Decimal(r['gross'])==Decimal(PI)*Decimal('1.188' if shape=='circular' else '.748');assert evaluate(r['formula_ast'],r['dimensions'])==Decimal(r['quantity'])
  for k,a in d['formula_asts'].items():assert evaluate(a,r['dimensions'])==Decimal(d[{'single_volume':'single_foundation_volume'}.get(k,k)])
  for v in d.get('foundation_levels',[]):
   for k,a in v['formula_asts'].items():assert evaluate(a,r['dimensions'])==Decimal(v[k])
 if shape=='stepped_circular':assert [(Decimal(v['from_z']),Decimal(v['to_z'])) for v in d['foundation_levels']]==[(Decimal(0),Decimal('.2')),(Decimal('.2'),Decimal('.6'))]
 else:assert 'foundation_levels' not in d and d['normalized_parameters']==dict(diameter='0.6',height='0.6')


@pytest.mark.parametrize('shape',SHAPES)
@pytest.mark.parametrize('make',[equal_body,fixed_body,fixed_added,equal_added,fixed_adjusted,equal_adjusted,combined,lambda s:combined(s,equal=True)])
def test_eight_final_post_modes(shape,make):
 s=source(make);r=calculate(body(s,shape=shape),s)
 with localcontext() as ctx:ctx.prec=80;assert Decimal(r['gross'])==Decimal(PI)*Decimal('.054' if shape=='circular' else '.034')*r['foundation_summary']['single_group_count']*2


@pytest.mark.parametrize('shape',SHAPES)
@pytest.mark.parametrize('unit,factor',[('m',Decimal('.001')),('mm',Decimal('1')),('cm',Decimal('.1')),('dm',Decimal('.01')),('km',Decimal('.000001'))])
def test_units_and_low_context(shape,unit,factor):
 s=source();b=body(s,shape=shape);d=b['foundation_items'][0];d['length_unit']=unit
 for v in d.get('levels',[d]):
  for k in v['parameters']:v['parameters'][k]=str(Decimal(v['parameters'][k])*factor)
 with localcontext() as ctx:ctx.prec=6;r=calculate(b,s)
 assert r==calculate(b,s) and r['quantity']==('3.732' if shape=='circular' else '2.350')


@pytest.mark.parametrize('diameter,accepted',[('600',True),('599.999999999',True),('600.000000001',False)])
def test_exact_non_overhang(diameter,accepted):
 s=source();b=body(s);b['foundation_items'][0]['levels'][1]['parameters']['diameter']=diameter
 if accepted:assert calculate(b,s)['foundation_summary']['total_count']=='22'
 else:
  with pytest.raises(DomainError,match='上层圆形台阶'):calculate(b,s)


def test_ten_layers_mixed_shapes_and_unselected_items():
 s=source();s['material_items']=[{**copy.deepcopy(s['material_items'][0]),'name':'合成柱'+str(i)} for i in range(5)];b=body(s);b['foundation_items'][0]['levels']=[level(height='100') for _ in range(10)];b['foundation_items']+=[row(shape='circular',index=2),rect_row(3),dict(rect_steps(index=4))];r=calculate(b,s);assert r['foundation_summary']['total_count']=='88' and r['foundation_summary']['unselected_post_item_count']==1 and Decimal(r['foundation_items'][0]['foundation_levels'][-1]['to_z'])==1;b['foundation_items'].reverse();assert calculate(b,s)==r
 with localcontext() as ctx:ctx.prec=80;assert Decimal(r['gross'])==Decimal(PI)*Decimal('3.168')+Decimal('5.016')


def test_independent_random_cylinder_layer_sums_and_rounding():
 rng=random.Random(8801)
 for shape in SHAPES:
  for _ in range(40):
   s=source();b=body(s,shape=shape);d=b['foundation_items'][0];diameter=1000;values=[];expected=Decimal(0)
   with localcontext() as ctx:
    ctx.prec=80
    for _ in range(rng.randint(2,10) if shape=='stepped_circular' else 1):
     diameter=rng.randint(1,diameter);height=rng.randint(1,999);values.append(level(str(diameter),str(height)));expected+=Decimal(PI)/4*Decimal(diameter)*diameter*height/10**9
    if shape=='circular':d['parameters']=values[0]['parameters']
    else:d['levels']=values
    digits=rng.randint(0,12);b['rounding_digits']=digits;r=calculate(b,s);assert Decimal(r['gross'])==expected*22 and Decimal(r['quantity'])==(expected*22).quantize(Decimal(1).scaleb(-digits),rounding=ROUND_HALF_UP)


@pytest.mark.parametrize('fault',['shape-null','shape-other','shape-list','missing-shape','extra-row','old-parameters','old-proofs','missing-row','unconfirmed','truthy','null-levels','dict-levels','empty','one','eleven','level-null','level-list','level-extra','level-missing','parameters-null','parameter-missing','parameter-extra','radius','proofs-null','proof-missing','proof-extra','proof-blank','proof-bool','negative','zero','bool','nan','infinite','precision','unit','overhang','reverse','scope-blank','scope-proof-blank','scope-proof-bool'])
def test_strict_tier_inputs(fault):
 s=source();b=body(s);d=b['foundation_items'][0];v=d['levels'][0]
 if fault=='shape-null':d['shape']=None
 elif fault=='shape-other':d['shape']='truncated_cone'
 elif fault=='shape-list':d['shape']=[]
 elif fault=='missing-shape':d.pop('shape')
 elif fault=='extra-row':d['quantity']='22'
 elif fault=='old-parameters':d['parameters']=dict(diameter='600',height='600')
 elif fault=='old-proofs':d['evidence']['diameter']='test'
 elif fault=='missing-row':d.pop('stacking_confirmed')
 elif fault=='unconfirmed':d['stacking_confirmed']=False
 elif fault=='truthy':d['stacking_confirmed']=1
 elif fault=='null-levels':d['levels']=None
 elif fault=='dict-levels':d['levels']={}
 elif fault=='empty':d['levels']=[]
 elif fault=='one':d['levels']=[v]
 elif fault=='eleven':d['levels']=[v]*11
 elif fault=='level-null':d['levels'][0]=None
 elif fault=='level-list':d['levels'][0]=[]
 elif fault=='level-extra':v['z']='0'
 elif fault=='level-missing':v.pop('evidence')
 elif fault=='parameters-null':v['parameters']=None
 elif fault=='parameter-missing':v['parameters'].pop('height')
 elif fault=='parameter-extra':v['parameters']['length']='600'
 elif fault=='radius':v['parameters']['radius']=v['parameters'].pop('diameter')
 elif fault=='proofs-null':v['evidence']=None
 elif fault=='proof-missing':v['evidence'].pop('height')
 elif fault=='proof-extra':v['evidence']['radius']='test'
 elif fault=='proof-blank':v['evidence']['height']=''
 elif fault=='proof-bool':v['evidence']['height']=True
 elif fault in ('negative','zero','bool','nan','infinite','precision'):v['parameters']['diameter']={'negative':'-1','zero':'0','bool':True,'nan':'NaN','infinite':'Infinity','precision':'.1234567890123456789'}[fault]
 elif fault=='unit':d['length_unit']='m³'
 elif fault=='overhang':d['levels'][1]['parameters']['diameter']='601'
 elif fault=='reverse':d['levels'].reverse()
 elif fault=='scope-blank':d['measurement_basis']=''
 elif fault=='scope-proof-blank':d['evidence']['measurement_basis']=''
 else:d['evidence']['measurement_basis']=True
 with pytest.raises(DomainError):calculate(b,s)


@pytest.mark.parametrize('fault',['radius','extra','levels','confirmation','missing-proof','bad-proof','parameter-null','parameter-extra','zero-height','negative-diameter','unit'])
def test_strict_cylinder_inputs(fault):
 s=source();b=body(s,shape='circular');d=b['foundation_items'][0]
 if fault=='radius':d['parameters']['radius']=d['parameters'].pop('diameter')
 elif fault=='extra':d['count']='22'
 elif fault=='levels':d['levels']=[level()]
 elif fault=='confirmation':d['stacking_confirmed']=True
 elif fault=='missing-proof':d['evidence'].pop('diameter')
 elif fault=='bad-proof':d['evidence']['diameter']=False
 elif fault=='parameter-null':d['parameters']=None
 elif fault=='parameter-extra':d['parameters']['inner_diameter']='1'
 elif fault=='zero-height':d['parameters']['height']='0'
 elif fault=='negative-diameter':d['parameters']['diameter']='-1'
 else:d['length_unit']='m²'
 with pytest.raises(DomainError):calculate(b,s)


@pytest.mark.parametrize('shape',SHAPES)
@pytest.mark.parametrize('fmt',['json','xlsx','pdf'])
def test_workflow_three_exports_old_quote_and_mapping(clients,shape,fmt):
 e,r,pid,_,quote=priced_project(clients);e,r,pid,s,b,eid=setup(clients,pid,shape);before=get(e,pid);assert post(e,pid,b,'/components/preview').json()['quantity']==('3.732' if shape=='circular' else '2.350') and get(e,pid)==before;c=post(e,pid,b).json();approve(r,pid,'components',c)
 for t in post(e,pid,{},'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
 boq=post(e,pid,{},'/boq/generate').json();assert any(v['quantity']==c['preview']['quantity'] and v['unit']=='m³' for v in boq) and any(v['quantity']=='79.200' for v in boq);url=f'/api/v1/projects/{pid}/exports';res=e.get(url,params=dict(format=fmt,unpriced=True));assert res.status_code==200,res.text;names={ROUND_TABLE}|({'台阶基础逐层体积','台阶基础逐层原件'} if shape=='stepped_circular' else set())
 if fmt=='json':assert names<={t[0] for t in res.json()['tables']} and PI in str(res.json()['tables'])
 elif fmt=='xlsx':
  w=load_workbook(io.BytesIO(res.content));assert names<=set(w.sheetnames);assert w['立柱基础逐项体积'].row_dimensions[3].height>48;assert w[ROUND_TABLE].max_row>3;assert PI in ''.join(str(row[1].value or '') for row in w[ROUND_TABLE].iter_rows(min_row=3))
  if shape=='stepped_circular':assert w['台阶基础逐层体积'].row_dimensions[3].height>48
 else:assert all(v in ''.join(p.extract_text() for p in PdfReader(io.BytesIO(res.content)).pages) for v in names)
 assert not any(t[0] in names for t in e.get(url,params=dict(format='json',quote_id=quote['id'])).json()['tables'])
 if fmt=='json':
  from tests.test_mapping_profiles import rule_body,save_rule,preview,selection
  p=get(e,pid);rule={**rule_body(p,'post_foundation','volume',eid),'code':'TEST-ROUND'};assert e.patch(f'/api/v1/projects/{pid}/settings',json=dict(boq_standard=rule['standard'],boq_edition=rule['edition'],reason='仅合成圆形基础口径验收')).status_code==200;save_rule(e,r,pid,rule);bid=next(v['id'] for v in p['boq'] if v['takeoff_id']==p['takeoffs'][-1]['id']);chosen=preview(e,pid,bid);assert chosen['state']=='unique',chosen;assert e.post(f'/api/v1/projects/{pid}/boq/apply-mappings',json=dict(selections=[selection(chosen)])).status_code==200;current=next(v for v in get(e,pid)['boq'] if v['id']==bid);assert current['quantity']==c['preview']['quantity'] and current['unit']=='m³';approve(r,pid,'boq',current)


@pytest.mark.parametrize('fault',['raw','proof','confirmation','shape','pi','z','volume','ast','count','book'])
def test_saved_record_tampering(clients,fault):
 e,r,pid,s,b,eid=setup(clients);c=post(e,pid,b).json()
 if fault=='book':approve(r,pid,'components',c);post(e,pid,{},'/takeoffs/calculate')
 def forge(p):
  c=p['components'][-1];d=c['preview']['foundation_items'][0]
  if fault=='raw':c['foundation_items'][0]['levels'][1]['parameters']['height']='500'
  elif fault=='proof':c['foundation_items'][0]['levels'][1]['evidence']['diameter']='forged'
  elif fault=='confirmation':c['foundation_items'][0]['stacking_confirmed']=False
  elif fault=='shape':d['shape']='circular'
  elif fault=='pi':c['preview']['foundation_pi_constant']='3.14'
  elif fault=='z':d['foundation_levels'][1]['from_z']='0'
  elif fault=='volume':d['foundation_levels'][1]['total_volume']='2'
  elif fault=='ast':d['foundation_levels'][1]['formula_asts']['single_level_volume']={'constant':'1'}
  elif fault=='count':d['total_count']='44'
  else:p['takeoffs'][-1]['foundation_pi_constant']='3.14'
 mutate(pid,forge);coll='takeoffs' if fault=='book' else 'components';approve(r,pid,coll,get(e,pid)[coll][-1],expected=404 if fault=='proof' else 422 if fault=='confirmation' else 409)


@pytest.mark.parametrize('shape',SHAPES)
@pytest.mark.parametrize('gate',['preview','save','review','calculate','book_review'])
def test_original_bytes_five_gates(clients,shape,gate):
 e,r,pid,s,b,eid=setup(clients,shape=shape)
 if gate in ('review','calculate','book_review'):c=post(e,pid,b).json()
 if gate in ('calculate','book_review'):approve(r,pid,'components',c)
 if gate=='book_review':t=post(e,pid,{},'/takeoffs/calculate').json()[-1]
 d=get(e,pid)['documents'][-1];(store.ROOT/'originals'/d['hash']).write_bytes(b'corrupt')
 if gate=='preview':post(e,pid,b,'/components/preview',409)
 elif gate=='save':post(e,pid,b,expected=409)
 elif gate=='review':approve(r,pid,'components',c,expected=409)
 elif gate=='calculate':post(e,pid,{},'/takeoffs/calculate',409)
 else:approve(r,pid,'takeoffs',t,expected=409)


@pytest.mark.parametrize('shape',SHAPES)
def test_roles_tenant_foreign_dimension_evidence(clients,shape):
 e,r,pid,s,b,eid=setup(clients,shape=shape);post(r,pid,b,expected=403);assert clients['other'].get(f'/api/v1/projects/{pid}').status_code==404;oid=create(e)['id'];upload(e,oid,'异项目尺寸原件');fid=get(e,oid)['evidence'][-1]['id'];d=b['foundation_items'][0];v=d['levels'][0] if d.get('levels') else d;v['evidence']['diameter']=fid;post(e,pid,b,expected=404)


def test_revisions_history_claims_shape_switch_source_invalidation(clients):
 e,r,pid,s,b,eid=setup(clients);c=post(e,pid,b).json();approve(r,pid,'components',c)
 for t in post(e,pid,{},'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
 post(e,pid,{},'/boq/generate');p=get(e,pid);old=p['components'][-1];new=copy.deepcopy(b);new['foundation_items'][0]['levels'][1]['parameters']['height']='500';new['object_hash']=store.digest(old);url=f'/api/v1/projects/{pid}/components/'+old['id'];res=e.put(url,json=new);assert res.status_code==200 and res.json()['preview']['quantity']=='2.626';assert get(e,pid)['component_history'][-1]==old and get(e,pid)['boq']==p['boq'] and all(t['state']=='stale' for t in get(e,pid)['takeoffs']);assert e.put(url,json=new).status_code==409;post(e,pid,{**b,'id':old['id']},expected=409);approve(r,pid,'components',res.json())
 for replacement in (row(eid,'circular'),rect_steps(eid),rect_row(eid=eid)):
  current=get(e,pid)['components'][-1];new=body(s,eid);new.update(foundation_items=[replacement],object_hash=store.digest(current));res=e.put(url,json=new);assert res.status_code==200,res.text;approve(r,pid,'components',res.json())
 s['object_hash']=store.digest(s);s['material_items'][0]['post_ranges'][0]['spacing']='3000';res=e.put(f'/api/v1/projects/{pid}/components/'+s['id'],json=s);assert res.status_code==200;approve(r,pid,'components',res.json());assert get(e,pid)['components'][-1]['state']=='stale'


def test_long_names_ten_circular_tiers_pdf(clients):
 e,r,pid,s,b,eid=setup(clients);b['name']='合成'+('圆形台阶长名称'*24)+'末';b['foundation_items'][0]['levels']=[level(height='100',eid=eid) for _ in range(10)];c=post(e,pid,b).json();approve(r,pid,'components',c)
 for t in post(e,pid,{},'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
 res=e.get(f'/api/v1/projects/{pid}/exports',params=dict(format='pdf',unpriced=True));assert res.status_code==200,res.text
