"""Synthetic circular hole-wall arcs, independent centers and periodic claims."""
import copy,io,random
from decimal import Decimal,localcontext,ROUND_HALF_UP
import pytest
from openpyxl import load_workbook
from pypdf import PdfReader
from backend import store
from backend.compound_takeoff import PI
from backend.engine import DomainError
from backend.foundation_void_formwork import calculate
from backend.foundation_void_formwork_reports import TITLES,ARC_TITLES,STRIP_TITLES
from tests.test_foundation_void_formwork import foundation,body as full_body,wall,setup as void_setup,check_asts
from tests.test_foundation_void_formwork_strips import selected
from tests.test_foundation_formwork_arcs import arc as outer_arc
from tests.test_foundation_openings import hole
from tests.test_foundation_recesses import interval
from tests.test_post_foundation import equal_body,fixed_body,fixed_added,equal_added,fixed_adjusted,equal_adjusted,combined
from tests.test_route_takeoff import evaluate
from tests.test_polygon_takeoff import post,mutate
from tests.test_workflow import get,approve,upload,create,priced_project

FRAME='合成角度基准：各自已审孔圆心为角度原点，零角沿基础共同局部+X，+X向图右、+Y向图上，从+Z俯视逆时针；孔心各异、轴向一致。仅软件验证。'


def arc(eid='test',start='0',end='90'):
 return {**outer_arc(eid,start,end),'coordinate_basis':FRAME}


def arc_wall(eid='test',start='0',end='90',z_start='300',z_end='600',level=1,hole_index=2,name='合成C1'):
 return {**wall(eid,face='circumference',start=z_start,end=z_end,level=level,hole_index=hole_index,name=name),'arc_range':arc(eid,start,end),'measurement_basis':'合成按来源实际孔径及该孔自身圆心，选入角度和共同标高内竖直孔壁；不以基础直径或基础圆心替代，来源全部个数一次，非真实工程。'}


def body(s,eid='test',level=1):
 b=full_body(s,eid);b.update(name='合成圆孔孔壁局部圆弧模板（非真实工程）',measurement_basis='合成H2圆孔实际Φ40mm，孔心(-80,0)mm；角0至90°及270至360°，Z300至600mm，22个基础，总净.132πm²→.415m²，仅总量修约3位，跨0°拆分、未选面不计。');b['contact_ranges']=[arc_wall(eid,level=level),arc_wall(eid,'270','360',level=level,name='合成C2')];return b


def setup(clients,pid=None,shape='rectangular'):
 e,r,pid,s,_,eid=void_setup(clients,pid,shape);return e,r,pid,s,body(s,eid,2 if 'stepped' in shape else 1),eid


def check(r):
 check_asts(r)
 with localcontext() as ctx:
  ctx.prec=80
  for c in r.get('void_arc_checks',[]):
   for k,a in c['formula_asts'].items():assert evaluate(a,r['dimensions'])==Decimal(c[k])


def test_actual_hole_diameter_offset_center_two_arcs_periodic_seam_no_mutation():
 s=foundation();b=body(s);before=copy.deepcopy((s,b));r=calculate(b,s);assert (s,b)==before and r['quantity']=='0.415' and r['rule_version']=='foundation-void-formwork-1.2';assert Decimal(r['deduction'])==0 and r['void_arc_checks'][0]['state']=='angular_tangent' and Decimal(r['void_arc_checks'][0]['angle_gap'])==0;assert all(Decimal(d['arc_diameter'])==Decimal('.04') and Decimal(d['arc_center_x'])==Decimal('-.08') and Decimal(d['arc_center_y'])==0 for d in r['contact_ranges']);check(r)
 with localcontext() as ctx:ctx.prec=80;assert Decimal(r['gross'])==Decimal('.132')*Decimal(PI)


@pytest.mark.parametrize('shape',['rectangular','stepped_rectangular','circular','stepped_circular'])
@pytest.mark.parametrize('make',[equal_body,fixed_body,fixed_added,equal_added,fixed_adjusted,equal_adjusted,combined,lambda s:combined(s,equal=True)])
def test_four_foundation_shapes_eight_final_counts_actual_hole_not_foundation(make,shape):
 s=foundation(make,shape);b=body(s,level=2 if 'stepped' in shape else 1);r=calculate(b,s);check(r)
 with localcontext() as ctx:ctx.prec=80;assert Decimal(r['gross'])==Decimal('.006')*Decimal(PI)*Decimal(r['contact_ranges'][0]['total_count'])


@pytest.mark.parametrize('unit,factor',[('mm','1'),('cm','.1'),('dm','.01'),('m','.001'),('km','.000001')])
def test_independent_z_unit_and_low_calling_precision(unit,factor):
 s=foundation();b=body(s)
 for d in b['contact_ranges']:d['unit']=unit;d['parameters']={k:str(Decimal(v)*Decimal(factor)) for k,v in d['parameters'].items()}
 with localcontext() as ctx:ctx.prec=6;r=calculate(b,s)
 assert r==calculate(b,s) and r['quantity']=='0.415' and len(r['unit_conversions'])==4;assert 'void_formwork.1.angle_start' in r['source_evidence'];check(r)


@pytest.mark.parametrize('left_end,right_start,accepted,state',[('90','90',True,'angular_tangent'),('89.999999999','90',True,'angular_disjoint'),('90.000000001','90',False,None),('360','0',False,None)])
def test_exact_angle_boundary_no_snap(left_end,right_start,accepted,state):
 s=foundation();b=body(s);b['contact_ranges']=[arc_wall(end=left_end),arc_wall(start=right_start,end='180',name='合成C2')]
 if accepted:r=calculate(b,s);assert r['void_arc_checks'][0]['state']==state;check(r)
 else:
  with pytest.raises(DomainError,match='重叠'):calculate(b,s)


@pytest.mark.parametrize('start,end',[('0','360'),('0','270'),('180','360'),('270','360'),('0.000000000000001','0.000000000000002')])
def test_full_large_and_tiny_arcs_do_not_take_smaller_arc(start,end):
 s=foundation();b=body(s);b['rounding_digits']=12;b['contact_ranges']=[arc_wall(start=start,end=end)];r=calculate(b,s);check(r)
 with localcontext() as ctx:ctx.prec=80;expected=Decimal('.04')*Decimal(PI)*(Decimal(end)-Decimal(start))/360*Decimal('.3')*22;assert Decimal(r['quantity'])==expected.quantize(Decimal('1e-12'),rounding=ROUND_HALF_UP)


@pytest.mark.parametrize('z_start,state',[('450','vertical_tangent'),('450.000000001','vertically_disjoint')])
def test_same_angles_allowed_only_if_z_is_disjoint(z_start,state):
 s=foundation();b=body(s);b['contact_ranges']=[arc_wall(end='300',z_end='450'),arc_wall(start='270',z_start=z_start,end='360',name='合成C2')];r=calculate(b,s);assert r['void_arc_checks'][0]['state']==state;check(r);b['contact_ranges'][1]['parameters']['z_start']='449.999999999'
 with pytest.raises(DomainError,match='重叠'):calculate(b,s)


def test_four_quarters_equal_whole_and_full_circle_conflicts():
 s=foundation();b=body(s);whole=wall(face='circumference',start='300',end='600',hole_index=2);b['contact_ranges']=[whole];r0=calculate(b,s);b['contact_ranges']=[arc_wall(start=str(i*90),end=str((i+1)*90),name='合成C'+str(i)) for i in range(4)];r=calculate(b,s);assert r['quantity']==r0['quantity']=='0.829' and r['gross']==r0['gross'];check(r);b['contact_ranges'].append({**whole,'name':'整周重复'})
 with pytest.raises(DomainError,match='重叠'):calculate(b,s)
 b['contact_ranges']=[arc_wall(end='360')];assert calculate(b,s)['quantity']==r0['quantity']


def test_two_different_holes_same_angles_are_independent_with_shared_axes():
 s=foundation();s['foundation_items'][0]['through_holes']['holes'].append({**hole(name='合成H4',x='-140',diameter='30'),'vertical_range':interval('300','600')});b=body(s);b['contact_ranges']=[arc_wall(),arc_wall(hole_index=4,name='另一圆孔')];r=calculate(b,s);assert not r['void_arc_checks'] and r['quantity']=='0.363';assert Decimal(r['contact_ranges'][1]['arc_center_x'])==Decimal('-.14') and Decimal(r['contact_ranges'][1]['arc_diameter'])==Decimal('.03');check(r)
 b['contact_ranges'][1]['arc_range']['coordinate_basis']='另一方向'
 with pytest.raises(DomainError,match='相同角度基准'):calculate(b,s)


def test_same_hole_across_layers_retains_hole_diameter_and_changes_only_available_z():
 s=foundation(shape='stepped_circular');s['foundation_items'][0]['through_holes']['holes'][1]['vertical_range']=interval('0','600');b=body(s);b['contact_ranges']=[arc_wall(z_start='0',z_end='200'),arc_wall(z_start='200',z_end='600',level=2,name='第二层')];r=calculate(b,s);assert r['quantity']=='0.415' and not r['void_arc_checks'];assert all(Decimal(d['arc_diameter'])==Decimal('.04') for d in r['contact_ranges']);check(r)


def test_mixed_rectangle_strip_circle_whole_and_arcs_keep_real_cuts():
 s=foundation();b=body(s);b['contact_ranges'].insert(0,selected(end='-10',name='合成矩形条带'));r=calculate(b,s);assert r['quantity']=='0.470' and len(r['contact_ranges'])==3 and len(r['void_arc_checks'])==1 and r['void_lateral_checks']==[];assert Decimal(r['contact_ranges'][0]['total_net_area'])==Decimal('.055');check(r)


def test_140_independent_arc_area_and_rounding_oracles():
 rng=random.Random(9801)
 for _ in range(140):
  s=foundation();d=rng.randint(1,80);s['foundation_items'][0]['through_holes']['holes'][1]['parameters']['diameter']=str(d);a=rng.randrange(359);b=rng.randint(a+1,360);lo=rng.randint(300,598);hi=rng.randint(lo+1,600);c=body(s);c['contact_ranges']=[arc_wall(start=str(a),end=str(b),z_start=str(lo),z_end=str(hi))];c['rounding_digits']=rng.randrange(13);r=calculate(c,s)
  with localcontext() as ctx:ctx.prec=80;expected=Decimal(d)/1000*Decimal(PI)*(b-a)/360*Decimal(hi-lo)/1000*22;assert Decimal(r['quantity'])==expected.quantize(Decimal(1).scaleb(-c['rounding_digits']),rounding=ROUND_HALF_UP)
  check(r)


@pytest.mark.parametrize('fault',['null','empty','extra','missing','parameters','proofs','start-proof','end-proof','frame','frame-proof','unconfirmed','truthy','unit','radian','negative','reverse','zero','over360','nan','inf','bool','precision','length-instead','rectangular','mixed','frame-different','frame-proof-different','outside-z','circle-center-input','foundation-diameter-input'])
def test_strict_arc_config_and_geometry(fault):
 s=foundation();b=body(s);d=b['contact_ranges'][0];c=d['arc_range']
 if fault=='null':d['arc_range']=None
 elif fault=='empty':d['arc_range']={}
 elif fault=='extra':c['radius']='20'
 elif fault=='missing':c.pop('unit')
 elif fault=='parameters':c['parameters'].pop('angle_start')
 elif fault=='proofs':c['evidence'].pop('angle_start')
 elif fault in ('start-proof','end-proof'):c['evidence']['angle_start' if fault=='start-proof' else 'angle_end']=' '
 elif fault=='frame':c['coordinate_basis']=' '
 elif fault=='frame-proof':c['coordinate_evidence_id']=' '
 elif fault in ('unconfirmed','truthy'):c['scope_confirmed']=False if fault=='unconfirmed' else 1
 elif fault in ('unit','radian'):c['unit']='mm' if fault=='unit' else 'rad'
 elif fault in ('negative','reverse','zero','over360','nan','inf','bool','precision'):c['parameters']['angle_start' if fault=='negative' else 'angle_end']={'negative':'-1','reverse':'-10','zero':'0','over360':'360.000000001','nan':'NaN','inf':'Infinity','bool':True,'precision':'90.0000000000000000001'}[fault]
 elif fault=='length-instead':c['parameters']['arc_length']='30'
 elif fault=='rectangular':d.update(hole_index=1,face='x_minus',parameters=dict(z_start='0',z_end='150'))
 elif fault=='mixed':d['lateral_range']={}
 elif fault in ('frame-different','frame-proof-different'):b['contact_ranges'][1]['arc_range']['coordinate_basis' if fault=='frame-different' else 'coordinate_evidence_id']='other'
 elif fault=='outside-z':d['parameters']['z_start']='299'
 elif fault=='circle-center-input':c['parameters']['x']='0'
 elif fault=='foundation-diameter-input':c['parameters']['diameter']='400'
 with pytest.raises(DomainError):calculate(b,s)


def test_volume_threshold_never_restores_or_removes_arc():
 values=[]
 for threshold in ('0','100'):
  s=foundation();s['foundation_items'][0]['through_holes']['deduction_rule']['threshold']=threshold;r=calculate(body(s),s);values.append((r['gross'],r['quantity'],r['deduction']));check(r)
 assert values[0]==values[1]


def test_preview_current_source_claims_periodic_seam_actual_id_and_full_circle(clients):
 e,r,pid,s,b,eid=setup(clients);before=get(e,pid);assert post(e,pid,b,'/components/preview').json()['quantity']=='0.415' and get(e,pid)==before;first=copy.deepcopy(b);first['contact_ranges']=first['contact_ranges'][:1];c=post(e,pid,first).json();approve(r,pid,'components',c);second=copy.deepcopy(b);second['contact_ranges']=second['contact_ranges'][1:];post(e,pid,second);post(e,pid,{**first,'id':c['id']},expected=409);second['contact_ranges'][0]['arc_range']['parameters']['angle_start']='89.999999999';post(e,pid,second,expected=409);whole=copy.deepcopy(b);whole['contact_ranges']=[wall(eid,face='circumference',start='300',end='600',hole_index=2)];post(e,pid,whole,expected=409);second['contact_ranges']=[arc_wall(eid,'90','180',name='合成C3')];second['contact_ranges'][0]['arc_range']['coordinate_basis']='另一个方向';post(e,pid,second,expected=409);second['contact_ranges'][0]['arc_range']['coordinate_basis']=FRAME;post(e,pid,second)


@pytest.mark.parametrize('fault',['angle','frame','proof','confirm','width','diameter','center','span','periodic-gap','formula','book'])
def test_saved_arc_tamper_rebuild(clients,fault):
 e,r,pid,s,b,eid=setup(clients);c=post(e,pid,b).json()
 if fault=='book':approve(r,pid,'components',c);post(e,pid,{},'/takeoffs/calculate')
 def forge(p):
  c=p['components'][-1];v=c['preview']['contact_ranges'][0]
  if fault=='angle':c['contact_ranges'][0]['arc_range']['parameters']['angle_end']='120'
  elif fault=='frame':
   for d in c['contact_ranges']:d['arc_range']['coordinate_basis']='forged'
  elif fault=='proof':c['contact_ranges'][0]['arc_range']['evidence']['angle_start']='forged'
  elif fault=='confirm':c['contact_ranges'][0]['arc_range']['scope_confirmed']=False
  elif fault=='width':v['width']='999'
  elif fault=='diameter':v['arc_diameter']='999'
  elif fault=='center':v['arc_center_x']='999'
  elif fault=='span':v['angle_span']='999'
  elif fault=='periodic-gap':c['preview']['void_arc_checks'][0]['angle_gap']='999'
  elif fault=='formula':v['formula_asts']['angle_span']={'constant':'0'}
  else:p['takeoffs'][-1]['quantity']='999'
 mutate(pid,forge);coll='takeoffs' if fault=='book' else 'components';approve(r,pid,coll,get(e,pid)[coll][-1],expected=404 if fault=='proof' else 422 if fault=='confirm' else 409)


@pytest.mark.parametrize('gate',['preview','save','review','calculate','book_review'])
def test_angle_original_bytes_five_gates(clients,gate):
 e,r,pid,s,b,eid=setup(clients)
 if gate in ('review','calculate','book_review'):c=post(e,pid,b).json()
 if gate in ('calculate','book_review'):approve(r,pid,'components',c)
 if gate=='book_review':t=post(e,pid,{},'/takeoffs/calculate').json()[-1]
 doc=get(e,pid)['documents'][-1];(store.ROOT/'originals'/doc['hash']).write_bytes(b'corrupt')
 if gate=='preview':post(e,pid,b,'/components/preview',409)
 elif gate=='save':post(e,pid,b,expected=409)
 elif gate=='review':approve(r,pid,'components',c,expected=409)
 elif gate=='calculate':post(e,pid,{},'/takeoffs/calculate',409)
 else:approve(r,pid,'takeoffs',t,expected=409)


def test_whole_to_arc_revision_history_boq_concurrency_disable_and_source_stale(clients):
 e,r,pid,s,b,eid=setup(clients);whole=copy.deepcopy(b);whole['contact_ranges']=[wall(eid,face='circumference',start='300',end='600',hole_index=2)];c=post(e,pid,whole).json();approve(r,pid,'components',c)
 for t in post(e,pid,{},'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
 post(e,pid,{},'/boq/generate');p=get(e,pid);old=p['components'][-1];b['object_hash']=store.digest(old);url=f'/api/v1/projects/{pid}/components/'+old['id'];res=e.put(url,json=b);assert res.status_code==200 and res.json()['preview']['quantity']=='0.415';assert get(e,pid)['component_history'][-1]==old and get(e,pid)['boq']==p['boq'] and all(t['state']=='stale' for t in get(e,pid)['takeoffs']);assert e.put(url,json=b).status_code==409;approve(r,pid,'components',res.json());whole['object_hash']=store.digest(get(e,pid)['components'][-1]);res=e.put(url,json=whole);assert res.status_code==200 and res.json()['preview']['quantity']=='0.829' and res.json()['preview']['rule_version']=='foundation-void-formwork-1.0';approve(r,pid,'components',res.json());s['object_hash']=store.digest(s);s['foundation_items'][0]['parameters']['height']='700';res=e.put(f'/api/v1/projects/{pid}/components/'+s['id'],json=s);assert res.status_code==200;approve(r,pid,'components',res.json());assert get(e,pid)['components'][-1]['state']=='stale'


def test_roles_tenants_and_foreign_angle_evidence(clients):
 e,r,pid,s,b,eid=setup(clients);post(r,pid,b,expected=403);assert clients['other'].get(f'/api/v1/projects/{pid}').status_code==404;other=create(e)['id'];upload(e,other,'异项目角度依据');b['contact_ranges'][0]['arc_range']['evidence']['angle_start']=get(e,other)['evidence'][-1]['id'];post(e,pid,b,expected=404)


@pytest.mark.parametrize('fmt',['json','xlsx','pdf'])
def test_three_formats_mixed_strip_arc_and_old_quote_isolation(clients,fmt):
 e,r,pid,_,q=priced_project(clients);e,r,pid,s,b,eid=setup(clients,pid);b['contact_ranges'].insert(0,selected(eid,end='-10',name='合成矩形条带'));c=post(e,pid,b).json();approve(r,pid,'components',c)
 for t in post(e,pid,{},'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
 boq=post(e,pid,{},'/boq/generate').json();assert any(d['quantity']=='0.470' and d['unit']=='m²' for d in boq);url=f'/api/v1/projects/{pid}/exports';res=e.get(url,params=dict(format=fmt,unpriced=True));assert res.status_code==200,res.text
 if fmt=='json':assert set(TITLES+ARC_TITLES+STRIP_TITLES)<={d[0] for d in res.json()['tables']}
 elif fmt=='xlsx':assert set(TITLES+ARC_TITLES+STRIP_TITLES)<=set(load_workbook(io.BytesIO(res.content)).sheetnames)
 else:assert all(v in ''.join(p.extract_text() for p in PdfReader(io.BytesIO(res.content)).pages) for v in TITLES+ARC_TITLES+STRIP_TITLES)
 assert not any(d[0] in TITLES+ARC_TITLES+STRIP_TITLES for d in e.get(url,params=dict(format='json',quote_id=q['id'])).json()['tables'])


@pytest.mark.parametrize('kind',['exterior','neighbor'])
def test_circle_point_tangency_has_no_positive_arc_area_cut(kind):
 s=foundation();holes=s['foundation_items'][0]['through_holes']['holes']
 if kind=='exterior':holes[1]['parameters']['x']='-180'
 else:holes[2].update(parameters=dict(x='-10',y='0',length='100',width='60'),vertical_range=interval('300','600'))
 b=body(s);b['contact_ranges']=[arc_wall(end='360')];r=calculate(b,s);assert r['quantity']=='0.829' and Decimal(r['deduction'])==0;check(r)


def test_legacy_through_circle_without_explicit_z_can_select_arcs():
 from tests.test_foundation_openings import body as through_body
 from tests.test_post_foundation import source
 s={**through_body(source(),shape='stepped_circular'),'id':'legacy-foundation','revision':1,'source_geometry':source()};b=body(s);b['contact_ranges']=[arc_wall(hole_index=1,z_start='0',z_end='200')];r=calculate(b,s);assert Decimal(r['quantity'])>0 and r['rule_version']=='foundation-void-formwork-1.2';check(r)


def test_interleaved_different_holes_compare_only_their_own_periodic_ranges():
 s=foundation();s['foundation_items'][0]['through_holes']['holes'].append({**hole(name='合成H4',x='-140',diameter='30'),'vertical_range':interval('300','600')});b=body(s);b['contact_ranges']=[arc_wall(),arc_wall(hole_index=4,name='另一孔首弧'),arc_wall(start='270',end='360',name='原孔末弧'),arc_wall(start='270',end='360',hole_index=4,name='另一孔末弧')];r=calculate(b,s);assert [(c['left_index'],c['right_index']) for c in r['void_arc_checks']]==[(1,3),(2,4)];assert all(c['state']=='angular_tangent' for c in r['void_arc_checks']);check(r)
