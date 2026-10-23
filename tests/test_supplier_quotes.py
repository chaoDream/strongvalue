import copy
import io
from datetime import date,timedelta
import pytest
from openpyxl import load_workbook
from pypdf import PdfReader
from backend import store
from tests.test_workflow import priced_project,get,approve,upload
from tests.test_unpriced_resources import add_pending


def fixture(clients):
    c,r,pid,comp,old=priced_project(clients)
    uploaded=upload(c,pid,'合成供应商A，TEST-MAT，仅数学测试规格，单价88.005元/m³，不含税，含运输不含安装，测试许可。','supplier.txt','合同计价')
    b=add_pending(c,pid)
    # Uploading new source material invalidates the original drawing chain.
    for coll in ('components','takeoffs'):
        for item in get(c,pid)[coll]:approve(r,pid,coll,item)
    eid=next(e['id'] for e in get(c,pid)['evidence'] if e['document_id']==uploaded['document']['id'])
    row=c.get(f'/api/v1/projects/{pid}/resource-prices').json()[0]
    payload={'request_id':'synthetic-quote-request-001','key':row['key'],'boq_hash':row['boq_hash'],'resource_hash':row['resource_hash'],
             'supplier':'合成供应商A','reference':'SYNTHETIC-A','received_on':date.today().isoformat(),'terms':'含运输，不含安装；验收后付款；仅软件测试',
             'reason':'核对合成报价与待询价资源规格单位一致','price':'88.005','license':'合成测试数据仅用于软件验收',
             'period':'合成2026-10','region':'河南省','tax_mode':'不含税',
             'valid_from':(date.today()-timedelta(days=1)).isoformat(),'valid_until':(date.today()+timedelta(days=365)).isoformat(),'evidence_id':eid}
    return c,r,pid,old,payload


def register(c,pid,body,expected=200):
    result=c.post(f'/api/v1/projects/{pid}/supplier-quotes',json=body)
    assert result.status_code==expected,result.text
    return result.json()


def read(c,pid,query='',expected=200):
    result=c.get(f'/api/v1/projects/{pid}/supplier-quotes'+query)
    assert result.status_code==expected,result.text
    return result.json()


def adopt(c,pid,offer,expected=200,**extra):
    response=c.post(f'/api/v1/projects/{pid}/supplier-quotes/{offer["id"]}/adopt',json={
        'object_hash':store.digest(offer),'reason':'明确选择A：规格与供货条件满足合成测试要求，不按最低价自动选取','confirm_conditions':True,**extra})
    assert response.status_code==expected,response.text
    return response.json()


def mutate(pid,fn):
    with store.transaction() as conn:
        import json
        p=json.loads(conn.execute('SELECT data FROM projects WHERE id=?',(pid,)).fetchone()['data']);fn(p);store.save(conn,p)


def test_register_read_only_review_choose_and_independent_boq_review(clients):
    c,r,pid,old,body=fixture(clients);before=get(c,pid)
    o=register(c,pid,body)['offer'];assert o['state']=='draft'
    assert register(c,pid,body)['duplicate'] is True
    listed=read(c,pid);assert listed==read(r,pid)
    assert listed['items'][0]['available'] is False
    for k in ['boq','components','takeoffs','quotes','stale_quotes']:
        assert get(c,pid)[k]==before[k]
    assert len(get(c,pid)['pricing_tasks'])==1
    adopt(c,pid,o,409)
    approve(c,pid,'supplier_quotes',o,expected=403);approve(r,pid,'supplier_quotes',o)
    o=get(c,pid)['supplier_quotes'][0]
    assert read(c,pid)['items'][0]['available'] is True
    adoption=adopt(c,pid,o)['adoption'];after=get(c,pid)
    assert after['boq'][0]['state']=='draft' and after['pricing_tasks']==[]
    resource=after['boq'][0]['quotas'][0]['resources'][0]
    assert resource['price']=='88.005' and resource['pricing_status']=='priced'
    for key in ['quantity','code','features']:assert after['boq'][0][key]==before['boq'][0][key]
    for key in ['name','unit','consumption','type','resource_code','specification','fee_allocation']:
        assert resource[key]==adoption['before'][key]
    assert after['boq'][0]['quotas'][0]['coefficient']==before['boq'][0]['quotas'][0]['coefficient']
    assert after['quotes']==before['quotes'] and old['id'] in after['stale_quotes']
    assert 'impact' not in adoption and adoption['before']['price'] is None
    assert adopt(c,pid,o)['duplicate'] is True
    adopt(c,pid,o,409,reason='不能改写选择理由')
    approve(c,pid,'boq',after['boq'][0],expected=403);approve(r,pid,'boq',after['boq'][0])
    response=c.post(f'/api/v1/projects/{pid}/quotes/calculate');assert response.status_code==200,response.text
    # 88.005 × 1.05 = 92.40525, rounded unit price 92.41 × 15 = 1386.15.
    assert response.json()['subtotal']=='1386.15'
    events=c.get(f'/api/v1/projects/{pid}/audit').json()
    assert [e['action'] for e in events].count('SUPPLIER_QUOTE_REGISTERED')==1
    assert [e['action'] for e in events].count('SUPPLIER_QUOTE_ADOPTED')==1


def test_multiple_offers_retained_no_lowest_auto_selection_or_tax_conversion(clients):
    c,r,pid,old,body=fixture(clients)
    a=register(c,pid,body)['offer']
    b=register(c,pid,{**body,'request_id':'synthetic-request-0002','supplier':'合成低价但另有条件B','reference':'B','price':'75.123456','tax_mode':'含税','terms':'运费安装另计'})['offer']
    assert len(get(c,pid)['pricing_tasks'])==1
    for offer in [a,b]:approve(r,pid,'supplier_quotes',offer)
    rows=read(c,pid,'?key='+body['key'])['items'];assert len(rows)==2 and all(x['available'] for x in rows)
    assert {x['offer']['metadata']['price'] for x in rows}=={'88.005','75.123456'}
    a=next(x['offer'] for x in rows if x['offer']['id']==a['id'])
    adopt(c,pid,a)
    assert get(c,pid)['boq'][0]['quotas'][0]['resources'][0]['price']=='88.005'
    rows=read(c,pid)['items'];assert len(rows)==2 and not any(x['available'] for x in rows)


@pytest.mark.parametrize('extra',[
    {'supplier':' '},{'price':'NaN'},{'price':'1e2'},{'price':88},{'price':'-1'},
    {'terms':''},{'received_on':'2026-02-30'},{'valid_until':'2026-1-1'},
    {'tax_mode':'未知'},{'state':'approved'},{'license':'x'*2001}])
def test_invalid_registration_is_atomic(clients,extra):
    c,r,pid,old,body=fixture(clients);before=get(c,pid)
    register(c,pid,{**body,**extra},422)
    assert get(c,pid)==before


def test_roles_hash_lease_and_idempotency(clients):
    c,r,pid,old,body=fixture(clients)
    register(r,pid,body,403);register(clients['other'],pid,body,404);read(clients['other'],pid,expected=404)
    register(c,pid,{**body,'boq_hash':'0'*64},409)
    o=register(c,pid,body)['offer']
    assert register(c,pid,body)['duplicate']
    register(c,pid,{**body,'price':'1'},409)
    approve(r,pid,'supplier_quotes',o);o=get(c,pid)['supplier_quotes'][0]
    adopt(r,pid,o,403);adopt(clients['other'],pid,o,404)
    adopt(c,pid,o,409,object_hash='0'*64)
    for extra in [{'confirm_conditions':False},{'confirm_conditions':'true'},{'reason':' '},{'price':'1'}]:adopt(c,pid,o,422,**extra)


@pytest.mark.parametrize('change',['target','proof','bytes','deleted','date','region'])
def test_source_target_and_context_changes_block_review_and_adoption(clients,change):
    c,r,pid,old,body=fixture(clients);o=register(c,pid,body)['offer']
    approve(r,pid,'supplier_quotes',o);o=get(c,pid)['supplier_quotes'][0]
    def action(p):
        if change=='target':p['boq'][0]['quotas'][0]['resources'][0]['specification']='changed'
        elif change=='proof':next(e for e in p['evidence'] if e['id']==body['evidence_id'])['text']+='changed'
        elif change=='date':p['pricing_date']='2099-01-01'
        elif change=='region':p['region']='其他地区'
    if change in ('bytes','deleted'):
        path=store.ROOT/'originals'/o['proof']['document_hash']
        if change=='bytes':path.write_bytes(b'corrupt')
        else:path.unlink()
    else:mutate(pid,action)
    before=get(c,pid)
    assert not read(c,pid)['items'][0]['available']
    adopt(c,pid,o,409);approve(r,pid,'supplier_quotes',o,expected=409)
    assert get(c,pid)==before


def test_expired_quote_kept_but_cannot_be_approved_or_used(clients):
    c,r,pid,old,body=fixture(clients)
    o=register(c,pid,{**body,'valid_from':'2000-01-01','valid_until':'2000-12-31'})['offer']
    assert '有效期' in read(c,pid)['items'][0]['issue']
    approve(r,pid,'supplier_quotes',o,expected=409)
    assert get(c,pid)['pricing_tasks']


def test_withdraw_selected_quote_blocks_current_boq_and_new_calculation(clients):
    c,r,pid,old,body=fixture(clients);o=register(c,pid,body)['offer'];approve(r,pid,'supplier_quotes',o)
    o=get(c,pid)['supplier_quotes'][0];adopt(c,pid,o)
    approve(r,pid,'boq',get(c,pid)['boq'][0])
    current=c.post(f'/api/v1/projects/{pid}/quotes/calculate').json()
    approve(r,pid,'supplier_quotes',o,'reject')
    p=get(c,pid);assert p['boq'][0]['state']=='stale' and current['id'] in p['stale_quotes']
    approve(r,pid,'boq',p['boq'][0],expected=409)
    blocked=c.post(f'/api/v1/projects/{pid}/quotes/calculate');assert blocked.status_code==409
    assert '供应商报价' in str(p['blockers'])
    # Historical quote payloads and the recorded alternatives remain unchanged.
    assert next(q for q in p['quotes'] if q['id']==current['id'])==current


def test_snapshot_three_formats_and_later_quotes_do_not_enter_old_versions(clients):
    c,r,pid,old,body=fixture(clients);o=register(c,pid,body)['offer'];approve(r,pid,'supplier_quotes',o)
    o=get(c,pid)['supplier_quotes'][0];adopt(c,pid,o);approve(r,pid,'boq',get(c,pid)['boq'][0])
    q=c.post(f'/api/v1/projects/{pid}/quotes/calculate').json();frozen=copy.deepcopy(q)
    base=f'/api/v1/projects/{pid}'
    for fmt in ['json','xlsx','pdf']:
        response=c.get(base+'/exports',params={'format':fmt,'quote_id':q['id']});assert response.status_code==200
        if fmt=='json':
            data=response.json();assert data['quote']==frozen
            assert {'供应商报价登记','供应商报价采用'}<={t[0] for t in data['tables']}
        elif fmt=='xlsx':
            wb=load_workbook(io.BytesIO(response.content));assert '供应商报价采用' in wb.sheetnames
        else:assert '供应商报价采用' in ''.join(p.extract_text() for p in PdfReader(io.BytesIO(response.content)).pages)
    add_pending(c,pid)
    newbody={**body,'request_id':'synthetic-later-quote-0001','supplier':'后续合成报价'}
    row=c.get(base+'/resource-prices').json()[0];newbody.update(boq_hash=row['boq_hash'],resource_hash=row['resource_hash'])
    register(c,pid,newbody)
    data=c.get(base+'/exports',params={'format':'json','quote_id':q['id']}).json()
    assert data['quote']==frozen and '后续合成报价' not in str(data['tables'])
    olddata=c.get(base+'/exports',params={'format':'json','quote_id':old['id']}).json()
    assert not any(t[0]=='供应商报价登记' for t in olddata['tables'])


def test_pagination_and_priced_resource_cannot_enter_missing_price_flow(clients):
    c,r,pid,old,body=fixture(clients)
    for n in range(3):register(c,pid,{**body,'request_id':f'synthetic-page-{n:04}'})
    assert read(c,pid,'?offset=2&limit=2')['total']==3
    assert len(read(c,pid,'?offset=2&limit=2')['items'])==1
    for query in ['?offset=-1','?limit=201','?limit=0']:read(c,pid,query,422)
    o=get(c,pid)['supplier_quotes'][0];approve(r,pid,'supplier_quotes',o);o=get(c,pid)['supplier_quotes'][0];adopt(c,pid,o)
    row=c.get(f'/api/v1/projects/{pid}/resource-prices').json()[0]
    register(c,pid,{**body,'request_id':'synthetic-priced-0001','boq_hash':row['boq_hash'],'resource_hash':row['resource_hash']},409)


@pytest.mark.parametrize('change',['price','specification','bytes','proof'])
def test_selected_quote_integrity_rechecked_at_boq_review_and_calculation(clients,change):
    c,r,pid,old,body=fixture(clients);o=register(c,pid,body)['offer'];approve(r,pid,'supplier_quotes',o)
    o=get(c,pid)['supplier_quotes'][0];adopt(c,pid,o)
    def action(p):
        if change in ('price','specification'):p['boq'][0]['quotas'][0]['resources'][0][change]='1' if change=='price' else 'other specification'
        elif change=='proof':next(e for e in p['evidence'] if e['id']==body['evidence_id'])['text']+='changed'
    if change=='bytes':(store.ROOT/'originals'/o['proof']['document_hash']).write_bytes(b'corrupt')
    else:mutate(pid,action)
    approve(r,pid,'boq',get(c,pid)['boq'][0],expected=409)
    assert get(c,pid)['blockers']
    assert c.post(f'/api/v1/projects/{pid}/quotes/calculate').status_code==409


def test_unselected_offer_rejection_does_not_invalidate_chosen_price(clients):
    c,r,pid,old,body=fixture(clients);a=register(c,pid,body)['offer']
    b=register(c,pid,{**body,'request_id':'synthetic-rejection-0002','supplier':'合成B'})['offer']
    approve(r,pid,'supplier_quotes',a);a=get(c,pid)['supplier_quotes'][0];adopt(c,pid,a)
    before=get(c,pid)
    approve(r,pid,'supplier_quotes',b,'reject')
    after=get(c,pid)
    assert after['boq']==before['boq'] and after['stale_quotes']==before['stale_quotes']


def test_restricted_project_offer_access_is_checked_for_existing_sessions(clients):
    from tests.test_project_access import restrict
    c,r,pid,old,body=fixture(clients);o=register(c,pid,body)['offer']
    restrict(clients['admin'],pid,[])
    read(c,pid,expected=404);register(c,pid,body,404);adopt(c,pid,o,404)
    read(r,pid,expected=404)
