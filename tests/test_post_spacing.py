"""Synthetic local post spacing caps; no global engineering compliance claim."""
import copy,io,random
from decimal import Decimal,localcontext
from fractions import Fraction
import pytest
from openpyxl import load_workbook
from pypdf import PdfReader
from backend import store
from backend.engine import DomainError
from backend.railing_material import calculate
from backend.railing_takeoff import snapshot
from tests.test_railing_posts import geometry,body as equal_body
from tests.test_fixed_posts import body as fixed_body
from tests.test_additional_posts import body as fixed_added,line_source
from tests.test_equal_additional_posts import body as equal_added
from tests.test_post_adjustments import body as fixed_adjusted
from tests.test_equal_post_adjustments import body as equal_adjusted,setup as source_setup
from tests.test_combined_post_adjustments import body as combined_body
from tests.test_route_takeoff import point,evaluate
from tests.test_route_railing import payload as railing_payload
from tests.test_railing_takeoff import interval
from tests.test_polygon_takeoff import post,mutate
from tests.test_workflow import get,approve,upload,create,priced_project


def check(limit='2500',unit='mm',eid='test'):
 return dict(max_spacing=str(limit),spacing_unit=unit,measurement_basis='合成仅核同一计入区间内实际相邻柱，端部、跨区间和跨边另核，不是全局设计合规',evidence=dict(max_spacing=eid,measurement_basis=eid))


def body(s,eid='test'):
 b=equal_adjusted(s,eid);b['name']='合成实际柱距核查（非真实工程）';b['material_items'][0]['spacing_check']=check(eid=eid);return b


def setup(clients,pid=None):
 e,r,pid,s,_,_=source_setup(clients,pid)
 upload(e,pid,'合成实际柱距核查验收，非真实工程。AB三维13m最大2m均分7跨，BC5m最大2m均分3跨，共11根，来源重复2。AB原j1 S13/7移2.5m，BC原j2 S10/3移3.5m。仅逐计入区间内最终实际相邻柱，独立上限2500mm；AB最大2.5m，BC最大11/6m，不默认采用原最大2m。等于上限允许，超限不能送审。BC起端未归属的0到5/3m只是边界距离，不是已核实柱距；跨区间、缺口、转角及闭环首尾另核。13m/7的长小数须保留精确表达式。每根1200mm、3000g/m，总79.2kg；核查上限改2600mm后质量仍79.2kg但须重新复核。','post-spacing-basis.txt')
 for c in get(e,pid)['components']:approve(r,pid,'components',c)
 s=get(e,pid)['components'][-1];eid=get(e,pid)['evidence'][-1]['id'];return e,r,pid,s,body(s,eid),eid


@pytest.mark.parametrize('make',[equal_body,fixed_body,fixed_added,equal_added,fixed_adjusted,equal_adjusted,combined_body,lambda s:combined_body(s,equal=True)])
def test_all_eight_layout_modes_count_mass_and_original_layout_stay_unchanged(make):
 s=geometry();b=make(s);old=calculate(copy.deepcopy(b),s);b['material_items'][0]['spacing_check']=check();before=copy.deepcopy((b,s));v=calculate(b,s);assert before==(b,s)
 assert v['quantity']==old['quantity'] and v['formula_ast']==old['formula_ast'] and v['rule_version']=='railing-material-mass-1.10'
 item=v['material_items'][0];g=item['post_spacing'];assert item['post_layout']==old['material_items'][0]['post_layout'];assert g['pair_count']==sum(max(0,r['post_count']-1) for r in item['post_layout']['post_ranges']) and g['checked_range_count']==2
 with localcontext() as c:
  c.prec=80;assert evaluate(g['limit_formula_ast'],v['dimensions'])==Decimal('2.5');assert evaluate(v['formula_ast'],v['dimensions'])==Decimal(v['gross'])
  for r in g['ranges']:
   for p in r['pairs']:
    assert abs(evaluate(p['distance_formula_ast'],v['dimensions'])-Decimal(p['distance']))<Decimal('1e-76')
    assert p['comparison']=='distance <= limit' and p['status']=='within_limit'
 assert len([k for k in v['source_evidence'] if '.spacing_check.' in k])==2


def test_post_move_gap_and_unowned_endpoint_distances_are_separate():
 s=geometry();b=body(s);v=calculate(b,s);g=v['material_items'][0]['post_spacing'];a,c=g['ranges'];assert a['max_distance']=='2.5' and a['max_pair_index']==1 and a['pair_count']==7 and a['unposted_start_distance']=='0' and a['unposted_end_distance']=='0'
 with localcontext() as ctx:
  ctx.prec=80;assert abs(Decimal(c['unposted_start_distance'])-Decimal(5)/3)<Decimal('1e-78');assert abs(Decimal(c['max_distance'])-Decimal(11)/6)<Decimal('1e-78')
 assert c['pair_count']==2 and c['pairs'][0]['from_post']==9 and c['pairs'][0]['to_post']==10
 b['material_items'][0]['spacing_check']['max_spacing']='2000'
 with pytest.raises(DomainError,match='柱1→柱2.*超过'):calculate(b,s)
 b['material_items'][0]['spacing_check']['max_spacing']='2499.999999999999'
 with pytest.raises(DomainError,match='超过'):calculate(b,s)


@pytest.mark.parametrize('basis',['spatial','horizontal'])
@pytest.mark.parametrize('unit,factor',[('m',Decimal(1)),('mm',Decimal(1000)),('cm',Decimal(100)),('dm',Decimal(10)),('km',Decimal('.001'))])
def test_units_and_selected_length_basis(basis,unit,factor):
 s=geometry(basis=basis);b=fixed_body(s);b['material_items'][0]['spacing_check']=check(str(Decimal(2)*factor),unit);v=calculate(b,s);assert v['material_items'][0]['post_spacing']['normalized_limit']=='2' and v['quantity']==('79.200' if basis=='spatial' else '50.400')


@pytest.mark.parametrize('fault',['null','list','missing_limit','missing_unit','missing_basis','extra_field','missing_proof','extra_proof','blank_proof','blank_basis','zero','negative','nan','infinity','bool','bad_unit','too_precise','area_mode'])
def test_strict_optional_check_schema(fault):
 s=geometry();b=body(s);m=b['material_items'][0];c=m['spacing_check']
 if fault=='null':m['spacing_check']=None
 elif fault=='list':m['spacing_check']=[]
 elif fault.startswith('missing_') and fault!='missing_proof':c.pop({'missing_limit':'max_spacing','missing_unit':'spacing_unit','missing_basis':'measurement_basis'}[fault])
 elif fault=='extra_field':c['auto_move']=True
 elif fault=='missing_proof':c['evidence'].pop('max_spacing')
 elif fault=='extra_proof':c['evidence']['spacing_unit']='test'
 elif fault=='blank_proof':c['evidence']['max_spacing']=' '
 elif fault=='blank_basis':c['measurement_basis']=' '
 elif fault in ('zero','negative','nan','infinity','bool','too_precise'):c['max_spacing']={'zero':'0','negative':'-2','nan':'NaN','infinity':'Infinity','bool':True,'too_precise':'2.0000000000000000001'}[fault]
 elif fault=='bad_unit':c['spacing_unit']='ft'
 else:
  from tests.test_railing_material import payload
  b=payload(s);b['material_items'][0]['spacing_check']=c
 with pytest.raises(DomainError):calculate(b,s)


def test_exact_irrational_cap_above_below_and_low_caller_precision():
 s=line_source(1);src=s['source_geometry'];src['route_points']=[point('A',0,0,0),point('B',1,1,0)];s['source_component_hash']=store.digest(snapshot(src));s['edge_decisions']=railing_payload(src)['edge_decisions'];b=equal_body(s);r=b['material_items'][0]['post_ranges'][0];r.update(spacing='1',spacing_unit='m');b['material_items'][0]['spacing_check']=check('.70710678118654752','m')
 with localcontext() as c:
  c.prec=5
  with pytest.raises(DomainError,match='超过'):calculate(b,s)
  b['material_items'][0]['spacing_check']['max_spacing']='.70710678118654753';g=calculate(b,s)['material_items'][0]['post_spacing']
 p=g['ranges'][0]['pairs'][0];assert p['exact_distance']==dict(constant=dict(numerator='0',denominator='1'),root_coefficient=dict(numerator='1',denominator='2'),squared_edge_length=dict(numerator='2',denominator='1'))


def test_no_bridge_over_gaps_and_unpaired_ranges_not_called_passed():
 s=line_source(8);s['edge_decisions'][0].update(mode='ranges',ranges=[interval(0,1),interval(7,8)]);b=equal_body(s);m=b['material_items'][0];m['post_ranges'][0]['end_post']=False;m['post_ranges'][1]['start_post']=True;m['spacing_check']=check('1','m');g=calculate(b,s)['material_items'][0]['post_spacing'];assert (g['pair_count'],g['checked_range_count'],g['unpaired_range_count'])==(1,1,1)
 a,c=g['ranges'];assert a['status']=='no_pairs' and a['unposted_end_distance']=='1' and a['max_distance'] is None and c['pairs'][0]['distance']=='1'
 m['post_ranges'][0].update(start_post=False,end_post=False);g=calculate(b,s)['material_items'][0]['post_spacing'];assert g['ranges'][0]['fully_unposted'] and g['ranges'][0]['unposted_start_distance'] is None
 m['post_ranges'][1]['end_post']=False
 with pytest.raises(DomainError,match='没有区间内实际相邻'):calculate(b,s)


@pytest.mark.parametrize('direction',['from_start','from_end'])
def test_both_anchors_partial_offsets_and_added_order(direction):
 s=line_source(12);s['edge_decisions'][0].update(mode='ranges',ranges=[interval(2,10)]);b=fixed_added(s);m=b['material_items'][0];r=m['post_ranges'][0];r['direction']=direction;r['additional_posts'][0]['s']='2500';m['spacing_check']=check('2','m');g=calculate(b,s)['material_items'][0]['post_spacing'];assert g['pair_count']==5 and g['ranges'][0]['max_distance']=='2';assert g['ranges'][0]['pairs'][0]['distance']=='0.5'


@pytest.mark.parametrize('starts',[False,True])
def test_closed_polygon_and_hole_do_not_imply_wrap_pairs(starts):
 from tests.test_railing_takeoff import geometry as pg,payload as pp
 src=pg();s=pp(src);s.update(id='polygon',revision=1,source_geometry=snapshot(src))
 for d in s['edge_decisions']:d.update(mode='full',ranges=[])
 b=equal_body(s);m=b['material_items'][0]
 for r in m['post_ranges']:r.update(start_post=starts,end_post=not starts)
 m['spacing_check']=check('2','m');v=calculate(b,s);g=v['material_items'][0]['post_spacing'];assert g['pair_count']==sum(max(0,r['post_count']-1) for r in v['material_items'][0]['post_layout']['post_ranges']);assert '闭环首尾' in g['notice']


def test_seeded_independent_rational_distance_oracle():
 rng=random.Random(7701)
 for _ in range(40):
  L=rng.randint(8,50);step=rng.randint(1,5);s=line_source(L);b=equal_body(s);m=b['material_items'][0];m['post_ranges'][0].update(spacing=str(step),spacing_unit='m');m['spacing_check']=check(str(step),'m');g=calculate(b,s)['material_items'][0]['post_spacing'];n=(L+step-1)//step;expected=Fraction(L,n)
  for p in g['ranges'][0]['pairs']:
   exact=p['exact_distance'];q=exact['squared_edge_length'];coef=exact['root_coefficient'];const=exact['constant'];assert Fraction(int(const['numerator']),int(const['denominator']))+Fraction(int(coef['numerator']),int(coef['denominator']))*L==expected and Fraction(int(q['numerator']),int(q['denominator']))==L*L


@pytest.mark.parametrize('fmt',['json','xlsx','pdf'])
def test_workflow_three_exports_and_old_quote_isolation(clients,fmt):
 e,r,pid,_,quote=priced_project(clients);e,r,pid,s,b,eid=setup(clients,pid);before=get(e,pid);post(e,pid,b,'/components/preview');assert get(e,pid)==before;c=post(e,pid,b).json();approve(r,pid,'components',c)
 for t in post(e,pid,{},'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
 assert any(v['quantity']=='79.200' for v in post(e,pid,{},'/boq/generate').json());url=f'/api/v1/projects/{pid}/exports';res=e.get(url,params=dict(format=fmt,unpriced=True));assert res.status_code==200,res.text;titles={'立柱区间实际柱距','立柱相邻柱距明细','立柱柱距核查依据'}
 if fmt=='json':assert titles<={t[0] for t in res.json()['tables']}
 elif fmt=='xlsx':assert titles<=set(load_workbook(io.BytesIO(res.content)).sheetnames)
 else:assert all(t in ''.join(p.extract_text() for p in PdfReader(io.BytesIO(res.content)).pages) for t in titles)
 old=e.get(url,params=dict(format='json',quote_id=quote['id'])).json();assert old['quote']==quote and not any(t[0] in titles for t in old['tables'])


@pytest.mark.parametrize('fault',['limit','proof','pair','exact','boundary','count','ast','book'])
def test_saved_inputs_and_derived_records_rebuilt(clients,fault):
 e,r,pid,s,b,eid=setup(clients);c=post(e,pid,b).json()
 if fault=='book':approve(r,pid,'components',c);post(e,pid,{},'/takeoffs/calculate')
 def forge(p):
  c=p['components'][-1];m=c['material_items'][0];g=c['preview']['material_items'][0]['post_spacing']
  if fault=='limit':m['spacing_check']['max_spacing']='2600'
  elif fault=='proof':m['spacing_check']['evidence']['max_spacing']='forged'
  elif fault=='pair':g['ranges'][0]['pairs'][0]['distance']='2'
  elif fault=='exact':g['ranges'][0]['pairs'][0]['exact_distance']['constant']['numerator']='2'
  elif fault=='boundary':g['ranges'][1]['unposted_start_distance']='0'
  elif fault=='count':g['pair_count']=99
  elif fault=='ast':g['ranges'][0]['pairs'][0]['distance_formula_ast']={'constant':'2.5'}
  else:p['takeoffs'][-1]['material_items'][0]['post_spacing']['ranges'][0]['max_distance']='2'
 mutate(pid,forge);coll='takeoffs' if fault=='book' else 'components';approve(r,pid,coll,get(e,pid)[coll][-1],expected=404 if fault=='proof' else 409)


def test_same_mass_limit_revision_removal_history_concurrency_and_source(clients):
 e,r,pid,s,b,eid=setup(clients);c=post(e,pid,b).json();approve(r,pid,'components',c)
 for t in post(e,pid,{},'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
 post(e,pid,{},'/boq/generate');p=get(e,pid);old=p['components'][-1];old_boq=copy.deepcopy(p['boq']);new=copy.deepcopy(b);new['material_items'][0]['spacing_check']['max_spacing']='2600';new['object_hash']=store.digest(old);url=f'/api/v1/projects/{pid}/components/'+c['id'];res=e.put(url,json=new);assert res.status_code==200,res.text;p=get(e,pid);assert p['component_history'][-1]==old and all(t['state']=='stale' for t in p['takeoffs']) and p['boq']==old_boq and res.json()['preview']['quantity']=='79.200';assert e.put(url,json=new).status_code==409;approve(r,pid,'components',res.json())
 for t in post(e,pid,{},'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
 assert get(e,pid)['takeoffs'][-1]['material_items'][0]['post_spacing']['normalized_limit']=='2.6'
 assert any(v['quantity']=='79.200' for v in post(e,pid,{},'/boq/generate').json())
 current=get(e,pid)['components'][-1];new['object_hash']=store.digest(current);new['material_items'][0].pop('spacing_check');res=e.put(url,json=new);assert res.status_code==200 and res.json()['preview']['rule_version']=='railing-material-mass-1.8' and 'post_spacing' not in res.json()['preview']['material_items'][0];assert get(e,pid)['component_history'][-1]==current;approve(r,pid,'components',res.json())
 s['object_hash']=store.digest(s);s['edge_decisions'][0].update(mode='ranges',ranges=[interval(0,10,eid)]);res=e.put(f'/api/v1/projects/{pid}/components/'+s['id'],json=s);assert res.status_code==200;approve(r,pid,'components',res.json());assert get(e,pid)['components'][-1]['state']=='stale';post(e,pid,b,'/components/preview',409)


@pytest.mark.parametrize('gate',['preview','save','review','calculate','book_review'])
def test_original_bytes_gates(clients,gate):
 e,r,pid,s,b,eid=setup(clients)
 if gate in ('review','calculate','book_review'):c=post(e,pid,b).json()
 if gate in ('calculate','book_review'):approve(r,pid,'components',c)
 if gate=='book_review':t=post(e,pid,{},'/takeoffs/calculate').json()[-1]
 doc=get(e,pid)['documents'][-1];(store.ROOT/'originals'/doc['hash']).write_bytes(b'corrupt')
 if gate=='preview':post(e,pid,b,'/components/preview',409)
 elif gate=='save':post(e,pid,b,expected=409)
 elif gate=='review':approve(r,pid,'components',c,expected=409)
 elif gate=='calculate':post(e,pid,{},'/takeoffs/calculate',409)
 else:approve(r,pid,'takeoffs',t,expected=409)


def test_roles_tenant_and_foreign_limit_evidence(clients):
 e,r,pid,s,b,eid=setup(clients);post(r,pid,b,expected=403);assert clients['other'].get(f'/api/v1/projects/{pid}').status_code==404
 other=create(e)['id'];upload(e,other,'其他工程原件');b['material_items'][0]['spacing_check']['evidence']['max_spacing']=get(e,other)['evidence'][-1]['id'];post(e,pid,b,expected=404)


def test_max_pair_and_range_limits_with_mixed_materials():
 s=line_source(1998);b=fixed_body(s);m=b['material_items'][0];m['spacing_check']=check('2','m');assert calculate(b,s)['material_items'][0]['post_spacing']['pair_count']==999
 b['material_items']=[{**copy.deepcopy(m),'name':str(i)} for i in range(5)];v=calculate(b,s);assert sum(i['post_spacing']['pair_count'] for i in v['material_items'])==4995
 b['material_items'].append({**copy.deepcopy(m),'name':'6'})
 with pytest.raises(DomainError,match='5000'):calculate(b,s)


def test_pdf_long_names_and_long_s_pairs_remain_exportable(clients):
 e,r,pid,s,b,eid=setup(clients);b['name']='合成'+('长名称验收'*39)+'末';b['material_items'][0]['name']='合成'+('长材料名称'*39)+'末';c=post(e,pid,b).json();approve(r,pid,'components',c)
 for t in post(e,pid,{},'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
 res=e.get(f'/api/v1/projects/{pid}/exports',params=dict(format='pdf',unpriced=True));assert res.status_code==200,res.text
 txt=''.join(p.extract_text() for p in PdfReader(io.BytesIO(res.content)).pages);assert '立柱相邻柱距明细' in txt and '立柱柱距核查依据' in txt
