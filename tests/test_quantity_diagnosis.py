import copy
import io
import pytest
from openpyxl import load_workbook
from backend import store
from backend.reports import tables
from tests.test_workflow import get, upload, approve, priced_project


def failed_baseline(e,pid,quantity='14',name='baseline.csv'):
    text='构件名称,单位,基准工程量,绝对容差,相对容差\n测试底板,m³,'+quantity+',0,0\n'
    r=e.post(f'/api/v1/projects/{pid}/takeoffs/reconcile',data={'basis':'独立合成底稿，仅用于测试'},files={'file':(name,text.encode(),'text/csv')})
    assert r.status_code==200,r.text
    return r.json()


def body(r,eid,category='geometry'):
    return {'reconciliation_id':r['id'],'target_key':'row:2','category':category,'evidence_id':eid,
            'observed':'系统侧参数与原始精度记录','reference':'独立人工基准参数及条款记录',
            'reason':'合成差异分析，位置为原图测试构件，不代表专业真值','action':'核查原图与底稿，修正后重算并重新对账','previous_hash':''}


def submit(e,pid,data):
    response=e.post(f'/api/v1/projects/{pid}/quantity-diagnoses',json=data)
    assert response.status_code==200,response.text
    return response.json()


@pytest.mark.parametrize('category',['geometry','coverage','method','standard','precision'])
def test_categories_review_does_not_approve_failed_quantities(clients,category):
    e,r,pid,comp,q=priced_project(clients);base=f'/api/v1/projects/{pid}'
    baseline=failed_baseline(e,pid);p=get(e,pid)
    d=submit(e,pid,body(baseline,p['evidence'][0]['id'],category))
    assert d['target']['actual']=='15.0' and d['target']['expected']=='14'
    assert d['target']['delta']=='1.0'
    approve(r,pid,'quantity_diagnoses',d)
    p=get(e,pid)
    assert p['reconciliations'][-1]==baseline and not baseline['passed']
    assert p['quotes'][0]==q
    approve(r,pid,'reconciliations',baseline,expected=409)
    assert e.patch(base+'/settings',json={'require_golden_baseline':True,'reason':'启用专业底稿门禁测试'}).status_code==200
    assert '黄金底稿' in '；'.join(get(e,pid)['blockers'])


def test_version_guards_and_tenant_role_boundaries(clients):
    e,r,pid,comp,q=priced_project(clients);base=f'/api/v1/projects/{pid}'
    baseline=failed_baseline(e,pid);data=body(baseline,get(e,pid)['evidence'][0]['id'])
    assert clients['other'].post(base+'/quantity-diagnoses',json=data).status_code==404
    assert r.post(base+'/quantity-diagnoses',json=data).status_code==403
    d=submit(e,pid,data)
    assert e.post(base+'/quantity-diagnoses',json=data).status_code==409
    approve(e,pid,'quantity_diagnoses',d,expected=403)
    data.update(previous_hash=store.digest(d),reason='修订分析，保留前一版')
    d2=submit(e,pid,data)
    assert d2['version']==2 and get(e,pid)['quantity_diagnoses'][0]==d
    approve(r,pid,'quantity_diagnoses',d,expected=409)
    targets=r.get(base+'/review-batch-targets?collection=quantity_diagnoses').json()['targets']
    assert [t['object']['id'] for t in targets]==[d2['id']]
    wrong=r.post(base+'/reviews',json={'collection':'quantity_diagnoses','object_id':d2['id'],'object_hash':'outdated','decision':'approve','reason':'test'})
    assert wrong.status_code==409
    approve(r,pid,'quantity_diagnoses',d2)
    data['previous_hash']=store.digest(d2)
    assert e.post(base+'/quantity-diagnoses',json=data).status_code==409


@pytest.mark.parametrize('mutation',[
    lambda d:d.update(category='unknown'),lambda d:d.update(observed=''),lambda d:d.update(reference=[]),
    lambda d:d.update(action=''),lambda d:d.update(evidence_id=None),lambda d:d.update(target_key='row:999'),
])
def test_invalid_diagnosis_has_no_partial_write(clients,mutation):
    e,r,pid,comp,q=priced_project(clients);baseline=failed_baseline(e,pid)
    data=body(baseline,get(e,pid)['evidence'][0]['id']);mutation(data)
    response=e.post(f'/api/v1/projects/{pid}/quantity-diagnoses',json=data)
    assert response.status_code==422,response.text
    assert get(e,pid)['quantity_diagnoses']==[]


def test_missing_coverage_and_unmatched_row_can_be_diagnosed(clients):
    e,r,pid,comp,q=priced_project(clients)
    baseline=failed_baseline(e,pid,name='unmatched.csv')
    # Use a different component name so the baseline row is unmatched and the takeoff uncovered.
    text='构件名称,单位,基准工程量,绝对容差,相对容差\n缺少对应构件,m³,14,0,0\n'
    response=e.post(f'/api/v1/projects/{pid}/takeoffs/reconcile',data={'basis':'合成底稿覆盖测试'},files={'file':('missing.csv',text.encode(),'text/csv')})
    baseline=response.json();eid=get(e,pid)['evidence'][0]['id']
    row=submit(e,pid,body(baseline,eid,'coverage'))
    assert 'actual' not in row['target']
    data=body(baseline,eid,'coverage');data['target_key']='missing:'+baseline['missing_takeoffs'][0]
    missing=submit(e,pid,data)
    assert missing['target']['kind']=='missing_takeoff'
    approve(r,pid,'quantity_diagnoses',missing)


def test_changed_calculation_context_requires_new_diagnosis(clients):
    e,r,pid,comp,q=priced_project(clients);baseline=failed_baseline(e,pid)
    d=submit(e,pid,body(baseline,get(e,pid)['evidence'][0]['id']))
    # Change the computational precision/rule context without changing quantity.
    with store.transaction() as c:
        from backend.main import project
        p=project(c,pid,{'tenant':'local'});p['takeoffs'][0]['rule_version']='different-test-rule';store.save(c,p)
    approve(r,pid,'quantity_diagnoses',d,expected=409)
    assert get(e,pid)['quantity_diagnoses'][0]==d
    data=body(baseline,d['evidence_id']);data['previous_hash']=store.digest(d)
    assert e.post(f'/api/v1/projects/{pid}/quantity-diagnoses',json=data).status_code==409


def test_legacy_baseline_needs_complete_context_before_diagnosis(clients):
    e,r,pid,comp,q=priced_project(clients);baseline=failed_baseline(e,pid)
    with store.transaction() as c:
        from backend.main import project
        p=project(c,pid,{'tenant':'local'});p['reconciliations'][0].pop('calculation_hash');store.save(c,p)
    res=e.post(f'/api/v1/projects/{pid}/quantity-diagnoses',json=body(baseline,get(e,pid)['evidence'][0]['id']))
    assert res.status_code==409 and '历史底稿' in res.text


def test_golden_gate_checks_calculation_context_even_when_quantity_is_same(clients):
    e,r,pid,comp,q=priced_project(clients);base=f'/api/v1/projects/{pid}'
    baseline=failed_baseline(e,pid,quantity='15');approve(r,pid,'reconciliations',baseline)
    e.patch(base+'/settings',json={'require_golden_baseline':True,'reason':'测试完整计算口径'})
    assert not any('黄金底稿' in message for message in get(e,pid)['blockers'])
    with store.transaction() as c:
        from backend.main import project
        p=project(c,pid,{'tenant':'local'});p['takeoffs'][0]['rule_version']='different-rule';store.save(c,p)
    assert any('黄金底稿' in message for message in get(e,pid)['blockers'])
    approve(r,pid,'reconciliations',get(e,pid)['reconciliations'][0],expected=409)
    fresh=failed_baseline(e,pid,quantity='15',name='new-context.csv');approve(r,pid,'reconciliations',fresh)
    assert not any('黄金底稿' in message for message in get(e,pid)['blockers'])


def test_missing_corrupt_and_superseded_sources_block_review(clients):
    e,r,pid,comp,q=priced_project(clients);baseline=failed_baseline(e,pid)
    p=get(e,pid);eid=p['evidence'][0]['id'];d=submit(e,pid,body(baseline,eid))
    source=next(doc for doc in p['documents'] if doc['id']==next(ev for ev in p['evidence'] if ev['id']==baseline['evidence_id'])['document_id'])
    path=store.ROOT/'originals'/source['hash'];raw=path.read_bytes();path.unlink()
    approve(r,pid,'quantity_diagnoses',d,expected=409)
    path.write_bytes(b'corrupt');approve(r,pid,'quantity_diagnoses',d,expected=409)
    path.write_bytes(raw);upload(e,pid,'新版合成原图，尺寸须重核')
    approve(r,pid,'quantity_diagnoses',d,expected=409)


def test_passed_rows_are_not_variance_targets_and_snapshot_exports_are_fixed(clients):
    e,r,pid,comp,old=priced_project(clients);base=f'/api/v1/projects/{pid}'
    passed=failed_baseline(e,pid,quantity='15',name='passed.csv')
    data=body(passed,get(e,pid)['evidence'][0]['id'])
    assert e.post(base+'/quantity-diagnoses',json=data).status_code==422
    baseline=failed_baseline(e,pid);long_body=body(baseline,data['evidence_id']);long_body['reason']='长篇合成分析核查原图与计量口径。'*80
    d=submit(e,pid,long_body)
    approve(r,pid,'quantity_diagnoses',d)
    # Adding baseline documents does not invalidate the source geometry or prices.
    response=e.post(base+'/quotes/calculate');assert response.status_code==200,response.text
    q=response.json();assert len(q['snapshot']['quantity_diagnoses'])==1
    x=e.get(base+'/exports',params={'format':'xlsx','quote_id':q['id']});assert x.status_code==200
    w=load_workbook(io.BytesIO(x.content));assert '工程量差异诊断' in w.sheetnames
    rows=list(w['工程量差异诊断'].values);assert any('诊断已复核' in str(v) for row in rows for v in row)
    assert any('不改变容差' in str(v) for row in rows for v in row)
    assert '长篇合成分析核查原图与计量口径。'*80 in ''.join(str(row[5] or '') for row in rows[2:])
    assert all(48<=w['工程量差异诊断'].row_dimensions[n].height<=409 for n in range(3,len(rows)+1))
    pdf=e.get(base+'/exports',params={'format':'pdf','quote_id':q['id']});assert pdf.status_code==200
    current=get(e,pid)['quantity_diagnoses'][-1]
    data=body(baseline,data['evidence_id']);data.update(previous_hash=store.digest(current),observed='新版系统侧记录')
    submit(e,pid,data)
    assert get(e,pid)['quotes'][-1]==q
    assert q['snapshot']['quantity_diagnoses'][0]['observed']!=data['observed']
    assert get(e,pid)['quotes'][0]==old
