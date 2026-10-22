"""Evidence-backed partial face measurement, using synthetic design coordinates."""
import copy
import io
import random
from decimal import Decimal, localcontext
import pytest
from openpyxl import load_workbook
from pypdf import PdfReader
from backend import store
from backend.engine import DomainError
from backend.surface_takeoff import boundary, partition, calculate
from tests.test_surface_takeoff import geometry, payload as base_payload
from tests.test_spatial_takeoff import region
from tests.test_compound_takeoff import ast_value
from tests.test_workflow import create, get, upload, approve


def divisions(source,eid='test'):
    front=next(f for f in boundary(source)['faces'] if f['normal']=='y-' and f['plane']=='0')
    return [{'face_id':front['face_id'],'u_cuts':[],'v_cuts':[{'value':'1200','unit':'mm','evidence_id':eid}]}]


def payload(source,eid='test',cut_eid='test',cuts=None):
    groups=divisions(source,cut_eid) if cuts is None else cuts
    result=partition(source,groups)
    body=base_payload(source,eid,included=lambda f:False)
    body.update(name='合成局部防水（非真实工程）',purpose='waterproof',face_subdivisions=groups,
                measurement_basis='仅合成：正面仅计Z=0至1.2m范围，其余排除，非专业计量规则',
                face_decisions=[{'face_id':f['face_id'],'included':bool(f.get('parent_index')) and Decimal(f['v_to'])<=Decimal('1.2'),
                                 'basis':'合成局部分区核定：仅计正面下部，其他排除','evidence_id':eid} for f in result['faces']])
    return body


def setup(clients):
    e,r=clients['estimator'],clients['reviewer'];pid=create(e)['id']
    upload(e,pid,'仅合成：同一XYZ基准，外块(0,0,0)4×3×2m，净空(1,1,0.5)2×1×1.5m，整组1。')
    upload(e,pid,'仅合成做法：Y=0正面沿Z=1200mm分区，仅下部0至1200mm计入防水，其余全部排除。无专业效力。',name='cut-basis.txt')
    p=get(e,pid);eid=p['evidence'][0]['id'];cut_eid=p['evidence'][1]['id']
    s=geometry(eid);s.pop('id');s.pop('revision')
    response=e.post('/api/v1/projects/'+pid+'/components',json=s);assert response.status_code==200,response.text
    approve(r,pid,'components',response.json());source=get(e,pid)['components'][0]
    return e,r,pid,source,payload(source,cut_eid,cut_eid),cut_eid


def test_partial_height_preserves_area_and_raw_ast():
    s=geometry();body=payload(s);before=copy.deepcopy(body);result=calculate(body,s)
    assert body==before and result['rule_version']=='solid-boundary-1.1'
    assert Decimal(result['quantity'])==Decimal('4.8') and Decimal(result['surface_summary']['total_area'])==61
    children=[f for f in result['surface_faces'] if f.get('parent_index')]
    assert len(children)==2 and {Decimal(f['area']) for f in children}=={Decimal('4.8'),Decimal('3.2')}
    assert Decimal(result['surface_summary']['excluded_area'])==Decimal('56.2')
    cut=result['surface_subdivisions'][0]
    assert cut['original_value']=='1200' and cut['factor']=='0.001' and cut['canonical_value']=='1.2' and cut['coordinate_axis']=='z'
    with localcontext() as ctx:
        ctx.prec=80
        assert ast_value(result['formula_ast'],result['dimensions'])==Decimal('4.8')
        for f in result['surface_faces']:assert ast_value(f['formula_ast'],result['dimensions'])==Decimal(f['area'])


def test_two_axes_negative_coordinates_units_and_repeat_count():
    s=geometry();s['regions']=[region('negative',-2,-3,-4,4,3,4)];s['multiplicity']='2'
    face=next(f for f in boundary(s)['faces'] if f['normal']=='y-')
    cuts=[{'face_id':face['face_id'],'u_cuts':[{'value':'-500','unit':'mm','evidence_id':'test'}],
           'v_cuts':[{'value':'-250','unit':'cm','evidence_id':'test'}]}]
    body=payload(s,cuts=cuts);body['face_decisions']=[{**d,'included':True} for d in body['face_decisions']]
    result=calculate(body,s)
    assert len([f for f in result['surface_faces'] if f.get('parent_index')])==4
    assert sum(Decimal(f['area']) for f in result['surface_faces'] if f.get('parent_index'))==16
    assert Decimal(result['quantity'])==160  # 2*(4*3+4*4+3*4)*2


def test_random_partitions_close_and_match_independent_decimal_rectangles():
    rng=random.Random(2801)
    for run in range(30):
        s=geometry();s['regions']=[region('box',0,0,0,4,3,2)]
        face=next(f for f in boundary(s)['faces'] if f['normal']=='y-')
        xs=sorted(rng.sample(range(1,40),rng.randint(1,4)));zs=sorted(rng.sample(range(1,20),rng.randint(1,4)))
        cuts=[{'face_id':face['face_id'],'u_cuts':[{'value':str(x*100),'unit':'mm','evidence_id':'test'} for x in xs],
               'v_cuts':[{'value':str(z*10),'unit':'cm','evidence_id':'test'} for z in zs]}]
        result=partition(s,cuts);children=[f for f in result['faces'] if f.get('parent_index')]
        expected=[(Decimal(b-a)/10)*(Decimal(d-c)/10) for a,b in zip([0]+xs,xs+[40]) for c,d in zip([0]+zs,zs+[20])]
        assert sorted(Decimal(f['area']) for f in children)==sorted(expected)
        assert sum(Decimal(f['area']) for f in children)==8
        with localcontext() as ctx:
            ctx.prec=80
            for f in children:assert ast_value(f['formula_ast'],result['dimensions'])==Decimal(f['area'])


def test_cut_order_does_not_change_child_geometry_or_ids():
    s=geometry();cuts=divisions(s);cuts[0]['v_cuts'].append({'value':'50','unit':'cm','evidence_id':'test'})
    a=partition(s,cuts);cuts[0]['v_cuts'].reverse();b=partition(s,cuts)
    keys=['face_id','area','u_from','u_to','v_from','v_to']
    assert [{k:f[k] for k in keys} for f in a['faces']]==[{k:f[k] for k in keys} for f in b['faces']]


def test_absent_or_empty_partitions_keep_legacy_calculation_identical():
    s=geometry();body=base_payload(s);old=calculate(body,s)
    assert calculate({**body,'face_subdivisions':[]},s)==old
    assert old['rule_version']=='solid-boundary-1.0' and 'surface_subdivisions' not in old


@pytest.mark.parametrize('case',['null','fields','unknown','duplicate_face','empty','too_many','outside','boundary','duplicate_coordinate','unit','nan','precision','proof','not_list','extra_cut_field'])
def test_invalid_subdivisions_rejected(case):
    s=geometry();cuts=divisions(s);group=cuts[0];cut=group['v_cuts'][0]
    if case=='null':cuts=None
    elif case=='fields':group['rotation']='30'
    elif case=='unknown':group['face_id']='unknown'
    elif case=='duplicate_face':cuts.append(copy.deepcopy(group))
    elif case=='empty':group['v_cuts']=[]
    elif case=='too_many':group['v_cuts']*=9
    elif case=='outside':cut['value']='2100'
    elif case=='boundary':cut['value']='0'
    elif case=='duplicate_coordinate':group['v_cuts'].append({'value':'1.2','unit':'m','evidence_id':'test'})
    elif case=='unit':cut['unit']='m²'
    elif case=='nan':cut['value']='NaN'
    elif case=='precision':cut['value']='.000000000000000001'
    elif case=='proof':cut['evidence_id']=''
    elif case=='not_list':group['u_cuts']={}
    else:cut['factor']='999'
    with pytest.raises(DomainError):partition(s,cuts)


def test_partition_face_limit_is_enforced():
    s=geometry();s['regions']=[region(str(i),i*3,0,0,1,1,1) for i in range(10)]
    cuts=[{'face_id':f['face_id'],**{a+'_cuts':[{'value':str(Decimal(f[a+'_from'])+Decimal(k)/10),'unit':'m','evidence_id':'test'} for k in range(1,9)] for a in ['u','v']}} for f in boundary(s)['faces']]
    with pytest.raises(DomainError,match='1000'):partition(s,cuts)


def test_partition_workflow_and_exports(clients):
    e,r,pid,source,body,eid=setup(clients);base='/api/v1/projects/'+pid;before=get(e,pid)
    response=e.post(base+'/surface-face-candidates/partition',json=body);assert response.status_code==200,response.text
    assert len(response.json()['faces'])==15 and get(e,pid)==before
    preview=e.post(base+'/components/preview',json=body);assert preview.status_code==200,preview.text
    assert Decimal(preview.json()['quantity'])==Decimal('4.8') and get(e,pid)==before
    comp=e.post(base+'/components',json=body);assert comp.status_code==200,comp.text
    approve(r,pid,'components',comp.json());books=e.post(base+'/takeoffs/calculate');assert books.status_code==200,books.text
    for t in books.json():approve(r,pid,'takeoffs',t)
    boq=e.post(base+'/boq/generate').json();assert {b['unit']:Decimal(b['quantity']) for b in boq}=={'m³':Decimal(21),'m²':Decimal('4.8')}
    data=e.get(base+'/exports?unpriced=true&format=json').json();assert data['unpriced']['takeoffs'][1]['surface_subdivisions']==preview.json()['surface_subdivisions']
    x=e.get(base+'/exports?unpriced=true&format=xlsx');assert x.status_code==200,x.text
    wb=load_workbook(io.BytesIO(x.content));assert '表面分区坐标' in wb.sheetnames
    assert all(c.data_type!='f' for ws in wb for row in ws for c in row)
    pdf=e.get(base+'/exports?unpriced=true&format=pdf');assert pdf.status_code==200,pdf.text
    assert '表面分区坐标' in '\n'.join(p.extract_text() or '' for p in PdfReader(io.BytesIO(pdf.content)).pages)


def test_decisions_for_old_unsplit_faces_are_rejected_atomically(clients):
    e,r,pid,source,body,eid=setup(clients);base='/api/v1/projects/'+pid
    old=base_payload(source,eid);comp=e.post(base+'/components',json=old).json();approve(r,pid,'components',comp)
    books=e.post(base+'/takeoffs/calculate').json()
    for t in books:approve(r,pid,'takeoffs',t)
    before=get(e,pid);body['object_hash']=e.get(base+'/hashes').json()['components'][comp['id']]
    bad={**body,'face_decisions':old['face_decisions']}
    assert e.put(base+'/components/'+comp['id'],json=bad).status_code==422 and get(e,pid)==before
    revised=e.put(base+'/components/'+comp['id'],json=body);assert revised.status_code==200,revised.text
    assert revised.json()['revision']==2 and get(e,pid)['components'][0]['state']=='approved'
    assert all(t['state']=='stale' for t in get(e,pid)['takeoffs'])
    assert e.put(base+'/components/'+comp['id'],json=body).status_code==409
    assert get(e,pid)['component_history'][0]['preview']['rule_version']=='solid-boundary-1.0'


@pytest.mark.parametrize('gate',['partition','save','component_review','calculate','book_review'])
def test_cut_original_corruption_blocks_every_gate(clients,gate):
    e,r,pid,source,body,eid=setup(clients);base='/api/v1/projects/'+pid
    if gate not in ('partition','save'):comp=e.post(base+'/components',json=body).json()
    if gate in ('calculate','book_review'):approve(r,pid,'components',comp)
    if gate=='book_review':book=e.post(base+'/takeoffs/calculate').json()[1]
    p=get(e,pid);doc=next(d for d in p['documents'] if d['name']=='cut-basis.txt');(store.ROOT/'originals'/doc['hash']).write_bytes(b'corrupt')
    if gate=='partition':assert e.post(base+'/surface-face-candidates/partition',json=body).status_code==409
    elif gate=='save':assert e.post(base+'/components',json=body).status_code==409
    elif gate=='component_review':approve(r,pid,'components',comp,expected=409)
    elif gate=='calculate':assert e.post(base+'/takeoffs/calculate').status_code==409
    else:approve(r,pid,'takeoffs',book,expected=409)


def test_permissions_foreign_cut_source_and_source_revision(clients):
    e,r,pid,source,body,eid=setup(clients);base='/api/v1/projects/'+pid
    assert r.post(base+'/surface-face-candidates/partition',json=body).status_code==403
    assert clients['other'].post(base+'/surface-face-candidates/partition',json=body).status_code==404
    second=create(e)['id'];upload(e,second,'其他项目分区');foreign=get(e,second)['evidence'][0]['id']
    bad=copy.deepcopy(body);bad['face_subdivisions'][0]['v_cuts'][0]['evidence_id']=foreign
    assert e.post(base+'/surface-face-candidates/partition',json=bad).status_code==404
    source['object_hash']=e.get(base+'/hashes').json()['components'][source['id']];source['regions'][0]['parameters']['height']='3'
    revised=e.put(base+'/components/'+source['id'],json=source);assert revised.status_code==200,revised.text
    approve(r,pid,'components',revised.json())
    assert e.post(base+'/surface-face-candidates/partition',json=body).status_code==409
