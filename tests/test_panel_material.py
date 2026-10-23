"""Synthetic panels; explicit areal mass is not a material acceptance."""
import copy,io,random
from decimal import Decimal,localcontext
import pytest
from openpyxl import load_workbook
from pypdf import PdfReader
from backend.engine import DomainError
from backend import store
from backend.railing_material import calculate
from tests.test_railing_material import geometry,payload,setup as material_setup
from tests.test_polygon_takeoff import post,mutate
from tests.test_workflow import get,approve,priced_project
from tests.test_route_takeoff import evaluate


def opening(name,x,y,L,W,eid='test',unit='m'):
 return dict(name=name,parameters=dict(x=str(x),y=str(y),length=str(L),width=str(W)),unit=unit,measurement_basis='合成展开矩形开口；与面密度净面积口径匹配',evidence={k:eid for k in ('x','y','length','width','measurement_basis')})


def panel(eid='test'):
 return dict(name='合成展开网片',mode='rect_panels',quantity='3',unit_weight='5000',weight_unit='g/m²',length='4000',width='2000',length_unit='mm',openings=[opening('开口A',1,.5,2,1,eid),opening('开口B',2,.5,1.5,1,eid)],measurement_basis='每组独立3片；每片8m²，开口并集2.5m²，净5.5m²；面密度按此净面积5kg/m²，微孔已含于面密度，不再次扣除',evidence={k:eid for k in ('quantity','unit_weight','length','width','measurement_basis')})


def body(s,eid='test',mixed=False):
 b=payload(s,eid);b['material_items']=(b['material_items'] if mixed else [])+[panel(eid)];b['measurement_basis']='合成板网片净展开面积质量验收；不推算片数、密度、损耗';return b


def setup(clients,pid=None):
 e,r,pid,s,_,eid=material_setup(clients,pid)
 # Use a separate original describing these exact synthetic dimensions and area basis.
 from tests.test_workflow import upload
 upload(e,pid,'合成网片每组3片，每片4000×2000mm，开口A(1,.5,2,1)m，B(2,.5,1.5,1)m；开口并集2.5m²，净5.5m²，每平方米5000g（匹配净面积口径，微孔不另扣），来源重复2，总165kg，修约3位；改4片220kg。软件合成验收，非真实工程。','panel-basis.txt')
 for c in get(e,pid)['components']:approve(r,pid,'components',c)
 s=get(e,pid)['components'][-1];eid=get(e,pid)['evidence'][-1]['id'];return e,r,pid,s,body(s,eid),eid


def test_union_panel_ast_source_count_once_and_legacy_unchanged():
 s=geometry();b=body(s);before=copy.deepcopy(b);v=calculate(b,s);assert b==before and v['quantity']=='165.000'
 row=v['material_items'][0];g=row['panel_geometry'];a=g['planar_summary']
 assert Decimal(a['net_area'])==Decimal('5.5') and Decimal(a['deduct_overlap_area'])==1
 assert Decimal(row['single_group_mass'])==Decimal('82.5') and v['rule_version']=='railing-material-mass-1.1'
 with localcontext() as c:
  c.prec=80
  assert evaluate(v['formula_ast'],v['dimensions'])==165
  assert evaluate(row['formula_ast'],v['dimensions'])==Decimal('82.5')
 assert v['dimensions']['multiplicity']=='2' and v['dimensions']['material.1.panel.multiplicity']=='1'
 assert calculate(payload(s),s)['rule_version']=='railing-material-mass-1.0'
 assert Decimal(calculate(body(s,mixed=True),s)['quantity'])==Decimal('352.2')


@pytest.mark.parametrize('unit,raw',[('kg/m²','5'),('g/m²','5000')])
@pytest.mark.parametrize('basis',['spatial','horizontal'])
def test_area_units_and_source_projection_independence(unit,raw,basis):
 s=geometry(basis=basis);b=body(s);b['material_items'][0].update(weight_unit=unit,unit_weight=raw);assert calculate(b,s)['quantity']=='165.000'


def test_nested_touching_no_openings_and_micro_precision():
 s=geometry();b=body(s);v=b['material_items'][0];v['openings'].append(opening('内嵌开口',1.2,.6,.2,.2));assert calculate(b,s)['quantity']=='165.000'
 v['openings']=[opening('全高边缘',0,0,1,2)];assert calculate(b,s)['quantity']=='180.000'
 v['openings']=[];assert calculate(b,s)['quantity']=='240.000'
 v.update(length='0.000001',width='0.000001',length_unit='m',quantity='1',unit_weight='1',weight_unit='kg/m²');b['rounding_digits']=12
 with localcontext() as c:
  c.prec=8;r=calculate(b,s)
 assert Decimal(r['quantity'])==Decimal('2e-12')


@pytest.mark.parametrize('fault',['outside','negative_x','zero_length','negative_width','unit','missing_dimension','missing_proof','extra_proof','extra_field','density','line_mass','no_openings','too_many','missing_opening','extra_opening','opening_no_proof','opening_blank_proof','opening_blank_basis','duplicate','outer_name','all_deducted','overlap_all','boolean_length','nan','precision','zero_mass','fraction_count'])
def test_strict_inputs(fault):
 s=geometry();b=body(s);v=b['material_items'][0];o=v['openings'][0]
 if fault=='outside':o['parameters']['x']='3'
 elif fault=='negative_x':o['parameters']['x']='-.1'
 elif fault=='zero_length':v['length']='0'
 elif fault=='negative_width':v['width']='-1'
 elif fault=='unit':v['length_unit']='m²'
 elif fault=='missing_dimension':v.pop('width')
 elif fault=='missing_proof':v['evidence'].pop('width')
 elif fault=='extra_proof':v['evidence']['density']='test'
 elif fault=='extra_field':v['thickness']='.003'
 elif fault=='density':v['weight_unit']='kg/m³'
 elif fault=='line_mass':v['weight_unit']='kg/m'
 elif fault=='no_openings':v['openings']=None
 elif fault=='too_many':v['openings']*=11
 elif fault=='missing_opening':o.pop('parameters')
 elif fault=='extra_opening':o['operation']='add'
 elif fault=='opening_no_proof':o['evidence'].pop('x')
 elif fault=='opening_blank_proof':o['evidence']['x']=''
 elif fault=='opening_blank_basis':o['measurement_basis']=''
 elif fault=='duplicate':v['openings'][1]['name']=o['name']
 elif fault=='outer_name':o['name']='外轮廓'
 elif fault in ('all_deducted','overlap_all'):v['openings']=[opening('整片',0,0,4,2)]+([opening('重叠',0,0,1,1)] if fault=='overlap_all' else [])
 elif fault=='boolean_length':v['length']=True
 elif fault=='nan':o['parameters']['x']='NaN'
 elif fault=='precision':v['unit_weight']='1e-18';v['weight_unit']='g/m²'
 elif fault=='zero_mass':v['unit_weight']='0'
 elif fault=='fraction_count':v['quantity']='1.5'
 with pytest.raises(DomainError):calculate(b,s)


def test_seeded_independent_cell_union_oracle():
 rng=random.Random(6801)
 for _ in range(35):
  s=geometry();b=body(s);v=b['material_items'][0];v.update(length='8',width='6',length_unit='m',quantity=str(rng.randint(1,10)),unit_weight='3',weight_unit='kg/m²');v['openings']=[];cells=set()
  for i in range(rng.randint(0,10)):
   x,y=rng.randrange(7),rng.randrange(5);L,W=rng.randint(1,8-x),rng.randint(1,6-y);v['openings'].append(opening(str(i),x,y,L,W));cells.update((a,z) for a in range(x,x+L) for z in range(y,y+W))
  if len(cells)==48:continue
  r=calculate(b,s);assert Decimal(r['quantity'])==(48-len(cells))*int(v['quantity'])*3*2


@pytest.mark.parametrize('fmt',['json','xlsx','pdf'])
def test_workflow_exports_and_old_quote(clients,fmt):
 e,r,pid,_,quote=priced_project(clients);e,r,pid,s,b,eid=setup(clients,pid);before=get(e,pid);post(e,pid,b,'/components/preview');assert before==get(e,pid)
 c=post(e,pid,b).json();approve(r,pid,'components',c)
 for t in post(e,pid,{},'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
 assert any(x['quantity']=='165.000' for x in post(e,pid,{},'/boq/generate').json())
 url=f'/api/v1/projects/{pid}/exports';res=e.get(url,params=dict(format=fmt,unpriced=True));assert res.status_code==200,res.text
 titles=['板网片净面积与质量','板网片原始尺寸与开口','板网片面积分条闭合']
 if fmt=='json':assert set(titles)<={a[0] for a in res.json()['tables']}
 elif fmt=='xlsx':assert set(titles)<=set(load_workbook(io.BytesIO(res.content)).sheetnames)
 else:assert all(v in ''.join(p.extract_text() for p in PdfReader(io.BytesIO(res.content)).pages) for v in titles)
 old=e.get(url,params=dict(format='json',quote_id=quote['id'])).json();assert old['quote']==quote and not any(t[0] in titles for t in old['tables'])


@pytest.mark.parametrize('fault',['area','strip','formula','count','opening','proof','source','book'])
def test_forgery_rebuilt_and_original_gate(clients,fault):
 e,r,pid,s,b,eid=setup(clients);c=post(e,pid,b).json()
 if fault=='book':approve(r,pid,'components',c);post(e,pid,{},'/takeoffs/calculate')
 def forge(p):
  c=p['components'][-1];v=c['preview']['material_items'][0]
  if fault=='area':v['single_panel_net_area']='999'
  elif fault=='strip':v['panel_geometry']['planar_strips'][0]['net_area']='999'
  elif fault=='formula':v['formula_ast']={'constant':'82.5'}
  elif fault=='count':c['material_items'][0]['quantity']='4'
  elif fault=='opening':c['material_items'][0]['openings'][0]['parameters']['x']='2'
  elif fault=='proof':c['material_items'][0]['openings'][0]['evidence']['x']='forged'
  elif fault=='source':c['source_geometry']['edge_decisions'][0]['mode']='exclude'
  else:p['takeoffs'][-1]['material_items'][0]['panel_geometry']['quantity']='999'
 mutate(pid,forge);coll='takeoffs' if fault=='book' else 'components';approve(r,pid,coll,get(e,pid)[coll][-1],expected=404 if fault=='proof' else 409)


def test_revision_exact_history_stale_source_and_corrupt_opening_evidence(clients):
 e,r,pid,s,b,eid=setup(clients);c=post(e,pid,b).json();approve(r,pid,'components',c)
 for t in post(e,pid,{},'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
 post(e,pid,{},'/boq/generate');old=get(e,pid)['components'][-1];new=copy.deepcopy(b);new['material_items'][0]['quantity']='4';new['object_hash']=store.digest(old)
 url=f'/api/v1/projects/{pid}/components/'+c['id'];res=e.put(url,json=new);assert res.status_code==200,res.text
 p=get(e,pid);assert p['component_history'][-1]==old and all(t['state']=='stale' for t in p['takeoffs']) and any(v['quantity']=='165.000' for v in p['boq']);assert res.json()['preview']['quantity']=='220.000' and e.put(url,json=new).status_code==409
 doc=p['documents'][-1];(store.ROOT/'originals'/doc['hash']).write_bytes(b'corrupt');post(e,pid,b,'/components/preview',409);approve(r,pid,'components',res.json(),expected=409)
