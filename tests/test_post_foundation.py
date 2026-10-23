"""Synthetic independent bases inherit final post layouts, never material mass."""
import copy,io,random
from decimal import Decimal,localcontext,ROUND_HALF_UP
import pytest
from fastapi import HTTPException
from openpyxl import load_workbook
from pypdf import PdfReader
from backend import store
from backend.engine import DomainError
from backend.post_foundation import calculate
from backend.railing_material import calculate as material_quantity,snapshot
from tests.test_railing_posts import geometry,body as equal_body,setup as post_setup
from tests.test_fixed_posts import body as fixed_body
from tests.test_additional_posts import body as fixed_added
from tests.test_equal_additional_posts import body as equal_added
from tests.test_post_adjustments import body as fixed_adjusted
from tests.test_equal_post_adjustments import body as equal_adjusted
from tests.test_combined_post_adjustments import body as combined
from tests.test_post_spacing import check
from tests.test_polygon_takeoff import post,mutate
from tests.test_route_takeoff import evaluate
from tests.test_workflow import get,approve,upload,create,priced_project

TITLES={'立柱基础体积汇总','立柱基础逐项体积','立柱基础原始尺寸与依据','立柱基础材料范围与修约'}


def source(make=equal_body):
 s=geometry();b=make(s);b.update(id='synthetic-post-material',revision=1,source_geometry=s);return b


def row(i=1,eid='test',length='400',width='500',height='600'):
 return dict(material_index=i,parameters=dict(length=length,width=width,height=height),length_unit='mm',measurement_basis='合成本项全部最终保留立柱，每根一个水平矩形独立基础，同一材料同尺寸且不重叠，无既有基础清单重复',evidence={k:eid for k in ('length','width','height','measurement_basis')})


def body(s,eid='test'):
 return dict(name='合成立柱基础（非真实工程）',kind='post_foundation',parameters={},quantity_type='volume',source_component_id=s['id'],source_component_hash=store.digest(snapshot(s)),foundation_items=[row(eid=eid)],scope_confirmed=True,material='合成C30混凝土独立矩形基础，仅软件验收，不作真实材料设计',material_evidence_id=eid,measurement_basis='合成一柱一基础，按来源最终数量，重复一次，基础独立不交叠，无孔及预埋扣减，不含垫层土方钢筋模板，未与既有构件或清单重复，非真实工程',basis_evidence_id=eid,rounding_digits=3,rounding_evidence_id=eid)


def setup(clients,pid=None):
 e,r,pid,s,b,post_eid=post_setup(clients,pid)
 upload(e,pid,'合成软件立柱基础验收，非真实工程。来源均分立柱单组11根，重复2共22根，原材质量79.2kg另计。每根一个水平矩形独立基础，400×500×600mm，同一合成C30做法，各基础独立不交叠，不含锚栓孔、预埋扣减、垫层土方钢筋模板，未与既有基础构件或清单重复，净基础2.64m3，总量修约3位。长改450mm后2.97m3，柱数及材料质量不变。来源AB间距改3000mm后单组9根、重复2共18根，400mm长基础2.16m3；须重新读取并复核。','post-foundation-basis.txt')
 for c in get(e,pid)['components']:approve(r,pid,'components',c)
 s=get(e,pid)['components'][-1];m=post(e,pid,equal_body(s,post_eid)).json();approve(r,pid,'components',m);m=get(e,pid)['components'][-1];eid=get(e,pid)['evidence'][-1]['id'];return e,r,pid,m,body(m,eid),eid


def test_count_once_exact_asts_and_source_unchanged():
 s=source();b=body(s);before=copy.deepcopy((s,b));r=calculate(b,s);d=r['foundation_items'][0];assert (s,b)==before and r['quantity']=='2.640' and r['unit']=='m³' and r['deduction']=='0' and r['rule_version']=='railing-post-foundation-1.0'
 assert r['foundation_summary']['single_group_count']==11 and r['foundation_summary']['total_count']=='22' and d['single_foundation_volume']=='0.120' and d['single_group_volume']=='1.320' and material_quantity(s)['quantity']=='79.200'
 assert evaluate(r['formula_ast'],r['dimensions'])==Decimal(r['quantity'])
 for k,a in d['formula_asts'].items():assert evaluate(a,r['dimensions'])==Decimal(d[{'single_volume':'single_foundation_volume'}.get(k,k)])
 assert r['source_geometry']==snapshot(s) and all('foundation.'+k in r['source_evidence'] for k in ('basis','material','rounding'))


@pytest.mark.parametrize('make',[equal_body,fixed_body,fixed_added,equal_added,fixed_adjusted,equal_adjusted,combined,lambda s:combined(s,equal=True)])
def test_all_eight_post_modes_inherit_final_retained_count(make):
 s=source(make);m=material_quantity(s);n=m['material_items'][0]['post_layout']['post_count'];r=calculate(body(s),s);assert r['foundation_summary']['single_group_count']==n and Decimal(r['gross'])==Decimal('.12')*n*2 and r['foundation_items'][0]['post_mode']==m['material_items'][0]['mode']


@pytest.mark.parametrize('scope',[None,'connected_path'])
def test_spacing_checks_unchanged(scope):
 s=source();s['material_items'][0]['spacing_check']=check()
 if scope:s['material_items'][0]['spacing_check']['scope']=scope
 before=copy.deepcopy(s);r=calculate(body(s),s);assert s==before and r['foundation_summary']['total_count']=='22'


@pytest.mark.parametrize('unit,factor',[('m',Decimal('.001')),('mm',Decimal('1')),('cm',Decimal('.1')),('dm',Decimal('.01')),('km',Decimal('.000001'))])
def test_length_units_and_low_context(unit,factor):
 s=source();b=body(s);d=b['foundation_items'][0];d['length_unit']=unit
 for k in d['parameters']:d['parameters'][k]=str(Decimal(d['parameters'][k])*factor)
 with localcontext() as ctx:ctx.prec=6;r=calculate(b,s)
 assert r==calculate(b,s) and r['quantity']=='2.640'


def test_partial_selection_mixed_sizes_order_and_only_final_rounding():
 s=source();s['material_items'].append({**copy.deepcopy(s['material_items'][0]),'name':'合成第二做法立柱'});b=body(s);r=calculate(b,s);assert r['foundation_summary']['unselected_post_item_count']==1
 b['foundation_items']=[row(2,length='333',width='111',height='777'),row()];r=calculate(b,s);b['foundation_items'].reverse();assert calculate(b,s)==r
 with localcontext() as ctx:
  ctx.prec=80;expected=(Decimal('.4')*Decimal('.5')*Decimal('.6')+Decimal('.333')*Decimal('.111')*Decimal('.777'))*22;assert Decimal(r['gross'])==expected and Decimal(r['quantity'])==expected.quantize(Decimal('.001'),rounding=ROUND_HALF_UP)


def test_independent_seeded_volume_oracle():
 rng=random.Random(8601)
 for _ in range(50):
  s=source();b=body(s);vals={k:str(rng.randint(1,999)) for k in ('length','width','height')};b['foundation_items'][0]['parameters']=vals;b['rounding_digits']=rng.randint(0,12)
  with localcontext() as ctx:
   ctx.prec=80;v=Decimal(vals['length'])*Decimal(vals['width'])*Decimal(vals['height'])/Decimal(10)**9*22;r=calculate(b,s);assert Decimal(r['gross'])==v and Decimal(r['quantity'])==v.quantize(Decimal(1).scaleb(-b['rounding_digits']),rounding=ROUND_HALF_UP)


@pytest.mark.parametrize('fault',['none','empty','dict','too-many','row-null','row-extra','row-missing','duplicate','unknown','bool-index','non-post','parameters-null','param-missing','param-extra','evidence-null','proof-missing','proof-extra','proof-blank','proof-bool','row-basis','negative','zero','bool','nan','infinite','precision','unit','confirm','truthy','quantity-type','top-parameters','multiplicity','rounding-bool','rounding-range','name','material','basis','material-proof','rounding-proof'])
def test_invalid_inputs(fault):
 s=source();b=body(s);d=b['foundation_items'][0]
 if fault=='none':b['foundation_items']=None
 elif fault=='empty':b['foundation_items']=[]
 elif fault=='dict':b['foundation_items']={}
 elif fault=='too-many':b['foundation_items']=[d]*51
 elif fault=='row-null':b['foundation_items']=[None]
 elif fault=='row-extra':d['quantity']='11'
 elif fault=='row-missing':d.pop('length_unit')
 elif fault=='duplicate':b['foundation_items'].append(copy.deepcopy(d))
 elif fault=='unknown':d['material_index']=2
 elif fault=='bool-index':d['material_index']=True
 elif fault=='non-post':s['material_items'][0]=dict(name='普通定长件',mode='cut_pieces',quantity='11',cut_length='1200',length_unit='mm',unit_weight='3000',weight_unit='g/m',measurement_basis='合成普通件',evidence={k:'test' for k in ('quantity','cut_length','unit_weight','measurement_basis')});b['source_component_hash']=store.digest(snapshot(s))
 elif fault=='parameters-null':d['parameters']=None
 elif fault=='param-missing':d['parameters'].pop('height')
 elif fault=='param-extra':d['parameters']['deduction']='0'
 elif fault=='evidence-null':d['evidence']=None
 elif fault=='proof-missing':d['evidence'].pop('height')
 elif fault=='proof-extra':d['evidence']['quantity']='test'
 elif fault=='proof-blank':d['evidence']['height']=''
 elif fault=='proof-bool':d['evidence']['length']=True
 elif fault=='row-basis':d['measurement_basis']=''
 elif fault in ('negative','zero','bool','nan','infinite','precision'):d['parameters']['length']={'negative':'-1','zero':'0','bool':True,'nan':'NaN','infinite':'Infinity','precision':'0.1234567890123456789'}[fault]
 elif fault=='unit':d['length_unit']='m²'
 elif fault=='confirm':b['scope_confirmed']=False
 elif fault=='truthy':b['scope_confirmed']=1
 elif fault=='quantity-type':b['quantity_type']='mass'
 elif fault=='top-parameters':b['parameters']={'count':'22'}
 elif fault=='multiplicity':b['multiplicity']='2'
 elif fault=='rounding-bool':b['rounding_digits']=True
 elif fault=='rounding-range':b['rounding_digits']=13
 elif fault=='name':b['name']=''
 elif fault=='material':b['material']=''
 elif fault=='basis':b['measurement_basis']=''
 elif fault=='material-proof':b['material_evidence_id']=''
 else:b['rounding_evidence_id']=''
 with pytest.raises(DomainError):calculate(b,s)


def test_source_hash_and_no_stale_snapshot():
 s=source();b=body(s);s['revision']=2
 with pytest.raises(HTTPException) as err:calculate(b,s)
 assert err.value.status_code==409


@pytest.mark.parametrize('fmt',['json','xlsx','pdf'])
def test_candidates_preview_review_boq_exports_and_old_quote(clients,fmt):
 e,r,pid,_,quote=priced_project(clients);e,r,pid,s,b,eid=setup(clients,pid);before=get(e,pid);url=f'/api/v1/projects/{pid}';c=e.get(url+'/post-foundation-candidates',params=dict(component_id=s['id']));assert c.status_code==200 and c.json()['post_items'][0]['post_count']==11 and c.json()['multiplicity']=='2';res=post(e,pid,b,'/components/preview');assert res.json()['quantity']=='2.640' and get(e,pid)==before
 comp=post(e,pid,{**b,'source_geometry':{'forged':True}}).json();assert comp['source_geometry']==snapshot(s) and get(e,pid)['components'][-2]==s;post(e,pid,{},'/takeoffs/calculate',409);approve(r,pid,'components',comp)
 for t in post(e,pid,{},'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
 boq=post(e,pid,{},'/boq/generate').json();assert boq[-1]['quantity']=='2.640' and boq[-1]['unit']=='m³';res=e.get(url+'/exports',params=dict(format=fmt,unpriced=True));assert res.status_code==200,res.text
 if fmt=='json':assert TITLES<={t[0] for t in res.json()['tables']}
 elif fmt=='xlsx':assert TITLES<=set(load_workbook(io.BytesIO(res.content)).sheetnames)
 else:assert all(v in ''.join(p.extract_text() for p in PdfReader(io.BytesIO(res.content)).pages) for v in TITLES)
 old=e.get(url+'/exports',params=dict(format='json',quote_id=quote['id'])).json();assert old['quote']==quote and not any(t[0] in TITLES for t in old['tables'])


def test_duplicate_claims_cannot_bypass_with_forged_id_and_revision_can_keep_claim(clients):
 e,r,pid,s,b,eid=setup(clients);c=post(e,pid,b).json();before=get(e,pid);post(e,pid,b,expected=409);post(e,pid,{**b,'id':c['id']},expected=409);assert get(e,pid)==before
 new=copy.deepcopy(b);new.update(id='forged',object_hash=store.digest(c));new['foundation_items'][0]['parameters']['length']='450';res=e.put(f'/api/v1/projects/{pid}/components/'+c['id'],json=new);assert res.status_code==200,res.text;assert res.json()['preview']['quantity']=='2.970'


def test_same_quantity_revision_and_source_chain_invalidation(clients):
 e,r,pid,s,b,eid=setup(clients);c=post(e,pid,b).json();approve(r,pid,'components',c)
 for t in post(e,pid,{},'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
 post(e,pid,{},'/boq/generate');before=get(e,pid);old=before['components'][-1];new=copy.deepcopy(b);new.update(object_hash=store.digest(old));new['foundation_items'][0]['parameters'].update(length='500',width='400');url=f'/api/v1/projects/{pid}/components/'+c['id'];res=e.put(url,json=new);assert res.status_code==200 and res.json()['preview']['quantity']=='2.640';p=get(e,pid);assert p['component_history'][-1]==old and p['boq']==before['boq'] and all(t['state']=='stale' for t in p['takeoffs']) and p['components'][-2]==s;assert e.put(url,json=new).status_code==409;approve(r,pid,'components',res.json())
 s['object_hash']=store.digest(s);s['material_items'][0]['post_ranges'][0]['spacing']='3000';res=e.put(f'/api/v1/projects/{pid}/components/'+s['id'],json=s);assert res.status_code==200,res.text;assert get(e,pid)['components'][-1]['state']=='stale';approve(r,pid,'components',res.json());post(e,pid,b,'/components/preview',409);approve(r,pid,'components',get(e,pid)['components'][-1],expected=409)
 latest=get(e,pid)['components'][-2];ca=e.get(f'/api/v1/projects/{pid}/post-foundation-candidates',params=dict(component_id=latest['id']));assert ca.json()['post_items'][0]['post_count']==9
 fresh=body(latest,eid);fresh.update(id=c['id'],object_hash=store.digest(get(e,pid)['components'][-1]));res=e.put(url,json=fresh);assert res.status_code==200 and res.json()['preview']['quantity']=='2.160'


@pytest.mark.parametrize('fault',['raw','proof','scope','source','count','shape','ast','quantity','book'])
def test_rebuild_rejects_tamper(clients,fault):
 e,r,pid,s,b,eid=setup(clients);c=post(e,pid,b).json()
 if fault=='book':approve(r,pid,'components',c);post(e,pid,{},'/takeoffs/calculate')
 def forge(p):
  c=p['components'][-1];r=c['preview']
  if fault=='raw':c['foundation_items'][0]['parameters']['height']='650'
  elif fault=='proof':c['foundation_items'][0]['evidence']['height']='forged'
  elif fault=='scope':c['scope_confirmed']=False
  elif fault=='source':c['source_geometry']['material_items'][0]['post_ranges'][0]['spacing']='3000'
  elif fault=='count':r['foundation_items'][0]['single_group_count']=12
  elif fault=='shape':r['foundation_items'][0]['normalized_parameters']['length']='.45'
  elif fault=='ast':r['formula_ast']={'constant':'2.64'}
  elif fault=='quantity':r['quantity']='2.65'
  else:p['takeoffs'][-1]['foundation_summary']['total_count']='24'
 mutate(pid,forge);coll='takeoffs' if fault=='book' else 'components';approve(r,pid,coll,get(e,pid)[coll][-1],expected=404 if fault=='proof' else 422 if fault=='scope' else 409)


@pytest.mark.parametrize('gate',['preview','save','review','calculate','book_review'])
def test_original_bytes_five_gates(clients,gate):
 e,r,pid,s,b,eid=setup(clients)
 if gate in ('review','calculate','book_review'):c=post(e,pid,b).json()
 if gate in ('calculate','book_review'):approve(r,pid,'components',c)
 if gate=='book_review':t=post(e,pid,{},'/takeoffs/calculate').json()[-1]
 d=get(e,pid)['documents'][-1];(store.ROOT/'originals'/d['hash']).write_bytes(b'corrupt')
 if gate=='preview':post(e,pid,b,'/components/preview',409)
 elif gate=='save':post(e,pid,b,expected=409)
 elif gate=='review':approve(r,pid,'components',c,expected=409)
 elif gate=='calculate':post(e,pid,{},'/takeoffs/calculate',409)
 else:approve(r,pid,'takeoffs',t,expected=409)


def test_roles_tenant_foreign_source_evidence_and_mapping_profile(clients):
 e,r,pid,s,b,eid=setup(clients);url=f'/api/v1/projects/{pid}/post-foundation-candidates';assert r.get(url,params=dict(component_id=s['id'])).status_code==200;assert clients['other'].get(url,params=dict(component_id=s['id'])).status_code==404;post(r,pid,b,expected=403)
 other=create(e)['id'];upload(e,other,'异项目');foreign=get(e,other)['evidence'][-1]['id'];bad=copy.deepcopy(b);bad['foundation_items'][0]['evidence']['height']=foreign;post(e,pid,bad,expected=404);bad=copy.deepcopy(b);bad['material_evidence_id']=foreign;post(e,pid,bad,expected=404)
 from backend.mapping_types import component_profile
 assert component_profile(b,calculate(b,s))==dict(id='volume',label='立柱独立基础实体体积（矩形/圆柱/同心台阶）',unit='m³')


def test_unapproved_and_non_post_sources_rejected(clients):
 e,r,pid,s,b,eid=setup(clients);url=f'/api/v1/projects/{pid}/post-foundation-candidates';p=get(e,pid);assert e.get(url,params=dict(component_id=p['components'][0]['id'])).status_code==409
 s['object_hash']=store.digest(s);s['measurement_basis']+='合成修订';res=e.put(f'/api/v1/projects/{pid}/components/'+s['id'],json=s);assert res.status_code==200;assert e.get(url,params=dict(component_id=s['id'])).status_code==409


def test_long_names_pdf(clients):
 e,r,pid,s,b,eid=setup(clients);b['name']='合成'+('长基础名称'*39)+'末';b['material']='合成'+('材料范围'*300);c=post(e,pid,b).json();approve(r,pid,'components',c)
 for t in post(e,pid,{},'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
 res=e.get(f'/api/v1/projects/{pid}/exports',params=dict(format='pdf',unpriced=True));assert res.status_code==200,res.text


def test_fifty_selected_items_and_1000_post_source_limit():
 s=source();s['material_items']=[{**copy.deepcopy(s['material_items'][0]),'name':'合成布置项'+str(i)} for i in range(50)];b=body(s);b['foundation_items']=[row(i+1) for i in range(50)];r=calculate(b,s);assert r['foundation_summary']['single_group_count']==550 and r['foundation_summary']['total_count']=='1100' and r['quantity']=='132.000'


def test_disjoint_claims_allowed_and_overlap_revision_rejected(clients):
 e,r,pid,s,b,eid=setup(clients);s['object_hash']=store.digest(s);s['material_items'].append({**copy.deepcopy(s['material_items'][0]),'name':'合成另一立柱做法'});res=e.put(f'/api/v1/projects/{pid}/components/'+s['id'],json=s);assert res.status_code==200;approve(r,pid,'components',res.json());s=get(e,pid)['components'][-1];b=body(s,eid);c=post(e,pid,b).json();second=body(s,eid);second['name']='合成第二基础';second['foundation_items']=[row(2,eid)];d=post(e,pid,second).json();second.update(object_hash=store.digest(d));second['foundation_items'].append(row(1,eid));assert e.put(f'/api/v1/projects/{pid}/components/'+d['id'],json=second).status_code==409


def test_ancestor_change_invalidates_material_and_foundation_transitively(clients):
 e,r,pid,s,b,eid=setup(clients);c=post(e,pid,b).json();approve(r,pid,'components',c);p=get(e,pid);route=p['components'][0];route['object_hash']=store.digest(route);route['route_points'][1]['parameters']['z']='9000';res=e.put(f'/api/v1/projects/{pid}/components/'+route['id'],json=route);assert res.status_code==200;current=get(e,pid);assert all(c['state']=='stale' for c in current['components'][1:]);assert e.get(f'/api/v1/projects/{pid}/post-foundation-candidates',params=dict(component_id=s['id'])).status_code==409


def test_tampered_source_layout_and_source_original_rejected(clients):
 e,r,pid,s,b,eid=setup(clients)
 def forge(p):p['components'][-1]['preview']['material_items'][0]['post_layout']['post_count']=12
 mutate(pid,forge);assert e.get(f'/api/v1/projects/{pid}/post-foundation-candidates',params=dict(component_id=s['id'])).status_code==409


def test_candidate_without_laid_out_posts_and_empty_read_has_no_write(clients):
 e,r,pid,s,b,eid=setup(clients);s['object_hash']=store.digest(s);s['material_items']=[dict(name='合成普通件',mode='cut_pieces',quantity='11',cut_length='1200',length_unit='mm',unit_weight='3000',weight_unit='g/m',measurement_basis='合成普通件不能推断基础',evidence={k:eid for k in ('quantity','cut_length','unit_weight','measurement_basis')})];res=e.put(f'/api/v1/projects/{pid}/components/'+s['id'],json=s);assert res.status_code==200;approve(r,pid,'components',res.json());before=get(e,pid);assert e.get(f'/api/v1/projects/{pid}/post-foundation-candidates',params=dict(component_id=s['id'])).status_code==422;assert get(e,pid)==before
