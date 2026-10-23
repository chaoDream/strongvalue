"""Rectangular recesses and mixed circle/rectangle void geometry and provenance."""
import copy,io,random
from decimal import Decimal,localcontext,ROUND_HALF_UP
import pytest
from openpyxl import load_workbook
from pypdf import PdfReader
from backend import store
from backend.compound_takeoff import PI
from backend.engine import DomainError
from backend.post_foundation import calculate
from tests.test_foundation_recesses import body as round_body,setup as round_setup,interval
from tests.test_foundation_openings import hole,SHAPES
from tests.test_post_foundation import source,equal_body,fixed_body,fixed_added,equal_added,fixed_adjusted,equal_adjusted,combined
from tests.test_route_takeoff import evaluate
from tests.test_polygon_takeoff import post,mutate
from tests.test_workflow import get,approve,upload,create,priced_project

TITLES={'基础孔洞毛扣净','基础孔洞原始依据','基础孔洞校验','基础孔洞范围与常量','基础孔洞标高依据','基础孔洞逐层相交'}
QUANTITIES=dict(zip(SHAPES,('2.612','2.348','3.704','2.322')))


def rectangle(name='合成H1',x='-80',y='0',length='100',width='60',eid='test'):
 return dict(shape='rectangular',name=name,parameters=dict(x=x,y=y,length=length,width=width),measurement_basis='合成轴向矩形竖直孔，长沿X宽沿Y，实际长宽，不旋转不重复，非真实工程。',evidence={k:eid for k in ('x','y','length','width','measurement_basis')})


def body(s,eid='test',shape='stepped_circular'):
 b=round_body(s,eid,shape);b['name']='合成基础矩形凹槽与圆孔（非真实工程）';b['measurement_basis']='合成22基础，每基础H1矩形孔长100宽60mm，长沿X宽沿Y，孔心(-80,0)mm，起止Z0至150mm；H2圆孔实际直径40mm孔心(80,0)mm，Z300至600mm，基础底Z=0向上，按实际相交层扣减，非真实工程。';d=b['foundation_items'][0];h={**rectangle(eid=eid),'vertical_range':interval(eid=eid)};h['measurement_basis']=b['measurement_basis'];c=d['through_holes']['holes'][1];c['parameters']['x']='80';c['measurement_basis']=b['measurement_basis'];d['through_holes']['holes']=[h,c];d['measurement_basis']=b['measurement_basis'];d['through_holes']['coordinate_basis']='合成共同基础中心为XY原点，X沿长度正向、Y沿宽度正向，Z为基础实体底面0向上。矩形孔长沿X宽沿Y不旋转，尺寸和标高逐项据原件。';return b


def setup(clients,pid=None,shape='stepped_circular'):
 e,r,pid,s,_,_=round_setup(clients,pid,shape)
 upload(e,pid,'合成矩形孔与圆孔软件验收，非真实工程。已审立柱单组11根、重复2，共22基础，材料79.2kg不变。基础下层圆形实际直径600mm净高200mm，上层实际直径400mm净高400mm，矩形及其它形态外尺寸沿原件。每基础H1轴向矩形孔长100宽60mm，长沿X宽沿Y不旋转，孔心(-80,0)mm，Z0至150mm；H2圆孔实际孔径40mm孔心(80,0)mm，Z300至600mm。共同中心局部XY，基础实体底面Z=0且向上；共44孔，按实际相交高度扣减.0198+.00264πm3。矩形基础净2.612m3、矩形台阶净2.348m3、圆柱净3.704m3、圆台阶净2.322m3。圆台阶验收H1可改孔心(200,0)mm，仅相交下层，矩形最远角点(250,30)mm在下层Φ600mm内；终止Z200mm仅接触层间界面允许，200.000000001mm进入Φ400mm上层且越界须拒。H1宽改80mm且仍高150mm，合计扣.0264+.00264πm3，圆台阶净2.315m3。边界验收：上层Φ400内矩形长120宽240、孔心(100,0)，最远角点(160,120)mm在圆周；100.000000001mm越界。矩形长100宽60中心0，圆孔直径40中心(70,0)相切，69.999999999为重叠；圆孔直径100中心(80,70)与矩形角点(50,30)相切。矩形100×60与40×40孔心相距70mm在X方向相切。独立软件验收不代表结构净距、真实工程或合同扣减规则。TEST-VOIDS仅合成映射编码，非官方条目。','foundation-voids-basis.txt')
 for c in get(e,pid)['components']:approve(r,pid,'components',c)
 s=get(e,pid)['components'][-1];eid=get(e,pid)['evidence'][-1]['id'];return e,r,pid,s,body(s,eid,shape),eid


@pytest.mark.parametrize('shape',SHAPES)
def test_mixed_shape_balances_all_formula_asts_and_immutability(shape):
 s=source();b=body(s,shape=shape);before=copy.deepcopy((b,s));r=calculate(b,s);assert (b,s)==before and r['quantity']==QUANTITIES[shape] and r['rule_version']=='railing-post-foundation-1.5' and r['foundation_summary']['total_count']=='22' and r['foundation_opening_summary']['total_hole_count']=='44'
 with localcontext() as ctx:
  ctx.prec=80;assert Decimal(r['deduction'])==Decimal('.0198')+Decimal(PI)*Decimal('.00264');assert Decimal(r['gross'])-Decimal(r['deduction'])==Decimal(r['foundation_summary']['unrounded_quantity']);assert evaluate(r['formula_ast'],r['dimensions'])==Decimal(r['quantity'])
  for d in r['foundation_items']:
   for k,a in d['formula_asts'].items():assert evaluate(a,r['dimensions'])==Decimal(d[{'single_volume':'single_foundation_volume'}.get(k,k)])
   for layer in d.get('foundation_levels',[]):
    for k,a in layer['formula_asts'].items():assert evaluate(a,r['dimensions'])==Decimal(layer[k])
   for h in d['opening_holes']:
    assert evaluate(h['height_ast'],r['dimensions'])==Decimal(h['through_height']) and evaluate(h['formula_ast'],r['dimensions'])==Decimal(h['single_volume'])
    for k,a in h['vertical_formula_asts'].items():assert evaluate(a,r['dimensions'])==Decimal(h[k])
    for seg in h['segments']:
     for k,a in seg['formula_asts'].items():assert evaluate(a,r['dimensions'])==Decimal(seg[k])
    for c in h['containments']:
     for k,a in c['formula_asts'].items():assert evaluate(a,r['dimensions'])==Decimal(c['margins'][k])>=0
   for c in d['opening_checks']:
    assert evaluate(c['formula_ast'],r['dimensions'])==Decimal(c['margin']) and evaluate(c['vertical_gap_ast'],r['dimensions'])==Decimal(c['vertical_gap'])
    for k,a in c.get('nearest_formula_asts',{}).items():assert evaluate(a,r['dimensions'])==Decimal(c['nearest_distances'][k])


@pytest.mark.parametrize('shape',SHAPES)
@pytest.mark.parametrize('make',[equal_body,fixed_body,fixed_added,equal_added,fixed_adjusted,equal_adjusted,combined,lambda s:combined(s,equal=True)])
def test_final_post_count_modes(shape,make):
 s=source(make);r=calculate(body(s,shape=shape),s)
 with localcontext() as ctx:ctx.prec=80;assert Decimal(r['deduction'])==(Decimal('.0009')+Decimal(PI)*Decimal('.00012'))*r['foundation_summary']['single_group_count']*2


@pytest.mark.parametrize('shape',SHAPES)
@pytest.mark.parametrize('unit,factor',[('m',Decimal('.001')),('mm',Decimal(1)),('cm',Decimal('.1')),('dm',Decimal('.01')),('km',Decimal('.000001'))])
def test_independent_units_and_low_decimal_context(shape,unit,factor):
 s=source();b=body(s,shape=shape);c=b['foundation_items'][0]['through_holes'];c['unit']=unit
 for h in c['holes']:
  for p in (h['parameters'],h['vertical_range']['parameters']):
   for k in p:p[k]=str(Decimal(p[k])*factor)
 with localcontext() as ctx:ctx.prec=6;r=calculate(b,s)
 assert r==calculate(b,s) and r['quantity']==QUANTITIES[shape]


@pytest.mark.parametrize('shape',SHAPES)
@pytest.mark.parametrize('delta,allowed',[('0',True),('-.000000001',True),('.000000001',False)])
def test_rectangular_containment_in_round_and_rectangular_outer(shape,delta,allowed):
 s=source();b=body(s,shape=shape);d=b['foundation_items'][0];h=d['through_holes']['holes'][0];d['through_holes']['holes']=[h]
 if shape in ('circular','stepped_circular'):
  h['parameters'].update(x=str(Decimal('100')+Decimal(delta)),length='280' if shape=='circular' else '120',width='360' if shape=='circular' else '240')
 else:h['parameters']['x']=str(Decimal('150')+Decimal(delta))
 if shape.startswith('stepped'):h['vertical_range']=interval('300','400')
 if allowed:
  row=calculate(b,s)['foundation_items'][0]['opening_holes'][0];assert row['containments'][0]['state']==('tangent' if delta=='0' else 'inside')
 else:
  with pytest.raises(DomainError,match='越出'):calculate(b,s)


@pytest.mark.parametrize('end,allowed',[('200',True),('199.999999999',True),('200.000000001',False)])
def test_lower_only_rectangular_recess_does_not_require_upper_containment(end,allowed):
 s=source();b=body(s);h=b['foundation_items'][0]['through_holes']['holes'][0];h['parameters']['x']='200';h['vertical_range']['parameters']['z_end']=end
 if allowed:assert [v['level_index'] for v in calculate(b,s)['foundation_items'][0]['opening_holes'][0]['segments']]==[1]
 else:
  with pytest.raises(DomainError,match='第2层'):calculate(b,s)


@pytest.mark.parametrize('kind',['rectangle_rectangle','circle_rectangle_side','circle_rectangle_corner'])
@pytest.mark.parametrize('reverse',[False,True])
@pytest.mark.parametrize('delta,allowed',[('0',True),('.000000001',True),('-.000000001',False)])
def test_pair_contact_exact_disjointness_and_overlap(kind,reverse,delta,allowed):
 s=source();b=body(s);d=b['foundation_items'][0];a=rectangle(x='0');x=Decimal('80' if kind.endswith('corner') else '70')+Decimal(delta)
 q=rectangle('H2',x=str(x),length='40',width='40') if kind=='rectangle_rectangle' else hole('H2',x=str(x),y='70' if kind.endswith('corner') else '0',diameter='100' if kind.endswith('corner') else '40');d['through_holes']['holes']=[a,q][::(-1 if reverse else 1)]
 if allowed:
  r=calculate(b,s);c=r['foundation_items'][0]['opening_checks'][0];assert c['state']==('tangent' if delta=='0' else 'disjoint')
  with localcontext() as ctx:
   ctx.prec=80;assert evaluate(c['formula_ast'],r['dimensions'])==Decimal(c['margin'])>=0
   for k,node in c.get('axis_formula_asts',{}).items():assert evaluate(node,r['dimensions'])==Decimal(c['axis_margins'][k])
 else:
  with pytest.raises(DomainError,match='重叠'):calculate(b,s)


@pytest.mark.parametrize('sx,sy',[(-1,-1),(-1,1),(1,-1),(1,1)])
def test_circle_rectangle_nearest_corner_all_quadrants(sx,sy):
 s=source();b=body(s);d=b['foundation_items'][0];d['through_holes']['holes']=[rectangle(x='0'),hole('C',x=str(sx*80),y=str(sy*70),diameter='100')];r=calculate(b,s);c=r['foundation_items'][0]['opening_checks'][0];assert c['state']=='tangent' and Decimal(c['margin'])==0


@pytest.mark.parametrize('end,state',[('299.999999999','vertically_disjoint'),('300','vertical_tangent'),('300.000000001',None)])
def test_overlapping_xy_only_blocked_for_positive_height_overlap(end,state):
 s=source();b=body(s);a,q=b['foundation_items'][0]['through_holes']['holes'];a['parameters']['x']=q['parameters']['x']='0';a['vertical_range']['parameters']['z_end']=end
 if state:assert calculate(b,s)['foundation_items'][0]['opening_checks'][0]['state']==state
 else:
  with pytest.raises(DomainError,match='重叠'):calculate(b,s)


def test_only_rectangles_have_no_pi_and_mix_with_solid_and_legacy_rows():
 s=source();b=body(s,shape='rectangular');d=b['foundation_items'][0];d['through_holes']['holes'][1]={**rectangle('H2',x='80',length='40',width='50'),'vertical_range':interval('300','600')};r=calculate(b,s);assert r['quantity']=='2.607' and 'foundation_pi_constant' not in r
 from tests.test_post_foundation import body as solid_body
 s['material_items'] += [{**copy.deepcopy(s['material_items'][0]),'name':name} for name in ('合成旧圆孔项','合成无孔项')];b=body(s,shape='rectangular');b['foundation_items'][0]['through_holes']=copy.deepcopy(d['through_holes']);old=round_body(s)['foundation_items'][0];old['material_index']=2;solid=solid_body(s)['foundation_items'][0];solid['material_index']=3;b['foundation_items'] += [old,solid];r=calculate(b,s);assert r['foundation_summary']['total_count']=='66' and r['foundation_opening_summary']['total_hole_count']=='88' and 'foundation_pi_constant' in r and 'shape' not in r['foundation_items'][1]['opening_holes'][0] and 'opening_holes' not in r['foundation_items'][2]


def test_twenty_finite_rectangles_ten_layers_and_only_total_rounding():
 s=source();b=body(s);d=b['foundation_items'][0];d['levels']=[dict(parameters=dict(diameter='400',height='100'),evidence=dict(diameter='test',height='test')) for _ in range(10)];d['through_holes']['holes']=[{**rectangle('H'+str(i),x='0',length='10',width='20'),'vertical_range':interval(str(i*50),str((i+1)*50))} for i in range(20)];r=calculate(b,s);assert r['foundation_opening_summary']['total_hole_count']=='440' and len(r['foundation_items'][0]['opening_checks'])==190
 with localcontext() as ctx:ctx.prec=80;assert Decimal(r['deduction'])==Decimal('.0044')


def test_seeded_independent_rectangular_clipping_and_rounding():
 rng=random.Random(9101)
 for shape in SHAPES:
  for _ in range(30):
   s=source();b=body(s,shape=shape);d=b['foundation_items'][0];a=rng.randint(0,598);end=rng.randint(a+1,600);length=rng.randint(1,100);width=rng.randint(1,100);h=d['through_holes']['holes'][0];h['vertical_range']=interval(str(a),str(end));h['parameters'].update(length=str(length),width=str(width));d['through_holes']['holes']=[h];digits=rng.randint(0,12);b['rounding_digits']=digits;plain=copy.deepcopy(b);plain['foundation_items'][0].pop('through_holes');r=calculate(b,s);gross=calculate(plain,s)['gross']
   with localcontext() as ctx:ctx.prec=80;ded=Decimal(length)*width*(end-a)/1000000000*22;assert Decimal(r['deduction'])==ded and Decimal(r['quantity'])==(Decimal(gross)-ded).quantize(Decimal(1).scaleb(-digits),rounding=ROUND_HALF_UP)


@pytest.mark.parametrize('fault',['shape','shape-null','shape-bool','angle','diameter','length-missing','width-zero','length-negative','length-bool','width-nan','width-infinite','precision','proof-missing','proof-extra','proof-blank','proof-bool','range-null','range-extra','range-proof-blank','negative-z','reverse-z','outside-z','scope','unit','empty','duplicate','extra-config'])
def test_strict_schema_actual_dimensions_and_evidence(fault):
 s=source();b=body(s);c=b['foundation_items'][0]['through_holes'];h=c['holes'][0]
 if fault=='shape':h['shape']='rotated_rectangle'
 elif fault=='shape-null':h['shape']=None
 elif fault=='shape-bool':h['shape']=True
 elif fault=='angle':h['parameters']['angle']='0'
 elif fault=='diameter':h['parameters']['diameter']='60'
 elif fault=='length-missing':h['parameters'].pop('length')
 elif fault in ('width-zero','length-negative','length-bool','width-nan','width-infinite','precision'):h['parameters']['width' if fault.startswith('width') else 'length']={'width-zero':'0','length-negative':'-1','length-bool':True,'width-nan':'NaN','width-infinite':'Infinity','precision':'.1234567890123456789'}[fault]
 elif fault=='proof-missing':h['evidence'].pop('length')
 elif fault=='proof-extra':h['evidence']['angle']='test'
 elif fault=='proof-blank':h['evidence']['width']=''
 elif fault=='proof-bool':h['evidence']['length']=True
 elif fault=='range-null':h['vertical_range']=None
 elif fault=='range-extra':h['vertical_range']['depth']='150'
 elif fault=='range-proof-blank':h['vertical_range']['evidence']['z_end']=''
 elif fault=='negative-z':h['vertical_range']['parameters']['z_start']='-1'
 elif fault=='reverse-z':h['vertical_range']['parameters'].update(z_start='200',z_end='150')
 elif fault=='outside-z':h['vertical_range']['parameters']['z_end']='600.000000001'
 elif fault=='scope':c['scope_confirmed']=1
 elif fault=='unit':c['unit']={}
 elif fault=='empty':c['holes']=[]
 elif fault=='duplicate':c['holes'][1]['name']=h['name']
 else:c['count']='22'
 with pytest.raises(DomainError):calculate(b,s)


@pytest.mark.parametrize('shape',SHAPES)
@pytest.mark.parametrize('fmt',['json','xlsx','pdf'])
def test_workflow_net_boq_exports_mapping_and_old_quote(clients,shape,fmt):
 e,r,pid,_,quote=priced_project(clients);e,r,pid,s,b,eid=setup(clients,pid,shape);before=get(e,pid);assert post(e,pid,b,'/components/preview').json()['quantity']==QUANTITIES[shape] and get(e,pid)==before;c=post(e,pid,b).json();approve(r,pid,'components',c)
 for t in post(e,pid,{},'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
 boq=post(e,pid,{},'/boq/generate').json();assert any(v['quantity']==QUANTITIES[shape] and v['unit']=='m³' for v in boq) and any(v['quantity']=='79.200' for v in boq);url=f'/api/v1/projects/{pid}/exports';res=e.get(url,params=dict(format=fmt,unpriced=True));assert res.status_code==200,res.text
 if fmt=='json':assert TITLES<={v[0] for v in res.json()['tables']}
 elif fmt=='xlsx':
  w=load_workbook(io.BytesIO(res.content));assert TITLES<=set(w.sheetnames) and w['基础孔洞标高依据'].row_dimensions[3].height>48
 else:assert all(v in ''.join(p.extract_text() for p in PdfReader(io.BytesIO(res.content)).pages) for v in TITLES)
 assert not any(v[0] in TITLES for v in e.get(url,params=dict(format='json',quote_id=quote['id'])).json()['tables'])
 if fmt=='json':
  from tests.test_mapping_profiles import rule_body,save_rule,preview,selection
  p=get(e,pid);rule={**rule_body(p,'post_foundation','volume',eid),'code':'TEST-VOIDS'};assert e.patch(f'/api/v1/projects/{pid}/settings',json=dict(boq_standard=rule['standard'],boq_edition=rule['edition'],reason='仅合成矩形凹槽与圆孔验收')).status_code==200;save_rule(e,r,pid,rule);bid=next(v['id'] for v in p['boq'] if v['takeoff_id']==p['takeoffs'][-1]['id']);chosen=preview(e,pid,bid);assert chosen['state']=='unique';assert e.post(f'/api/v1/projects/{pid}/boq/apply-mappings',json=dict(selections=[selection(chosen)])).status_code==200;approve(r,pid,'boq',next(v for v in get(e,pid)['boq'] if v['id']==bid))


@pytest.mark.parametrize('fault',['shape','length','width-proof','corner','segment','pair','pair-ast','unit','book'])
def test_tamper_reconstruction(clients,fault):
 e,r,pid,s,b,eid=setup(clients);c=post(e,pid,b).json()
 if fault=='book':approve(r,pid,'components',c);post(e,pid,{},'/takeoffs/calculate')
 def forge(p):
  c=p['components'][-1];d=c['preview']['foundation_items'][0];h=d['opening_holes'][0]
  if fault=='shape':c['foundation_items'][0]['through_holes']['holes'][0]['shape']='rotated'
  elif fault=='length':c['foundation_items'][0]['through_holes']['holes'][0]['parameters']['length']='120'
  elif fault=='width-proof':c['foundation_items'][0]['through_holes']['holes'][0]['evidence']['width']='forged'
  elif fault=='corner':h['containments'][0]['margins']['corner_squared']='0'
  elif fault=='segment':h['segments'][0]['single_volume']='0'
  elif fault=='pair':d['opening_checks'][0]['margin']='0'
  elif fault=='pair-ast':d['opening_checks'][0]['formula_ast']={'constant':'0'}
  elif fault=='unit':d['opening_checks'][0]['margin_unit']='m'
  else:p['takeoffs'][-1]['deduction']='0'
 mutate(pid,forge);collection='takeoffs' if fault=='book' else 'components';approve(r,pid,collection,get(e,pid)[collection][-1],expected=404 if fault=='width-proof' else 422 if fault=='shape' else 409)


@pytest.mark.parametrize('gate',['preview','save','review','calculate','book_review'])
def test_rectangular_original_bytes_gates(clients,gate):
 e,r,pid,s,b,eid=setup(clients)
 if gate in ('review','calculate','book_review'):c=post(e,pid,b).json()
 if gate in ('calculate','book_review'):approve(r,pid,'components',c)
 if gate=='book_review':t=post(e,pid,{},'/takeoffs/calculate').json()[-1]
 doc=get(e,pid)['documents'][-1];(store.ROOT/'originals'/doc['hash']).write_bytes(b'corrupt')
 if gate=='preview':post(e,pid,b,'/components/preview',409)
 elif gate=='save':post(e,pid,b,expected=409)
 elif gate=='review':approve(r,pid,'components',c,expected=409)
 elif gate=='calculate':post(e,pid,{},'/takeoffs/calculate',409)
 else:approve(r,pid,'takeoffs',t,expected=409)


def test_width_revision_history_conflict_switch_back_and_source_stale(clients):
 e,r,pid,s,b,eid=setup(clients);c=post(e,pid,b).json();approve(r,pid,'components',c)
 for t in post(e,pid,{},'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
 post(e,pid,{},'/boq/generate');p=get(e,pid);old=p['components'][-1];new=copy.deepcopy(b);new['foundation_items'][0]['through_holes']['holes'][0]['parameters']['width']='80';new['object_hash']=store.digest(old);url=f'/api/v1/projects/{pid}/components/'+old['id'];res=e.put(url,json=new);assert res.status_code==200 and res.json()['preview']['quantity']=='2.315';assert get(e,pid)['component_history'][-1]==old and get(e,pid)['boq']==p['boq'] and all(t['state']=='stale' for t in get(e,pid)['takeoffs']);assert e.put(url,json=new).status_code==409;post(e,pid,{**b,'id':old['id']},expected=409);approve(r,pid,'components',res.json())
 new=round_body(s,eid);new['object_hash']=store.digest(get(e,pid)['components'][-1]);res=e.put(url,json=new);assert res.status_code==200 and res.json()['preview']['quantity']=='2.337' and res.json()['preview']['rule_version']=='railing-post-foundation-1.4';approve(r,pid,'components',res.json())
 s['object_hash']=store.digest(s);s['material_items'][0]['post_ranges'][0]['spacing']='3000';res=e.put(f'/api/v1/projects/{pid}/components/'+s['id'],json=s);assert res.status_code==200;approve(r,pid,'components',res.json());assert get(e,pid)['components'][-1]['state']=='stale'


@pytest.mark.parametrize('field',['length','width','z_end'])
def test_roles_tenants_and_foreign_rectangle_evidence(clients,field):
 e,r,pid,s,b,eid=setup(clients);post(r,pid,b,expected=403);assert clients['other'].get(f'/api/v1/projects/{pid}').status_code==404;other=create(e)['id'];upload(e,other,'异项目矩形孔原件');fid=get(e,other)['evidence'][-1]['id'];h=b['foundation_items'][0]['through_holes']['holes'][0];(h['vertical_range']['evidence'] if field=='z_end' else h['evidence'])[field]=fid;post(e,pid,b,expected=404)


def test_long_names_twenty_rectangular_depths_pdf(clients):
 e,r,pid,s,b,eid=setup(clients);b['name']='合成'+('矩形凹槽长名称'*20);d=b['foundation_items'][0];d['through_holes']['holes']=[{**rectangle('凹槽'+str(i)+('原件说明'*12),x='0',length='10',width='20',eid=eid),'vertical_range':interval(str(i*30),str((i+1)*30),eid)} for i in range(20)];c=post(e,pid,b).json();approve(r,pid,'components',c)
 for t in post(e,pid,{},'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
 res=e.get(f'/api/v1/projects/{pid}/exports',params=dict(format='pdf',unpriced=True));assert res.status_code==200,res.text


@pytest.mark.parametrize('delta,allowed',[('0',True),('.000000001',True),('-.000000001',False)])
def test_rectangle_y_axis_contact_and_circle_pairs_inside_mixed_table(delta,allowed):
 s=source();b=body(s);d=b['foundation_items'][0];d['through_holes']['holes']=[rectangle(x='0'),rectangle('H2',x='0',y=str(Decimal('50')+Decimal(delta)),length='40',width='40')]
 if allowed:assert calculate(b,s)['foundation_items'][0]['opening_checks'][0]['state']==('tangent' if delta=='0' else 'disjoint')
 else:
  with pytest.raises(DomainError,match='重叠'):calculate(b,s)
 d['through_holes']['holes']=[rectangle(x='-130',length='20',width='20'),hole('C1',x='-50'),hole('C2',x=str(Decimal('-10')+Decimal(delta)))]
 if allowed:
  r=calculate(b,s);c=next(v for v in r['foundation_items'][0]['opening_checks'] if v['separation_kind']=='circle_circle');assert c['state']==('tangent' if delta=='0' else 'disjoint')
 else:
  with pytest.raises(DomainError,match='重叠'):calculate(b,s)


@pytest.mark.parametrize('fmt',['json','xlsx','pdf'])
def test_all_rectangular_geometry_exports_without_pi(clients,fmt):
 e,r,pid,s,b,eid=setup(clients,shape='rectangular');d=b['foundation_items'][0];d['through_holes']['holes'][1]={**rectangle('H2',x='80',length='40',width='50',eid=eid),'vertical_range':interval('300','600',eid)};c=post(e,pid,b).json();assert 'foundation_pi_constant' not in c['preview'];approve(r,pid,'components',c)
 for t in post(e,pid,{},'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
 res=e.get(f'/api/v1/projects/{pid}/exports',params=dict(format=fmt,unpriced=True));assert res.status_code==200,res.text
 if fmt=='json':assert TITLES<={v[0] for v in res.json()['tables']} and any('仅矩形几何' in str(v) for v in res.json()['tables'])


def test_rectangular_full_layer_erasure_and_partial_height_recess():
 s=source();b=body(s,shape='rectangular');d=b['foundation_items'][0];d['through_holes']['holes']=[rectangle(x='0',length='400',width='500')]
 with pytest.raises(DomainError,match='每层净体积'):calculate(b,s)
 d['through_holes']['holes'][0]['vertical_range']=interval('0','150');r=calculate(b,s);assert r['quantity']=='1.980'
