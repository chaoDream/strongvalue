"""Synthetic developed polygons; no real material or engineering acceptance."""
import copy,io,random
from decimal import Decimal,localcontext,ROUND_HALF_UP
import pytest
from openpyxl import load_workbook
from pypdf import PdfReader
from backend.engine import DomainError
from backend import store
from backend.railing_material import calculate
from tests.test_railing_material import geometry,payload,setup as material_setup
from tests.test_panel_material import panel as rectangular_panel
from tests.test_polygon_takeoff import ring,OUTER,TRIANGLE,DIAMOND,post,mutate
from tests.test_workflow import get,approve,upload,priced_project,create
from tests.test_route_takeoff import evaluate


def panel(eid='test'):
 return dict(name='合成异形展开板片',mode='polygon_panels',quantity='3',unit_weight='2500',weight_unit='g/m²',polygon_rings=[ring('凹形外轮廓',OUTER,eid),ring('三角孔',TRIANGLE,eid,'hole'),ring('菱形孔',DIAMOND,eid,'hole')],coordinate_basis='每片独立展开平面；原点外轮廓首点，X右Y上，非建筑水平投影，非测量图',measurement_basis='独立3片，展开26m²扣独立三角孔和菱形孔各0.5m²，净25m²；2.5kg/m²匹配此净面积，微孔不另扣',evidence={k:eid for k in ('quantity','unit_weight','coordinate_basis','measurement_basis')})


def body(s,eid='test',mixed=False):
 b=payload(s,eid);b['material_items']=(b['material_items']+[rectangular_panel(eid)] if mixed else [])+[panel(eid)];b['measurement_basis']='合成异形板片净展开面积质量，无损耗、搭接或折弯推算';return b


def setup(clients,pid=None):
 e,r,pid,s,_,eid=material_setup(clients,pid)
 upload(e,pid,'合成异形板片软件验收，非真实工程：每片独立展开局部XY，原点(0,0)，X右Y上。外轮廓(0,0)(6,0)(6,3)(4,3)(4,5)(0,5)m，面积26m²；三角孔(1,1)(2,1)(1,2)m，0.5m²；菱形孔(3,2.5)(3.5,3)(3,3.5)(2.5,3)m，0.5m²，孔严格在内且彼此分离。净25m²，每组3片，2500g/m²匹配净面积，微孔不另扣，重复2，总375kg，修约3位；改4片500kg。不推断损耗或折弯。','polygon-panel-basis.txt')
 for c in get(e,pid)['components']:approve(r,pid,'components',c)
 s=get(e,pid)['components'][-1];eid=get(e,pid)['evidence'][-1]['id'];return e,r,pid,s,body(s,eid),eid


def test_concave_sloped_holes_unrounded_net_area_and_namespaced_asts():
 s=geometry();b=body(s);before=copy.deepcopy((s,b));v=calculate(b,s);assert (s,b)==before and v['quantity']=='375.000' and v['rule_version']=='railing-material-mass-1.2'
 row=v['material_items'][0];g=row['panel_geometry'];a=g['polygon_summary'];assert (Decimal(a['outer_area']),Decimal(a['hole_area']),Decimal(a['net_area']))==(26,1,25)
 assert a['rounding_digits'] is None and a['multiplicity']=='1' and v['dimensions']['multiplicity']=='2' and v['dimensions']['material.1.panel.multiplicity']=='1'
 with localcontext() as c:
  c.prec=80;assert evaluate(v['formula_ast'],v['dimensions'])==375
  assert evaluate(g['formula_ast'],v['dimensions'])==25
  assert evaluate(row['formula_ast'],v['dimensions'])==Decimal('187.5')
  for ring in g['polygon_rings']:
   assert evaluate(ring['formula_ast'],v['dimensions'])==Decimal(ring['area'])
   for edge in ring['edges']:assert evaluate(edge['formula_ast'],v['dimensions'])==Decimal(edge['cross_product'])
 assert Decimal(calculate(body(s,mixed=True),s)['quantity'])==Decimal('727.2')


@pytest.mark.parametrize('basis',['spatial','horizontal'])
@pytest.mark.parametrize('unit,raw',[('kg/m²','2.5'),('g/m²','2500')])
def test_source_projection_and_areal_mass_units(basis,unit,raw):
 s=geometry(basis=basis);b=body(s);b['material_items'][0].update(weight_unit=unit,unit_weight=raw);assert calculate(b,s)['quantity']=='375.000'


def test_microscopic_area_not_rounded_before_multiplication_and_large_origin():
 s=geometry();b=body(s);v=b['material_items'][0];v.update(polygon_rings=[ring('tiny',[(0,0),('1e-7',0),(0,'1e-7')])],quantity='1000000',unit_weight='1',weight_unit='kg/m²');b['rounding_digits']=12
 assert Decimal(calculate(b,s)['quantity'])==Decimal('1e-8')
 with localcontext() as c:
  c.prec=80;o=Decimal('1000000000');d=Decimal('1e-8');v['polygon_rings']=[ring('tiny offset',[(o,o),(o+d,o),(o+d,o+d),(o,o+d)])]
 with localcontext() as c:c.prec=8;r=calculate(b,s)
 assert Decimal(r['quantity'])==Decimal('2e-10')


@pytest.mark.parametrize('mutation',['reverse','cycle','translate','rotate','mixed_units'])
def test_geometry_invariants(mutation):
 s=geometry();b=body(s)
 for r in b['material_items'][0]['polygon_rings']:
  if mutation=='reverse':r['points'].reverse()
  if mutation=='cycle':r['points']=r['points'][1:]+r['points'][:1]
  for p in r['points']:
   x,y=(Decimal(p['parameters'][k]) for k in ('x','y'))
   if mutation=='translate':x,y=x-100,y+50
   if mutation=='rotate':x,y=-y,x
   if mutation=='mixed_units':x,y=x*1000,y*1000;p['unit']='mm'
   p['parameters']=dict(x=str(x),y=str(y))
 assert calculate(b,s)['quantity']=='375.000'


@pytest.mark.parametrize('fault',['missing_rings','rings_type','malformed_ring','empty','holes_limit','point_limit','total_points','self_cross','duplicate_point','touch_outer','outside','cross_concavity','touch_holes','overlap_holes','nested_holes','second_outer','first_hole','missing_coordinate','missing_coordinate_proof','blank_coordinate','extra_rect_field','line_mass','density','missing_point_proof','foreign_point_axis','bad_point_unit','nonfinite','zero_area','duplicate_ring','duplicate_point_name','string_confirmation'])
def test_strict_input_and_geometry_gates(fault):
 s=geometry();b=body(s);v=b['material_items'][0];r=v['polygon_rings'][0];p=r['points'][0]
 if fault=='missing_rings':v.pop('polygon_rings')
 elif fault=='rings_type':v['polygon_rings']={}
 elif fault=='malformed_ring':v['polygon_rings']=[True]
 elif fault=='empty':v['polygon_rings']=[]
 elif fault=='holes_limit':v['polygon_rings']*=8
 elif fault=='point_limit':r['points']*=17
 elif fault=='total_points':v['polygon_rings']=[ring(str(i),[(0,0),(1,0),(1,1),(0,1)]*20,role='outer' if i==0 else 'hole') for i in range(3)]
 elif fault=='self_cross':v['polygon_rings']=[ring('bowtie',[(0,0),(2,2),(2,0),(0,2)])]
 elif fault=='duplicate_point':r['points'][-1]['parameters']=p['parameters'].copy()
 elif fault=='touch_outer':v['polygon_rings'][1]=ring('touch',[(0,1),(1,1),(1,2)],role='hole')
 elif fault=='outside':v['polygon_rings'][1]=ring('outside',[(8,1),(9,1),(8,2)],role='hole')
 elif fault=='cross_concavity':v['polygon_rings'][1]=ring('cross',[(5,2),(5,4),(3,4)],role='hole')
 elif fault=='touch_holes':v['polygon_rings'][2]=ring('touch',[(1,1),(.5,1),(1,.5)],role='hole')
 elif fault=='overlap_holes':v['polygon_rings'][2]=ring('overlap',[(1,1),(2,1),(2,2),(1,2)],role='hole')
 elif fault=='nested_holes':v['polygon_rings'][2]=ring('nested',[(1.1,1.1),(1.3,1.1),(1.1,1.3)],role='hole')
 elif fault=='second_outer':v['polygon_rings'][1]['role']='outer'
 elif fault=='first_hole':r['role']='hole'
 elif fault=='missing_coordinate':v.pop('coordinate_basis')
 elif fault=='missing_coordinate_proof':v['evidence'].pop('coordinate_basis')
 elif fault=='blank_coordinate':v['coordinate_basis']=''
 elif fault=='extra_rect_field':v['length']='100'
 elif fault=='line_mass':v['weight_unit']='kg/m'
 elif fault=='density':v['weight_unit']='kg/m³'
 elif fault=='missing_point_proof':p['evidence'].pop('x')
 elif fault=='foreign_point_axis':p['parameters']['z']='0'
 elif fault=='bad_point_unit':p['unit']='m²'
 elif fault=='nonfinite':p['parameters']['x']='NaN'
 elif fault=='zero_area':v['polygon_rings']=[ring('flat',[(0,0),(1,0),(2,0)])]
 elif fault=='duplicate_ring':v['polygon_rings'][1]['name']=r['name']
 elif fault=='duplicate_point_name':r['points'][1]['name']=p['name']
 elif fault=='string_confirmation':b['scope_confirmed']='true'
 with pytest.raises(DomainError):calculate(b,s)


def test_polygon_point_budget_across_material_rows():
 s=geometry();b=body(s);v=b['material_items'][0];coords=[(i,0) for i in range(50)]+[(i,2) for i in reversed(range(50))];v['polygon_rings']=[ring('many',coords)]
 b['material_items']=[{**copy.deepcopy(v),'name':str(i)} for i in range(10)];assert Decimal(calculate(b,s)['quantity'])>0
 b['material_items'].append({**copy.deepcopy(v),'name':'11'})
 with pytest.raises(DomainError,match='1000'):calculate(b,s)


def test_seeded_independent_triangle_mass_oracle():
 rng=random.Random(6901)
 for _ in range(35):
  s=geometry();b=body(s);v=b['material_items'][0];L,W=rng.randint(3,100),rng.randint(3,100);x,y=rng.randint(-100,100),rng.randint(-100,100);v.update(polygon_rings=[ring('triangle',[(x,y),(x+L,y),(x,y+W)])],quantity=str(rng.randint(1,30)),unit_weight=str(rng.randint(1,100)),weight_unit='kg/m²')
  expect=Decimal(L)*W/2*Decimal(v['quantity'])*Decimal(v['unit_weight'])*2
  assert Decimal(calculate(b,s)['quantity'])==expect


@pytest.mark.parametrize('fmt',['json','xlsx','pdf'])
def test_workflow_mixed_exports_and_old_quote_isolation(clients,fmt):
 e,r,pid,_,quote=priced_project(clients);e,r,pid,s,b,eid=setup(clients,pid);b=body(s,eid,mixed=True);before=get(e,pid);post(e,pid,b,'/components/preview');assert get(e,pid)==before
 c=post(e,pid,b).json();approve(r,pid,'components',c)
 for t in post(e,pid,{},'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
 assert any(v['quantity']=='727.200' for v in post(e,pid,{},'/boq/generate').json())
 url=f'/api/v1/projects/{pid}/exports';res=e.get(url,params=dict(format=fmt,unpriced=True));assert res.status_code==200,res.text
 titles=['异形板网片面积与质量','异形板网片原始顶点','异形板网片逐边面积','异形板网片孔边界校验','板网片净面积与质量','栏杆材料逐项计算']
 if fmt=='json':assert set(titles)<={t[0] for t in res.json()['tables']}
 elif fmt=='xlsx':assert set(titles)<=set(load_workbook(io.BytesIO(res.content)).sheetnames)
 else:assert all(v in ''.join(p.extract_text() for p in PdfReader(io.BytesIO(res.content)).pages) for v in titles)
 old=e.get(url,params=dict(format='json',quote_id=quote['id'])).json();assert old['quote']==quote and not any(t[0] in titles for t in old['tables'])


@pytest.mark.parametrize('fault',['area','point','ring_ast','edge_ast','check','count','source','book'])
def test_saved_and_book_tampering_rebuilt(clients,fault):
 e,r,pid,s,b,eid=setup(clients);c=post(e,pid,b).json()
 if fault=='book':approve(r,pid,'components',c);post(e,pid,{},'/takeoffs/calculate')
 def forge(p):
  c=p['components'][-1];v=c['preview']['material_items'][0];g=v['panel_geometry']
  if fault=='area':v['single_panel_net_area']='999'
  elif fault=='point':c['material_items'][0]['polygon_rings'][0]['points'][1]['parameters']['x']='7'
  elif fault=='ring_ast':g['polygon_rings'][0]['formula_ast']={'constant':'26'}
  elif fault=='edge_ast':g['polygon_rings'][0]['edges'][0]['formula_ast']={'constant':'0'}
  elif fault=='check':g['polygon_checks'][0]['passed']=False
  elif fault=='count':c['material_items'][0]['quantity']='4'
  elif fault=='source':c['source_geometry']['edge_decisions'][0]['mode']='exclude'
  else:p['takeoffs'][-1]['material_items'][0]['panel_geometry']['quantity']='999'
 mutate(pid,forge);coll='takeoffs' if fault=='book' else 'components';approve(r,pid,coll,get(e,pid)[coll][-1],expected=409)


def test_revision_source_invalidation_history_concurrency_and_new_boq(clients):
 e,r,pid,s,b,eid=setup(clients);c=post(e,pid,b).json();approve(r,pid,'components',c)
 for t in post(e,pid,{},'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
 post(e,pid,{},'/boq/generate');old=get(e,pid)['components'][-1];new=copy.deepcopy(b);new['material_items'][0]['quantity']='4';new['object_hash']=store.digest(old);url=f'/api/v1/projects/{pid}/components/'+c['id'];res=e.put(url,json=new);assert res.status_code==200,res.text
 p=get(e,pid);assert p['component_history'][-1]==old and all(t['state']=='stale' for t in p['takeoffs']) and any(v['quantity']=='375.000' for v in p['boq']);assert res.json()['preview']['quantity']=='500.000' and e.put(url,json=new).status_code==409;approve(r,pid,'components',res.json())
 for t in post(e,pid,{},'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
 assert any(v['quantity']=='500.000' for v in post(e,pid,{},'/boq/generate').json())
 from tests.test_railing_takeoff import interval
 s['object_hash']=store.digest(s);s['edge_decisions'][0].update(mode='ranges',ranges=[interval(0,10,eid)]);res=e.put(f'/api/v1/projects/{pid}/components/'+s['id'],json=s);assert res.status_code==200,res.text;approve(r,pid,'components',res.json());assert get(e,pid)['components'][-1]['state']=='stale';post(e,pid,b,'/components/preview',409)


@pytest.mark.parametrize('gate',['preview','save','review','calculate','book_review'])
def test_original_bytes_gates(clients,gate):
 e,r,pid,s,b,eid=setup(clients)
 if gate in ('review','calculate','book_review'):c=post(e,pid,b).json()
 if gate in ('calculate','book_review'):approve(r,pid,'components',c)
 if gate=='book_review':t=post(e,pid,{},'/takeoffs/calculate').json()[-1]
 doc=get(e,pid)['documents'][-1];(store.ROOT/'originals'/doc['hash']).write_bytes(b'corrupt')
 if gate=='preview':post(e,pid,b,'/components/preview',409)
 elif gate=='save':post(e,pid,b,expected=409)
 elif gate=='review':approve(r,pid,'components',c,expected=409)
 elif gate=='calculate':post(e,pid,{},'/takeoffs/calculate',409)
 else:approve(r,pid,'takeoffs',t,expected=409)


def test_roles_tenant_foreign_point_and_coordinate_proofs(clients):
 e,r,pid,s,b,eid=setup(clients);post(r,pid,b,expected=403);assert clients['other'].get(f'/api/v1/projects/{pid}').status_code==404
 other=create(e)['id'];upload(e,other,'其他工程合成原件');foreign=get(e,other)['evidence'][-1]['id']
 for target in ('point','coordinate'):
  bad=copy.deepcopy(b)
  if target=='point':bad['material_items'][0]['polygon_rings'][0]['points'][0]['evidence']['x']=foreign
  else:bad['material_items'][0]['evidence']['coordinate_basis']=foreign
  post(e,pid,bad,expected=404)
