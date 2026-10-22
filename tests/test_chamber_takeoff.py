"""Synthetic circular chamber barrels; no real design, site or material suitability claim."""
import copy
import io
import json
import random
from decimal import Decimal, localcontext
import pytest
from fastapi import HTTPException
from openpyxl import load_workbook
from pypdf import PdfReader
from backend import store
from backend.compound_takeoff import PI
from backend.chamber_takeoff import calculate
from backend.engine import DomainError
from backend.route_takeoff import snapshot
from backend.volume_balance import snapshot as balance_snapshot, calculate as balance_calculate
from tests.test_trench_takeoff import geometry, setup as trench_setup, evaluate
from tests.test_volume_balance import payload as balance_payload
from tests.test_workflow import get, upload, approve, create, priced_project


def well(point=1,eid='test',outer='2',inner='1.6',height='3',count='1',unit='m'):
    return {'name':'合成W'+str(point)+'井筒段','point_index':point,'parameters':{'outer_diameter':outer,'inner_diameter':inner,'height':height},'unit':unit,
            'multiplicity':count,'evidence':{k:eid for k in ['outer_diameter','inner_diameter','height','multiplicity','measurement_basis']},
            'measurement_basis':'仅合成：关联点作定位索引，实际井筒高度及独立井数有据核定，不从管线高程推断；非整井'}


def payload(source,eid='test',purpose='outer_envelope'):
    return {'name':'合成圆形井筒（非真实工程）','kind':'circular_chamber','parameters':{},'quantity_type':'volume','purpose':purpose,
            'source_component_id':source['id'],'source_component_hash':store.digest(snapshot(source)),
            'chamber_wells':[well(eid=eid)],'locations_confirmed':True,'material':'合成井筒材料（非工程选材）','material_evidence_id':eid,
            'construction_basis':'仅合成竖直等内外径井筒，无收口、底板、盖板或洞口；不代表施工选型','construction_evidence_id':eid,
            'measurement_basis':'仅合成井筒段，材料与含空腔外包用途分别核定，数量独立，不从管线重复量推算，回填范围另审','basis_evidence_id':eid,
            'rounding_digits':6,'rounding_evidence_id':eid}


def setup(clients,save=False,purpose='outer_envelope'):
    e,r,pid,route,trench,eid=trench_setup(clients);base='/api/v1/projects/'+pid
    upload(e,pid,'合成W1井筒段：站点P1仅作定位索引，实际外径2m、内径1.6m、高3m、独立井数1。竖直等径，材料仅合成，不含底盖板收口和洞口。按实际内外径计算材料/含空腔外包体积，总量修约6位；回填范围需另审。',name='chamber.txt')
    approve(r,pid,'components',get(e,pid)['components'][0]);route=get(e,pid)['components'][0]
    eid=get(e,pid)['evidence'][-1]['id'];b=payload(route,eid,purpose)
    if save:
        response=e.post(base+'/components',json=b);assert response.status_code==200,response.text
    return e,r,pid,route,trench,b,eid


@pytest.mark.parametrize('purpose,expected',[('wall_material','3.392920'),('outer_envelope','9.424778')])
def test_distinct_purposes_raw_parameters_and_ast(purpose,expected):
    s=geometry();b=payload(s,purpose=purpose);before=copy.deepcopy((s,b));r=calculate(b,s)
    assert (s,b)==before and r['quantity']==expected and r['unit']=='m³'
    assert r['chamber_wells'][0]['location']['name']==s['route_points'][0]['name']
    with localcontext() as ctx:
        ctx.prec=80
        assert Decimal(r['chamber_summary']['wall_volume'])==Decimal(PI)*Decimal('1.08')
        assert Decimal(r['chamber_summary']['outer_volume'])==Decimal(PI)*3
        assert evaluate(r['formula_ast'],r['dimensions'])==Decimal(expected)
        assert evaluate(r['chamber_wells'][0]['formula_ast'],r['dimensions'])==Decimal(r['gross'])


def test_mixed_units_multiwell_counts_do_not_inherit_route_count_or_high_elevation():
    s=geometry();s['multiplicity']='7';s['route_points'][1]['parameters']['z']='99000';b=payload(s)
    b['chamber_wells'][0]['multiplicity']='2';b['chamber_wells'].append(well(2,outer='1200',inner='1000',height='2000',unit='mm'))
    r=calculate(b,s);assert r['quantity']=='21.111503' and r['chamber_summary']['well_count']=='3' and r['chamber_summary']['route_multiplicity_not_inherited']=='7'
    assert r['chamber_wells'][1]['normalized_parameters']=={'outer_diameter':'1.2','inner_diameter':'1','height':'2'}
    b['purpose']='wall_material';assert calculate(b,s)['quantity']=='7.476991'


def test_seeded_cylinders_match_independent_decimal_volume_difference():
    rng=random.Random(3401)
    for _ in range(30):
        s=geometry();s['multiplicity']=str(rng.randint(2,9));b=payload(s)
        d=Decimal(rng.randint(1,50))/10;t=Decimal(rng.randint(1,50))/100;outer=d+2*t;h=Decimal(rng.randint(1,99))/10;n=rng.randint(1,5)
        b['chamber_wells']=[well(outer=str(outer),inner=str(d),height=str(h),count=str(n))];b['rounding_digits']=12
        with localcontext() as ctx:
            ctx.prec=80
            for purpose in ['wall_material','outer_envelope']:
                b['purpose']=purpose;r=calculate(b,s)
                expected=(Decimal(PI)*outer*outer/4*h-(Decimal(PI)*d*d/4*h if purpose=='wall_material' else 0))*n
                assert Decimal(r['quantity'])==expected.quantize(Decimal('.000000000001'))
                assert evaluate(r['formula_ast'],r['dimensions'])==Decimal(r['quantity'])


@pytest.mark.parametrize('case',['empty','many','duplicate_name','duplicate_point','unknown_point','bool_point','extra','missing_height','missing_proof','zero_inner','negative','nan','inverted','unit','precision','count_zero','count_fraction','count_over','no_material','no_construction','blank_scope','unconfirmed','string_confirmation','global_count','rounding','mode','parameters','purpose','purpose_object','overflow'])
def test_invalid_inputs_fail(case):
    s=geometry();b=payload(s);w=b['chamber_wells'][0]
    if case=='empty':b['chamber_wells']=[]
    elif case=='many':b['chamber_wells']*=51
    elif case=='duplicate_name':b['chamber_wells'].append({**well(2),'name':w['name']})
    elif case=='duplicate_point':b['chamber_wells'].append({**well(),'name':'other'})
    elif case=='unknown_point':w['point_index']=99
    elif case=='bool_point':w['point_index']=True
    elif case=='extra':w['depth_from_pipe']='1'
    elif case=='missing_height':del w['parameters']['height']
    elif case=='missing_proof':del w['evidence']['multiplicity']
    elif case=='zero_inner':w['parameters']['inner_diameter']='0'
    elif case=='negative':w['parameters']['height']='-1'
    elif case=='nan':w['parameters']['outer_diameter']='NaN'
    elif case=='inverted':w['parameters']['inner_diameter']='2'
    elif case=='unit':w['unit']='m³'
    elif case=='precision':w['unit']='mm';w['parameters']['height']='.000000000000000001'
    elif case=='count_zero':w['multiplicity']='0'
    elif case=='count_fraction':w['multiplicity']='1.5'
    elif case=='count_over':w['multiplicity']='1000001'
    elif case=='no_material':b['material']=''
    elif case=='no_construction':b['construction_basis']=''
    elif case=='blank_scope':w['measurement_basis']=''
    elif case=='unconfirmed':b['locations_confirmed']=False
    elif case=='string_confirmation':b['locations_confirmed']='true'
    elif case=='global_count':b['multiplicity']='7'
    elif case=='rounding':b['rounding_digits']=True
    elif case=='mode':b['quantity_type']='area'
    elif case=='parameters':b['parameters']={'volume':'1'}
    elif case=='purpose':b['purpose']='entire_well'
    elif case=='purpose_object':b['purpose']={}
    elif case=='overflow':w['parameters']={'outer_diameter':'1000000000000','inner_diameter':'1','height':'1000000000000'}
    with pytest.raises(DomainError):calculate(b,s)


def test_old_source_version_fails():
    s=geometry();b=payload(s);s['revision']=2
    with pytest.raises(HTTPException) as err:calculate(b,s)
    assert err.value.status_code==409


def test_candidates_preview_review_books_boq_and_exports(clients):
    e,r,pid,route,trench,b,eid=setup(clients);base='/api/v1/projects/'+pid;before=get(e,pid)
    url=base+'/chamber-location-candidates?component_id='+route['id'];c=e.get(url);assert c.status_code==200,c.text
    assert c.json()['route_points'][0]['index']==1 and get(e,pid)==before
    response=e.post(base+'/components/preview',json=b);assert response.status_code==200,response.text
    assert response.json()['quantity']=='9.424778' and get(e,pid)==before
    response=e.post(base+'/components',json={**b,'source_geometry':{'forged':True}});assert response.status_code==200,response.text
    comp=response.json();assert comp['source_geometry']==snapshot(route) and get(e,pid)['components'][0]['state']=='approved'
    assert e.post(base+'/takeoffs/calculate').status_code==409;approve(r,pid,'components',comp)
    response=e.post(base+'/takeoffs/calculate');assert response.status_code==200,response.text
    for book in response.json():approve(r,pid,'takeoffs',book)
    response=e.post(base+'/boq/generate');assert response.status_code==200,response.text
    assert response.json()[-1]['quantity']=='9.424778'
    j=e.get(base+'/exports?unpriced=true&format=json').json()['unpriced']['takeoffs'][-1]
    assert j['chamber_wells']==comp['preview']['chamber_wells'] and j['source_geometry']==snapshot(route)
    response=e.get(base+'/exports?unpriced=true&format=xlsx');assert response.status_code==200,response.text
    wb=load_workbook(io.BytesIO(response.content));assert {'检查井井筒计量汇总','井筒逐井尺寸与数量','井筒定位及原始依据','井筒材料做法与规则'}<=set(wb.sheetnames)
    assert all(c.data_type!='f' for s in wb for row in s for c in row)
    response=e.get(base+'/exports?unpriced=true&format=pdf');assert response.status_code==200,response.text
    text='\n'.join(p.extract_text() for p in PdfReader(io.BytesIO(response.content)).pages)
    assert '检查井井筒计量汇总' in text and '井筒逐井尺寸与数量' in text and '9.424778' in text


def test_material_rejected_as_backfill_occupation_outer_envelope_links_and_invalidates(clients):
    e,r,pid,route,trench,b,eid=setup(clients,save=True);base='/api/v1/projects/'+pid;outer=get(e,pid)['components'][-1];approve(r,pid,'components',outer)
    material=copy.deepcopy(b);material.update(name='仅合成井筒材料',purpose='wall_material');m=e.post(base+'/components',json=material);assert m.status_code==200,m.text
    approve(r,pid,'components',m.json());t=e.post(base+'/components',json=trench);assert t.status_code==200,t.text;approve(r,pid,'components',t.json())
    p=get(e,pid);ex=balance_snapshot(p['components'][-1]);ma=balance_snapshot(p['components'][2]);out=balance_snapshot(p['components'][1])
    selector={'excavation_component_id':ex['id'],'occupation_component_ids':[ma['id']]}
    assert e.post(base+'/volume-balance-candidates',json=selector).status_code==409
    bad_src={'excavation':ex,'occupations':[ma]};bad=balance_payload(bad_src,eid)
    with pytest.raises(DomainError):balance_calculate(bad,bad_src)
    src={'excavation':ex,'occupations':[out]};balance=balance_payload(src,eid);balance['rounding_digits']=6
    response=e.post(base+'/components',json=balance);assert response.status_code==200,response.text
    assert response.json()['preview']['quantity']=='38.075222';approve(r,pid,'components',response.json())
    books=e.post(base+'/takeoffs/calculate');assert books.status_code==200,books.text
    for book in books.json():approve(r,pid,'takeoffs',book)
    before=get(e,pid);assert e.put(base+'/components/'+outer['id'],json=b).status_code==409 and get(e,pid)==before
    b['object_hash']=e.get(base+'/hashes').json()['components'][outer['id']];bad=copy.deepcopy(b);bad['chamber_wells'][0]['parameters']['inner_diameter']='3'
    assert e.put(base+'/components/'+outer['id'],json=bad).status_code==422 and get(e,pid)==before
    b['chamber_wells'][0]['parameters']['height']='4';response=e.put(base+'/components/'+outer['id'],json=b);assert response.status_code==200,response.text
    assert response.json()['revision']==2 and response.json()['preview']['quantity']=='12.566371'
    p=get(e,pid);assert [c['state'] for c in p['components']]==['approved','draft','approved','approved','stale']
    assert all(t['state']=='stale' for t in p['takeoffs']) and e.post(base+'/takeoffs/calculate').status_code==409
    approve(r,pid,'components',p['components'][1]);approve(r,pid,'components',p['components'][-1],expected=409)


@pytest.mark.parametrize('gate',['preview','save','component_review','calculate','book_review'])
def test_original_corruption_blocks_all_gates(clients,gate):
    e,r,pid,route,trench,b,eid=setup(clients);base='/api/v1/projects/'+pid
    if gate not in ('preview','save'):comp=e.post(base+'/components',json=b).json()
    if gate in ('calculate','book_review'):approve(r,pid,'components',comp)
    if gate=='book_review':book=e.post(base+'/takeoffs/calculate').json()[-1]
    doc=get(e,pid)['documents'][-1];(store.ROOT/'originals'/doc['hash']).write_bytes(b'corrupt')
    if gate=='preview':assert e.post(base+'/components/preview',json=b).status_code==409
    elif gate=='save':assert e.post(base+'/components',json=b).status_code==409
    elif gate=='component_review':approve(r,pid,'components',comp,expected=409)
    elif gate=='calculate':assert e.post(base+'/takeoffs/calculate').status_code==409
    else:approve(r,pid,'takeoffs',book,expected=409)


def test_permissions_cross_project_and_modified_source(clients):
    e,r,pid,route,trench,b,eid=setup(clients);base='/api/v1/projects/'+pid;url=base+'/chamber-location-candidates?component_id='+route['id']
    assert r.get(url).status_code==200 and r.post(base+'/components',json=b).status_code==403 and clients['other'].get(url).status_code==404
    other=create(e)['id'];upload(e,other,'其他项目井筒依据');bad=copy.deepcopy(b);bad['chamber_wells'][0]['evidence']['height']=get(e,other)['evidence'][0]['id']
    assert e.post(base+'/components',json=bad).status_code==404
    comp=e.post(base+'/components',json=b).json();approve(r,pid,'components',comp)
    route['object_hash']=e.get(base+'/hashes').json()['components'][route['id']];route['route_points'][0]['parameters']['z']='-5'
    assert e.put(base+'/components/'+route['id'],json=route).status_code==200
    assert e.get(url).status_code==409 and e.post(base+'/components/preview',json=b).status_code==409
    approve(r,pid,'components',get(e,pid)['components'][0]);assert e.post(base+'/components/preview',json=b).status_code==409
    changed={'kind':'ring','parameters':{},'object_hash':e.get(base+'/hashes').json()['components'][comp['id']]}
    assert e.put(base+'/components/'+comp['id'],json=changed).status_code==409


@pytest.mark.parametrize('target',['component','book'])
def test_saved_output_tampering_blocks_review(clients,target):
    e,r,pid,route,trench,b,eid=setup(clients,save=True);base='/api/v1/projects/'+pid
    if target=='book':approve(r,pid,'components',get(e,pid)['components'][-1]);e.post(base+'/takeoffs/calculate')
    with store.transaction() as db:
        p=json.loads(db.execute('SELECT data FROM projects WHERE id=?',(pid,)).fetchone()['data'])
        if target=='component':p['components'][-1]['preview']['chamber_wells'][0]['single_wall_volume']='999'
        else:p['takeoffs'][-1]['chamber_wells'][0]['single_wall_volume']='999'
        store.save(db,p)
    col='components' if target=='component' else 'takeoffs';approve(r,pid,col,get(e,pid)[col][-1],expected=409)


def test_new_chamber_does_not_rewrite_historical_quote(clients):
    e,r,pid,_,quote=priced_project(clients);old=copy.deepcopy(quote);base='/api/v1/projects/'+pid;eid=get(e,pid)['evidence'][0]['id']
    route=geometry(eid);route.pop('id');route.pop('revision');response=e.post(base+'/components',json=route);assert response.status_code==200,response.text
    approve(r,pid,'components',response.json());source=get(e,pid)['components'][-1]
    response=e.post(base+'/components',json=payload(source,eid));assert response.status_code==200,response.text
    p=get(e,pid);assert p['quotes'][0]==old and old['id'] in p['stale_quotes']
