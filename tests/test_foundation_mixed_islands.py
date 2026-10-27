"""Synthetic mixed circle/polygon islands; independent geometry and workflow evidence."""
import copy,io,itertools,random
from decimal import Decimal,localcontext
from fractions import Fraction
import pytest
from openpyxl import load_workbook
from pypdf import PdfReader
from backend import store
from backend.engine import DomainError
from backend.compound_takeoff import PI
from backend.foundation_horizontal_formwork import calculate
from backend.foundation_mixed_islands import VERSION,METHOD,validate,separation,part_config,topology
from tests.test_foundation_circular_islands import block as cbk,parent,config as cic,setup as old_setup
from tests.test_foundation_circular_holes_windows import inner,config as hc,replay as replay_circles
from tests.test_foundation_circular_windows import body as cbody,config as cc,independent_area
from tests.test_foundation_eccentric_windows import config as ec
from tests.test_foundation_annular_windows import config as ac
from tests.test_foundation_horizontal_windows import source,window
from tests.test_foundation_hole_end_formwork import source as end_source
from tests.test_foundation_multi_windows import replay_blocks,config as mc,part
from tests.test_foundation_nested_windows import config as nc,block as pbk,box
from tests.test_foundation_perforated_windows import config as pc,excluded
from tests.test_foundation_simple_windows import C_SHAPE
from tests.test_foundation_openings import hole
from tests.test_foundation_voids import rectangle
from tests.test_foundation_recesses import interval
from tests.test_post_foundation import equal_body,fixed_body,fixed_added,equal_added,fixed_adjusted,equal_adjusted,combined
from tests.test_polygon_takeoff import post,mutate
from tests.test_workflow import get,approve,upload,create,priced_project


def circle(eid='test',**kw):return dict(cbk(eid,**kw),shape='circular')
def poly(pts,eid='test',name='合成直边块',holes=(),relation=None):return dict(pbk(pts,eid,name,holes,relation),shape='polygon')
def config(eid='test',direction='circle_parent',parts=None,unit='mm'):
    c=cic(eid,unit=unit);c['shape']='nested_mixed';c['parts']=parts if parts is not None else ([circle(eid,holes=[inner(eid,10,0,120,'父圆W')]),poly(box(-20,20,-20,20),eid,'直边岛',relation=parent())] if direction=='circle_parent' else [poly(box(-100,100,-100,100),eid,holes=[box(-60,60,-60,60)]),circle(eid,x=10,diameter=40,name='圆形岛',parent=parent())]);return c

def body(s,eid='test',face='bottom',direction='circle_parent',parts=None):
    b=cbody(s,eid,face);b['contact_ranges'][0]['plan_range']=config(eid,direction,parts);return b

def external_config(c,p):
    cfg=part_config(c,p)
    if p['shape']!='circular' or len(p['holes'])>=2:return cfg
    q=p['parameters']
    if not p['holes']:return {k:v for k,v in cfg.items() if k!='holes'}|dict(shape='circular')
    h=p['holes'][0];v=h['parameters'];e=p['evidence'];f=h['evidence']
    return {k:cfg[k] for k in ('unit','coordinate_basis','coordinate_evidence_id','scope_confirmed')}|dict(shape='eccentric_circular',parameters=dict(x=q['x'],y=q['y'],outer_diameter=q['diameter'],inner_x=v['x'],inner_y=v['y'],inner_diameter=v['diameter']),evidence=dict(x=e['x'],y=e['y'],outer_diameter=e['diameter'],inner_x=f['x'],inner_y=f['y'],inner_diameter=f['diameter']),exclusion_basis=h['measurement_basis'],exclusion_evidence_id=h['evidence_id'])

def exact(n,dims):
    if 'constant' in n:return Fraction(n['constant'])
    if 'dimension' in n:return Fraction(dims[n['dimension']])
    vals=[exact(v,dims) for v in n['args']]
    if n['op']=='add':return sum(vals,Fraction(0))
    if n['op']=='subtract':return vals[0]-vals[1]
    if n['op']=='multiply':
        r=Fraction(1)
        for v in vals:r*=v
        return r
    assert n['op']=='divide';return vals[0]/vals[1]

def replay(r):
    replay_blocks(r);replay_circles({**r,'contact_ranges':[{**d,**p} for d in r['contact_ranges'] for p in d['window_parts']]})
    for d in r['contact_ranges']:
        for q in d['window_part_checks']:
            for v in q['proof']['conditions']:
                for k,n in v['formula_asts'].items():assert exact(n,r['dimensions'])==Fraction(v[k])
                a,b=Fraction(v['value']),Fraction(v['limit']);assert a<b if v['operator']=='<' else a>b
                w=v.get('edge_witness')
                if w:
                    for k,n in w['formula_asts'].items():assert exact(n,r['dimensions'])==Fraction(w[k])
                    assert 0<=Fraction(w['segment_parameter'])<=1


@pytest.mark.parametrize('direction',['circle_parent','polygon_parent'])
@pytest.mark.parametrize('shape',['rectangular','circular','stepped_rectangular','stepped_circular'])
@pytest.mark.parametrize('face',['top','bottom','hole_top','hole_bottom'])
@pytest.mark.parametrize('make',[equal_body,fixed_body,fixed_added,equal_added,fixed_adjusted,equal_adjusted,combined,lambda s:combined(s,equal=True)])
def test_shapes_ends_counts_immutable_all_source_and_rational_ast_replay(direction,shape,face,make):
    s=end_source(shape,make) if face.startswith('hole') else source(shape,make);b=body(s,face=face,direction=direction);before=copy.deepcopy((s,b));r=calculate(b,s);d=r['contact_ranges'][0]
    assert before==(s,b) and r['rule_version']==VERSION and d['nested_mixed_window'];assert d['window_nesting']==[dict(block_index=1,parent=None,depth=0),dict(block_index=2,parent=parent(),depth=1)]
    with localcontext() as ctx:
        ctx.prec=80;expected=Decimal(PI)*Decimal('.0064')+Decimal('.0016') if direction=='circle_parent' else Decimal('.0256')+Decimal(PI)*Decimal('.0004');assert abs(Decimal(d['window_area'])-expected)<Decimal('1e-78')
    replay(r)


@pytest.mark.parametrize('direction',['circle_parent','polygon_parent'])
@pytest.mark.parametrize('shape,face',[(s,f) for s in ['rectangular','circular','stepped_rectangular','stepped_circular'] for f in ['top','bottom','hole_bottom','hole_top']]+[('stepped_rectangular','step'),('stepped_circular','step')])
def test_independent_old_modes_sum_full_gross_hole_net_without_double_count(direction,shape,face):
    s=end_source(shape) if face.startswith('hole') else source(shape);b=body(s,face=face,direction=direction);r=calculate(b,s);c=b['contact_ranges'][0]['plan_range'];parts=[]
    for p in c['parts']:
        child=copy.deepcopy(b);child['contact_ranges'][0]['plan_range']=external_config(c,p);parts.append(calculate(child,s)['contact_ranges'][0])
    with localcontext() as ctx:
        ctx.prec=80
        for k in ['single_gross_area','single_opening_area','single_net_area','total_net_area']:assert abs(Decimal(r['contact_ranges'][0][k])-sum(Decimal(p[k]) for p in parts))<Decimal('1e-73')
    replay(r)


@pytest.mark.parametrize('direction',['circle_parent','polygon_parent'])
@pytest.mark.parametrize('shape',['rectangular','circular','stepped_rectangular','stepped_circular'])
@pytest.mark.parametrize('x',[-300,-140,0,140,300])
def test_40_independent_numerical_integrals_with_actual_hole(direction,shape,x):
    s=source(shape);s['foundation_items'][0]['through_holes']['holes']=[{**rectangle(x=150,length=100,width=60),'vertical_range':interval('0','150')}]
    ps=[circle(x=x,y=30,diameter=260,holes=[inner('test',x+15,20,140,'W')]),poly(box(x-20,x+20,10,50),name='直边岛',relation=parent())] if direction=='circle_parent' else [poly(box(x-130,x+130,-100,160),holes=[box(x-70,x+70,-50,90)]),circle(x=x+10,y=20,diameter=60,name='圆形岛',parent=parent())]
    r=calculate(body(s,parts=ps),s);ds=[(0.,0.,.3)] if 'circular' in shape else [(0.,0.,2.)];bs=[] if 'circular' in shape else [(-.2,.2,-.25,.25)] if shape=='rectangular' else [(-.3,.3,-.25,.25)];H=[(.1,.2,-.03,.03)]
    if direction=='circle_parent':
        a=independent_area(ds+[(x/1000,.03,.13)],bs,excluded_disks=[((x+15)/1000,.02,.07)],excluded_boxes=H);b=independent_area(ds,bs+[((x-20)/1000,(x+20)/1000,.01,.05)],excluded_boxes=H)
    else:
        a=independent_area(ds,bs+[((x-130)/1000,(x+130)/1000,-.1,.16)],excluded_boxes=H+[((x-70)/1000,(x+70)/1000,-.05,.09)]);b=independent_area(ds+[((x+10)/1000,.02,.03)],bs,excluded_boxes=H)
    assert abs(float(r['contact_ranges'][0]['single_net_area'])-a-b)<3e-10;replay(r)


@pytest.mark.parametrize('direction',['circle_parent','polygon_parent'])
@pytest.mark.parametrize('shape',['stepped_rectangular','stepped_circular'])
@pytest.mark.parametrize('x',[0,170,280,700])
def test_16_platform_numeric_integrals(direction,shape,x):
    s=source(shape);s['foundation_items'][0].pop('through_holes',None);ps=[circle(x=x,diameter=220,holes=[inner('test',x,0,120,'W')]),poly(box(x-20,x+20,-20,20),name='直边岛',relation=parent())] if direction=='circle_parent' else [poly(box(x-110,x+110,-110,110),holes=[box(x-60,x+60,-60,60)]),circle(x=x,diameter=40,name='圆形岛',parent=parent())]
    r=calculate(body(s,face='step',parts=ps),s);ds=[(0.,0.,.3)] if shape=='stepped_circular' else [(0.,0.,2.)];bs=[] if shape=='stepped_circular' else [(-.3,.3,-.25,.25)];xs=[(0.,0.,.2)] if shape=='stepped_circular' else [];xb=[] if shape=='stepped_circular' else [(-.2,.2,-.15,.15)]
    if direction=='circle_parent':a=independent_area(ds+[(x/1000,0.,.11)],bs,excluded_disks=xs+[(x/1000,0.,.06)],excluded_boxes=xb);b=independent_area(ds,bs+[((x-20)/1000,(x+20)/1000,-.02,.02)],excluded_disks=xs,excluded_boxes=xb)
    else:a=independent_area(ds,bs+[((x-110)/1000,(x+110)/1000,-.11,.11)],excluded_disks=xs,excluded_boxes=xb+[((x-60)/1000,(x+60)/1000,-.06,.06)]);b=independent_area(ds+[(x/1000,0.,.02)],bs,excluded_disks=xs,excluded_boxes=xb)
    assert abs(float(r['contact_ranges'][0]['single_net_area'])-a-b)<3e-10;replay(r)


def chain():
    return [circle(holes=[inner('test',0,0,160,'W1')]),poly(box(-50,50,-50,50),name='B2',holes=[box(-30,30,-30,30)],relation=parent()),circle(diameter=40,name='B3',holes=[inner('test',0,0,30,'W3')],parent=parent(2)),poly(box(-5,5,-5,5),name='B4',relation=parent(3))]


def test_four_alternating_levels_all_24_orders_and_rational_dependencies():
    s=source();base=chain();quantity=None
    for order in itertools.permutations(range(4)):
        ps=copy.deepcopy([base[i] for i in order]);indices={old+1:new+1 for new,old in enumerate(order)}
        for p in ps:
            if p['parent']:p['parent']['block_index']=indices[p['parent']['block_index']]
        r=calculate(body(s,parts=ps),s);assert [h['depth'] for h in r['contact_ranges'][0]['window_nesting']]==list(order)
        if quantity is None:quantity=r['quantity']
        assert r['quantity']==quantity;replay(r)


@pytest.mark.parametrize('gap',['-0.000000001','0','0.000000001'])
def test_polygon_vertex_on_circle_3_4_5_boundary_strict_no_tolerance(gap):
    v=Decimal(30)+Decimal(gap);c=config(parts=[circle(holes=[inner('test',0,0,100,'W')]),poly(box(-v,v,-40,40),name='直边岛',relation=parent())])
    if Decimal(gap)<0:validate(c)
    else:
        with pytest.raises(DomainError):validate(c)


@pytest.mark.parametrize('gap',['-0.000000001','0','0.000000001'])
def test_circle_against_sloping_finite_polygon_edge_strict_no_tolerance(gap):
    ring=[(92,-44),(-68,76),(-140,-20),(20,-140)];c=config(parts=[poly(box(-200,200,-200,200),holes=[ring]),circle(diameter=str(Decimal(40)+Decimal(gap)),name='圆岛',parent=parent())])
    if Decimal(gap)<0:validate(c);r=calculate(body(source(),parts=c['parts']),source());replay(r)
    else:
        with pytest.raises(DomainError):validate(c)


def test_concave_parent_hole_requires_every_segment_and_center_inside():
    c=config(parts=[poly(box(-150,150,-150,150),holes=[C_SHAPE]),circle(x=-70,diameter=40,name='圆岛',parent=parent())]);validate(c)
    for x in [-55,0]:
        c['parts'][1]['parameters']['x']=str(x)
        with pytest.raises(DomainError):validate(c)
    # A polygon around the notch with all vertices in a concave parent is still rejected if an edge crosses the void.
    c=config(parts=[poly(box(-150,150,-150,150),holes=[C_SHAPE]),poly([(-80,-80),(80,-80),(-80,80)],name='跨凹口',relation=parent()),circle(x=400,name='分离圆根')])
    with pytest.raises(DomainError):validate(c)


@pytest.mark.parametrize('direction',['circle_parent','polygon_parent'])
@pytest.mark.parametrize('unit,factor',[('mm','1'),('cm','.1'),('dm','.01'),('m','.001'),('km','.000001')])
def test_units_all_inputs_and_low_context_precision(direction,unit,factor):
    s=source();b=body(s,direction=direction);c=b['contact_ranges'][0]['plan_range'];c['unit']=unit
    for p in c['parts']:
        for h in [p]+p['holes']:h['parameters']={k:str(Decimal(v)*Decimal(factor)) for k,v in h['parameters'].items()}
    with localcontext() as ctx:ctx.prec=6;r=calculate(b,s)
    assert r['quantity']==calculate(body(s,direction=direction),s)['quantity'];replay(r)


@pytest.mark.parametrize('fault',['shape','hole_kind','root','wrong_hole','self','cycle','skip','bool','missing_parent','no_circle','no_polygon','no_island','name','basis','original','extra','one','five','overlap','contact','too_many_holes','too_many_points'])
def test_schema_family_relationship_and_limits(fault):
    c=config();p=c['parts'][1]
    if fault=='shape':p['shape']='ellipse'
    elif fault=='hole_kind':c['parts'][0]['holes'][0]=excluded(box(-10,10,-10,10))
    elif fault=='root':p['parent']=None
    elif fault=='wrong_hole':p['parent']['hole_index']=2
    elif fault=='self':p['parent']=parent(2)
    elif fault=='cycle':c['parts'][0]['parent']=parent(2);p['holes']=[excluded(box(-5,5,-5,5))]
    elif fault=='skip':c['parts']=chain();c['parts'][2]['parent']=parent(1)
    elif fault=='bool':p['parent']['block_index']=True
    elif fault=='missing_parent':del p['parent']
    elif fault=='no_circle':c['parts']=[poly(box(-100,100,-100,100),holes=[box(-60,60,-60,60)]),poly(box(-10,10,-10,10),name='岛',relation=parent())]
    elif fault=='no_polygon':c['parts']=cic()['parts'];c['parts']=[dict(p,shape='circular') for p in c['parts']]
    elif fault=='no_island':c['parts'][1]=poly(box(300,320,-10,10),name='分离直边')
    elif fault=='name':p['name']=c['parts'][0]['name']
    elif fault=='basis':p['measurement_basis']=''
    elif fault=='original':p['evidence']['x1']=''
    elif fault=='extra':p['count']=2
    elif fault=='one':c['parts']=c['parts'][:1]
    elif fault=='five':c['parts']*=3
    elif fault in ['overlap','contact']:c['parts'].append(circle(x=200 if fault=='contact' else 199,name='相邻根'))
    elif fault=='too_many_holes':c['parts'][0]['holes']*=5
    else:c['parts'][1]['parameters'].update(x13='1',y13='1');c['parts'][1]['evidence'].update(x13='test',y13='test')
    with pytest.raises(DomainError):validate(c)


def test_mixed_whole_window_pairs_against_all_old_families_and_own_islands():
    a=dict(plan_range=config());cases=[(cc(x=0,y=0,diameter=10),True),(cc(x=45,y=0,diameter=10),False),(window('test','-10','10','-5','5'),True),(window('test','35','45','-5','5'),False),(ac(x=0,y=0,outer=120,inner=110),True),(ec(x=10,y=0,outer=120,inner_x=0,inner_y=0,inner=70),False),(hc(),True),(nc(),True),(cic(),True)]
    for cfg,expected in cases:
        for left,right in [(a,dict(plan_range=cfg)),(dict(plan_range=cfg),a)]:
            q=separation(left,right);assert q['overlap']==expected
            if not expected:assert q['method']==METHOD and len(q['pairs'])==2
    b=dict(plan_range=config(parts=[circle(x=400,name='远根',holes=[inner('test',400,0,120,'W')]),poly(box(380,420,-20,20),name='远岛',relation=parent())]));q=separation(a,b);assert not q['overlap'] and len(q['pairs'])==4


@pytest.mark.parametrize('direction',['circle_parent','polygon_parent'])
def test_80_seeded_disk_intersections_independent_boolean_integration(direction):
    rng=random.Random(11701);a=dict(plan_range=config(direction=direction))
    for _ in range(80):
        x,y,d=rng.randint(-140,140),rng.randint(-140,140),rng.randint(2,60);base=[(x/1000,y/1000,d/2000)]
        if direction=='circle_parent':area=independent_area(base+[(0.,0.,.1)],excluded_disks=[(.01,0.,.06)])+independent_area(base,[(-.02,.02,-.02,.02)])
        else:area=independent_area(base,[(-.1,.1,-.1,.1)],excluded_boxes=[(-.06,.06,-.06,.06)])+independent_area(base+[(.01,0.,.02)])
        assert separation(a,dict(plan_range=cc(x=x,y=y,diameter=d)))['overlap']==(area>1e-11)


@pytest.mark.parametrize('direction',['circle_parent','polygon_parent'])
@pytest.mark.parametrize('gap',['0','0.000000001'])
def test_all_blocks_zero_or_tiny_positive_contact_not_rounded_to_proof(direction,gap):
    v=str(Decimal(25)+Decimal(gap));hs=[{**hole(x=0,diameter=80),'vertical_range':interval('100','300')},{**rectangle(name='H2',x='-'+v,length=50,width=100),'vertical_range':interval('300','500')},{**rectangle(name='H3',x=v,length=50,width=100),'vertical_range':interval('300','500')}]
    ps=[circle(diameter=40,holes=[inner('test',0,0,26,'W')]),poly(box(-5,5,-5,5),name='直边岛',relation=parent())] if direction=='circle_parent' else [poly(box(-20,20,-20,20),holes=[box(-13,13,-13,13)]),circle(x=2,diameter=12,name='圆形岛',parent=parent())]
    s=end_source(holes=hs);r=calculate(body(s,face='hole_top',parts=ps),s);d=r['contact_ranges'][0];assert ('polygon_zero_contact' in d)==(gap=='0') and (Decimal(d['single_net_area'])==0)==(gap=='0');replay(r)
    if gap=='0':assert {q['block_index'] for q in d['polygon_zero_contact']['net_trace']['partitions']}=={1,2}


def setup(clients,pid=None,zero=False,direction='circle_parent'):
    e,r,pid,s,b,eid=old_setup(clients,pid,zero)
    if direction=='circle_parent':ps=[circle(eid,diameter=40,holes=[inner(eid,0,0,26,'父圆W')]),poly(box(-5,5,-5,5),eid,'直边支模岛',relation=parent())] if zero else [circle(eid,x=196,y=8,diameter=42,holes=[inner(eid,201,10,16,'父圆W')]),poly(box(198,204,7,13),eid,'直边支模岛',relation=parent())]
    else:ps=[poly(box(-20,20,-20,20),eid,holes=[box(-13,13,-13,13)]),circle(eid,x=2,diameter=12,name='圆形支模岛',parent=parent())] if zero else [poly(box(175,215,-15,25),eid,holes=[box(185,205,-5,15)]),circle(eid,x=195,y=5,diameter=10,name='圆形支模岛',parent=parent())]
    b['contact_ranges'][0]['plan_range']=config(eid,parts=ps);return e,r,pid,s,b,eid


def test_roles_foreign_polygon_original_history_concurrency_source_stale(clients):
    e,r,pid,s,b,eid=setup(clients);before=get(e,pid);post(e,pid,b,'/components/preview');assert get(e,pid)==before;post(r,pid,b,expected=403)
    foreign=create(e)['id'];upload(e,foreign,'异工程合成');bad=copy.deepcopy(b);bad['contact_ranges'][0]['plan_range']['parts'][1]['evidence']['x1']=get(e,foreign)['evidence'][-1]['id'];post(e,pid,bad,expected=404)
    c=post(e,pid,b).json();approve(r,pid,'components',c);c=next(x for x in get(e,pid)['components'] if x['id']==c['id']);b['object_hash']=store.digest(c);b['contact_ranges'][0]['plan_range']['parts'][1]['parameters']['x2']='204.1';res=e.put(f'/api/v1/projects/{pid}/components/'+c['id'],json=b);assert res.status_code==200,res.text;assert e.put(f'/api/v1/projects/{pid}/components/'+c['id'],json=b).status_code==409;assert get(e,pid)['component_history'][-1]==c;approve(r,pid,'components',res.json())
    s['object_hash']=store.digest(s);s['foundation_items'][0]['levels'][1]['parameters']['diameter']='420';res=e.put(f'/api/v1/projects/{pid}/components/'+s['id'],json=s);assert res.status_code==200;approve(r,pid,'components',res.json());assert next(x for x in get(e,pid)['components'] if x['id']==c['id'])['state']=='stale'


def test_unused_parent_space_claim_allowed_island_overlap_or_frame_refused(clients):
    e,r,pid,s,b,eid=setup(clients);post(e,pid,b);other=copy.deepcopy(b);other['name']='父孔剩余局部另项';other['contact_ranges'][0]['plan_range']=window(eid,'195','197','9','11');post(e,pid,other)
    bad=copy.deepcopy(other);bad['contact_ranges'][0]['plan_range']['parameters']['x_end']='198.000000001';before=get(e,pid);post(e,pid,bad,expected=409);assert get(e,pid)==before
    bad=copy.deepcopy(other);bad['contact_ranges'][0]['plan_range']['coordinate_basis']='其他基准';post(e,pid,bad,expected=409)


@pytest.mark.parametrize('fault',['input_parent','family','relation_source','hierarchy','distance','distance_ast','edge_witness','geometry','zero','expression'])
def test_saved_mixed_topology_or_source_tamper_rejected(clients,fault):
    e,r,pid,s,b,eid=setup(clients,zero=True,direction='polygon_parent');c=post(e,pid,b).json()
    def forge(p):
        c=p['components'][-1];d=c['preview']['contact_ranges'][0];q=d['window_part_checks'][0]['proof']
        if fault=='input_parent':c['contact_ranges'][0]['plan_range']['parts'][1]['parent']=None
        elif fault=='family':c['contact_ranges'][0]['plan_range']['parts'][1]['shape']='polygon'
        elif fault=='relation_source':c['contact_ranges'][0]['plan_range']['parts'][1]['evidence_id']='fake'
        elif fault=='hierarchy':d['window_nesting'][1]['depth']=0
        elif fault=='distance':q['conditions'][0]['value']='999'
        elif fault=='distance_ast':q['conditions'][0]['formula_asts']['value']={'constant':'0'}
        elif fault=='edge_witness':q['conditions'][0]['edge_witness']['segment_parameter']='999'
        elif fault=='geometry':q['container']['points'][0][0]='999'
        elif fault=='zero':d['polygon_zero_contact']['net_trace']['partitions'][0]['remaining']=[{}]
        else:d['window_expression_definitions'][0]['value']='999'
    mutate(pid,forge);before=get(e,pid);approve(r,pid,'components',c,expected=422 if fault in ('input_parent','family') else 404 if fault=='relation_source' else 409);assert get(e,pid)==before


@pytest.mark.parametrize('gate',['preview','save','review','calculate','book_review'])
def test_mixed_island_relation_original_bytes_five_gates(clients,gate):
    e,r,pid,s,b,eid=setup(clients);upload(e,pid,'仅合成：直边岛在圆W内实际支模及父关系','mixed-relation.txt')
    for c in get(e,pid)['components']:approve(r,pid,'components',c)
    b['source_component_hash']=e.get(f'/api/v1/projects/{pid}/foundation-horizontal-formwork-candidates',params={'component_id':s['id']}).json()['source_component_hash'];b['contact_ranges'][0]['plan_range']['parts'][1]['evidence_id']=get(e,pid)['evidence'][-1]['id']
    if gate in ('review','calculate','book_review'):c=post(e,pid,b).json()
    if gate in ('calculate','book_review'):approve(r,pid,'components',c)
    if gate=='book_review':t=post(e,pid,{},'/takeoffs/calculate').json()[-1]
    doc=get(e,pid)['documents'][-1];(store.ROOT/'originals'/doc['hash']).write_bytes(b'corrupted-synthetic-mixed-relation-original')
    if gate=='preview':post(e,pid,b,'/components/preview',409)
    elif gate=='save':post(e,pid,b,expected=409)
    elif gate=='review':approve(r,pid,'components',c,expected=409)
    elif gate=='calculate':post(e,pid,{},'/takeoffs/calculate',409)
    else:approve(r,pid,'takeoffs',t,expected=409)


@pytest.mark.parametrize('direction',['circle_parent','polygon_parent'])
@pytest.mark.parametrize('zero',[False,True])
@pytest.mark.parametrize('fmt',['json','xlsx','pdf'])
def test_three_formats_both_mixed_directions_and_old_quote_isolation(clients,direction,zero,fmt):
    e,r,pid,_,q=priced_project(clients);e,r,pid,s,b,eid=setup(clients,pid,zero,direction);c=post(e,pid,b).json();approve(r,pid,'components',c)
    for t in post(e,pid,{},'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
    post(e,pid,{},'/boq/generate');url=f'/api/v1/projects/{pid}/exports';res=e.get(url,params=dict(format=fmt,unpriced=True));assert res.status_code==200,res.text
    titles={'基础混合岛区净块','基础混合岛区尺寸原件','基础混合岛区严格拓扑','基础混合岛区分段积分'}|({'基础混合岛区整体覆盖'} if zero else set())
    if fmt=='json':assert titles<={t[0] for t in res.json()['tables']}
    elif fmt=='xlsx':assert titles<=set(load_workbook(io.BytesIO(res.content)).sheetnames)
    else:assert all(any(title in p.extract_text() for p in PdfReader(io.BytesIO(res.content)).pages) for title in titles)
    assert not any(t[0] in titles for t in e.get(url,params=dict(format='json',quote_id=q['id'])).json()['tables'])


@pytest.mark.parametrize('fault',['block_index','state','edge_distance'])
def test_complete_mixed_block_pair_certificate_tamper_rejected(clients,fault):
    e,r,pid,s,b,eid=setup(clients);d=copy.deepcopy(b['contact_ranges'][0]);d.update(name='父孔未用空区',plan_range=window(eid,'195','197','9','11'));b['contact_ranges'].append(d);c=post(e,pid,b).json()
    def forge(p):
        q=p['components'][-1]['preview']['horizontal_window_checks'][0]['pairs'][0]
        if fault=='block_index':q['left_block_index']=2
        elif fault=='state':q['proof']['state']='boundary_touch'
        else:q['proof']['components'][0]['minimum_outer_squared']='999'
    mutate(pid,forge);before=get(e,pid);approve(r,pid,'components',c,expected=409);assert get(e,pid)==before


def test_48_straight_vertex_limit_and_nonterminating_projection_ast_are_exact():
    import math
    def ring(radius):return [(str(round(radius*math.cos(i*math.pi/6),9)),str(round(radius*math.sin(i*math.pi/6),9))) for i in range(12)]
    ps=[poly(ring(100),holes=[ring(70)]),poly(ring(50),name='直边岛',holes=[ring(30)],relation=parent()),circle(diameter=20,name='圆岛',parent=parent(2))];validate(config(parts=ps));s=source();r=calculate(body(s,parts=ps),s);replay(r)
    q=next(q for q in r['contact_ranges'][0]['window_part_checks'] if q['left_index']==2 and q['right_index']==3)['proof'];assert any(Fraction(v['edge_witness']['segment_parameter']).denominator not in (1,2) for v in q['conditions'])
    ps.append(poly([(300,0),(320,0),(310,20)],name='额外根块'))
    with pytest.raises(DomainError,match='48'):validate(config(parts=ps))


def test_eight_same_family_holes_across_mixed_blocks_have_usable_limit():
    positions=[(-40,-40),(40,-40),(40,40),(-40,40)];ps=[circle(holes=[inner('test',x,y,40,str(i)) for i,(x,y) in enumerate(positions)]),poly(box(300,500,-100,100),name='直边根',holes=[box(400+x-15,400+x+15,y-15,y+15) for x,y in positions]),poly(box(-45,-35,-45,-35),name='直边岛',relation=parent()),circle(x=360,y=-40,diameter=10,name='圆岛',parent=parent(2))];validate(config(parts=ps));s=source();replay(calculate(body(s,parts=ps),s))


def test_mixed_defense_exports_keep_circle_edges_and_straight_scan_block_ids():
    from backend.foundation_mixed_island_reports import tables,TITLES
    s=source();b=body(s);first=b['contact_ranges'][0]
    for name,cfg in [('父孔空区直边另项',window('test','-5','5','35','45')),('父孔空区圆形另项',cc(x=-30,y=20,diameter=6)),('分离直边窗',window('test','200','210','-5','5'))]:b['contact_ranges'].append({**copy.deepcopy(first),'name':name,'plan_range':cfg})
    r=calculate(b,s);ts=tables({'takeoffs':[{**b,**r}]});assert set(TITLES[5:])<={t[0] for t in ts};overview=next(v for title,_,v in ts if title==TITLES[5]);assert len(overview)==6 and all('C1 B' in q[0] for q in overview)
