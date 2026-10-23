"""Finite-height circular voids: exact clipping, 3D separation, evidence and workflows."""
import copy,io,random
from decimal import Decimal,localcontext,ROUND_HALF_UP
import pytest
from openpyxl import load_workbook
from pypdf import PdfReader
from backend import store
from backend.compound_takeoff import PI
from backend.engine import DomainError
from backend.post_foundation import calculate
from tests.test_foundation_openings import body as old_body,setup as old_setup,hole,SHAPES
from tests.test_post_foundation import source,equal_body,fixed_body,fixed_added,equal_added,fixed_adjusted,equal_adjusted,combined
from tests.test_route_takeoff import evaluate
from tests.test_polygon_takeoff import post,mutate
from tests.test_workflow import get,approve,upload,create,priced_project

TITLES={'基础圆孔毛扣净','基础圆孔原始依据','基础圆孔校验','基础圆孔范围与常量','基础圆孔标高依据','基础圆孔逐层相交'}
QUANTITIES=dict(zip(SHAPES,('2.628','2.364','3.720','2.337')))


def interval(start='0',end='150',eid='test'):
 return dict(parameters=dict(z_start=start,z_end=end),evidence=dict(z_start=eid,z_end=eid))


def body(s,eid='test',shape='stepped_circular'):
 b=old_body(s,eid,shape);b['name']='合成基础有限标高圆孔（非真实工程）';b['measurement_basis']='合成22基础，每基础两竖直圆孔，基础底面Z=0向上，H1起止0至150mm、H2起止300至600mm，实际直径40mm；逐层只扣正相交高度，相离不重复，净总量修约3位，非真实工程。';d=b['foundation_items'][0];d['measurement_basis']=b['measurement_basis']
 for h,v in zip(d['through_holes']['holes'],[interval(eid=eid),interval('300','600',eid)]):h.update(vertical_range=v);h['parameters'].update(x='0',y='0');h['measurement_basis']=b['measurement_basis']
 d['through_holes']['coordinate_basis']='合成各层共同中心为XY原点、X沿长度正向Y沿宽度正向，Z=0为基础实体底面且向上，全部孔位及标高有原件；有限孔不推定贯通高度。';return b


def setup(clients,pid=None,shape='stepped_circular'):
 e,r,pid,s,_,old_eid=old_setup(clients,pid,shape)
 upload(e,pid,'合成有限标高圆孔软件验收，非真实工程。已审最终立柱11根、重复2共22基础，原材料79.2kg不变。基础下圆层实际直径600mm净高200mm，上圆层实际直径400mm净高400mm，总高600mm；同心紧接同材料。局部XY以共同中心为原点、X沿长度正向Y沿宽度正向；Z=0为基础实体底面向上。每基础H1实际孔径40mm，孔心(250,0)mm，Z从0至150mm，仅在下层，距上层轮廓外但不进入上层；H2实际孔径40mm，孔心(0,0)mm，Z从300至600mm，仅在上层；共44孔，合计扣减.00396πm3，圆台阶净2.337m3。H1终止Z改190mm，合计扣减.004312πm3，净2.336m3。H1终止Z改200mm仅端面接触上层仍可；200.000000001mm进入上层且越界须拒。同XY两孔，Z0至150与300至600标高相离可扣；H1终止Z300与H2起始300仅端面接触可扣；300.000000001产生正体积重叠须拒。其它形状盲孔验收孔心均0，矩形净2.628m3、矩形台阶净2.364m3、圆柱净3.720m3。所有标高、孔径、坐标及计量范围须逐参数据原件，深度不从柱脚推断，非结构净距或专业真实验收。TEST-RECESS仅合成映射编码，非官方条目。','foundation-recess-basis.txt')
 for c in get(e,pid)['components']:approve(r,pid,'components',c)
 s=get(e,pid)['components'][-1];eid=get(e,pid)['evidence'][-1]['id'];return e,r,pid,s,body(s,eid,shape),eid


@pytest.mark.parametrize('shape',SHAPES)
def test_clipped_holes_exact_balances_all_asts_and_input_immutability(shape):
 s=source();b=body(s,shape=shape);before=copy.deepcopy((b,s));r=calculate(b,s);assert (b,s)==before and r['quantity']==QUANTITIES[shape] and r['rule_version']=='railing-post-foundation-1.4' and r['foundation_summary']['total_count']=='22' and r['foundation_opening_summary']['total_hole_count']=='44'
 with localcontext() as ctx:
  ctx.prec=80;assert Decimal(r['deduction'])==Decimal(PI)*Decimal('.00396');assert Decimal(r['gross'])-Decimal(r['deduction'])==Decimal(r['foundation_summary']['unrounded_quantity']);assert evaluate(r['formula_ast'],r['dimensions'])==Decimal(r['quantity'])
  for d in r['foundation_items']:
   for k,a in d['formula_asts'].items():assert evaluate(a,r['dimensions'])==Decimal(d[{'single_volume':'single_foundation_volume'}.get(k,k)])
   for layer in d.get('foundation_levels',[]):
    for k,a in layer['formula_asts'].items():assert evaluate(a,r['dimensions'])==Decimal(layer[k])
   for h in d['opening_holes']:
    assert evaluate(h['height_ast'],r['dimensions'])==Decimal(h['through_height']) and evaluate(h['formula_ast'],r['dimensions'])==Decimal(h['single_volume'])
    for k,a in h['vertical_formula_asts'].items():assert evaluate(a,r['dimensions'])==Decimal(h[k])
    for seg in h['segments']:
     for k,a in seg['formula_asts'].items():assert evaluate(a,r['dimensions'])==Decimal(seg[k])
    for c in h['containments']:
     for k,a in c['formula_asts'].items():assert evaluate(a,r['dimensions'])==Decimal(c['margins'][k])>=0
   for c in d['opening_checks']:assert evaluate(c['formula_ast'],r['dimensions'])==Decimal(c['margin']) and evaluate(c['vertical_gap_ast'],r['dimensions'])==Decimal(c['vertical_gap'])


@pytest.mark.parametrize('shape',SHAPES)
@pytest.mark.parametrize('make',[equal_body,fixed_body,fixed_added,equal_added,fixed_adjusted,equal_adjusted,combined,lambda s:combined(s,equal=True)])
def test_final_post_count_modes(shape,make):
 s=source(make);r=calculate(body(s,shape=shape),s)
 with localcontext() as ctx:ctx.prec=80;assert Decimal(r['deduction'])==Decimal(PI)*Decimal('.00018')*r['foundation_summary']['single_group_count']*2


@pytest.mark.parametrize('shape',SHAPES)
@pytest.mark.parametrize('unit,factor',[('m',Decimal('.001')),('mm',Decimal(1)),('cm',Decimal('.1')),('dm',Decimal('.01')),('km',Decimal('.000001'))])
def test_independent_units_and_low_decimal_context(shape,unit,factor):
 s=source();b=body(s,shape=shape);c=b['foundation_items'][0]['through_holes'];c['unit']=unit
 for h in c['holes']:
  for p in (h['parameters'],h['vertical_range']['parameters']):
   for k in p:p[k]=str(Decimal(p[k])*factor)
 with localcontext() as ctx:ctx.prec=6;r=calculate(b,s)
 assert r==calculate(b,s) and r['quantity']==QUANTITIES[shape]


@pytest.mark.parametrize('shape',['stepped_circular','stepped_rectangular'])
@pytest.mark.parametrize('end,allowed',[('150',True),('200',True),('199.999999999',True),('200.000000001',False)])
def test_check_only_actual_intersecting_layers_and_exact_boundary(shape,end,allowed):
 s=source();b=body(s,shape=shape);h=b['foundation_items'][0]['through_holes']['holes'][0];h['parameters']['x']='250';h['vertical_range']['parameters']['z_end']=end
 if allowed:
  d=calculate(b,s)['foundation_items'][0];assert [v['level_index'] for v in d['opening_holes'][0]['segments']]==[1];assert [v['level_index'] for v in d['opening_holes'][0]['containments']]==[1]
 else:
  with pytest.raises(DomainError,match='第2层'):calculate(b,s)


@pytest.mark.parametrize('end,state',[('299.999999999','vertically_disjoint'),('300','vertical_tangent'),('300.000000001',None)])
def test_coincident_xy_depth_separation_and_positive_intersection(end,state):
 s=source();b=body(s);b['foundation_items'][0]['through_holes']['holes'][0]['vertical_range']['parameters']['z_end']=end
 if state:assert calculate(b,s)['foundation_items'][0]['opening_checks'][0]['state']==state
 else:
  with pytest.raises(DomainError,match='正相交高度内重叠'):calculate(b,s)


@pytest.mark.parametrize('x,state',[('40','tangent'),('40.000000001','disjoint'),('39.999999999',None)])
def test_positive_depth_overlap_xy_separation(x,state):
 s=source();b=body(s);holes=b['foundation_items'][0]['through_holes']['holes'];holes[0]['vertical_range']['parameters']['z_end']='400';holes[1]['parameters']['x']=x
 if state:assert calculate(b,s)['foundation_items'][0]['opening_checks'][0]['state']==state
 else:
  with pytest.raises(DomainError,match='重叠'):calculate(b,s)


def test_cross_layer_clipping_unaffected_layers_and_mixed_through_hole():
 s=source();b=body(s);d=b['foundation_items'][0];h=d['through_holes']['holes'][0];h['vertical_range']['parameters'].update(z_start='100',z_end='300');d['through_holes']['holes'][1].pop('vertical_range');d['through_holes']['holes'][1]['parameters']['x']='50';r=calculate(b,s);rows=r['foundation_items'][0]['opening_holes'];assert [(v['level_index'],Decimal(v['height'])) for v in rows[0]['segments']]==[(1,Decimal('.1')),(2,Decimal('.1'))] and rows[1]['through_height']=='0.6'
 with localcontext() as ctx:ctx.prec=80;assert Decimal(r['deduction'])==Decimal(PI)*Decimal('.00704')
 d['through_holes']['holes']=[hole(eid='test')];d['through_holes']['holes'][0]['vertical_range']=interval('300','400');r=calculate(b,s);assert r['foundation_items'][0]['foundation_levels'][0]['single_opening_volume']=='0'


def test_mixed_legacy_and_finite_rows_and_ten_layers_twenty_depths():
 s=source();s['material_items'].append({**copy.deepcopy(s['material_items'][0]),'name':'合成贯通项'});b=body(s);old=old_body(s)['foundation_items'][0];old['material_index']=2;b['foundation_items'].append(old);r=calculate(b,s);assert r['foundation_summary']['total_count']=='44' and r['foundation_opening_summary']['total_hole_count']=='88' and 'segments' not in r['foundation_items'][1]['opening_holes'][0]
 b['foundation_items'].reverse();assert calculate(b,s)==r
 b=body(s);d=b['foundation_items'][0];d['levels']=[dict(parameters=dict(diameter='400',height='100'),evidence=dict(diameter='test',height='test')) for _ in range(10)];d['through_holes']['holes']=[{**hole('z'+str(i),x='0',diameter='10'),'vertical_range':interval(str(i*50),str((i+1)*50))} for i in range(20)];r=calculate(b,s);assert r['foundation_opening_summary']['total_hole_count']=='440' and len(r['foundation_items'][0]['opening_checks'])==190
 with localcontext() as ctx:ctx.prec=80;assert Decimal(r['deduction'])==Decimal(PI)*Decimal('.00055')


def test_seeded_independent_interval_intersections_and_rounding():
 rng=random.Random(9001)
 for shape in SHAPES:
  for _ in range(30):
   s=source();b=body(s,shape=shape);d=b['foundation_items'][0];a=rng.randint(0,598);end=rng.randint(a+1,600);diam=rng.randint(1,80);h=d['through_holes']['holes'][0];h['vertical_range']=interval(str(a),str(end));h['parameters']['diameter']=str(diam);d['through_holes']['holes']=[h];digits=rng.randint(0,12);b['rounding_digits']=digits;plain=copy.deepcopy(b);plain['foundation_items'][0].pop('through_holes');r=calculate(b,s);gross=calculate(plain,s)['gross']
   with localcontext() as ctx:
    ctx.prec=80;ded=Decimal(PI)/4*diam*diam*Decimal(end-a)/1000000000*22;assert Decimal(r['deduction'])==ded and Decimal(r['quantity'])==(Decimal(gross)-ded).quantize(Decimal(1).scaleb(-digits),rounding=ROUND_HALF_UP)


@pytest.mark.parametrize('fault',['null','list','extra','missing','params-null','params-extra','params-missing','proof-null','proof-extra','proof-missing','proof-empty','proof-bool','negative','reverse','equal','outside','zero-end','nan','infinite','bool','precision','unit','unconfirmed'])
def test_strict_finite_schema_and_depth_limits(fault):
 s=source();b=body(s);c=b['foundation_items'][0]['through_holes'];h=c['holes'][0];v=h['vertical_range'];p=v['parameters']
 if fault=='null':h['vertical_range']=None
 elif fault=='list':h['vertical_range']=[]
 elif fault=='extra':v['depth']='150'
 elif fault=='missing':v.pop('evidence')
 elif fault=='params-null':v['parameters']=None
 elif fault=='params-extra':p['depth']='150'
 elif fault=='params-missing':p.pop('z_end')
 elif fault=='proof-null':v['evidence']=None
 elif fault=='proof-extra':v['evidence']['depth']='test'
 elif fault=='proof-missing':v['evidence'].pop('z_start')
 elif fault=='proof-empty':v['evidence']['z_start']=''
 elif fault=='proof-bool':v['evidence']['z_end']=True
 elif fault=='negative':p['z_start']='-.000000001'
 elif fault=='reverse':p.update(z_start='151',z_end='150')
 elif fault=='equal':p.update(z_start='150',z_end='150')
 elif fault=='outside':p['z_end']='600.000000001'
 elif fault=='zero-end':p['z_end']='0'
 elif fault in ('nan','infinite','bool','precision'):p['z_start']={'nan':'NaN','infinite':'Infinity','bool':True,'precision':'.1234567890123456789'}[fault]
 elif fault=='unit':c['unit']={}
 else:c['scope_confirmed']=False
 with pytest.raises(DomainError):calculate(b,s)


@pytest.mark.parametrize('shape',SHAPES)
@pytest.mark.parametrize('fmt',['json','xlsx','pdf'])
def test_workflow_net_boq_three_exports_mapping_and_old_quote(clients,shape,fmt):
 e,r,pid,_,quote=priced_project(clients);e,r,pid,s,b,eid=setup(clients,pid,shape);before=get(e,pid);assert post(e,pid,b,'/components/preview').json()['quantity']==QUANTITIES[shape] and get(e,pid)==before;c=post(e,pid,b).json();approve(r,pid,'components',c)
 for t in post(e,pid,{},'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
 boq=post(e,pid,{},'/boq/generate').json();assert any(v['quantity']==QUANTITIES[shape] and v['unit']=='m³' for v in boq) and any(v['quantity']=='79.200' for v in boq);url=f'/api/v1/projects/{pid}/exports';res=e.get(url,params=dict(format=fmt,unpriced=True));assert res.status_code==200,res.text
 if fmt=='json':assert TITLES<={v[0] for v in res.json()['tables']}
 elif fmt=='xlsx':
  w=load_workbook(io.BytesIO(res.content));assert TITLES<=set(w.sheetnames) and w['基础圆孔逐层相交'].max_row>2 and w['基础圆孔标高依据'].row_dimensions[3].height>48
 else:assert all(v in ''.join(p.extract_text() for p in PdfReader(io.BytesIO(res.content)).pages) for v in TITLES)
 assert not any(v[0] in TITLES for v in e.get(url,params=dict(format='json',quote_id=quote['id'])).json()['tables'])
 if fmt=='json':
  from tests.test_mapping_profiles import rule_body,save_rule,preview,selection
  p=get(e,pid);rule={**rule_body(p,'post_foundation','volume',eid),'code':'TEST-RECESS'};assert e.patch(f'/api/v1/projects/{pid}/settings',json=dict(boq_standard=rule['standard'],boq_edition=rule['edition'],reason='仅合成有限标高圆孔验收')).status_code==200;save_rule(e,r,pid,rule);bid=next(v['id'] for v in p['boq'] if v['takeoff_id']==p['takeoffs'][-1]['id']);chosen=preview(e,pid,bid);assert chosen['state']=='unique';assert e.post(f'/api/v1/projects/{pid}/boq/apply-mappings',json=dict(selections=[selection(chosen)])).status_code==200;approve(r,pid,'boq',next(v for v in get(e,pid)['boq'] if v['id']==bid))


@pytest.mark.parametrize('fault',['raw','proof','segment-height','segment-volume','segment-ast','z','pair','rule','notice','book'])
def test_tamper_reconstruction(clients,fault):
 e,r,pid,s,b,eid=setup(clients);c=post(e,pid,b).json()
 if fault=='book':approve(r,pid,'components',c);post(e,pid,{},'/takeoffs/calculate')
 def forge(p):
  c=p['components'][-1];r=c['preview'];d=r['foundation_items'][0];h=d['opening_holes'][0]
  if fault=='raw':c['foundation_items'][0]['through_holes']['holes'][0]['vertical_range']['parameters']['z_end']='190'
  elif fault=='proof':c['foundation_items'][0]['through_holes']['holes'][0]['vertical_range']['evidence']['z_end']='forged'
  elif fault=='segment-height':h['segments'][0]['height']='0'
  elif fault=='segment-volume':h['segments'][0]['single_volume']='0'
  elif fault=='segment-ast':h['segments'][0]['formula_asts']['height']={'constant':'0'}
  elif fault=='z':h['z_start']='.01'
  elif fault=='pair':d['opening_checks'][0]['vertical_gap']='0'
  elif fault=='rule':r['rule_version']='railing-post-foundation-1.3'
  elif fault=='notice':r['notice']='changed'
  else:p['takeoffs'][-1]['deduction']='0'
 mutate(pid,forge);collection='takeoffs' if fault=='book' else 'components';approve(r,pid,collection,get(e,pid)[collection][-1],expected=404 if fault=='proof' else 409)


@pytest.mark.parametrize('gate',['preview','save','review','calculate','book_review'])
def test_z_original_bytes_gates(clients,gate):
 e,r,pid,s,b,eid=setup(clients)
 if gate in ('review','calculate','book_review'):c=post(e,pid,b).json()
 if gate in ('calculate','book_review'):approve(r,pid,'components',c)
 if gate=='book_review':t=post(e,pid,{},'/takeoffs/calculate').json()[-1]
 doc=get(e,pid)['documents'][-1];(store.ROOT/'originals'/doc['hash']).write_bytes(b'corrupt')
 if gate=='preview':post(e,pid,b,'/components/preview',409)
 elif gate=='save':post(e,pid,b,expected=409)
 elif gate=='review':approve(r,pid,'components',c,expected=409)
 elif gate=='calculate':post(e,pid,{},'/takeoffs/calculate',409)
 else:approve(r,pid,'takeoffs',t,expected=409)


def test_depth_revision_exact_history_old_boq_conflict_and_switch_back(clients):
 e,r,pid,s,b,eid=setup(clients);c=post(e,pid,b).json();approve(r,pid,'components',c)
 for t in post(e,pid,{},'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
 post(e,pid,{},'/boq/generate');p=get(e,pid);old=p['components'][-1];new=copy.deepcopy(b);new['foundation_items'][0]['through_holes']['holes'][0]['vertical_range']['parameters']['z_end']='190';new['object_hash']=store.digest(old);url=f'/api/v1/projects/{pid}/components/'+old['id'];res=e.put(url,json=new);assert res.status_code==200 and res.json()['preview']['quantity']=='2.336';assert get(e,pid)['component_history'][-1]==old and get(e,pid)['boq']==p['boq'] and all(t['state']=='stale' for t in get(e,pid)['takeoffs']);assert e.put(url,json=new).status_code==409;post(e,pid,{**b,'id':old['id']},expected=409);approve(r,pid,'components',res.json())
 new=old_body(s,eid);new['object_hash']=store.digest(get(e,pid)['components'][-1]);res=e.put(url,json=new);assert res.status_code==200 and res.json()['preview']['quantity']=='2.317' and res.json()['preview']['rule_version']=='railing-post-foundation-1.3';approve(r,pid,'components',res.json())
 s['object_hash']=store.digest(s);s['material_items'][0]['post_ranges'][0]['spacing']='3000';res=e.put(f'/api/v1/projects/{pid}/components/'+s['id'],json=s);assert res.status_code==200;approve(r,pid,'components',res.json());assert get(e,pid)['components'][-1]['state']=='stale'


@pytest.mark.parametrize('field',['z_start','z_end'])
def test_roles_tenants_and_foreign_depth_evidence(clients,field):
 e,r,pid,s,b,eid=setup(clients);post(r,pid,b,expected=403);assert clients['other'].get(f'/api/v1/projects/{pid}').status_code==404;other=create(e)['id'];upload(e,other,'异项目标高原件');fid=get(e,other)['evidence'][-1]['id'];b['foundation_items'][0]['through_holes']['holes'][0]['vertical_range']['evidence'][field]=fid;post(e,pid,b,expected=404)


def test_eighteen_decimal_normalized_z_and_full_height_explicit_range():
 s=source();b=body(s);holes=b['foundation_items'][0]['through_holes']['holes'];holes[0]['vertical_range']=interval('0.000000000000001','150.000000000000001');r=calculate(b,s)
 with localcontext() as ctx:ctx.prec=80;assert Decimal(r['deduction'])==Decimal(PI)*Decimal('.00396')
 holes[0]['vertical_range']=interval('0','600');holes[1]['parameters']['x']='50';r=calculate(b,s);assert r['foundation_items'][0]['opening_holes'][0]['through_height']=='0.6'


def test_long_named_twenty_depths_pdf(clients):
 e,r,pid,s,b,eid=setup(clients);b['name']='合成'+('有限标高长名称'*20);d=b['foundation_items'][0];d['through_holes']['holes']=[{**hole('深度'+str(i)+('原件说明'*12),x='0',diameter='10',eid=eid),'vertical_range':interval(str(i*30),str((i+1)*30),eid)} for i in range(20)];c=post(e,pid,b).json();approve(r,pid,'components',c)
 for t in post(e,pid,{},'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
 res=e.get(f'/api/v1/projects/{pid}/exports',params=dict(format='pdf',unpriced=True));assert res.status_code==200,res.text
