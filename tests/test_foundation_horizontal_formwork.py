"""Synthetic whole end-face contact; independent physical opening and count checks."""
import copy, io, random
from decimal import Decimal, localcontext, ROUND_HALF_UP
import pytest
from openpyxl import load_workbook
from pypdf import PdfReader
from backend import store
from backend.compound_takeoff import PI
from backend.engine import DomainError
from backend.foundation_horizontal_formwork import calculate
from backend.foundation_horizontal_formwork_reports import TITLES, tables
from backend.post_foundation import calculate as foundation_quantity
from tests.test_foundation_formwork import foundation, setup as outer_setup
from tests.test_foundation_openings import hole
from tests.test_foundation_voids import rectangle
from tests.test_foundation_recesses import interval
from tests.test_post_foundation import equal_body, fixed_body, fixed_added, equal_added, fixed_adjusted, equal_adjusted, combined
from tests.test_route_takeoff import evaluate
from tests.test_polygon_takeoff import post, mutate
from tests.test_workflow import get, approve, upload, create, priced_project
from backend.railing_material import snapshot


def contact(eid='test', face='top', name='合成顶面'):
    return dict(name=name, material_index=1, face=face, measurement_basis='合成整端面支模，按实际贯通本端面的孔截面扣减；非真实工程施工依据。', evidence=dict(measurement_basis=eid))


def body(s, eid='test'):
    return dict(name='合成基础整体顶底模板（非真实工程）',kind='foundation_horizontal_formwork',parameters={},quantity_type='area',
                source_component_id=s['id'],source_component_hash=store.digest(snapshot(s)),contact_ranges=[contact(eid),contact(eid,'bottom','合成底面')],scope_confirmed=True,
                material='合成顶底模板材料及施工做法，仅软件验证',material_evidence_id=eid,measurement_basis='合成22基础整体最顶和最底整面，按来源真实尺寸与实际孔截面计量，不因体积阈值填实孔；不含台阶中间平台、孔底、竖直面和措施，不重复既有清单。',basis_evidence_id=eid,rounding_digits=3,rounding_evidence_id=eid)


def setup(clients,pid=None,shape='rectangular'):
    e,r,pid,s,_,_=outer_setup(clients,pid,shape)
    upload(e,pid,'合成基础整体顶底模板软件验收，非真实工程或工法。22个400×500×600mm矩形基础，H1矩形100×60mm、孔心(150,0)、Z0至150mm，仅底面贯通。体积阈值100m²不扣孔体积，但实际孔仍存在。整体顶毛净.2m²/个，底毛.2扣.006净.194m²/个，共8.668m²；只选顶面4.400m²，只选底面4.268m²。源改高700mm仍保持同顶底面积，模板须失效重审。台阶中间平台不计；材料/实际支模须独立核定。TEST-HORIZONTAL仅软件映射码。','foundation-horizontal-basis.txt')
    for c in get(e,pid)['components']: approve(r,pid,'components',c)
    s=get(e,pid)['components'][-1]
    return e,r,pid,s,body(s,get(e,pid)['evidence'][-1]['id']),get(e,pid)['evidence'][-1]['id']


def check(r):
    with localcontext() as ctx:
        ctx.prec=80
        assert evaluate(r['formula_ast'],r['dimensions'])==Decimal(r['quantity'])
        for d in r['contact_ranges']:
            for k,a in d['formula_asts'].items(): assert evaluate(a,r['dimensions'])==Decimal(d[k]),k
            for h in d['apertures']: assert evaluate(h['formula_ast'],r['dimensions'])==Decimal(h['single_area'])


def test_rectangular_end_faces_physical_hole_despite_volume_threshold_immutable():
    s=foundation();b=body(s);before=copy.deepcopy((s,b));r=calculate(b,s)
    assert (s,b)==before and r['quantity']=='8.668' and r['gross']=='8.80' and r['deduction']=='0.132'
    assert foundation_quantity(s)['quantity']=='2.640' and foundation_quantity(s)['deduction']=='0'
    assert r['contact_ranges'][0]['apertures']==[] and r['contact_ranges'][1]['apertures'][0]['hole_index']==1
    assert r['contact_ranges'][0]['z']=='0.6' and r['contact_ranges'][1]['z']=='0';check(r)


@pytest.mark.parametrize('shape',['rectangular','stepped_rectangular','circular','stepped_circular'])
@pytest.mark.parametrize('make',[equal_body,fixed_body,fixed_added,equal_added,fixed_adjusted,equal_adjusted,combined,lambda s:combined(s,equal=True)])
def test_four_shapes_eight_final_counts_end_layers_once(shape,make):
    s=foundation(make,shape);r=calculate(body(s),s);check(r)
    with localcontext() as ctx:
        ctx.prec=80;p=Decimal(PI);single={'rectangular':Decimal('.394'),'stepped_rectangular':Decimal('.414')-Decimal('.0004')*p,'circular':Decimal('.1796')*p-Decimal('.006'),'stepped_circular':Decimal('.1296')*p-Decimal('.006')}[shape]
        count=Decimal(r['contact_ranges'][0]['total_count']);assert Decimal(r['horizontal_formwork_summary']['unrounded_quantity'])==single*count
    if 'stepped' in shape: assert [d['level_index'] for d in r['contact_ranges']]==[2,1]


@pytest.mark.parametrize('lo,hi',[('0','600'),('0','150'),('300','600'),('100','500'),('0.000000001','599.999999999')])
def test_finite_holes_only_open_to_actual_end_no_snapping(lo,hi):
    s=foundation();c=s['foundation_items'][0]['through_holes'];c['holes']=[{**rectangle(x='0'),'vertical_range':interval(lo,hi)}];r=calculate(body(s),s);check(r)
    assert [len(d['apertures']) for d in r['contact_ranges']]==[int(hi=='600'),int(lo=='0')]
    assert Decimal(r['deduction'])==Decimal('.132')*(int(lo=='0')+int(hi=='600'))


def test_through_holes_legacy_and_mixed_profiles_multiple_disjoint():
    from tests.test_foundation_openings import body as old_body
    from tests.test_post_foundation import source
    s={**old_body(source(),shape='stepped_circular'),'id':'legacy','revision':1,'source_geometry':source()};r=calculate(body(s),s);check(r);assert all(len(d['apertures'])==2 for d in r['contact_ranges'])
    s=foundation();c=s['foundation_items'][0]['through_holes'];c['holes']=[rectangle(x='-80'),hole(name='合成H2',x='80',diameter='40')];r=calculate(body(s),s);assert all(len(d['apertures'])==2 for d in r['contact_ranges']);check(r)


@pytest.mark.parametrize('shape',['rectangular','stepped_rectangular','circular','stepped_circular'])
def test_no_holes_and_low_calling_precision(shape):
    s=foundation(shape=shape);s['foundation_items'][0].pop('through_holes');b=body(s)
    with localcontext() as ctx:ctx.prec=6;r=calculate(b,s)
    assert r==calculate(b,s) and r['deduction']=='0' and all(not d['apertures'] for d in r['contact_ranges']);check(r)


def test_seeded_independent_rectangular_holes_and_rounding_oracle():
    rng=random.Random(9901)
    for _ in range(100):
        s=foundation();c=s['foundation_items'][0]['through_holes'];length=rng.randint(1,100);width=rng.randint(1,100);lo=rng.choice([0,rng.randint(1,299)]);hi=rng.choice([600,rng.randint(300,599)]);c['holes']=[{**rectangle(x='0',length=str(length),width=str(width)),'vertical_range':interval(str(lo),str(hi))}];b=body(s);b['rounding_digits']=rng.randint(0,12);r=calculate(b,s)
        with localcontext() as ctx:ctx.prec=80;expected=(Decimal('.4')-Decimal(length*width)/1000000*(int(lo==0)+int(hi==600)))*22;assert Decimal(r['quantity'])==expected.quantize(Decimal(1).scaleb(-b['rounding_digits']),rounding=ROUND_HALF_UP)
        check(r)


@pytest.mark.parametrize('shape',['rectangular','stepped_rectangular','circular','stepped_circular'])
@pytest.mark.parametrize('unit,factor',[('mm','1'),('cm','.1'),('dm','.01'),('m','.001'),('km','.000001')])
def test_source_length_units_independent_from_holes_and_end_z(shape,unit,factor):
    s=foundation(shape=shape);expected=calculate(body(s),s);d=s['foundation_items'][0];d['length_unit']=unit
    for p in ([v['parameters'] for v in d['levels']] if 'levels' in d else [d['parameters']]):
        for k in p:p[k]=str(Decimal(p[k])*Decimal(factor))
    r=calculate(body(s),s);assert r['quantity']==expected['quantity'] and r['gross']==expected['gross'] and r['deduction']==expected['deduction'];check(r)


def test_full_footprint_recess_removes_bottom_contact_without_negative_area():
    s=foundation();c=s['foundation_items'][0]['through_holes'];c['holes']=[{**rectangle(x='0',length='400',width='500'),'vertical_range':interval('0','150')}];b=body(s);b['contact_ranges']=[contact(face='bottom')];r=calculate(b,s);assert r['quantity']=='0.000' and Decimal(r['gross'])==Decimal(r['deduction']) and r['contact_ranges'][0]['single_net_area']=='0.00';check(r)


def test_two_foundation_items_same_end_have_independent_claims_and_counts():
    s=foundation();original=copy.deepcopy(s['source_geometry']['material_items'][0]);original['name']='合成第二来源立柱项';s['source_geometry']['material_items'].append(original);s['source_component_hash']=store.digest(snapshot(s['source_geometry']));s['foundation_items'].append({**copy.deepcopy(s['foundation_items'][0]),'material_index':2});b=body(s);one=calculate(b,s);b['contact_ranges'] += [{**copy.deepcopy(d),'material_index':2,'name':d['name']+'第二项'} for d in b['contact_ranges']];r=calculate(b,s);assert Decimal(r['quantity'])==Decimal(one['quantity'])*2;check(r)


@pytest.mark.parametrize('fault',['empty','null','extra','missing','bool-index','unknown-index','face','face-type','step','hole-bottom','parameters','area','unit','count','z','evidence','blank-proof','blank-basis','name','duplicate-name','duplicate-face','confirm','truthy','rounding','material','too-many','type'])
def test_strict_inputs(fault):
    s=foundation();b=body(s);d=b['contact_ranges'][0]
    if fault=='empty':b['contact_ranges']=[]
    elif fault=='null':b['contact_ranges']=[None]
    elif fault=='extra':d['unknown']='x'
    elif fault=='missing':d.pop('face')
    elif fault=='bool-index':d['material_index']=True
    elif fault=='unknown-index':d['material_index']=99
    elif fault in ('face','step','hole-bottom'):d['face']={'face':'circumference','step':'step','hole-bottom':'hole_bottom'}[fault]
    elif fault=='face-type':d['face']=[]
    elif fault in ('parameters','area','unit','z'):d[fault]={} if fault=='parameters' else '1'
    elif fault=='count':b['multiplicity']='2'
    elif fault=='evidence':d['evidence']['z']='test'
    elif fault=='blank-proof':d['evidence']['measurement_basis']=' '
    elif fault=='blank-basis':d['measurement_basis']=' '
    elif fault=='name':d['name']=' '
    elif fault=='duplicate-name':b['contact_ranges'][1]['name']=d['name']
    elif fault=='duplicate-face':b['contact_ranges'][1]['face']='top'
    elif fault=='confirm':b['scope_confirmed']=False
    elif fault=='truthy':b['scope_confirmed']=1
    elif fault=='rounding':b['rounding_digits']=True
    elif fault=='material':b['material']=' '
    elif fault=='too-many':b['contact_ranges']=[copy.deepcopy(d)]*201
    elif fault=='type':b['quantity_type']='volume'
    with pytest.raises(DomainError):calculate(b,s)


def test_candidates_preview_no_mutation_and_actual_ids(clients):
    e,r,pid,s,b,eid=setup(clients);before=get(e,pid);res=e.get(f'/api/v1/projects/{pid}/foundation-horizontal-formwork-candidates',params=dict(component_id=s['id']));assert res.status_code==200,res.text
    f=res.json()['foundation_items'][0];assert [x['face'] for x in f['faces'] if x['face'] in ('top','bottom','step')]==['top','bottom'] and f['faces'][1]['single_opening_area']=='0.006'
    assert post(e,pid,b,'/components/preview').json()['quantity']=='8.668' and get(e,pid)==before
    b['contact_ranges']=[contact(eid)];c=post(e,pid,b).json();approve(r,pid,'components',c);post(e,pid,{**b,'id':c['id']},expected=409)
    b['contact_ranges']=[contact(eid,'bottom')];post(e,pid,b);b['contact_ranges']=[contact(eid)];post(e,pid,b,expected=409)


@pytest.mark.parametrize('fault',['face','proof','gross','hole','cut','z','formula','source','book'])
def test_saved_reconstruction_tamper_gates(clients,fault):
    e,r,pid,s,b,eid=setup(clients);c=post(e,pid,b).json()
    if fault=='book':approve(r,pid,'components',c);post(e,pid,{},'/takeoffs/calculate')
    def forge(p):
        c=p['components'][-1];d=c['preview']['contact_ranges'][1]
        if fault=='face':c['contact_ranges'][0]['face']='bottom'
        elif fault=='proof':c['contact_ranges'][0]['evidence']['measurement_basis']='forged'
        elif fault=='gross':d['single_gross_area']='999'
        elif fault=='hole':d['apertures'][0]['hole_index']=99
        elif fault=='cut':d['apertures'][0]['single_area']='0'
        elif fault=='z':d['z']='1'
        elif fault=='formula':d['formula_asts']['single_net_area']={'constant':'1'}
        elif fault=='source':c['source_geometry']['foundation_items'][0]['parameters']['height']='700'
        else:p['takeoffs'][-1]['quantity']='999'
    mutate(pid,forge);coll='takeoffs' if fault=='book' else 'components';approve(r,pid,coll,get(e,pid)[coll][-1],expected=422 if fault=='face' else 404 if fault=='proof' else 409)


@pytest.mark.parametrize('gate',['preview','save','review','calculate','book_review'])
def test_original_bytes_five_gates(clients,gate):
    e,r,pid,s,b,eid=setup(clients)
    if gate in ('review','calculate','book_review'):c=post(e,pid,b).json()
    if gate in ('calculate','book_review'):approve(r,pid,'components',c)
    if gate=='book_review':t=post(e,pid,{},'/takeoffs/calculate').json()[-1]
    doc=get(e,pid)['documents'][-1];(store.ROOT/'originals'/doc['hash']).write_bytes(b'corrupt')
    if gate=='preview':post(e,pid,b,'/components/preview',409)
    elif gate=='save':post(e,pid,b,expected=409)
    elif gate=='review':approve(r,pid,'components',c,expected=409)
    elif gate=='calculate':post(e,pid,{},'/takeoffs/calculate',409)
    else:approve(r,pid,'takeoffs',t,expected=409)


def test_revision_exact_history_old_boq_concurrency_and_source_dependency(clients):
    e,r,pid,s,b,eid=setup(clients);c=post(e,pid,b).json();approve(r,pid,'components',c)
    for t in post(e,pid,{},'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
    post(e,pid,{},'/boq/generate');p=get(e,pid);old=p['components'][-1];b['object_hash']=store.digest(old);b['contact_ranges']=[contact(eid)];url=f'/api/v1/projects/{pid}/components/'+old['id'];res=e.put(url,json=b);assert res.status_code==200,res.text;assert res.json()['preview']['quantity']=='4.400';p2=get(e,pid);assert p2['component_history'][-1]==old and p2['boq']==p['boq'] and all(t['state']=='stale' for t in p2['takeoffs']);assert e.put(url,json=b).status_code==409;approve(r,pid,'components',res.json())
    s['object_hash']=store.digest(s);s['foundation_items'][0]['parameters']['height']='700';res=e.put(f'/api/v1/projects/{pid}/components/'+s['id'],json=s);assert res.status_code==200,res.text;approve(r,pid,'components',res.json());assert get(e,pid)['components'][-1]['state']=='stale'


def test_roles_tenants_foreign_proof_source_and_no_empty_old_tables(clients):
    e,r,pid,s,b,eid=setup(clients);post(r,pid,b,expected=403);assert clients['other'].get(f'/api/v1/projects/{pid}/foundation-horizontal-formwork-candidates',params=dict(component_id=s['id'])).status_code==404
    other=create(e)['id'];upload(e,other,'异项目');b['contact_ranges'][0]['evidence']['measurement_basis']=get(e,other)['evidence'][-1]['id'];post(e,pid,b,expected=404);assert tables(dict(takeoffs=[]))==[]


@pytest.mark.parametrize('shape',['rectangular','stepped_rectangular','circular','stepped_circular'])
@pytest.mark.parametrize('fmt',['json','xlsx','pdf'])
def test_three_formats_and_old_quote_isolation(clients,shape,fmt):
    e,r,pid,_,q=priced_project(clients);e,r,pid,s,b,eid=setup(clients,pid,shape);c=post(e,pid,b).json();approve(r,pid,'components',c)
    for t in post(e,pid,{},'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
    boq=post(e,pid,{},'/boq/generate').json();assert any(d['quantity']==c['preview']['quantity'] and d['unit']=='m²' for d in boq);url=f'/api/v1/projects/{pid}/exports';res=e.get(url,params=dict(format=fmt,unpriced=True));assert res.status_code==200,res.text
    if fmt=='json':assert set(TITLES)<={d[0] for d in res.json()['tables']}
    elif fmt=='xlsx':assert set(TITLES)<=set(load_workbook(io.BytesIO(res.content)).sheetnames)
    else:assert all(v in ''.join(p.extract_text() for p in PdfReader(io.BytesIO(res.content)).pages) for v in TITLES)
    assert not any(d[0] in TITLES for d in e.get(url,params=dict(format='json',quote_id=q['id'])).json()['tables'])
