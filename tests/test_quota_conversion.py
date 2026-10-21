"""Synthetic conversion arithmetic and source/review/snapshot contracts; no official coefficients."""
import copy
from decimal import Decimal
import pytest
from backend.engine import calculate, DomainError
from backend.quote_comparison import resource_changes
from backend.reports import tables
from backend import store, main as core
from tests.test_workflow import priced_project, get, approve


def conversion(eid='synthetic'):
    return {'quota_unit': '10m³', 'unit_factor': '.1',
            'category_factors': {'人工': '1.2', '材料': '1.1', '机械': '1.3', '设备': '1'},
            'source': '纯数学换算测试', 'license': '仅用于测试', 'reason': '合成单位和系数，无专业效力', 'evidence_id': eid}


def example():
    return {'id': 'line', 'name': '合成清单', 'quantity': '15', 'quotas': [
        {'code': 'TEST', 'coefficient': '2', 'conversion': conversion(), 'resources': [
            {'name': k, 'type': k, 'consumption': '3', 'price': '10'} for k in ['人工', '材料', '机械', '设备']]}]}


def test_unit_and_category_factors_apply_once_and_feed_fees():
    item = example()
    result = calculate([item], {'fees': [{'name': '人工附加', 'rate': '.1', 'basis': 'labor'}]})
    assert result['lines'][0]['unit_price'] == '27.60'
    assert result['subtotal'] == '414.00'
    assert result['resource_bases']['labor'] == '108.00'
    assert result['fees'][0]['amount'] == '10.80'
    assert result['total'] == '424.80'
    row = result['lines'][0]['resources_snapshot'][0]
    assert Decimal(row['coefficient']) == Decimal('.24')
    assert Decimal(row['effective_consumption']) == Decimal('.72')
    assert row['conversion'] == conversion()
    assert item == example(), 'calculation must not mutate inputs'


@pytest.mark.parametrize('field,value', [('unit_factor','0'),('unit_factor','-1'),('unit_factor','NaN'),
    ('unit_factor','1e13'),('category_factors',{'人工':'1'}),('category_factors',[]),
    ('source',' '),('reason',''),('evidence_id',None),('quota_unit',True)])
def test_invalid_conversion_rejected(field,value):
    item = example(); item['quotas'][0]['conversion'][field] = value
    with pytest.raises(DomainError):
        calculate([item], {'fees': []})


@pytest.mark.parametrize('value', [None,[],{}, {'unexpected': 1}])
def test_malformed_conversion_rejected(value):
    item = example(); item['quotas'][0]['conversion'] = value
    with pytest.raises(DomainError):
        calculate([item], {'fees': []})


def test_legacy_and_converted_quote_review_export_and_comparison(clients):
    e,r,pid,comp,old = priced_project(clients)
    base = '/api/v1/projects/'+pid
    b = get(e,pid)['boq'][0]
    assert 'conversion' not in old['lines'][0]['resources_snapshot'][0]
    quotas = copy.deepcopy(b['quotas'])
    quotas[0]['conversion'] = conversion(quotas[0]['evidence_id'])
    response=e.put(base+'/boq/'+b['id'],json={'quotas':quotas})
    assert response.status_code==200,response.text
    assert e.post(base+'/quotes/calculate').status_code == 409
    approve(r,pid,'boq',response.json())
    new=e.post(base+'/quotes/calculate')
    assert new.status_code==200,new.text
    new=new.json()
    assert new['lines'][0]['unit_price']=='11.55'
    assert new['total']=='178.45'
    p=get(e,pid)
    assert next(q for q in p['quotes'] if q['id']==old['id'])==old
    changes=resource_changes(old['lines'][0],new['lines'][0])
    assert 'conversion' in changes[0]['fields']
    report=dict((title, rows) for title,headers,rows in tables(p,new,False)[1])
    assert ['定额单位','10m³'] in [row[2:] for row in report['定额换算明细']]
    assert ['材料系数','1.1'] in [row[2:] for row in report['定额换算明细']]
    assert str(new['lines'][0]['resources_snapshot'][0]['coefficient'])=='0.11'
    import io
    from openpyxl import load_workbook
    for fmt in ['xlsx', 'pdf', 'json']:
        exported=e.get(base+'/exports',params={'format':fmt,'quote_id':new['id']})
        assert exported.status_code==200,exported.text[:100] if fmt=='json' else exported.status_code
        if fmt=='xlsx':
            workbook=load_workbook(io.BytesIO(exported.content))
            assert '定额换算明细' in workbook.sheetnames
        if fmt=='pdf':
            assert exported.content.startswith(b'%PDF')


def test_foreign_or_superseded_conversion_evidence_cannot_be_saved_or_reviewed(clients):
    e,r,pid,comp,old=priced_project(clients)
    base='/api/v1/projects/'+pid
    b=get(e,pid)['boq'][0]
    quotas=copy.deepcopy(b['quotas']);quotas[0]['conversion']=conversion('not-in-project')
    before=get(e,pid)
    assert e.put(base+'/boq/'+b['id'],json={'quotas':quotas}).status_code==404
    assert get(e,pid)==before
    quotas[0]['conversion']['evidence_id']=quotas[0]['evidence_id']
    updated=e.put(base+'/boq/'+b['id'],json={'quotas':quotas}).json()
    # Simulate an original being superseded without rewriting the saved quota.
    with store.transaction() as c:
        p=core.task_project(c,pid,{'tenant':'local'})
        p['documents'][0]['superseded']=True
        store.save(c,p)
    approve(r,pid,'boq',updated,expected=409)
    assert any('原始证据' in x for x in get(e,pid)['blockers'])
