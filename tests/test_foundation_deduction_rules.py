"""Per-hole quantity thresholds keep physical geometry and measured deductions apart."""
import copy,io,random
from decimal import Decimal,localcontext,ROUND_HALF_UP
import pytest
from openpyxl import load_workbook
from pypdf import PdfReader
from backend import store
from backend.compound_takeoff import PI
from backend.engine import DomainError
from backend.post_foundation import calculate
from tests.test_foundation_voids import body as geometry_body,setup as geometry_setup,rectangle,SHAPES
from tests.test_foundation_openings import body as through_body,hole
from tests.test_foundation_recesses import body as recess_body,interval
from tests.test_post_foundation import source,equal_body,fixed_body,fixed_added,equal_added,fixed_adjusted,equal_adjusted,combined
from tests.test_route_takeoff import evaluate
from tests.test_polygon_takeoff import post,mutate
from tests.test_workflow import get,approve,upload,create,priced_project

TITLES={'基础孔洞扣减阈值','基础孔洞实际与计量'}
QUANTITIES=dict(zip(SHAPES,('2.620','2.356','3.712','2.330')))


def rule(eid='test',metric='section_area',comparison='gt',threshold='.002',unit=None):
 return dict(metric=metric,comparison=comparison,threshold=threshold,unit=unit or ('m²' if metric=='section_area' else 'm³'),basis='仅合成规则：按本基础单孔未修约指标与阈值比较，符合扣全部实际相交体积，否则不扣，实体孔保留；非官方标准或真实合同。',evidence_id=eid,confirmed=True)


def body(s,eid='test',shape='stepped_circular'):
 b=geometry_body(s,eid,shape);b['name']='合成基础孔洞阈值计量（非真实工程）';b['measurement_basis']='仅合成单孔阈值验收：孔洞几何沿原件，每基础单孔截面积严格大于.002m²才扣全部实际孔体积；矩形.006m²扣，圆孔.0004πm²不扣，实际孔和实体净量保留，不以层或基础合计凑阈值。';d=b['foundation_items'][0];d['measurement_basis']=b['measurement_basis'];d['through_holes']['deduction_rule']=rule(eid);return b


def setup(clients,pid=None,shape='stepped_circular'):
 e,r,pid,s,_,_=geometry_setup(clients,pid,shape)
 upload(e,pid,'合成孔洞扣减阈值软件验收，非真实工程、非官方标准或真实合同。22基础44孔，来源材料79.2kg保持。下圆层Φ600×净高200mm，上圆层Φ400×净高400mm。H1矩形100×60mm，中心(-80,0)mm，Z0至150mm；H2圆孔Φ40mm，中心(80,0)mm，Z300至600mm。基础底Z=0向上。H1面积.006m2、实际单孔体积.0009m3；H2面积.0004πm2、实际单孔体积.00012πm3。每基础单孔截面积严格大于.002m2时扣全部实际孔体积，等于不扣：仅H1扣，全部计量扣减.0198m3，圆台阶计量2.330m3；全部实际孔体积.0198+.00264πm3、实际实体净2.321817500…m3不变。阈值改.001m2则两孔扣，计量2.322m3，物理孔不变。等值边界验收阈值.006m2：严格大于时两孔不扣，计量2.350m3；大于等于时H1扣，计量2.330m3。体积指标可按单孔实际相交体积比较，例如.0009m3等值；不得用22基础合计、层合计、四舍五入后值或半径代替。面积支持m2/cm2/mm2；体积支持m3/cm3/mm3，不能跨维度。几何越界重叠即使不扣也阻断。其它形态同孔样按.002m2阈值，矩形计量2.620、矩形台阶2.356、圆柱3.712m3。TEST-THRESHOLD仅合成净量映射编码，非官方条目。','foundation-threshold-basis.txt')
 for c in get(e,pid)['components']:approve(r,pid,'components',c)
 s=get(e,pid)['components'][-1];eid=get(e,pid)['evidence'][-1]['id'];return e,r,pid,s,body(s,eid,shape),eid


@pytest.mark.parametrize('shape',SHAPES)
def test_geometry_and_measured_balances_all_asts_unchanged_inputs(shape):
 s=source();b=body(s,shape=shape);before=copy.deepcopy((b,s));r=calculate(b,s);plain=geometry_body(s,shape=shape);physical=calculate(plain,s);assert (b,s)==before and r['quantity']==QUANTITIES[shape] and r['rule_version']=='railing-post-foundation-1.6' and r['foundation_opening_summary']['total_hole_count']=='44'
 with localcontext() as ctx:
  ctx.prec=80;assert Decimal(r['deduction'])==Decimal('.0198') and r['foundation_opening_summary']['geometric_opening_volume']==physical['foundation_opening_summary']['opening_volume'];assert Decimal(r['foundation_opening_summary']['geometric_net_volume'])==Decimal(physical['foundation_summary']['unrounded_quantity']);assert evaluate(r['formula_ast'],r['dimensions'])==Decimal(r['quantity'])
  for d in r['foundation_items']:
   assert d['opening_deduction_rule']['eligible_count']==1 and d['opening_deduction_rule']['excluded_count']==1
   for k,a in d['formula_asts'].items():assert evaluate(a,r['dimensions'])==Decimal(d[{'single_volume':'single_foundation_volume'}.get(k,k)])
   for v in d.get('foundation_levels',[]):
    for k,a in v['formula_asts'].items():assert evaluate(a,r['dimensions'])==Decimal(v[k])
   for h in d['opening_holes']:
    dec=h['deduction_decision'];assert evaluate(dec['metric_formula_ast'],r['dimensions'])==Decimal(dec['metric_value']);assert evaluate(dec['threshold_difference_ast'],r['dimensions'])==Decimal(dec['threshold_difference']);assert evaluate(h['deduction_formula_ast'],r['dimensions'])==Decimal(h['deducted_volume']);assert evaluate(h['formula_ast'],r['dimensions'])==Decimal(h['single_volume'])
    for seg in h['segments']:
     for k,a in seg['formula_asts'].items():assert evaluate(a,r['dimensions'])==Decimal(seg[k])


@pytest.mark.parametrize('metric,threshold',[('section_area','.006'),('single_volume','.0009')])
@pytest.mark.parametrize('comparison',['gt','gte'])
@pytest.mark.parametrize('delta',['0','.000000000001','-.000000000001'])
def test_exact_threshold_boundaries(metric,threshold,comparison,delta):
 s=source();b=body(s);c=b['foundation_items'][0]['through_holes'];c['deduction_rule']=rule(metric=metric,comparison=comparison,threshold=str(Decimal(threshold)+Decimal(delta)));r=calculate(b,s);d=r['foundation_items'][0];dec=d['opening_holes'][0]['deduction_decision'];expected=Decimal(delta)<0 or delta=='0' and comparison=='gte';assert dec['eligible']==expected and Decimal(r['deduction'])==(Decimal('.0198') if expected else Decimal(0))


@pytest.mark.parametrize('shape',SHAPES)
@pytest.mark.parametrize('make',[equal_body,fixed_body,fixed_added,equal_added,fixed_adjusted,equal_adjusted,combined,lambda s:combined(s,equal=True)])
def test_threshold_compares_single_hole_before_count_and_repeat(shape,make):
 s=source(make);r=calculate(body(s,shape=shape),s)
 with localcontext() as ctx:ctx.prec=80;assert Decimal(r['deduction'])==Decimal('.0009')*r['foundation_summary']['single_group_count']*2


@pytest.mark.parametrize('metric,unit,threshold',[('section_area','m²','.002'),('section_area','cm²','20'),('section_area','mm²','2000'),('single_volume','m³','.0005'),('single_volume','cm³','500'),('single_volume','mm³','500000')])
@pytest.mark.parametrize('shape',SHAPES)
def test_threshold_units_do_not_depend_on_hole_or_foundation_units(metric,unit,threshold,shape):
 s=source();b=body(s,shape=shape);b['foundation_items'][0]['through_holes']['deduction_rule']=rule(metric=metric,unit=unit,threshold=threshold)
 with localcontext() as ctx:ctx.prec=6;r=calculate(b,s)
 assert r==calculate(b,s) and r['quantity']==QUANTITIES[shape]


@pytest.mark.parametrize('make',[through_body,recess_body,geometry_body])
@pytest.mark.parametrize('metric',['section_area','single_volume'])
@pytest.mark.parametrize('comparison',['gt','gte'])
def test_zero_threshold_all_included_and_large_threshold_none(make,metric,comparison):
 s=source();b=make(s);old=calculate(b,s);c=b['foundation_items'][0]['through_holes'];c['deduction_rule']=rule(metric=metric,comparison=comparison,threshold='0');r=calculate(b,s);assert r['quantity']==old['quantity'] and Decimal(r['deduction'])==Decimal(old['deduction']);c['deduction_rule']['threshold']='100';r=calculate(b,s);assert Decimal(r['deduction'])==0 and r['foundation_opening_summary']['geometric_opening_volume']==old['foundation_opening_summary']['opening_volume'] and all(Decimal(h['single_volume'])>0 and h['deducted_volume']=='0' for h in r['foundation_items'][0]['opening_holes'])


def test_total_hole_volume_not_split_layers_and_mixed_ruled_unruled_solid():
 s=source();b=body(s);c=b['foundation_items'][0]['through_holes'];h=c['holes'][0];h['vertical_range']=interval('100','300');c['holes']=[h];c['deduction_rule']=rule(metric='single_volume',threshold='.001');r=calculate(b,s);assert r['foundation_items'][0]['opening_holes'][0]['deduction_decision']['eligible'] and Decimal(r['deduction'])==Decimal('.0264')
 from tests.test_post_foundation import body as solid_body
 s['material_items'] += [{**copy.deepcopy(s['material_items'][0]),'name':n} for n in ('合成全扣孔项','合成无孔项')];b=body(s);old=geometry_body(s)['foundation_items'][0];old['material_index']=2;solid=solid_body(s)['foundation_items'][0];solid['material_index']=3;b['foundation_items'] += [old,solid];r=calculate(b,s);assert r['foundation_summary']['total_count']=='66' and r['foundation_opening_summary']['total_hole_count']=='88'
 with localcontext() as ctx:ctx.prec=80;assert Decimal(r['deduction'])==Decimal('.0198')*2+Decimal(PI)*Decimal('.00264')
 b['foundation_items'].reverse();assert calculate(b,s)==r


@pytest.mark.parametrize('fault',['outside','overlap','full-layer'])
def test_non_deducted_holes_still_have_geometry_validation(fault):
 s=source();b=body(s,shape='rectangular' if fault=='full-layer' else 'stepped_circular');c=b['foundation_items'][0]['through_holes'];c['deduction_rule']['threshold']='100'
 if fault=='outside':c['holes'][0]['parameters']['x']='500'
 elif fault=='overlap':c['holes'][0]['parameters']['x']=c['holes'][1]['parameters']['x']='0';c['holes'][0]['vertical_range']['parameters']['z_end']='400'
 else:c['holes']=[rectangle(x='0',length='400',width='500')]
 with pytest.raises(DomainError,match='越出|重叠|净体积'):calculate(b,s)


def test_seeded_independent_decisions_and_measurement_rounding():
 rng=random.Random(9201)
 for _ in range(60):
  s=source();b=body(s);c=b['foundation_items'][0]['through_holes'];h=c['holes'][0];length=rng.randint(1,100);width=rng.randint(1,100);h['parameters'].update(length=str(length),width=str(width));threshold=Decimal(rng.randint(0,10000))/1000000;c['deduction_rule']=rule(threshold=str(threshold));b['rounding_digits']=rng.randint(0,12);physical=copy.deepcopy(b);physical['foundation_items'][0]['through_holes'].pop('deduction_rule');r=calculate(b,s);pr=calculate(physical,s)
  with localcontext() as ctx:
   ctx.prec=80;ded=(Decimal(length*width)/1000000*Decimal('.15') if Decimal(length*width)/1000000>threshold else Decimal(0))+(Decimal(PI)*Decimal('.00012') if Decimal(PI)*Decimal('.0004')>threshold else Decimal(0));net=Decimal(pr['gross'])-ded*22;assert Decimal(r['deduction'])==ded*22 and Decimal(r['quantity'])==net.quantize(Decimal(1).scaleb(-b['rounding_digits']),rounding=ROUND_HALF_UP)


@pytest.mark.parametrize('fault',['null','list','extra','missing','metric','metric-dict','comparison','comparison-bool','negative','bool','nan','infinite','precision','converted-precision','unit','cross-unit','basis','proof','proof-bool','confirm','truthy'])
def test_strict_rule_schema_precision_dimension_and_confirmation(fault):
 s=source();b=body(s);c=b['foundation_items'][0]['through_holes'];r=c['deduction_rule']
 if fault=='null':c['deduction_rule']=None
 elif fault=='list':c['deduction_rule']=[]
 elif fault=='extra':r['factor']='1'
 elif fault=='missing':r.pop('confirmed')
 elif fault=='metric':r['metric']='all_holes_volume'
 elif fault=='metric-dict':r['metric']={}
 elif fault=='comparison':r['comparison']='lt'
 elif fault=='comparison-bool':r['comparison']=True
 elif fault in ('negative','bool','nan','infinite','precision'):r['threshold']={'negative':'-1','bool':True,'nan':'NaN','infinite':'Infinity','precision':'.1234567890123456789'}[fault]
 elif fault=='converted-precision':r.update(unit='mm²',threshold='.0000000000001')
 elif fault=='unit':r['unit']={}
 elif fault=='cross-unit':r['unit']='m³'
 elif fault=='basis':r['basis']=''
 elif fault=='proof':r['evidence_id']=''
 elif fault=='proof-bool':r['evidence_id']=True
 elif fault=='confirm':r['confirmed']=False
 else:r['confirmed']=1
 with pytest.raises(DomainError):calculate(b,s)


@pytest.mark.parametrize('shape',SHAPES)
@pytest.mark.parametrize('fmt',['json','xlsx','pdf'])
def test_workflow_measured_boq_three_exports_mapping_and_old_quote(clients,shape,fmt):
 e,r,pid,_,quote=priced_project(clients);e,r,pid,s,b,eid=setup(clients,pid,shape);before=get(e,pid);assert post(e,pid,b,'/components/preview').json()['quantity']==QUANTITIES[shape] and get(e,pid)==before;c=post(e,pid,b).json();approve(r,pid,'components',c)
 for t in post(e,pid,{},'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
 boq=post(e,pid,{},'/boq/generate').json();assert any(v['quantity']==QUANTITIES[shape] and v['unit']=='m³' for v in boq) and any(v['quantity']=='79.200' for v in boq);url=f'/api/v1/projects/{pid}/exports';res=e.get(url,params=dict(format=fmt,unpriced=True));assert res.status_code==200,res.text
 if fmt=='json':assert TITLES<={v[0] for v in res.json()['tables']}
 elif fmt=='xlsx':
  w=load_workbook(io.BytesIO(res.content));assert TITLES<=set(w.sheetnames) and w['基础孔洞扣减阈值'].row_dimensions[3].height>48
 else:assert all(v in ''.join(p.extract_text() for p in PdfReader(io.BytesIO(res.content)).pages) for v in TITLES)
 assert not any(v[0] in TITLES for v in e.get(url,params=dict(format='json',quote_id=quote['id'])).json()['tables'])
 if fmt=='json':
  from tests.test_mapping_profiles import rule_body,save_rule,preview,selection
  p=get(e,pid);mapping={**rule_body(p,'post_foundation','volume',eid),'code':'TEST-THRESHOLD'};assert e.patch(f'/api/v1/projects/{pid}/settings',json=dict(boq_standard=mapping['standard'],boq_edition=mapping['edition'],reason='仅合成孔阈值计量验收')).status_code==200;save_rule(e,r,pid,mapping);bid=next(v['id'] for v in p['boq'] if v['takeoff_id']==p['takeoffs'][-1]['id']);chosen=preview(e,pid,bid);assert chosen['state']=='unique';assert e.post(f'/api/v1/projects/{pid}/boq/apply-mappings',json=dict(selections=[selection(chosen)])).status_code==200;approve(r,pid,'boq',next(v for v in get(e,pid)['boq'] if v['id']==bid))


@pytest.mark.parametrize('fault',['threshold','comparison','proof','confirm','decision','metric','deducted','geometric','layer','formula','book'])
def test_tamper_reconstruction(clients,fault):
 e,r,pid,s,b,eid=setup(clients);c=post(e,pid,b).json()
 if fault=='book':approve(r,pid,'components',c);post(e,pid,{},'/takeoffs/calculate')
 def forge(p):
  c=p['components'][-1];d=c['preview']['foundation_items'][0];h=d['opening_holes'][0]
  if fault in ('threshold','comparison','proof','confirm'):c['foundation_items'][0]['through_holes']['deduction_rule'][{'proof':'evidence_id','confirm':'confirmed'}.get(fault,fault)]={'threshold':'.001','comparison':'gte','proof':'forged','confirm':False}[fault]
  elif fault=='decision':h['deduction_decision']['eligible']=False
  elif fault=='metric':h['deduction_decision']['metric_value']='0'
  elif fault=='deducted':h['deducted_volume']='0'
  elif fault=='geometric':d['single_geometric_net_volume']='0'
  elif fault=='layer':d['foundation_levels'][1]['geometric_opening_volume']='0'
  elif fault=='formula':h['deduction_formula_ast']={'constant':'0'}
  else:p['takeoffs'][-1]['deduction']='0'
 mutate(pid,forge);collection='takeoffs' if fault=='book' else 'components';approve(r,pid,collection,get(e,pid)[collection][-1],expected=404 if fault=='proof' else 422 if fault=='confirm' else 409)


@pytest.mark.parametrize('gate',['preview','save','review','calculate','book_review'])
def test_rule_original_bytes_gates(clients,gate):
 e,r,pid,s,b,eid=setup(clients)
 if gate in ('review','calculate','book_review'):c=post(e,pid,b).json()
 if gate in ('calculate','book_review'):approve(r,pid,'components',c)
 if gate=='book_review':t=post(e,pid,{},'/takeoffs/calculate').json()[-1]
 doc=get(e,pid)['documents'][-1];(store.ROOT/'originals'/doc['hash']).write_bytes(b'corrupt')
 if gate=='preview':post(e,pid,b,'/components/preview',409)
 elif gate=='save':post(e,pid,b,expected=409)
 elif gate=='review':approve(r,pid,'components',c,expected=409)
 elif gate=='calculate':post(e,pid,{},'/takeoffs/calculate',409)
 else:approve(r,pid,'takeoffs',t,expected=409)


def test_threshold_revision_history_boq_conflict_switch_off_source_stale(clients):
 e,r,pid,s,b,eid=setup(clients);c=post(e,pid,b).json();approve(r,pid,'components',c)
 for t in post(e,pid,{},'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
 post(e,pid,{},'/boq/generate');p=get(e,pid);old=p['components'][-1];new=copy.deepcopy(b);new['foundation_items'][0]['through_holes']['deduction_rule']['threshold']='.001';new['object_hash']=store.digest(old);url=f'/api/v1/projects/{pid}/components/'+old['id'];res=e.put(url,json=new);assert res.status_code==200 and res.json()['preview']['quantity']=='2.322';assert get(e,pid)['component_history'][-1]==old and get(e,pid)['boq']==p['boq'] and all(t['state']=='stale' for t in get(e,pid)['takeoffs']);assert e.put(url,json=new).status_code==409;post(e,pid,{**b,'id':old['id']},expected=409);approve(r,pid,'components',res.json())
 new=geometry_body(s,eid);new['object_hash']=store.digest(get(e,pid)['components'][-1]);res=e.put(url,json=new);assert res.status_code==200 and res.json()['preview']['quantity']=='2.322' and res.json()['preview']['rule_version']=='railing-post-foundation-1.5';approve(r,pid,'components',res.json())
 s['object_hash']=store.digest(s);s['material_items'][0]['post_ranges'][0]['spacing']='3000';res=e.put(f'/api/v1/projects/{pid}/components/'+s['id'],json=s);assert res.status_code==200;approve(r,pid,'components',res.json());assert get(e,pid)['components'][-1]['state']=='stale'


def test_roles_tenants_and_foreign_rule_evidence(clients):
 e,r,pid,s,b,eid=setup(clients);post(r,pid,b,expected=403);assert clients['other'].get(f'/api/v1/projects/{pid}').status_code==404;other=create(e)['id'];upload(e,other,'异项目扣孔阈值原件');b['foundation_items'][0]['through_holes']['deduction_rule']['evidence_id']=get(e,other)['evidence'][-1]['id'];post(e,pid,b,expected=404)


def test_twenty_holes_long_rule_pdf(clients):
 e,r,pid,s,b,eid=setup(clients);d=b['foundation_items'][0];d['through_holes']['holes']=[{**rectangle('H'+str(i)+('长名称'*15),x='0',length='10',width='20',eid=eid),'vertical_range':interval(str(i*30),str((i+1)*30),eid)} for i in range(20)];d['through_holes']['deduction_rule']['basis']='合成'+('逐孔阈值适用说明'*150);c=post(e,pid,b).json();approve(r,pid,'components',c)
 for t in post(e,pid,{},'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
 res=e.get(f'/api/v1/projects/{pid}/exports',params=dict(format='pdf',unpriced=True));assert res.status_code==200,res.text


@pytest.mark.parametrize('metric,unit,threshold',[('section_area','mm²','2000'),('single_volume','cm³','500')])
def test_foundation_m_hole_cm_threshold_independent_dimension_units(metric,unit,threshold):
 s=source();b=body(s);d=b['foundation_items'][0];d['length_unit']='m'
 for layer in d['levels']:
  for k,v in layer['parameters'].items():layer['parameters'][k]=str(Decimal(v)*Decimal('.001'))
 c=d['through_holes'];c['unit']='cm'
 for h in c['holes']:
  for params in (h['parameters'],h['vertical_range']['parameters']):
   for k,v in params.items():params[k]=str(Decimal(v)*Decimal('.1'))
 c['deduction_rule']=rule(metric=metric,unit=unit,threshold=threshold);assert calculate(b,s)['quantity']=='2.330'


@pytest.mark.parametrize('fmt',['json','xlsx','pdf'])
def test_rectangular_only_rule_geometry_exports_without_pi(clients,fmt):
 e,r,pid,s,b,eid=setup(clients,shape='rectangular');b['foundation_items'][0]['through_holes']['holes'].pop();c=post(e,pid,b).json();assert 'foundation_pi_constant' not in c['preview'];approve(r,pid,'components',c)
 for t in post(e,pid,{},'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
 res=e.get(f'/api/v1/projects/{pid}/exports',params=dict(format=fmt,unpriced=True));assert res.status_code==200,res.text
 if fmt=='json':assert TITLES<={v[0] for v in res.json()['tables']}
