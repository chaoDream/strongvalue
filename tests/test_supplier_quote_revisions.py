import copy
import io
from openpyxl import load_workbook
from pypdf import PdfReader
import pytest
from backend import store
from tests.test_supplier_quotes import fixture,register,read,adopt,mutate
from tests.test_workflow import upload,get,approve


def setup(clients,selected=True):
    c,r,pid,old,body=fixture(clients)
    d=upload(c,pid,'合成供应商A撤回A-001。修订A-002单价91.123456元/m³，不含税，含运输不含安装，验收后付款。仅软件测试。','revision.txt','合同计价')['document']
    eid=next(e['id'] for e in get(c,pid)['evidence'] if e['document_id']==d['id'])
    for coll in ['components','takeoffs']:
        for item in get(c,pid)[coll]:approve(r,pid,coll,item)
    row=c.get(f'/api/v1/projects/{pid}/resource-prices').json()[0]
    body.update(boq_hash=row['boq_hash'],resource_hash=row['resource_hash'])
    o=register(c,pid,body)['offer'];approve(r,pid,'supplier_quotes',o);o=get(c,pid)['supplier_quotes'][0]
    if selected:adopt(c,pid,o)
    return c,r,pid,old,body,o,eid


def reopen_body(c,pid,offer,eid,**extra):
    row=next(x for x in read(c,pid)['items'] if x['offer']['id']==offer['id'])
    return {'request_id':'synthetic-reopen-request-001','object_hash':row['object_hash'],**(row['reopen_target'] or {}),
            'reason':'对照合成撤回说明，撤出已选报价并重新询价；数量和消耗不变','evidence_id':eid,'confirm_reopen':True,**extra}


def reopen(c,pid,qid,body,expected=200):
    response=c.post(f'/api/v1/projects/{pid}/supplier-quotes/{qid}/reopen-inquiry',json=body)
    assert response.status_code==expected,response.text
    return response.json()


def revision_body(c,pid,body,offer,eid,**extra):
    row=c.get(f'/api/v1/projects/{pid}/resource-prices').json()[0]
    current=next(o for o in get(c,pid)['supplier_quotes'] if o['id']==offer['id'])
    return {**body,'request_id':'synthetic-revised-offer-001','boq_hash':row['boq_hash'],'resource_hash':row['resource_hash'],
            'revision_of':offer['id'],'revision_hash':store.digest(current),'price':'91.123456','reference':'SYNTHETIC-A-R2',
            'evidence_id':eid,'reason':'对照修订报价逐项核对规格、单位、有效期及新单价',**extra}


def test_reopen_revision_reapprove_reprice_preserves_history_and_quantities(clients):
    c,r,pid,old,body,o,eid=setup(clients)
    approve(r,pid,'boq',get(c,pid)['boq'][0])
    q=c.post(f'/api/v1/projects/{pid}/quotes/calculate').json();before=get(c,pid)
    request=reopen_body(c,pid,o,eid);rec=reopen(c,pid,o['id'],request)['reopening']
    p=get(c,pid);assert p['boq'][0]['state']=='draft' and len(p['pricing_tasks'])==1
    pending=p['boq'][0]['quotas'][0]['resources'][0]
    assert pending['price'] is None and pending['pricing_status']=='unpriced'
    assert not any(k in pending for k in ['source','license','period','tax_mode','valid_from','valid_until','supplier_quote_ref'])
    for k in ['quantity','code','features']:assert p['boq'][0][k]==before['boq'][0][k]
    for k in ['resource_code','name','type','unit','specification','fee_allocation','consumption']:
        assert pending[k]==rec['before'][k]
    assert p['supplier_quotes']==before['supplier_quotes'] and p['supplier_quote_adoptions']==before['supplier_quote_adoptions']
    assert q['id'] in p['stale_quotes'] and next(x for x in p['quotes'] if x['id']==q['id'])==q
    frozen=copy.deepcopy(p);assert reopen(c,pid,o['id'],request)['duplicate'] and get(c,pid)==frozen
    rb=revision_body(c,pid,body,o,eid);new=register(c,pid,rb)['offer']
    assert new['revision']==2 and new['family_id']==o['id'] and new['previous_quote_hash']==store.digest(o)
    assert new['state']=='draft' and register(c,pid,rb)['duplicate']
    assert len(get(c,pid)['pricing_tasks'])==1
    approve(c,pid,'supplier_quotes',new,expected=403);approve(r,pid,'supplier_quotes',new)
    new=get(c,pid)['supplier_quotes'][-1];adopt(c,pid,new)
    p=get(c,pid);assert p['pricing_tasks']==[] and p['boq'][0]['state']=='draft'
    assert p['boq'][0]['quotas'][0]['resources'][0]['price']=='91.123456'
    assert 'supplier_inquiry_ref' not in p['boq'][0]['quotas'][0]['resources'][0]
    approve(r,pid,'boq',p['boq'][0])
    result=c.post(f'/api/v1/projects/{pid}/quotes/calculate');assert result.status_code==200,result.text
    # 91.123456 × 1.05 = 95.6796288, rounded 95.68 × 15 = 1435.20.
    assert result.json()['subtotal']=='1435.20'
    events=c.get(f'/api/v1/projects/{pid}/audit').json()
    assert [x['action'] for x in events].count('SUPPLIER_INQUIRY_REOPENED')==1


def test_pending_quote_revision_disables_old_new_adoption_without_mutating_old_row(clients):
    c,r,pid,old,body,o,eid=setup(clients,False);original=copy.deepcopy(o);p=get(c,pid)
    new=register(c,pid,revision_body(c,pid,body,o,eid))['offer']
    assert next(x for x in get(c,pid)['supplier_quotes'] if x['id']==o['id'])==original
    assert get(c,pid)['boq']==p['boq'] and get(c,pid)['stale_quotes']==p['stale_quotes']
    items=read(c,pid)['items'];oldrow=next(x for x in items if x['offer']['id']==o['id'])
    assert not oldrow['latest_revision'] and not oldrow['available'] and not oldrow['can_revise']
    assert '修订版本' in oldrow['issue']
    approve(r,pid,'supplier_quotes',o,expected=409);adopt(c,pid,o,409)
    approve(r,pid,'supplier_quotes',new);new=get(c,pid)['supplier_quotes'][-1];adopt(c,pid,new)


def test_revoked_or_missing_old_source_can_recover_with_new_evidence(clients):
    c,r,pid,old,body,o,eid=setup(clients)
    approve(r,pid,'supplier_quotes',o,'reject')
    (store.ROOT/'originals'/o['proof']['document_hash']).unlink()
    request=reopen_body(c,pid,o,eid);assert request['boq_hash']
    reopen(c,pid,o['id'],request)
    new=register(c,pid,revision_body(c,pid,body,o,eid))['offer']
    approve(r,pid,'supplier_quotes',new);new=get(c,pid)['supplier_quotes'][-1];adopt(c,pid,new)
    approve(r,pid,'boq',get(c,pid)['boq'][0])
    response=c.post(f'/api/v1/projects/{pid}/quotes/calculate');assert response.status_code==200,response.text


@pytest.mark.parametrize('extra',[
    {'reason':' '},{'confirm_reopen':False},{'confirm_reopen':'true'},{'confirm_reopen':1},
    {'price':'0'},{'resource_hash':'bad'},{'request_id':'short'},{'reason':'x'*2001}])
def test_invalid_reopening_is_atomic(clients,extra):
    c,r,pid,old,body,o,eid=setup(clients);request=reopen_body(c,pid,o,eid,**extra);before=get(c,pid)
    reopen(c,pid,o['id'],request,422)
    assert get(c,pid)==before


def test_reopen_permissions_idempotency_and_stale_hashes(clients):
    c,r,pid,old,body,o,eid=setup(clients);request=reopen_body(c,pid,o,eid)
    reopen(r,pid,o['id'],request,403);reopen(clients['other'],pid,o['id'],request,404)
    for k in ['object_hash','boq_hash','resource_hash']:reopen(c,pid,o['id'],{**request,k:'0'*64},409)
    reopen(c,pid,o['id'],request)
    reopen(c,pid,o['id'],{**request,'reason':'不能复用请求标识改写原原因'},409)
    reopen(c,pid,o['id'],{**request,'request_id':'synthetic-another-reopen-001'},409)


@pytest.mark.parametrize('change',['price','specification','reference','missing_resource','missing_proof','proof_bytes','proof_text'])
def test_changed_resource_or_reopen_evidence_blocks_atomic_reopening(clients,change):
    c,r,pid,old,body,o,eid=setup(clients);request=reopen_body(c,pid,o,eid)
    def action(p):
        resource=p['boq'][0]['quotas'][0]['resources'][0]
        if change=='price':resource['price']='1'
        elif change=='specification':resource['specification']='other'
        elif change=='reference':resource['supplier_quote_ref']['adoption_id']='foreign'
        elif change=='missing_resource':p['boq'][0]['quotas'][0]['resources']=[]
        elif change=='proof_text':next(e for e in p['evidence'] if e['id']==eid)['document_hash']='0'*64
    if change in ('missing_proof','proof_bytes'):
        ev=next(e for e in get(c,pid)['evidence'] if e['id']==eid);path=store.ROOT/'originals'/ev['document_hash']
        if change=='missing_proof':path.unlink()
        else:path.write_bytes(b'bad source')
    else:mutate(pid,action)
    before=get(c,pid);reopen(c,pid,o['id'],request,409);assert get(c,pid)==before


@pytest.mark.parametrize('extra',[
    {'revision_of':None},{'revision_hash':None},{'supplier':'另一供应商'},
    {'revision_hash':'0'*64},{'revision_of':'foreign-quote'}])
def test_invalid_revision_parent_or_identity_is_atomic(clients,extra):
    c,r,pid,old,body,o,eid=setup(clients,False);request=revision_body(c,pid,body,o,eid,**extra)
    before=get(c,pid);register(c,pid,request,422 if None in extra.values() else 404 if extra.get('revision_of') else 409)
    assert get(c,pid)==before


def test_revision_is_linear_and_old_rejection_does_not_stale_new_selected_price(clients):
    c,r,pid,old,body,o,eid=setup(clients,False)
    rb=revision_body(c,pid,body,o,eid);new=register(c,pid,rb)['offer']
    register(c,pid,{**rb,'request_id':'synthetic-branch-revision-0002'},409)
    third=register(c,pid,revision_body(c,pid,body,new,eid,request_id='synthetic-revision-0003'))['offer']
    assert third['revision']==3 and third['previous_quote_id']==new['id']
    approve(r,pid,'supplier_quotes',third);third=get(c,pid)['supplier_quotes'][-1];adopt(c,pid,third)
    approve(r,pid,'boq',get(c,pid)['boq'][0]);q=c.post(f'/api/v1/projects/{pid}/quotes/calculate').json()
    before=get(c,pid);approve(r,pid,'supplier_quotes',o,'reject');after=get(c,pid)
    assert before['boq']==after['boq'] and q['id'] not in after['stale_quotes']


@pytest.mark.parametrize('change',['record','proof','bytes','resource'])
def test_reopening_provenance_checked_during_revision_review_and_adoption(clients,change):
    c,r,pid,old,body,o,eid=setup(clients);reopen(c,pid,o['id'],reopen_body(c,pid,o,eid))
    new=register(c,pid,revision_body(c,pid,body,o,eid))['offer'];approve(r,pid,'supplier_quotes',new);new=get(c,pid)['supplier_quotes'][-1]
    def action(p):
        if change=='record':p['supplier_inquiry_reopenings'][0]['reason']='tampered'
        elif change=='proof':next(e for e in p['evidence'] if e['id']==eid)['text']+='changed'
        elif change=='resource':p['boq'][0]['quotas'][0]['resources'][0]['consumption']='2'
    if change=='bytes':(store.ROOT/'originals'/new['proof']['document_hash']).write_bytes(b'corrupt')
    else:mutate(pid,action)
    before=get(c,pid);approve(r,pid,'supplier_quotes',new,expected=409);adopt(c,pid,new,409)
    assert get(c,pid)==before


def test_legacy_quotes_revision_and_prior_request_idempotency(clients):
    c,r,pid,old,body,o,eid=setup(clients,False)
    mutate(pid,lambda p:[x.pop(k,None) for x in p['supplier_quotes'] for k in ['family_id','revision','previous_quote_id','previous_quote_hash']])
    assert register(c,pid,body)['duplicate']
    o=get(c,pid)['supplier_quotes'][0];new=register(c,pid,revision_body(c,pid,body,o,eid))['offer']
    assert new['family_id']==o['id'] and new['revision']==2


def test_revision_and_reopening_three_format_snapshots_stay_frozen(clients):
    c,r,pid,old,body,o,eid=setup(clients);before=c.get(f'/api/v1/projects/{pid}/exports',params={'format':'json','quote_id':old['id']}).json()
    reopen(c,pid,o['id'],reopen_body(c,pid,o,eid));new=register(c,pid,revision_body(c,pid,body,o,eid))['offer']
    approve(r,pid,'supplier_quotes',new);new=get(c,pid)['supplier_quotes'][-1];adopt(c,pid,new);approve(r,pid,'boq',get(c,pid)['boq'][0])
    q=c.post(f'/api/v1/projects/{pid}/quotes/calculate').json();frozen=copy.deepcopy(q)
    for fmt in ['json','xlsx','pdf']:
        response=c.get(f'/api/v1/projects/{pid}/exports',params={'format':fmt,'quote_id':q['id']});assert response.status_code==200
        if fmt=='json':assert {'供应商报价修订','供应商重新询价'}<={t[0] for t in response.json()['tables']}
        elif fmt=='xlsx':assert '供应商重新询价' in load_workbook(io.BytesIO(response.content)).sheetnames
        else:assert '供应商报价修订' in ''.join(p.extract_text() for p in PdfReader(io.BytesIO(response.content)).pages)
    again=reopen_body(c,pid,new,eid,request_id='synthetic-second-reopening-0002');reopen(c,pid,new['id'],again)
    assert next(x for x in get(c,pid)['quotes'] if x['id']==q['id'])==frozen
    assert c.get(f'/api/v1/projects/{pid}/exports',params={'format':'json','quote_id':old['id']}).json()['tables']==before['tables']


def test_unselected_and_historical_adoption_cannot_be_reopened(clients):
    c,r,pid,old,body,o,eid=setup(clients,False)
    request={'request_id':'synthetic-unselected-reopen-001','object_hash':store.digest(o),'boq_hash':body['boq_hash'],
             'resource_hash':body['resource_hash'],'evidence_id':eid,'reason':'合成验证','confirm_reopen':True}
    before=get(c,pid);reopen(c,pid,o['id'],request,409);assert get(c,pid)==before
    adopt(c,pid,o);first=reopen_body(c,pid,o,eid);reopen(c,pid,o['id'],first)
    new=register(c,pid,revision_body(c,pid,body,o,eid))['offer'];approve(r,pid,'supplier_quotes',new)
    new=get(c,pid)['supplier_quotes'][-1];adopt(c,pid,new)
    request.update(request_id='synthetic-historical-reopen-001')
    before=get(c,pid);reopen(c,pid,o['id'],request,409);assert get(c,pid)==before


def test_project_restriction_applies_to_reopening_and_revision(clients):
    from tests.test_project_access import restrict
    c,r,pid,old,body,o,eid=setup(clients);request=reopen_body(c,pid,o,eid)
    restrict(clients['admin'],pid,[])
    reopen(c,pid,o['id'],request,404);register(c,pid,body,404)


def test_reopening_provenance_survives_new_adoption_and_is_rechecked_for_pricing(clients):
    c,r,pid,old,body,o,eid=setup(clients);reopen(c,pid,o['id'],reopen_body(c,pid,o,eid))
    new=register(c,pid,revision_body(c,pid,body,o,eid))['offer'];approve(r,pid,'supplier_quotes',new)
    new=get(c,pid)['supplier_quotes'][-1];adopt(c,pid,new)
    mutate(pid,lambda p:p['supplier_inquiry_reopenings'][0].update(reason='tampered after adoption'))
    approve(r,pid,'boq',get(c,pid)['boq'][0],expected=409)
    assert c.post(f'/api/v1/projects/{pid}/quotes/calculate').status_code==409


def test_changed_specification_is_not_a_revision_of_another_resource(clients):
    c,r,pid,old,body,o,eid=setup(clients,False)
    mutate(pid,lambda p:p['boq'][0]['quotas'][0]['resources'][0].update(specification='different specification'))
    request=revision_body(c,pid,body,o,eid);before=get(c,pid)
    register(c,pid,request,409);assert get(c,pid)==before
