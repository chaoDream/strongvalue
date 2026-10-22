"""Synthetic actual-OD envelopes, traceable arithmetic and backfill dependencies."""
import copy
import io
import json
import random
from decimal import Decimal, localcontext
import pytest
from fastapi import HTTPException
from openpyxl import load_workbook
from pypdf import PdfReader
from backend import store
from backend.engine import DomainError
from backend.pipe_envelope import calculate
from backend.route_takeoff import snapshot
from backend.compound_takeoff import PI
from backend.volume_balance import snapshot as balance_snapshot
from tests.test_route_takeoff import payload as route_payload, point
from tests.test_trench_takeoff import setup as trench_setup, evaluate
from tests.test_volume_balance import payload as balance_payload
from tests.test_workflow import get, upload, approve, create


def geometry():
    s=route_payload();s.update(id='synthetic-route',revision=1);return s


def payload(source,eid='test'):
    return {'name':'合成管道外包（非真实工程）','kind':'pipe_envelope','parameters':{},'quantity_type':'volume',
            'source_component_id':source['id'],'source_component_hash':store.digest(snapshot(source)),
            'pipe_sections':[{'segment_index':i+1,'outer_diameter':'200' if i==0 else '400','diameter_unit':'mm','evidence_id':eid,
                              'measurement_basis':'仅合成直圆管外包占位，含空腔，非管壁材料，不含管件与交接重叠','basis_evidence_id':eid} for i in range(len(source['route_points'])-1)],
            'outside_confirmed':True,'measurement_basis':'合成实际外径，按未修约三维直段长度计算外包占位，不推定DN或施工附加量','basis_evidence_id':eid,
            'rounding_digits':6,'rounding_evidence_id':eid}


def setup(clients,save=False):
    e,r,pid,route,trench,eid=trench_setup(clients);base='/api/v1/projects/'+pid
    upload(e,pid,'合成实际外径200mm，直段三维13m水平5m，重复1；外包含空腔，非管壁，不含管件交接重叠。总体积修约6位。用于同域回填须另审范围。',name='actual-od.txt')
    approve(r,pid,'components',get(e,pid)['components'][0]);route=get(e,pid)['components'][0]
    eid=get(e,pid)['evidence'][-1]['id'];b=payload(route,eid)
    if save:
        response=e.post(base+'/components',json=b);assert response.status_code==200,response.text
    return e,r,pid,route,trench,b,eid


def test_vertical_mixed_units_raw_lengths_and_inherited_count():
    s=geometry();s['multiplicity']='2';b=payload(s);before=copy.deepcopy((s,b));r=calculate(b,s)
    assert (s,b)==before and r['quantity']=='2.073451' and r['unit']=='m³'
    assert [x['spatial_length'] for x in r['pipe_sections']]==['13','5']
    assert r['pipe_sections'][1]['horizontal_length']=='0'
    assert [x['normalized_outer_diameter'] for x in r['pipe_sections']]==['0.2','0.4']
    with localcontext() as ctx:
        ctx.prec=80
        assert Decimal(r['gross'])==Decimal(PI)*Decimal('.66')
        assert evaluate(r['formula_ast'],r['dimensions'])==Decimal(r['quantity'])
        for row in r['pipe_sections']:assert evaluate(row['formula_ast'],r['dimensions'])==Decimal(row['single_volume'])
    b['pipe_sections'].reverse();assert calculate(b,s)==r


def test_source_rounded_length_is_never_reused():
    s=geometry();s['route_points']=[point('a',0,0,0),point('b',1,1,0)];s['rounding_digits']=0;b=payload(s)
    r=calculate(b,s);assert r['quantity']=='0.044429'
    assert Decimal(r['pipe_sections'][0]['spatial_length'])>1


def test_seeded_envelopes_match_independent_pythagorean_calculation():
    rng=random.Random(3201)
    for _ in range(30):
        s=geometry();k=rng.randint(1,8);s['route_points']=[point('a',-3*k,-4*k,-12*k),point('b',0,0,0)];s['multiplicity']=str(rng.randint(1,4))
        b=payload(s);d=Decimal(rng.randint(1,999))/1000;b['pipe_sections'][0].update(outer_diameter=str(d),diameter_unit='m');b['rounding_digits']=12
        with localcontext() as ctx:
            ctx.prec=80
            expected=Decimal(PI)/4*d*d*13*k*Decimal(s['multiplicity']);r=calculate(b,s)
            assert Decimal(r['quantity'])==expected.quantize(Decimal('.000000000001'))
            assert evaluate(r['formula_ast'],r['dimensions'])==Decimal(r['quantity'])


@pytest.mark.parametrize('case',['missing','duplicate','unknown','bool_index','extra','zero','negative','dn','nan','unit','precision','missing_proof','blank_basis','unconfirmed','string_confirmation','multiplicity','rounding','mode','parameters'])
def test_invalid_inputs_fail(case):
    s=geometry();b=payload(s);d=b['pipe_sections'][0]
    if case=='missing':b['pipe_sections'].pop()
    elif case=='duplicate':b['pipe_sections'][1]['segment_index']=1
    elif case=='unknown':d['segment_index']=99
    elif case=='bool_index':d['segment_index']=True
    elif case=='extra':d['wall_thickness']='2'
    elif case=='zero':d['outer_diameter']='0'
    elif case=='negative':d['outer_diameter']='-200'
    elif case=='dn':d['outer_diameter']='DN200'
    elif case=='nan':d['outer_diameter']='NaN'
    elif case=='unit':d['diameter_unit']='m³'
    elif case=='precision':d['outer_diameter']='.000000000000000001'
    elif case=='missing_proof':del d['evidence_id']
    elif case=='blank_basis':d['measurement_basis']=''
    elif case=='unconfirmed':b['outside_confirmed']=False
    elif case=='string_confirmation':b['outside_confirmed']='true'
    elif case=='multiplicity':b['multiplicity']='2'
    elif case=='rounding':b['rounding_digits']=True
    elif case=='mode':b['quantity_type']='length'
    elif case=='parameters':b['parameters']={'volume':'1'}
    with pytest.raises(DomainError):calculate(b,s)


def test_old_route_hash_fails():
    s=geometry();b=payload(s);s['revision']=2
    with pytest.raises(HTTPException) as err:calculate(b,s)
    assert err.value.status_code==409


def test_candidates_preview_review_books_boq_exports(clients):
    e,r,pid,route,trench,b,eid=setup(clients);base='/api/v1/projects/'+pid;before=get(e,pid)
    url=base+'/pipe-envelope-candidates?component_id='+route['id']
    response=e.get(url);assert response.status_code==200,response.text
    assert response.json()['route_segments'][0]['length']=='13' and get(e,pid)==before
    response=e.post(base+'/components/preview',json=b);assert response.status_code==200,response.text
    assert response.json()['quantity']=='0.408407' and get(e,pid)==before
    response=e.post(base+'/components',json={**b,'source_geometry':{'forged':True}});assert response.status_code==200,response.text
    c=response.json();assert c['source_geometry']==snapshot(route)
    assert get(e,pid)['components'][0]['state']=='approved' and e.post(base+'/takeoffs/calculate').status_code==409
    approve(r,pid,'components',c);response=e.post(base+'/takeoffs/calculate');assert response.status_code==200,response.text
    for book in response.json():approve(r,pid,'takeoffs',book)
    response=e.post(base+'/boq/generate');assert response.status_code==200,response.text
    assert response.json()[-1]['quantity']=='0.408407'
    j=e.get(base+'/exports?unpriced=true&format=json').json()['unpriced']['takeoffs'][-1]
    assert j['pipe_sections']==c['preview']['pipe_sections'] and j['source_geometry']==snapshot(route)
    response=e.get(base+'/exports?unpriced=true&format=xlsx');assert response.status_code==200,response.text
    wb=load_workbook(io.BytesIO(response.content));assert {'管道外包体积汇总','管道逐段外径与占位','管道外包口径与常量','管道外包来源坐标'}<=set(wb.sheetnames)
    assert all(c.data_type!='f' for s in wb for row in s for c in row)
    response=e.get(base+'/exports?unpriced=true&format=pdf');assert response.status_code==200,response.text
    text='\n'.join(p.extract_text() for p in PdfReader(io.BytesIO(response.content)).pages)
    assert '管道外包体积汇总' in text and '管道逐段外径与占位' in text and '0.408407' in text


def test_pipe_deduction_and_dependent_invalidation(clients):
    e,r,pid,route,trench,b,eid=setup(clients,save=True);base='/api/v1/projects/'+pid
    pipe=get(e,pid)['components'][-1];approve(r,pid,'components',pipe)
    t=e.post(base+'/components',json=trench);assert t.status_code==200,t.text
    approve(r,pid,'components',t.json());p=get(e,pid)
    src={'excavation':balance_snapshot(p['components'][-1]),'occupations':[balance_snapshot(p['components'][1])]}
    balance=balance_payload(src,eid);balance['rounding_digits']=6
    response=e.post(base+'/components',json=balance);assert response.status_code==200,response.text
    assert response.json()['preview']['quantity']=='47.091593';approve(r,pid,'components',response.json())
    books=e.post(base+'/takeoffs/calculate');assert books.status_code==200,books.text
    for book in books.json():approve(r,pid,'takeoffs',book)
    p=get(e,pid);b['object_hash']=e.get(base+'/hashes').json()['components'][pipe['id']];b['pipe_sections'][0]['outer_diameter']='250'
    bad=copy.deepcopy(b);bad['outside_confirmed']=False
    assert e.put(base+'/components/'+pipe['id'],json=bad).status_code==422 and get(e,pid)==p
    response=e.put(base+'/components/'+pipe['id'],json=b);assert response.status_code==200,response.text
    current=get(e,pid);assert [c['state'] for c in current['components']]==['approved','draft','approved','stale']
    assert all(t['state']=='stale' for t in current['takeoffs'])
    assert e.post(base+'/takeoffs/calculate').status_code==409
    approve(r,pid,'components',current['components'][1]);approve(r,pid,'components',current['components'][-1],expected=409)
    # Re-read the occupation version; the existing backfill inputs cannot pass merely by re-reviewing.
    selector={'excavation_component_id':src['excavation']['id'],'occupation_component_ids':[pipe['id']]}
    c=e.post(base+'/volume-balance-candidates',json=selector);assert c.status_code==200,c.text
    assert c.json()['occupations'][0]['quantity']=='0.638136'


@pytest.mark.parametrize('gate',['preview','save','component_review','calculate','book_review'])
def test_corrupt_diameter_original_blocks_all_gates(clients,gate):
    e,r,pid,route,trench,b,eid=setup(clients);base='/api/v1/projects/'+pid
    if gate not in ('preview','save'):comp=e.post(base+'/components',json=b).json()
    if gate in ('calculate','book_review'):approve(r,pid,'components',comp)
    if gate=='book_review':book=e.post(base+'/takeoffs/calculate').json()[-1]
    doc=get(e,pid)['documents'][-1];(store.ROOT/'originals'/doc['hash']).write_bytes(b'corrupt')
    if gate=='preview':assert e.post(base+'/components/preview',json=b).status_code==409
    elif gate=='save':assert e.post(base+'/components',json=b).status_code==409
    elif gate=='component_review':approve(r,pid,'components',comp,expected=409)
    elif gate=='calculate':assert e.post(base+'/takeoffs/calculate').status_code==409
    else:approve(r,pid,'takeoffs',book,expected=409)


def test_permissions_cross_project_and_source_change(clients):
    e,r,pid,route,trench,b,eid=setup(clients);base='/api/v1/projects/'+pid;url=base+'/pipe-envelope-candidates?component_id='+route['id']
    assert r.get(url).status_code==200 and r.post(base+'/components',json=b).status_code==403
    assert clients['other'].get(url).status_code==404
    other=create(e)['id'];upload(e,other,'其他项目依据');bad=copy.deepcopy(b);bad['pipe_sections'][0]['evidence_id']=get(e,other)['evidence'][0]['id']
    assert e.post(base+'/components',json=bad).status_code==404
    comp=e.post(base+'/components',json=b).json();approve(r,pid,'components',comp)
    route['object_hash']=e.get(base+'/hashes').json()['components'][route['id']];route['route_points'][1]['parameters']['z']='15000'
    assert e.put(base+'/components/'+route['id'],json=route).status_code==200
    assert e.get(url).status_code==409 and e.post(base+'/components/preview',json=b).status_code==409
    approve(r,pid,'components',get(e,pid)['components'][0]);assert e.post(base+'/components/preview',json=b).status_code==409


@pytest.mark.parametrize('target',['component','book'])
def test_saved_output_tampering_cannot_pass_review(clients,target):
    e,r,pid,route,trench,b,eid=setup(clients,save=True);base='/api/v1/projects/'+pid
    if target=='book':approve(r,pid,'components',get(e,pid)['components'][-1]);e.post(base+'/takeoffs/calculate')
    with store.transaction() as db:
        p=json.loads(db.execute('SELECT data FROM projects WHERE id=?',(pid,)).fetchone()['data'])
        if target=='component':p['components'][-1]['preview']['pipe_sections'][0]['single_volume']='999'
        else:p['takeoffs'][-1]['pipe_sections'][0]['single_volume']='999'
        store.save(db,p)
    col='components' if target=='component' else 'takeoffs';approve(r,pid,col,get(e,pid)[col][-1],expected=409)
