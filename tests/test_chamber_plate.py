"""Synthetic horizontal plates, concentric openings and reviewed well-count dependencies."""
import copy
import io
import json
import random
from decimal import Decimal, localcontext
import pytest
from fastapi import HTTPException
from openpyxl import load_workbook
from pypdf import PdfReader
from backend import store
from backend.compound_takeoff import PI
from backend.chamber_plate import calculate
from backend.chamber_takeoff import calculate as chamber_quantity
from backend.engine import DomainError
from backend.volume_balance import snapshot
from tests.test_chamber_takeoff import setup as chamber_setup, payload as chamber_payload, well
from tests.test_trench_takeoff import geometry, evaluate
from tests.test_volume_balance import payload as balance_payload
from tests.test_workflow import get, upload, approve, create, priced_project


def seal(source):
    source['preview']=chamber_quantity(source,source['source_geometry']);source['evidence']=source['preview']['source_evidence'];return source


def parent():
    route=geometry();s=chamber_payload(route);s.update(id='synthetic-chamber',revision=1,state='approved',source_geometry=route);return seal(s)


def plate(eid='test',index=1,diameter='2.4',opening='.6',thickness='.2',bottom='-.2',unit='m'):
    return {'name':'合成W'+str(index)+'盖板','source_well_index':index,'parameters':{'diameter':diameter,'opening_diameter':opening,'thickness':thickness,'bottom_z':bottom},'unit':unit,
            'evidence':{k:eid for k in ['diameter','opening_diameter','thickness','bottom_z','measurement_basis']},
            'measurement_basis':'仅合成水平等厚圆板，同心圆形贯通孔；板底标高独立核定，不与井筒重复扣减'}


def payload(source,eid='test'):
    return {'name':'合成圆形井盖板（非真实工程）','kind':'chamber_plate','parameters':{},'quantity_type':'volume','plate_role':'cover',
            'source_component_id':source['id'],'source_component_hash':store.digest(snapshot(source)),
            'plate_sections':[plate(eid)],'scope_confirmed':True,'material':'合成板材料（非工程选材）','material_evidence_id':eid,
            'construction_basis':'仅合成等厚圆板，标高以独立基准为准；不由井筒或管线推定，不含井盖框和钢筋','construction_evidence_id':eid,
            'measurement_basis':'仅合成同层圆板、同心通孔及净实体，已审井数继承一次；回填同域与不重复范围另审','basis_evidence_id':eid,
            'rounding_digits':6,'rounding_evidence_id':eid}


def setup(clients,save=False):
    e,r,pid,route,trench,chamber,eid=chamber_setup(clients,save=True);base='/api/v1/projects/'+pid
    upload(e,pid,'合成W1水平等厚圆盖板：实际板径2.4m、同心贯通孔径0.6m、板厚0.2m、独立板底标高-0.2m；每井本层一块板，继承井筒核定井数一次；材料及做法仅合成，不含井盖框和钢筋。净体积总体修约6位，回填范围另审。',name='plate.txt')
    for c in get(e,pid)['components']:approve(r,pid,'components',c)
    source=get(e,pid)['components'][-1];eid=get(e,pid)['evidence'][-1]['id'];b=payload(source,eid)
    if save:
        response=e.post(base+'/components',json=b);assert response.status_code==200,response.text
    return e,r,pid,source,trench,b,eid


def test_concentric_opening_net_volume_signed_elevation_and_ast():
    s=parent();b=payload(s);before=copy.deepcopy((s,b));r=calculate(b,s)
    assert (s,b)==before and r['quantity']=='0.848230' and r['unit']=='m³'
    row=r['plate_sections'][0];assert row['normalized_parameters']['bottom_z']=='-0.2' and Decimal(row['top_z'])==0
    with localcontext() as ctx:
        ctx.prec=80
        assert Decimal(r['gross'])==Decimal(PI)*Decimal('.288')
        assert Decimal(r['deduction'])==Decimal(PI)*Decimal('.018')
        assert evaluate(r['formula_ast'],r['dimensions'])==Decimal(r['quantity'])
        assert evaluate(row['formula_ast'],r['dimensions'])==Decimal(row['volume'])


def test_solid_base_explicit_zero_hole_and_elevation_never_infers_thickness():
    s=parent();b=payload(s);b['plate_role']='base';b['plate_sections'][0]['parameters'].update(opening_diameter='0',bottom_z='-800')
    r=calculate(b,s);assert r['quantity']=='0.904779' and Decimal(r['deduction'])==0
    assert r['plate_sections'][0]['top_z']=='-799.8'


def test_multiple_wells_mixed_units_inherited_count_and_parent_quantity_not_reused():
    s=parent();s['chamber_wells'][0]['multiplicity']='2';s['chamber_wells'].append(well(2));s['rounding_digits']=0;seal(s)
    b=payload(s);b['plate_sections'].append(plate(index=2,diameter='1200',opening='0',thickness='100',bottom='1000',unit='mm'))
    r=calculate(b,s);assert r['quantity']=='1.809557' and r['plate_summary']['plate_count']=='3'
    assert r['plate_sections'][1]['normalized_parameters']=={'diameter':'1.2','opening_diameter':'0','thickness':'0.1','bottom_z':'1'}
    s['purpose']='wall_material';seal(s);b['source_component_hash']=store.digest(snapshot(s));assert calculate(b,s)['quantity']=='1.809557'


def test_round_only_total_after_inherited_counts():
    s=parent();s['chamber_wells'][0]['multiplicity']='7';seal(s);b=payload(s);b['plate_sections']=[plate(diameter='1',opening='0',thickness='.1')];b['rounding_digits']=0
    assert calculate(b,s)['quantity']=='1'


def test_seeded_plates_match_independent_cylinder_difference():
    rng=random.Random(3501)
    for _ in range(30):
        s=parent();count=rng.randint(1,8);s['chamber_wells'][0]['multiplicity']=str(count);seal(s);b=payload(s)
        hole=Decimal(rng.randint(0,50))/100;diameter=hole+Decimal(rng.randint(1,50))/10;t=Decimal(rng.randint(1,80))/1000
        b['plate_sections']=[plate(diameter=str(diameter),opening=str(hole),thickness=str(t),bottom=str(-rng.randint(1,99)))];b['rounding_digits']=12
        with localcontext() as ctx:
            ctx.prec=80
            expected=(Decimal(PI)/4*diameter*diameter*t-Decimal(PI)/4*hole*hole*t)*count;r=calculate(b,s)
            assert Decimal(r['quantity'])==expected.quantize(Decimal('.000000000001'))
            assert evaluate(r['formula_ast'],r['dimensions'])==Decimal(r['quantity'])


@pytest.mark.parametrize('case',['empty','many','duplicate','unknown','bool_index','extra_count','missing_height','missing_zero_proof','zero_diameter','zero_thickness','negative_hole','nan','whole_hole','over_hole','unit','precision','count_override','blank_scope','no_material','no_construction','unconfirmed','string_confirmation','rounding','mode','parameters','role','role_object','overflow'])
def test_invalid_inputs_fail(case):
    s=parent();b=payload(s);d=b['plate_sections'][0]
    if case=='empty':b['plate_sections']=[]
    elif case=='many':b['plate_sections']*=51
    elif case=='duplicate':b['plate_sections'].append(copy.deepcopy(d))
    elif case=='unknown':d['source_well_index']=99
    elif case=='bool_index':d['source_well_index']=True
    elif case=='extra_count':d['multiplicity']='2'
    elif case=='missing_height':del d['parameters']['bottom_z']
    elif case=='missing_zero_proof':d['parameters']['opening_diameter']='0';del d['evidence']['opening_diameter']
    elif case=='zero_diameter':d['parameters']['diameter']='0'
    elif case=='zero_thickness':d['parameters']['thickness']='0'
    elif case=='negative_hole':d['parameters']['opening_diameter']='-.1'
    elif case=='nan':d['parameters']['bottom_z']='NaN'
    elif case=='whole_hole':d['parameters']['opening_diameter']='2.4'
    elif case=='over_hole':d['parameters']['opening_diameter']='3'
    elif case=='unit':d['unit']='m²'
    elif case=='precision':d['unit']='mm';d['parameters']['thickness']='.000000000000000001'
    elif case=='count_override':b['multiplicity']='2'
    elif case=='blank_scope':d['measurement_basis']=''
    elif case=='no_material':b['material']=''
    elif case=='no_construction':b['construction_basis']=''
    elif case=='unconfirmed':b['scope_confirmed']=False
    elif case=='string_confirmation':b['scope_confirmed']='true'
    elif case=='rounding':b['rounding_digits']=True
    elif case=='mode':b['quantity_type']='area'
    elif case=='parameters':b['parameters']={'volume':'1'}
    elif case=='role':b['plate_role']='entire_well'
    elif case=='role_object':b['plate_role']={}
    elif case=='overflow':d['parameters'].update(diameter='1000000000000',thickness='1000000000000')
    with pytest.raises(DomainError):calculate(b,s)


def test_old_source_version_fails():
    s=parent();b=payload(s);s['revision']=2
    with pytest.raises(HTTPException) as err:calculate(b,s)
    assert err.value.status_code==409


def test_candidates_preview_review_books_boq_and_exports(clients):
    e,r,pid,source,trench,b,eid=setup(clients);base='/api/v1/projects/'+pid;before=get(e,pid)
    c=e.get(base+'/chamber-plate-candidates?component_id='+source['id']);assert c.status_code==200,c.text
    assert c.json()['chamber_wells'][0]['multiplicity']=='1' and get(e,pid)==before
    response=e.post(base+'/components/preview',json=b);assert response.status_code==200,response.text
    assert response.json()['quantity']=='0.848230' and get(e,pid)==before
    response=e.post(base+'/components',json={**b,'source_geometry':{'forged':True}});assert response.status_code==200,response.text
    comp=response.json();assert comp['source_geometry']==snapshot(source) and all(c['state']=='approved' for c in get(e,pid)['components'][:-1])
    assert e.post(base+'/takeoffs/calculate').status_code==409;approve(r,pid,'components',comp)
    response=e.post(base+'/takeoffs/calculate');assert response.status_code==200,response.text
    for book in response.json():approve(r,pid,'takeoffs',book)
    response=e.post(base+'/boq/generate');assert response.status_code==200,response.text
    assert response.json()[-1]['quantity']=='0.848230'
    j=e.get(base+'/exports?unpriced=true&format=json').json()['unpriced']['takeoffs'][-1]
    assert j['plate_sections']==comp['preview']['plate_sections'] and j['source_geometry']==snapshot(source)
    response=e.get(base+'/exports?unpriced=true&format=xlsx');assert response.status_code==200,response.text
    wb=load_workbook(io.BytesIO(response.content));assert {'井底盖板净实体汇总','井板逐项尺寸与孔扣减','井板标高范围及原件','井板材料做法与规则'}<=set(wb.sheetnames)
    assert all(c.data_type!='f' for s in wb for row in s for c in row)
    response=e.get(base+'/exports?unpriced=true&format=pdf');assert response.status_code==200,response.text
    text='\n'.join(p.extract_text() for p in PdfReader(io.BytesIO(response.content)).pages)
    assert '井底盖板净实体汇总' in text and '井板逐项尺寸与孔扣减' in text and '0.848230' in text


def test_net_plate_backfill_and_transitive_source_invalidation(clients):
    e,r,pid,source,trench,b,eid=setup(clients,save=True);base='/api/v1/projects/'+pid;plate_obj=get(e,pid)['components'][-1];approve(r,pid,'components',plate_obj)
    t=e.post(base+'/components',json=trench);assert t.status_code==200,t.text;approve(r,pid,'components',t.json());p=get(e,pid)
    src={'excavation':snapshot(p['components'][-1]),'occupations':[snapshot(c) for c in p['components'][1:3]]}
    balance=balance_payload(src,eid);balance['rounding_digits']=6
    for d in balance['occupations']:d['factor']='1'
    response=e.post(base+'/components',json=balance);assert response.status_code==200,response.text
    assert response.json()['preview']['quantity']=='37.226992';approve(r,pid,'components',response.json())
    books=e.post(base+'/takeoffs/calculate');assert books.status_code==200,books.text
    for book in books.json():approve(r,pid,'takeoffs',book)
    before=get(e,pid);assert e.put(base+'/components/'+plate_obj['id'],json=b).status_code==409 and get(e,pid)==before
    b['object_hash']=e.get(base+'/hashes').json()['components'][plate_obj['id']];bad=copy.deepcopy(b);bad['plate_sections'][0]['parameters']['opening_diameter']='3'
    assert e.put(base+'/components/'+plate_obj['id'],json=bad).status_code==422 and get(e,pid)==before
    b['plate_sections'][0]['parameters']['opening_diameter']='0';response=e.put(base+'/components/'+plate_obj['id'],json=b);assert response.status_code==200,response.text
    assert response.json()['revision']==2 and response.json()['preview']['quantity']=='0.904779'
    p=get(e,pid);assert [c['state'] for c in p['components']]==['approved','approved','draft','approved','stale']
    assert all(t['state']=='stale' for t in p['takeoffs']);approve(r,pid,'components',p['components'][2]);approve(r,pid,'components',p['components'][-1],expected=409)
    # An upstream well revision invalidates both its plate and the downstream balance transitively.
    source['object_hash']=e.get(base+'/hashes').json()['components'][source['id']];source['chamber_wells'][0]['multiplicity']='2'
    response=e.put(base+'/components/'+source['id'],json=source);assert response.status_code==200,response.text
    p=get(e,pid);assert [c['state'] for c in p['components']]==['approved','draft','stale','approved','stale']
    approve(r,pid,'components',p['components'][1]);assert e.post(base+'/components/preview',json=b).status_code==409
    assert e.post(base+'/takeoffs/calculate').status_code==409


@pytest.mark.parametrize('gate',['preview','save','component_review','calculate','book_review'])
def test_original_corruption_blocks_all_gates(clients,gate):
    e,r,pid,source,trench,b,eid=setup(clients);base='/api/v1/projects/'+pid
    if gate not in ('preview','save'):comp=e.post(base+'/components',json=b).json()
    if gate in ('calculate','book_review'):approve(r,pid,'components',comp)
    if gate=='book_review':book=e.post(base+'/takeoffs/calculate').json()[-1]
    doc=get(e,pid)['documents'][-1];(store.ROOT/'originals'/doc['hash']).write_bytes(b'corrupt')
    if gate=='preview':assert e.post(base+'/components/preview',json=b).status_code==409
    elif gate=='save':assert e.post(base+'/components',json=b).status_code==409
    elif gate=='component_review':approve(r,pid,'components',comp,expected=409)
    elif gate=='calculate':assert e.post(base+'/takeoffs/calculate').status_code==409
    else:approve(r,pid,'takeoffs',book,expected=409)


def test_permissions_foreign_proof_wrong_parent_and_corrupt_parent_original(clients):
    e,r,pid,source,trench,b,eid=setup(clients);base='/api/v1/projects/'+pid;url=base+'/chamber-plate-candidates?component_id='+source['id']
    assert r.get(url).status_code==200 and r.post(base+'/components',json=b).status_code==403 and clients['other'].get(url).status_code==404
    route=get(e,pid)['components'][0];assert e.get(base+'/chamber-plate-candidates?component_id='+route['id']).status_code==409
    other=create(e)['id'];upload(e,other,'其他项目孔径依据');bad=copy.deepcopy(b);bad['plate_sections'][0]['evidence']['opening_diameter']=get(e,other)['evidence'][0]['id']
    assert e.post(base+'/components',json=bad).status_code==404
    comp=e.post(base+'/components',json=b).json();changed={'kind':'ring','parameters':{},'object_hash':e.get(base+'/hashes').json()['components'][comp['id']]}
    assert e.put(base+'/components/'+comp['id'],json=changed).status_code==409
    doc=get(e,pid)['documents'][-2];(store.ROOT/'originals'/doc['hash']).write_bytes(b'corrupt parent')
    assert e.get(url).status_code==409 and e.post(base+'/components/preview',json=b).status_code==409


@pytest.mark.parametrize('target',['component','book'])
def test_saved_output_tampering_blocks_review(clients,target):
    e,r,pid,source,trench,b,eid=setup(clients,save=True);base='/api/v1/projects/'+pid
    if target=='book':approve(r,pid,'components',get(e,pid)['components'][-1]);e.post(base+'/takeoffs/calculate')
    with store.transaction() as db:
        p=json.loads(db.execute('SELECT data FROM projects WHERE id=?',(pid,)).fetchone()['data'])
        if target=='component':p['components'][-1]['preview']['plate_sections'][0]['single_opening_volume']='999'
        else:p['takeoffs'][-1]['plate_sections'][0]['single_opening_volume']='999'
        store.save(db,p)
    col='components' if target=='component' else 'takeoffs';approve(r,pid,col,get(e,pid)[col][-1],expected=409)


def test_new_plate_preserves_historical_quote(clients):
    e,r,pid,_,quote=priced_project(clients);old=copy.deepcopy(quote);base='/api/v1/projects/'+pid;eid=get(e,pid)['evidence'][0]['id']
    route=geometry(eid);route.pop('id');route.pop('revision');response=e.post(base+'/components',json=route);assert response.status_code==200,response.text
    approve(r,pid,'components',response.json());route=get(e,pid)['components'][-1]
    response=e.post(base+'/components',json=chamber_payload(route,eid));assert response.status_code==200,response.text;approve(r,pid,'components',response.json())
    response=e.post(base+'/components',json=payload(get(e,pid)['components'][-1],eid));assert response.status_code==200,response.text
    p=get(e,pid);assert p['quotes'][0]==old and old['id'] in p['stale_quotes']
