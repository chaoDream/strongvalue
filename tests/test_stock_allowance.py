"""Synthetic linear stock plans are annexes, not billable net quantities."""
import copy,io,random
from decimal import Decimal,localcontext
import pytest
from openpyxl import load_workbook
from pypdf import PdfReader
from backend import store
from backend.engine import DomainError
from backend.railing_material import calculate
from tests.test_railing_material import geometry,payload,setup as material_setup
from tests.test_railing_posts import body as equal_body
from tests.test_fixed_posts import body as fixed_body
from tests.test_additional_posts import body as fixed_added
from tests.test_equal_additional_posts import body as equal_added
from tests.test_post_adjustments import body as fixed_adjusted
from tests.test_equal_post_adjustments import body as equal_adjusted
from tests.test_combined_post_adjustments import body as combined
from tests.test_path_post_spacing import body as path_body
from tests.test_post_spacing import body as local_body
from tests.test_polygon_material import panel as polygon_panel
from tests.test_panel_material import panel as rect_panel
from tests.test_route_takeoff import evaluate
from tests.test_railing_takeoff import interval
from tests.test_polygon_takeoff import post,mutate
from tests.test_workflow import get,approve,upload,create,priced_project


def config(eid='test',count='4',lap='150',loss='3',unit='mm'):
 return dict(joint_count=count,lap_length=lap,length_unit=unit,loss_percent=loss,base_excludes_allowance=True,measurement_basis='合成每组接头总数独立核定，原净下料不含本次搭接及损耗，损耗基数为净长加搭接，备料仅附表不增加净工程量或组价',evidence={k:eid for k in ('joint_count','lap_length','loss_percent','measurement_basis')})


def body(s,eid='test'):
 b=payload(s,eid);b['name']='合成线材搭接损耗备料（非真实工程）';b['material_items'][0]['stock_allowance']=config(eid);b['material_items'][1]['stock_allowance']=config(eid,'2','50','2');return b


def setup(clients,pid=None):
 e,r,pid,s,_,_=material_setup(clients,pid)
 upload(e,pid,'合成软件备料验收，非真实工程。来源单组18m重复2。两道扶手净36m，单位2kg/m，每组接头总4处（不再乘道数），每处搭接150mm，原净下料不含该搭接和损耗。搭接0.6m，损耗3%按36.6m计1.098m，每组备料37.698m/75.396kg，全部75.396m/150.792kg。独立六件1200mm、3000g/m，单组净7.2m/21.6kg，每组接头总2处×50mm，损耗2%，备料7.446m/22.338kg，全部14.892m/44.676kg。两项合计净187.2kg，备料195.468kg，备料仅附表，不增加净清单或组价消耗；4处改5处后净量不变但仍须复核。','stock-allowance-basis.txt')
 for c in get(e,pid)['components']:approve(r,pid,'components',c)
 s=get(e,pid)['components'][-1];eid=get(e,pid)['evidence'][-1]['id'];return e,r,pid,s,body(s,eid),eid


def test_two_materials_exact_ledger_net_formula_and_no_mutation():
 s=geometry();b=body(s);before=copy.deepcopy((s,b));v=calculate(b,s);old=calculate(payload(s),s);assert (s,b)==before and v['quantity']==old['quantity']=='187.200' and v['formula_ast']==old['formula_ast'] and v['rule_version']=='railing-material-mass-1.12'
 a,c=[x['stock_planning'] for x in v['material_items']]
 assert a['single_group']==dict(net_length='36',lap_length='0.60',loss_length='1.098',procurement_length='37.698',net_mass='72',lap_mass='1.20',loss_mass='2.196',procurement_mass='75.396')
 assert Decimal(c['total']['procurement_mass'])==Decimal('44.676') and Decimal(v['stock_summary']['total']['procurement_mass'])==Decimal('195.468') and v['stock_summary']['item_count']==2
 assert len([k for k in v['source_evidence'] if '.stock_allowance.' in k])==8
 assert v['stock_summary']['unplanned_item_count']==0 and v['stock_summary']['planned_item_indices']==[1,2]
 for k,ast in v['stock_summary']['total_formula_asts'].items():assert evaluate(ast,v['dimensions'])==Decimal(v['stock_summary']['total'][k])
 for item in v['material_items']:
  p=item['stock_planning']
  with localcontext() as ctx:
   ctx.prec=80
   for k,ast in p['single_group_formula_asts'].items():assert evaluate(ast,v['dimensions'])==Decimal(p['single_group'][k])
   for k,ast in p['total_formula_asts'].items():assert evaluate(ast,v['dimensions'])==Decimal(p['total'][k])


@pytest.mark.parametrize('make',[equal_body,fixed_body,fixed_added,equal_added,fixed_adjusted,equal_adjusted,combined,lambda s:combined(s,equal=True),local_body,path_body])
def test_all_post_modes_and_spacing_checks_keep_layout_and_net_mass(make):
 s=geometry();b=make(s);old=calculate(copy.deepcopy(b),s);b['material_items'][0]['stock_allowance']=config();v=calculate(b,s);row=v['material_items'][0];assert v['quantity']==old['quantity'] and row['post_layout']==old['material_items'][0]['post_layout'] and row.get('post_spacing')==old['material_items'][0].get('post_spacing')
 p=row['stock_planning'];assert Decimal(p['single_group']['lap_length'])==Decimal('.6') and Decimal(p['total']['net_mass'])==Decimal(v['quantity'])


@pytest.mark.parametrize('unit,lap',[('m','.15'),('mm','150'),('cm','15'),('dm','1.5'),('km','.00015')])
def test_length_units_and_no_extra_count_multiplication(unit,lap):
 s=geometry();b=body(s);b['material_items'][0]['stock_allowance']=config(unit=unit,lap=lap);g=calculate(b,s)['material_items'][0]['stock_planning'];assert g['normalized_lap_length']=='0.15' and Decimal(g['total']['lap_length'])==Decimal('1.2')


@pytest.mark.parametrize('weight,unit',[('2','kg/m'),('2000','g/m'),('.002','kg/mm'),('2','g/mm')])
def test_weight_units_remain_independent(weight,unit):
 s=geometry();b=body(s);b['material_items'][0].update(unit_weight=weight,weight_unit=unit);assert Decimal(calculate(b,s)['material_items'][0]['stock_planning']['total']['procurement_mass'])==Decimal('150.792')


@pytest.mark.parametrize('count,lap,loss',[('0','0','0'),('0','0','100'),('4','150','0'),('1000000','1','100')])
def test_explicit_zero_and_maximums(count,lap,loss):
 s=geometry();b=body(s);b['material_items'][0]['stock_allowance']=config(count=count,lap=lap,loss=loss);p=calculate(b,s)['material_items'][0]['stock_planning'];assert Decimal(p['single_group']['procurement_length'])==(Decimal(36)+Decimal(count)*Decimal(lap)/1000)*(1+Decimal(loss)/100)


@pytest.mark.parametrize('fault',['null','list','missing','extra','no-proof','extra-proof','empty-proof','empty-basis','unconfirmed','truthy-confirmation','negative-count','fractional-count','large-count','negative-lap','zero-lap','unused-lap','negative-loss','large-loss','nan','infinity','unit','precision','bool'])
def test_strict_schema_and_parameter_bounds(fault):
 s=geometry();b=body(s);m=b['material_items'][0];c=m['stock_allowance']
 if fault=='null':m['stock_allowance']=None
 elif fault=='list':m['stock_allowance']=[]
 elif fault=='missing':c.pop('joint_count')
 elif fault=='extra':c['apply_to_boq']=True
 elif fault=='no-proof':c['evidence'].pop('lap_length')
 elif fault=='extra-proof':c['evidence']['scope']='test'
 elif fault=='empty-proof':c['evidence']['loss_percent']=''
 elif fault=='empty-basis':c['measurement_basis']=''
 elif fault=='unconfirmed':c['base_excludes_allowance']=False
 elif fault=='truthy-confirmation':c['base_excludes_allowance']=1
 elif fault=='negative-count':c['joint_count']='-1'
 elif fault=='fractional-count':c['joint_count']='1.1'
 elif fault=='large-count':c['joint_count']='1000001'
 elif fault=='negative-lap':c['lap_length']='-1'
 elif fault=='zero-lap':c['lap_length']='0'
 elif fault=='unused-lap':c['joint_count']='0'
 elif fault=='negative-loss':c['loss_percent']='-1'
 elif fault=='large-loss':c['loss_percent']='100.00001'
 elif fault=='nan':c['loss_percent']='NaN'
 elif fault=='infinity':c['lap_length']='Infinity'
 elif fault=='unit':c['length_unit']='m²'
 elif fault=='precision':c['lap_length']='.1234567890123456789'
 else:c['joint_count']=True
 with pytest.raises(DomainError):calculate(b,s)


@pytest.mark.parametrize('make',[rect_panel,polygon_panel])
def test_panel_modes_reject_linear_allowance(make):
 s=geometry();b=payload(s);b['material_items']=[make()];b['material_items'][0]['stock_allowance']=config()
 with pytest.raises(DomainError,match='字段'):calculate(b,s)


def test_partial_coverage_and_legacy_outputs_have_no_hidden_allowances():
 s=geometry();b=body(s);b['material_items'][1].pop('stock_allowance');v=calculate(b,s);assert v['stock_summary']['item_count']==1 and Decimal(v['stock_summary']['total']['net_mass'])==144 and v['quantity']=='187.200'
 b['material_items'][0].pop('stock_allowance');v=calculate(b,s);assert 'stock_summary' not in v and all('stock_planning' not in x for x in v['material_items']) and v['rule_version']=='railing-material-mass-1.0'


def test_seeded_independent_integer_planning_oracle_low_caller_precision():
 rng=random.Random(7901)
 for _ in range(50):
  s=geometry();b=payload(s);m=b['material_items'][1];m['quantity']=str(rng.randint(1,100));m['cut_length']=str(rng.randint(1,3000));m['unit_weight']=str(rng.randint(1,8000));count=rng.randint(1,20);lap=rng.randint(1,500);loss=rng.randint(0,100);m['stock_allowance']=config(count=str(count),lap=str(lap),loss=str(loss))
  with localcontext() as c:c.prec=8;v=calculate(b,s)
  with localcontext() as c:
   c.prec=80;expected=(Decimal(m['quantity'])*Decimal(m['cut_length'])/1000+Decimal(count)*lap/1000)*(1+Decimal(loss)/100)*(Decimal(m['unit_weight'])/1000)*2;assert Decimal(v['material_items'][1]['stock_planning']['total']['procurement_mass'])==expected


@pytest.mark.parametrize('fmt',['json','xlsx','pdf'])
def test_full_workflow_annex_exports_do_not_increase_boq_or_old_quote(clients,fmt):
 e,r,pid,_,quote=priced_project(clients);e,r,pid,s,b,eid=setup(clients,pid);before=get(e,pid);post(e,pid,b,'/components/preview');assert get(e,pid)==before;c=post(e,pid,b).json();approve(r,pid,'components',c)
 for t in post(e,pid,{},'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
 boq=post(e,pid,{},'/boq/generate').json();assert any(v['quantity']=='187.200' and v['unit']=='kg' for v in boq) and all(v['quantity']!='195.468' for v in boq)
 url=f'/api/v1/projects/{pid}/exports';res=e.get(url,params=dict(format=fmt,unpriced=True));assert res.status_code==200,res.text;titles={'线材备料汇总','线材搭接损耗备料','线材备料参数依据','线材备料规则与范围'}
 if fmt=='json':assert titles<={x[0] for x in res.json()['tables']}
 elif fmt=='xlsx':assert titles<=set(load_workbook(io.BytesIO(res.content)).sheetnames)
 else:assert all(t in ''.join(p.extract_text() for p in PdfReader(io.BytesIO(res.content)).pages) for t in titles)
 old=e.get(url,params=dict(format='json',quote_id=quote['id'])).json();assert old['quote']==quote and not any(t[0] in titles for t in old['tables'])


@pytest.mark.parametrize('fault',['input','proof','length','mass','ast','summary','book'])
def test_saved_parameters_ledger_ast_and_summary_are_rebuilt(clients,fault):
 e,r,pid,s,b,eid=setup(clients);c=post(e,pid,b).json()
 if fault=='book':approve(r,pid,'components',c);post(e,pid,{},'/takeoffs/calculate')
 def forge(p):
  c=p['components'][-1];v=c['preview'];g=v['material_items'][0]['stock_planning']
  if fault=='input':c['material_items'][0]['stock_allowance']['joint_count']='5'
  elif fault=='proof':c['material_items'][0]['stock_allowance']['evidence']['lap_length']='forged'
  elif fault=='length':g['single_group']['procurement_length']='36'
  elif fault=='mass':g['total']['procurement_mass']='144'
  elif fault=='ast':g['total_formula_asts']['lap_mass']={'constant':'0'}
  elif fault=='summary':v['stock_summary']['total']['procurement_mass']='187.2'
  else:p['takeoffs'][-1]['material_items'][0]['stock_planning']['total']['procurement_mass']='0'
 mutate(pid,forge);coll='takeoffs' if fault=='book' else 'components';approve(r,pid,coll,get(e,pid)[coll][-1],expected=404 if fault=='proof' else 409)


def test_same_net_revision_removal_history_concurrency_and_source_staleness(clients):
 e,r,pid,s,b,eid=setup(clients);c=post(e,pid,b).json();approve(r,pid,'components',c)
 for t in post(e,pid,{},'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
 post(e,pid,{},'/boq/generate');before=get(e,pid);old=before['components'][-1];new=copy.deepcopy(b);new['material_items'][0]['stock_allowance']['joint_count']='5';new['object_hash']=store.digest(old);url=f'/api/v1/projects/{pid}/components/'+old['id'];res=e.put(url,json=new);assert res.status_code==200,res.text;p=get(e,pid);assert p['component_history'][-1]==old and p['boq']==before['boq'] and all(t['state']=='stale' for t in p['takeoffs']) and res.json()['preview']['quantity']=='187.200';assert e.put(url,json=new).status_code==409;approve(r,pid,'components',res.json())
 for t in post(e,pid,{},'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
 assert any(v['quantity']=='187.200' for v in post(e,pid,{},'/boq/generate').json())
 current=get(e,pid)['components'][-1];new['object_hash']=store.digest(current)
 for m in new['material_items']:m.pop('stock_allowance')
 res=e.put(url,json=new);assert res.status_code==200 and 'stock_summary' not in res.json()['preview'];assert get(e,pid)['component_history'][-1]==current;approve(r,pid,'components',res.json())
 s['object_hash']=store.digest(s);s['edge_decisions'][0].update(mode='ranges',ranges=[interval(0,10,eid)]);res=e.put(f'/api/v1/projects/{pid}/components/'+s['id'],json=s);assert res.status_code==200;approve(r,pid,'components',res.json());assert get(e,pid)['components'][-1]['state']=='stale'


@pytest.mark.parametrize('gate',['preview','save','review','calculate','book_review'])
def test_allowance_original_byte_gate(clients,gate):
 e,r,pid,s,b,eid=setup(clients)
 if gate in ('review','calculate','book_review'):c=post(e,pid,b).json()
 if gate in ('calculate','book_review'):approve(r,pid,'components',c)
 if gate=='book_review':t=post(e,pid,{},'/takeoffs/calculate').json()[-1]
 d=get(e,pid)['documents'][-1];(store.ROOT/'originals'/d['hash']).write_bytes(b'corrupt')
 if gate=='preview':post(e,pid,b,'/components/preview',409)
 elif gate=='save':post(e,pid,b,expected=409)
 elif gate=='review':approve(r,pid,'components',c,expected=409)
 elif gate=='calculate':post(e,pid,{},'/takeoffs/calculate',409)
 else:approve(r,pid,'takeoffs',t,expected=409)


def test_allowance_role_tenant_and_foreign_evidence(clients):
 e,r,pid,s,b,eid=setup(clients);post(r,pid,b,expected=403);assert clients['other'].get(f'/api/v1/projects/{pid}').status_code==404
 other=create(e)['id'];upload(e,other,'异项目');b['material_items'][0]['stock_allowance']['evidence']['loss_percent']=get(e,other)['evidence'][-1]['id'];post(e,pid,b,expected=404)


def test_long_names_pdf_exports_without_row_overflow(clients):
 e,r,pid,s,b,eid=setup(clients);b['name']='合成'+('长名称验收'*39)+'末';b['material_items'][0]['name']='合成'+('长材料名'*39)+'末';c=post(e,pid,b).json();approve(r,pid,'components',c)
 for t in post(e,pid,{},'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
 res=e.get(f'/api/v1/projects/{pid}/exports',params=dict(format='pdf',unpriced=True));assert res.status_code==200,res.text
