import copy
import io
import json
import zipfile
from datetime import date,timedelta
import pytest
from openpyxl import load_workbook
from backend import store
from tests.test_workflow import priced_project,get,approve


def source(clients,converted=False):
    e,r,pid,comp,q=priced_project(clients);base=f'/api/v1/projects/{pid}';b=get(e,pid)['boq'][0]
    quotas=copy.deepcopy(b['quotas']);quotas[0]['resources'][0].update(resource_code='SYNTHETIC-MAT',specification='仅软件合成测试规格',fee_allocation='材料费')
    if converted:
        quotas[0]['conversion']={'quota_unit':'10m³','unit_factor':'0.1','category_factors':{'人工':'1','材料':'1.2','机械':'1','设备':'1'},
                                 'source':'合成换算依据','license':'仅软件测试','reason':'每10m³原始消耗，分类系数仅测试','evidence_id':get(e,pid)['evidence'][0]['id']}
        quotas[0]['coefficient']='1.5'
        second=copy.deepcopy(quotas[0]);second['code']='SYNTHETIC-OTHER';second['resources'][0]['resource_code']='SYNTHETIC-SECOND';quotas.append(second)
    response=e.put(base+'/boq/'+b['id'],json={'quotas':quotas});assert response.status_code==200;approve(r,pid,'boq',response.json())
    b=get(e,pid)['boq'][0];ev=get(e,pid)['evidence'][0]['id']
    body={'boq_id':b['id'],'object_hash':store.digest(b),'discipline':'建筑','quota_edition':'math-v1','applicability':'合成底板，通用系数1，非专业定额',
          'reuse_basis':'软件合成测试本项目可复用','reuse_evidence_id':ev,'valid_from':(date.today()-timedelta(days=1)).isoformat(),'valid_until':(date.today()+timedelta(days=30)).isoformat()}
    response=e.post(base+'/catalog',json=body);assert response.status_code==200,response.text
    candidate=response.json();approve(r,pid,'catalog',candidate);candidate=get(e,pid)['catalog'][0]
    publish={'catalog_id':candidate['id'],'object_hash':store.digest(candidate),'reuse_basis':'仅本企业合成测试项目可复用组成并共享相关原件；非真实商业许可',
             'reuse_evidence_id':ev,'confirm_share_originals':True}
    response=e.post(base+'/quota-library/publish',json=publish);assert response.status_code==200,response.text
    return e,r,pid,response.json(),publish,q


def enterprise_review(r,item,decision='approve',expected=200):
    response=r.post('/api/v1/organization/quota-library/'+item['entry']['id']+'/reviews',json={'object_hash':item['object_hash'],'decision':decision,'reason':'仅合成软件测试，核对跨项目许可及原件，无真实签审效力'})
    assert response.status_code==expected,response.text
    return response.json()


def request_body(e,pid,item):
    b=get(e,pid)['boq'][0]
    return b,{'boq_hash':store.digest(b),'object_hash':item['object_hash'],'discipline':'建筑','quota_edition':'math-v1'}


def endpoint(pid,b,item,action):return f"/api/v1/projects/{pid}/boq/{b['id']}/enterprise-quota-candidates/{item['entry']['id']}/{action}"


def adopt(e,pid,item):
    b,body=request_body(e,pid,item);response=e.post(endpoint(pid,b,item,'preview'),json=body);assert response.status_code==200,response.text
    body.update(preview_hash=response.json()['preview_hash'],reason='核查合成构造与复用许可，所有资源在目标重新询价',confirm_conditions=True,request_id=store.uid())
    response=e.post(endpoint(pid,b,item,'adopt'),json=body);assert response.status_code==200,response.text
    return response.json(),body


def fill_prices(e,r,pid):
    p=get(e,pid);b=p['boq'][0];quotas=copy.deepcopy(b['quotas'])
    local=next(ev for ev in p['evidence'] if not ev.get('quota_library_origin'))
    for q in quotas:
        for resource in q['resources']:
            resource.update(pricing_status='priced',price='120.005',source_type='供应商报价',source='本项目合成报价单',license='仅软件测试',period='2026-10',region='河南省',tax_mode='不含税',
                            valid_from=(date.today()-timedelta(days=1)).isoformat(),valid_until=(date.today()+timedelta(days=30)).isoformat(),evidence_id=local['id'])
    response=e.put(f'/api/v1/projects/{pid}/boq/'+b['id'],json={'object_hash':store.digest(b),'quotas':quotas})
    assert response.status_code==200,response.text
    assert response.json().get('enterprise_quota_ref')
    approve(r,pid,'boq',response.json())
    return e.post(f'/api/v1/projects/{pid}/quotes/calculate')


def test_publish_review_search_adopt_inquire_and_report_provenance(clients):
    e,r,sid,item,publish,old_source=source(clients)
    assert not e.get('/api/v1/organization/quota-library').json()['items'][0]['available']
    item=enterprise_review(r,item)
    e,r,pid,comp,old=priced_project(clients);base=f'/api/v1/projects/{pid}';b,body=request_body(e,pid,item)
    result=e.get(base+'/boq/'+b['id']+'/enterprise-quota-candidates?discipline=建筑&quota_edition=math-v1').json()
    assert len(result)==1 and '全部资源待询价' in result[0]['match_reason']
    preview=e.post(endpoint(pid,b,item,'preview'),json=body).json();assert preview['quotas'][0]['resources'][0]['price'] is None
    before=get(e,pid);record,body=adopt(e,pid,item);p=get(e,pid);after=p['boq'][0]
    assert after['state']=='draft' and p['status']=='BOQ_DRAFT' and p['quotes']==before['quotes']
    for field in ['code','name','unit','quantity','features','node','takeoff_id']:assert after[field]==b[field]
    resource=after['quotas'][0]['resources'][0]
    assert resource['consumption']=='1.05' and resource['price'] is None and resource['pricing_status']=='unpriced'
    assert 'source' not in resource and 'period' not in resource
    assert len(p['pricing_tasks'])==1 and e.post(base+'/quotes/calculate').status_code==409
    assert len(p['documents'])==len(before['documents'])+1
    assert resource['evidence_id'] in record['proof_map'].values()
    unpriced=e.get(base+'/exports?format=zip&unpriced=true');assert unpriced.status_code==200,unpriced.text[:200]
    with zipfile.ZipFile(io.BytesIO(unpriced.content)) as archive:
        snapshot=json.loads(archive.read('unpriced-snapshot.json'))
        assert snapshot['enterprise_quota_adoptions']==[record] and 'quota_adoptions' in snapshot and 'catalog' in snapshot
        assert any(ev.get('quota_library_origin') for ev in snapshot['evidence'])
    assert e.post(endpoint(pid,b,item,'adopt'),json=body).json()==record
    assert e.post(endpoint(pid,b,item,'adopt'),json={**body,'reason':'another'}).status_code==409
    q=fill_prices(e,r,pid);assert q.status_code==200,q.text
    assert q.json()['total']=='1946.85' and q.json()['snapshot']['enterprise_quota_adoptions'][0]['entry']==item['entry']
    j=e.get(base+'/exports?format=json').json();assert '企业定额复用台账' in [s[0] for s in j['tables']]
    assert '企业定额复用台账' in load_workbook(io.BytesIO(e.get(base+'/exports?format=xlsx').content)).sheetnames
    assert e.get(base+'/exports?format=pdf').content.startswith(b'%PDF')
    bundle=e.get(base+'/exports?format=zip');assert bundle.status_code==200,bundle.text[:200]
    assert get(e,pid)['quotes'][0]==old
    approve(r,pid,'quotes',q.json());saved=copy.deepcopy(get(e,pid)['quotes'])
    enterprise_review(r,item,'reject')
    assert any('企业定额候选不可用' in x for x in get(e,pid)['blockers'])
    assert e.post(base+'/quotes/calculate').status_code==409 and get(e,pid)['quotes']==saved
    assert '未完成成果草稿' in e.get(base+'/exports?format=json').json()['status']


@pytest.mark.parametrize('change',[{'confirm_share_originals':False},{'reuse_basis':''},{'reuse_evidence_id':'foreign'},{'object_hash':'stale'}])
def test_publication_requires_permission_and_current_source(clients,change):
    e,r,sid,item,body,old=source(clients);before=e.get('/api/v1/organization/quota-library').json()
    response=e.post(f'/api/v1/projects/{sid}/quota-library/publish',json={**body,**change})
    assert response.status_code in (404,409,422),response.text
    assert e.get('/api/v1/organization/quota-library').json()==before


def test_roles_tenant_and_project_access_boundaries(clients):
    e,r,sid,item,publish,old=source(clients);eid=item['entry']['id']
    assert r.post(f'/api/v1/projects/{sid}/quota-library/publish',json=publish).status_code==403
    assert clients['other'].get('/api/v1/organization/quota-library').json()['total']==0
    assert clients['other'].post('/api/v1/organization/quota-library/'+eid+'/reviews',json={'decision':'approve','reason':'test','object_hash':item['object_hash']}).status_code==403
    enterprise_review(e,item,expected=403);item=enterprise_review(r,item)
    e,r,pid,comp,q=priced_project(clients);b,body=request_body(e,pid,item)
    assert clients['other'].post(endpoint(pid,b,item,'preview'),json=body).status_code==404
    assert r.post(endpoint(pid,b,item,'adopt'),json={**body,'reason':'test','confirm_conditions':True,'request_id':'test'}).status_code==403


@pytest.mark.parametrize('change',[{'discipline':'电气'},{'quota_edition':'other'},{'boq_hash':'stale'},{'object_hash':'stale'}])
def test_target_filter_and_hash_guards(clients,change):
    e,r,sid,item,publish,old=source(clients);item=enterprise_review(r,item)
    e,r,pid,comp,q=priced_project(clients);b,body=request_body(e,pid,item)
    before=get(e,pid);response=e.post(endpoint(pid,b,item,'preview'),json={**body,**change})
    assert response.status_code==409,response.text
    assert get(e,pid)==before
    assert e.get(f'/api/v1/projects/{pid}/boq/{b["id"]}/enterprise-quota-candidates?discipline=建筑&quota_edition=bad').json()==[]


def test_source_restriction_hides_entries_and_blocks_existing_references(clients):
    e,r,sid,item,publish,old=source(clients);item=enterprise_review(r,item)
    e,r,pid,comp,q=priced_project(clients);record,body=adopt(e,pid,item)
    admin=clients['admin'];base=f'/api/v1/projects/{sid}';access=admin.get(base+'/access').json()
    response=admin.put(base+'/access',json={'mode':'restricted','member_ids':[],'reason':'测试停止共享','object_hash':access['object_hash']});assert response.status_code==200
    assert e.get('/api/v1/organization/quota-library').json()['total']==0
    assert any('原项目已限制共享' in x for x in get(e,pid)['blockers'])
    # Prior copies remain in the already-authorized target's history; no new access is granted.
    assert get(e,pid)['enterprise_quota_adoptions'][0]==record
    b=get(e,pid)['boq'][0]
    assert e.post(endpoint(pid,b,item,'preview'),json={**body,'boq_hash':store.digest(b)}).status_code==409


def test_source_edit_or_candidate_withdrawal_blocks_adoption(clients):
    e,r,sid,item,publish,old=source(clients);item=enterprise_review(r,item)
    e,r,pid,comp,q=priced_project(clients);b,body=request_body(e,pid,item)
    preview=e.post(endpoint(pid,b,item,'preview'),json=body).json();before=get(e,pid)
    source_b=get(e,sid)['boq'][0];e.put(f'/api/v1/projects/{sid}/boq/'+source_b['id'],json={'features':'原工程条件修改'})
    body.update(preview_hash=preview['preview_hash'],request_id=store.uid(),reason='test',confirm_conditions=True)
    assert e.post(endpoint(pid,b,item,'adopt'),json=body).status_code==409
    assert get(e,pid)==before
    assert not e.get('/api/v1/organization/quota-library').json()['items'][0]['available']


@pytest.mark.parametrize('pricing_offset',[0,1])
def test_stale_preview_date_and_confirmation_are_atomic(clients,pricing_offset):
    e,r,sid,item,publish,old=source(clients);item=enterprise_review(r,item)
    e,r,pid,comp,q=priced_project(clients);b,body=request_body(e,pid,item)
    # Pin the project date: host UTC tomorrow may equal Shanghai today.
    initial_day=date.today()+timedelta(days=pricing_offset)
    response=e.patch(f'/api/v1/projects/{pid}/settings',json={'pricing_date':initial_day.isoformat(),'reason':'固定测试起始计价日期'})
    assert response.status_code==200,response.text
    preview=e.post(endpoint(pid,b,item,'preview'),json=body).json()
    body.update(preview_hash=preview['preview_hash'],request_id=store.uid(),reason='test',confirm_conditions=False)
    assert e.post(endpoint(pid,b,item,'adopt'),json=body).status_code==422
    e.patch(f'/api/v1/projects/{pid}/settings',json={'pricing_date':(initial_day+timedelta(days=1)).isoformat(),'reason':'测试日期变化'})
    before=get(e,pid);body['confirm_conditions']=True
    assert e.post(endpoint(pid,b,item,'adopt'),json=body).status_code==409
    assert get(e,pid)==before


def test_tamper_and_missing_original_block_review_preview_and_approval(clients):
    e,r,sid,item,publish,old=source(clients);doc=get(e,sid)['documents'][0];path=store.ROOT/'originals'/doc['hash'];raw=path.read_bytes()
    path.write_bytes(b'bad');enterprise_review(r,item,expected=409)
    path.write_bytes(raw);item=enterprise_review(r,item)
    e,r,pid,comp,q=priced_project(clients);b,body=request_body(e,pid,item);record,request=adopt(e,pid,item)
    quote=fill_prices(e,r,pid);assert quote.status_code==200
    path.write_bytes(b'bad');approve(r,pid,'quotes',quote.json(),expected=409)
    assert e.post(f'/api/v1/projects/{pid}/quotes/calculate').status_code==409
    path.unlink();assert e.post(endpoint(pid,b,item,'preview'),json={**body,'boq_hash':store.digest(get(e,pid)['boq'][0])}).status_code==409


def test_composition_reuse_does_not_circulate_and_manual_detach_is_audited(clients):
    e,r,sid,item,publish,old=source(clients);item=enterprise_review(r,item)
    e,r,pid,comp,q=priced_project(clients);record,request=adopt(e,pid,item);assert fill_prices(e,r,pid).status_code==200
    b=get(e,pid)['boq'][0];ev=get(e,pid)['evidence'][0]['id']
    body={'boq_id':b['id'],'object_hash':store.digest(b),'discipline':'建筑','quota_edition':'math-v1','applicability':'test','reuse_basis':'test','reuse_evidence_id':ev,
          'valid_from':date.today().isoformat(),'valid_until':(date.today()+timedelta(days=1)).isoformat()}
    assert e.post(f'/api/v1/projects/{pid}/catalog',json=body).status_code==409
    response=e.put(f'/api/v1/projects/{pid}/boq/'+b['id'],json={'object_hash':store.digest(b),'detach_catalog':True,'reason':'本项目独立核查和维护定额来源'})
    assert response.status_code==200 and 'enterprise_quota_ref' not in response.json()
    assert get(e,pid)['enterprise_quota_adoptions'][0]==record
    assert any(x['action']=='ENTERPRISE_QUOTA_DETACHED' for x in e.get(f'/api/v1/projects/{pid}/audit').json())


def test_source_missing_resource_identity_cannot_be_published(clients):
    e,r,sid,item,publish,old=source(clients)
    with store.transaction() as c:
        p=json.loads(c.execute('SELECT data FROM projects WHERE id=?',(sid,)).fetchone()['data']);b=p['boq'][0]
        del b['quotas'][0]['resources'][0]['resource_code']
        candidate=p['catalog'][0];candidate['quotas']=copy.deepcopy(b['quotas']);candidate['source_hash']=store.digest(b);store.save(c,p)
    response=e.post(f'/api/v1/projects/{sid}/quota-library/publish',json={**publish,'object_hash':store.digest(candidate)})
    assert response.status_code==422,response.text


def test_multiple_subitems_conversions_and_original_dedup_survive_copy(clients):
    e,r,sid,item,publish,old=source(clients,True);item=enterprise_review(r,item)
    e,r,pid,comp,q=priced_project(clients);before=get(e,pid);record,body=adopt(e,pid,item);p=get(e,pid)
    assert len(p['boq'][0]['quotas'])==2 and len(p['pricing_tasks'])==2
    assert len(p['documents'])==len(before['documents'])+1
    for quota in p['boq'][0]['quotas']:
        assert quota['conversion']['unit_factor']=='0.1' and quota['conversion']['category_factors']['材料']=='1.2'
        assert quota['coefficient']=='1.5' and quota['conversion']['evidence_id'] in record['proof_map'].values()
        assert quota['resources'][0]['consumption']=='1.05' and quota['resources'][0]['price'] is None
    result=fill_prices(e,r,pid);assert result.status_code==200,result.text
    assert result.json()['engine']=='pricing-2.1'


@pytest.mark.parametrize('field,value',[('region','河北省'),('scope',['电气']),('pricing_date','2040-01-01'),('rule_version','v2'),('unit','m²')])
def test_cross_project_applicability_never_uses_approximate_fallback(clients,field,value):
    e,r,sid,item,publish,old=source(clients);item=enterprise_review(r,item)
    e,r,pid,comp,q=priced_project(clients)
    with store.transaction() as c:
        p=json.loads(c.execute('SELECT data FROM projects WHERE id=?',(pid,)).fetchone()['data'])
        if field=='unit':p['boq'][0]['unit']=value
        elif field=='rule_version':p['rules'][0]['version']=value
        else:p[field]=value
        store.save(c,p)
    b,body=request_body(e,pid,item);before=get(e,pid)
    assert e.post(endpoint(pid,b,item,'preview'),json=body).status_code==409
    assert e.get(f'/api/v1/projects/{pid}/boq/{b["id"]}/enterprise-quota-candidates?discipline=建筑&quota_edition=math-v1').json()==[]
    assert get(e,pid)==before


def test_restricted_source_cannot_publish_even_for_allowed_member(clients):
    e,r,sid,item,publish,old=source(clients);admin=clients['admin'];base=f'/api/v1/projects/{sid}'
    access=admin.get(base+'/access').json();members=admin.get('/api/v1/organization/members').json()
    ids=[m['id'] for m in members if m['username'] in ('estimator','reviewer')]
    response=admin.put(base+'/access',json={'mode':'restricted','member_ids':ids,'reason':'测试受限来源','object_hash':access['object_hash']});assert response.status_code==200
    assert e.post(base+'/quota-library/publish',json=publish).status_code==409
    enterprise_review(r,item,expected=409)


def test_reapproving_source_never_reactivates_old_enterprise_permission(clients):
    e,r,sid,item,publish,old=source(clients);item=enterprise_review(r,item)
    candidate=get(e,sid)['catalog'][0]
    approve(r,sid,'catalog',candidate,'reject');approve(r,sid,'catalog',candidate)
    assert not e.get('/api/v1/organization/quota-library').json()['items'][0]['available']
    enterprise_review(r,item,expected=409)
    response=e.post(f'/api/v1/projects/{sid}/quota-library/publish',json=publish)
    assert response.status_code==200,response.text
    renewed=enterprise_review(r,response.json())
    assert renewed['entry']['source_approval_ids']!=item['entry']['source_approval_ids']
