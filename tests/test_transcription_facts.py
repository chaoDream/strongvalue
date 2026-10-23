import copy
import io
import pytest
from openpyxl import load_workbook
from pypdf import PdfReader
from backend import store,ocr
from tests.test_ocr import setup,submit,pdf
from tests.test_ocr_line_locations import selected
from tests.test_ocr_drawing_candidates import result,mutate
from tests.test_ocr_transcriptions import save,body_for
from tests.test_workflow import get,approve,priced_project

TEXT='处理规模：0.05万m³/d；施工工期45日历天；监理费率1.49%；施工合同金额298.24858万元；合同金额：300万元；建设单位：合成业主'


def fixture(clients,monkeypatch,text=TEXT):
    c,pid,d,url=setup(clients,monkeypatch)
    monkeypatch.setattr(ocr,'recognize',lambda *a:result(text))
    submit(c,url,d);e,_=selected(c,pid)
    record=save(c,pid,e,body_for(e,text=text))['evidence']
    return c,clients['reviewer'],pid,d,e,record


def read(c,pid,record,query='',expected=200):
    response=c.get(f'/api/v1/projects/{pid}/evidence/{record["id"]}/transcription-facts'+query)
    assert response.status_code==expected,response.text
    return response.json()


def adopt(c,pid,record,row,expected=200,**extra):
    response=c.post(f'/api/v1/projects/{pid}/evidence/{record["id"]}/transcription-facts/{row["id"]}/adopt',json={
        'object_hash':row['object_hash'],'reason':'合成测试：核对抄录字段、原始数值和单位，待独立复核',**extra})
    assert response.status_code==expected,response.text
    return response.json()


def test_candidates_are_read_only_and_preserve_original_units_and_distinct_fields(clients,monkeypatch):
    c,r,pid,d,e,record=fixture(clients,monkeypatch)
    before=get(c,pid);audit=c.get(f'/api/v1/projects/{pid}/audit').json()
    data=read(c,pid,record);rows={v['key']:v for v in data['candidates']}
    assert rows['capacity']['value']=='500.00' and rows['capacity']['original_value']=='0.05'
    assert rows['capacity']['original_unit']=='万m³/日' and rows['capacity']['unit']=='m³/日'
    assert rows['duration']['value']=='45' and rows['supervision_rate']['value']=='1.49'
    assert rows['construction_amount']['value']=='2982485.80000' and rows['contract_amount']['value']=='3000000'
    assert rows['client']['value']=='合成业主'
    assert all(v['raw_text'] in record['text'] for v in rows.values())
    assert data==read(r,pid,record) and get(c,pid)==before and c.get(f'/api/v1/projects/{pid}/audit').json()==audit


def test_manual_adoption_requires_independent_review_and_prevents_duplicates(clients,monkeypatch):
    c,r,pid,d,e,record=fixture(clients,monkeypatch)
    row=next(v for v in read(c,pid,record)['candidates'] if v['key']=='capacity')
    f=adopt(c,pid,record,row);assert f['state']=='draft' and f['method']=='manual'
    assert f['evidence_id']==record['id'] and f['original_value']=='0.05'
    assert f['transcription_adoption']['candidate_hash']==row['object_hash']
    assert read(c,pid,record)['candidates'][0]['active_fact_id']==f['id']
    adopt(c,pid,record,row,409)
    approve(c,pid,'facts',f,expected=403);approve(r,pid,'facts',f)
    assert not get(c,pid)['components'] and not get(c,pid)['takeoffs']
    assert [a['action'] for a in c.get(f'/api/v1/projects/{pid}/audit').json()].count('TRANSCRIPTION_FACT_ADOPTED')==1
    approve(r,pid,'facts',f,'reject');again=adopt(c,pid,record,row)
    assert again['id']!=f['id'] and again['state']=='draft'


@pytest.mark.parametrize('text,conflicts',[
    ('合同金额：1万元；合同金额：10000元',False),
    ('合同金额：1万元；合同金额：2万元',True),
    ('施工工期45日历天；施工工期60日历天',True)])
def test_different_values_preserved_without_majority_or_false_unit_conflict(clients,monkeypatch,text,conflicts):
    c,r,pid,d,e,record=fixture(clients,monkeypatch,text)
    rows=read(c,pid,record)['candidates'];assert len(rows)==2 and rows[0]['id']!=rows[1]['id']
    for row in rows:adopt(c,pid,record,row)
    p=get(c,pid);assert bool(p['fact_conflicts']) is conflicts
    if conflicts:assert all(f['state']=='conflict' for f in p['facts'])


@pytest.mark.parametrize('extra',[{'object_hash':'0'*64},{'value':'999'},{'reason':'   '},{'reason':42}])
def test_stale_hash_or_modified_payload_rejected_atomically(clients,monkeypatch,extra):
    c,r,pid,d,e,record=fixture(clients,monkeypatch);row=read(c,pid,record)['candidates'][0];before=get(c,pid)
    adopt(c,pid,record,row,409 if extra.get('object_hash') else 422,**extra)
    assert get(c,pid)==before


@pytest.mark.parametrize('change',['task','source','transcription','bytes','missing'])
def test_invalidated_source_blocks_read_and_adoption(clients,monkeypatch,change):
    c,r,pid,d,e,record=fixture(clients,monkeypatch);row=read(c,pid,record)['candidates'][0]
    def action(p,ev,j):
        if change=='task':j['state']='failed'
        elif change=='source':ev['ocr']['lines'][0]['confidence']=.1
        elif change=='transcription':next(v for v in p['evidence'] if v['id']==record['id'])['text']='合同金额：2万元'
    if change=='bytes':(store.ROOT/'originals'/d['hash']).write_bytes(b'changed')
    elif change=='missing':(store.ROOT/'originals'/d['hash']).unlink()
    else:mutate(pid,action)
    read(c,pid,record,expected=409);adopt(c,pid,record,row,409)
    assert not get(c,pid)['facts']


def test_permissions_source_kind_and_pagination(clients,monkeypatch):
    c,r,pid,d,e,record=fixture(clients,monkeypatch,'；'.join('合同金额：'+str(i)+'元' for i in range(1,56)))
    data=read(c,pid,record,'?limit=10');assert data['total']==55 and data['next_offset']==10
    assert read(c,pid,record,'?offset=50&limit=10')['next_offset'] is None
    row=data['candidates'][0]
    adopt(r,pid,record,row,403);adopt(clients['other'],pid,record,row,404)
    read(clients['other'],pid,record,expected=404);read(c,pid,e,expected=422)
    for query in ['?limit=0','?limit=201','?offset=-1']:read(c,pid,record,query,422)


def test_no_fact_inferred_from_drawing_length_or_supervision_duration(clients,monkeypatch):
    c,r,pid,d,e,record=fixture(clients,monkeypatch,'长度：1200 mm；监理服务45日历天')
    assert read(c,pid,record)['total']==0 and not get(c,pid)['facts']


def test_adoption_invalidates_current_quote_and_exports_frozen_source_trace(clients,monkeypatch):
    c,r,pid,_,old=priced_project(clients);base=f'/api/v1/projects/{pid}'
    monkeypatch.setattr(ocr,'capability',lambda:{'available':True});monkeypatch.setattr(ocr,'recognize',lambda *a:result('施工工期45日历天'))
    d=c.post(base+'/documents',data={'specialty':'建筑'},files={'file':('fact-scan.pdf',pdf(),'application/pdf')}).json()['document']
    submit(c,base+'/documents/'+d['id']+'/ocr',d);e,_=selected(c,pid)
    record=save(c,pid,e,body_for(e,text='施工工期45日历天'))['evidence'];row=read(c,pid,record)['candidates'][0]
    f=adopt(c,pid,record,row);assert old['id'] in get(c,pid)['stale_quotes']
    approve(r,pid,'facts',f)
    for coll in ['components','takeoffs','boq']:
        for item in get(c,pid)[coll]:approve(r,pid,coll,item)
    response=c.post(base+'/quotes/calculate');assert response.status_code==200,response.text
    q=response.json();frozen=copy.deepcopy(q)
    save(c,pid,e,body_for(e,text='施工工期60日历天',reason='后续抄录（合成）'))
    data=c.get(base+'/exports',params={'format':'json','quote_id':q['id']}).json()
    assert data['quote']==frozen and any(t[0]=='人工抄录事实采用' for t in data['tables'])
    assert '后续抄录' not in str(data['tables'])
    raw=c.get(base+'/exports',params={'format':'xlsx','quote_id':q['id']}).content
    assert '人工抄录事实采用' in load_workbook(io.BytesIO(raw)).sheetnames
    raw=c.get(base+'/exports',params={'format':'pdf','quote_id':q['id']}).content
    assert '人工抄录事实采用' in ''.join(p.extract_text() for p in PdfReader(io.BytesIO(raw)).pages)
    assert not any(f.get('transcription_adoption') for f in old['snapshot']['facts'])
