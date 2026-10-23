"""Synthetic per-stock processing, exact geometry and workflow evidence."""
import copy,io,random
from decimal import Decimal,localcontext
from fractions import Fraction
from collections import Counter
import pytest
from openpyxl import load_workbook
from pypdf import PdfReader
from backend import store,multi_stock_cutting as multi,limited_stock_cutting as limited,mixed_stock_cutting as mixed
from backend.engine import DomainError
from backend.railing_material import calculate
from tests.test_multi_stock_cutting import body as multi_body,setup as multi_setup,geometry,member,equal_body,fixed_body,fixed_added,equal_added,fixed_adjusted,equal_adjusted,combined,local_body,path_body
from tests.test_limited_stock_cutting import limits
from tests.test_mixed_stock_cutting import config as mixed_config
from tests.test_route_takeoff import evaluate
from tests.test_railing_takeoff import interval
from tests.test_polygon_takeoff import post,mutate
from tests.test_workflow import get,approve,upload,create,priced_project

PARAMS=('trim_start','trim_end','reusable_min_length')
TITLES={'逐规格加工参数','逐规格加工范围依据'}


def row(eid='test',start='4',end='6',threshold='10'):
 return dict(trim_start=start,trim_end=end,reusable_min_length=threshold,evidence={k:eid for k in PARAMS})


def processing(eid='test'):
 return dict(alternatives=[row(eid),row(eid,'1','2','21')],scope_confirmed=True,measurement_basis='合成S1沿用本组，S2起端4终端6mm留用10mm；S3起端1终端2mm留用21mm，按逐规格原件核定，锯缝仍每件3mm含最后一件，未重复扣切除，非真实工程',evidence_id=eid)


def body(s,eid='test'):
 b=multi_body(s,eid);b['name']='合成逐规格加工（非真实工程）';c=b['mixed_cutting'][0]['stock_choices'];c.update(processing=processing(eid),measurement_basis='合成同规格方管A，三种定尺可按需供应，每件锯缝统一，端切和余料阈值按独立逐规格加工依据，不增加净量或报价');return b


def setup(clients,pid=None):
 e,r,pid,s,b,eid=multi_setup(clients,pid)
 upload(e,pid,'合成逐规格加工软件验收，非真实工程。同规格方管A每米3000g，长短每组各2件3500/2400mm，重复2共8件净70.8kg。S1为6000mm，起终端各10mm，留用阈值50mm；S2为3523mm，起4终6mm，留用阈值10mm；S3为2423mm，起1终2mm，留用阈值21mm。所有每件锯缝3mm含最后一件且端切不重复，各规格可按需供应。选S2/S3各4根，共23.784m71.352kg，S2余10mm等于阈值待留用，S3余17mm低于21mm报废；总端切.052m、锯缝.024m、报废.068m、待留用.04m，平衡0。S2起端改15mm后单根无法放3500mm长料加3mm锯缝，改选四根6m，24m72kg，净70.8kg不变须重新复核。S3阈值17mm时余料等值留用。只列方案，不计采购价或自动入库。','stock-processing-basis.txt')
 for c in get(e,pid)['components']:approve(r,pid,'components',c)
 s=get(e,pid)['components'][-1];eid=get(e,pid)['evidence'][-1]['id'];return e,r,pid,s,body(s,eid),eid


def group(s,b):return calculate(b,s)['mixed_stock_cutting']['groups'][0]


def test_per_stock_positions_thresholds_balances_and_asts():
 s=geometry();b=body(s);before=copy.deepcopy((s,b));v=calculate(b,s);g=v['mixed_stock_cutting']['groups'][0];base=calculate(multi_body(s),s)
 assert (s,b)==before and v['quantity']==base['quantity']=='70.800' and v['material_items']==base['material_items'] and v['formula_ast']==base['formula_ast'] and v['rule_version']=='railing-material-mass-1.17'
 assert g['stock_selection']['processing_independent'] and g['proof']['processing_model']=='per_stock_end_trims' and g['proof']['status']=='length_optimal_search' and g['total']['purchased_mass']=='71.352'
 assert g['total']['trim_length']=='0.052' and g['total']['scrap_length']=='0.068' and g['total']['retained_length']=='0.04' and g['balance_difference']=='0'
 assert len([k for k in v['source_evidence'] if '.mixed_cutting.' in k])==17
 for o in g['stock_selection']['options']:
  for k,ast in o['processing']['formula_asts'].items():assert evaluate(ast,v['dimensions'])==Decimal(o['processing'][k])
 for bar in g['bars']:
  if bar['stock_choice_index']==2:assert bar['pieces'][0]['from_s']=='0.004' and bar['offcut_length']=='0.01' and bar['offcut_disposition']=='retain'
  else:assert bar['stock_choice_index']==3 and bar['pieces'][0]['from_s']=='0.001' and bar['offcut_length']=='0.017' and bar['offcut_disposition']=='scrap'
  for k,ast in bar['formula_asts'].items():assert evaluate(ast,v['dimensions'])==Decimal(bar['values'][k])
 for k,ast in g['total_formula_asts'].items():assert evaluate(ast,v['dimensions'])==Decimal(g['total'][k])


def test_end_trim_changes_selection_without_changing_net():
 s=geometry();b=body(s);b['mixed_cutting'][0]['stock_choices']['processing']['alternatives'][0]['trim_start']='15';g=group(s,b)
 assert g['stock_count']==4 and g['total']['purchased_mass']=='72' and [o['stock_count'] for o in g['stock_selection']['options']]==[4,0,0] and g['total']['retained_length']=='0.296'


@pytest.mark.parametrize('threshold,status',[('17','retain'),('17.000000001','scrap'),('16.999999999','retain')])
def test_exact_threshold_per_stock(threshold,status):
 s=geometry();b=body(s);b['mixed_cutting'][0]['stock_choices']['processing']['alternatives'][1]['reusable_min_length']=threshold;g=group(s,b);assert all(x['offcut_disposition']==status for x in g['bars'] if x['stock_choice_index']==3)


def test_zero_end_trims_and_smaller_than_group_trim_alternative():
 s=geometry();b=body(s);c=b['mixed_cutting'][0];c['trim_start']='1500';c['trim_end']='1500';c['stock_choices']['alternatives'][1]['stock_length']='2420';p=c['stock_choices']['processing'];p['alternatives']=[row(start='0',end='0'),row(start='0',end='0',threshold='17')];g=group(s,b);assert g['stock_selection']['options'][2]['processing']['usable_length']=='2.42' and g['total']['trim_length']=='0'


def test_finite_supply_combines_with_processing_and_capacity_rejection():
 s=geometry();b=body(s);c=b['mixed_cutting'][0]['stock_choices'];c.update(supply_limits=limits(),supply_confirmed=False);g=group(s,b)
 assert g['stock_selection']['supply_limited'] and g['total']['purchased_mass']=='71.676' and g['total']['trim_length']=='0.066' and g['total']['retained_length']=='0.168' and g['total']['scrap_length']=='0.034'
 c['processing']['alternatives'][0]['trim_start']='15'
 with pytest.raises(DomainError,match='已由容量约束或穷尽搜索证明'):group(s,b)


@pytest.mark.parametrize('fault',['null','list','empty','extra','missing','unconfirmed','truthy','no-basis','no-proof','bad-proof','rows-null','rows-dict','rows-missing','rows-extra','row-null','row-list','row-missing','row-extra','row-proof-missing','row-proof-extra','row-proof-blank','row-proof-bool','negative','zero-threshold','nan','infinite','bool','precision','trim-exhausted','no-piece','no-long-fit'])
def test_strict_processing_schema_and_boundaries(fault):
 s=geometry();b=body(s);c=b['mixed_cutting'][0]['stock_choices'];p=c['processing'];r=p['alternatives'][0]
 if fault=='null':c['processing']=None
 elif fault=='list':c['processing']=[]
 elif fault=='empty':c['processing']={}
 elif fault=='extra':p['kerf']='3'
 elif fault=='missing':p.pop('scope_confirmed')
 elif fault=='unconfirmed':p['scope_confirmed']=False
 elif fault=='truthy':p['scope_confirmed']=1
 elif fault=='no-basis':p['measurement_basis']=''
 elif fault=='no-proof':p['evidence_id']=''
 elif fault=='bad-proof':p['evidence_id']=False
 elif fault=='rows-null':p['alternatives']=None
 elif fault=='rows-dict':p['alternatives']={}
 elif fault=='rows-missing':p['alternatives'].pop()
 elif fault=='rows-extra':p['alternatives'].append(r)
 elif fault=='row-null':p['alternatives'][0]=None
 elif fault=='row-list':p['alternatives'][0]=[]
 elif fault=='row-missing':r.pop('trim_start')
 elif fault=='row-extra':r['kerf_per_piece']='1'
 elif fault=='row-proof-missing':r['evidence'].pop('trim_end')
 elif fault=='row-proof-extra':r['evidence']['scope']='test'
 elif fault=='row-proof-blank':r['evidence']['trim_end']=''
 elif fault=='row-proof-bool':r['evidence']['trim_end']=True
 elif fault=='negative':r['trim_start']='-1'
 elif fault=='zero-threshold':r['reusable_min_length']='0'
 elif fault=='nan':r['trim_start']='NaN'
 elif fault=='infinite':r['trim_end']='Infinity'
 elif fault=='bool':r['trim_end']=True
 elif fault=='precision':r['trim_start']='.1234567890123456789'
 elif fault=='trim-exhausted':r['trim_start']='3517'
 elif fault=='no-piece':r['trim_start']='1200'
 else:
  b['mixed_cutting'][0]['trim_start']='3000';b['mixed_cutting'][0]['trim_end']='1';r['trim_start']='20'
 with pytest.raises(DomainError):calculate(b,s)


@pytest.mark.parametrize('unit,factor',[('m',Decimal('.001')),('mm',Decimal('1')),('cm',Decimal('.1')),('dm',Decimal('.01')),('km',Decimal('.000001'))])
def test_length_units_and_low_context(unit,factor):
 s=geometry();b=body(s);g=b['mixed_cutting'][0];g['length_unit']=unit
 for k in ('stock_length','kerf_per_piece','trim_start','trim_end','reusable_min_length'):g[k]=str(Decimal(g[k])*factor)
 for o in g['stock_choices']['alternatives']:o['stock_length']=str(Decimal(o['stock_length'])*factor)
 for o in g['stock_choices']['processing']['alternatives']:
  for k in PARAMS:o[k]=str(Decimal(o[k])*factor)
 with localcontext() as ctx:ctx.prec=6;v=calculate(b,s)
 assert v==calculate(b,s) and v['mixed_stock_cutting']['total']['purchased_mass']=='71.352'


@pytest.mark.parametrize('weight,unit',[('3','kg/m'),('3000','g/m'),('.003','kg/mm'),('3','g/mm')])
def test_weight_units(weight,unit):
 s=geometry();b=body(s)
 for m in b['material_items']:m.update(unit_weight=weight,weight_unit=unit)
 assert group(s,b)['total']['purchased_mass']=='71.352'


@pytest.mark.parametrize('make',[equal_body,fixed_body,fixed_added,equal_added,fixed_adjusted,equal_adjusted,combined,lambda s:combined(s,equal=True),local_body,path_body])
def test_post_modes_and_spacing_not_mutated(make):
 s=geometry();b=make(s);b['material_items'][0]['name']='合成立柱';b['material_items'].append({**member('合成短料','cut_pieces'),'quantity':'2','cut_length':'2400'});old=calculate(copy.deepcopy(b),s);g=mixed_config(members=['合成立柱','合成短料']);g['stock_choices']=body(s)['mixed_cutting'][0]['stock_choices'];b['mixed_cutting']=[g];v=calculate(b,s);assert v['quantity']==old['quantity'] and v['material_items']==old['material_items']


def test_seeded_independent_variable_trim_oracle():
 rng=random.Random(8401)
 def oracle(sizes,stocks,trims,counts):
  best=None;bins=[];used=[0]*len(stocks)
  def walk(i):
   nonlocal best
   if i==len(sizes):
    score=(sum(stocks[j] for j,_ in bins),len(bins));best=score if best is None else min(best,score);return
   for k,(j,load) in enumerate(bins):
    if load+sizes[i]<=stocks[j]-trims[j]:bins[k]=(j,load+sizes[i]);walk(i+1);bins[k]=(j,load)
   for j,stock in enumerate(stocks):
    if used[j]<counts[j] and sizes[i]<=stock-trims[j]:used[j]+=1;bins.append((j,sizes[i]));walk(i+1);bins.pop();used[j]-=1
  walk(0);return best
 for _ in range(70):
  stocks=sorted(rng.sample(range(5,15),3));trims=[rng.randint(0,3) for _ in stocks];sizes=sorted([rng.randint(1,max(s-t for s,t in zip(stocks,trims))) for _ in range(rng.randint(2,6))],reverse=True)
  for counts in ([len(sizes)]*3,[rng.randint(0,3) for _ in stocks]):
   expected=oracle(sizes,stocks,trims,counts)
   if expected is None:
    with pytest.raises(DomainError,match='已由容量约束或穷尽搜索证明'):limited.pack(sizes,stocks,trims,counts)
   else:
    plan,p=limited.pack(sizes,stocks,trims,counts);assert (p['purchased_length'],p['stock_count'])==expected and p['status']!='length_feasible_unproven'
    assert all(sum(sizes[i] for i in ps)<=stocks[j]-trims[j] for j,ps in plan) and all(Counter(j for j,_ in plan)[j]<=counts[j] for j in range(3))
   if counts==[len(sizes)]*3:
    plan,p=multi.pack(sizes,stocks,trims);assert (p['purchased_length'],p['stock_count'])==expected and p['status']!='length_feasible_unproven'


def test_max_200_pieces_five_stock_options():
 s=geometry();b=body(s)
 for m in b['material_items']:m['quantity']='50'
 c=b['mixed_cutting'][0]['stock_choices'];c['alternatives']+=[dict(stock_length=x,evidence_id='test') for x in ('5926','8000')];c['processing']['alternatives']+=[row(),row()];assert group(s,b)['required_piece_count']==200


@pytest.mark.parametrize('fmt',['json','xlsx','pdf'])
def test_full_workflow_export_and_old_quote_isolation(clients,fmt):
 e,r,pid,_,quote=priced_project(clients);e,r,pid,s,b,eid=setup(clients,pid);before=get(e,pid);post(e,pid,b,'/components/preview');assert get(e,pid)==before;c=post(e,pid,b).json();approve(r,pid,'components',c)
 for t in post(e,pid,{},'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
 assert any(t['quantity']=='70.800' for t in post(e,pid,{},'/boq/generate').json());url=f'/api/v1/projects/{pid}/exports';res=e.get(url,params=dict(format=fmt,unpriced=True));assert res.status_code==200,res.text
 if fmt=='json':assert TITLES<={t[0] for t in res.json()['tables']}
 elif fmt=='xlsx':assert TITLES<=set(load_workbook(io.BytesIO(res.content)).sheetnames)
 else:assert all(t in ''.join(p.extract_text() for p in PdfReader(io.BytesIO(res.content)).pages) for t in TITLES)
 assert not any(t[0] in TITLES for t in e.get(url,params=dict(format='json',quote_id=quote['id'])).json()['tables'])


@pytest.mark.parametrize('fault',['raw','proof','scope','normalized','usable','ast','position','offcut','book'])
def test_saved_rebuild_rejects_processing_tamper(clients,fault):
 e,r,pid,s,b,eid=setup(clients);c=post(e,pid,b).json()
 if fault=='book':approve(r,pid,'components',c);post(e,pid,{},'/takeoffs/calculate')
 def forge(p):
  c=p['components'][-1];g=c['preview']['mixed_stock_cutting']['groups'][0]
  if fault=='raw':c['mixed_cutting'][0]['stock_choices']['processing']['alternatives'][0]['trim_start']='5'
  elif fault=='proof':c['mixed_cutting'][0]['stock_choices']['processing']['alternatives'][0]['evidence']['trim_end']='forged'
  elif fault=='scope':c['mixed_cutting'][0]['stock_choices']['processing']['scope_confirmed']=False
  elif fault=='normalized':g['stock_selection']['options'][1]['processing']['trim_start']='0.005'
  elif fault=='usable':g['stock_selection']['options'][1]['processing']['usable_length']='3.5'
  elif fault=='ast':g['stock_selection']['options'][1]['processing']['formula_asts']['usable_length']={'constant':'3.5'}
  elif fault=='position':g['bars'][0]['pieces'][0]['from_s']='0.01'
  elif fault=='offcut':g['bars'][0]['offcut_disposition']='scrap'
  else:p['takeoffs'][-1]['mixed_stock_cutting']['groups'][0]['proof']['processing_model']='forged'
 mutate(pid,forge);coll='takeoffs' if fault=='book' else 'components';approve(r,pid,coll,get(e,pid)[coll][-1],expected=404 if fault=='proof' else 422 if fault=='scope' else 409)


def test_same_net_revision_removal_history_concurrency_and_source(clients):
 e,r,pid,s,b,eid=setup(clients);c=post(e,pid,b).json();approve(r,pid,'components',c)
 for t in post(e,pid,{},'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
 post(e,pid,{},'/boq/generate');before=get(e,pid);old=before['components'][-1];new=copy.deepcopy(b);new['mixed_cutting'][0]['stock_choices']['processing']['alternatives'][0]['trim_start']='15';new['object_hash']=store.digest(old);url=f'/api/v1/projects/{pid}/components/'+old['id'];res=e.put(url,json=new);assert res.status_code==200,res.text
 p=get(e,pid);assert p['component_history'][-1]==old and p['boq']==before['boq'] and all(t['state']=='stale' for t in p['takeoffs']) and res.json()['preview']['quantity']=='70.800' and res.json()['preview']['mixed_stock_cutting']['total']['purchased_mass']=='72';assert e.put(url,json=new).status_code==409;approve(r,pid,'components',res.json())
 current=get(e,pid)['components'][-1];new['object_hash']=store.digest(current);new['mixed_cutting'][0]['stock_choices'].pop('processing');res=e.put(url,json=new);assert res.status_code==200 and res.json()['preview']['rule_version']=='railing-material-mass-1.15';approve(r,pid,'components',res.json())
 s['object_hash']=store.digest(s);s['edge_decisions'][0].update(mode='ranges',ranges=[interval(0,10,eid)]);res=e.put(f'/api/v1/projects/{pid}/components/'+s['id'],json=s);assert res.status_code==200;approve(r,pid,'components',res.json());assert get(e,pid)['components'][-1]['state']=='stale'


@pytest.mark.parametrize('gate',['preview','save','review','calculate','book_review'])
def test_processing_original_bytes_gate(clients,gate):
 e,r,pid,s,b,eid=setup(clients)
 if gate in ('review','calculate','book_review'):c=post(e,pid,b).json()
 if gate in ('calculate','book_review'):approve(r,pid,'components',c)
 if gate=='book_review':t=post(e,pid,{},'/takeoffs/calculate').json()[-1]
 d=get(e,pid)['documents'][-1];(store.ROOT/'originals'/d['hash']).write_bytes(b'corrupt')
 if gate=='preview':post(e,pid,b,'/components/preview',409)
 elif gate=='save':post(e,pid,b,expected=409)
 elif gate=='review':approve(r,pid,'components',c,expected=409)
 elif gate=='calculate':post(e,pid,{},'/takeoffs/calculate',409)
 else:approve(r,pid,'takeoffs',t,expected=409)


def test_roles_tenant_foreign_processing_evidence(clients):
 e,r,pid,s,b,eid=setup(clients);post(r,pid,b,expected=403);assert clients['other'].get(f'/api/v1/projects/{pid}').status_code==404
 other=create(e)['id'];upload(e,other,'异项目');fid=get(e,other)['evidence'][-1]['id'];original=copy.deepcopy(b);b['mixed_cutting'][0]['stock_choices']['processing']['alternatives'][0]['evidence']['trim_end']=fid;post(e,pid,b,expected=404)
 b=original;b['mixed_cutting'][0]['stock_choices']['processing']['evidence_id']=fid;post(e,pid,b,expected=404)


def test_unproven_and_long_names_export(clients,monkeypatch):
 monkeypatch.setattr(mixed,'SEARCH_LIMIT',0);e,r,pid,s,b,eid=setup(clients);b['name']='合成'+('长名称验收'*39)+'末';b['mixed_cutting'][0]['name']='合成'+('长混排组名'*39)+'末';c=post(e,pid,b).json();assert c['preview']['mixed_stock_cutting']['groups'][0]['proof']['status']=='length_feasible_unproven';approve(r,pid,'components',c)
 for t in post(e,pid,{},'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
 res=e.get(f'/api/v1/projects/{pid}/exports',params=dict(format='pdf',unpriced=True));assert res.status_code==200,res.text
