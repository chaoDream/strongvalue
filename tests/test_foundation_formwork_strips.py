"""Synthetic signed planar strips, with physical side aperture intersections."""
import copy,io,random
from decimal import Decimal,localcontext,ROUND_HALF_UP
import pytest
from openpyxl import load_workbook
from pypdf import PdfReader
from backend import store
from backend.engine import DomainError
from backend.foundation_formwork import calculate
from tests.test_foundation_formwork import foundation,body as full_body,contact,setup as full_setup
from tests.test_post_foundation import equal_body,fixed_body,fixed_added,equal_added,fixed_adjusted,equal_adjusted,combined
from tests.test_route_takeoff import evaluate
from tests.test_polygon_takeoff import post,mutate
from tests.test_workflow import get,approve,upload,create,priced_project

TITLES={'基础模板横向条带','基础模板条带基准与原件','基础模板条带分离校验','基础模板条带开口交集'}
FRAME='合成基础中心为0，局部X向设计东、Y向设计北，X两侧横向均沿Y正向，Y两侧沿X正向，不随观察面翻转；仅软件验收。'

def strip(eid='test',start='-100',end='0'):
 return dict(parameters=dict(lateral_start=start,lateral_end=end),unit='mm',evidence=dict(lateral_start=eid,lateral_end=eid),coordinate_basis=FRAME,coordinate_evidence_id=eid,scope_confirmed=True)

def strip_contact(eid='test',start='-100',end='0',z_start='0',z_end='600',face='x_plus',level=1,name='合成C1'):
 return {**contact(eid,face,z_start,z_end,level,name),'lateral_range':strip(eid,start,end)}

def body(s,eid='test'):
 b=full_body(s,eid);b.update(name='合成矩形条带模板（非真实工程）',measurement_basis='合成22基础X正侧，沿Y选C1负100至0及C2正50至150mm，均Z0至600mm；仅C1与侧开口Y负30至30、Z0至150交叠，单基础扣.0045m²，合计净2.541m²。未选范围不计，不再乘重复。');b['contact_ranges']=[strip_contact(eid),strip_contact(eid,'50','150',name='合成C2')];return b

def setup(clients,pid=None):
 e,r,pid,s,_,_=full_setup(clients,pid)
 upload(e,pid,FRAME+'22个矩形基础400×500×600mm；X正侧侧开口Y负30至30、Z0至150mm，体积阈值不扣孔但侧开口仍扣。模板C1沿Y负100至0、C2正50至150，均Z0至600，每基础毛.12、实际扣.0045、净.1155m2，全部2.541m2。C1止坐标改20则全部2.739m2；止坐标60与C2正范围重叠应拒，越本层半宽250mm拒，基准不同拒。TEST-STRIP仅软件清单编码，不是官方条目。','foundation-formwork-strips-basis.txt')
 for c in get(e,pid)['components']:approve(r,pid,'components',c)
 s=get(e,pid)['components'][-1];eid=get(e,pid)['evidence'][-1]['id'];return e,r,pid,s,body(s,eid),eid

def check_asts(r):
 with localcontext() as ctx:
  ctx.prec=80;assert evaluate(r['formula_ast'],r['dimensions'])==Decimal(r['quantity'])
  for d in r['contact_ranges']:
   for k,a in d['formula_asts'].items():assert evaluate(a,r['dimensions'])==Decimal(d['normalized_parameters'][k] if k in ('z_start','z_end') else d[k]),k
   for h in d['apertures']:
    for k,a in h['formula_asts'].items():assert evaluate(a,r['dimensions'])==Decimal(h[k]),k
  for c in r['lateral_checks']:
   for k,a in c['formula_asts'].items():assert evaluate(a,r['dimensions'])==Decimal(c[k]),k

def test_partial_width_clips_real_hole_not_volume_threshold_and_asts():
 s=foundation();b=body(s);before=copy.deepcopy((s,b));r=calculate(b,s);assert (s,b)==before and r['quantity']=='2.541' and r['rule_version']=='foundation-outer-formwork-1.2';assert Decimal(r['gross'])==Decimal('2.64') and Decimal(r['deduction'])==Decimal('.099');h=r['contact_ranges'][0]['apertures'][0];assert tuple(Decimal(h[k]) for k in ('lateral_start','lateral_end'))==(Decimal('-.03'),Decimal(0)) and Decimal(h['width'])==Decimal('.03');assert not r['contact_ranges'][1]['apertures'];check_asts(r)

@pytest.mark.parametrize('make',[equal_body,fixed_body,fixed_added,equal_added,fixed_adjusted,equal_adjusted,combined,lambda s:combined(s,equal=True)])
def test_inherits_final_counts_once(make):
 s=foundation(make);r=calculate(body(s),s);assert Decimal(r['formwork_summary']['unrounded_quantity'])==Decimal('.1155')*Decimal(r['contact_ranges'][0]['total_count']);check_asts(r)

@pytest.mark.parametrize('unit,factor',[('mm','1'),('cm','.1'),('dm','.01'),('m','.001'),('km','.000001')])
def test_independent_length_units_and_low_precision(unit,factor):
 s=foundation();b=body(s)
 for d in b['contact_ranges']:
  c=d['lateral_range'];c['unit']=unit
  for k in c['parameters']:c['parameters'][k]=str(Decimal(c['parameters'][k])*Decimal(factor))
 with localcontext() as ctx:ctx.prec=6;r=calculate(b,s)
 assert r==calculate(b,s) and r['quantity']=='2.541';check_asts(r)

@pytest.mark.parametrize('face,x,y,start,end,cut',[('x_plus','150','0','-10','10','.003'),('x_minus','-150','0','-10','10','.003'),('y_plus','0','220','-10','10','.003'),('y_minus','0','-220','-10','10','.003')])
def test_four_face_coordinate_axes_never_mirror(face,x,y,start,end,cut):
 s=foundation();s['foundation_items'][0]['through_holes']['holes'][0]['parameters'].update(x=x,y=y);b=body(s);b['contact_ranges']=[strip_contact(face=face,start=start,end=end)];r=calculate(b,s);d=r['contact_ranges'][0];assert d['lateral_axis']==('Y' if face.startswith('x_') else 'X') and Decimal(d['single_opening_area'])==Decimal(cut);check_asts(r)

@pytest.mark.parametrize('start,end,expected',[('-30','30','.009'),('-10','10','.003'),('-100','100','.009'),('-100','0','.0045'),('0','100','.0045'),('30','100','0'),('-100','-30','0'),('30.000000001','100','0')])
def test_exact_hole_horizontal_intersections(start,end,expected):
 s=foundation();b=body(s);b['contact_ranges']=[strip_contact(start=start,end=end)];r=calculate(b,s);assert Decimal(r['contact_ranges'][0]['single_opening_area'])==Decimal(expected);check_asts(r)

@pytest.mark.parametrize('start,end,z_start,z_end,expected',[('-10','10','100','600','.001'),('-10','10','150','600','0'),('-10','10','0','100','.002'),('30','100','0','100','0')])
def test_two_dimensional_clipping(start,end,z_start,z_end,expected):
 s=foundation();b=body(s);b['contact_ranges']=[strip_contact(start=start,end=end,z_start=z_start,z_end=z_end)];r=calculate(b,s);assert Decimal(r['contact_ranges'][0]['single_opening_area'])==Decimal(expected);check_asts(r)

@pytest.mark.parametrize('end,state',[('50','lateral_tangent'),('49.999999999','laterally_disjoint')])
def test_exact_lateral_gap_and_tangency(end,state):
 s=foundation();b=body(s);b['contact_ranges'][0]['lateral_range']['parameters']['lateral_end']=end;r=calculate(b,s);assert r['lateral_checks'][0]['state']==state;check_asts(r)
 b['contact_ranges'][0]['lateral_range']['parameters']['lateral_end']='50.000000001'
 with pytest.raises(DomainError,match='重叠'):calculate(b,s)

@pytest.mark.parametrize('start,state',[('100','vertical_tangent'),('100.000000001','vertically_disjoint')])
def test_lateral_overlap_allowed_with_separate_heights(start,state):
 s=foundation();b=body(s);b['contact_ranges']=[strip_contact(z_end='100'),strip_contact(z_start=start,name='合成C2')];r=calculate(b,s);assert r['lateral_checks'][0]['state']==state;check_asts(r)
 b['contact_ranges'][1]['parameters']['z_start']='99.999999999'
 with pytest.raises(DomainError,match='重叠'):calculate(b,s)

def test_four_adjacent_strips_equal_full_and_whole_claim_participates():
 s=foundation();b=body(s);b['contact_ranges']=[strip_contact(start=str(a),end=str(z),name=str(a)) for a,z in [(-250,-30),(-30,0),(0,30),(30,250)]];r=calculate(b,s);assert r['quantity']==calculate(full_body(s),s)['quantity']=='6.402';check_asts(r)
 b['contact_ranges'].append(contact(name='整宽重复'))
 with pytest.raises(DomainError,match='重叠'):calculate(b,s)
 # Whole width and strip with only vertical tangency is valid and has reproducible gaps.
 b['contact_ranges']=[{**contact(end='100')},strip_contact(z_start='100',name='合成C2')];check_asts(calculate(b,s))

def test_tiers_use_own_width_and_z_and_common_frame():
 s=foundation(shape='stepped_rectangular');b=body(s);b['contact_ranges']=[strip_contact(start='-100',end='100',z_end='200'),strip_contact(start='-50',end='50',z_start='200',level=2,name='合成C2')];r=calculate(b,s);assert r['quantity']=='1.760' and len({d['full_face_width'] for d in r['contact_ranges']})==2;check_asts(r)
 b['contact_ranges'][1]['lateral_range']['parameters']['lateral_end']='200'
 with pytest.raises(DomainError,match='半宽'):calculate(b,s)

def test_seeded_oracle_150_intersections_and_final_rounding():
 rng=random.Random(9501);s=foundation()
 for _ in range(150):
  a=rng.randint(-250,249);z=rng.randint(a+1,250);lo=rng.randint(0,599);hi=rng.randint(lo+1,600);b=body(s);b['contact_ranges']=[strip_contact(start=str(a),end=str(z),z_start=str(lo),z_end=str(hi))];b['rounding_digits']=rng.randint(0,12);r=calculate(b,s)
  with localcontext() as ctx:ctx.prec=80;gross=Decimal((z-a)*(hi-lo))/1000000*22;cut=Decimal(max(0,min(z,30)-max(a,-30))*max(0,min(hi,150)-max(lo,0)))/1000000*22;assert Decimal(r['gross'])==gross and Decimal(r['deduction'])==cut and Decimal(r['quantity'])==(gross-cut).quantize(Decimal(1).scaleb(-b['rounding_digits']),rounding=ROUND_HALF_UP)
  check_asts(r)

@pytest.mark.parametrize('fault',['null','empty','extra','missing','params','proofs','start-proof','end-proof','frame','frame-proof','unconfirmed','truthy','unit','reverse','zero','nan','inf','bool','precision','start-outside','end-outside','width-instead','circular','both-modes','frame-different','frame-proof-different'])
def test_strict_configuration(fault):
 s=foundation();b=body(s);c=b['contact_ranges'][0]['lateral_range']
 if fault=='null':b['contact_ranges'][0]['lateral_range']=None
 elif fault=='empty':b['contact_ranges'][0]['lateral_range']={}
 elif fault=='extra':c['width']='100'
 elif fault=='missing':c.pop('unit')
 elif fault=='params':c['parameters'].pop('lateral_start')
 elif fault=='proofs':c['evidence'].pop('lateral_start')
 elif fault in ('start-proof','end-proof'):c['evidence']['lateral_start' if fault=='start-proof' else 'lateral_end']=' '
 elif fault=='frame':c['coordinate_basis']=''
 elif fault=='frame-proof':c['coordinate_evidence_id']=''
 elif fault in ('unconfirmed','truthy'):c['scope_confirmed']=False if fault=='unconfirmed' else 1
 elif fault=='unit':c['unit']='°'
 elif fault in ('reverse','zero','nan','inf','bool','precision','start-outside','end-outside'):c['parameters']['lateral_start' if fault=='start-outside' else 'lateral_end']={'reverse':'-101','zero':'-100','nan':'NaN','inf':'Infinity','bool':True,'precision':'0.0000000000000000001','start-outside':'-250.000000001','end-outside':'250.000000001'}[fault]
 elif fault=='width-instead':c['parameters']['width']='100'
 elif fault=='circular':
  s=foundation(shape='circular');b['source_component_hash']=store.digest({k:v for k,v in s.items() if k!='state'});b['contact_ranges'][0]['face']='circumference'
 elif fault=='both-modes':b['contact_ranges'][0]['arc_range']={}
 elif fault=='frame-different':b['contact_ranges'][1]['lateral_range']['coordinate_basis']='不同坐标基准'
 elif fault=='frame-proof-different':b['contact_ranges'][1]['lateral_range']['coordinate_evidence_id']='other'
 with pytest.raises(DomainError):calculate(b,s)

def test_preview_no_mutation_claims_frames_and_forged_id(clients):
 e,r,pid,s,b,eid=setup(clients);before=get(e,pid);assert post(e,pid,b,'/components/preview').json()['quantity']=='2.541' and get(e,pid)==before
 first=copy.deepcopy(b);first['contact_ranges']=first['contact_ranges'][:1];c=post(e,pid,first).json();approve(r,pid,'components',c);second=copy.deepcopy(b);second['contact_ranges']=second['contact_ranges'][1:];d=post(e,pid,second).json();approve(r,pid,'components',d);post(e,pid,{**first,'id':c['id']},expected=409)
 second['contact_ranges'][0]['lateral_range']['parameters']['lateral_start']='-0.000000001';post(e,pid,second,expected=409)
 second['contact_ranges'][0]['lateral_range']['parameters']['lateral_start']='0';second['contact_ranges'][0]['lateral_range']['parameters']['lateral_end']='50';second['contact_ranges'][0]['lateral_range']['coordinate_basis']='改变基准';post(e,pid,second,expected=409)
 second['contact_ranges'][0]['lateral_range']['coordinate_basis']=FRAME;assert post(e,pid,second).status_code==200

@pytest.mark.parametrize('fault',['coordinate','frame','proof','confirm','width','full-width','aperture-coordinate','gap','formula','book'])
def test_saved_tamper_rebuilt(clients,fault):
 e,r,pid,s,b,eid=setup(clients);c=post(e,pid,b).json()
 if fault=='book':approve(r,pid,'components',c);post(e,pid,{},'/takeoffs/calculate')
 def forge(p):
  c=p['components'][-1];d=c['contact_ranges'][0];v=c['preview']['contact_ranges'][0]
  if fault=='coordinate':d['lateral_range']['parameters']['lateral_end']='20'
  elif fault=='frame':
   for t in c['contact_ranges']:t['lateral_range']['coordinate_basis']='篡改基准'
  elif fault=='proof':d['lateral_range']['evidence']['lateral_start']='forged'
  elif fault=='confirm':d['lateral_range']['scope_confirmed']=False
  elif fault=='width':v['width']='9'
  elif fault=='full-width':v['full_face_width']='9'
  elif fault=='aperture-coordinate':v['apertures'][0]['lateral_start']='9'
  elif fault=='gap':c['preview']['lateral_checks'][0]['lateral_gap']='9'
  elif fault=='formula':v['formula_asts']['width']={'constant':'9'}
  else:p['takeoffs'][-1]['quantity']='9'
 mutate(pid,forge);coll='takeoffs' if fault=='book' else 'components';approve(r,pid,coll,get(e,pid)[coll][-1],expected=404 if fault=='proof' else 422 if fault=='confirm' else 409)

@pytest.mark.parametrize('gate',['preview','save','review','calculate','book_review'])
def test_coordinate_original_five_gates(clients,gate):
 e,r,pid,s,b,eid=setup(clients)
 if gate in ('review','calculate','book_review'):c=post(e,pid,b).json()
 if gate in ('calculate','book_review'):approve(r,pid,'components',c)
 if gate=='book_review':t=post(e,pid,{},'/takeoffs/calculate').json()[-1]
 doc=get(e,pid)['documents'][-1];(store.ROOT/'originals'/doc['hash']).write_bytes(b'corrupt')
 if gate=='preview':post(e,pid,b,'/components/preview',409)
 elif gate=='save':post(e,pid,b,expected=409)
 elif gate=='review':approve(r,pid,'components',c,expected=409)
 elif gate=='calculate':post(e,pid,{},'/takeoffs/calculate',409)
 else:approve(r,pid,'takeoffs',t,expected=409)

def test_revision_history_concurrency_disable_source_stale(clients):
 e,r,pid,s,b,eid=setup(clients);c=post(e,pid,b).json();approve(r,pid,'components',c)
 for t in post(e,pid,{},'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
 post(e,pid,{},'/boq/generate');p=get(e,pid);old=p['components'][-1];new=copy.deepcopy(b);new['object_hash']=store.digest(old);new['contact_ranges'][0]['lateral_range']['parameters']['lateral_end']='20';url=f'/api/v1/projects/{pid}/components/'+old['id'];res=e.put(url,json=new);assert res.status_code==200 and res.json()['preview']['quantity']=='2.739';assert get(e,pid)['component_history'][-1]==old and get(e,pid)['boq']==p['boq'] and all(t['state']=='stale' for t in get(e,pid)['takeoffs']);assert e.put(url,json=new).status_code==409;approve(r,pid,'components',res.json())
 new['contact_ranges']=[contact(eid)];new['object_hash']=store.digest(get(e,pid)['components'][-1]);res=e.put(url,json=new);assert res.status_code==200 and res.json()['preview']['quantity']=='6.402' and res.json()['preview']['rule_version']=='foundation-outer-formwork-1.0';approve(r,pid,'components',res.json())
 s['object_hash']=store.digest(s);s['foundation_items'][0]['parameters']['height']='700';res=e.put(f'/api/v1/projects/{pid}/components/'+s['id'],json=s);assert res.status_code==200;approve(r,pid,'components',res.json());assert get(e,pid)['components'][-1]['state']=='stale'

def test_roles_tenants_foreign_proof(clients):
 e,r,pid,s,b,eid=setup(clients);post(r,pid,b,expected=403);assert clients['other'].get(f'/api/v1/projects/{pid}').status_code==404;other=create(e)['id'];upload(e,other,'异项目');b['contact_ranges'][0]['lateral_range']['evidence']['lateral_start']=get(e,other)['evidence'][-1]['id'];post(e,pid,b,expected=404)

@pytest.mark.parametrize('fmt',['json','xlsx','pdf'])
def test_three_exports_boq_old_quote(clients,fmt):
 e,r,pid,_,q=priced_project(clients);e,r,pid,s,b,eid=setup(clients,pid);c=post(e,pid,b).json();approve(r,pid,'components',c)
 for t in post(e,pid,{},'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
 assert any(d['quantity']=='2.541' and d['unit']=='m²' for d in post(e,pid,{},'/boq/generate').json());url=f'/api/v1/projects/{pid}/exports';res=e.get(url,params=dict(format=fmt,unpriced=True));assert res.status_code==200,res.text
 if fmt=='json':assert TITLES<={d[0] for d in res.json()['tables']}
 elif fmt=='xlsx':assert TITLES<=set(load_workbook(io.BytesIO(res.content)).sheetnames)
 else:assert all(v in ''.join(p.extract_text() for p in PdfReader(io.BytesIO(res.content)).pages) for v in TITLES)
 assert not any(d[0] in TITLES for d in e.get(url,params=dict(format='json',quote_id=q['id'])).json()['tables'])

def test_strip_fully_in_side_void_zero_net_and_touching_holes_no_double_cut():
 s=foundation();b=body(s);b['contact_ranges']=[strip_contact(start='-30',end='30',z_end='150')];r=calculate(b,s);assert r['quantity']=='0.000' and Decimal(r['gross'])==Decimal(r['deduction'])>0;check_asts(r)
 c=s['foundation_items'][0]['through_holes'];h=c['holes'][0];h['parameters']['y']='-30';second=copy.deepcopy(h);second.update(name='合成相接孔');second['parameters']['y']='30';c['holes'].append(second);b['source_component_hash']=store.digest({k:v for k,v in s.items() if k!='state'});b['contact_ranges']=[strip_contact(start='-20',end='20',z_end='150')];r=calculate(b,s);assert r['quantity']=='0.000' and len(r['contact_ranges'][0]['apertures'])==2;check_asts(r)

def test_multi_item_rectangle_and_circle_arc_mixed_keep_independent_frames():
 from tests.test_foundation_formwork_arcs import arc_contact
 s=foundation();m=s['source_geometry'];m['material_items'].append({**copy.deepcopy(m['material_items'][0]),'name':'第二立柱项'});s['source_component_hash']=store.digest({k:v for k,v in m.items() if k!='state'});second=foundation(shape='circular')['foundation_items'][0];second['material_index']=2;s['foundation_items'].append(second);b=body(s);d=arc_contact(z_end='600',name='合成圆弧');d['material_index']=2;b['contact_ranges'].append(d);r=calculate(b,s);assert r['rule_version']=='foundation-outer-formwork-1.2' and 'arc_checks' in r and 'lateral_checks' in r;check_asts(r)
 for c in r['unit_conversions']:assert c['parameter'] in r['source_evidence'] and c['parameter'] in r['dimensions']

def test_twenty_adjacent_strips_long_basis_pdf(clients):
 e,r,pid,s,b,eid=setup(clients);b['contact_ranges']=[strip_contact(eid,str(-250+i*25),str(-225+i*25),name='合成C'+str(i)) for i in range(20)]
 for d in b['contact_ranges']:d['lateral_range']['coordinate_basis']='合成'+('局部坐标及实际接触范围依据'*100)
 c=post(e,pid,b).json();assert c['preview']['quantity']=='6.402';approve(r,pid,'components',c)
 for t in post(e,pid,{},'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
 post(e,pid,{},'/boq/generate');res=e.get(f'/api/v1/projects/{pid}/exports',params=dict(format='pdf',unpriced=True));assert res.status_code==200,res.text
