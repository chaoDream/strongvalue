"""Synthetic inner-window exclusions: independent geometry and workflow gates."""
import copy
import io
import random
from decimal import Decimal, localcontext
from fractions import Fraction
import pytest
from openpyxl import load_workbook
from pypdf import PdfReader
from backend import store
from backend.engine import DomainError
from backend.foundation_horizontal_formwork import calculate
from backend.foundation_perforated_windows import VERSION, METHOD, separation, validate
from backend.foundation_horizontal_windows import checks
from tests.test_foundation_polygon_windows import polygon, body as pb, replay, clip, quadrature
from tests.test_foundation_simple_windows import setup as simple_setup, C_SHAPE
from tests.test_foundation_horizontal_windows import source, window
from tests.test_foundation_hole_end_formwork import source as end_source
from tests.test_foundation_openings import hole
from tests.test_foundation_voids import rectangle
from tests.test_foundation_recesses import interval
from tests.test_post_foundation import equal_body, fixed_body, fixed_added, equal_added, fixed_adjusted, equal_adjusted, combined
from tests.test_polygon_takeoff import post, mutate
from tests.test_workflow import get, approve, upload, create, priced_project
from tests.test_route_takeoff import evaluate

OUTER=[(-100,-100),(100,-100),(100,100),(-100,100)]
INNER=[(-20,-20),(20,-20),(20,20),(-20,20)]


def excluded(pts=INNER,eid='test',name='合成不支模内区'):
    c=polygon(pts,eid)
    return dict(name=name,parameters=c['parameters'],evidence=c['evidence'],measurement_basis='合成窗口内部不支模范围；非基础实体孔',evidence_id=eid)


def config(outer=OUTER,inners=None,eid='test',unit='mm'):
    c=polygon(outer,eid,unit);c.update(shape='polygon_with_holes',holes=[excluded(p,eid,'合成W'+str(i)) for i,p in enumerate(inners if inners is not None else [INNER],1)])
    return c


def body(s,eid='test',face='bottom',outer=OUTER,inners=None):
    b=pb(s,eid,face,outer);b['contact_ranges'][0]['plan_range']=config(outer,inners,eid);return b


def replay_rings(r):
    replay(r)
    for d in r['contact_ranges']:
        for ring in d['window_rings']:assert evaluate(ring['formula_ast'],r['dimensions'])==Decimal(ring['area'])


@pytest.mark.parametrize('shape',['rectangular','circular','stepped_rectangular','stepped_circular'])
@pytest.mark.parametrize('face',['top','bottom','hole_top','hole_bottom'])
@pytest.mark.parametrize('make',[equal_body,fixed_body,fixed_added,equal_added,fixed_adjusted,equal_adjusted,combined,lambda s:combined(s,equal=True)])
def test_shapes_ends_final_counts_input_immutable_and_replay(shape,face,make):
    s=end_source(shape,make) if face.startswith('hole') else source(shape,make)
    b=body(s,face=face);before=copy.deepcopy((s,b));r=calculate(b,s);d=r['contact_ranges'][0]
    assert before==(s,b) and r['rule_version']==VERSION
    assert Decimal(d['window_outer_area'])==Decimal('.04') and Decimal(d['window_exclusion_area'])==Decimal('.0016')
    assert Decimal(d['window_area'])==Decimal('.0384') and d['window_ring_checks']==[dict(left='外轮廓',right='W1',relation='strictly_inside')]
    replay_rings(r)


@pytest.mark.parametrize('shape,face',[(s,f) for s in ['rectangular','circular','stepped_rectangular','stepped_circular'] for f in ['top','bottom','hole_bottom','hole_top']]+[('stepped_rectangular','step'),('stepped_circular','step')])
def test_direct_window_net_equals_outer_minus_inner_contact_without_double_cut(shape,face):
    s=end_source(shape) if face.startswith('hole') else source(shape)
    outer=[(100,-100),(270,-100),(270,100),(100,100)] if face=='step' else OUTER
    inner=[(170,-20),(200,-20),(200,20),(170,20)] if face=='step' else INNER
    r=calculate(body(s,face=face,outer=outer,inners=[inner]),s);d=r['contact_ranges'][0]
    def simple_result(pts):
        b=pb(s,face=face,pts=pts);b['contact_ranges'][0]['plan_range']['shape']='simple_polygon';return calculate(b,s)['contact_ranges'][0]
    a,h=simple_result(outer),simple_result(inner)
    with localcontext() as ctx:
        ctx.prec=80
        for k in ['single_gross_area','single_opening_area','single_net_area']:
            assert abs(Decimal(d[k])-(Decimal(a[k])-Decimal(h[k])))<Decimal('1e-75')
    replay_rings(r)


def test_40_seeded_inner_diamonds_against_independent_exact_rectangle_clip():
    rng=random.Random(10901);s=source();outer=[(-300,-300),(300,-300),(300,300),(-300,300)]
    for _ in range(40):
        x,y=rng.randint(-120,120),rng.randint(-150,150);rx,ry=rng.randint(5,120),rng.randint(5,100)
        inner=[(x-rx,y),(x,y-ry),(x+rx,y),(x,y+ry)]
        r=calculate(body(s,outer=outer,inners=[inner]),s);d=r['contact_ranges'][0]
        def area(pts):
            pts=[tuple(map(Fraction,p)) for p in pts]
            return clip(pts,tuple(map(Fraction,(-200,200,-250,250))))-clip(pts,tuple(map(Fraction,(100,200,-30,30))))
        expected=(area(outer)-area(inner))/1000000
        assert abs(float(d['single_net_area'])-float(expected))<2e-16
    replay_rings(r)


def test_window_exclusion_over_physical_hole_does_not_double_deduct():
    s=source();outer=[(-300,-300),(300,-300),(300,300),(-300,300)]
    inner=[(120,-10),(160,-10),(160,10),(120,10)]
    r=calculate(body(s,outer=outer,inners=[inner]),s);d=r['contact_ranges'][0]
    assert Decimal(d['single_gross_area'])==Decimal('.1992')
    assert Decimal(d['single_opening_area'])==Decimal('.0052') and Decimal(d['single_net_area'])==Decimal('.194')
    replay_rings(r)
    covering=[(90,-40),(210,-40),(210,40),(90,40)]
    r=calculate(body(s,outer=outer,inners=[covering]),s);d=r['contact_ranges'][0]
    assert Decimal(d['single_opening_area'])==0 and Decimal(d['single_net_area'])==Decimal('.1912')
    replay_rings(r)


@pytest.mark.parametrize('inner,expected',[
    ([(-400,-400),(400,-400),(400,400),(-400,400)],'0'),
    ([(600,-100),(700,-100),(700,100),(600,100)],None)])
def test_exclusion_covering_entire_physical_profile_or_entirely_outside(inner,expected):
    s=source('circular');s['foundation_items'][0].pop('through_holes',None)
    r=calculate(body(s,face='top',outer=[(-800,-800),(800,-800),(800,800),(-800,800)],inners=[inner]),s);d=r['contact_ranges'][0]
    if expected is not None:assert Decimal(d['single_gross_area'])==Decimal(expected) and 'polygon_zero_contact' not in d
    else:
        from backend.compound_takeoff import PI
        with localcontext() as ctx:
            ctx.prec=80;assert abs(Decimal(d['single_gross_area'])-Decimal(PI)*Decimal('.09'))<Decimal('1e-75')
    replay_rings(r)


def test_nested_window_net_inside_another_window_exclusion_is_separate():
    a={'plan_range':config()}
    b={'plan_range':config([(-15,-15),(15,-15),(15,15),(-15,15)],[[(-5,-5),(5,-5),(5,5),(-5,5)]])}
    assert not separation(a,b)['overlap'] and not separation(b,a)['overlap']
    b['plan_range']['parameters']=config(INNER)['parameters']
    assert separation(a,b)['state']=='boundary_touch'


@pytest.mark.parametrize('offset',[-350,-250,-100,0,100,250,350])
def test_curved_inner_exclusion_independent_adaptive_integral(offset):
    s=source('circular');s['foundation_items'][0].pop('through_holes',None);outer=[(-400,-400),(400,-400),(400,400),(-400,400)]
    inner=[(offset-30,-70),(offset+30,-10),(offset+20,80),(offset-40,20)]
    r=calculate(body(s,face='top',outer=outer,inners=[inner]),s);d=r['contact_ranges'][0]
    # Circular source top has radius300mm and no through-top physical hole.
    meters=lambda pts:[(x/1000,y/1000) for x,y in pts]
    expected=quadrature(meters(outer),[(0,0,.3)])-quadrature(meters(inner),[(0,0,.3)])
    assert abs(float(d['single_net_area'])-expected)<3e-11
    replay_rings(r)


@pytest.mark.parametrize('unit,factor',[('mm','1'),('cm','.1'),('dm','.01'),('m','.001'),('km','.000001')])
@pytest.mark.parametrize('reverse',[False,True])
def test_inherited_units_orientation_low_caller_precision(unit,factor,reverse):
    s=source();scale=lambda pts:[tuple(str(Decimal(str(v))*Decimal(factor)) for v in p) for p in (list(reversed(pts)) if reverse else pts)]
    b=body(s);b['contact_ranges'][0]['plan_range']=config(scale(OUTER),[scale(INNER)],unit=unit)
    with localcontext() as ctx:ctx.prec=7;r=calculate(b,s)
    assert Decimal(r['contact_ranges'][0]['window_area'])==Decimal('.0384');replay_rings(r)


def test_four_holes_36_points_and_five_vertical_pieces():
    outer=[(-100,-100),(100,-100),(100,100),(-100,100)]
    octagon=[(-10,-6),(-6,-10),(6,-10),(10,-6),(10,6),(6,10),(-6,10),(-10,6)]
    inners=[[(x,y+cy) for x,y in octagon] for cy in [-65,-25,25,65]]
    s=source();r=calculate(body(s,outer=outer,inners=inners),s);d=r['contact_ranges'][0]
    assert len(d['window_rings'])==5 and any(q['window_region_index']==5 for q in d['polygon_net_trace']['partitions'])
    assert len(d['window_ring_checks'])==10;replay_rings(r)
    b=body(s,outer=[(-100,-100),(100,-100),(100,100),(0,110),(-100,100)],inners=inners)
    with pytest.raises(DomainError,match='36'):calculate(b,s)


@pytest.mark.parametrize('inners',[
    [[(-100,-10),(-90,-10),(-90,10),(-100,10)]],
    [[(90,-10),(110,-10),(110,10),(90,10)]],
    [[(110,-10),(120,-10),(120,10),(110,10)]],
    [INNER,INNER],
    [INNER,[(20,-10),(40,-10),(40,10),(20,10)]],
    [INNER,[(-10,-10),(10,-10),(10,10),(-10,10)]],
    [[(-20,-20),(20,20),(-20,20),(20,-20)]],
    [[(-20,-20),(0,-20),(20,-20),(20,20),(-20,20)]],
    [[(-20,-20),(20,-20),(-20,-20),(20,20)]]])
def test_invalid_hole_topology_is_rejected(inners):
    with pytest.raises(DomainError):validate(config(inners=inners))


def test_concave_outer_chord_leaves_boundary_and_concave_inner_multiple_spans():
    with pytest.raises(DomainError):validate(config(C_SHAPE,[[(-70,-70),(70,-70),(-70,70)]]))
    outer=[(-150,-150),(150,-150),(150,150),(-150,150)]
    s=source();r=calculate(body(s,outer=outer,inners=[C_SHAPE]),s)
    assert Decimal(r['contact_ranges'][0]['window_area'])==Decimal('.0612');replay_rings(r)


@pytest.mark.parametrize('fault',['empty','too_many','shape','unit','evidence','basis','extra','extra_inner','confirm','hole_unit','missing_xy','not_list'])
def test_strict_schema_and_inner_evidence(fault):
    c=config()
    if fault=='empty':c['holes']=[]
    elif fault=='too_many':c['holes']*=5
    elif fault=='shape':c['holes'][0]=None
    elif fault=='unit':c['unit']='px'
    elif fault=='evidence':c['holes'][0]['evidence']['x1']=''
    elif fault=='basis':c['holes'][0]['evidence_id']=''
    elif fault=='extra':c['z']='0'
    elif fault=='extra_inner':c['holes'][0]['multiplicity']='2'
    elif fault=='confirm':c['scope_confirmed']=1
    elif fault=='hole_unit':c['holes'][0]['unit']='mm'
    elif fault=='missing_xy':c['holes'][0]['parameters'].pop('x1')
    else:c['holes']={}
    with pytest.raises(DomainError):validate(c)


@pytest.mark.parametrize('delta',['0','0.000000001'])
def test_inner_empty_area_other_window_allowed_boundary_and_tiny_overlap(delta):
    a=dict(index=1,name='带孔区',material_index=1,face='top',plan_range=config())
    b=dict(index=2,name='内孔内单独支模',material_index=1,face='top',plan_range=window('test','-20',str(Decimal(20)+Decimal(delta)),'-20','20'))
    r=separation(a,b);assert r['overlap']==(delta!='0')
    if delta=='0':
        assert r['state']=='boundary_touch' and r['method']==METHOD
        q=checks([a,b])[0];assert q['partitions']==r['partitions'] and q['left_index']==1
        assert any(len(p['left'])==2 and len(p['right'])==1 for p in q['partitions'])
    assert separation(a,a)['overlap']


def test_100_seeded_net_window_overlap_against_independent_fraction_clipping():
    rng=random.Random(10902);a={'plan_range':config()}
    for _ in range(100):
        x,y=rng.randint(-130,130),rng.randint(-130,130);box=tuple(map(Fraction,(x,x+rng.randint(1,130),y,y+rng.randint(1,130))))
        expected=clip([tuple(map(Fraction,p)) for p in OUTER],box)-clip([tuple(map(Fraction,p)) for p in INNER],box)>0
        b={'plan_range':window('test',*[str(v) for v in box])}
        assert separation(a,b)['overlap']==expected


@pytest.mark.parametrize('gap',['0','0.000000001'])
def test_zero_requires_all_remaining_window_fragments_no_tolerance(gap):
    v=str(Decimal(25)+Decimal(gap));hs=[{**hole(x='0',diameter='80'),'vertical_range':interval('100','300')},
        {**rectangle(name='H2',x='-'+v,length='50',width='100'),'vertical_range':interval('300','500')},
        {**rectangle(name='H3',x=v,length='50',width='100'),'vertical_range':interval('300','500')}]
    s=end_source(holes=hs);r=calculate(body(s,face='hole_top'),s);d=r['contact_ranges'][0]
    assert ('polygon_zero_contact' in d)==(gap=='0') and (Decimal(d['single_net_area'])==0)==(gap=='0')
    assert any(q['window_region_index']==2 for q in d['polygon_net_trace']['partitions']);replay_rings(r)


def setup(clients,pid=None,zero=False):
    e,r,pid,s,b,eid=simple_setup(clients,pid,zero)
    c=b['contact_ranges'][0]['plan_range'];c['shape']='polygon_with_holes'
    pts=[(-37,-4),(-30,-4),(-30,4),(-37,4)] if zero else [(171,-3),(174,-3),(174,3),(171,3)]
    c['holes']=[excluded(pts,eid)];return e,r,pid,s,b,eid


def test_workflow_roles_foreign_inner_evidence_readonly_history_source_stale(clients):
    e,r,pid,s,b,eid=setup(clients);before=get(e,pid);post(e,pid,b,'/components/preview');assert get(e,pid)==before
    post(r,pid,b,expected=403);assert clients['other'].get(f'/api/v1/projects/{pid}').status_code==404
    foreign=create(e)['id'];upload(e,foreign,'异项目合成');bad=copy.deepcopy(b);bad['contact_ranges'][0]['plan_range']['holes'][0]['evidence']['x1']=get(e,foreign)['evidence'][-1]['id'];post(e,pid,bad,expected=404)
    c=post(e,pid,b).json();approve(r,pid,'components',c);post(e,pid,b,expected=409)
    c=next(x for x in get(e,pid)['components'] if x['id']==c['id']);b['object_hash']=store.digest(c);b['contact_ranges'][0]['plan_range']['holes'][0]['parameters']['x1']='170.5'
    res=e.put(f'/api/v1/projects/{pid}/components/'+c['id'],json=b);assert res.status_code==200,res.text
    assert get(e,pid)['component_history'][-1]==c;approve(r,pid,'components',res.json())
    s['object_hash']=store.digest(s);s['foundation_items'][0]['levels'][1]['parameters']['diameter']='420'
    res=e.put(f'/api/v1/projects/{pid}/components/'+s['id'],json=s);assert res.status_code==200,res.text;approve(r,pid,'components',res.json())
    assert next(x for x in get(e,pid)['components'] if x['id']==c['id'])['state']=='stale'


def test_cross_record_inner_empty_region_allowed_true_overlap_and_frame_blocked(clients):
    e,r,pid,s,b,eid=setup(clients);post(e,pid,b)
    other=copy.deepcopy(b);other['name']='窗口内孔内单独支模';other['contact_ranges'][0]['plan_range']=window(eid,'171','174','-3','3');post(e,pid,other)
    bad=copy.deepcopy(other);bad['contact_ranges'][0]['plan_range']['parameters']['x_end']='174.000000001';before=get(e,pid);post(e,pid,bad,expected=409);assert get(e,pid)==before
    bad=copy.deepcopy(other);bad['contact_ranges'][0]['plan_range']['coordinate_basis']='其他正向';post(e,pid,bad,expected=409)


@pytest.mark.parametrize('fault',['inner_input','inner_basis','ring_area','ring_vertex','net_area','topology','region','proof','expression'])
def test_inner_input_and_stored_proof_tamper_rejected(clients,fault):
    e,r,pid,s,b,eid=setup(clients,zero=True);c=post(e,pid,b).json()
    def forge(p):
        c=p['components'][-1];d=c['preview']['contact_ranges'][0]
        if fault=='inner_input':c['contact_ranges'][0]['plan_range']['holes'][0]['parameters']['x1']='-36'
        elif fault=='inner_basis':c['contact_ranges'][0]['plan_range']['holes'][0]['measurement_basis']='篡改排除范围'
        elif fault=='ring_area':d['window_rings'][1]['area']='999'
        elif fault=='ring_vertex':d['window_rings'][1]['vertices'][0]['x']='999'
        elif fault=='net_area':d['window_area']='999'
        elif fault=='topology':d['window_ring_checks']=[]
        elif fault=='region':d['polygon_net_trace']['partitions'][0]['window_region_index']=99
        elif fault=='proof':d['polygon_zero_contact']['net_trace']['partitions'][0]['remaining']=[{}]
        else:d['window_expression_definitions'][0]['value']='999'
    mutate(pid,forge);approve(r,pid,'components',c,expected=409)


@pytest.mark.parametrize('gate',['preview','save','review','calculate','book_review'])
def test_inner_evidence_original_bytes_all_five_gates(clients,gate):
    e,r,pid,s,b,eid=setup(clients)
    upload(e,pid,'合成独立内孔排除及逐坐标依据','inner-window-basis.txt')
    for c in get(e,pid)['components']:approve(r,pid,'components',c)
    b['source_component_hash']=e.get(f'/api/v1/projects/{pid}/foundation-horizontal-formwork-candidates',params={'component_id':s['id']}).json()['source_component_hash']
    inner=b['contact_ranges'][0]['plan_range']['holes'][0];inner['evidence_id']=get(e,pid)['evidence'][-1]['id'];inner['evidence']={k:inner['evidence_id'] for k in inner['evidence']}
    if gate in ('review','calculate','book_review'):c=post(e,pid,b).json()
    if gate in ('calculate','book_review'):approve(r,pid,'components',c)
    if gate=='book_review':t=post(e,pid,{},'/takeoffs/calculate').json()[-1]
    doc=get(e,pid)['documents'][-1];(store.ROOT/'originals'/doc['hash']).write_bytes(b'corrupted-synthetic-inner-window-original')
    if gate=='preview':post(e,pid,b,'/components/preview',409)
    elif gate=='save':post(e,pid,b,expected=409)
    elif gate=='review':approve(r,pid,'components',c,expected=409)
    elif gate=='calculate':post(e,pid,{},'/takeoffs/calculate',409)
    else:approve(r,pid,'takeoffs',t,expected=409)


@pytest.mark.parametrize('fmt',['json','xlsx','pdf'])
@pytest.mark.parametrize('zero',[False,True])
def test_three_formats_and_old_quote_isolation(clients,fmt,zero):
    from backend.foundation_horizontal_formwork_reports import PERFORATED_TITLES,PERFORATED_NET_TITLE
    e,r,pid,_,q=priced_project(clients);e,r,pid,s,b,eid=setup(clients,pid,zero);c=post(e,pid,b).json();approve(r,pid,'components',c)
    for t in post(e,pid,{},'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
    post(e,pid,{},'/boq/generate');url=f'/api/v1/projects/{pid}/exports';res=e.get(url,params=dict(format=fmt,unpriced=True));assert res.status_code==200,res.text
    titles=set(PERFORATED_TITLES if zero else PERFORATED_TITLES[:2])|{PERFORATED_NET_TITLE}
    if fmt=='json':assert titles<={t[0] for t in res.json()['tables']}
    elif fmt=='xlsx':assert titles<=set(load_workbook(io.BytesIO(res.content)).sheetnames)
    else:assert all(any(t in page.extract_text() for page in PdfReader(io.BytesIO(res.content)).pages) for t in titles)
    assert not any(t[0] in titles for t in e.get(url,params=dict(format='json',quote_id=q['id'])).json()['tables'])
