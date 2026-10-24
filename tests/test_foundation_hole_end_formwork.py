"""Internal void ends, concrete-side ledges and independently checked contact cuts."""
import copy, io, math, random
from decimal import Decimal, localcontext, ROUND_HALF_UP
import pytest
from fastapi import HTTPException
from openpyxl import load_workbook
from pypdf import PdfReader
from backend import store
from backend.engine import DomainError
from backend.compound_takeoff import PI
from backend.foundation_horizontal_formwork import calculate
from backend.foundation_hole_end_formwork import FACES, VERSION
from backend.foundation_horizontal_formwork_reports import END_TITLES, END_CURVED_TITLE, tables
from backend.post_foundation import calculate as foundation_quantity
from backend.railing_material import snapshot
from tests.test_foundation_formwork import foundation
from tests.test_foundation_horizontal_formwork import body as base_body, contact as base_contact, check as base_check, setup as base_setup
from tests.test_foundation_step_formwork import body as step_body
from tests.test_foundation_voids import rectangle
from tests.test_foundation_openings import hole
from tests.test_foundation_recesses import interval
from tests.test_post_foundation import equal_body, fixed_body, fixed_added, equal_added, fixed_adjusted, equal_adjusted, combined
from tests.test_route_takeoff import evaluate
from tests.test_polygon_takeoff import post, mutate
from tests.test_workflow import get, approve, upload, priced_project, create


def contact(eid='test', face='hole_bottom', index=1, name='合成内部孔端'):
    return {**base_contact(eid,face,name),'hole_index':index,'measurement_basis':'仅软件验收：所选孔水平端面实际支模，毛面积按混凝土侧层裁切，扣相接孔，不重复平台或其他模板。'}


def source(shape='rectangular', make=equal_body, holes=None):
    s=foundation(make,shape)
    s['foundation_items'][0]['through_holes']['holes']=holes if holes is not None else [{**rectangle(x='0'),'vertical_range':interval('100','150')}]
    return s


def body(s,eid='test',face='hole_bottom',index=1):
    return {**base_body(s,eid),'name':'合成基础孔端水平模板（非真实工程）','contact_ranges':[contact(eid,face,index)],'measurement_basis':'仅软件验证已审孔端实际水平接触，不由孔体积、土面或投影推断模板；尺寸、起止标高和个数来自已审基础。'}


def check(r):
    base_check(r)
    with localcontext() as ctx:
        ctx.prec=80
        for d in r['contact_ranges']:
            for h in d['apertures']:
                for k,a in h.get('formula_asts',{}).items(): assert evaluate(a,r['dimensions'])==Decimal(h[k])
            for m in [d.get('hole_end_curved_intersection')]+[h.get('curved_intersection') for h in d['apertures']]:
                if not m: continue
                for t in m['terms']:
                    assert evaluate(t['formula_ast'],r['dimensions'])==Decimal(t['area'])
                    for k,a in t['formula_asts'].items(): assert evaluate(a,r['dimensions'])==Decimal(t[k])


@pytest.mark.parametrize('shape',['rectangular','stepped_rectangular','circular','stepped_circular'])
@pytest.mark.parametrize('face',list(FACES))
@pytest.mark.parametrize('make',[equal_body,fixed_body,fixed_added,equal_added,fixed_adjusted,equal_adjusted,combined,lambda s:combined(s,equal=True)])
def test_shapes_both_internal_ends_final_counts_inherited_once_immutable(shape,face,make):
    s=source(shape,make);b=body(s,face=face);before=copy.deepcopy((s,b));r=calculate(b,s);d=r['contact_ranges'][0]
    assert (s,b)==before and r['rule_version']==VERSION and Decimal(d['single_gross_area'])==Decimal('.006') and r['deduction']=='0'
    assert d['concrete_side']==('below' if face=='hole_bottom' else 'above') and d['level_index']==1 and Decimal(d['z'])==Decimal('.100' if face=='hole_bottom' else '.150')
    with localcontext() as ctx: ctx.prec=80;assert Decimal(r['gross'])==Decimal('.006')*Decimal(d['total_count'])
    check(r)


@pytest.mark.parametrize('face,lo,hi,valid',[('hole_bottom','0','150',False),('hole_top','100','600',False),('hole_bottom','0.000000001','150',True),('hole_top','100','599.999999999',True),('hole_bottom','0','600',False),('hole_top','0','600',False)])
def test_external_openings_blocked_but_tiny_internal_distance_not_snapped(face,lo,hi,valid):
    s=source(holes=[{**rectangle(x='0'),'vertical_range':interval(lo,hi)}]);b=body(s,face=face)
    if not valid:
        with pytest.raises(DomainError,match='通向基础外部'):calculate(b,s)
    else:check(calculate(b,s))


@pytest.mark.parametrize('shape',['stepped_rectangular','stepped_circular'])
@pytest.mark.parametrize('kind',['rect','circle'])
def test_roof_at_step_uses_above_footprint_and_complements_platform(shape,kind):
    h=rectangle(x='200') if kind=='rect' else hole(x='190')
    s=source(shape,holes=[{**h,'vertical_range':interval('0','200')}]);r=calculate(body(s,face='hole_top'),s);platform=calculate(step_body(s),s);d=r['contact_ranges'][0];cut=platform['contact_ranges'][0]['apertures'][0]
    assert d['level_index']==2 and Decimal(d['z'])==Decimal('.200') and d['single_gross_area']==cut['covered_area'] and d['hole_end_intersection_rule']==cut['intersection_rule'].replace('上层覆盖','混凝土侧层截面')
    with localcontext() as ctx:ctx.prec=80;assert Decimal(d['single_gross_area'])+Decimal(cut['single_area'])==Decimal(cut['full_area'])
    check(r)


def test_bottom_at_step_uses_below_layer_and_equal_footprint_layer_boundary():
    s=source('stepped_rectangular',holes=[{**rectangle(x='0'),'vertical_range':interval('200','400')}]);r=calculate(body(s),s);d=r['contact_ranges'][0];assert d['level_index']==1 and Decimal(d['z'])==Decimal('.200') and Decimal(d['single_gross_area'])==Decimal('.006');check(r)
    s['foundation_items'][0]['levels'][1]['parameters'].update(length='600',width='500');s['foundation_items'][0]['through_holes']['holes'][0]['vertical_range']=interval('0','200');r=calculate(body(s,face='hole_top'),s);assert r['contact_ranges'][0]['level_index']==2;check(r)


@pytest.mark.parametrize('x,expected',[('250','0'),('200','.003'),('0','.006')])
def test_roof_outside_partial_or_inside_upper_footprint(x,expected):
    s=source('stepped_rectangular',holes=[{**rectangle(x=x),'vertical_range':interval('0','200')}]);r=calculate(body(s,face='hole_top'),s);assert Decimal(r['contact_ranges'][0]['single_gross_area'])==Decimal(expected);check(r)


@pytest.mark.parametrize('face',list(FACES))
@pytest.mark.parametrize('offset,expected',[('0','0'),('40','.0024'),('100','.006'),('140','.006')])
def test_vertically_meeting_rectangles_only_actual_intersection_and_zero_contact(face,offset,expected):
    a=interval('100','300') if face=='hole_top' else interval('300','500');b=interval('300','500') if face=='hole_top' else interval('100','300')
    s=source(holes=[{**rectangle(x='10'),'vertical_range':a},{**rectangle(name='相接H2',x=str(Decimal(10)+Decimal(offset))),'vertical_range':b}]);r=calculate(body(s,face=face),s);d=r['contact_ranges'][0];assert len(d['apertures'])==1 and Decimal(d['single_net_area'])==Decimal(expected);check(r)


@pytest.mark.parametrize('face',list(FACES))
@pytest.mark.parametrize('kind',['circle-circle','circle-rect','rect-circle'])
def test_shifted_partial_curved_adjacent_holes_independent_lens_or_half_disk(face,kind):
    a=interval('100','300') if face=='hole_top' else interval('300','500');b=interval('300','500') if face=='hole_top' else interval('100','300')
    h1=hole(x='-10',diameter='40') if kind!='rect-circle' else rectangle(x='-10',length='40',width='60')
    h2=hole(name='相接H2',x='10',diameter='40') if kind!='circle-rect' else rectangle(name='相接H2',x='10',length='40',width='60')
    s=source(holes=[{**h1,'vertical_range':a},{**h2,'vertical_range':b}]);r=calculate(body(s,face=face),s);d=r['contact_ranges'][0];h=d['apertures'][0];assert h['curved_intersection']
    if kind=='circle-circle':expected=.02**2*(2*math.pi/3-math.sqrt(3)/2)
    else:expected=math.pi*.02**2/2
    assert float(h['single_area'])==pytest.approx(expected,rel=1e-13)
    assert all('上层圆弓形'!=t['name'] for t in h['curved_intersection']['terms']);check(r)


def test_multiple_disjoint_same_side_cuts_and_wrong_side_holes_not_deducted():
    s=source(holes=[{**rectangle(x='0'),'vertical_range':interval('300','500')},{**rectangle(name='左侧相接',x='-25',length='50'),'vertical_range':interval('100','300')},{**rectangle(name='右侧相接',x='25',length='50'),'vertical_range':interval('200','300')}]);r=calculate(body(s),s);assert r['quantity']=='0.000' and len(r['contact_ranges'][0]['apertures'])==2;check(r)
    # A hole entirely on the void side cannot create a cut at the floor.
    s=source(holes=[{**rectangle(x='-60'),'vertical_range':interval('300','500')},{**rectangle(name='另一同侧孔',x='60'),'vertical_range':interval('300','500')}]);r=calculate(body(s),s);assert not r['contact_ranges'][0]['apertures'];check(r)


@pytest.mark.parametrize('unit,factor',[('mm','1'),('cm','.1'),('dm','.01'),('m','.001'),('km','.000001')])
def test_independent_hole_units_source_units_and_low_context(unit,factor):
    s=source('stepped_circular',holes=[{**hole(x='190'),'vertical_range':interval('0','200')}]);expected=calculate(body(s,face='hole_top'),s);item=s['foundation_items'][0];item['length_unit']=unit
    for level in item['levels']:
        for k in level['parameters']:level['parameters'][k]=str(Decimal(level['parameters'][k])*Decimal(factor))
    with localcontext() as ctx:ctx.prec=6;r=calculate(body(s,face='hole_top'),s)
    assert r['quantity']==expected['quantity'] and r['gross']==expected['gross'];check(r)
    c=item['through_holes'];c['unit']=unit
    for h in c['holes']:
        for p in (h['parameters'],h['vertical_range']['parameters']):
            for k in p:p[k]=str(Decimal(p[k])*Decimal(factor))
    r=calculate(body(s,face='hole_top'),s);assert r['quantity']==expected['quantity'];check(r)


def test_160_seeded_translated_rectangles_independent_integer_overlap_and_rounding():
    rng=random.Random(10201)
    for _ in range(160):
        l,w=rng.randint(2,110),rng.randint(2,100);a,b=rng.randint(-80,80),rng.randint(-80,80);x,y=rng.randint(-80,80),rng.randint(-80,80)
        hs=[{**rectangle(x=str(a),y=str(b),length=str(l),width=str(w)),'vertical_range':interval('100','300')},{**rectangle(name='合成相接孔',x=str(x),y=str(y),length='100',width='60'),'vertical_range':interval('300','500')}]
        s=source(holes=hs);c=body(s,face='hole_top');c['rounding_digits']=rng.randint(0,12);r=calculate(c,s)
        with localcontext() as ctx:
            ctx.prec=80;D=Decimal;overlap=max(D(0),min(D(2*a+l),D(2*x+100))-max(D(2*a-l),D(2*x-100)))*max(D(0),min(D(2*b+w),D(2*y+60))-max(D(2*b-w),D(2*y-60)))/4;net=(D(l*w)-overlap)/1000000*22;assert D(r['quantity'])==net.quantize(D(1).scaleb(-c['rounding_digits']),rounding=ROUND_HALF_UP)
        check(r)


@pytest.mark.parametrize('fault',['missing-hole','bool-hole','unknown-hole','float-hole','str-hole','list-hole','dict-hole','manual-level','manual-z','manual-area','manual-profile','duplicate','top-with-hole','missing-face'])
def test_strict_hole_end_schema(fault):
    s=source();b=body(s);d=b['contact_ranges'][0]
    if fault=='missing-hole':d.pop('hole_index')
    elif fault=='missing-face':d.pop('face')
    elif fault in ('bool-hole','unknown-hole','float-hole','str-hole','list-hole','dict-hole'):d['hole_index']={'bool-hole':True,'unknown-hole':99,'float-hole':1.0,'str-hole':'1','list-hole':[],'dict-hole':{}}[fault]
    elif fault=='manual-level':d['level_index']=1
    elif fault=='top-with-hole':d['face']='top'
    elif fault=='duplicate':b['contact_ranges'].append({**d,'name':'重复孔端'})
    else:d[fault]='1'
    with pytest.raises(DomainError):calculate(b,s)


def setup(clients,pid=None):
    e,r,pid,s,_,_=base_setup(clients,pid,'stepped_circular')
    upload(e,pid,'仅合成软件验收：22个同心圆台阶基础，下Φ600×200，上Φ400×400mm。H1圆孔Φ40，中心(190,0)mm，Z0至200mm。孔顶Z200mm取上侧L2圆截面交集.001001962168042696…m²/个，22个.0220431676969393…m²，修约3位.022m²。H2矩形孔100×60，中心(-80,0)mm，Z300至500mm，孔底孔顶各.006m²/个。三面合计.286043167…→.286m²。体积阈值100m²不扣孔体积，材料79.200kg不变。孔端实际支模、施工与修约仅合成，不代表真实工法。TEST-HOLE-END仅软件映射码。','foundation-hole-end-basis.txt')
    for component in get(e,pid)['components']:approve(r,pid,'components',component)
    s=next(c for c in get(e,pid)['components'] if c['id']==s['id']);eid=get(e,pid)['evidence'][-1]['id'];s['object_hash']=store.digest(s);c=s['foundation_items'][0]['through_holes'];c['deduction_rule']['threshold']='100';c['holes']=[{**hole(eid=eid,x='190'),'vertical_range':interval('0','200',eid)},{**rectangle(eid=eid,name='合成H2',x='-80'),'vertical_range':interval('300','500',eid)}]
    res=e.put(f'/api/v1/projects/{pid}/components/'+s['id'],json=s);assert res.status_code==200,res.text;approve(r,pid,'components',res.json());s=next(c for c in get(e,pid)['components'] if c['id']==s['id']);return e,r,pid,s,body(s,eid,'hole_top'),eid


def test_candidates_preview_claim_identity_and_distinct_holes_no_mutation(clients):
    e,r,pid,s,b,eid=setup(clients);before=get(e,pid);res=e.get(f'/api/v1/projects/{pid}/foundation-horizontal-formwork-candidates',params=dict(component_id=s['id']));assert res.status_code==200,res.text;faces=res.json()['foundation_items'][0]['faces'];ends=[f for f in faces if f['face'] in FACES];assert len(ends)==4 and ends[0]['available'] is False and all(f['available'] for f in ends[1:]);assert ends[1]['hole_end_curved_intersection'] and ends[1]['level_index']==2
    assert post(e,pid,b,'/components/preview').json()['quantity']=='0.022' and get(e,pid)==before;c=post(e,pid,b).json();approve(r,pid,'components',c);post(e,pid,{**b,'id':c['id']},expected=409)
    b['contact_ranges']=[contact(eid,'hole_top',2)];post(e,pid,b);b['contact_ranges']=[contact(eid,'hole_bottom',2)];post(e,pid,b);post(e,pid,base_body(s,eid));post(e,pid,step_body(s,eid))


@pytest.mark.parametrize('fault',['index','side','layer','z','gross','full','trace','formula','source','book'])
def test_saved_end_geometry_reconstruction_tamper_gates(clients,fault):
    e,r,pid,s,b,eid=setup(clients);c=post(e,pid,b).json()
    if fault=='book':approve(r,pid,'components',c);post(e,pid,{},'/takeoffs/calculate')
    def forge(p):
        c=p['components'][-1];d=c['preview']['contact_ranges'][0]
        if fault=='index':d['hole_index']=2
        elif fault=='side':d['concrete_side']='below'
        elif fault=='layer':d['level_index']=1
        elif fault=='z':d['z']='0.201'
        elif fault=='gross':d['single_gross_area']='0'
        elif fault=='full':d['single_full_hole_area']='0'
        elif fault=='trace':d['hole_end_curved_intersection']['terms'][0]['area']='0'
        elif fault=='formula':d['formula_asts']['z']={'constant':'0'}
        elif fault=='source':c['source_geometry']['foundation_items'][0]['through_holes']['holes'][0]['parameters']['x']='170'
        else:p['takeoffs'][-1]['quantity']='999'
    mutate(pid,forge);coll='takeoffs' if fault=='book' else 'components';approve(r,pid,coll,get(e,pid)[coll][-1],expected=409)


@pytest.mark.parametrize('gate',['preview','save','review','calculate','book_review'])
def test_original_integrity_five_gates(clients,gate):
    e,r,pid,s,b,eid=setup(clients)
    if gate in ('review','calculate','book_review'):c=post(e,pid,b).json()
    if gate in ('calculate','book_review'):approve(r,pid,'components',c)
    if gate=='book_review':t=post(e,pid,{},'/takeoffs/calculate').json()[-1]
    doc=get(e,pid)['documents'][-1];(store.ROOT/'originals'/doc['hash']).write_bytes(b'corrupt')
    if gate=='preview':post(e,pid,b,'/components/preview',409)
    elif gate=='save':post(e,pid,b,expected=409)
    elif gate=='review':approve(r,pid,'components',c,expected=409)
    elif gate=='calculate':post(e,pid,{},'/takeoffs/calculate',409)
    else:approve(r,pid,'takeoffs',t,expected=409)


def test_revision_history_same_end_source_stale_and_permissions(clients):
    e,r,pid,s,b,eid=setup(clients);post(r,pid,b,expected=403);assert clients['other'].get(f'/api/v1/projects/{pid}/foundation-horizontal-formwork-candidates',params=dict(component_id=s['id'])).status_code==404
    foreign=create(e)['id'];upload(e,foreign,'异项目孔端依据');bad=copy.deepcopy(b);bad['contact_ranges'][0]['evidence']['measurement_basis']=get(e,foreign)['evidence'][-1]['id'];post(e,pid,bad,expected=404)
    c=post(e,pid,b).json();approve(r,pid,'components',c)
    for t in post(e,pid,{},'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
    post(e,pid,{},'/boq/generate');before=get(e,pid);old=before['components'][-1];b.update(object_hash=store.digest(old));b['contact_ranges'] += [contact(eid,'hole_bottom',2,'合成第二孔底'),contact(eid,'hole_top',2,'合成第二孔顶')];url=f'/api/v1/projects/{pid}/components/'+old['id'];res=e.put(url,json=b);assert res.status_code==200,res.text;after=get(e,pid);assert res.json()['preview']['quantity']=='0.286' and after['component_history'][-1]==old and after['boq']==before['boq'] and all(t['state']=='stale' for t in after['takeoffs']) and e.put(url,json=b).status_code==409;approve(r,pid,'components',res.json())
    s.update(object_hash=store.digest(s));s['foundation_items'][0]['through_holes']['holes'][0]['parameters']['x']='170';res=e.put(f'/api/v1/projects/{pid}/components/'+s['id'],json=s);assert res.status_code==200,res.text;approve(r,pid,'components',res.json());assert get(e,pid)['components'][-1]['state']=='stale'


@pytest.mark.parametrize('fmt',['json','xlsx','pdf'])
def test_three_formats_end_trace_tables_exact_quantity_old_quote_excluded(clients,fmt):
    e,r,pid,_,q=priced_project(clients);e,r,pid,s,b,eid=setup(clients,pid);c=post(e,pid,b).json();approve(r,pid,'components',c)
    for t in post(e,pid,{},'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
    boq=post(e,pid,{},'/boq/generate').json();assert any(d['quantity']=='0.022' and d['unit']=='m²' for d in boq);url=f'/api/v1/projects/{pid}/exports';res=e.get(url,params=dict(format=fmt,unpriced=True));assert res.status_code==200,res.text;titles=set(END_TITLES)|{END_CURVED_TITLE}
    if fmt=='json':assert titles<={d[0] for d in res.json()['tables']}
    elif fmt=='xlsx':assert titles<=set(load_workbook(io.BytesIO(res.content)).sheetnames)
    else:assert all(v in ''.join(p.extract_text() for p in PdfReader(io.BytesIO(res.content)).pages) for v in titles)
    assert not any(d[0] in titles for d in e.get(url,params=dict(format='json',quote_id=q['id'])).json()['tables'])
    assert tables(dict(takeoffs=[]))==[]
