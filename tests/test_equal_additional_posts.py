"""Synthetic construction-post additions; no real engineering acceptance."""
import copy,io,random
from decimal import Decimal,localcontext
import pytest
from openpyxl import load_workbook
from pypdf import PdfReader
from backend import store
from backend.engine import DomainError
from backend.railing_material import calculate
from backend.railing_takeoff import snapshot
from tests.test_railing_posts import body as fixed_body,setup as fixed_setup
from tests.test_fixed_posts import stations
from tests.test_railing_posts import geometry
from tests.test_route_railing import payload as railing_payload
from tests.test_route_takeoff import point,evaluate
from tests.test_railing_takeoff import interval
from tests.test_polygon_takeoff import post,mutate
from tests.test_workflow import get,approve,upload,priced_project,create


def addition(s='1500',unit='mm',eid='test',label='合成支撑构造柱'):
 return dict(label=label,s=str(s),s_unit=unit,measurement_basis='按合成原件独立指定原段端距，加柱不移原均分柱',evidence=dict(s=eid,measurement_basis=eid))


def body(s,eid='test'):
 b=fixed_body(s,eid);m=b['material_items'][0];m.update(mode='spaced_added_posts',name='合成均分及构造加柱')
 for i,r in enumerate(m['post_ranges']):r['additional_posts']=[addition('1500' if i==0 else '2500',eid=eid)]
 b['measurement_basis']='原均分11根，加AB S1.5m与BC S2.5m两根；13×1.2×3×来源2=93.6kg，非真实工程';return b


def setup(clients,pid=None):
 e,r,pid,s,_,_=fixed_setup(clients,pid)
 upload(e,pid,'合成构造加柱验收，非真实工程：原均分AB13m最大间距2m分7跨共8根，BC5m最大间距2m分3跨共3根，共享B只归AB。追加同规格构造柱：AB原段S1500mm一根、BC原段S2500mm一根。不得移动原均分柱或按加柱重新分段均分。每根1200mm、3000g/m，单组13根，总93.6kg，修约3位。修订移除BC加柱后12根、总86.4kg。端部只能由两端归属决定，重合及越界拒绝。','equal-additional-posts-basis.txt')
 for c in get(e,pid)['components']:approve(r,pid,'components',c)
 s=get(e,pid)['components'][-1];eid=get(e,pid)['evidence'][-1]['id'];return e,r,pid,s,body(s,eid),eid


def line_source(length):
 s=geometry();src=s['source_geometry'];src['route_points']=[point('A',0,0,0),point('B',length,0,0)];s['source_component_hash']=store.digest(snapshot(src));s['edge_decisions']=railing_payload(src)['edge_decisions'];return s


def test_added_positions_counts_mass_ast_and_preserved_base():
 s=geometry();b=body(s);before=copy.deepcopy((b,s));v=calculate(b,s);assert (b,s)==before
 assert v['quantity']=='93.600' and v['rule_version']=='railing-material-mass-1.6'
 g=v['material_items'][0]['post_layout'];base=calculate(fixed_body(s),s)['material_items'][0]['post_layout']
 assert (g['base_post_count'],g['additional_post_count'],g['post_count'])==(11,2,13)
 for label,extra in [('S1 整段',Decimal('1.5')),('S2 整段',Decimal('2.5'))]:assert stations(g,label)==sorted(stations(base,label)+[extra])
 assert [p['index'] for p in g['post_stations']]==list(range(1,14))
 for old,new in zip(base['post_ranges'],g['post_ranges']):
  for k in ('interval_count','actual_spacing','ceil_certificate'):assert old[k]==new[k]
 with localcontext() as c:
  c.prec=80;assert evaluate(v['formula_ast'],v['dimensions'])==Decimal('93.6')
  for p in g['post_stations']:assert evaluate(p['station_formula_ast'],v['dimensions'])==Decimal(p['s'])
  for r in g['post_ranges']:assert evaluate(r['formula_ast'],v['dimensions'])==r['post_count']
 for p in base['post_stations']:
  assert any(all(q.get(k)==p[k] for k in p if k!='index') for q in g['post_stations'])


@pytest.mark.parametrize('basis',['spatial','horizontal'])
@pytest.mark.parametrize('unit,s',[('m','1.5'),('mm','1500'),('cm','150'),('dm','15'),('km','.0015')])
def test_units_basis_and_direction_independent_s(basis,unit,s):
 src=geometry(basis=basis);b=body(src)
 for r in b['material_items'][0]['post_ranges']:r['additional_posts']=[addition(s,unit)]
 v=calculate(b,src);g=v['material_items'][0]['post_layout'];assert v['quantity']==('93.600' if basis=='spatial' else '64.800')
 assert [p['s'] for p in g['post_stations'] if p['position_kind']=='additional']==['1.5','1.5']


@pytest.mark.parametrize('fault',['missing_list','not_list','extra_field','missing_s','negative','zero','endpoint','beyond','nan','bool','bad_unit','blank_label','blank_note','missing_proof','extra_proof','blank_proof','duplicate_s','same_s_units','duplicate_label','coincident_start','coincident_end','manual_count','no_additions','bad_range_id'])
def test_invalid_added_posts(fault):
 s=geometry();b=body(s);m=b['material_items'][0];r=m['post_ranges'][0];a=r['additional_posts'][0]
 if fault=='missing_list':r.pop('additional_posts')
 elif fault=='not_list':r['additional_posts']={}
 elif fault=='extra_field':a['quantity']='2'
 elif fault=='missing_s':a.pop('s')
 elif fault in ('negative','zero','endpoint','beyond','nan','bool'):a['s']={'negative':'-1','zero':'0','endpoint':'13000','beyond':'14000','nan':'NaN','bool':True}[fault]
 elif fault=='bad_unit':a['s_unit']='ft'
 elif fault=='blank_label':a['label']=' '
 elif fault=='blank_note':a['measurement_basis']=' '
 elif fault=='missing_proof':a['evidence'].pop('s')
 elif fault=='extra_proof':a['evidence']['quantity']='test'
 elif fault=='blank_proof':a['evidence']['s']=''
 elif fault in ('duplicate_s','same_s_units','duplicate_label'):
  r['additional_posts'].append(addition('1500' if fault=='duplicate_s' else '1.5' if fault=='same_s_units' else '1600',unit='m' if fault=='same_s_units' else 'mm',label='其他' if fault!='duplicate_label' else a['label']))
 elif fault=='coincident_start':r['spacing']='1000';a['s']='2000'
 elif fault=='coincident_end':m['post_ranges'][1]['spacing']='1000';m['post_ranges'][1]['additional_posts'][0]['s']='3000'
 elif fault=='manual_count':m['quantity']='13'
 elif fault=='no_additions':
  for p in m['post_ranges']:p['additional_posts']=[]
 else:r['range_id']=[]
 with pytest.raises(DomainError):calculate(b,s)


def test_local_range_gaps_and_original_s():
 s=geometry();s['edge_decisions'][0].update(mode='ranges',ranges=[interval(4,10)]);b=body(s);r=b['material_items'][0]['post_ranges'][0];r['additional_posts']=[addition('5500')];v=calculate(b,s);assert Decimal('5.5') in stations(v['material_items'][0]['post_layout'],'S1 区间1')
 for value in ('1.5','4','10','11'):
  r['additional_posts']=[addition(value,'m')]
  with pytest.raises(DomainError,match='严格位于'):calculate(b,s)


def test_exact_equal_collision_and_nearby_irrational_root():
 s=line_source(10);b=body(s);r=b['material_items'][0]['post_ranges'][0];r['additional_posts']=[addition('8000')]
 with pytest.raises(DomainError,match='重合'):calculate(b,s)
 src=s['source_geometry'];src['route_points']=[point('A',0,0,0),point('B',1,1,0)];s['source_component_hash']=store.digest(snapshot(src));s['edge_decisions']=railing_payload(src)['edge_decisions'];b=body(s);r=b['material_items'][0]['post_ranges'][0];r.update(spacing='1',spacing_unit='m');r['additional_posts']=[addition('.707106781186547524','m')]
 with localcontext() as c:c.prec=8;g=calculate(b,s)['material_items'][0]['post_layout']
 assert g['post_count']==4 and len({p['s'] for p in g['post_stations']})==4
 r['additional_posts']=[addition('1.41421356237309505','m')]
 with pytest.raises(DomainError,match='严格位于'):calculate(b,s)


def test_added_per_range_item_and_table_limits():
 s=geometry();b=body(s);r=b['material_items'][0]['post_ranges'][0];r['additional_posts']=[addition(str(i),'cm',label='柱'+str(i)) for i in range(1,101)];assert calculate(b,s)['material_items'][0]['post_layout']['additional_post_count']==101
 r['additional_posts'].append(addition('101','cm',label='柱101'))
 with pytest.raises(DomainError,match='最多100'):calculate(b,s)
 s=line_source(998);b=body(s);r=b['material_items'][0]['post_ranges'][0];r.update(spacing='1',spacing_unit='m');r['additional_posts']=[addition('.5','m')];assert calculate(b,s)['material_items'][0]['post_layout']['post_count']==1000
 r['additional_posts'].append(addition('1.5','m',label='第二'))
 with pytest.raises(DomainError,match='最多1000'):calculate(b,s)
 r['additional_posts'].pop();b['material_items']=[dict(copy.deepcopy(b['material_items'][0]),name='材料'+str(i)) for i in range(5)];assert calculate(b,s)['quantity']=='36000.000'
 b['material_items'].append(dict(copy.deepcopy(b['material_items'][0]),name='第六项'))
 with pytest.raises(DomainError,match='最多5000'):calculate(b,s)


def test_seeded_independent_equal_station_oracle():
 from fractions import Fraction
 rng=random.Random(7301)
 for _ in range(35):
  L=rng.randint(5,90);d=rng.randint(2,8);s=line_source(L);b=body(s);r=b['material_items'][0]['post_ranges'][0];r.update(spacing=str(d),spacing_unit='m');n=(L+d-1)//d
  base={Fraction(L*j,n) for j in range(n+1)};extra=sorted({Fraction(rng.randint(0,L-1))+Fraction(123,1000) for _ in range(3)});r['additional_posts']=[addition(str(Decimal(v.numerator)/Decimal(v.denominator)),'m',label=str(i)) for i,v in enumerate(extra)]
  g=calculate(b,s)['material_items'][0]['post_layout'];actual=stations(g,'S1 整段');expected=sorted(base|set(extra));assert len(actual)==len(expected)
  with localcontext() as c:
   c.prec=80
   for a,e in zip(actual,expected):assert abs(a-Decimal(e.numerator)/Decimal(e.denominator))<Decimal('1e-75')


@pytest.mark.parametrize('fmt',['json','xlsx','pdf'])
def test_workflow_exports_and_old_quote(clients,fmt):
 e,r,pid,_,quote=priced_project(clients);e,r,pid,s,b,eid=setup(clients,pid)
 if fmt=='pdf':b['name']='材料表'+('长名称验收'*39)[:197];b['material_items'][0]['name']='立柱'+('长名称验收'*40)[:198]
 before=get(e,pid);post(e,pid,b,'/components/preview');assert get(e,pid)==before
 c=post(e,pid,b).json();approve(r,pid,'components',c)
 for t in post(e,pid,{},'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
 assert any(t['quantity']=='93.600' for t in post(e,pid,{},'/boq/generate').json())
 url=f'/api/v1/projects/{pid}/exports';res=e.get(url,params=dict(format=fmt,unpriced=True));assert res.status_code==200,res.text;titles=['构造加柱数量组成','构造加柱端距及依据','立柱逐根端距']
 if fmt=='json':assert set(titles)<={v[0] for v in res.json()['tables']}
 elif fmt=='xlsx':assert set(titles)<=set(load_workbook(io.BytesIO(res.content)).sheetnames)
 else:
  pages=[p.extract_text() for p in PdfReader(io.BytesIO(res.content)).pages];assert all(v in ''.join(pages) for v in titles)
  page_words=[''.join(p.split()) for p in pages]
  for station in c['preview']['material_items'][0]['post_layout']['post_stations']:
   assert any(station['s'] in p for p in page_words),'一根柱的完整端距不能拆到两页'
 old=e.get(url,params=dict(format='json',quote_id=quote['id'])).json();assert old['quote']==quote and not any(t[0] in titles for t in old['tables'])


@pytest.mark.parametrize('fault',['s','addition_record','station','count','formula','base_position','source','book'])
def test_saved_records_rebuild(clients,fault):
 e,r,pid,s,b,eid=setup(clients);c=post(e,pid,b).json()
 if fault=='book':approve(r,pid,'components',c);post(e,pid,{},'/takeoffs/calculate')
 def forge(p):
  c=p['components'][-1];g=c['preview']['material_items'][0]['post_layout']
  if fault=='s':c['material_items'][0]['post_ranges'][0]['additional_posts'][0]['s']='1600'
  elif fault=='addition_record':g['post_ranges'][0]['additional_posts'][0]['normalized_s']='1.6'
  elif fault=='station':g['post_stations'][1]['s']='1.6'
  elif fault=='count':g['additional_post_count']=99
  elif fault=='formula':g['formula_ast']={'constant':'13'}
  elif fault=='base_position':g['post_stations'][2]['s']='2.1'
  elif fault=='source':c['source_geometry']['edge_decisions'][0]['mode']='exclude'
  else:p['takeoffs'][-1]['material_items'][0]['post_layout']['additional_post_count']=99
 mutate(pid,forge);coll='takeoffs' if fault=='book' else 'components';approve(r,pid,coll,get(e,pid)[coll][-1],expected=409)


def test_revision_history_concurrency_and_source(clients):
 e,r,pid,s,b,eid=setup(clients);c=post(e,pid,b).json();approve(r,pid,'components',c)
 for t in post(e,pid,{},'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
 post(e,pid,{},'/boq/generate');old=get(e,pid)['components'][-1];new=copy.deepcopy(b);new['material_items'][0]['post_ranges'][1]['additional_posts']=[];new['object_hash']=store.digest(old);url=f'/api/v1/projects/{pid}/components/'+c['id'];res=e.put(url,json=new);assert res.status_code==200,res.text
 p=get(e,pid);assert p['component_history'][-1]==old and all(t['state']=='stale' for t in p['takeoffs']) and any(v['quantity']=='93.600' for v in p['boq']);assert res.json()['preview']['quantity']=='86.400' and e.put(url,json=new).status_code==409;approve(r,pid,'components',res.json())
 for t in post(e,pid,{},'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
 assert any(v['quantity']=='86.400' for v in post(e,pid,{},'/boq/generate').json())
 s['object_hash']=store.digest(s);s['edge_decisions'][0].update(mode='ranges',ranges=[interval(0,10,eid)]);res=e.put(f'/api/v1/projects/{pid}/components/'+s['id'],json=s);assert res.status_code==200,res.text;approve(r,pid,'components',res.json());assert get(e,pid)['components'][-1]['state']=='stale';post(e,pid,b,'/components/preview',409)


@pytest.mark.parametrize('gate',['preview','save','review','calculate','book_review'])
def test_original_bytes_gates(clients,gate):
 e,r,pid,s,b,eid=setup(clients)
 if gate in ('review','calculate','book_review'):c=post(e,pid,b).json()
 if gate in ('calculate','book_review'):approve(r,pid,'components',c)
 if gate=='book_review':t=post(e,pid,{},'/takeoffs/calculate').json()[-1]
 doc=get(e,pid)['documents'][-1];(store.ROOT/'originals'/doc['hash']).write_bytes(b'corrupt')
 if gate=='preview':post(e,pid,b,'/components/preview',409)
 elif gate=='save':post(e,pid,b,expected=409)
 elif gate=='review':approve(r,pid,'components',c,expected=409)
 elif gate=='calculate':post(e,pid,{},'/takeoffs/calculate',409)
 else:approve(r,pid,'takeoffs',t,expected=409)


def test_roles_tenant_and_foreign_s_evidence(clients):
 e,r,pid,s,b,eid=setup(clients);post(r,pid,b,expected=403);assert clients['other'].get(f'/api/v1/projects/{pid}').status_code==404
 other=create(e)['id'];upload(e,other,'其他工程原件');b['material_items'][0]['post_ranges'][0]['additional_posts'][0]['evidence']['s']=get(e,other)['evidence'][-1]['id'];post(e,pid,b,expected=404)


def test_mixed_modes_have_separate_parameters_and_sum_asts():
 s=geometry();b=body(s);b['material_items'].append(fixed_body(s)['material_items'][0]);v=calculate(b,s)
 assert v['quantity']=='172.800' and v['material_items'][1]['post_layout']['rule_version']=='railing-equal-bay-posts-1.0'
 with localcontext() as c:
  c.prec=80;assert evaluate(v['formula_ast'],v['dimensions'])==Decimal('172.8')
  for item in v['material_items']:assert evaluate(item['formula_ast'],v['dimensions'])==Decimal(item['single_group_mass'])


@pytest.mark.parametrize('starts',[True,False])
def test_closed_polygon_and_hole_additions_keep_endpoint_ownership(starts):
 from tests.test_railing_takeoff import geometry as pg,payload as pp
 src=pg();s=pp(src);s.update(id='polygon',revision=1,source_geometry=snapshot(src))
 for d in s['edge_decisions']:d.update(mode='full',ranges=[])
 b=body(s)
 for i,r in enumerate(b['material_items'][0]['post_ranges']):r.update(start_post=starts,end_post=not starts,additional_posts=[addition('.1','m',label='加柱'+str(i))])
 v=calculate(b,s);g=v['material_items'][0]['post_layout'];assert g['base_post_count']==15 and g['additional_post_count']==6 and g['post_count']==21
 b['material_items'][0]['post_ranges'][0]['end_post' if starts else 'start_post']=True
 with pytest.raises(DomainError,match='共享端点'):calculate(b,s)


def test_max_spacing_is_not_the_actual_equal_bay_spacing():
 s=line_source(12);b=body(s);r=b['material_items'][0]['post_ranges'][0];r.update(spacing='5',spacing_unit='m',additional_posts=[addition('4','m')])
 with pytest.raises(DomainError,match='均分柱重合'):calculate(b,s)
 r['additional_posts']=[addition('4.00000000000000001','m')];v=calculate(b,s);g=v['material_items'][0]['post_layout'];assert g['base_post_count']==4 and g['post_count']==5
 assert stations(g,'S1 整段')==[0,4,Decimal('4.00000000000000001'),8,12]


def test_partial_equal_collision_uses_original_s_offset():
 s=geometry();s['edge_decisions'][0].update(mode='ranges',ranges=[interval(2,8)]);b=body(s);r=b['material_items'][0]['post_ranges'][0];r['additional_posts']=[addition('4','m')]
 with pytest.raises(DomainError,match='均分柱重合'):calculate(b,s)
 r['additional_posts']=[addition('4.5','m')];g=calculate(b,s)['material_items'][0]['post_layout'];assert stations(g,'S1 区间1')==[2,4,Decimal('4.5'),6,8]


@pytest.mark.parametrize('location',['parameter','evidence'])
def test_fixed_direction_fields_cannot_leak_into_equal_mode(location):
 s=geometry();b=body(s);r=b['material_items'][0]['post_ranges'][0]
 (r if location=='parameter' else r['evidence'])['direction']='from_start' if location=='parameter' else 'test'
 with pytest.raises(DomainError):calculate(b,s)


def test_equal_and_fixed_additions_can_coexist_with_separate_dimensions():
 from tests.test_additional_posts import body as fixed_added_body
 s=geometry();b=body(s);b['material_items']+=fixed_added_body(s)['material_items'];v=calculate(b,s);assert v['quantity']=='187.200'
 assert [i['post_layout']['rule_version'] for i in v['material_items']]==['railing-equal-added-posts-1.0','railing-fixed-added-posts-1.0']
 with localcontext() as c:
  c.prec=80;assert evaluate(v['formula_ast'],v['dimensions'])==Decimal('187.2')
