"""Synthetic sloped railing bases and evidence gates, not professional acceptance."""
import copy,io,json,random
from decimal import Decimal,localcontext
import pytest
from fastapi import HTTPException
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pypdf import PdfReader
from backend import store
from backend.main import app
from backend.engine import DomainError
from backend.railing_takeoff import calculate,boundary,snapshot
from backend.mapping_types import component_profile
from tests.test_railing_takeoff import interval
from tests.test_route_takeoff import payload as route_payload,point,evaluate
from tests.test_polygon_takeoff import post,mutate
from tests.test_workflow import create,get,upload,approve,priced_project


def geometry(eid='test'):
 s=route_payload(eid);s.update(id='synthetic-sloped-line',revision=1,name='合成坡段坐标线（非真实工程）',multiplicity='2',rounding_digits=0)
 s['route_points']=[point('A',0,0,0,eid),point('B',3,4,12,eid),point('C',6,8,12,eid)]
 return s


def payload(s,eid='test',basis='spatial'):
 es=boundary(s,basis)['railing_edges']
 return dict(name='合成坡段栏杆（非真实工程）',kind='polygon_railing',parameters={},quantity_type='length',source_component_id=s['id'],source_component_hash=store.digest(snapshot(s)),
 length_basis=basis,length_basis_evidence_id=eid,edge_decisions=[dict(edge_id=e['edge_id'],mode='full',ranges=[],basis='合成每原段整段计入，按选定三维或水平口径，非专业核定',evidence_id=eid) for e in es],
 scope_confirmed=True,material='合成栏杆材料做法，无实际计价效力',material_evidence_id=eid,installation_basis='合成确认借用设计坐标线的几何范围为本栏杆中心线，不从管线用途推断安装做法',installation_evidence_id=eid,
 measurement_basis='合成整段三维13+5=18m，水平5+5=10m，重复2；局部S只按所选口径，修约3位',basis_evidence_id=eid,rounding_digits=3,rounding_evidence_id=eid)


def setup(clients,pid=None,basis='spatial'):
 e,r=clients['estimator'],clients['reviewer'];pid=pid or create(e)['id']
 upload(e,pid,'合成软件验收：共同设计XYZ，A(0,0,0)、B(3,4,12)、C(6,8,12)m，重复2；AB三维13m水平5m，BC三维5m水平5m。明确本样例栏杆沿同一坐标线安装，材料做法仅合成，不将管线用途自动认定为栏杆；整段三维总36m，水平总20m。局部三维S0→2及4→10m、BC整5m总26m；水平AB S0→1及2→4m、BC整5m总16m，端距依据随口径核定。来源改B.Z为9m时须重新核定范围。栏杆仅总长修约3位，非专业签审。','slope-railing-basis.txt')
 eid=get(e,pid)['evidence'][-1]['id'];s=geometry(eid);s.pop('id');s.pop('revision');post(e,pid,s)
 for current in get(e,pid)['components']:approve(r,pid,'components',current)
 s=get(e,pid)['components'][-1];return e,r,pid,s,payload(s,eid,basis),eid


@pytest.mark.parametrize('basis,total,raws',[('spatial',36,[13,5]),('horizontal',20,[5,5])])
def test_full_sloped_and_horizontal_bases_ast_and_multiplicity_once(basis,total,raws):
 s=geometry();b=payload(s,basis=basis);before=copy.deepcopy((s,b));v=calculate(b,s)
 assert (s,b)==before and Decimal(v['quantity'])==total and v['rule_version']=='design-line-railing-1.0'
 assert [Decimal(e['length']) for e in v['railing_edges']]==raws and v['railing_edges'][-1]['to_index']==3
 assert v['source_geometry']==snapshot(s) and v['railing_summary']['source_kind']=='design_route'
 assert 'railing.length_basis' in v['source_evidence'] and component_profile(b,v)['id']==('spatial_length' if basis=='spatial' else 'horizontal_projection')
 with localcontext() as ctx:
  ctx.prec=80
  assert evaluate(v['formula_ast'],v['dimensions'])==total
  for e in v['railing_edges']:assert evaluate(e['selected_formula_ast'],v['dimensions'])==Decimal(e['selected_length'])
 b['edge_decisions'].reverse();assert calculate(b,s)==v


@pytest.mark.parametrize('basis,total,ranges',[('spatial',26,[(0,2),(4,10)]),('horizontal',16,[(0,1),(2,4)])])
def test_local_ranges_respect_selected_basis(basis,total,ranges):
 s=geometry();b=payload(s,basis=basis);d=b['edge_decisions'][0];d.update(mode='ranges',ranges=[interval(a*100,z*100,unit='cm') for a,z in ranges]);v=calculate(b,s)
 assert Decimal(v['quantity'])==total and evaluate(v['formula_ast'],v['dimensions'])==total
 # All endpoints are original raw distance measurements, not derived percentages.
 assert all(e['source_unit']=='cm' for e in v['unit_conversions'] if e['parameter'].startswith('railing.1.'))


def test_vertical_spatial_allowed_horizontal_blocked_not_silently_excluded():
 s=geometry();s['route_points']=[point('A',0,0,0),point('B',0,0,5)];b=payload(s);assert Decimal(calculate(b,s)['quantity'])==10
 with pytest.raises(DomainError,match='竖直'):boundary(s,'horizontal')


def test_irrational_limit_and_raw_source_length_not_rounded():
 s=geometry();s['route_points']=[point('A',0,0,0),point('B',1,1,1)];b=payload(s);b['rounding_digits']=12
 with localcontext() as c:
  c.prec=80;v=calculate(b,s);assert Decimal(v['quantity'])==(Decimal(3).sqrt()*2).quantize(Decimal('1e-12'))
 d=b['edge_decisions'][0];d.update(mode='ranges',ranges=[interval(0,'1.73205080756887729')]);assert calculate(b,s)
 d['ranges'][0]['to']='1.73205080756887730'
 with pytest.raises(DomainError,match='实际长度'):calculate(b,s)


@pytest.mark.parametrize('transform',['translate','reverse','unit','swap_xy','scale'])
def test_design_line_transforms_and_raw_points_keep_expected_length(transform):
 s=geometry();factor=1
 if transform=='reverse':s['route_points'].reverse()
 for p in s['route_points']:
  x,y,z=map(Decimal,p['parameters'].values())
  if transform=='translate':x+=123;y-=99;z+=50
  if transform=='swap_xy':x,y=y,x
  if transform=='unit':x*=1000;y*=1000;z*=1000;p['unit']='mm'
  if transform=='scale':x*=7;y*=7;z*=7;factor=7
  p['parameters']=dict(zip('xyz',map(str,[x,y,z])))
 assert Decimal(calculate(payload(s),s)['quantity'])==36*factor
 assert Decimal(calculate(payload(s,basis='horizontal'),s)['quantity'])==20*factor


def test_seeded_independent_sloped_oracle():
 rng=random.Random(6601)
 for _ in range(30):
  k=rng.randint(1,20);s=geometry();s['route_points']=[point('A',0,0,0),point('B',3*k,4*k,12*k)];s['multiplicity']=str(rng.randint(1,8))
  for basis,n in [('spatial',13),('horizontal',5)]:
   v=calculate(payload(s,basis=basis),s);assert Decimal(v['quantity'])==n*k*int(s['multiplicity'])


@pytest.mark.parametrize('fault',['missing_basis','bool_basis','unknown_basis','missing_proof','blank_proof','old_edges','overlap','outside','missing_decision','unknown_edge','count','scope','no_z','stale_hash'])
def test_invalid_route_basis_and_range_inputs(fault):
 s=geometry();b=payload(s)
 if fault=='missing_basis':b.pop('length_basis')
 elif fault=='bool_basis':b['length_basis']=True
 elif fault=='unknown_basis':b['length_basis']='slope_factor'
 elif fault=='missing_proof':b.pop('length_basis_evidence_id')
 elif fault=='blank_proof':b['length_basis_evidence_id']=''
 elif fault=='old_edges':b['length_basis']='horizontal'
 elif fault=='overlap':b['edge_decisions'][0].update(mode='ranges',ranges=[interval(0,2),interval(1,3)])
 elif fault=='outside':b['edge_decisions'][0].update(mode='ranges',ranges=[interval(0,'13.0000000000000001')])
 elif fault=='missing_decision':b['edge_decisions'].pop()
 elif fault=='unknown_edge':b['edge_decisions'][0]['edge_id']='unknown'
 elif fault=='count':b['multiplicity']='2'
 elif fault=='scope':b['scope_confirmed']=False
 elif fault=='no_z':s['route_points'][1]['parameters'].pop('z');b['source_component_hash']=store.digest(snapshot(s))
 elif fault=='stale_hash':b['source_component_hash']='old'
 with pytest.raises((DomainError,HTTPException)):calculate(b,s)


def test_legacy_polygon_outputs_and_mode_fields_remain_strict():
 from tests.test_railing_takeoff import geometry as pg,payload as pp
 s=pg();b=pp(s);v=calculate(b,s);assert v['rule_version']=='polygon-railing-length-1.0' and 'length_basis' not in v['railing_summary']
 b['length_basis']='horizontal'
 with pytest.raises(DomainError):calculate(b,s)


def test_candidates_preview_independent_review_books_boq_mapping(clients):
 e,r,pid,s,b,eid=setup(clients);before=get(e,pid);url=f'/api/v1/projects/{pid}/railing-edge-candidates?component_id='+s['id']
 assert e.get(url).status_code==422 and get(e,pid)==before
 for basis in ('spatial','horizontal'):
  res=e.get(url+'&length_basis='+basis);assert res.status_code==200,res.text
  assert res.json()['length_basis']==basis and get(e,pid)==before
 post(e,pid,b,'/components/preview');assert get(e,pid)==before
 c=post(e,pid,{**b,'source_geometry':{'forged':True}}).json();assert c['source_geometry']==snapshot(s)
 post(e,pid,{},'/takeoffs/calculate',409);approve(r,pid,'components',c)
 for t in post(e,pid,{},'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
 assert any(v['unit']=='m' and Decimal(v['quantity'])==36 for v in post(e,pid,{},'/boq/generate').json())


@pytest.mark.parametrize('fmt',['json','xlsx','pdf'])
def test_three_exports_numeric_ranges_and_old_quote_isolation(clients,fmt):
 e,r,pid,_,quote=priced_project(clients);e,r,pid,s,b,eid=setup(clients,pid)
 b['edge_decisions'][0].update(mode='ranges',ranges=[{**interval(0,2,eid),'from':0,'to':2}]);c=post(e,pid,b).json();approve(r,pid,'components',c)
 for t in post(e,pid,{},'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
 url=f'/api/v1/projects/{pid}/exports';res=e.get(url,params={'format':fmt,'unpriced':True});assert res.status_code==200,res.text
 titles=['坐标线栏杆口径与段长','坐标线栏杆原始XYZ']
 if fmt=='json':assert all(any(t[0]==title for t in res.json()['tables']) for title in titles)
 elif fmt=='xlsx':assert set(titles)<=set(load_workbook(io.BytesIO(res.content)).sheetnames)
 else:assert all(t in ''.join(p.extract_text() for p in PdfReader(io.BytesIO(res.content)).pages) for t in titles)
 old=e.get(url,params={'format':'json','quote_id':quote['id']}).json();assert old['quote']==quote and not any(t[0] in titles for t in old['tables'])


@pytest.mark.parametrize('fault',['mode','mode_proof','xyz','squared_length','length','range','ast','source','book'])
def test_saved_outputs_and_basis_cannot_be_forged_at_review(clients,fault):
 e,r,pid,s,b,eid=setup(clients);c=post(e,pid,b).json()
 if fault=='book':approve(r,pid,'components',c);post(e,pid,{},'/takeoffs/calculate')
 def forge(p):
  x=p['components'][-1];v=x['preview'];edge=v['railing_edges'][0]
  if fault=='mode':x['length_basis']='horizontal'
  elif fault=='mode_proof':x['length_basis_evidence_id']='forged'
  elif fault=='xyz':edge['from_xyz']['z']='999'
  elif fault=='squared_length':edge['squared_length']['numerator']='999'
  elif fault=='length':edge['spatial_length']='999'
  elif fault=='range':edge['selected_length']='999'
  elif fault=='ast':v['formula_ast']={'constant':'36'}
  elif fault=='source':x['source_geometry']['route_points'][1]['parameters']['z']='999'
  else:p['takeoffs'][-1]['railing_summary']['length_basis']='horizontal'
 mutate(pid,forge);coll='takeoffs' if fault=='book' else 'components';approve(r,pid,coll,get(e,pid)[coll][-1],expected=404 if fault=='mode_proof' else 422 if fault=='mode' else 409)


def test_mode_change_needs_new_edges_history_stale_books_and_boq(clients):
 e,r,pid,s,b,eid=setup(clients);c=post(e,pid,b).json();approve(r,pid,'components',c)
 for t in post(e,pid,{},'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
 post(e,pid,{},'/boq/generate');old=get(e,pid)['components'][-1];url=f'/api/v1/projects/{pid}/components/'+c['id'];bad={**b,'length_basis':'horizontal','object_hash':store.digest(old)}
 assert e.put(url,json=bad).status_code==422
 new={**payload(s,eid,'horizontal'),'object_hash':store.digest(old)};res=e.put(url,json=new);assert res.status_code==200,res.text
 p=get(e,pid);assert p['component_history'][-1]==old and all(t['state']=='stale' for t in p['takeoffs']) and any(v['quantity']=='36.000' for v in p['boq'])
 assert e.put(url,json=new).status_code==409;approve(r,pid,'components',res.json())
 for t in post(e,pid,{},'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
 assert any(v['quantity']=='20.000' for v in post(e,pid,{},'/boq/generate').json())
 # Source Z revision invalidates the dependent even for horizontal basis; fresh installation review required.
 s['object_hash']=store.digest(s);s['route_points'][1]['parameters']['z']='9';res=e.put(f'/api/v1/projects/{pid}/components/'+s['id'],json=s);assert res.status_code==200,res.text
 assert get(e,pid)['components'][-1]['state']=='stale';approve(r,pid,'components',res.json());post(e,pid,new,'/components/preview',409)


@pytest.mark.parametrize('gate',['candidate','preview','save','review','calculate','book_review'])
def test_original_integrity_at_all_gates(clients,gate):
 e,r,pid,s,b,eid=setup(clients)
 if gate in ('review','calculate','book_review'):c=post(e,pid,b).json()
 if gate in ('calculate','book_review'):approve(r,pid,'components',c)
 if gate=='book_review':t=post(e,pid,{},'/takeoffs/calculate').json()[-1]
 doc=get(e,pid)['documents'][-1];(store.ROOT/'originals'/doc['hash']).write_bytes(b'corrupt')
 if gate=='candidate':assert e.get(f'/api/v1/projects/{pid}/railing-edge-candidates?component_id='+s['id']+'&length_basis=spatial').status_code==409
 elif gate=='preview':post(e,pid,b,'/components/preview',409)
 elif gate=='save':post(e,pid,b,expected=409)
 elif gate=='review':approve(r,pid,'components',c,expected=409)
 elif gate=='calculate':post(e,pid,{},'/takeoffs/calculate',409)
 else:approve(r,pid,'takeoffs',t,expected=409)


def test_access_foreign_proofs_project_members_and_source_state(clients):
 e,r,pid,s,b,eid=setup(clients);url=f'/api/v1/projects/{pid}/railing-edge-candidates?component_id='+s['id']+'&length_basis=spatial'
 assert r.get(url).status_code==200 and clients['other'].get(url).status_code==404 and TestClient(app).get(url).status_code==401
 post(r,pid,b,expected=403);other=create(e)['id'];upload(e,other,'其他项目的依据');bad=copy.deepcopy(b);bad['length_basis_evidence_id']=get(e,other)['evidence'][-1]['id'];post(e,pid,bad,expected=404)
 assert e.get(f'/api/v1/projects/{other}/railing-edge-candidates?component_id='+s['id']+'&length_basis=spatial').status_code==404
 mutate(pid,lambda p:p['components'][0].update(state='draft'));assert e.get(url).status_code==409
 from tests.test_project_access import restrict
 restrict(clients['admin'],pid,[]);assert e.get(url).status_code==404


def test_large_origin_small_xyz_edges_and_low_caller_precision():
 s=geometry()
 with localcontext() as ctx:
  ctx.prec=80
  for p in s['route_points']:p['parameters']={k:str(Decimal('1000000000')+Decimal(v)*Decimal('1e-8')) for k,v in p['parameters'].items()}
 b=payload(s)
 with localcontext() as ctx:
  ctx.prec=10;v=calculate(b,s)
 assert Decimal(v['gross'])==Decimal('36e-8')


def test_horizontal_projection_overlap_does_not_infer_physical_contact_or_deduction():
 s=geometry();s['route_points'][2]=point('C',0,0,20);b=payload(s,basis='horizontal');v=calculate(b,s)
 assert Decimal(v['quantity'])==20 and all(Decimal(e['horizontal_length'])==5 for e in v['railing_edges'])
 assert v['railing_edges'][1]['to_xyz']['z']=='20' and v['railing_edges'][1]['to_index']==3
