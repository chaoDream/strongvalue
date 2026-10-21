"""Synthetic field comparisons, not actual GTJ/GQI/GCCP compatibility tests."""
import copy
import io
import json
from uuid import uuid4
import pytest
from fastapi import HTTPException
from openpyxl import Workbook, load_workbook
from backend import store
from backend.integration_trials import HEADERS, compare, parse_table
from tests.test_workflow import create, get, upload, approve, priced_project


def workbook(rows, header=None):
    w=Workbook();s=w.active;s.title='字段对照';s.append(header or HEADERS)
    for row in rows:s.append(row)
    raw=io.BytesIO();w.save(raw)
    return raw.getvalue()


def baseline_rows():
    return [['清单','B01','工程量','15.2','m³','数值','0'],['定额','Q01','清单关联','B01','','文本',''],
            ['资源','R01','单价','100.005','元','数值','.005'],['费用','F01','税制','不含税','','文本','']]


def returned_rows():
    return [row[:5]+['',''] for row in baseline_rows()]


def setup(clients, differences=False, priced=False):
    e,r=clients['estimator'],clients['reviewer']
    if priced:e,r,pid,_,_=priced_project(clients)
    else:pid=create(e)['id']
    base='/api/v1/projects/'+pid
    files=[]
    for name,rows in [('baseline.xlsx',baseline_rows()),('returned.xlsx',returned_rows())]:
        if differences and name=='returned.xlsx':rows[0][3]='16';rows.pop(1);rows[-1][4]='含税标识'
        response=e.post(base+'/documents',data={'specialty':'合同计价'},files={'file':(name,workbook(rows),'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')})
        assert response.status_code==200,response.text
        files.append(response.json()['document']['id'])
    upload(e,pid,'合成软件许可和执行日志：仅用于功能测试，未运行真实厂商软件。',name='trial-proof.txt',specialty='合同计价')
    eid=get(e,pid)['evidence'][-1]['id']
    body={'trial_key':'SYNTHETIC-01','name':'合成软件字段往返试验（非厂商验收）','provider':'合成测试工具','product':'Fixture',
          'software_version':'test-1.0','return_software_version':'test-1.0','adapter_version':'manual-v1',
          'input_format':'人工整理XLSX','output_format':'人工整理XLSX','legal_scope':'仅合成测试，无厂商软件许可主张',
          'execution_summary':'仅模拟运行记录，不代表真实软件回导','model_fidelity':'仅列示字段，不含原生几何拓扑',
          'fee_fidelity':'仅列示字段，不推断费用规则保真','failure_recovery':'合成测试，不操作实际软件','fallback':'不支持字段改由人工核查并重新试验',
          'tested_on':'2026-10-21','baseline_document_id':files[0],'returned_document_id':files[1],
          'license_evidence_id':eid,'execution_evidence_id':eid,'run_status':'imported','previous_hash':'','request_id':str(uuid4())}
    return e,r,pid,body


def submit(e,pid,body,expected=200):
    base='/api/v1/projects/'+pid
    preview=e.post(base+'/integration-trials/preview',json=body)
    assert preview.status_code==200,preview.text
    body={**body,'preview_hash':preview.json()['preview_hash']}
    response=e.post(base+'/integration-trials',json=body)
    assert response.status_code==expected,response.text
    return response.json(),body


def test_same_fields_review_and_capabilities_remain_unconfigured(clients):
    e,r,pid,body=setup(clients);base='/api/v1/projects/'+pid
    before=get(e,pid)
    preview=e.post(base+'/integration-trials/preview',json=body)
    assert preview.status_code==200,preview.text
    assert get(e,pid)==before
    record,_=submit(e,pid,body)
    assert record['comparison']['passed'] and record['comparison']['coverage_percent']=='100.00'
    assert record['comparison']['counts']['equal']==4
    approve(r,pid,'integration_trials',record)
    current=e.get(base+'/integration-trials/status').json()['items'][0]
    assert current['review_valid'] and not current['can_import'] and not current['can_export']
    assert all(not x['can_import'] and not x['can_export'] for x in e.get('/api/v1/integrations/capabilities').json())


def test_losses_units_and_value_changes_are_independent(clients):
    e,r,pid,body=setup(clients,True)
    record,_=submit(e,pid,body)
    report=record['comparison']
    assert not report['passed'] and report['coverage_percent']=='75.00'
    assert report['counts']=={'equal':1,'within_tolerance':0,'changed':1,'missing':1,'unit_mismatch':1,'added':0}
    assert report['rows'][0]['delta']=='0.8'
    # Reviewing the truthful failed result does not turn it into a passing compatibility claim.
    approve(r,pid,'integration_trials',record)
    assert not get(e,pid)['integration_trials'][0]['comparison']['passed']


def test_tolerance_additions_empty_text_and_signed_numbers():
    old=parse_table({'name':'a.xlsx'},workbook([['资源','R','价','100','元','数值','.01'],['标高','Z','值','-2','m','数值','0'],['模型','M','特征','','','文本','']]),True)
    new=parse_table({'name':'b.xlsx'},workbook([['资源','R','价','100.005','元','',''],['标高','Z','值','-1','m','',''],['模型','M','特征',' ','','',''],['模型','N','父项','M','','','']]),False)
    result=compare(old,new,'imported')
    assert result['counts']['within_tolerance']==1 and result['counts']['changed']==2 and result['counts']['added']==1
    assert result['rows'][1]['delta']=='1' and not result['passed']


def test_returned_invalid_number_is_difference_not_silent_zero():
    old=parse_table({'name':'a.xlsx'},workbook(baseline_rows()),True)
    rows=returned_rows();rows[0][3]='缺失数值'
    new=parse_table({'name':'b.xlsx'},workbook(rows),False)
    result=compare(old,new,'imported')
    assert result['rows'][0]['status']=='changed' and result['rows'][0]['delta'] is None
    assert result['rows'][0]['error'] and not result['passed']


@pytest.mark.parametrize('case',['duplicate','formula','error','header','blank','mode','tolerance','negative_tolerance','nan','oversized','extra_column','over_rows'])
def test_strict_baseline_parser(case):
    rows=baseline_rows();header=None
    if case=='duplicate':rows.append(rows[0])
    elif case=='formula':rows[0][3]='=1+2'
    elif case=='error':rows[0][3]='#DIV/0!'
    elif case=='header':header=['错误']+HEADERS[1:]
    elif case=='blank':rows=[]
    elif case=='mode':rows[0][5]='自动'
    elif case=='tolerance':rows[0][6]=''
    elif case=='negative_tolerance':rows[0][6]='-1'
    elif case=='nan':rows[0][3]='NaN'
    elif case=='oversized':rows[0][3]='1e100'
    elif case=='extra_column':rows[0].append('隐藏信息')
    else:rows=[['类型',str(i),'字段','1','','数值','0'] for i in range(5001)]
    with pytest.raises(HTTPException):parse_table({'name':'source.xlsx'},workbook(rows,header),True)


def test_failed_import_can_archive_without_a_return_file(clients):
    e,r,pid,body=setup(clients);body.update(run_status='failed',returned_document_id='')
    record,_=submit(e,pid,body)
    assert record['comparison']['counts']['missing']==4 and not record['comparison']['passed']
    approve(r,pid,'integration_trials',record)
    assert get(e,pid)['integration_trials'][0]['state']=='approved'


def test_roles_tenants_foreign_files_and_evidence(clients):
    e,r,pid,body=setup(clients);base='/api/v1/projects/'+pid
    assert r.post(base+'/integration-trials/preview',json=body).status_code==403
    assert r.post(base+'/integration-trials',json=body).status_code==403
    assert clients['other'].get(base+'/integration-trials/template').status_code==404
    assert clients['other'].post(base+'/integration-trials/preview',json=body).status_code==404
    _,_,pid2,foreign=setup(clients)
    for key in ['baseline_document_id','returned_document_id','license_evidence_id','execution_evidence_id']:
        invalid={**body,key:foreign[key]}
        assert e.post(base+'/integration-trials/preview',json=invalid).status_code==404


def test_versions_preview_guard_idempotence_and_original_protection(clients):
    e,r,pid,body=setup(clients);base='/api/v1/projects/'+pid
    record,sent=submit(e,pid,body)
    assert e.post(base+'/integration-trials',json=sent).json()['id']==record['id']
    assert e.post(base+'/integration-trials',json={**sent,'name':'别的内容'}).status_code==409
    approve(r,pid,'integration_trials',record)
    next_body={**body,'request_id':str(uuid4()),'name':'第二次合成试验'}
    assert e.post(base+'/integration-trials',json={**next_body,'preview_hash':sent['preview_hash']}).status_code==409
    next_body['previous_hash']=e.get(base+'/hashes').json()['integration_trials'][record['id']]
    second,_=submit(e,pid,next_body)
    assert second['version']==2
    approve(r,pid,'integration_trials',record,expected=409)
    assert e.get(base+'/integration-trials/status').json()['items'][0]['issue']
    assert get(e,pid)['integration_trials'][0]['comparison']==record['comparison']


def test_stale_preview_is_atomic_and_same_software_version_required(clients):
    e,_,pid,body=setup(clients);base='/api/v1/projects/'+pid
    preview=e.post(base+'/integration-trials/preview',json=body).json()
    before=get(e,pid)
    response=e.post(base+'/integration-trials',json={**body,'model_fidelity':'改变口径','preview_hash':preview['preview_hash']})
    assert response.status_code==409 and get(e,pid)==before
    assert e.post(base+'/integration-trials/preview',json={**body,'return_software_version':'another-version'}).status_code==422


@pytest.mark.parametrize('stage',['preview','save','review','export'])
def test_original_tampering_blocks_gates_and_changes_new_export_label(clients,stage):
    e,r,pid,body=setup(clients);base='/api/v1/projects/'+pid
    if stage in ('review','export'):record,_=submit(e,pid,body)
    if stage=='export':approve(r,pid,'integration_trials',record)
    if stage=='save':preview=e.post(base+'/integration-trials/preview',json=body).json()
    doc=next(d for d in get(e,pid)['documents'] if d['id']==body['baseline_document_id'])
    path=store.ROOT/'originals'/doc['hash']
    if stage=='export':path.unlink()
    else:path.write_bytes(b'tampered')
    if stage=='preview':assert e.post(base+'/integration-trials/preview',json=body).status_code==409
    elif stage=='save':assert e.post(base+'/integration-trials',json={**body,'preview_hash':preview['preview_hash']}).status_code==409
    elif stage=='review':approve(r,pid,'integration_trials',record,expected=409)
    else:
        result=e.get(base+'/integration-trials/'+record['id']+'/exports').json()
        assert '草稿' in result['status'] and result['issue'] and not result['can_import']


def test_replaced_sources_cannot_be_reviewed_again(clients):
    e,r,pid,body=setup(clients);base='/api/v1/projects/'+pid
    record,_=submit(e,pid,body);approve(r,pid,'integration_trials',record)
    upload(e,pid,'新的许可说明',name='trial-proof.txt',specialty='合同计价')
    approve(r,pid,'integration_trials',record,expected=409)
    assert not e.get(base+'/integration-trials/status').json()['items'][0]['review_valid']


def test_exports_template_formula_safety_revocation_and_immutable_artifacts(clients):
    e,r,pid,body=setup(clients);base='/api/v1/projects/'+pid
    template=e.get(base+'/integration-trials/template')
    assert list(load_workbook(io.BytesIO(template.content))['字段对照'].values)[0]==tuple(HEADERS)
    body['name']='=HYPERLINK("unsafe","合成试验")'
    record,_=submit(e,pid,body);approve(r,pid,'integration_trials',record)
    output=e.get(base+'/integration-trials/'+record['id']+'/exports')
    assert '已复核' in output.json()['status'] and output.json()['reviews']
    sheet=e.get(base+'/integration-trials/'+record['id']+'/exports?format=xlsx')
    assert sheet.status_code==200,sheet.text
    wb=load_workbook(io.BytesIO(sheet.content),data_only=False)
    assert '逐字段对照' in wb.sheetnames
    assert all(cell.data_type!='f' for s in wb for row in s for cell in row)
    artifact=next(a for a in e.get(base+'/artifacts').json() if a['format']=='json')
    assert artifact['purpose']=='integration_trial'
    assert clients['other'].get(base+'/artifacts/'+artifact['id']+'/file').status_code==404
    approve(r,pid,'integration_trials',record,'reject')
    assert '草稿' in e.get(base+'/integration-trials/'+record['id']+'/exports').json()['status']
    downloaded=e.get(base+'/artifacts/'+artifact['id']+'/file')
    assert downloaded.content==output.content


def test_trial_review_and_rejection_do_not_change_quote_or_approvals(clients):
    e,r,pid,body=setup(clients,priced=True);base='/api/v1/projects/'+pid
    before=get(e,pid)
    fields=['quotes','boq','takeoffs','rules','report_approval','stale_quotes','status']
    record,_=submit(e,pid,body)
    approve(r,pid,'integration_trials',record)
    approve(r,pid,'integration_trials',record,'reject')
    after=get(e,pid)
    assert all(after[k]==before[k] for k in fields)


def test_content_integrity_and_metadata_evidence_changes(clients):
    e,r,pid,body=setup(clients);base='/api/v1/projects/'+pid
    record,_=submit(e,pid,body)
    with store.transaction() as c:
        p=json.loads(c.execute('SELECT data FROM projects WHERE id=?',(pid,)).fetchone()['data'])
        p['integration_trials'][0]['comparison']['passed']=False
        store.save(c,p)
    approve(r,pid,'integration_trials',record,expected=409)
    assert e.get(base+'/integration-trials/'+record['id']+'/exports').status_code==409


def test_changed_evidence_invalidates_previously_reviewed_trial(clients):
    e,r,pid,body=setup(clients);base='/api/v1/projects/'+pid
    record,_=submit(e,pid,body);approve(r,pid,'integration_trials',record)
    with store.transaction() as c:
        p=json.loads(c.execute('SELECT data FROM projects WHERE id=?',(pid,)).fetchone()['data'])
        next(x for x in p['evidence'] if x['id']==body['license_evidence_id'])['text']='依据摘录已改变'
        store.save(c,p)
    approve(r,pid,'integration_trials',record,expected=409)
    assert not e.get(base+'/integration-trials/status').json()['items'][0]['review_valid']


def test_project_member_scope_covers_trial_and_artifact_routes(clients):
    e,r,pid,body=setup(clients);base='/api/v1/projects/'+pid
    record,_=submit(e,pid,body)
    a=clients['admin'];access=a.get(base+'/access').json()
    response=a.put(base+'/access',json={'mode':'restricted','member_ids':[],'reason':'合成权限验收','object_hash':access['object_hash']})
    assert response.status_code==200,response.text
    for path in ['/integration-trials/status','/integration-trials/template','/integration-trials/'+record['id']+'/exports']:
        assert e.get(base+path).status_code==404
    assert e.post(base+'/integration-trials/preview',json=body).status_code==404
    assert r.get(base+'/review-target?collection=integration_trials&object_id='+record['id']).status_code==404
