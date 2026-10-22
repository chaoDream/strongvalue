import copy
import io
import json
import pytest
from reportlab.pdfgen import canvas
from openpyxl import load_workbook
from pypdf import PdfReader
from backend import store,main as core
from tests.test_workflow import create,get,approve,priced_project


def pdf(c,pid,name='synthetic-drawing.pdf'):
    stream=io.BytesIO();out=canvas.Canvas(stream,pagesize=(600,400))
    for line in ['J-01 REV A PLAN - SECTION 1-1 see J-02','J-02 REV A SECTION 1-1 - software test only']:
        out.drawString(30,350,line);out.showPage()
    out.save()
    response=c.post(f'/api/v1/projects/{pid}/documents',data={'specialty':'建筑'},files={'file':(name,stream.getvalue(),'application/pdf')})
    assert response.status_code==200,response.text
    return response.json()['document']


def body(c,pid,doc,page=1,**extra):
    e=next(e for e in get(c,pid)['evidence'] if e['document_id']==doc['id'] and e['page']==page and e['method']=='local_text')
    return {'document_id':doc['id'],'page':page,'drawing_no':'J-01' if page==1 else 'J-02','title':'合成平面' if page==1 else '合成剖面',
            'edition':'A','view_type':'plan' if page==1 else 'section','discipline':'建筑',
            'evidence_id':e['id'],'raw_text':e['text'].strip(),'reason':'仅合成图号身份核对，不作专业验收',**extra}


def add(c,pid,data,expected=200):
    response=c.post(f'/api/v1/projects/{pid}/drawing-sheets',json=data)
    assert response.status_code==expected,response.text
    return response.json()


def setup(clients):
    e,r=clients['estimator'],clients['reviewer'];pid=create(e)['id'];doc=pdf(e,pid)
    a=add(e,pid,body(e,pid,doc));b=add(e,pid,body(e,pid,doc,2))
    approve(r,pid,'drawing_sheets',a);approve(r,pid,'drawing_sheets',b)
    return e,r,pid,doc,a,b


def reference(c,pid,a,b,**extra):
    return {'from_sheet_id':a['id'],'to_sheet_id':b['id'],'component_name':'合成池壁','kind':'section','marker':'1-1',
            'evidence_id':a['evidence_id'],'raw_text':a['raw_text'],'reason':'按平面1-1索引核对剖面，仅合成软件验收',**extra}


def ref(c,pid,data,expected=200):
    response=c.post(f'/api/v1/projects/{pid}/drawing-references',json=data)
    assert response.status_code==expected,response.text
    return response.json()


def revise(c,pid,collection,row,data,expected=200):
    h=c.get(f'/api/v1/projects/{pid}/hashes').json()[collection][row['id']]
    url='drawing-sheets' if collection=='drawing_sheets' else 'drawing-references'
    response=c.post(f'/api/v1/projects/{pid}/{url}/{row["id"]}/revisions',json={**data,'object_hash':h})
    assert response.status_code==expected,response.text
    return response.json()


def test_reviewed_plan_section_reference_is_source_bound_and_never_creates_quantities(clients):
    e,r,pid,doc,a,b=setup(clients);p=get(e,pid)
    assert p['drawing_register']['issues']==[]
    record=ref(e,pid,reference(e,pid,a,b));assert record['state']=='draft'
    assert record['from_sheet_hash'] and record['to_sheet_hash'] and record['evidence_hash']
    assert any('互参' in x for x in get(e,pid)['blockers'])
    approve(r,pid,'drawing_references',record)
    p=get(e,pid);assert p['drawing_register']['issues']==[]
    assert not p['components'] and not p['takeoffs'] and not p['boq']
    assert [x['action'] for x in e.get(f'/api/v1/projects/{pid}/audit').json()].count('DRAWING_REFERENCE_RECORDED')==1


@pytest.mark.parametrize('change',[
    {'page':2},{'page':True},{'page':0},{'page':3},{'page':None},
    {'view_type':'invented'},{'discipline':'unknown'},{'edition':''},
    {'raw_text':'没有这段原文'},{'state':'approved'},{'drawing_no':{}},
])
def test_bad_sheet_fields_pages_and_source_are_rejected(clients,change):
    e=clients['estimator'];pid=create(e)['id'];doc=pdf(e,pid)
    add(e,pid,{**body(e,pid,doc),**change},422)
    assert not get(e,pid)['drawing_sheets']


@pytest.mark.parametrize('change',[
    {'marker':'不存在'},{'kind':'detail'},{'kind':'unknown'},
    {'raw_text':'原文不匹配'},{'state':'approved'},
])
def test_bad_reference_marker_type_or_source_is_rejected(clients,change):
    e,r,pid,doc,a,b=setup(clients)
    ref(e,pid,reference(e,pid,a,b,**change),422)
    assert not get(e,pid)['drawing_references']


def test_reference_requires_two_distinct_current_approved_sheets(clients):
    e,r,pid,doc,a,b=setup(clients)
    ref(e,pid,reference(e,pid,a,a),422)
    approve(r,pid,'drawing_sheets',b,'request_more_info')
    ref(e,pid,reference(e,pid,a,b),409)


def test_duplicate_number_candidates_require_explicit_rejection(clients):
    e,r,pid,doc,a,b=setup(clients)
    duplicate=add(e,pid,body(e,pid,doc,drawing_no='Ｊ－０１',edition='B'))
    p=get(e,pid);groups=p['drawing_register']['conflicts'];assert len(groups)==1
    assert set(groups[0]['sheet_ids'])=={a['id'],duplicate['id']}
    approve(r,pid,'drawing_sheets',a,expected=409)
    approve(r,pid,'drawing_sheets',duplicate,'reject');approve(r,pid,'drawing_sheets',a)
    assert get(e,pid)['drawing_register']['conflicts']==[]


def test_drawing_revision_preserves_history_and_invalidates_exact_reference(clients):
    e,r,pid,doc,a,b=setup(clients)
    link=ref(e,pid,reference(e,pid,a,b));approve(r,pid,'drawing_references',link)
    changed=revise(e,pid,'drawing_sheets',b,body(e,pid,doc,2,edition='B'))
    p=get(e,pid);assert p['drawing_sheets'][1]['state']=='approved'
    assert changed['previous_id']==b['id'] and changed['series_id']==b['series_id'] and changed['revision']==2
    approve(r,pid,'drawing_sheets',b,expected=409)
    revise(e,pid,'drawing_sheets',b,body(e,pid,doc,2),409)
    approve(r,pid,'drawing_sheets',changed)
    assert any('互参' in x for x in get(e,pid)['drawing_register']['issues'])
    assert link['id'] in get(e,pid)['drawing_register']['source_issues']['drawing_references']
    approve(r,pid,'drawing_references',link,expected=409)
    corrected=revise(e,pid,'drawing_references',link,reference(e,pid,a,changed))
    approve(r,pid,'drawing_references',corrected)
    assert get(e,pid)['drawing_register']['issues']==[]
    assert len(get(e,pid)['drawing_references'])==2
    approve(r,pid,'drawing_references',link,'reject',409)


def test_reapproval_of_unchanged_sheet_still_requires_reference_recheck(clients):
    e,r,pid,doc,a,b=setup(clients);link=ref(e,pid,reference(e,pid,a,b));approve(r,pid,'drawing_references',link)
    approve(r,pid,'drawing_sheets',b,'reject');approve(r,pid,'drawing_sheets',b)
    approve(r,pid,'drawing_references',link,expected=409)
    assert any('审批已变化' in s for s in get(e,pid)['drawing_register']['issues'])


@pytest.mark.parametrize('mutation',['bytes','missing','evidence','replacement'])
def test_source_changes_block_sheet_review_and_report_gate(clients,mutation):
    e,r,pid,doc,a,b=setup(clients)
    if mutation=='bytes':(store.ROOT/'originals'/doc['hash']).write_bytes(b'changed')
    elif mutation=='missing':(store.ROOT/'originals'/doc['hash']).unlink()
    elif mutation=='replacement':pdf(e,pid)
    else:
        with store.transaction() as c:
            p=core.project(c,pid,{'tenant':'local','role':'estimator','username':'estimator'})
            next(e for e in p['evidence'] if e['id']==a['evidence_id'])['text']+='changed'
            store.save(c,p)
    approve(r,pid,'drawing_sheets',a,expected=409)
    assert get(e,pid)['drawing_register']['issues']
    assert e.post(f'/api/v1/projects/{pid}/quotes/calculate').status_code==409


def test_permissions_foreign_evidence_and_stale_revision_hash_do_not_mutate_ledger(clients):
    e,r,pid,doc,a,b=setup(clients);base=f'/api/v1/projects/{pid}'
    assert r.post(base+'/drawing-sheets',json=body(e,pid,doc)).status_code==403
    assert clients['other'].post(base+'/drawing-sheets',json=body(e,pid,doc)).status_code==404
    pid2=create(e)['id'];foreign=pdf(e,pid2)
    add(e,pid,body(e,pid2,foreign),404)
    data=body(e,pid,doc);data['object_hash']='stale'
    assert e.post(base+'/drawing-sheets/'+a['id']+'/revisions',json=data).status_code==409
    assert len(get(e,pid)['drawing_sheets'])==2
    target=r.get(base+'/review-batch-targets?collection=drawing_sheets').json();assert target['total']==0


def test_catalog_rejection_resolves_optional_gate_without_erasing_record(clients):
    e=clients['estimator'];r=clients['reviewer'];pid=create(e)['id'];doc=pdf(e,pid)
    a=add(e,pid,body(e,pid,doc));approve(r,pid,'drawing_sheets',a,'reject')
    p=get(e,pid);assert not p['drawing_register']['issues'] and len(p['drawing_sheets'])==1


def test_cad_keeps_no_physical_page_and_manual_evidence(clients):
    e=clients['estimator'];pid=create(e)['id'];base=f'/api/v1/projects/{pid}'
    response=e.post(base+'/documents',data={'specialty':'结构'},files={'file':('test.dwg',b'CAD software test only')})
    doc=response.json()['document']
    ev=e.post(base+'/evidence',json={'document_id':doc['id'],'page':1,'text':'合成CAD图号S-1，A版，仅人工核对'}).json()
    data={'document_id':doc['id'],'page':None,'drawing_no':'S-1','title':'合成CAD','edition':'A','view_type':'plan','discipline':'结构','evidence_id':ev['id'],'raw_text':ev['text'],'reason':'仅合成登记，无授权CAD解析声明'}
    record=add(e,pid,data);assert record['page'] is None
    approve(clients['reviewer'],pid,'drawing_sheets',record)
    add(e,pid,{**data,'page':1},422)


def test_reports_use_frozen_drawing_register_and_new_registration_invalidates_quote(clients):
    e,r,pid,_,q0=priced_project(clients);base=f'/api/v1/projects/{pid}'
    doc=pdf(e,pid);a=add(e,pid,body(e,pid,doc));b=add(e,pid,body(e,pid,doc,2))
    approve(r,pid,'drawing_sheets',a);approve(r,pid,'drawing_sheets',b)
    link=ref(e,pid,reference(e,pid,a,b));approve(r,pid,'drawing_references',link)
    p=get(e,pid)
    for coll in ['components','takeoffs','boq']:
        for row in p[coll]:approve(r,pid,coll,row)
    q=e.post(base+'/quotes/calculate');assert q.status_code==200,q.text
    q=q.json();assert len(q['snapshot']['drawing_sheets'])==2
    frozen=copy.deepcopy(q)
    revise(e,pid,'drawing_sheets',b,body(e,pid,doc,2,edition='B',title='新版合成剖面'))
    p=get(e,pid);assert q['id'] in p['stale_quotes'] and all(b['state']=='stale' for b in p['boq'])
    data=e.get(base+'/exports',params={'format':'json','quote_id':q['id']}).json()
    assert '新版合成剖面' not in str(data['tables']) and '平立剖图纸互参' in str(data['tables'])
    assert data['quote']==frozen
    x=e.get(base+'/exports',params={'format':'xlsx','quote_id':q['id']});wb=load_workbook(io.BytesIO(x.content))
    assert {'图纸目录与版次','平立剖图纸互参'}<=set(wb.sheetnames)
    out=e.get(base+'/exports',params={'format':'pdf','quote_id':q['id']});text=''.join(p.extract_text() for p in PdfReader(io.BytesIO(out.content)).pages)
    assert '平立剖图纸互参' in text and '新版合成剖面' not in text
    assert get(e,pid)['quotes'][-1]==frozen
    assert q0['snapshot']['drawing_sheets']==[]
    old=e.get(base+'/exports',params={'format':'json','quote_id':q0['id']}).json()
    assert not any(t[0]=='图纸目录与版次' for t in old['tables'])
    draft=e.get(base+'/exports',params={'format':'json','unpriced':True}).json()
    assert any(t[0]=='图纸校核待办' for t in draft['tables'])


def test_batch_targets_only_latest_and_historical_batch_rolls_back(clients):
    e,r,pid,doc,a,b=setup(clients)
    new=revise(e,pid,'drawing_sheets',b,body(e,pid,doc,2,edition='B'))
    base=f'/api/v1/projects/{pid}'
    targets=r.get(base+'/review-batch-targets?collection=drawing_sheets').json()
    assert [x['object']['id'] for x in targets['targets']]==[new['id']]
    hashes=r.get(base+'/hashes').json()['drawing_sheets']
    response=r.post(base+'/reviews/batch',json={'collection':'drawing_sheets','decision':'approve','reason':'合成批审原子性核对',
        'targets':[{'object_id':x['id'],'object_hash':hashes[x['id']]} for x in [new,b]]})
    assert response.status_code==409,response.text
    assert next(x for x in get(e,pid)['drawing_sheets'] if x['id']==new['id'])['state']=='draft'


def test_foreign_sheet_reference_and_nonmember_write_are_rejected(clients):
    e,r,pid,doc,a,b=setup(clients);pid2=create(e)['id'];doc2=pdf(e,pid2)
    foreign=add(e,pid2,body(e,pid2,doc2,2));approve(r,pid2,'drawing_sheets',foreign)
    ref(e,pid,reference(e,pid,a,foreign),404)
    with store.transaction() as c:
        p=core.project(c,pid,{'tenant':'local','role':'admin','username':'admin'})
        p['access']={'mode':'restricted','member_ids':[],'version':1};store.save(c,p)
    response=e.post(f'/api/v1/projects/{pid}/drawing-references',json=reference(e,pid,a,b))
    assert response.status_code==404
