"""Synthetic polygon-volume ledgers; no inferred 3D containment or professional approval."""
import copy
import io
import json
from decimal import Decimal, localcontext
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pypdf import PdfReader
from backend import store
from backend.main import app
from backend.volume_balance import snapshot, reference
from tests.test_polygon_takeoff import setup as polygon_setup, payload as polygon_body, ring
from tests.test_volume_balance import payload as balance_body, primitive
from tests.test_route_takeoff import evaluate
from tests.test_workflow import get, upload, approve, create, priced_project


def post(e,pid,endpoint,body,expected=200):
    res=e.post(f'/api/v1/projects/{pid}'+endpoint,json=body);assert res.status_code==expected,res.text;return res


def mutation(pid,action):
    with store.transaction() as c:
        p=json.loads(c.execute('select data from projects where id=?',(pid,)).fetchone()['data']);action(p);store.save(c,p)


def setup(clients,pid=None,excavation_kind='polygon_region',occupation_kind='polygon_region',tiny=False):
    e,r,pid,b=polygon_setup(clients,pid);eid=b['basis_evidence_id'];b.update(name='合成凹形开挖（非真实工程）',thickness='2',thickness_unit='m')
    if excavation_kind!='polygon_region':
        b=primitive('excavation','trench',{'length':'5','bottom_width':'4','top_width':'6','depth':'4'},eid)
        for k in ('id','revision','preview'):b.pop(k)
    ex=post(e,pid,'/components',b).json()
    b=polygon_body(eid);b['name']='合成斜孔底板占位（非真实工程）'
    if occupation_kind!='polygon_region':
        b=primitive('occupation','slab',{'length':'5','width':'2','thickness':'1'},eid)
        for k in ('id','revision','preview'):b.pop(k)
    elif tiny:
        b.update(polygon_rings=[ring('合成微量三角形',[(0,0),(1,0),(0,'.005')],eid)],thickness='1',thickness_unit='m',rounding_digits=2)
    oc=post(e,pid,'/components',b).json()
    area=polygon_body(eid,'area');area['name']='合成面积对照（不可作为体积来源）';area=post(e,pid,'/components',area).json()
    upload(e,pid,'合成回填依据（非真实工程）：凹形开挖净面积25m²、统一竖向深度2m、重复2组，开挖100m³；斜孔底板净面积25m²、统一竖向厚200mm、重复2组，净占位10m³。本次明确两者属于同一计量范围，孔范围已在各来源内扣除，来源组数不再重乘；本次占位按1全部扣除，无重复，回填90m³，总量三位。部分扣除须另核定比例及范围；不证明真实空间包含。','polygon-balance-basis.txt')
    for c in get(e,pid)['components']:approve(r,pid,'components',c)
    p=get(e,pid);by_id={c['id']:c for c in p['components']};src={'excavation':snapshot(by_id[ex['id']]),'occupations':[snapshot(by_id[oc['id']])]}
    body=balance_body(src,p['evidence'][-1]['id']);body['name']='合成多边形关联回填（非真实工程）'
    return e,r,pid,src,body,by_id[area['id']]


def selection(body):
    return {'excavation_component_id':body['excavation_component_id'],'occupation_component_ids':[v['component_id'] for v in body['occupations']]}


@pytest.mark.parametrize('ex,oc',[('polygon_region','polygon_region'),('polygon_region','slab'),('trench','polygon_region')])
def test_polygon_sources_in_both_roles_and_mixed_legacy_roles_use_reviewed_quantities(clients,ex,oc):
    e,r,pid,src,b,area=setup(clients,excavation_kind=ex,occupation_kind=oc);before=get(e,pid)
    candidate=post(e,pid,'/volume-balance-candidates',selection(b)).json();assert candidate==post(r,pid,'/volume-balance-candidates',selection(b)).json()
    assert candidate['excavation']['kind']==ex and candidate['occupations'][0]['kind']==oc
    assert Decimal(candidate['excavation']['quantity'])==100 and Decimal(candidate['occupations'][0]['quantity'])==10
    view=post(e,pid,'/components/preview',b).json();assert view['quantity']=='90.000' and get(e,pid)==before
    assert view['source_geometry']==src and view['excavation_reference']['component_hash']==b['excavation_component_hash']
    assert Decimal(view['balance_deductions'][0]['deduction'])==10
    with localcontext() as ctx:ctx.prec=80;assert evaluate(view['formula_ast'],view['dimensions'])==90
    forged=copy.deepcopy(b);forged['source_geometry']={'excavation':{'quantity':'1'},'occupations':[]}
    comp=post(e,pid,'/components',forged).json();assert comp['source_geometry']==src and comp['preview']['quantity']=='90.000'
    assert all(c['state']=='approved' for c in get(e,pid)['components'][:-1])
    post(e,pid,'/takeoffs/calculate',{},409);approve(r,pid,'components',comp)
    books=post(e,pid,'/takeoffs/calculate',{}).json()
    for t in books:approve(r,pid,'takeoffs',t)
    items=post(e,pid,'/boq/generate',{}).json();assert next(v for v in items if v['name']==b['name'])['quantity']=='90.000'


def test_source_rounding_holes_and_multiplicity_are_applied_exactly_once(clients):
    e,r,pid,src,b,area=setup(clients,tiny=True)
    oc=src['occupations'][0];assert oc['preview']['quantity']=='0.01' and Decimal(oc['preview']['polygon_summary']['unrounded_quantity'])==Decimal('.005')
    b['occupations'][0]['factor']='0.5'
    view=post(e,pid,'/components/preview',b).json()
    assert view['quantity']=='99.995' and Decimal(view['deduction'])==Decimal('.005')
    assert view['dimensions']['occupied.1.quantity']=='0.01' and view['balance_deductions'][0]['rule_version']=='horizontal-polygon-1.0'
    assert 'multiplicity' not in view['dimensions']


@pytest.mark.parametrize('target',['excavation','occupation'])
def test_area_polygon_never_becomes_volume_even_with_forged_display_unit(clients,target):
    e,r,pid,src,b,area=setup(clients);sel=selection(b)
    if target=='excavation':sel['excavation_component_id']=area['id']
    else:sel['occupation_component_ids']=[area['id']]
    post(e,pid,'/volume-balance-candidates',sel,409)
    mutation(pid,lambda p:next(c for c in p['components'] if c['id']==area['id'])['preview'].update(unit='m³'))
    post(e,pid,'/volume-balance-candidates',sel,409)


@pytest.mark.parametrize('target',['excavation','occupation'])
@pytest.mark.parametrize('fault',['bytes','missing','superseded','proof','coordinate','preview','mapping','stamp'])
def test_actual_polygon_source_integrity_is_checked_for_candidate_and_save(clients,target,fault):
    e,r,pid,src,b,area=setup(clients);cid=src['excavation']['id'] if target=='excavation' else src['occupations'][0]['id']
    if fault in ('bytes','missing'):
        sha=get(e,pid)['documents'][0]['hash'];path=store.ROOT/'originals'/sha
        if fault=='bytes':path.write_bytes(b'corrupted polygon source')
        else:path.unlink()
    else:
        def damage(p):
            c=next(c for c in p['components'] if c['id']==cid)
            if fault=='superseded':p['documents'][0]['superseded']=True
            elif fault=='proof':p['evidence'][0]['text']+=' changed'
            elif fault=='coordinate':c['polygon_rings'][0]['points'][1]['parameters']['x']='7'
            elif fault=='preview':c['preview']['quantity']='999'
            elif fault=='mapping':c['preview']['polygon_rings'][0]['edges'][0]['cross_product']='999'
            else:c['source_stamp']='0'*64
        mutation(pid,damage)
    before=get(e,pid);post(e,pid,'/volume-balance-candidates',selection(b),409);post(e,pid,'/components/preview',b,409);post(e,pid,'/components',b,409)
    assert get(e,pid)==before


@pytest.mark.parametrize('gate',['component_review','calculate','book_review'])
def test_saved_balance_rechecks_polygon_sources_at_every_downstream_gate(clients,gate):
    e,r,pid,src,b,area=setup(clients);comp=post(e,pid,'/components',b).json()
    if gate!='component_review':approve(r,pid,'components',comp)
    if gate=='book_review':book=post(e,pid,'/takeoffs/calculate',{}).json()[-1]
    mutation(pid,lambda p:next(c for c in p['components'] if c['id']==src['occupations'][0]['id'])['polygon_rings'][0]['points'][1]['parameters'].update(x='7'))
    if gate=='component_review':approve(r,pid,'components',comp,expected=409)
    elif gate=='calculate':post(e,pid,'/takeoffs/calculate',{},409)
    else:approve(r,pid,'takeoffs',book,expected=409)


@pytest.mark.parametrize('target',['component','book'])
def test_saved_polygon_source_snapshot_cannot_be_replaced_by_a_forgery(clients,target):
    e,r,pid,src,b,area=setup(clients);comp=post(e,pid,'/components',b).json()
    if target=='book':approve(r,pid,'components',comp);post(e,pid,'/takeoffs/calculate',{})
    def damage(p):
        item=p['components'][-1] if target=='component' else p['takeoffs'][-1]
        item['source_geometry']['occupations'][0]['polygon_rings'][0]['points'][0]['parameters']['x']='999'
    mutation(pid,damage);collection='components' if target=='component' else 'takeoffs'
    approve(r,pid,collection,get(e,pid)[collection][-1],expected=409)


def test_polygon_revision_invalidates_old_balance_and_requires_new_source_hashes(clients):
    e,r,pid,src,b,area=setup(clients);comp=post(e,pid,'/components',b).json();approve(r,pid,'components',comp)
    for t in post(e,pid,'/takeoffs/calculate',{}).json():approve(r,pid,'takeoffs',t)
    post(e,pid,'/boq/generate',{})
    old=next(c for c in get(e,pid)['components'] if c['id']==src['occupations'][0]['id']);change={**copy.deepcopy(old),'object_hash':store.digest(old),'thickness':'300'}
    res=e.put(f'/api/v1/projects/{pid}/components/'+old['id'],json=change);assert res.status_code==200,res.text
    p=get(e,pid);assert next(c for c in p['components'] if c['id']==comp['id'])['state']=='stale' and all(t['state']=='stale' for t in p['takeoffs'])
    assert next(v for v in p['boq'] if v['name']==b['name'])['quantity']=='90.000'
    post(e,pid,'/components/preview',b,409)
    for c in p['components']:
        if c['id']!=comp['id']:approve(r,pid,'components',c)
    post(e,pid,'/components/preview',b,409)
    fresh=post(e,pid,'/volume-balance-candidates',selection(b)).json();assert fresh['occupations'][0]['quantity']=='15.000'
    b['excavation_component_hash']=fresh['excavation']['component_hash'];b['occupations'][0]['component_hash']=fresh['occupations'][0]['component_hash']
    current=next(c for c in get(e,pid)['components'] if c['id']==comp['id']);b['object_hash']=store.digest(current)
    res=e.put(f'/api/v1/projects/{pid}/components/'+comp['id'],json=b);assert res.status_code==200,res.text
    revised=res.json();assert revised['preview']['quantity']=='85.000' and revised['revision']==2
    assert get(e,pid)['component_history'][-1]==current
    approve(r,pid,'components',revised)
    for t in post(e,pid,'/takeoffs/calculate',{}).json():approve(r,pid,'takeoffs',t)
    assert next(v for v in post(e,pid,'/boq/generate',{}).json() if v['name']==b['name'])['quantity']=='85.000'


def test_role_tenant_membership_cross_project_and_recursive_source_guards(clients):
    e,r,pid,src,b,area=setup(clients);sel=selection(b)
    post(r,pid,'/components',b,403);post(clients['other'],pid,'/volume-balance-candidates',sel,404)
    post(TestClient(app),pid,'/volume-balance-candidates',sel,401)
    other=create(e)['id'];post(e,other,'/volume-balance-candidates',sel,404)
    post(e,pid,'/volume-balance-candidates',{**sel,'occupation_component_ids':[sel['occupation_component_ids'][0]]*2},422)
    post(e,pid,'/volume-balance-candidates',{**sel,'occupation_component_ids':[sel['excavation_component_id']]},422)
    comp=post(e,pid,'/components',b).json();approve(r,pid,'components',comp)
    post(e,pid,'/volume-balance-candidates',{**sel,'occupation_component_ids':[comp['id']]},409)
    from tests.test_project_access import restrict
    restrict(clients['admin'],pid,[]);post(e,pid,'/volume-balance-candidates',sel,404)


@pytest.mark.parametrize('fmt',['json','xlsx','pdf'])
def test_exports_keep_polygon_net_quantities_and_original_quote_snapshot(clients,fmt):
    e,r,pid,_,quote=priced_project(clients);e,r,pid,src,b,area=setup(clients,pid=pid)
    comp=post(e,pid,'/components',b).json();approve(r,pid,'components',comp)
    for t in post(e,pid,'/takeoffs/calculate',{}).json():approve(r,pid,'takeoffs',t)
    base=f'/api/v1/projects/{pid}/exports';res=e.get(base,params={'format':fmt,'unpriced':True});assert res.status_code==200,res.text
    titles=['多边形净量汇总','多边形外轮廓与孔','关联回填平衡','回填逐项扣减','回填来源计算索引']
    if fmt=='json':
        data=res.json();assert all(any(t[0]==title for t in data['tables']) for title in titles)
        book=next(t for t in data['unpriced']['takeoffs'] if t['name']==b['name']);assert book['source_geometry']==src and book['quantity']=='90.000'
    elif fmt=='xlsx':assert all(title in load_workbook(io.BytesIO(res.content)).sheetnames for title in titles)
    else:
        text=''.join(p.extract_text() for p in PdfReader(io.BytesIO(res.content)).pages)
        assert all(title in text for title in titles) and '90.000' in text
    historic=e.get(base,params={'format':'json','quote_id':quote['id']}).json();assert historic['quote']==quote
    assert not any(t[0] in titles for t in historic['tables'])
