"""Synthetic boundary railing lengths; not a professional engineering signoff."""
import copy,io,json,random
from decimal import Decimal,localcontext
import pytest
from fastapi import HTTPException
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pypdf import PdfReader
from backend import store
from backend.engine import DomainError
from backend.main import app
from backend.railing_takeoff import calculate,boundary
from backend.surface_takeoff import snapshot
from tests.test_polygon_takeoff import payload as polygon,ring,setup as polygon_setup,post,mutate
from tests.test_route_takeoff import evaluate
from tests.test_workflow import approve,get,create,upload,priced_project


def geometry(eid='test',mode='volume'):
 s=polygon(eid,mode);s.update(id='synthetic-polygon',revision=1);s['polygon_rings']=[ring('三四五外轮廓',[(0,0),(6,0),(6,8)],eid),ring('三角孔',[(4,1),(5,1),(5,2)],eid,'hole')];return s


def interval(a,b,eid='test',unit='m'):
 return {'from':str(a),'to':str(b),'unit':unit,'from_evidence_id':eid,'to_evidence_id':eid}


def payload(s,eid='test'):
 es=boundary(s)['railing_edges']
 return dict(name='合成边线栏杆（非真实工程）',kind='polygon_railing',parameters={},quantity_type='length',source_component_id=s['id'],source_component_hash=store.digest(snapshot(s)),
 edge_decisions=[dict(edge_id=e['edge_id'],mode='full' if e['index']==3 else 'ranges' if e['index']==1 else 'exclude',ranges=[interval(0,1,eid),interval(2,5,eid)] if e['index']==1 else [],basis='合成只计外边E1两段与闭合斜边E3；孔边须明确排除，非专业认可',evidence_id=eid) for e in es],scope_confirmed=True,material='合成栏杆材料做法，仅软件验收',material_evidence_id=eid,installation_basis='合成核对水平原边及起终方向，非实际安装图；不计高度与多道扶手',installation_evidence_id=eid,measurement_basis='单组14m，来源重复2=28m，仅总长修约3位，其他原边及孔边排除',basis_evidence_id=eid,rounding_digits=3,rounding_evidence_id=eid)


def setup(clients,pid=None):
 e,r,pid,_=polygon_setup(clients,pid);upload(e,pid,'合成软件验收：水平外三角(0,0)(6,0)(6,8)m，原边6、8、10m；孔(4,1)(5,1)(5,2)m。仅外E1取S0→1、2→5m，共4m，加闭合外E3整10m，重复2得28m，孔边及外E2排除。材料与安装按原边水平范围核定，不计高度或多道扶手，总长修约3位。修订外E1取S0→2、3→5m仍4m；改仅S0→3m时总长26m。非真实工程或专业签审。','railing-basis.txt')
 eid=get(e,pid)['evidence'][-1]['id'];s=geometry(eid);s.pop('id');s.pop('revision');s=post(e,pid,s).json();
 for current in get(e,pid)['components']:approve(r,pid,'components',current)
 s=get(e,pid)['components'][-1]
 return e,r,pid,s,payload(s,eid),eid


def test_exact_closing_edge_holes_partial_gaps_and_inherited_count_ast():
 s=geometry();b=payload(s);before=copy.deepcopy((s,b));r=calculate(b,s)
 assert (s,b)==before and r['quantity']=='28.000' and r['unit']=='m'
 assert [e['length'] for e in r['railing_edges'][:3]]==['6','8','10']
 assert r['railing_edges'][2]['to_index']==1 and r['railing_edges'][3]['role']=='hole'
 assert [Decimal(e['selected_length']) for e in r['railing_edges']]==[4,0,10,0,0,0]
 with localcontext() as c:
  c.prec=80
  assert evaluate(r['formula_ast'],r['dimensions'])==28
  for e in r['railing_edges']:assert evaluate(e['selected_formula_ast'],r['dimensions'])==Decimal(e['selected_length'])
 # User ordering of decisions is irrelevant; raw interval order retained in dimensional keys.
 b['edge_decisions'].reverse();assert calculate(b,s)==r


@pytest.mark.parametrize('mode',['area','volume'])
def test_polygon_area_or_volume_source_preserves_raw_lengths(mode):
 s=geometry(mode=mode);s['rounding_digits']=0;r=calculate(payload(s),s)
 assert r['quantity']=='28.000' and r['source_geometry']==snapshot(s)


def test_units_touching_ranges_zero_selection_and_source_not_rounded():
 s=geometry();b=payload(s);d=b['edge_decisions'][0];d['ranges']=[interval(0,100,unit='cm'),interval(1000,2000,unit='mm')]
 assert calculate(b,s)['quantity']=='24.000'
 d['ranges']=[interval(0,600,unit='cm')];assert calculate(b,s)['quantity']=='32.000'
 for d in b['edge_decisions']:d.update(mode='exclude',ranges=[])
 assert calculate(b,s)['quantity']=='0.000'


def test_irrational_length_limit_is_compared_exactly_not_to_displayed_sqrt():
 s=geometry();s['polygon_rings']=[ring('小三角',[(0,0),(1,1),(2,0)])];b=payload(s)
 d=b['edge_decisions'][0];d['ranges']=[interval(0,'1.41421356237309504')]
 assert calculate(b,s)
 d['ranges']=[interval(0,'1.41421356237309505')]
 with pytest.raises(DomainError):calculate(b,s)
 b['edge_decisions'][0].update(mode='full',ranges=[]);b['rounding_digits']=12
 with localcontext() as c:
  c.prec=80;r=calculate(b,s);assert Decimal(r['quantity'])==((Decimal(2).sqrt()+2)*2).quantize(Decimal('1e-12'))


@pytest.mark.parametrize('transform',['translate','reverse','rotate','swap','unit','scale'])
def test_perimeter_and_selected_lengths_under_design_transforms(transform):
 s=geometry();factor=1
 for ring_ in s['polygon_rings']:
  if transform=='reverse':ring_['points'].reverse()
  if transform=='rotate':ring_['points']=ring_['points'][1:]+ring_['points'][:1]
  for p in ring_['points']:
   x,y=map(Decimal,p['parameters'].values())
   if transform=='translate':x+=123;y-=45
   if transform=='swap':x,y=y,x
   if transform=='unit':x*=1000;y*=1000;p['unit']='mm'
   if transform=='scale':x*=7;y*=7;factor=7
   p['parameters']={'x':str(x),'y':str(y)}
 b=payload(s)
 for d in b['edge_decisions']:d.update(mode='full' if next(e for e in boundary(s)['railing_edges'] if e['edge_id']==d['edge_id'])['role']=='outer' else 'exclude',ranges=[])
 assert Decimal(calculate(b,s)['quantity'])==48*factor


def test_seeded_independent_pythagorean_and_gap_oracle():
 rng=random.Random(6501)
 for _ in range(35):
  k=rng.randint(1,20);s=geometry();s['polygon_rings']=[ring('三角',[(0,0),(3*k,0),(3*k,4*k)])];s['multiplicity']=str(rng.randint(1,9));b=payload(s)
  end=Decimal(rng.randint(1,3*k));b['edge_decisions'][0]['ranges']=[interval(0,end)];r=calculate(b,s)
  assert Decimal(r['quantity'])==(end+5*k)*int(s['multiplicity'])


@pytest.mark.parametrize('case',['missing','duplicate','unknown','bool_mode','range_missing','range_extra','zero','negative','reverse','outside','overlap','duplicate_range','full_with_range','exclude_with_range','empty_ranges','range_count','unit','nan','precision','blank_basis','missing_proof','unconfirmed','str_confirm','count','digits','parameters','mode','blank_install','blank_material','old_hash'])
def test_strict_invalid_inputs(case):
 s=geometry();b=payload(s);d=b['edge_decisions'][0];v=d['ranges'][0]
 if case=='missing':b['edge_decisions'].pop()
 elif case=='duplicate':b['edge_decisions'][1]=copy.deepcopy(d)
 elif case=='unknown':d['edge_id']='unknown'
 elif case=='bool_mode':d['mode']=True
 elif case=='range_missing':v.pop('from_evidence_id')
 elif case=='range_extra':v['epsilon']='.01'
 elif case=='zero':v['to']='0'
 elif case=='negative':v['from']='-1'
 elif case=='reverse':v.update({'from':'2','to':'1'})
 elif case=='outside':v['to']='6.000000000000000001'
 elif case=='overlap':d['ranges'][1]['from']='.5'
 elif case=='duplicate_range':d['ranges']=[v,copy.deepcopy(v)]
 elif case=='full_with_range':d['mode']='full'
 elif case=='exclude_with_range':d['mode']='exclude'
 elif case=='empty_ranges':d['ranges']=[]
 elif case=='range_count':d['ranges']=[v]*21
 elif case=='unit':v['unit']='m²'
 elif case=='nan':v['to']='NaN'
 elif case=='precision':v['to']='.0000000000000000001'
 elif case=='blank_basis':d['basis']=' '
 elif case=='missing_proof':v['to_evidence_id']=''
 elif case=='unconfirmed':b['scope_confirmed']=False
 elif case=='str_confirm':b['scope_confirmed']='true'
 elif case=='count':b['multiplicity']='2'
 elif case=='digits':b['rounding_digits']=True
 elif case=='parameters':b['parameters']={'length':1}
 elif case=='mode':b['quantity_type']='area'
 elif case=='blank_install':b['installation_basis']=''
 elif case=='blank_material':b['material']=''
 elif case=='old_hash':b['source_component_hash']='old'
 with pytest.raises((DomainError,HTTPException)):calculate(b,s)


def test_read_only_candidates_preview_review_boq_and_mapping_profile(clients):
 e,r,pid,s,b,eid=setup(clients);before=get(e,pid);url=f'/api/v1/projects/{pid}/railing-edge-candidates?component_id='+s['id'];res=e.get(url)
 assert res.status_code==200 and len(res.json()['railing_edges'])==6 and get(e,pid)==before
 post(e,pid,b,'/components/preview');assert get(e,pid)==before
 bad=copy.deepcopy(b);bad['edge_decisions'][0]['ranges'][0]['to']='7';post(e,pid,bad,expected=422);assert get(e,pid)==before
 c=post(e,pid,{**b,'source_geometry':{'forged':True}}).json();assert c['source_geometry']==snapshot(s)
 post(e,pid,{},'/takeoffs/calculate',409);approve(r,pid,'components',c)
 for t in post(e,pid,{},'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
 assert any(x['unit']=='m' and Decimal(x['quantity'])==28 for x in post(e,pid,{},'/boq/generate').json())
 from backend.mapping_types import component_profile
 assert component_profile(c,c['preview'])['id']=='length'


@pytest.mark.parametrize('fmt',['json','xlsx','pdf'])
def test_exports_and_old_quote_snapshot_isolation(clients,fmt):
 e,r,pid,_,quote=priced_project(clients);e,r,pid,s,b,eid=setup(clients,pid);c=post(e,pid,b).json();approve(r,pid,'components',c)
 for t in post(e,pid,{},'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
 url=f'/api/v1/projects/{pid}/exports';res=e.get(url,params={'format':fmt,'unpriced':True});assert res.status_code==200,res.text
 titles=['栏杆边线长度汇总','栏杆逐边范围核定','栏杆局部计长区间','栏杆来源原始坐标']
 if fmt=='json':assert all(any(t[0]==title for t in res.json()['tables']) for title in titles)
 elif fmt=='xlsx':assert set(titles)<=set(load_workbook(io.BytesIO(res.content)).sheetnames)
 else:assert all(t in ''.join(p.extract_text() for p in PdfReader(io.BytesIO(res.content)).pages) for t in titles)
 old=e.get(url,params={'format':'json','quote_id':quote['id']}).json();assert old['quote']==quote and not any(t[0] in titles for t in old['tables'])


@pytest.mark.parametrize('fault',['edge','range','ast','source','basis','evidence','stamp','book'])
def test_forged_saved_records_are_rebuilt_at_review(clients,fault):
 e,r,pid,s,b,eid=setup(clients);c=post(e,pid,b).json()
 if fault=='book':approve(r,pid,'components',c);post(e,pid,{},'/takeoffs/calculate')
 def forge(p):
  x=p['components'][-1];v=x['preview']
  if fault=='edge':v['railing_edges'][0]['length']='999'
  elif fault=='range':v['railing_edges'][0]['normalized_ranges'][0]['s_to']='999'
  elif fault=='ast':v['formula_ast']={'constant':'28.000'}
  elif fault=='source':x['source_geometry']['polygon_rings'][0]['points'][0]['parameters']['x']='99'
  elif fault=='basis':x['edge_decisions'][0]['basis']='其他说明'
  elif fault=='evidence':x['evidence']['railing.basis']='forged'
  elif fault=='stamp':x['source_stamp']='forged'
  else:p['takeoffs'][-1]['railing_edges'][0]['selected_length']='999'
 mutate(pid,forge);coll='takeoffs' if fault=='book' else 'components';approve(r,pid,coll,get(e,pid)[coll][-1],expected=409)


def test_revision_history_concurrency_source_stale_and_rebuild(clients):
 e,r,pid,s,b,eid=setup(clients);c=post(e,pid,b).json();approve(r,pid,'components',c)
 for t in post(e,pid,{},'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
 post(e,pid,{},'/boq/generate');old=get(e,pid)['components'][-1];new=copy.deepcopy(b);new['edge_decisions'][0]['ranges']=[interval(0,3,eid)];new['object_hash']=store.digest(old)
 url=f'/api/v1/projects/{pid}/components/'+c['id'];res=e.put(url,json=new);assert res.status_code==200,res.text
 p=get(e,pid);assert Decimal(res.json()['preview']['quantity'])==26 and p['component_history'][-1]==old and all(x['state']=='stale' for x in p['takeoffs'])
 assert any(x['unit']=='m' and Decimal(x['quantity'])==28 for x in p['boq']) and e.put(url,json=new).status_code==409
 approve(r,pid,'components',res.json())
 for t in post(e,pid,{},'/takeoffs/calculate').json():approve(r,pid,'takeoffs',t)
 assert any(x['unit']=='m' and Decimal(x['quantity'])==26 for x in post(e,pid,{},'/boq/generate').json())
 # Revision of the geometry invalidates the dependent railing; old hash never revalidates.
 s['object_hash']=store.digest(s);s['polygon_rings'][0]['points'][1]['parameters']['x']='7';res=e.put(f'/api/v1/projects/{pid}/components/'+s['id'],json=s);assert res.status_code==200,res.text
 assert get(e,pid)['components'][-1]['state']=='stale';approve(r,pid,'components',res.json());post(e,pid,b,'/components/preview',409)


@pytest.mark.parametrize('gate',['candidate','preview','save','review','calculate','book_review'])
def test_original_bytes_corruption_blocks_all_gates(clients,gate):
 e,r,pid,s,b,eid=setup(clients)
 if gate in ('review','calculate','book_review'):c=post(e,pid,b).json()
 if gate in ('calculate','book_review'):approve(r,pid,'components',c)
 if gate=='book_review':t=post(e,pid,{},'/takeoffs/calculate').json()[-1]
 doc=get(e,pid)['documents'][-1];(store.ROOT/'originals'/doc['hash']).write_bytes(b'corrupt')
 if gate=='candidate':assert e.get(f'/api/v1/projects/{pid}/railing-edge-candidates?component_id='+s['id']).status_code==409
 elif gate=='preview':post(e,pid,b,'/components/preview',409)
 elif gate=='save':post(e,pid,b,expected=409)
 elif gate=='review':approve(r,pid,'components',c,expected=409)
 elif gate=='calculate':post(e,pid,{},'/takeoffs/calculate',409)
 else:approve(r,pid,'takeoffs',t,expected=409)


def test_permissions_tenant_project_foreign_proofs_and_unapproved_source(clients):
 e,r,pid,s,b,eid=setup(clients);url=f'/api/v1/projects/{pid}/railing-edge-candidates?component_id='+s['id'];assert r.get(url).status_code==200
 assert clients['other'].get(url).status_code==404 and TestClient(app).get(url).status_code==401
 post(r,pid,b,expected=403);other=create(e)['id'];upload(e,other,'另一个项目的依据');foreign=get(e,other)['evidence'][-1]['id'];bad=copy.deepcopy(b);bad['edge_decisions'][0]['ranges'][0]['from_evidence_id']=foreign;post(e,pid,bad,expected=404)
 assert e.get(f'/api/v1/projects/{other}/railing-edge-candidates?component_id='+s['id']).status_code==404
 mutate(pid,lambda p:p['components'][0].update(state='draft'));assert e.get(url).status_code==409


def test_large_origin_small_edges_and_low_caller_precision():
 s=geometry()
 with localcontext() as c:
  c.prec=80
  for ring_ in s['polygon_rings']:
   for p in ring_['points']:
    p['parameters']={k:str(Decimal('1000000000')+Decimal(v)*Decimal('.00000001')) for k,v in p['parameters'].items()}
 b=payload(s)
 for d in b['edge_decisions']:d.update(mode='full',ranges=[])
 with localcontext() as c:
  c.prec=10;r=calculate(b,s)
 with localcontext() as c:
  c.prec=80;assert Decimal(r['gross'])==(Decimal(24)+2+Decimal(2).sqrt())*Decimal('2e-8')


def test_total_local_interval_limit_blocks_before_geometric_evaluation():
 s=geometry();s['polygon_rings']=[ring('百边矩形',[(i,0) for i in range(50)]+[(i,1) for i in reversed(range(50))])];b=payload(s)
 for d in b['edge_decisions']:d.update(mode='ranges',ranges=[interval(0,'.001')]*20)
 with pytest.raises(DomainError,match='1000'):calculate(b,s)


def test_membership_revocation_hides_candidates(clients):
 e,r,pid,s,b,eid=setup(clients)
 from tests.test_project_access import restrict
 restrict(clients['admin'],pid,[])
 assert e.get(f'/api/v1/projects/{pid}/railing-edge-candidates?component_id='+s['id']).status_code==404
 post(e,pid,b,'/components/preview',404)
